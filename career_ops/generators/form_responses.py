"""Form response generator — pre-computed answers for ATS application forms.

Uses Claude to generate personalized answers to common application-form
questions (Workday, Lever, Greenhouse, etc.), tailored to each job posting.

Outputs:
  - ``output/FORM_Paula_{Company}_{Role}.json``  (machine-readable, for Phase B)
  - ``output/FORM_Paula_{Company}_{Role}.pdf``   (human-readable, copy-paste)
"""
from __future__ import annotations

import json
import logging
import re
from dataclasses import asdict, dataclass, field
from datetime import date
from pathlib import Path

from docx import Document

from .. import llm
from ..analyzer import JobAnalysis
from ..config import settings
from ..discovery.normalize import Job
from ._paths import job_subdir

logger = logging.getLogger(__name__)

_TEMPLATE_PATH = settings.templates_dir / "form_responses_template.docx"


# -------------------------------------------------------------------
# Data model
# -------------------------------------------------------------------

@dataclass
class FormResponses:
    """All pre-computed answers for a job application form."""

    # Personal (straight from profile.yaml)
    full_name: str = ""
    email: str = ""
    phone: str = ""
    linkedin_url: str = ""
    location: str = ""
    nationality: str = ""
    work_authorization: str = ""

    # Professional (Claude-adapted per job)
    current_title: str = ""
    years_of_experience: str = ""
    current_company: str = ""
    notice_period: str = ""

    # Narrative answers (Claude-generated, job-specific)
    why_interested: str = ""
    why_good_fit: str = ""
    greatest_achievement: str = ""
    salary_expectation: str = ""
    availability: str = ""

    # Dropdown-style answers
    education_level: str = ""
    willing_to_relocate: str = ""
    requires_sponsorship: str = ""
    preferred_work_model: str = ""

    # Additional Q&A (Claude generates 5-8 per ATS platform)
    additional_qa: list[dict] = field(default_factory=list)


# -------------------------------------------------------------------
# ATS platform detection + common questions
# -------------------------------------------------------------------

_ATS_PATTERNS: dict[str, list[str]] = {
    "workday": ["myworkdayjobs.com", "myworkdaysite.com", "wd3.myworkday", "wd5.myworkday"],
    "lever": ["jobs.lever.co"],
    "greenhouse": ["boards.greenhouse.io", "greenhouse.io"],
    "smartrecruiters": ["jobs.smartrecruiters.com"],
    "bamboohr": ["bamboohr.com"],
    "icims": ["icims.com"],
    "taleo": ["taleo.net"],
    "successfactors": ["successfactors.com", "successfactors.eu"],
}

_ATS_COMMON_QUESTIONS: dict[str, list[str]] = {
    "workday": [
        "How did you hear about this position?",
        "Have you previously worked for this company?",
        "Are you legally authorized to work in this country?",
        "Will you now or in the future require sponsorship for employment visa status?",
        "What is your desired compensation?",
        "Are you willing to undergo a background check?",
        "Do you have experience with [industry-specific tool/skill]?",
        "Please describe your experience managing teams or cross-functional projects.",
    ],
    "lever": [
        "How did you hear about this role?",
        "Are you authorized to work in the UAE?",
        "What is your expected salary?",
        "LinkedIn profile URL",
        "What excites you most about this opportunity?",
    ],
    "greenhouse": [
        "How did you hear about this job?",
        "Are you legally authorized to work in this location?",
        "Do you now or will you in the future require visa sponsorship?",
        "Desired salary",
        "Earliest start date",
        "Are you willing to relocate?",
        "Do you have experience in the [sector] industry?",
    ],
    "smartrecruiters": [
        "How did you find out about this position?",
        "What is your current notice period?",
        "Expected annual salary",
        "Are you eligible to work in this country?",
        "Do you have relevant industry experience?",
    ],
    "generic": [
        "How did you hear about this position?",
        "Are you legally authorized to work in the UAE?",
        "Do you require visa sponsorship now or in the future?",
        "What is your expected salary range?",
        "What is your earliest available start date?",
        "What is your current notice period?",
        "Do you have experience in this industry?",
        "Describe a professional achievement you are most proud of.",
    ],
}


def _detect_ats_platform(url: str) -> str | None:
    """Detect ATS platform from job URL domain."""
    if not url:
        return None
    url_lower = url.lower()
    for platform, patterns in _ATS_PATTERNS.items():
        if any(p in url_lower for p in patterns):
            return platform
    return None


# -------------------------------------------------------------------
# Claude prompt
# -------------------------------------------------------------------

_SYSTEM = """\
You generate pre-computed application form responses for Paula De Francisco,
a Brand & Marketing Manager applying to roles in Dubai's FMCG, Beauty, and
E-Commerce sectors. Your answers will be copy-pasted into ATS application
forms (Workday, Lever, Greenhouse, etc.).

Rules:
1. Answer as Paula — first person, professional but warm tone.
2. Use REAL metrics from her profile: +30% GMV QoQ, 42 key accounts, 6 NPD
   launches end-to-end, 50+ countries, 25-50 influencers per campaign.
3. For "why interested" — reference something SPECIFIC about the company or
   role from the job description. Never be generic.
4. For salary — state a range in AED/month. Paula's floor is 20,000 AED/month.
   Suggest a range appropriate to the role seniority (e.g., 20,000-25,000 for
   mid-level, 25,000-35,000 for senior).
5. For years of experience — Paula has 4+ years (since Aug 2021).
6. Narrative answers: 2-4 sentences each. Concise, specific, metric-driven.
7. NEVER invent achievements not listed in the profile.
8. Generate 5-8 additional Q&A pairs tailored to the ATS platform and role.
9. Output via the submit_form_responses tool.
"""

_FORM_TOOL = {
    "name": "submit_form_responses",
    "description": "Submit pre-computed form responses for a job application",
    "input_schema": {
        "type": "object",
        "properties": {
            "current_title": {
                "type": "string",
                "description": "Paula's current job title, adapted if needed for the application",
            },
            "years_of_experience": {
                "type": "string",
                "description": "Total years of professional experience (e.g., '4+ years')",
            },
            "current_company": {
                "type": "string",
                "description": "Current employer name",
            },
            "notice_period": {
                "type": "string",
                "description": "Current notice period (e.g., '30 days', '2 weeks', 'Immediately available')",
            },
            "why_interested": {
                "type": "string",
                "description": "Answer to 'Why are you interested in this role/company?' — 2-4 sentences, company-specific",
            },
            "why_good_fit": {
                "type": "string",
                "description": "Answer to 'What makes you a good fit?' — 2-4 sentences with metrics",
            },
            "greatest_achievement": {
                "type": "string",
                "description": "Answer to 'Describe your greatest professional achievement' — 2-4 sentences",
            },
            "salary_expectation": {
                "type": "string",
                "description": "Salary expectation as a range in AED/month",
            },
            "availability": {
                "type": "string",
                "description": "Answer to 'When can you start?' (e.g., 'Available to start within 30 days')",
            },
            "additional_qa": {
                "type": "array",
                "description": "5-8 additional Q&A pairs for common ATS questions",
                "items": {
                    "type": "object",
                    "properties": {
                        "question": {"type": "string"},
                        "answer": {"type": "string"},
                    },
                    "required": ["question", "answer"],
                },
            },
        },
        "required": [
            "current_title", "years_of_experience", "current_company",
            "notice_period", "why_interested", "why_good_fit",
            "greatest_achievement", "salary_expectation", "availability",
            "additional_qa",
        ],
    },
}


# -------------------------------------------------------------------
# Consistency cache
# -------------------------------------------------------------------

def _load_response_history() -> list[dict]:
    """Load previous form responses for consistency context."""
    path = settings.form_response_history_path
    if not path.exists():
        return []
    try:
        with path.open("r", encoding="utf-8") as f:
            data = json.load(f)
        return data.get("responses", [])
    except (json.JSONDecodeError, KeyError):
        return []


def _save_to_history(job: Job, responses: FormResponses) -> None:
    """Append form responses to history for future consistency."""
    path = settings.form_response_history_path
    history = _load_response_history()
    history.append({
        "job_id": job.id,
        "company": job.company,
        "role": job.title,
        "generated_at": date.today().isoformat(),
        "responses": asdict(responses),
    })
    with path.open("w", encoding="utf-8") as f:
        json.dump({"responses": history}, f, indent=2, ensure_ascii=False)


def _build_consistency_context() -> str:
    """Build a context block from the 3 most recent responses."""
    history = _load_response_history()
    if not history:
        return ""

    recent = history[-3:]
    lines = ["## Previous answers (maintain consistency across applications):\n"]
    for entry in recent:
        r = entry.get("responses", {})
        lines.append(f"- {entry.get('company', '?')} / {entry.get('role', '?')}:")
        lines.append(f"  Years of experience: {r.get('years_of_experience', '?')}")
        lines.append(f"  Salary expectation: {r.get('salary_expectation', '?')}")
        lines.append(f"  Notice period: {r.get('notice_period', '?')}")
        lines.append(f"  Availability: {r.get('availability', '?')}")
        lines.append("")
    return "\n".join(lines)


# -------------------------------------------------------------------
# Claude API call
# -------------------------------------------------------------------

def _generate_responses(job: Job, analysis: JobAnalysis, ats_platform: str | None) -> dict:
    """Call the LLM to generate personalized form responses."""
    p = settings.profile

    platform = ats_platform or _detect_ats_platform(job.url) or "generic"
    platform_questions = _ATS_COMMON_QUESTIONS.get(platform, _ATS_COMMON_QUESTIONS["generic"])
    consistency = _build_consistency_context()

    user_msg = f"""\
Generate form responses for this job application.

## Target Job
- Title: {job.title}
- Company: {job.company}
- Location: {job.location}
- URL: {job.url}
- ATS Platform: {platform}

### Job Description
{(job.description or 'No description available')[:4000]}

### Scorer Notes
- Score: {analysis.score}/100 ({analysis.tier})
- Skills match: {', '.join(analysis.skills_match)}
- Missing skills: {', '.join(analysis.missing_skills)}
- Key ATS terms: {', '.join(analysis.ats_keywords)}

## Candidate Profile
{p['personal']['name']} — {p['headline']}
{p['professional_summary']}

Current role: {p['experience'][0]['role']} at {p['experience'][0]['company']} ({p['experience'][0]['dates']})
Previous: {p['experience'][1]['role']} at {p['experience'][1]['company']} ({p['experience'][1]['dates']})

Key metrics:
- +30% GMV growth QoQ managing 42 key accounts at Alibaba Group (Miravia)
- 6 NPD launches end-to-end across 50+ countries at DoFreeze
- Scaled influencer programme from zero to 25-50 creators per campaign
- UAE quick-commerce platforms: Noon, Talabat, Careem, Deliveroo
- Already in Dubai with UAE residence visa — no sponsorship needed

Education: {p['education'][0]['degree']} — {p['education'][0]['school']} ({p['education'][0]['dates']})
Languages: {', '.join(f"{l['lang']} ({l['level']})" for l in p['languages'])}
Salary floor: {p['preferences']['min_salary_aed_month']} AED/month
Work model: {'Hybrid OK' if p['preferences']['hybrid_ok'] else 'On-site only'}

{consistency}

## ATS Platform Typical Questions ({platform})
Generate additional Q&A for these common questions:
{chr(10).join(f'- {q}' for q in platform_questions)}

Use the submit_form_responses tool to return all answers."""

    data = llm.generate_structured(
        system=_SYSTEM,
        user=user_msg,
        tool_schema=_FORM_TOOL,
        tier="sonnet",
        max_tokens=2048,
        timeout=60.0,
    )

    if not data:
        logger.warning("LLM did not return form responses")
    return data


# -------------------------------------------------------------------
# Build full FormResponses (profile data + Claude output)
# -------------------------------------------------------------------

def _build_form_responses(claude_output: dict) -> FormResponses:
    """Merge static profile data with Claude-generated answers."""
    p = settings.profile["personal"]
    prefs = settings.profile["preferences"]

    return FormResponses(
        # Static from profile
        full_name=p["name"],
        email=p["email"],
        phone=p["phone"],
        linkedin_url=p["linkedin"],
        location=p["location"],
        nationality=p["nationality"],
        work_authorization=prefs.get("work_authorization", p.get("visa", "")),
        # Dropdown-style (static)
        education_level="Bachelor's Degree",
        willing_to_relocate="No — already based in Dubai, UAE",
        requires_sponsorship="No",
        preferred_work_model="Hybrid" if prefs.get("hybrid_ok") else "On-site",
        # Claude-generated
        current_title=claude_output.get("current_title", ""),
        years_of_experience=claude_output.get("years_of_experience", ""),
        current_company=claude_output.get("current_company", ""),
        notice_period=claude_output.get("notice_period", ""),
        why_interested=claude_output.get("why_interested", ""),
        why_good_fit=claude_output.get("why_good_fit", ""),
        greatest_achievement=claude_output.get("greatest_achievement", ""),
        salary_expectation=claude_output.get("salary_expectation", ""),
        availability=claude_output.get("availability", ""),
        additional_qa=claude_output.get("additional_qa", []),
    )


# -------------------------------------------------------------------
# Output: JSON
# -------------------------------------------------------------------

def _sanitize(s: str) -> str:
    return re.sub(r'[<>:"/\\|?*]', "", s).strip().replace(" ", "_")[:50]


def _save_json(responses: FormResponses, job: Job) -> Path:
    """Save form responses as JSON for machine consumption / Phase B."""
    company = _sanitize(job.company)
    role = _sanitize(job.title)
    path = job_subdir(job, "application") / f"FORM_Paula_{company}_{role}.json"

    data = {
        "job_id": job.id,
        "company": job.company,
        "role": job.title,
        "url": job.url,
        "generated_at": date.today().isoformat(),
        "responses": asdict(responses),
    }
    with path.open("w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

    return path


# -------------------------------------------------------------------
# Output: DOCX → PDF
# -------------------------------------------------------------------

def _replace_in_paragraphs(doc: Document, placeholder: str, value: str) -> None:
    """Replace placeholder in all paragraphs, handling split runs."""
    for para in doc.paragraphs:
        full_text = "".join(run.text for run in para.runs)
        if placeholder not in full_text:
            continue
        new_text = full_text.replace(placeholder, value)
        if para.runs:
            para.runs[0].text = new_text
            for run in para.runs[1:]:
                run.text = ""
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    full_text = "".join(run.text for run in para.runs)
                    if placeholder not in full_text:
                        continue
                    new_text = full_text.replace(placeholder, value)
                    if para.runs:
                        para.runs[0].text = new_text
                        for run in para.runs[1:]:
                            run.text = ""


def _fill_template(responses: FormResponses, job: Job) -> Path:
    """Fill the DOCX template with form responses and save."""
    doc = Document(str(_TEMPLATE_PATH))

    _replace_in_paragraphs(doc, "{{COMPANY}}", job.company)
    _replace_in_paragraphs(doc, "{{ROLE}}", job.title)
    _replace_in_paragraphs(doc, "{{DATE}}", date.today().strftime("%d %B %Y"))

    # Personal info
    _replace_in_paragraphs(doc, "{{FULL_NAME}}", responses.full_name)
    _replace_in_paragraphs(doc, "{{EMAIL}}", responses.email)
    _replace_in_paragraphs(doc, "{{PHONE}}", responses.phone)
    _replace_in_paragraphs(doc, "{{LINKEDIN}}", responses.linkedin_url)
    _replace_in_paragraphs(doc, "{{LOCATION}}", responses.location)
    _replace_in_paragraphs(doc, "{{NATIONALITY}}", responses.nationality)
    _replace_in_paragraphs(doc, "{{WORK_AUTH}}", responses.work_authorization)

    # Professional
    _replace_in_paragraphs(doc, "{{CURRENT_TITLE}}", responses.current_title)
    _replace_in_paragraphs(doc, "{{YEARS_EXP}}", responses.years_of_experience)
    _replace_in_paragraphs(doc, "{{CURRENT_COMPANY}}", responses.current_company)
    _replace_in_paragraphs(doc, "{{NOTICE_PERIOD}}", responses.notice_period)

    # Narrative
    _replace_in_paragraphs(doc, "{{WHY_INTERESTED}}", responses.why_interested)
    _replace_in_paragraphs(doc, "{{WHY_GOOD_FIT}}", responses.why_good_fit)
    _replace_in_paragraphs(doc, "{{GREATEST_ACHIEVEMENT}}", responses.greatest_achievement)
    _replace_in_paragraphs(doc, "{{SALARY_EXPECTATION}}", responses.salary_expectation)
    _replace_in_paragraphs(doc, "{{AVAILABILITY}}", responses.availability)

    # Dropdowns
    _replace_in_paragraphs(doc, "{{EDUCATION_LEVEL}}", responses.education_level)
    _replace_in_paragraphs(doc, "{{RELOCATE}}", responses.willing_to_relocate)
    _replace_in_paragraphs(doc, "{{SPONSORSHIP}}", responses.requires_sponsorship)
    _replace_in_paragraphs(doc, "{{WORK_MODEL}}", responses.preferred_work_model)

    # Additional Q&A — build a single text block
    qa_text = ""
    for qa in responses.additional_qa:
        qa_text += f"Q: {qa['question']}\nA: {qa['answer']}\n\n"
    _replace_in_paragraphs(doc, "{{ADDITIONAL_QA}}", qa_text.strip())

    company = _sanitize(job.company)
    role = _sanitize(job.title)
    out_path = job_subdir(job, "application") / f"FORM_Paula_{company}_{role}.docx"
    doc.save(str(out_path))
    return out_path


def _to_pdf(docx_path: Path) -> Path:
    """Convert DOCX to PDF via Word and remove the intermediate DOCX.

    Tolerates the post-success ``Word.Application.Quit`` COM teardown error
    from ``docx2pdf`` on Windows.
    """
    from docx2pdf import convert

    pdf_path = docx_path.with_suffix(".pdf")
    try:
        convert(str(docx_path), str(pdf_path))
    except Exception as exc:  # AttributeError on Quit, or pywintypes.com_error from RPC teardown
        if not pdf_path.exists():
            raise
        logger.info("docx2pdf reported COM teardown after success (%s) — PDF exists, continuing", type(exc).__name__)
    if pdf_path.exists():
        docx_path.unlink(missing_ok=True)
    return pdf_path


# -------------------------------------------------------------------
# Public API
# -------------------------------------------------------------------

def generate_form_responses(
    job: Job,
    analysis: JobAnalysis,
    ats_platform: str | None = None,
) -> tuple[Path, Path] | None:
    """Generate form responses for a specific job.

    Returns (json_path, pdf_path) or None on failure.
    """
    if not _TEMPLATE_PATH.exists():
        logger.error("Form responses template not found: %s", _TEMPLATE_PATH)
        return None

    available, reason = llm.is_available()
    if not available:
        logger.warning("LLM not available (%s) — cannot generate form responses", reason)
        return None

    logger.info("generating form responses for %s @ %s", job.title[:40], job.company)

    claude_output = _generate_responses(job, analysis, ats_platform)
    if not claude_output:
        return None

    responses = _build_form_responses(claude_output)

    # Save JSON (always)
    json_path = _save_json(responses, job)

    # Save DOCX → PDF
    docx_path = _fill_template(responses, job)
    pdf_path = _to_pdf(docx_path)

    # Update consistency cache
    _save_to_history(job, responses)

    logger.info("form responses saved: %s, %s", json_path, pdf_path)
    return json_path, pdf_path
