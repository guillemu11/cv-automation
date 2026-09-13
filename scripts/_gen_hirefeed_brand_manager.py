"""One-off: generate Paula's CV for the **Brand Manager (Remote)** role at
**Hire Feed** (staffing/recruiting company; remote, UAE; $85-130k USD; LinkedIn
Easy Apply, 100+ applicants, hidden end-client).

Classic brand-management JD — shape & protect how the brand shows up: own brand
positioning, messaging architecture & voice guidelines; lead brand campaigns
concept→launch; partner with creative on identity, design system & brand assets;
run brand research, tracking studies & competitive monitoring; educate teams on
brand standards; manage brand-health metrics and report to leadership. Asks 4+
yrs brand management (consumer or B2B), brand-strategy frameworks, integrated
brand campaigns end-to-end, creative sensibility, agency/in-house creative work,
comfort with brand research & tracking metrics.

This is Paula's #1 target title and the 4+ yrs asked matches her exactly. Fit is
strong and honest:
- Leads Brand & Marketing at DoFreeze — owns positioning, messaging & visual
  guidelines for Befit/Eurocake/Flair across 50+ markets; campaigns concept→launch;
  creative direction & agency/creator management; brand-health tracking & reporting.
- Miravia (Alibaba): built brand campaigns (Beauty Club, Hot on Social); competitive
  & trend monitoring; +30% GMV QoQ.
- Mondelez: category/brand planning, competitive & promo analysis, NPD (Milka).

Honesty guardrails: brand-health / tracking studies framed on her genuine Nielsen/
Kantar + KPI tracking (not a claim of running bespoke primary-research panels); NO
Arabic (not required here); factual "UAE Residence Visa" only — never the
"no-sponsorship" claim.

ONE-PAGE standard (since 2026-08-27): tight summary (~2 lines), 3 bullets current
role / 1-2 older, ~6 skills per row, SHORT skills-row labels (≤~20 chars) so they
don't wrap. Fills the real CV template, relabels the skills rows for a brand
manager (template stays pristine), converts to PDF via LibreOffice, registers the
job, verifies 1 page, and lands the package under output/2026-08-30/. Also drops a
short-named 'Paula De Francisco - CV.pdf' copy for portals.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Hire Feed"
TITLE = "Brand Manager"
DATE_FOLDER = "2026-08-30"

JOB_DESCRIPTION = """\
Brand Manager — Hire Feed (Remote, UAE; full-time; ~$85-130k USD).

We are hiring a Brand Manager to shape and protect how the brand shows up in the
world. The role spans positioning, voice, visual identity, and the campaigns that
make a brand feel like something — not just look like something.

Key Responsibilities:
- Own brand positioning, messaging architecture, and voice guidelines.
- Lead brand campaigns from concept through launch.
- Partner with creative on identity, design system, and brand assets.
- Conduct brand research, tracking studies, and competitive monitoring.
- Educate internal teams on brand standards and consistent application.
- Manage brand health metrics and report to leadership.

Required Skills:
- 4+ years in brand management at a consumer or B2B brand.
- Strong understanding of brand strategy frameworks.
- Track record of leading integrated brand campaigns end-to-end.
- Excellent written communication and creative sensibility.
- Experience working closely with creative agencies or in-house creative teams.
- Comfort with brand research methods and tracking metrics.

What You'll Bring: curiosity about audiences; storytelling instincts; a
test-and-learn mindset (ship campaigns fast, iterate on data); comfort working
asynchronously across time zones.
"""

ATS = [
    "brand manager", "brand management", "brand strategy", "brand positioning",
    "positioning", "messaging architecture", "messaging", "voice guidelines",
    "tone of voice", "visual identity", "brand identity", "design system",
    "brand assets", "brand standards", "brand guidelines", "brand campaigns",
    "integrated campaigns", "campaign", "concept to launch", "end-to-end",
    "creative brief", "creative direction", "creative agencies", "in-house creative",
    "brand research", "tracking studies", "brand health", "brand health metrics",
    "competitive monitoring", "competitive analysis", "consumer insight",
    "storytelling", "go-to-market", "NPD", "test-and-learn", "KPI", "ROI", "ROAS",
    "awareness", "sell-out", "Nielsen", "Kantar", "FMCG", "beauty", "fragrances",
    "e-commerce", "influencer", "social", "paid media", "Meta Ads", "Google Ads",
    "stakeholder management", "leadership reporting", "remote", "UAE", "async",
]

CONTENT = {
    "headline": (
        "Brand Manager · Positioning, Messaging & Voice · Integrated Campaigns "
        "End-to-End · Brand Health & Creative Direction"
    ),
    "professional_summary": (
        "Brand manager with 4+ years owning brands end-to-end across FMCG, beauty, fragrances and "
        "e-commerce — positioning, messaging, voice and visual identity through to integrated campaigns "
        "from concept to launch. Currently lead Brand & Marketing at DoFreeze across 50+ markets: brand "
        "strategy, creative direction, brand-health tracking and reporting to leadership. Data-led, "
        "AI-first and Dubai-based."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Multi-brand FMCG group | Brands: Befit, Eurocake, Flair | 50+ countries",
            "bullets": [
                "Own brand positioning, messaging architecture and voice/visual guidelines for Befit, Eurocake and Flair across 50+ GCC/MENA and international markets — keeping the brand consistent and protected at every touchpoint",
                "Lead integrated brand campaigns from concept through launch — brief, creative direction, paid/social/influencer and go-to-market — partnering with creative teams and agencies on identity, design system and brand assets",
                "Track brand-health and campaign metrics (awareness, sell-out, ROI/ROAS via Nielsen/Kantar and BI), run competitive monitoring, and report to leadership — with an AI (Claude/GPT) layer automating research and reporting",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Miravia — Alibaba's marketplace | Top-5 global e-commerce | 100K+ employees",
            "bullets": [
                "Created and led brand campaigns — Beauty Club and Hot on Social — from concept to activation, building brand love and loyalty across 42 accounts and driving +30% GMV QoQ",
                "Ran competitive and trend monitoring across beauty, fragrances and fashion to shape positioning, assortment and creative decisions",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Delivered bespoke brand activations for strategic partners on a quick-commerce platform, coordinating marketing, creative and operations to lift order volume and GMV",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category & Brand Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG multinational | €36B revenue | 90K+ employees",
            "bullets": [
                "Brand & category planning: promotional-effectiveness and competitive analysis and NPD support (Milka Spread, Mini Suchard) — early grounding in brand tracking and consumer insight",
            ],
        },
    ],
    # Skills-block ROW LABELS relabelled for a brand manager (SHORT so they don't wrap).
    "skills_brand": (  # label -> "Brand Strategy"
        "brand positioning, messaging architecture, voice & tone guidelines, visual identity & brand "
        "standards, brand architecture, go-to-market, NPD end-to-end"
    ),
    "skills_ecommerce": (  # label -> "Campaigns & Creative"
        "integrated brand campaigns (concept→launch), creative briefing & agency direction, design "
        "system & brand assets, influencer & UGC, social & content, paid media (Meta & Google Ads)"
    ),
    "skills_commercial": (  # label -> "Commercial"
        "brand & category management, key account management, pricing & assortment, trade & shopper "
        "marketing, A&P budget management, cross-functional & stakeholder management"
    ),
    "skills_data": (  # label -> "Brand Health & Data"
        "brand-health & tracking studies (Nielsen, Kantar), competitive monitoring, KPI tracking & "
        "reporting, awareness & sell-out, ROI/ROAS, AI-assisted analysis, Power BI, Tableau, Looker"
    ),
    "skills_tools": (  # label -> "Tools" (unchanged)
        "Nielsen, Kantar, Meta Ads, Google Ads, Adobe Creative Suite, Canva, Power BI, Tableau, "
        "Looker, Salesforce, Generative AI (Claude, ChatGPT), Microsoft Office (Expert)"
    ),
}

# Skills-block row-label overrides (keep SHORT — long labels wrap to 2 lines and push a 2nd page).
ROLE_LABELS = {
    "Brand & Marketing": "Brand Strategy",
    "E-Commerce & Digital": "Campaigns & Creative",
    "Commercial": "Commercial",
    "Data & Analytics": "Brand Health & Data",
}


def make_job() -> Job:
    return Job(
        id="hirefeed-brand-manager-2026-08",
        title=TITLE,
        company=COMPANY,
        location="Remote (UAE)",
        url="https://www.linkedin.com/jobs/",
        source="linkedin",
        description=JOB_DESCRIPTION,
        raw={
            "query": "Brand Manager Remote UAE Hire Feed positioning messaging brand campaigns",
            "note": "Staffing/recruiting-company Easy Apply, remote, hidden end-client, 100+ applicants. "
                    "Paula's #1 target title (Brand Manager); JD asks 4+ yrs = exact match. Strong honest "
                    "brand fit (positioning/messaging/voice, campaigns concept→launch, creative direction, "
                    "brand-health tracking). No Arabic required. Factual UAE Residence Visa only.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id,
        "title": job.title,
        "company": job.company,
        "location": job.location,
        "url": job.url,
        "source": job.source,
        "description": job.description,
        "salary_raw": "$85,000 – $130,000 USD + benefits",
        "salary_aed_min": 26000,
        "salary_aed_max": 40000,
        "posted_date": DATE_FOLDER,
        "raw": job.raw,
        "ai_score": 86,
        "ai_tier": "Hot",
        "skills_match": [
            "Paula's #1 target title (Brand Manager); JD asks 4+ yrs — exact match",
            "Owns brand positioning, messaging architecture & voice/visual guidelines across 50+ markets (DoFreeze)",
            "Leads integrated brand campaigns concept→launch; creative direction + agency/creator management",
            "Brand-health & campaign tracking (awareness, sell-out, ROI/ROAS, Nielsen/Kantar, BI); competitive monitoring; reports to leadership",
            "Built brand campaigns at Miravia/Alibaba (Beauty Club, Hot on Social), +30% GMV QoQ",
            "Creative sensibility: art direction & design (Adobe, Canva); strong storytelling",
            "Beauty/fragrances/FMCG/e-commerce breadth; Dubai-based; AI-first; comfortable async",
        ],
        "missing_skills": [
            "'Tracking studies' framed on genuine Nielsen/Kantar + KPI tracking, not bespoke primary-research panels (not fabricated)",
            "No Arabic (not required for this remote role)",
        ],
        "sector_fit": "excellent (classic brand management — positioning, messaging, voice, campaigns, brand health — her core)",
        "seniority_fit": "on-band (JD asks 4+ yrs; Paula 4+) — clean, no stretch",
        "red_flags": [
            "Staffing/recruiting-company Easy Apply, remote, hidden end-client, 100+ applicants — low signal; the tailored CV is the main lever",
            "Remote (Paula prefers hybrid) — salary band ($85-130k ≈ AED 26-40k/mo) clears her floor comfortably",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Strong, on-band fit for Paula's #1 target title. The JD is classic brand management — own brand "
            "positioning, messaging architecture and voice guidelines; lead brand campaigns concept→launch; "
            "partner with creative on identity/design system/brand assets; run brand research, tracking studies "
            "and competitive monitoring; manage brand-health metrics and report to leadership — and asks for 4+ "
            "yrs, which Paula matches exactly. She leads Brand & Marketing at DoFreeze (positioning, messaging & "
            "visual guidelines across 50+ markets, campaigns concept→launch, creative/agency direction, "
            "brand-health tracking and leadership reporting), built brand campaigns at Miravia/Alibaba (+30% GMV "
            "QoQ), and has creative sensibility (art direction & design). Brand-health/tracking is framed on her "
            "genuine Nielsen/Kantar + KPI work, not fabricated primary research. Caveats are about the posting, "
            "not the fit: staffing-company Easy Apply with a hidden client and 100+ applicants (low signal) and "
            "remote (she prefers hybrid); salary clears her floor. No Arabic required; factual UAE Residence Visa only."
        ),
        "scored_by": "manual:claude",
        "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00",
        "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    """Relabel skills rows for this brand-manager CV (template stays pristine)."""
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    """Convert DOCX -> PDF headlessly via LibreOffice (docx2pdf/Word fails on this Mac)."""
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    """Move output/<Company> - <Role>/ under output/<DATE_FOLDER>/."""
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS — trim a bullet/skills")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
