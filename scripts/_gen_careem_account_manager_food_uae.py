"""One-off: CV for Careem "Account Manager I", Food - UAE team, Dubai (seen on Talentmate, ref
2610-2049-620, posted 2026-10-09). CV only (no cover letter requested).

Restaurant-partner account management on a food-delivery platform: own a portfolio of restaurant
partners, run business reviews with data, work with Operations/Marketing/Product, resolve partner
issues, onboard and train new partners. 2-4 yrs in account management / partner success, tech,
e-commerce or F&B; Excel/Google Sheets; negotiation; Business degree. The posting also leads with
Careem's push for "AI talent" who build tools and automate workflows.

Paula's honest angle:
- Glovo Account Manager – XL Accounts: restaurant chains (KFC, Taco Bell, La Tagliatella, Sushi Shop)
  on a food-delivery platform. The same job.
- Miravia KAM: 42 accounts, +30% GMV QoQ, 30+ new stores onboarded in two months.
- DoFreeze: works with Careem, talabat and Noon from the partner side (listings, promotions,
  weekly stock-risk review by dark store) and built the Claude automation system (~40% less manual work).

Honesty guardrails:
- DoFreeze is NOT on Deliveroo (only Careem, talabat, Noon).
- Never "own the forecast"; the AI stock-risk dashboard was built as a team ("we built").
- No invented partner training / QBR claims; no Arabic; visa line comes from profile.yaml.
- Associate-level role: salary likely below her 20k AED/month floor (flagged in the dashboard entry).

ONE-PAGE standard. Lands under output/2026-10-12/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Careem"
TITLE = "Account Manager I - Food UAE"
DATE_FOLDER = "2026-10-12"

JOB_DESCRIPTION = """\
Account Manager I — Careem, Food - UAE team, Dubai. Associate level, full-time.
Own and nurture relationships with restaurant partners to help them grow on the Careem platform;
optimise partner performance, improve operational metrics and drive strategic initiatives.
Manage a portfolio of restaurant partners; drive partner performance through business reviews, data
insights and strategic recommendations; collaborate with Operations, Marketing and Product on growth
initiatives; resolve partner concerns swiftly; monitor and analyse partner performance metrics; support
onboarding and training of new partners; advocate internally for partner needs.
Needs: 2–4 years in account management, partner success, business development or similar, preferably in
tech, e-commerce or F&B; managing external stakeholders and influencing decisions; strong analytical
skills (Excel / Google Sheets); interpersonal, communication and negotiation skills; bias for action;
cross-functional work in a fast-paced environment; Bachelor's in Business, Economics or related.
Careem is hiring "AI talent": people who apply AI to build tools and automate workflows.
4 days office / 1 day home.
"""

ATS = [
    "Account Manager", "account management", "partner success", "business development",
    "restaurant partners", "portfolio management", "relationship management", "food delivery",
    "F&B", "e-commerce", "tech", "business reviews", "data insights", "strategic recommendations",
    "partner performance", "performance metrics", "operational metrics", "onboarding", "training",
    "cross-functional", "Operations", "Marketing", "Product", "stakeholder management",
    "influencing", "negotiation", "problem-solving", "bias for action", "Excel", "Google Sheets",
    "GMV", "orders", "promotions", "AI", "automation", "Careem", "talabat", "Noon", "Glovo",
    "Dubai", "UAE",
]

CV_CONTENT = {
    "headline": (
        "Account Management · Restaurant & Brand Partners · Food Delivery & Q-Commerce · "
        "Data & AI Automation"
    ),
    "professional_summary": (
        "Account manager with 5 years growing partners on delivery and e-commerce platforms. At Glovo I "
        "managed XL restaurant accounts (KFC, Taco Bell, Sushi Shop); at Alibaba's Miravia, 42 accounts to "
        "+30% GMV QoQ. Today I work with Careem, talabat and Noon from the partner side and build AI tools "
        "that automate reporting."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Homegrown UAE F&B / FMCG group | Befit, Eurocake, Flair | Careem, talabat, Noon | 50+ markets",
            "bullets": [
                "Partner-side view of Careem, talabat and Noon: onboarded our SKUs, manage listings and in-app promotions, and measure them against control: SMASH x talabat lifted daily sales +165% (control brand +4.7%)",
                "Run the weekly stock-risk review by dark store across the three platforms, using an AI dashboard we built with Claude to flag where platform POs will not cover demand",
                "Built a Claude-based system that automates campaign planning, market research and KPI reporting, cutting manual work ~40%",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Top-5 global e-commerce | 100K+ employees",
            "bullets": [
                "Managed 42 partner accounts to +30% GMV QoQ, turning sales, traffic and conversion data into pricing, assortment and promotion recommendations; reported the Flash Sales channel to the CEO",
                "Onboarded 30+ new stores in two months as PIC Fragrances, using launch promotions and trend-driven products",
                "Created and led the Beauty Club and Hot on Social programmes, giving partner brands extra visibility and positioning Miravia as a beauty destination",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Food-delivery & quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Managed XL restaurant accounts (KFC, Taco Bell, La Tagliatella, Sushi Shop), growing orders and GMV through data-led plans and bespoke marketing activations",
                "Negotiated commercial deals profitable for both Glovo and partners; worked with marketing, logistics and customer support to deliver campaigns, and helped build the Retail vertical",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | Chocolate category",
            "bullets": [
                "Sell-in/sell-out and promotional-effectiveness analysis and performance reports for the chocolate category",
            ],
        },
    ],
    "skills_brand": (  # -> "Partners"
        "portfolio & relationship management, negotiation, onboarding, promotions & activations"
    ),
    "skills_ecommerce": (  # -> "Platforms"
        "Careem, talabat, Noon, Glovo, Miravia: listings, in-app promotions, dark-store availability"
    ),
    "skills_commercial": (  # -> "Cross-functional"
        "Operations, Marketing, Logistics, Customer Support, Finance; reporting to C-level"
    ),
    "skills_data": (  # -> "Data & KPIs"
        "GMV, orders, conversion, promo uplift vs control, ROI, sell-in/sell-out, dashboards"
    ),
    "skills_tools": (
        "Excel (advanced), Google Sheets, Power BI, Looker, Salesforce, Claude / Claude Code, ChatGPT"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Partners",
    "E-Commerce & Digital": "Platforms",
    "Commercial": "Cross-functional",
    "Data & Analytics": "Data & KPIs",
    "Tools": "Tools & AI",
}


def make_job() -> Job:
    return Job(
        id="careem-account-manager-i-food-uae-2026-10",
        title=TITLE,
        company=COMPANY,
        location="Dubai, United Arab Emirates",
        url="https://boards.greenhouse.io/careem/jobs/8873949002",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={"team": "Careem Food - UAE (restaurant partners)",
             "via": "Talentmate ref 2610-2049-620 (posted 2026-10-09)",
             "workplace": "Dubai, 4 days office / 1 home",
             "level": "Associate"},
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not posted (Associate band — likely below 20k AED/month floor)",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": "2026-10-09", "raw": job.raw,
        "ai_score": 80, "ai_tier": "Hot",
        "skills_match": [
            "Glovo Account Manager – XL Accounts: restaurant chains on a food-delivery platform (same job)",
            "Miravia: 42 accounts, +30% GMV QoQ; 30+ new stores onboarded in two months",
            "Works with Careem, talabat and Noon from the partner side today (DoFreeze)",
            "AI automation (Claude system, ~40% less manual work; AI stock-risk dashboard), matching Careem's AI-talent push",
            "Excel / Google Sheets, promo measurement vs control, Business Administration degree",
        ],
        "missing_skills": [
            "Restaurant-partner work is from 2022-23 (Glovo); current role is supplier-side brand/marketing",
        ],
        "sector_fit": "excellent (food delivery account management; Glovo is the same model)",
        "seniority_fit": "slightly over (2–4 yrs asked; Paula 5) and a step down from Manager title",
        "red_flags": [
            "Associate band: salary likely under her 20k AED/month floor",
        ],
        "ats_keywords": ATS,
        "reasoning": "Direct fit from Glovo restaurant AM plus current Careem partner-side work; salary band is the main risk.",
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _pages(pdf: Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(pdf)).pages)


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    final_dir = settings.output_dir / DATE_FOLDER / f"{COMPANY} - {TITLE}"
    dest = final_dir / "01_CV_y_Carta"
    dest.mkdir(parents=True, exist_ok=True)

    cv_docx = cv._fill_template(CV_CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)

    moved = dest / cv_pdf.name
    shutil.move(str(cv_pdf), str(moved))
    shutil.copy(str(moved), str(dest / "Paula De Francisco - CV.pdf"))
    print("OK", moved.name, "PAGES", _pages(moved))

    stray = cv_pdf.parent.parent
    if stray.exists() and stray.resolve() != final_dir.resolve():
        shutil.rmtree(stray, ignore_errors=True)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
