"""One-off: CV de una página para **Brand Manager** de un cliente de **belleza** sin nombre (haircare, skincare y
otros productos de belleza). EAU, reporta al Brand Director. Oferta pegada por Guille el 2026-10-04 desde LinkedIn.

Es el **mismo JD** que el de Business Umbrella del 2026-10-01 (`_gen_business_umbrella_brand_manager_beauty.py`),
pero esta vez el anunciante añade dos requisitos de cribado: **grado universitario** y **más de 5 años de
experiencia en Market Research**. Por eso este CV sube la investigación de mercado (el bloque 2 del JD) y el
seguimiento de KPIs (bloque 10) frente a la versión del 1 de octubre. Además, aquel PDF salió con la foto cortada
y la línea de visado antigua; este coge la plantilla y el profile.yaml actuales.

Pide el JD: apoyar al Brand Director en estrategia y ejecución de marca; gestión diaria, identidad y guías visuales;
calendario de marca; **estudio de mercado y competencia** (tendencias, consumidor, precios, lanzamientos); campañas
en digital, social, retail, PR, eventos e influencers; posicionamiento y lanzamientos; apoyo a distribución y retail
(trade marketing, POS, sell-through); contenido con equipos creativos y agencias; comunicación de marca; trabajo
con ventas, educación, operaciones, finanzas y producto; relación con retailers, distribuidores e influencers;
**KPIs y reporting**; presupuesto (gasto, POs, facturas, ROI); formación de producto a ventas y educación.
Requisitos: grado en Marketing/Empresa; brand management, mejor en belleza, cosmética, haircare, skincare o FMCG;
digital, social, influencer y contenido; analítica. **No pide árabe.** Salario no publicado.

Ángulo honesto de la investigación de mercado (está en todos sus puestos, pero nunca fue su puesto principal):
- Mondelez: Nielsen, sell-in/sell-out, eficacia promocional e informes de categoría que alimentaban NPD.
- Miravia: estrategia de precios, productos de tendencia y análisis de ROI, ROAS, conversión y tráfico en 42 marcas.
- DoFreeze: investigación de mercado automatizada con el sistema de IA que montó; medición de campañas con grupo
  de control (SMASH × talabat frente a Befit) y auditoría del informe de reach de la agencia.

Guardarraíles de honestidad:
- **5+ años de Market Research**: no los tiene como puesto dedicado. No se escribe "5 years of market research"
  ni se presenta como researcher. La investigación aparece como parte del trabajo de marca y de cuenta.
- **Haircare**, **canal profesional / salones / equipo de educación**: sin experiencia. No se mencionan.
- **Brand-side en belleza**: no lo ha tenido. Su belleza es de marketplace (Miravia).
- **POS e in-store**: sin prueba sólida; se habla de promociones con ventas y distribuidores.
- **Forecast**: lo cierra la e-commerce manager; no aparece.
- **Deliveroo**: DoFreeze NO está. Plataformas: talabat, Noon, Careem.
- Creators: "103 activaciones en 3 campañas", nunca "una campaña con 103". Befit × Noon es contra baseline, no
  contra grupo de control.
- Sin árabe (aquí no se pide). Visa: la de profile.yaml ("UAE Employment Visa (employer-sponsored)").
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Confidential Beauty"
TITLE = "Brand Manager"
DATE_FOLDER = "2026-10-04"
JOB_ID = "confidential-beauty-brand-manager-research-2026-10"

JOB_DESCRIPTION = """\
Brand Manager. Confidential beauty company (haircare, skincare and other beauty products). Location: UAE.
Reports to: Brand Director.
Overview: The Brand Manager will support the Brand Director in managing and developing the company's portfolio of
beauty products, including haircare, skincare, and other beauty products. Responsible for the day-to-day management
and execution of brand strategies, marketing initiatives, product positioning, brand communications, and commercial
activities, with the objective of driving brand growth, awareness, and profitability. Works closely with internal
teams, distributors, retailers, suppliers, agencies, influencers and other external partners to ensure consistent
brand execution across all channels; monitors market trends, coordinates product launches, supports sales
activities, manages content and tracks brand performance.
Key responsibilities: 1. Brand management: support brand strategy for haircare, skincare and other beauty products;
manage day-to-day brand activities; ensure consistent brand identity, guidelines, messaging and visual standards;
maintain the brand calendar (campaigns, launches, activations, promotions). 2. Market research & competitive
analysis: industry trends, consumer needs, competitor activity, pricing, launches; regular market insights to the
Brand Director; recommend adjustments based on market and consumer insights. 3. Marketing campaigns across digital,
social media, retail, PR, events and influencer partnerships; coordinate timelines and deliverables; monitor
performance and optimize. 4. Product positioning & launches: positioning strategies, value propositions, launch
plans from planning through execution with sales, marketing and operations; post-launch performance.
5. Distribution & retail support: product distribution and visibility across key retail and professional channels;
promotional and trade marketing initiatives; brand standards in retail, displays and POS; sell-through.
6. Content creation & brand materials with creative teams, agencies and suppliers: product descriptions, campaign
visuals, social content, presentations, brochures, POS; review and approve content. 7. Brand communication across
digital, social, packaging, advertising, events and retail; support PR, influencer, media and partnerships.
8. Cross-functional collaboration with sales, marketing, education, operations, finance and product development;
coordinate agencies, suppliers and distributors. 9. Consumer engagement & relationship building with consumers,
retailers, distributors, influencers and educators; monitor reviews and sentiment. 10. Performance tracking &
reporting: KPIs on brand, sales, campaigns, launches and engagement; analyze sales and marketing data; regular
reports for the Brand Director. 11. Budget & marketing investment: support budgets, track expenses, purchase orders
and invoices, evaluate ROI. 12. Brand training & product knowledge for education and sales teams.
Skills & qualifications: Bachelor's degree in Marketing, Business, Communications or related; experience in brand
management or marketing, preferably within beauty, cosmetics, haircare, skincare or FMCG; brand strategy, consumer
behavior and commercial activities; project management; communication; cross-functional collaboration;
digital marketing, social media, influencer marketing and content development; analytical mindset; creative
thinking; presentation and reporting skills; knowledge of the beauty market is an advantage.
Requirements added by the job poster: Bachelor's Degree; 5+ years of work experience in Market Research.
"""

ATS = [
    "brand manager", "brand management", "brand strategy", "beauty", "cosmetics", "skincare", "haircare",
    "fragrances", "FMCG", "brand portfolio", "brand growth", "brand awareness", "brand identity",
    "brand guidelines", "brand calendar", "market research", "competitive analysis", "competitor analysis",
    "consumer insights", "market insights", "trends", "pricing", "Nielsen", "sell-in/sell-out",
    "promotional effectiveness", "control group", "campaigns", "launches", "activations", "promotions",
    "digital marketing", "social media", "influencer marketing", "product positioning", "product launches",
    "NPD", "retail", "trade marketing", "sell-through", "distributors", "content creation", "creative briefs",
    "agencies", "cross-functional", "KPIs", "performance tracking", "reporting", "budget", "ROI",
    "presentations", "Business Administration", "UAE", "Dubai",
]

CONTENT = {
    "headline": "Brand Manager · Beauty & FMCG · Market Research",
    "professional_summary": (
        "Brand manager in Dubai leading brand strategy, research, campaigns and launches for an FMCG group sold in "
        "50+ markets, with a team of two. Beauty grounding from Alibaba's Miravia: 42 beauty, fragrance and fashion "
        "brands, incl. KIKO Milano. Research-led since Nielsen at Mondelez; every campaign measured on sell-out."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG group | Befit, Eurocake, SMASH, Flair | talabat, Noon, Careem | 50+ markets",
            "bullets": [
                "Own the brand calendar for a multi-brand portfolio — seasonal campaigns (Ramadan, back to school, Fitness Month, New Year), activations and promotions — keeping identity, messaging and visual standards consistent across retail, e-commerce and social",
                "Run market, competitor and trend research — partly automated with a self-built AI system — and turn it into positioning for 6 product launches led end-to-end (brief, packaging, pricing, go-to-market)",
                "Lead a team of two (designer and social media executive), briefing and approving all campaign visuals, social content and product copy; manage four influencer agencies — 103 creator activations over three campaigns",
                "Measure every campaign on sell-out against control groups and baselines: SMASH × talabat +165% daily sales vs +4.7% for the control brand; Befit × Noon +31% over baseline; agency fees negotiated -30%, agency reach reports audited",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees",
            "bullets": [
                "Managed 42 beauty, fragrance and fashion brands — incl. KIKO Milano makeup & skincare — on assortment, pricing and promotional calendars, using pricing, trend and conversion analysis to deliver +30% GMV QoQ across the portfolio",
                "Created the Beauty Club and Hot on Social programmes to lift brand visibility; as PIC Fragrances, onboarded 30+ houses in two months (Arabian Oud, Lattafa, Swiss Arabian, Ajmal) around trend-driven products",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Grew XL accounts (KFC, Taco Bell, La Tagliatella) through data-led planning and joint marketing activations, and helped build the Retail vertical with fashion and lifestyle brands",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | chocolate category (Milka, Suchard)",
            "bullets": [
                "Market and category research with Nielsen — sell-in/sell-out and promotional-effectiveness analysis, category performance reports — identifying growth opportunities for NPD launches (Milka Spread, Mini Suchard)",
            ],
        },
    ],
    "skills_brand": (  # -> "Brand & Launches"
        "brand strategy & positioning, brand calendar, 6 NPD launches, campaigns, team leadership"
    ),
    "skills_ecommerce": (  # -> "Digital & Content"
        "social media, influencer & creator marketing, content & creative briefs, agency management"
    ),
    "skills_commercial": (  # -> "Retail & Trade"
        "trade promotions, distributors, modern trade, quick-commerce (talabat, Noon, Careem), ROI"
    ),
    "skills_data": (  # -> "Research & Insights"
        "market & competitor research, trends, Nielsen, sell-in/sell-out, control groups, KPI reporting"
    ),
    "skills_tools": (  # -> "Tools"
        "Excel, PowerPoint, Power BI, SAP, Salesforce, Meta & Google Ads, Canva, Adobe, Claude (AI)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Brand & Launches",
    "E-Commerce & Digital": "Digital & Content",
    "Commercial": "Retail & Trade",
    "Data & Analytics": "Research & Insights",
}


def make_job() -> Job:
    return Job(
        id=JOB_ID,
        title=TITLE,
        company=COMPANY,
        location="United Arab Emirates",
        url="https://www.linkedin.com/jobs/search/?keywords=Brand%20Manager%20haircare%20skincare%20UAE",
        source="linkedin",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Brand Manager beauty haircare skincare UAE",
            "note": "Mismo JD que Business Umbrella (2026-10-01), ahora con requisitos de cribado añadidos por el "
                    "anunciante: grado universitario y 5+ años en Market Research. Empresa no indicada. No pide "
                    "árabe. CV del 2026-10-04 con la investigación de mercado subida.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    existing = next((j for j in jobs if j.get("id") == job.id), None)
    entry = existing if existing is not None else {}
    entry.update({
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": entry.get("posted_date", DATE_FOLDER), "raw": job.raw,
        "ai_score": 74, "ai_tier": "Warm",
        "skills_match": [
            "Belleza real: 42 marcas de belleza, fragancia y moda en Miravia, incluida KIKO Milano (maquillaje y skincare)",
            "Fragancias árabes: incorporó Arabian Oud, Lattafa, Swiss Arabian y Ajmal como PIC Fragrances",
            "Investigación de mercado en todos sus puestos: Nielsen en Mondelez, precios y tendencias en Miravia, research con IA en DoFreeze",
            "Medición con grupo de control: SMASH × talabat +165% frente a +4,7% del control",
            "Calendario de marca por temporadas y 6 lanzamientos NPD end-to-end",
            "Contenido y agencias: dirige a diseñador y social media executive, 4 agencias, 103 activaciones en 3 campañas",
            "Grado en Administración de Empresas (CUNEF) y sin requisito de árabe",
        ],
        "missing_skills": [
            "5+ años de Market Research como puesto dedicado: no los tiene; la investigación es parte de sus roles (~4 años)",
            "Haircare: no lo ha trabajado",
            "Brand management en belleza desde la marca: su belleza es del lado marketplace (Miravia)",
            "Canal profesional (salones) y formación a equipos de educación: sin experiencia",
        ],
        "sector_fit": "bueno — belleza (skincare, fragancias) por Miravia y FMCG por DoFreeze",
        "seniority_fit": "en línea — Brand Manager que reporta a un Brand Director, frente a su Brand & Marketing Manager actual",
        "red_flags": [
            "La pregunta de cribado de 5+ años en Market Research puede descartarla automáticamente",
            "Mismo JD que Business Umbrella (2026-10-01): probablemente la misma vacante; no solicitar dos veces",
            "Empresa y salario no publicados",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Mismo Brand Manager de belleza que el de Business Umbrella, ahora con un filtro de 5+ años en Market "
            "Research. Paula encaja en el trabajo de marca (calendario, lanzamientos, contenido, agencias, "
            "influencers con venta medida) y trae belleza real del lado marketplace (Miravia, KIKO, fragancias "
            "árabes). La investigación de mercado está en todos sus puestos, pero nunca fue su puesto principal, "
            "así que la pregunta de cribado es el mayor riesgo. Le falta haircare, canal profesional y belleza "
            "desde la marca."
        ),
        "scored_by": "manual:claude", "freshness": entry.get("freshness", "fresh"),
        "discovered_at": entry.get("discovered_at", f"{DATE_FOLDER}T00:00:00+00:00"),
        "status": entry.get("status", "Reviewing"),
    })
    if existing is None:
        jobs.insert(0, entry)
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
