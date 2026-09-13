"""One-off: generate Paula's CV for **Product Marketing Manager (Crypto)** at **Revolut**
(remote, based in the UAE — LinkedIn + Indeed, posted 2026-08-11 / re-surfaced 2026-09-02).

JD ask: lead the crypto marketing roadmap. Long-term multi-channel marketing strategy for crypto
products; data-driven quarterly roadmaps that forecast, prioritise and deliver against departmental
KPIs; feedback loops turned into actionable performance insight; deep customer understanding in
core and expansion markets; launch strategies for new features with creative excellence and
flawless cross-functional execution; briefing creative teams, building audience strategies,
crafting value propositions, orchestrating PR. Needs B2B or B2C product/brand/consumer marketing
in a fast-paced environment (preferably a hyper-growth global tech company), GTM project
management across teams, channel fluency (in-product, digital, PR), and data → insight → strategy.
Nice to have: Tier-1 strategy consultancy, in-depth crypto knowledge.

Paula's honest angle:
- PRODUCT MARKETING IS HER DAY JOB, IN FMCG LANGUAGE: she leads NPD end-to-end — brief, positioning,
  pricing, packaging, go-to-market — for 6 launches across GCC, MENA, Asia, Europe, USA and Africa.
  That is launch strategy + value proposition + cross-functional execution, which is the spine of
  this role.
- HYPER-GROWTH GLOBAL TECH, TWICE: Miravia was Alibaba's brand-new marketplace launch in Spain
  (she owned Flash Sales reporting to the CEO and ran category expansion), and Glovo was
  quick-commerce at €500M+ scale where she helped build the retail vertical from zero.
- CORE AND EXPANSION MARKETS: she runs multi-market strategy across 50+ countries, so "how the
  product lands differently per market" is not theory for her.
- BRIEFING CREATIVE AT VOLUME: 25–50 creators briefed and negotiated per campaign, plus art
  direction, campaign messaging and content systems.
- DATA → ACTION: ROI, ROAS, conversion, traffic, retention and forecast-accuracy analysis feeding
  the next quarter's plan; built an AI reporting/planning system that cuts manual work ~40%.
- PROOF OF WORK: she already built and deployed a full fictional Revolut campaign, "No Borders"
  (https://guillemu11.github.io/revolut-no-borders/) — an unusually strong artefact for a PMM
  application and the single biggest differentiator to lead with in outreach.

Honesty guardrails: she has NOT worked in crypto, fintech or any regulated financial product, and
has NOT worked at a strategy consultancy. The CV claims consumer/product marketing, GTM and
hyper-growth platform experience — never crypto expertise, never fintech, never financial services.
No PR-agency ownership claim (she has earned-media and creator work, not orchestrated PR campaigns
at Revolut's scale) — phrased as "campaign and communications planning". No Arabic. Factual
"UAE Residence Visa" only — never "no sponsorship needed". Remote role based in the UAE.

ONE-PAGE standard: tight summary, 3/2/1/1 bullets, ~6-9 short skills per row.
Fills the real CV template, relabels the skills rows for this role, converts to PDF via
LibreOffice, registers the job, verifies 1 page, lands under output/2026-09-03/.
Also drops a short-named 'Paula De Francisco - CV.pdf' copy for portals / Easy Apply.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Revolut"
TITLE = "Product Marketing Manager (Crypto)"
DATE_FOLDER = "2026-09-03"

JOB_DESCRIPTION = """\
Product Marketing Manager (Crypto) — Revolut. Full-time. Remote, based in the United Arab Emirates.
About Revolut: 80+ million customers, 13,000+ people, products across spending, saving, investing,
exchanging and travelling. Marketing at Revolut is about clarity, creativity and commercial impact;
the Growth team turns bold ideas into campaigns that connect with millions.
About the role: lead the crypto marketing roadmap — articulate and drive marketing strategies that
promote crypto products, build trust in the industry, and bring features to life for customers.
What you'll be doing: create and oversee a long-term marketing strategy across channels for crypto
products; design impactful, data-driven quarterly roadmaps to forecast, prioritise and deliver
projects contributing to the Crypto department's KPIs and milestones; leverage feedback loops to
extract data-driven, actionable marketing performance insights that inform strategy; immerse
yourself in customers' needs and how crypto products function in core and expansion markets; shape
launch strategies for new crypto features with creative excellence and flawless cross-functional
execution; brief creative teams, develop audience strategies, craft value propositions and
orchestrate PR campaigns; develop solid cross-functional relationships in a fast-paced environment
to bring new products and campaigns to market.
What you'll need: experience in B2B or B2C product, brand or consumer marketing in a fast-paced
environment (preferably a hyper-growth global tech company); a passion for Revolut's products;
expertise setting strategic direction and project managing complex go-to-market campaigns requiring
input from various teams; a solid track record in consumer marketing and creative acumen, with an
understanding of channels including in-product, digital and PR; demonstrated ability to analyse and
distil data into actionable insights informing campaign strategy and product development; excellent
communication and interpersonal skills; charismatic people-person able to work in a dynamic,
international environment.
Nice to have: experience at a Tier 1 strategy consultancy; in-depth knowledge of crypto products
and the crypto market.
"""

ATS = [
    "product marketing", "product marketing manager", "PMM", "go-to-market", "GTM",
    "launch strategy", "product launch", "roadmap", "quarterly roadmap", "prioritisation",
    "value proposition", "positioning", "messaging", "audience strategy", "segmentation",
    "consumer marketing", "brand marketing", "B2C", "B2B", "campaign strategy",
    "creative briefing", "creative excellence", "art direction", "content strategy",
    "PR", "communications", "earned media", "influencer marketing", "UGC",
    "in-product", "digital", "paid social", "Meta Ads", "Google Ads", "CRM", "EDM",
    "cross-functional", "stakeholder management", "project management",
    "core and expansion markets", "multi-market", "international", "localisation",
    "customer insight", "feedback loops", "data-driven", "actionable insights",
    "KPI", "forecasting", "performance analysis", "ROI", "ROAS", "conversion", "retention",
    "GMV", "AOV", "A/B testing", "experimentation", "marketplace", "platform",
    "hyper-growth", "fast-paced", "e-commerce", "Shopify", "quick-commerce",
    "Dubai", "UAE", "remote", "generative AI",
]

CONTENT = {
    "headline": (
        "Product Marketing & Go-to-Market · Launch Strategy · Consumer Growth Across 50+ Markets · "
        "Hyper-Growth Marketplaces & E-Commerce"
    ),
    "professional_summary": (
        "Consumer product marketer with 5 years across hyper-growth platforms (Alibaba's Miravia, Glovo) "
        "and global consumer brands. Owns launches end-to-end — audience, positioning, pricing, channel "
        "plan, cross-functional execution — across 50+ core and expansion markets, and turns performance "
        "data into the next quarter's roadmap. Dubai-based, available remote."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Consumer brands (Befit, Eurocake, Flair) | 50+ markets | D2C + marketplaces",
            "bullets": [
                "Lead go-to-market end-to-end for 6 product launches — customer and category insight, positioning and value proposition, pricing, packaging brief, channel plan and launch phasing — across GCC, MENA, Asia, Europe, USA and Africa, adapting the story per core and expansion market",
                "Set the quarterly marketing roadmap against commercial KPIs and defend it with data: ROI/ROAS, conversion, traffic and retention feedback loops decide what gets scaled, reworked or cut, with results read back within weeks of launch",
                "Brief and direct creative at volume — 25–50 creators per campaign plus art direction, messaging and always-on content — and coordinate marketing, supply, sales and distributors so launches land on date in every market",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba Group's marketplace launch in Spain | Hyper-growth global tech | 100K+ employees",
            "bullets": [
                "Joined a brand-new Alibaba marketplace at launch and owned the Flash Sales channel for Beauty, Fashion and Home reporting to the CEO — building the commercial plan, the promotional calendar and the P&L case behind it",
                "Grew 42 accounts +30% GMV QoQ and led category expansion as PIC Fragrances (30+ stores onboarded in two months), plus the Beauty Club and Hot on Social programmes that gave the young marketplace a beauty identity",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce platform | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Helped build Glovo's retail vertical from zero — taking a food-delivery app into fashion, beauty and non-food — and ran launch activations with strategic partners (KFC, Taco Bell, Sushi Shop) across marketing, ops and support",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | Category planning & analytics",
            "bullets": [
                "Measured launch and promotional performance for the chocolate category and fed the read-outs into NPD decisions (Milka Spread, Mini Suchard)",
            ],
        },
    ],
    "skills_brand": (  # label -> "Product Marketing & GTM"
        "go-to-market strategy, product launches (NPD end-to-end), positioning & value propositions, "
        "audience strategy, quarterly roadmaps, pricing, brand strategy"
    ),
    "skills_ecommerce": (  # label -> "Channels & Creative"
        "in-product & onsite merchandising, paid social (Meta), Google Ads, CRM/EDM, "
        "creative briefing & art direction, creator & UGC programmes, communications planning"
    ),
    "skills_commercial": (  # label -> "Markets & Stakeholders"
        "50+ core & expansion markets, cross-functional leadership, distributor & partner management, "
        "marketplace & quick-commerce platforms, C-level reporting, negotiation"
    ),
    "skills_data": (  # label -> "Data & Insight"
        "customer & category insight, feedback loops, KPI forecasting, ROI, ROAS, conversion, "
        "retention, GMV, AOV, A/B testing, launch performance read-outs"
    ),
    "skills_tools": (  # label -> "Tools"
        "Power BI, Tableau, Looker, Google Analytics, Meta Ads Manager, Google Ads, Salesforce, "
        "Shopify, Adobe, Canva, Generative AI (Claude, ChatGPT)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Product Marketing & GTM",
    "E-Commerce & Digital": "Channels & Creative",
    "Commercial": "Markets & Stakeholders",
    "Data & Analytics": "Data & Insight",
}

OUTREACH = """\
# Outreach — Revolut · Product Marketing Manager (Crypto)

**Rol:** Product Marketing Manager (Crypto) · Remoto, con base en EAU · publicado 11-ago-2026, resurfaceado 2-sep
**Aplicar:** solo por canales oficiales de Revolut (revolut.com/careers). Nunca por terceros — la propia oferta avisa de estafas.
**Indeed:** https://to.indeed.com/aaybj2j97hq7

## El diferenciador: ya existe una campaña Revolut hecha por Paula

**"No Borders"** — landing de campaña completa, desplegada y navegable:
https://guillemu11.github.io/revolut-no-borders/

Esto es lo que hay que poner delante. Un PMM que llega con una campaña de Revolut ya construida
(concepto, audiencia, propuesta de valor, ejecución visual) responde en un link a la mitad de la
JD: "shape launch strategies… providing top-notch creative excellence". Ir con el link en la
primera línea del mensaje, no en un adjunto.

## Contacto

LinkedIn marca **1 contacto trabajando en Revolut** en la vista de la oferta. Antes de aplicar:
abrir la oferta → "1 contacto trabaja aquí" → pedir referral. En Revolut el referral interno pesa
mucho más que el ATS, y el proceso es notoriamente duro (varias rondas + case study).

## Mensaje sugerido (LinkedIn, <300 caracteres si es nota de Top Choice)

> Hola [nombre], veo que estás en Revolut — acabo de aplicar al PMM de Crypto. Vengo de product
> marketing de consumo en Miravia (Alibaba) y Glovo, lanzando en 50+ mercados. Me hacía tanta
> ilusión el rol que monté una campaña Revolut de cero: [link]. ¿Te importaría referirme?

## Ángulo en la aplicación

1. **Product marketing, en idioma FMCG.** 6 lanzamientos end-to-end (brief, posicionamiento,
   pricing, packaging, GTM) en 50+ mercados = exactamente "shape launch strategies for new
   features" + "craft value propositions".
2. **Hyper-growth global tech, dos veces.** Miravia era el marketplace nuevo de Alibaba en España
   (canal Flash Sales reportando al CEO) y Glovo construyó su vertical retail desde cero. La JD
   pide "preferably in a hyper-growth global tech company" — lo tiene, aunque no sea fintech.
3. **Core and expansion markets.** No es teoría: gestiona GCC, MENA, Asia, Europa, USA y África.
4. **Data → insight → roadmap.** ROI/ROAS/conversión/retención alimentando el plan del trimestre
   siguiente, más el sistema de reporting con IA que se construyó ella.

## Lo que NO se dice

- **Cero experiencia en crypto, fintech o producto financiero regulado.** Es un "nice to have" en
  la JD, no un must — pero no se maquilla. Si sale en entrevista: se reconoce y se compensa con
  una lectura propia del mercado crypto de Revolut (preparar 2-3 opiniones reales antes).
- **Nada de consultoría estratégica Tier 1.** Otro nice-to-have que no tiene.
- **Ni "orchestrating PR campaigns" a escala Revolut.** Tiene earned media y creadores, no gabinete
  de prensa. En el CV va como "communications planning".
- **Nunca "no necesito sponsorship".** Visado de residencia EAU patrocinado por su empresa actual.

## Riesgos

- Revolut filtra durísimo por pedigrí (ex-consultoría, ex-tech). Sin referral, la probabilidad de
  pasar el CV screen es baja.
- Es remoto, fuera de su preferencia declarada.
- Sin salario publicado — confirmar contra el suelo de 20.000 AED/mes.
"""


def make_job() -> Job:
    return Job(
        id="revolut-product-marketing-manager-crypto-uae-remote-2026-09",
        title=TITLE,
        company=COMPANY,
        location="United Arab Emirates (Remote)",
        url="https://www.revolut.com/careers/",
        source="indeed",
        description=JOB_DESCRIPTION,
        raw={
            "query": "Revolut Product Marketing Manager Crypto remote UAE",
            "indeed_url": "https://to.indeed.com/aaybj2j97hq7",
            "note": "Posted 2026-08-11 on Indeed, re-surfaced on LinkedIn 2026-09-02. Remote, based in "
                    "the UAE. LinkedIn shows 1 first-degree contact at Revolut — chase the referral before "
                    "applying. Apply ONLY via revolut.com/careers (the JD warns about recruitment scams). "
                    "Paula already has a deployed fictional Revolut campaign, 'No Borders' "
                    "(https://guillemu11.github.io/revolut-no-borders/) — lead with it. Real gaps: no "
                    "crypto/fintech/regulated-finance experience and no Tier-1 consultancy (both listed as "
                    "nice-to-have). Revolut is a stated priority target for Paula.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": None, "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": "2026-08-11", "raw": job.raw,
        "ai_score": 74, "ai_tier": "Warm",
        "skills_match": [
            "Product marketing is her actual day job in FMCG language: 6 NPD launches end-to-end (brief, positioning, pricing, packaging, GTM) — the JD's 'shape launch strategies' and 'craft value propositions'",
            "Hyper-growth global tech twice over: Miravia was Alibaba's new marketplace launch in Spain (owned Flash Sales, reported to the CEO) and Glovo built its retail vertical from zero",
            "Core and expansion markets for real — multi-market strategy across GCC, MENA, Asia, Europe, USA and Africa (50+ countries)",
            "Briefs creative at volume: 25–50 creators per campaign plus art direction, messaging and content systems",
            "Data into strategy: ROI, ROAS, conversion, traffic and retention feedback loops driving the next quarter's plan, plus a self-built AI reporting system",
            "Already built and deployed a full fictional Revolut campaign ('No Borders') — a rare, concrete artefact for a PMM application",
            "Revolut is a stated priority target — genuine 'passion for Revolut's products' is credible here",
        ],
        "missing_skills": [
            "No crypto experience and no in-depth knowledge of crypto products or the crypto market (listed as nice-to-have)",
            "No fintech, banking or regulated financial-product marketing at all",
            "No Tier-1 strategy consultancy background (nice-to-have, and a common Revolut screen)",
            "No formal Product Marketing Manager title — PMM work sits inside brand/e-commerce and KAM roles",
            "PR is her weakest channel: earned media and creator work, not orchestrated PR campaigns at scale",
        ],
        "sector_fit": "adjacent — consumer marketing and hyper-growth platforms are a direct hit; crypto and fintech are entirely new",
        "seniority_fit": "good — 5 years consumer/product marketing, and the role is a hands-on IC manager rather than a people-leadership post",
        "red_flags": [
            "Revolut screens hard on pedigree (ex-consultancy, ex-big-tech); without the internal referral the CV screen is a real wall",
            "Multi-stage process with case study — long time-to-offer",
            "Remote role, outside Paula's stated preference (remote_ok: false in profile.yaml)",
            "No salary published — confirm against the AED 20K/month floor",
            "Apply only via revolut.com/careers; the JD explicitly warns about recruitment scams via third parties",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Strong-adjacent fit and worth the effort because of one asymmetry: Paula already has a "
            "deployed Revolut campaign ('No Borders'), which answers the JD's creative-excellence ask in a "
            "single link and is the kind of artefact almost no other applicant brings. On substance, the "
            "role's spine — launch strategy, value propositions, audience strategy, briefing creative, "
            "cross-functional GTM across core and expansion markets, data feeding the quarterly roadmap — "
            "is what she does today for 6 launches across 50+ countries, and she has done it inside two "
            "hyper-growth platforms (Alibaba's Miravia at launch, Glovo's retail vertical from zero). The "
            "honest gap is the category: zero crypto, zero fintech, zero regulated financial product, plus "
            "no Tier-1 consultancy. Both are nice-to-haves rather than must-haves, so the CV leads on "
            "consumer product marketing and platform growth and claims no crypto expertise whatsoever. "
            "The decisive move is the referral: LinkedIn shows one first-degree contact at Revolut, and at "
            "Revolut a referral is worth more than the ATS."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    outreach_dir = final_dir / "03_Outreach"
    outreach_dir.mkdir(parents=True, exist_ok=True)
    (outreach_dir / "Outreach.md").write_text(OUTREACH, encoding="utf-8")
    print("OK_OUTREACH", outreach_dir / "Outreach.md")

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
