"""One-off: CV de una página para **Brand Manager** en **Apparel Group** (Dubái; 85+ marcas de moda y lifestyle en
franquicia, 2.500+ tiendas en el GCC, India y Asia; dueño del programa de fidelización **Club Apparel**).
Oferta pegada por Guille el 2026-10-04; el texto no dice empresa, título ni ciudad, pero cita "Club Apparel", que es
de Apparel Group. Se asume Brand Manager en Dubái.

Es un Brand Manager de **retail en franquicia**, no de marketing: dueño de la venta y el margen de las tiendas de
una marca. Pide el JD: estrategia de lanzamiento y posicionamiento de la marca con el Brand GM y Marketing; control
de costes; productividad de tienda y **sell-through**; cumplir el plan financiero (**venta, markdowns, margen,
inventario medio**); **plan de merchandise** (precio, promoción, surtido), **plan y presupuesto de compra**,
selección de producto; gestión en temporada (stock balancing, markdowns, stock vs venta); awareness de Club
Apparel; estrategia de **visual merchandising**; investigación de mercado, tendencias y **mapa de precios de la
competencia**; búsqueda y evaluación de **locales nuevos** y coordinación de aperturas y re-fits; reclutar y formar.
No pide árabe ni años concretos.

Ángulo honesto de Paula:
- Surtido, precio y promociones: 42 marcas de beauty, fragancia y moda en Miravia (KIKO Milano incluida), +30% GMV
  QoQ; dueña del canal Flash Sales de Beauty, Fashion & Home reportando al CEO (selección de producto y mecánicas
  contra objetivos de P&L).
- Lanzamientos y posicionamiento: 6 lanzamientos end-to-end en DoFreeze (brief, packaging, precio, go-to-market),
  SMASH como marca nueva; abrió la categoría de fragancias en Miravia (30+ tiendas en dos meses).
- Venta y promoción medidas: sell-in/sell-out y eficacia promocional con Nielsen (Mondelez); SMASH × talabat +165%
  frente al +4,7% del control; agencias -30% (control de costes).
- Tienda de moda: un año en Massimo Dutti (Inditex) — venta, styling y estándares de VM de Inditex.
- Equipo: lidera a un diseñador y a un social media executive (formación y coaching, no contratación).

Guardarraíles de honestidad:
- **Compras / OTB / plan de compra**: no lo ha hecho. No se escribe "buying", "open-to-buy" ni "buying budget".
- **P&L de tiendas, markdowns, stock balancing, inventario medio**: no los ha gestionado. Nada de "sell-through",
  "markdown" ni "stock" como logro.
- **Aperturas, re-fits, evaluación de locales**: nunca. No aparecen.
- **VM**: lo vivió como dependienta (estándares de Inditex), no diseñó una estrategia de VM.
- **Mapa de precios de la competencia**: se habla de investigación de mercado y competencia para lanzamientos, no
  de un price-mapping formal.
- **Reclutar**: no consta. "Lead, train and coach", nunca "hired/recruited".
- **Forecast**: lo cierra la e-commerce manager de DoFreeze; no aparece.
- **Deliveroo**: DoFreeze NO está. Plataformas: talabat, Noon, Careem.
- Creators: "103 activaciones en 3 campañas", nunca "una campaña con 103".
- Visa: "UAE Employment Visa (employer-sponsored)", nunca "no sponsorship".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Apparel Group"
TITLE = "Brand Manager"
DATE_FOLDER = "2026-10-04"
JOB_ID = "apparel-group-brand-manager-2026-10"

JOB_DESCRIPTION = """\
Brand Manager (title assumed). Apparel Group (inferred from "Club Apparel"). Location not stated (assumed Dubai, UAE).
POSITION OBJECTIVE: The position is responsible for overall brand development and implementation for maximized sales
and profit. Strategize for brand penetration and positioning within the region to create brand image in the market.
Key Responsibilities.
Brand Growth and Profitability: Maximize sales and profitability of brand stores in line with Company targets. Develop
the brand launch strategy in consultation with the Brand General Manager. Control costs within budgetary guidelines for
the brand stores. Maximize productivity in stores through effective deployment of resources at brand stores to ensure
complete sell through as per brand product, style and assortment plans. Review feedback from territories on Brand
acceptance and customer expectations. Ensure adherence to financial plan in all key areas: sales, markdowns, margin and
average inventory. Create the brand positioning in consultation with the Marketing team and Brand Principal / Brand
General Manager. Involve in recruitment and training of team members.
Brand Merchandise Planning and Product Selection: Study merchandise requirements for assigned brand by setting sourcing
triggers based on actual sales, sales forecasts, company order parameters, inventory checks, forth coming events,
replenishment needs. Design the merchandise plan (pricing, promotion, assortment etc.) with the team for timely
availability of stocks at the stores. Determine buying requirements and formulate buying plan and budget. Conduct
product selection in coordination with the buying requirements and the product research received. Ensure effective
in-season management across territories (stock balancing, markdown management, stock analysis, stock vs. sales
performance).
Brand Image: Create brand awareness "Club Apparel". Coordinate with Retail Operations to collate the feedback on retail
and brand operations. Develop Strategy for Visual Merchandising at Stores and ensure implementation of the same.
Market and Competition Research: Acquire market intelligence through various sources and analyze trends that may impact
business. Research the brand acceptance within the region or any new territory and identify any inputs on brand
customizations required. Map competition prices / products and provide qualitative inputs to business. Research and
look at new sites for brand outlets and discuss the same with the Operations Manager and General Manager. Research on
the latest trends in products, brands, styles, designs, fits etc.
Projects & Administration: Evaluate the selected site on parameters like trade area, customer base, additional
merchandising considerations and seek internal approvals. Co-ordinate for new store openings and re-fits with the
operations and projects team. Ensure brand outlet locations are in line with the brand strategy and positioning.
Ensure adherence to the store opening plan.
"""

ATS = [
    "Brand Manager", "brand development", "brand positioning", "brand image", "brand awareness",
    "brand launch strategy", "brand penetration", "go-to-market", "launches", "sales and profitability",
    "retail", "fashion retail", "apparel", "beauty", "fragrances", "lifestyle", "Inditex", "Massimo Dutti",
    "KIKO Milano", "merchandise plan", "assortment", "assortment planning", "pricing", "promotions",
    "promotional calendar", "product selection", "sell-in/sell-out", "sales performance", "margin",
    "P&L targets", "cost control", "budget", "visual merchandising", "market intelligence", "market research",
    "competitor research", "trends", "new territory", "customer expectations", "Nielsen", "category planning",
    "team leadership", "training", "coaching", "cross-functional", "retail operations", "loyalty",
    "KPI reporting", "Excel", "PowerPoint", "SAP", "Power BI", "UAE", "Dubai", "GCC",
]

CONTENT = {
    "headline": "Brand Management · Fashion & Beauty Retail · Assortment, Pricing & Promotions · Brand Launches",
    "professional_summary": (
        "Brand manager with 4+ years across fashion, beauty and consumer brands. Owns brand positioning, launches "
        "and promotions for a homegrown UAE group with a team of two, measured on sales. Previously ran assortment, "
        "pricing and promotions for 42 beauty and fashion brands at Alibaba; started in Inditex retail."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Homegrown UAE F&B / FMCG group | Befit, Eurocake, SMASH, Flair | GCC + 50+ markets",
            "bullets": [
                "Own brand positioning and launch strategy for a multi-brand portfolio (Befit, Eurocake, the new SMASH brand, Flair): 6 launches end-to-end, from brief, packaging and pricing to the channel plan across retail, talabat, Noon, Careem and our Shopify store",
                "Run market, trend and competitor research before every launch and seasonal moment (Ramadan, back to school, Fitness Month, New Year) to set positioning, price points and the promotional calendar",
                "Control campaign and agency costs (fees negotiated -30%) and measure every promotion on sales: SMASH × talabat lifted daily sales +165% vs +4.7% for the control brand; Befit × Noon +31% over baseline",
                "Built brand awareness from zero with 103 creator activations over 3 campaigns, samplings in modern trade and on talabat and community events; lead, train and coach a team of two (graphic designer, social media executive)",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | beauty, fashion & lifestyle",
            "bullets": [
                "Managed 42 beauty, fragrance and fashion brands (incl. KIKO Milano) on assortment, pricing and promotional calendars, reviewing sales, conversion and traffic weekly to deliver +30% GMV QoQ",
                "Owned the Flash Sales channel for Beauty, Fashion & Home, reporting to the CEO: product selection, deal pricing and calendar against P&L targets; opened the fragrance category with 30+ new stores in two months",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue",
            "bullets": [
                "Helped build the Retail vertical, onboarding fashion and lifestyle brands, and grew XL accounts (KFC, Taco Bell) through data-led planning with marketing, operations and customer support",
            ],
        },
        {
            "company": "Mondelez International · Massimo Dutti (Inditex)",
            "role": "Trainee, Category Planning · Sales Associate",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG, Nielsen | Premium fashion retail, Las Rozas Village store",
            "bullets": [
                "At Mondelez, category planning with Nielsen (sell-in/sell-out, promotional effectiveness) and NPD launches (Milka Spread, Mini Suchard); at Massimo Dutti, shop-floor sales, styling and Inditex visual merchandising standards",
            ],
        },
    ],
    "skills_brand": (  # -> "Brand & Strategy"
        "brand positioning, launch strategy & go-to-market, brand awareness, seasonal calendar, loyalty projects"
    ),
    "skills_ecommerce": (  # -> "Merchandising"
        "assortment planning, product selection, pricing & promotions, Inditex visual merchandising standards"
    ),
    "skills_commercial": (  # -> "Commercial & Team"
        "sales & promotion analysis, cost & agency control, P&L targets, team leadership, training & coaching"
    ),
    "skills_data": (  # -> "Research & Data"
        "market, trend & competitor research, Nielsen, sell-in/sell-out, KPI reporting, Power BI"
    ),
    "skills_tools": (  # -> "Tools"
        "MS Office (Excel, PowerPoint), SAP, Salesforce, Power BI, Adobe (Photoshop, Illustrator), Canva, Claude (AI)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Brand & Strategy",
    "E-Commerce & Digital": "Merchandising",
    "Commercial": "Commercial & Team",
    "Data & Analytics": "Research & Data",
}


def make_job() -> Job:
    return Job(
        id=JOB_ID,
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE (assumed — not stated in the posting)",
        url="https://www.linkedin.com/jobs/search/?keywords=Apparel%20Group%20Brand%20Manager",
        source="linkedin",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Apparel Group Brand Manager",
            "note": "Empresa inferida por la mención a Club Apparel (fidelización de Apparel Group). El JD no dice "
                    "título ni ciudad. Brand Manager de retail en franquicia: venta, margen, markdowns, plan de "
                    "compra, VM y aperturas. No pide árabe. CV del 2026-10-04.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    existing = next((j for j in jobs if j.get("id") == job.id), None)
    entry = existing if existing is not None else {}
    entry.update({
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": entry.get("posted_date", DATE_FOLDER), "raw": job.raw,
        "ai_score": 62, "ai_tier": "Stretch",
        "skills_match": [
            "Surtido, precio y promociones para 42 marcas de beauty, fragancia y moda en Miravia (KIKO incluida), +30% GMV QoQ",
            "Canal Flash Sales reportando al CEO: selección de producto y precio contra objetivos de P&L",
            "Lanzamientos y posicionamiento: 6 lanzamientos end-to-end en DoFreeze, SMASH como marca nueva",
            "Abrir categoría: fragancias en Miravia, 30+ tiendas en dos meses",
            "Retail de moda: un año en tienda en Massimo Dutti (Inditex), estándares de VM",
            "Análisis de venta y promoción: sell-in/sell-out y Nielsen en Mondelez; SMASH × talabat +165% vs control",
            "Control de costes (agencias -30%) y equipo de dos al que forma",
        ],
        "missing_skills": [
            "Plan de compra, presupuesto de compra y open-to-buy: nunca lo ha hecho",
            "P&L de tiendas: venta, markdowns, margen e inventario medio de una red de tiendas",
            "Gestión en temporada: stock balancing, markdowns, stock vs venta entre territorios",
            "Aperturas, re-fits y evaluación de locales: sin experiencia",
            "Estrategia de visual merchandising: la conoce desde la tienda, no la ha diseñado",
            "Reclutar: forma a su equipo, no consta que haya contratado",
        ],
        "sector_fit": "bueno — moda y lifestyle en retail; su moda y beauty vienen de plataforma y de tienda",
        "seniority_fit": "en el límite — el puesto es de dueño de P&L de tiendas y compras, que Paula no ha llevado",
        "red_flags": [
            "Es un rol de merchandising y compras de retail, no de marketing: buena parte del JD (plan de compra, markdowns, stock, aperturas) queda fuera de su experiencia",
            "Empresa inferida por Club Apparel; título y ciudad no aparecen en el texto",
            "Salario no publicado: confirmar que llega al suelo de 20.000 AED/mes",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Encaje parcial. Apparel Group busca un Brand Manager de franquicia que responda de la venta y el margen "
            "de las tiendas de una marca: plan de merchandise, plan de compra, markdowns, stock en temporada, VM y "
            "aperturas. Paula cubre bien la parte de marca (posicionamiento, lanzamientos, awareness) y la de surtido, "
            "precio y promociones (42 marcas de beauty y moda en Miravia, Flash Sales con el CEO), y conoce la tienda "
            "de moda desde Massimo Dutti. No ha hecho compras ni open-to-buy, no ha llevado el P&L de una red de "
            "tiendas y no ha abierto locales. El CV se apoya en lo que sí ha hecho y no promete lo otro. Vale la "
            "pena enviarlo si la marca es española o de beauty, donde su perfil pesa más."
        ),
        "scored_by": "manual:claude", "freshness": entry.get("freshness", "fresh"),
        "discovered_at": entry.get("discovered_at", f"{DATE_FOLDER}T00:00:00+00:00"),
        "status": entry.get("status", "Reviewing"),
    })
    if existing is None:
        jobs.insert(0, entry)
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
