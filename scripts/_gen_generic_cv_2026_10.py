"""One-off: refreshed GENERIC, role-agnostic CV of Paula's full profile (2026-10-12).

Replaces the 2026-08-28 generic CV (scripts/_gen_generic_cv.py), which predates several fixes:
the photo-clipping bug (fixed 2026-10-01), the "UAE Employment Visa (employer-sponsored)" header,
the one-page standard, and facts learned since (team of two, KIKO Milano, NIVEA three months).
It also listed Deliveroo as a DoFreeze platform, which is wrong.

Balanced across Brand & Marketing + E-Commerce + Key Account Management, for postings with no JD
to tailor against (recruiter groups, referrals, speculative applications). Content from
profile.yaml and the standing memory notes — nothing invented.

Honesty guardrails: no Arabic; never "own the forecast"; DoFreeze is NOT on Deliveroo; no KIKO /
NIVEA-specific metrics (the +30% GMV QoQ is the whole 42-brand portfolio).

ONE-PAGE CV. Lands under output/2026-10-12/CV_Generico/ as PDF + editable DOCX.
No dashboard registration.
"""
from __future__ import annotations

import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

DATE_FOLDER = "2026-10-12"
TARGET_SUBDIR = "CV_Generico"
CLEAN_STEM = "Paula De Francisco - CV"

CV_CONTENT = {
    "headline": (
        "Brand & Marketing · E-Commerce · Key Account Management · Trade Marketing "
        "| FMCG · Beauty · Fashion"
    ),
    "professional_summary": (
        "Brand, e-commerce and commercial professional with 5 years across FMCG, beauty and fashion. "
        "Brand & Marketing Manager at DoFreeze in Dubai, leading a team of two across 50+ markets; "
        "previously grew 42 beauty and fashion brands +30% GMV QoQ at Alibaba's Miravia."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE F&B / FMCG group | Befit, Eurocake, Flair | sold in 50+ countries",
            "bullets": [
                "Lead brand and marketing across 50+ markets with a team of two (designer, social media exec); led 6 NPD launches end-to-end (brief, packaging, pricing, go-to-market)",
                "Built the influencer programme from zero to 25–50 creators per campaign, with sampling and seeding across modern trade and quick-commerce; run Meta and Google Ads on ROAS",
                "Onboarded the brands onto Noon, talabat and Careem and read every promotion against a control: SMASH x talabat lifted daily sales +165% vs +4.7% for the control brand",
                "Own the Shopify store end-to-end (catalogue, UX, checkout); built a Claude-based system that automates campaign planning, content and KPI reporting (~40% less manual work)",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace in Spain | 100K+ employees",
            "bullets": [
                "Managed 42 beauty, fragrance and fashion brands, incl. KIKO Milano and NIVEA (three months), to +30% GMV QoQ through pricing, assortment and promotions",
                "Owned the Flash Sales channel for Beauty, Fashion & Home, reporting directly to the CEO against P&L targets",
                "As PIC Fragrances, onboarded 30+ stores in two months (incl. the distributors of Arabian Oud, Lattafa, Swiss Arabian, Ajmal); created the Beauty Club and Hot on Social programmes",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Helped build the Retail vertical, onboarding fashion, beauty and lifestyle brands; managed XL accounts (KFC, Taco Bell, Sushi Shop) and negotiated their commercial deals",
            ],
        },
        {
            "company": "Mondelez International · Massimo Dutti (Inditex)",
            "role": "Category Planning Trainee · Sales Associate",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG category planning | Premium fashion retail floor",
            "bullets": [
                "Sell-in/sell-out and promo-effectiveness analysis for the chocolate category at Mondelez, supporting NPD (Milka Spread, Mini Suchard); Inditex retail-floor grounding",
            ],
        },
    ],
    "skills_brand": (
        "brand strategy, NPD, influencer & UGC, sampling, trade & shopper marketing, team leadership"
    ),
    "skills_ecommerce": (
        "Shopify, CRO, Noon, talabat, Careem, Miravia (Alibaba), Meta Ads, Google Ads, EDM"
    ),
    "skills_commercial": (
        "key account management, negotiation, pricing, assortment, distributors, modern trade"
    ),
    "skills_data": (
        "P&L targets, GMV, sell-in/sell-out, ROI, ROAS, promo uplift vs control, Power BI"
    ),
    "skills_tools": (
        "Claude / Claude Code, Shopify, Meta Ads Manager, SAP, Salesforce, Nielsen, Excel (Expert)"
    ),
}

ROLE_LABELS = {"Tools": "Tools & AI"}


def make_job() -> Job:
    # Synthetic "job" only drives the template's output path; renamed afterwards.
    return Job(
        id="generic-profile-cv-2026-10",
        title="Perfil General",
        company="Paula",
        location="Dubai, United Arab Emirates",
        url="",
        source="manual",
        description="Generic profile CV for circulation.",
        raw={},
    )


def _relabel(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _pages(pdf: Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(pdf)).pages)


def main() -> None:
    tmp_docx = cv._fill_template(CV_CONTENT, make_job())
    _relabel(tmp_docx)

    dest_dir = settings.output_dir / DATE_FOLDER / TARGET_SUBDIR
    dest_dir.mkdir(parents=True, exist_ok=True)
    clean_docx = dest_dir / f"{CLEAN_STEM}.docx"
    shutil.move(str(tmp_docx), str(clean_docx))

    # Clean up the synthetic job folder the template created at the output root.
    tmp_job_dir = tmp_docx.parent.parent
    if tmp_job_dir.exists() and tmp_job_dir.name.startswith("Paula - "):
        shutil.rmtree(tmp_job_dir, ignore_errors=True)

    pdf = _to_pdf_soffice(clean_docx)
    print("OK_CV", pdf, "PAGES", _pages(pdf))


if __name__ == "__main__":
    main()
