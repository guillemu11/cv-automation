"""One-off: CV de una página para **Precision Marketing** en **Beiersdorf** (NIVEA, Eucerin, La Prairie,
Hansaplast, Labello). **Contrato de 7 meses.** Contacto: Kartik Kulshrestha, Talent Acquisition Lead.
Oferta pegada por Guille el 2026-10-04 desde LinkedIn; el texto no da título exacto ni ciudad. Se asume Dubái
(Beiersdorf Middle East) y el título de trabajo "Precision Marketing Manager".

Pide el JD: planificar y ejecutar **campañas digitales full-funnel** con principios de Precision Marketing;
estrategia de audiencias, selección de canales, definición de KPIs y optimización; trabajar con la **agencia de
medios** y **gestionar y retar sus recomendaciones** contra objetivos y ROI; eficiencia y eficacia de medios;
integrar con categoría, **influencer**, **ecommerce** y agencia a lo largo del consumer journey; crecer en
retailers **omnicanal y pure-play**. Medición: frameworks, fuentes de datos, **atribución** y tracking; rutinas de
governance y revisión; análisis **post-campaña** con recomendaciones; cultura **test-and-learn** (audiencia,
canal, creatividad, medios). Mejora continua: cambios de plataformas y algoritmos, nuevas capacidades.
Perfil: **performance marketing en el lado marca como último puesto**; idealmente antes **agencia de medios**;
analítica; pasar objetivos de marketing a ejecución digital; **ad tech: GMP, DSP, DMP**.

Ángulo honesto de Paula:
- Lado marca hoy (DoFreeze): lleva Meta y Google Ads, creators, EDM y la tienda Shopify para varias marcas
  (Befit, Eurocake, SMASH, Flair), con KPIs de ROI y ROAS. Es performance dentro de un rol de marca.
- Medición con grupo de control (lo más cercano a "incrementality" y "test-and-learn"): SMASH × talabat +165%
  de venta diaria frente a +4,7% del control (Befit, misma plataforma y ventana); ~AED 32K incrementales sobre
  ~AED 8,7K. Befit × Noon +31% sobre baseline, +4.176 unidades incrementales.
- "Manage and challenge agency recommendations": auditó el informe de la agencia (reach estimado desde
  seguidores, no medido; doble conteo de cuentas únicas; mismo creador con 5× de diferencia entre oleadas) y
  sacó que los micro (<20K) rendían 10–40× la media de la campaña. Negoció -30% de fees (AED 81 por pieza).
- Influencer + ecommerce + plataformas integrados: sorteo in-app y samplings con talabat, creators que mandan a
  comprar a talabat, Noon, Careem y la web propia.
- Pure-play desde dentro: 2 años en Miravia (Alibaba), 42 marcas de beauty (KIKO incluida), +30% GMV QoQ,
  Flash Sales reportando al CEO. Beauty = la categoría de Beiersdorf.
- Reporting: sistema de IA (Claude) para planificación y reporting de campañas, ~40% menos trabajo manual.

Guardarraíles de honestidad:
- **Sin agencia de medios** en su carrera. No se insinúa.
- **Sin GMP / DV360 / DSP / DMP / programática.** No aparecen en el CV. Sus plataformas: Meta Ads y Google Ads
  (Search y Display), Google Analytics (certificada).
- **Atribución**: no ha montado modelos (MMM, MTA, geo-lift). Se dice "control groups, baselines, incremental
  sales", nunca "attribution modelling".
- **Retail media de pago en talabat/Noon/Careem**: no confirmado. Se habla de campañas con creators,
  promociones, sampling y sorteo, no de "sponsored placements".
- **Nunca tuvo el título "Performance Marketing"**: su título es Brand & Marketing Manager.
- **Forecast**: lo cierra la e-commerce manager; no aparece. **Deliveroo**: DoFreeze NO está.
- Creators: "103 activaciones en 3 campañas", nunca "una campaña con 103".
- Sin árabe (aquí no se pide). Visa: "UAE Employment Visa (employer-sponsored)", nunca "no sponsorship".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Beiersdorf"
TITLE = "Precision Marketing Manager (7-month contract)"
DATE_FOLDER = "2026-10-04"
JOB_ID = "beiersdorf-precision-marketing-contract-2026-10"

JOB_DESCRIPTION = """\
Precision Marketing (title not stated in the posting). Beiersdorf (NIVEA, Eucerin, La Prairie, Hansaplast,
Labello). Contract Duration: 7 months. Location not stated (assumed Dubai, Beiersdorf Middle East).
Orchestrate Precision Marketing Activities: Lead the planning and execution of full-funnel digital campaigns,
ensuring alignment with business objectives and Precision Marketing principles. Guide audience strategy, channel
selection, KPI definition, and campaign optimization to maximize performance. Work closely with media agency teams
to drive campaign excellence, optimization, and business results. Manage and challenge agency recommendations,
ensuring media strategies and plans deliver against business objectives and ROI expectations. Drive Precision
Marketing excellence through strong media efficiency & effectiveness. Partner with category, influencer,
ecommerce, and agency teams to ensure integrated campaign execution across the full consumer journey. Partner
closely with Ecommerce teams to drive growth across omnichannel and pure-play retailers.
Define, Monitor & Optimize Measurement Approach: Define campaign measurement frameworks and ensure the right data
sources, attribution models, and tracking capabilities are in place to measure success and support optimization.
Establish governance and performance review routines to monitor campaign effectiveness and drive continuous
improvement. Strengthen tracking and reporting capabilities to provide clearer visibility on campaign performance
& business impact. Conduct post-campaign analyses, identify key insights and opportunities, translating into
actionable recommendations. Drive a test-and-learn culture by identifying opportunities for audience, channel,
creative, and media optimization.
Drive Continuous Improvement & Capability Building: Stay up to date with digital platform, algorithm, and
measurement changes, translating implications into actionable recommendations for the business. Identify
opportunities to improve campaign performance through new capabilities, audience strategies, media solutions, and
emerging best practices. Ensure key processes, tools, and frameworks remain effective and fit for purpose across
the organization.
Your Profile: Previous experience in "performance marketing" on brand side (as last role). Prior to that, ideal if
experience from within media agency (before moving to brand). Analytical skills. Ability to effectively translate
marketing objectives into digital execution. Comprehensive knowledge on digital campaigns execution and planning.
Knowledge of advertising tech: GMP, DSP, DMP, etc.
For any queries, reach out to Kartik Kulshrestha - Talent Acquisition Lead.
"""

ATS = [
    "Precision Marketing", "performance marketing", "brand side", "full-funnel", "digital campaigns",
    "campaign planning", "campaign execution", "audience strategy", "channel selection", "KPI definition",
    "campaign optimization", "media efficiency", "media effectiveness", "media agency", "agency management",
    "challenge agency recommendations", "ROI", "ROAS", "business objectives", "consumer journey",
    "integrated campaigns", "influencer", "creators", "ecommerce", "omnichannel retailers", "pure-play retailers",
    "measurement framework", "data sources", "tracking", "reporting", "governance", "performance reviews",
    "post-campaign analysis", "insights", "actionable recommendations", "test-and-learn", "control groups",
    "incremental sales", "baseline", "creative testing", "A/B testing", "Meta Ads", "Google Ads",
    "Google Analytics", "paid social", "paid search", "EDM", "Shopify", "algorithm changes", "best practices",
    "analytical skills", "Power BI", "Looker", "Excel", "beauty", "skincare", "FMCG", "talabat", "Noon",
    "Careem", "Alibaba", "Miravia", "UAE", "Dubai",
]

CONTENT = {
    "headline": (
        "Performance & Precision Marketing · Full-Funnel Digital Campaigns · Measurement & Test-and-Learn · "
        "Agency Management"
    ),
    "professional_summary": (
        "Brand-side marketer with 4+ years across FMCG, beauty and marketplace e-commerce. Plans and optimises "
        "full-funnel digital campaigns for a UAE FMCG group (Meta and Google Ads, creators, EDM, Shopify, talabat, "
        "Noon, Careem), measured on incremental sales against control groups. Challenges agencies on their numbers. "
        "Previously managed 42 beauty brands inside Alibaba's Miravia marketplace."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Homegrown UAE F&B / FMCG group | Befit, Eurocake, SMASH, Flair | talabat, Noon, Careem + Shopify",
            "bullets": [
                "Plan and run full-funnel digital campaigns for four brands: audience strategy, channel mix (Meta, Google Ads, creators, EDM, Shopify), KPI definition and in-flight optimisation on ROI and ROAS",
                "Measure campaigns on incremental sales with control groups and baselines: SMASH × talabat lifted daily sales +165% vs +4.7% for the control brand (~AED 32K incremental on ~AED 8.7K spend); Befit × Noon +31% over baseline (+4,176 units)",
                "Manage and challenge creator agencies: negotiated fees -30% (AED 81 per content piece) and audited their reports, finding modelled reach and double-counted accounts, and that micro-creators (<20K) beat the campaign average 10–40×",
                "Integrate creator, e-commerce and platform activity across the consumer journey (talabat samplings, an in-app sweepstakes, creators linking to purchase); built AI (Claude) planning and reporting tools (~40% less manual work); lead a team of two",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's pure-play marketplace | beauty, fashion & lifestyle",
            "bullets": [
                "Managed 42 beauty, fragrance and fashion brands (incl. KIKO Milano) on onsite visibility, promotions and pricing, reading traffic, conversion, ROAS and retention weekly to deliver +30% GMV QoQ",
                "Owned the Flash Sales channel for Beauty, Fashion & Home reporting to the CEO, reviewing each campaign against P&L targets and feeding the learnings into the next one",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue",
            "bullets": [
                "Helped build the Retail vertical (fashion, beauty and lifestyle brands) and grew XL accounts (KFC, Taco Bell) through data-led in-app promotions and joint marketing activations",
            ],
        },
        {
            "company": "Mondelez International · Massimo Dutti (Inditex)",
            "role": "Trainee, Category Planning · Sales Associate",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG, Nielsen | Premium fashion retail, Las Rozas Village store",
            "bullets": [
                "At Mondelez, Nielsen sell-in/sell-out and promotional-effectiveness analysis for the chocolate category; at Massimo Dutti, shop-floor sales and Inditex retail standards",
            ],
        },
    ],
    "skills_brand": (  # -> "Precision Marketing"
        "full-funnel planning, audience strategy, channel selection, KPI definition, in-flight optimisation"
    ),
    "skills_ecommerce": (  # -> "Channels"
        "Meta Ads, Google Ads (Search & Display), creators & UGC, EDM, Shopify, talabat, Noon, Careem"
    ),
    "skills_commercial": (  # -> "Agencies & Partners"
        "agency briefing & challenge, fee negotiation, report audits, e-commerce & platform partners"
    ),
    "skills_data": (  # -> "Measurement"
        "control groups, baselines, incremental sales, ROI & ROAS, post-campaign analysis, test-and-learn"
    ),
    "skills_tools": (  # -> "Tools"
        "Meta Ads Manager, Google Ads, Google Analytics, Power BI, Looker, Excel, Shopify, Claude (AI)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Precision Marketing",
    "E-Commerce & Digital": "Channels",
    "Commercial": "Agencies & Partners",
    "Data & Analytics": "Measurement",
}


def make_job() -> Job:
    return Job(
        id=JOB_ID,
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE (assumed — not stated in the posting)",
        url="https://www.linkedin.com/jobs/search/?keywords=Beiersdorf%20Precision%20Marketing",
        source="linkedin",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Beiersdorf Precision Marketing",
            "note": "Contrato de 7 meses. El JD no da título ni ciudad. Contacto: Kartik Kulshrestha (Talent "
                    "Acquisition Lead). Pide performance marketing lado marca como último puesto, idealmente "
                    "agencia de medios antes, y ad tech (GMP, DSP, DMP). CV del 2026-10-04.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    existing = next((j for j in jobs if j.get("id") == job.id), None)
    entry = existing if existing is not None else {}
    entry.update({
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": entry.get("posted_date", DATE_FOLDER), "raw": job.raw,
        "ai_score": 60, "ai_tier": "Warm",
        "skills_match": [
            "Lado marca hoy: Meta y Google Ads, creators, EDM y Shopify para cuatro marcas, con KPIs de ROI y ROAS",
            "Test-and-learn con grupo de control: SMASH × talabat +165% frente a +4,7% del control",
            "Post-campaign con incrementales: Befit × Noon +31% sobre baseline, +4.176 unidades",
            "Retar a la agencia: auditó su informe (reach estimado, doble conteo) y negoció -30%",
            "Influencer + ecommerce + plataformas integrados: sorteo y samplings con talabat",
            "Pure-play desde dentro: 42 marcas de beauty en Miravia (Alibaba), +30% GMV QoQ",
            "Beauty, la categoría de Beiersdorf; ya trabajó con KIKO Milano",
        ],
        "missing_skills": [
            "Agencia de medios: no tiene, y el JD la pone como ideal antes del lado marca",
            "Ad tech (GMP, DSP, DMP, DV360, programática): no lo ha usado",
            "Atribución: mide con grupos de control y baselines, no ha montado modelos de atribución",
            "Título de performance: su puesto es Brand & Marketing Manager, el paid media va dentro",
            "Escala de medios: presupuestos de pyme, no de multinacional",
        ],
        "sector_fit": "bueno — beauty y FMCG; Beiersdorf es skincare de gran consumo",
        "seniority_fit": "en línea en años; corta en profundidad de medios frente a un perfil de agencia",
        "red_flags": [
            "Contrato de 7 meses: dejaría un puesto fijo en DoFreeze por uno temporal; confirmar renovación o paso a fijo",
            "Visado: en contratos temporales a veces sponsoriza una agencia de payroll, no Beiersdorf; preguntar quién",
            "El JD no dice ciudad: si no es EAU, queda fuera del filtro de ubicación",
            "Salario no publicado: confirmar que llega al suelo de 20.000 AED/mes",
            "Ya hay otras tres candidaturas a Beiersdorf en el dashboard (Global Performance, KAM Modern Trade, Eucerin ABM)",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Encaje medio. Beiersdorf busca a alguien de performance marketing en el lado marca que lleve campañas "
            "full-funnel con la agencia de medios, la rete y monte la medición. Paula hace eso a escala pequeña: "
            "lleva Meta y Google Ads, creators y ecommerce para cuatro marcas, mide en venta incremental con grupo "
            "de control (SMASH × talabat +165% frente a +4,7%) y ha auditado y renegociado a sus agencias. Suma el "
            "lado pure-play (Miravia) y beauty. Le faltan dos cosas que el JD nombra: no viene de agencia de medios "
            "y no ha tocado ad tech (GMP, DSP, DMP). Es un contrato de 7 meses: hay que pesar si compensa dejar "
            "DoFreeze, aunque puede servir de entrada a Beiersdorf, donde ya hay otras candidaturas abiertas."
        ),
        "scored_by": "manual:claude", "freshness": entry.get("freshness", "fresh"),
        "discovered_at": entry.get("discovered_at", f"{DATE_FOLDER}T00:00:00+00:00"),
        "status": entry.get("status", "Reviewing"),
    })
    if existing is None:
        jobs.insert(0, entry)
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
