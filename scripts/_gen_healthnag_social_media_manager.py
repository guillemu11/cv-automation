"""One-off: CV for "Social Media Manager" at Health Nag: Dubai supplement and functional-wellness brand
(founder: naturopath Natasha Rudatsenko; collagen jelly, liposomal vitamins, detox, matcha; in Waitrose;
@health_nag ~106K IG followers). AED 30,000/month, Dubai, visa + health insurance included.
Applied through a Google Form (CV + 3-min intro video + open questions); no cover letter.
Form answers live in output/2026-10-12/Health Nag - Social Media Manager/04_Aplicacion/.

The JD: turn their channels into an "editorial wellness destination". Storytelling, community, personality,
a point of view; own social while working with the team, creators and agencies. They screen hard for voice
("Nag us" one-liner, a wellness trend take) and on-camera personality.

Paula's honest angle:
- Befit is DoFreeze's fitness / better-for-you brand: the closest wellness adjacency she has.
- Leads social and content with a team of two (designer + social media executive).
- Built the creator programme from zero: seasonal moments (Fitness Month, New Year New Me, Ramadan, back to
  school), two SMASH launch campaigns, 103 creator activations across 3 campaigns, AED 81 per content piece.
- Measured on sales: SMASH x talabat +165% vs +4.7% control; Befit x Noon +4,176 incremental units (+31%).
  Micro-creators (<20K) beat the campaign's average engagement 10-40x.
- Community beyond influencers: running, padel and yoga clubs; talabat purchase-to-win giveaway; affiliate
  programme with an AI-agent-built landing page and a multilingual AI avatar briefing creators.
- Agencies: ran selection across four, negotiated 30% off, audited their reports.

Honesty guardrails:
- No follower-growth or organic engagement numbers exist for the brand pages: none claimed.
- Hands-on reel shooting/editing and on-camera presence are not claimed on the CV (Paula answers in the form).
- Wellness / supplements category is new to her: framed via Befit (fitness), never "wellness brand experience".
- DoFreeze is NOT on Deliveroo. Never "own the forecast". No Arabic (not asked here). Visa line from profile.yaml.

ONE-PAGE standard. Lands under output/2026-10-12/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Health Nag"
TITLE = "Social Media Manager"
DATE_FOLDER = "2026-10-12"

JOB_DESCRIPTION = """\
Social Media Manager — Health Nag. Dubai, UAE. AED 30,000/month. Visa and health insurance included.
One voice. A whole community to NAG. Health Nag is an opinionated wellness brand. It nags, it educates, it
innovates — and now we're growing. We're looking for a Social Media Manager to turn our channels into an
editorial wellness destination people actually want to follow. Think strong storytelling, community, personality
and content with a point of view. You'll own our social presence while working with our team, creators and
agencies to make Health Nag unmistakable online. We're not just looking for someone who loves wellness. We're
looking for personality, taste and an inner Health Nag.
Application form asks for: links to work, 3 pieces of content that outperformed expectations, personal/brand
pages managed, 2-3 sample posts/reels, platforms used day-to-day, whether the candidate can shoot and edit
short-form video (CapCut, Premiere, in-app editors), a take on a recent wellness trend, comfort on camera / on
voice as a brand personality, "AI content that looks and feels human" skills, a 1-3 minute intro video.
"""

ATS = [
    "Social Media Manager", "social media", "social media strategy", "editorial", "content strategy",
    "storytelling", "brand voice", "tone of voice", "point of view", "community", "community building",
    "content calendar", "Instagram", "TikTok", "reels", "short-form video", "UGC", "creators",
    "influencer marketing", "creator partnerships", "agency management", "wellness", "fitness", "health",
    "better-for-you", "seasonal campaigns", "product launches", "sampling", "affiliates", "AI content",
    "generative AI", "AI video", "copywriting", "engagement", "measurement", "ROI", "Meta Ads", "paid social",
    "team leadership", "Dubai", "UAE", "GCC",
]

CV_CONTENT = {
    "headline": "Social Media, Content & Creators · Hands-on AI Content",
    "professional_summary": (
        "Brand and social marketer in Dubai with 5 years across FMCG, marketplaces and quick-commerce (Mondelez, "
        "Alibaba, Glovo). I lead social and content for a UAE food group, including its fitness brand Befit: built "
        "the creator and community programme from zero, measured on sales, with AI tools I build myself."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE F&B / FMCG group | Befit (fitness), Eurocake, SMASH, Flair | sold in 50+ markets",
            "bullets": [
                "Lead social and content with a team of two (designer + social media executive): set the content calendar around seasonal moments (Fitness Month, New Year New Me, Ramadan, back to school) and brief and review every post, reel and creator asset",
                "Built the creator programme from zero: 103 creator activations across 3 campaigns plus two SMASH launches, at AED 81 per content piece; micro-creators (<20K) beat the campaign's average engagement 10–40×",
                "Measure content on what it sells: SMASH × talabat lifted daily sales +165% vs +4.7% for a no-creator control brand; Befit × Noon 'New Year, New Me' added 4,176 incremental units (+31% over baseline)",
                "Built community beyond influencers: running, padel and yoga clubs, a talabat purchase-to-win giveaway, seasonal sampling, and Befit's affiliate programme, launched on an AI-agent-built landing page with a multilingual AI avatar briefing creators",
                "Ran selection across four creator agencies, negotiated 30% off (103 creators delivered vs 75 contracted) and audited their reports; built a Claude system for content, planning and reporting (~40% less manual work)",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees",
            "bullets": [
                "Created Hot on Social and Beauty Club: social-first content and community projects built on trending formats and reviews, positioning Miravia as a beauty destination",
                "Grew 42 beauty, fragrance and fashion accounts (incl. KIKO Milano and NIVEA) +30% GMV QoQ; owned the Flash Sales channel, reporting to the CEO",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Managed XL accounts (KFC, Taco Bell, La Tagliatella) on GMV and joint marketing activations; helped build the Retail vertical, onboarding new brand partners",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | chocolate category (Milka, Suchard)",
            "bullets": [
                "Sell-in/sell-out and promo-effectiveness analysis; supported NPD launches (Milka Spread, Mini Suchard)",
            ],
        },
    ],
    "skills_brand": (  # -> "Social & Content"
        "content calendars, brand voice, Instagram & TikTok, reels & UGC briefs, EN/ES copywriting"
    ),
    "skills_ecommerce": (  # -> "Creators & Community"
        "creator programmes from zero, agency selection & audits, affiliates, sampling, run/padel/yoga clubs"
    ),
    "skills_commercial": (  # -> "AI Content"
        "Claude / Claude Code, ChatGPT, AI video (Higgsfield, Hailuo AI), AI avatars, Canva AI"
    ),
    "skills_data": (  # -> "Growth & Data"
        "Meta & Google Ads, control groups, incremental sales, cost per content, ROI & ROAS"
    ),
    "skills_tools": (
        "Canva & Adobe (Photoshop, Illustrator), Meta Business Suite, Shopify, Notion, Excel & PowerPoint"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Social & Content",
    "E-Commerce & Digital": "Creators & Community",
    "Commercial": "AI Content",
    "Data & Analytics": "Growth & Data",
    "Tools": "Tools",
}


def make_job() -> Job:
    return Job(
        id="healthnag-social-media-manager-2026-10",
        title=TITLE,
        company=COMPANY,
        location="Dubai, United Arab Emirates",
        url="https://healthnag.com",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={"via": "Google Form application (CV + 3-min intro video + open questions)",
             "salary": "AED 30,000/month, visa and health insurance included",
             "brand": "Supplements / functional wellness; @health_nag ~106K IG followers"},
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "AED 30,000/month + visa + health insurance",
        "salary_aed_min": 30000, "salary_aed_max": 30000,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 76, "ai_tier": "Warm",
        "skills_match": [
            "Leads social and content at DoFreeze with a designer + social media executive",
            "Befit is a fitness / better-for-you brand: seasonal wellness moments (Fitness Month, New Year New Me)",
            "Creator programme from zero: 103 activations, AED 81 per piece, sales lift vs control",
            "Community-led: running, padel and yoga clubs; affiliates; agencies selected and audited",
            "AI content hands-on (Claude, Higgsfield, Hailuo, AI avatar)",
        ],
        "missing_skills": [
            "No organic follower-growth / engagement numbers for the pages she runs",
            "Hands-on reel shooting and editing not evidenced; on-camera presence unknown",
            "Supplements / functional-wellness category is new",
        ],
        "sector_fit": "adjacent (fitness / better-for-you food → supplements & wellness)",
        "seniority_fit": "match (Manager, owns social with team, creators and agencies)",
        "red_flags": [
            "Form asks date, time and place of birth (unusual; personal data)",
            "Heavy screen on personality and on-camera presence (3-min video required)",
        ],
        "ats_keywords": ATS,
        "reasoning": "Social + creators + community on a fitness brand is close to her day job and the salary clears her floor comfortably; the screen is about voice and on-camera personality, which the video and the open answers must carry.",
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _pages(pdf: Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(pdf)).pages)


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    final_dir = settings.output_dir / DATE_FOLDER / f"{COMPANY} - {TITLE}"
    dest = final_dir / "01_CV_y_Carta"
    dest.mkdir(parents=True, exist_ok=True)

    cv_docx = cv._fill_template(CV_CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)

    moved = dest / cv_pdf.name
    shutil.move(str(cv_pdf), str(moved))
    shutil.copy(str(moved), str(dest / "Paula De Francisco - CV.pdf"))
    print("OK", moved.name, "PAGES", _pages(moved))

    stray = cv_pdf.parent.parent
    if stray.exists() and stray.resolve() != final_dir.resolve():
        shutil.rmtree(stray, ignore_errors=True)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
