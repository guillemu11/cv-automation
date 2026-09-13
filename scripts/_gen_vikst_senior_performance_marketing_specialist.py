"""One-off: generate Paula's ONE-PAGE CV for **Senior Performance Marketing Specialist** at
**VIKST** (UAE performance-marketing agency for DTC eCommerce brands; full-time, remote).

JD ask: lead and scale Meta Ads (Instagram & Facebook) — "spent millions, not thousands", Meta
mastery non-negotiable; end-to-end customer journeys (ad creative -> landing page -> LTV); audit,
optimise and scale on data; work with creative specialists and developers; client strategy calls
advising founders/CMOs; proactive insights; Snapchat/TikTok/Google a bonus.

Paula's honest angle: in-house DTC owner — Shopify web end-to-end (CRO, AOV) fed by Meta + Google
paid media with creative A/B tests; landing pages built with AI agents; works with a designer on
creative; audited agency reporting (bottleneck-finding); control-group sell-out measurement;
client-facing advisory from Glovo/Miravia KAM (brands, CEO reporting).

Honesty guardrails:
- NO spend figures, NO "millions" — her Meta budgets are in-house FMCG scale (the likely blocker).
- LTV not claimed. TikTok = creator content, not TikTok ads. No Snapchat.
- Forecast not owned. Factual "UAE Residence Visa" only.

ONE-PAGE standard. Lands under output/2026-09-13/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "VIKST"
TITLE = "Senior Performance Marketing Specialist"
DATE_FOLDER = "2026-09-13"

JOB_DESCRIPTION = """\
Senior Performance Marketing Specialist — VIKST (performance marketing agency), full-time, remote,
UAE & GCC clients. Someone who lives and breathes Meta Ads, thrives in eCommerce and is obsessed with
scaling DTC brands; a growth partner who moves the revenue needle and guides clients with confidence.

Work: lead and scale Meta Ads campaigns (Instagram & Facebook); strategise end-to-end customer journeys
from ad creative to landing pages and LTV; audit, optimise and scale campaigns based on data; work with
creative specialists and developers to align messaging and performance; join client strategy calls;
offer proactive insights, not just reports; explore scaling on Snapchat, TikTok, Google if relevant.

Profile: serious Meta Ads experience (spent millions, not thousands); proven eCommerce performance
marketing track record; comfortable advising founders and CMOs; break down data, isolate bottlenecks,
find scalable wins; clear, confident communication; fast-moving teams, direct feedback, accountability.
Bonus: Snapchat, TikTok or Google Ads. Meta Ads mastery is non-negotiable. Start ASAP.
"""

ATS = [
    "performance marketing", "Meta Ads", "Facebook Ads", "Instagram Ads", "eCommerce", "DTC",
    "scaling", "customer journey", "ad creative", "landing pages", "conversion rate optimisation",
    "CRO", "AOV", "audit", "optimise", "ROAS", "ROI", "A/B testing", "creative testing",
    "Shopify", "Google Ads", "TikTok", "client strategy", "founders", "CMOs", "insights",
    "bottlenecks", "UAE", "GCC",
]

CONTENT = {
    "headline": "Performance Marketing · Meta Ads · DTC eCommerce & Shopify",
    "professional_summary": (
        "Performance marketer with 5 years in eCommerce (Alibaba, Glovo) and FMCG, now running Meta and Google Ads into a "
        "Shopify DTC store in Dubai, owning the journey from ad creative to landing page and checkout."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG snacking group (Befit, Eurocake, Smash) | DTC Shopify + quick-commerce | GCC",
            "bullets": [
                "Run Meta Ads (Facebook & Instagram) and Google Ads for the brands: audience building, creative A/B tests and optimisation against ROAS and ROI",
                "Own the Shopify DTC store end-to-end (catalogue, collections, discounts, checkout), lifting conversion rate and AOV; built the affiliate-programme landing page with AI agents",
                "Work with a designer and social media executive to match ad creative to performance; 103 creator activations at AED 81 per content piece, with micro-creators beating the engagement average 10-40x",
                "Measure incrementality, not vanity metrics: SMASH x talabat +165% daily sales vs +4.7% for a control brand; audited agency reports and caught reach estimated from followers",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's eCommerce marketplace | 100K+ employees",
            "bullets": [
                "Advised 42 brand accounts on promotions, pricing and assortment for +30% GMV QoQ, analysing ROAS, conversion, traffic and retention",
                "Owned the Flash Sales channel, reporting performance and recommendations directly to the CEO",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Delivery & quick-commerce tech leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Managed XL partners (KFC, Taco Bell, Sushi Shop), growing orders with data-led promotional activations",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | Chocolate category (Milka, Suchard) | €36B revenue",
            "bullets": [
                "Built promotional-effectiveness and sell-in/sell-out reports with Nielsen for the chocolate category",
            ],
        },
    ],
    "skills_brand": (  # label -> "Performance"
        "Meta Ads (Facebook & Instagram), Google Ads, creative A/B testing, audiences, creators & UGC"
    ),
    "skills_ecommerce": (  # label -> "eCommerce"
        "Shopify DTC, landing pages, conversion rate (CRO), AOV, checkout, talabat / Noon / Careem"
    ),
    "skills_commercial": (  # label -> "Client & Team"
        "brand & partner advisory, CEO reporting, agency audits, designer & social team, negotiation"
    ),
    "skills_data": (  # label -> "Analytics"
        "ROAS / ROI, funnel & conversion analysis, control groups, KPI dashboards, Looker, Power BI"
    ),
    "skills_tools": (  # label -> "Tools"
        "Meta Ads Manager, Meta Business Suite, Google Ads, Shopify, Excel, Canva, Claude (AI agents)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Performance",
    "E-Commerce & Digital": "eCommerce",
    "Commercial": "Client & Team",
    "Data & Analytics": "Analytics",
}


def make_job() -> Job:
    return Job(
        id="vikst-senior-performance-marketing-specialist-2026-09",
        title=TITLE,
        company=COMPANY,
        location="UAE (Remote)",
        url="https://vikst.com/",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "VIKST Senior Performance Marketing Specialist",
            "note": "Moderate fit: in-house Meta + Google Ads into Shopify DTC, CRO/AOV, creative testing. "
                    "Blocker: 'spent millions' Meta mastery non-negotiable — her budgets are in-house FMCG scale.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed (competitive, performance-based)",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 58, "ai_tier": "Warm",
        "skills_match": [
            "Meta Ads (FB & IG) + Google Ads hands-on with creative A/B testing and ROAS optimisation",
            "Owns a Shopify DTC store end-to-end (CRO, AOV, checkout) + AI-built landing pages",
            "Incrementality mindset: control-group sell-out measurement, agency report audit",
            "Client-facing advisory from Glovo / Miravia KAM (42 brands, CEO reporting)",
        ],
        "missing_skills": [
            "Meta spend at millions scale (non-negotiable in JD)",
            "Specialist agency performance-marketing tenure across multiple DTC clients",
            "LTV modelling, Snapchat and TikTok Ads",
        ],
        "sector_fit": "adjacent — in-house FMCG/eCommerce vs agency DTC performance",
        "seniority_fit": "partial — generalist brand manager, not a senior Meta specialist",
        "red_flags": ["'Spent millions, not thousands' on Meta is non-negotiable", "Remote agency role"],
        "ats_keywords": ATS,
        "reasoning": (
            "Paula has real Meta/Google Ads + Shopify DTC ownership and a strong measurement mindset, but the role "
            "wants a senior Meta specialist with millions in spend across agency DTC clients. Likely screened out on "
            "spend scale; worth a shot only as a low-effort application."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
