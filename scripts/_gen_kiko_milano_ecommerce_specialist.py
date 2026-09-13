"""One-off: generate Paula's ONE-PAGE CV for **E-Commerce Specialist** at **KIKO Milano Middle East**
(Dubai; reports to Head of E-Commerce; IMEA region).

JD ask: support e-commerce sales/growth across IMEA; coordinate HQ + local teams on launches, campaigns,
seasonal activations; digital merchandising, product content, online CX; KPIs (sales, conversion,
engagement, sell-out); support forecasting/demand planning/inventory/replenishment; product setup and
catalogue accuracy; work with Marketing, Supply Chain, Finance, HQ. Beauty preferred; Excel/PowerPoint;
fluent English AND Arabic.

Paula's honest angle: at Miravia she was the Key Account for KIKO Milano's makeup & skincare account
(confirmed by Paula 2026-09-13) — direct brand familiarity. Plus Miravia beauty catalogue/promotions for
42 brands, Beauty Club; DoFreeze Shopify + talabat/Noon/Careem listings, launches, dark-store stock review.

Honesty guardrails:
- NO Arabic (JD requires it) — hard gap, flagged; never implied.
- Forecast: contributes, doesn't own. No KIKO-specific metrics invented.
- Factual "UAE Residence Visa" only.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "KIKO Milano"
TITLE = "E-Commerce Specialist"
DATE_FOLDER = "2026-09-13"

JOB_DESCRIPTION = """\
E-Commerce Specialist — KIKO Milano Middle East, Dubai, UAE. Reporting to Head of E-Commerce.
Support e-commerce sales and growth across the IMEA region. Coordinate with HQ and local teams on e-commerce
activities, product launches, campaigns and seasonal activations. Manage and optimize digital merchandising,
product content and the online customer experience. Monitor key e-commerce KPIs, including sales, conversion,
customer engagement and sell-out, and identify opportunities for improvement. Support forecasting, demand
planning, inventory and replenishment to ensure strong product availability. Manage product setup and catalogue
accuracy across e-commerce platforms. Work closely with Marketing, Supply Chain, Finance, HQ and local market teams.
Looking for: experience in e-commerce, digital merchandising or related; beauty/cosmetics preferred; analytical and
commercial mindset with data and KPIs; online customer journeys, digital merchandising, digital marketing;
organisation, attention to detail, multiple teams and markets; Excel and PowerPoint; e-commerce, analytics or
content management tools a plus; fast-paced international environment. Fluent in English and Arabic.
"""

ATS = [
    "e-commerce", "digital merchandising", "product content", "online customer experience", "customer journey",
    "product setup", "catalogue accuracy", "product launches", "seasonal activations", "campaigns", "HQ",
    "local markets", "IMEA", "KPIs", "sales", "conversion", "customer engagement", "sell-out", "forecasting",
    "demand planning", "inventory", "replenishment", "product availability", "marketing", "supply chain",
    "finance", "beauty", "cosmetics", "makeup", "skincare", "Excel", "PowerPoint", "Shopify", "marketplaces",
    "analytics", "digital marketing",
]

CONTENT = {
    "headline": "E-Commerce · Digital Merchandising · Beauty & Marketplaces",
    "professional_summary": (
        "E-commerce professional with 5 years across Alibaba's Miravia (Key Account for KIKO Milano makeup & skincare), "
        "Glovo and FMCG, now running Shopify and quick-commerce channels in Dubai: catalogue, launches, KPIs and availability."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG group (Befit, Eurocake, Smash) | Shopify D2C, talabat, Noon, Careem | GCC & 50+ markets",
            "bullets": [
                "Own the Shopify store end-to-end: product setup, content, collections, discounts and checkout, lifting conversion rate and AOV through data-led merchandising",
                "Run launches and seasonal activations (6 NPD launches, Ramadan, New Year) across web, talabat, Noon and Careem: listings, imagery, promotions and catalogue accuracy",
                "Track sales, conversion and sell-out in KPI dashboards; run a weekly stock-risk review by dark store and contribute to the monthly forecast with Sales and Finance",
                "Coordinate marketing, supply chain and 4 agencies on campaigns across markets; lead a team of two (designer + social media executive)",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's eCommerce marketplace | 100K+ employees",
            "bullets": [
                "Key Account for KIKO Milano (makeup & skincare) among 42 beauty, fragrance and fashion brands: catalogue, assortment, pricing and promotions (+30% GMV QoQ)",
                "Owned the Flash Sales channel reporting to the CEO; created the Beauty Club project; onboarded 30+ fragrance stores in two months",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Delivery & quick-commerce app | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Managed XL partners (KFC, Taco Bell, Sushi Shop) on menus, promotions and GMV; helped build the Retail vertical incl. beauty",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | Chocolate category (Milka, Suchard) | €36B revenue",
            "bullets": [
                "Built sell-in/sell-out and promotional-effectiveness reports with Nielsen for the chocolate category",
            ],
        },
    ],
    "skills_brand": (  # label -> "E-Commerce"
        "product setup, catalogue accuracy, launches, seasonal activations, promotions"
    ),
    "skills_ecommerce": (  # label -> "Merchandising & CX"
        "digital merchandising, product content, customer journey, CRO, Shopify, marketplaces"
    ),
    "skills_commercial": (  # label -> "Beauty & Markets"
        "KIKO Milano & beauty brands, HQ & local teams, marketing, supply chain, finance"
    ),
    "skills_data": (  # label -> "KPIs & Planning"
        "sales, conversion, sell-out, stock-risk & availability, forecast support"
    ),
    "skills_tools": (  # label -> "Tools"
        "Excel, PowerPoint, Shopify, Looker, Power BI, Meta & Google Ads, Canva, Claude (AI)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "E-Commerce",
    "E-Commerce & Digital": "Merchandising & CX",
    "Commercial": "Beauty & Markets",
    "Data & Analytics": "KPIs & Planning",
}


def make_job() -> Job:
    return Job(
        id="kiko-milano-me-ecommerce-specialist-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE",
        url="https://www.linkedin.com/jobs/search/?keywords=KIKO%20Milano%20E-Commerce%20Specialist%20Dubai",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "KIKO Milano E-Commerce Specialist Dubai",
            "note": "Strong function/brand fit: Paula was KIKO Milano's Key Account (makeup & skincare) at Miravia. "
                    "Hard gap: fluent Arabic required. Specialist level may pay below AED 20k floor.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 66, "ai_tier": "Hot",
        "skills_match": [
            "Was KIKO Milano's Key Account (makeup & skincare) at Miravia — knows the brand and catalogue",
            "Beauty marketplace ops for 42 brands: catalogue, assortment, pricing, promotions (+30% GMV QoQ)",
            "Owns Shopify store end-to-end: product setup, content, merchandising, CRO/AOV",
            "Launches and seasonal activations across web + talabat/Noon/Careem; catalogue accuracy",
            "KPI tracking, dark-store stock-risk review, contributes to forecast",
        ],
        "missing_skills": [
            "Fluent Arabic (required in JD) — Paula has none",
            "Formal demand planning / replenishment ownership",
            "IMEA multi-country brand-side e-commerce role",
        ],
        "sector_fit": "strong — beauty (KIKO account at Miravia) + e-commerce",
        "seniority_fit": "step down — Specialist vs her Manager title",
        "red_flags": ["Fluent Arabic required", "Specialist level (salary may sit below AED 20k floor)"],
        "ats_keywords": ATS,
        "reasoning": (
            "Excellent brand and function fit: she managed KIKO Milano's makeup & skincare account at Miravia and runs "
            "Shopify/quick-commerce catalogue, launches and KPIs in Dubai. Arabic is required and she has none — "
            "apply leaning on the KIKO relationship; check salary early."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
