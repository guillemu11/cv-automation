"""One-off: CV de una página **SIN FOTO** para **Ecommerce Shopper and Customer Marketing Manager -
Middle East** en **Beiersdorf Middle East FZCO** (NIVEA, Eucerin, Labello, Hansaplast...). Dubái.
Reclutador: Akash Sharma. Plazo: 19-oct-2026 vía Beiersdorf Career Page. Quinta candidatura a
Beiersdorf (antes: Eucerin ABM, KAM Modern Trade, Global Perf. Marketing & Retail Media, Precision
Marketing), todas en "Reviewing".

**Sin foto**: la oferta pide expresamente subir el CV sin foto para evitar sesgos. Se elimina la
imagen de la cabecera y se fusionan las dos celdas para que el bloque de nombre ocupe todo el ancho.

Qué es el puesto: especialista del equipo de e-commerce de Oriente Medio que traduce estrategias de
marca, categoría y cliente en experiencias de compra online y planes por cliente. Dueño del Digital
Shelf (contenido, disponibilidad, ratings y reviews, búsqueda en el retailer, conversión; Profitero) y
de Salsify (sindicación de contenido). Planes de activación por cliente alineados con los JBP, sell-in
con Sales y KAM, surtido, posicionamiento y promociones, RGM. Promociones, lanzamientos, bundles y
grandes eventos de compra online; pilotos de Live y Social Commerce. Cuota de mercado online, KPIs y
dashboards. Presupuestos, agencias y coordinación multifuncional. Pide 3-4 años en shopper, customer,
trade marketing, category management o e-commerce en gran consumo. Inglés; árabe solo ventaja.

Ángulo honesto de Paula (encaje bueno: 5 años frente a 3-4):
- **Lado marca en e-retailers y quick commerce de EAU**: talabat, Noon y Careem (listings, precio,
  mecánicas promocionales, activaciones conjuntas). SMASH × talabat +165% venta diaria vs +4,7% del
  control; Befit × Noon "New Year, New Me" +4.176 uds incrementales, +31% sobre baseline.
- **Ya conoce Beiersdorf**: en Miravia llevó la cuenta de **NIVEA durante 3 meses** (dato de Guille,
  2026-10-06). Sin métricas propias de NIVEA; el +30% GMV QoQ es de toda la cartera.
- **Planes de cliente desde el otro lado**: Miravia (Alibaba), 42 cuentas de beauty y fragancias
  (incl. KIKO Milano) con planes de visibilidad, promociones, precio y surtido; +30% GMV QoQ; canal
  Flash Sales reportando al CEO.
- **Disponibilidad**: revisión semanal de riesgo de stock por dark store con dashboard hecho con IA.
- **Contenido y conversión**: Shopify de punta a punta (fichas, colecciones, descuentos, checkout).
- **Social commerce real**: programa de creators de cero a 25-50 por campaña, sampling y seeding,
  medido en sell-out. 4 agencias de influencers (negoció -30%). Equipo de dos.
- **Análisis**: Mondelez, sell-in/sell-out y eficacia promocional en chocolate.

Guardarraíles de honestidad:
- **Salsify, Profitero, PIM**: cero experiencia. No se mencionan como herramienta usada.
- **GEO / AI search / keyword strategy en retailer**: sin evidencia. Se habla de flujos con IA
  (Claude) para contenido y reporting, que sí tiene; no de optimización para buscadores generativos.
- **Ratings, reviews, share of search, cuota de mercado online**: sin evidencia como KPI gestionado.
- **Retail media de pago** con presupuesto propio: sin confirmar → "platform campaigns".
- **Live commerce / TikTok Shop**: no. Social commerce = creators + sampling + amplificación.
- **RGM formal**: no. Precio y promociones sí.
- Árabe no (ventaja aquí). Deliveroo: DoFreeze NO está. Forecast: "contribute".
- Visa: "UAE Employment Visa (employer-sponsored)".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Beiersdorf"
TITLE = "Ecommerce Shopper and Customer Marketing Manager - Middle East"
DATE_FOLDER = "2026-10-06"

JOB_DESCRIPTION = """\
Ecommerce Shopper and Customer Marketing Manager - Middle East. Beiersdorf Middle East FZCO. Dubai, UAE.
Unlimited / Full-Time. Job function: Sales. Recruiter: Akash Sharma. Apply via Beiersdorf Career Page until 19 October 2026.
Role purpose: translates brand, category and customer strategies into winning online shopper experiences and customer
plans. Owns Digital Shelf excellence and works across eCommerce, Sales, Marketing, Category and external partners to
improve search, conversion and profitable growth. Specialist role delivering through cross-functional influence, part of
the Middle East E-commerce Team.
1) Digital Shelf and Content Excellence: own Digital Shelf standards and execution across e-retailers, marketplaces,
omnichannel customers and quick commerce (content quality, availability, ratings and reviews, retailer search visibility,
conversion; tracking via Profitero); lead Salsify and content syndication; improve search, AI optimization and
generative-engine optimization through keyword strategy, structured product data and content; governance and
performance routines.
2) Shopper and Customer Planning: customer-specific strategies and annual activation plans aligned with joint business
plans; customer stories and sell-in materials; partner with Sales and Key Account teams to secure visibility, activation
support and customer investment; support assortment, placement, content and promotional choices; drive profitable growth
(Revenue Growth Management).
3) Online activation and Future Ecommerce: lead online promotions, launches, bundles and key shopping events from concept
through execution and post-evaluation; innovation and portfolio priorities translated into online activation, content
readiness and availability; build Live Commerce and Social Commerce capabilities through pilots, formats, playbooks;
test-and-learn.
4) Market Share, Insights and Performance: track online market share, category trends, competitor activity and customer
performance; define customer- and platform-specific actions; KPIs and dashboards across Digital Shelf, shopper activation,
media and sales; turn data into action.
5) Delivery and cross-functional collaboration: manage budgets, agencies, timelines and deliverables; work across
eCommerce, Sales, Marketing, Category, Media, Supply Chain, Finance and Legal; build tools, standards and capabilities.
Profile: university degree in Marketing, Business, Commerce or related; 3-4 years in shopper marketing, revenue growth
management, customer marketing, trade marketing, category management or eCommerce within consumer goods or retail;
digital shelf expertise (content, retailer search, availability, ratings and reviews, retail media, performance
measurement); hands-on Salsify or comparable PIM / content-syndication platform; working knowledge of AI-led search,
structured product data and GEO; commercial and analytical capability; budget, project and agency management;
fluent business English, Arabic or another regional language an advantage.
Note: Beiersdorf encourages candidates to upload the CV without a picture.
"""

ATS = [
    "eCommerce", "e-commerce", "shopper marketing", "customer marketing", "trade marketing", "category management",
    "digital shelf", "content quality", "availability", "conversion", "retailer search", "e-retailers", "marketplaces",
    "quick commerce", "omnichannel", "talabat", "Noon", "Careem", "Miravia", "Glovo", "customer plans",
    "joint business plans", "activation plans", "sell-in", "sell-out", "key accounts", "assortment", "placement",
    "promotions", "pricing", "online promotions", "launches", "key shopping events", "post-evaluation",
    "social commerce", "creators", "sampling", "test-and-learn", "market share", "category trends", "KPIs",
    "dashboards", "budgets", "agencies", "cross-functional", "supply chain", "finance", "FMCG", "consumer goods",
    "beauty", "skincare", "KIKO Milano", "Shopify", "AI", "Claude", "Excel", "Power BI", "Middle East", "Dubai",
]

CONTENT = {
    "headline": "eCommerce Shopper & Customer Marketing · Digital Shelf · Quick Commerce · FMCG & Beauty · UAE",
    "professional_summary": (
        "Shopper and customer marketer with 5 years across FMCG, beauty and marketplaces. In Dubai, activates FMCG "
        "brands on talabat, Noon and Careem with campaigns measured on sell-out; before, built visibility and promotion "
        "plans for 42 beauty and fragrance accounts at Miravia (Alibaba), including NIVEA and KIKO Milano, at +30% GMV QoQ."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG group (Befit, Eurocake, Flair) | talabat, Noon & Careem + Shopify D2C | 50+ markets",
            "bullets": [
                "Build the online activation plan with talabat, Noon and Careem — listings, pricing, promotional mechanics and joint activations (purchase-linked talabat cashback giveaway, seasonal sampling); Befit × Noon \"New Year, New Me\" added +4,176 incremental units, +31% over baseline",
                "Launched SMASH, a new Eurocake brand, on talabat with a control-group read: daily sales +165% vs +4.7% for a brand left out of the campaign; 6 product launches end-to-end (brief, packaging, pricing, go-to-market)",
                "Run the weekly stock-risk review by dark store with a dashboard built with AI, checking PO coverage against demand to protect availability; own the Shopify store end-to-end — product content, collections, discounts, checkout — merchandised on conversion",
                "Built the creator programme from zero to 25–50 creators per campaign with sampling and seeding, measured on sell-out; manage 4 influencer agencies (negotiated -30%) and lead a team of two",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees | retailer side",
            "bullets": [
                "Planned visibility, campaign slots, promotions, pricing and assortment with 42 beauty, fragrance and fashion brands (incl. NIVEA for three months and KIKO Milano) — the retailer side of a joint business plan — delivering +30% GMV QoQ",
                "Owned the Flash Sales channel for Beauty, Fashion & Home reporting to the CEO, reading traffic, conversion, ROI and ROAS to move promotional investment; created the Beauty Club and Hot on Social projects",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce & delivery platform | €500M+ revenue | platform side",
            "bullets": [
                "Managed XL accounts (KFC, Taco Bell, La Tagliatella, Sushi Shop) with data-led growth plans and in-app activations, aligning marketing, logistics and customer support on execution",
            ],
        },
        {
            "company": "Mondelez International · Massimo Dutti (Inditex)",
            "role": "Trainee, Category Planning · Sales Associate",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG (chocolate category) | Premium fashion retail",
            "bullets": [
                "At Mondelez, sell-in/sell-out and promotional-effectiveness analysis for chocolate, supporting NPD launches (Milka Spread, Mini Suchard); at Inditex, store operations and customer service",
            ],
        },
    ],
    "skills_brand": (  # -> "Shopper & Customer"
        "customer activation plans, JBP input, promotions & key events, launches, sampling, creators"
    ),
    "skills_ecommerce": (  # -> "Digital Shelf"
        "product content, assortment, availability & stock risk, conversion, Shopify · talabat, Noon, Careem"
    ),
    "skills_commercial": (  # -> "Commercial"
        "pricing, promo mechanics, sell-in stories, budgets, agency management (-30%), forecast input"
    ),
    "skills_data": (  # -> "Insights & KPIs"
        "sell-in/sell-out, promo effectiveness, control-group reads, GMV, ROI, ROAS, dashboards"
    ),
    "skills_tools": (  # -> "Tools"
        "Excel (Advanced), Power BI, Tableau, Looker, Salesforce, SAP, Shopify, Claude (AI workflows)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Shopper & Customer",
    "E-Commerce & Digital": "Digital Shelf",
    "Commercial": "Commercial",
    "Data & Analytics": "Insights & KPIs",
}


def make_job() -> Job:
    return Job(
        id="beiersdorf-eshopper-customer-marketing-me-dubai-2026-10",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE",
        url="https://www.beiersdorf.com/careers",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={
            "query": "Beiersdorf Ecommerce Shopper and Customer Marketing Manager Middle East",
            "note": "Beiersdorf ME, equipo de e-commerce. Digital shelf (Profitero, Salsify), planes por cliente, "
                    "promociones y eventos online, social commerce. Pide 3-4 años. CV SIN FOTO (lo pide la oferta). "
                    "Reclutador Akash Sharma; plazo 19-oct-2026. Gaps: Salsify/Profitero, GEO, ratings/reviews.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": None, "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": "2026-10-05", "raw": job.raw,
        "ai_score": 80, "ai_tier": "Hot",
        "skills_match": [
            "Lado marca en e-retailers y quick commerce de EAU: talabat, Noon y Careem",
            "Activaciones online medidas en sell-out: SMASH × talabat +165% vs control, Befit × Noon +31%",
            "Ya conoce Beiersdorf: llevó la cuenta de NIVEA en Miravia durante 3 meses",
            "Planes por cliente desde el lado retailer: Miravia, 42 cuentas beauty/fragancias (KIKO), +30% GMV QoQ",
            "Disponibilidad: revisión semanal de riesgo de stock por dark store",
            "Contenido y conversión: Shopify de punta a punta",
            "Social commerce: programa de creators de 0 a 25-50 por campaña, sampling y seeding",
            "Presupuestos y agencias: 4 agencias de influencers (-30% negociado), equipo de dos",
            "FMCG (DoFreeze, Mondelez) y beauty (Miravia)",
            "3-4 años pedidos; tiene 5",
        ],
        "missing_skills": [
            "Salsify o PIM / sindicación de contenido (lo piden hands-on)",
            "Profitero u otra herramienta de digital shelf",
            "GEO, AI search y keyword strategy en retailer",
            "Ratings, reviews y share of search como KPI gestionado",
            "Cuota de mercado online y RGM formal",
            "Live commerce",
            "Árabe (ventaja, no requisito)",
        ],
        "sector_fit": "muy bueno — FMCG y beauty/skincare, e-commerce en Oriente Medio",
        "seniority_fit": "acorde — 5 años frente a 3-4 pedidos; rol especialista sin equipo a cargo",
        "red_flags": [
            "CV sin foto (lo pide la oferta) — usar la versión de este paquete",
            "Quinta candidatura a Beiersdorf con el mismo reclutador; mantener el discurso coherente",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Encaje fuerte: es justo lo que hace hoy (activar marcas FMCG en talabat, Noon y Careem con resultados en "
            "sell-out) más dos años en el lado retailer de los planes por cliente en Miravia con beauty. Los años "
            "encajan (5 frente a 3-4). El hueco real son las herramientas: Salsify, Profitero y GEO, que el JD pide "
            "hands-on; se compensa con contenido, disponibilidad y flujos con IA, sin inventarlas."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _remove_photo(docx_path: Path) -> None:
    """Drop the header photo and merge its cell into the name block (full width)."""
    doc = Document(str(docx_path))
    header = doc.tables[0]
    name_cell, photo_cell = header.rows[0].cells[0], header.rows[0].cells[1]
    for drawing in photo_cell._tc.xpath(".//w:drawing"):
        run = drawing.getparent()
        run.getparent().remove(run)
    merged = name_cell.merge(photo_cell)
    # merge() appends the (now empty) photo paragraph — drop it so the header keeps its height
    paras = merged.paragraphs
    if len(paras) > 1 and not paras[-1].text.strip():
        paras[-1]._p.getparent().remove(paras[-1]._p)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    _remove_photo(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
