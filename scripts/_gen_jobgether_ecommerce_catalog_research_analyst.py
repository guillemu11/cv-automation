"""One-off: CV + cover letter for "Ecommerce Product Catalog & Research Analyst", listed by
Jobgether on behalf of an undisclosed partner (UAE-based, fully remote). 2026-09-13.

Paula's honest angle: catalogue accuracy is part of her week. At DoFreeze she runs the Shopify
store (product setup, pricing, content, images, collections) and keeps listings, content and
images correct on Careem, talabat and Noon. At Miravia she onboarded 30+ fragrance stores in two
months (product data, pricing, presentation). She built a Claude-based automation that does market
research and reporting (~40% less manual work): a real match for the "AI-powered research tools"
requirement. No competitor-benchmarking claim (not in profile.yaml).

Honesty guardrails:
- NO Magento. Shopify is the "comparable ecommerce platform"; never claim Magento.
- NO technology-product catalogue (IT hardware, software, cloud). Her categories are food, beauty,
  fashion. Not claimed; the letter names it.
- Team of two stated exactly; never "own the forecast"; no Arabic; no sponsorship claim.
- Analyst role, remote via Jobgether: salary likely below the 20k AED/month floor. Flagged.

ONE-PAGE standard. Lands under output/2026-09-13/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cover_letter as cl
from career_ops.generators import cv_generator as cv

COMPANY = "Jobgether"
TITLE = "Ecommerce Product Catalog & Research Analyst"
DATE_FOLDER = "2026-09-13"
CONTACT = None

JOB_DESCRIPTION = """\
Ecommerce Product Catalog & Research Analyst — listed by Jobgether for a partner company, United Arab
Emirates, fully remote.
Research newly launched technology products from manufacturer websites and trusted sources; create and
maintain ecommerce listings (manufacturer numbers, pricing, technical specifications, descriptions,
images); update existing listings for new models, spec and price changes; run catalog audits for
missing, outdated, inconsistent or inaccurate data; ensure data meets quality standards, marketplace
requirements and catalog guidelines; use AI-powered research and data extraction tools; monitor
competitor websites for trends, pricing and positioning; manage multiple assignments under deadlines;
collaborate with teams to resolve product-data issues.
Needs: hands-on Magento or comparable ecommerce platform; research and analytical skills, attention to
detail; AI research tools and automated data extraction; understanding of IT hardware/software, cloud
and AI products; Excel and Word; interpret technical specs into clear content; organization and time
management; proactive data-quality mindset; independent work across a large catalog.
"""

ATS = [
    "ecommerce", "product catalog", "catalog management", "product listings", "product data",
    "product information", "data quality", "catalog audits", "listing accuracy", "product research",
    "manufacturer", "technical specifications", "pricing", "descriptions", "images", "product content",
    "Magento", "Shopify", "ecommerce platform", "marketplace", "marketplace requirements",
    "AI-powered research tools", "AI tools", "data extraction", "automation", "competitor monitoring",
    "competitor analysis", "price benchmarking", "market trends", "technology products", "Excel",
    "Microsoft Office", "Word", "attention to detail", "analytical", "deadlines", "remote", "UAE",
]

CV_CONTENT = {
    "headline": (
        "E-Commerce Catalogue & Product Data · Marketplace Listings · AI-Assisted Research · "
        "Listing Audits"
    ),
    "professional_summary": (
        "E-commerce professional with 5 years across marketplaces, quick-commerce and FMCG. I keep "
        "product catalogues accurate on Shopify, Careem, talabat and Noon, and built a Claude-based "
        "automation for market research and reporting. Marketplace background at Alibaba's Miravia."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG food group (Befit, Eurocake, Smash) | Shopify, Careem, talabat, Noon | 50+ markets",
            "bullets": [
                "Run the Shopify store catalogue: product setup, pricing, descriptions, images and collections, keeping product data accurate through 6 new launches",
                "Maintain listings, content and images on Careem, talabat and Noon to each platform's requirements; update them as products and prices change",
                "Built an AI-powered automation (Claude) for market research, content and KPI reporting, cutting manual workload ~40%",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Top-5 global e-commerce marketplace | 100K+ employees",
            "bullets": [
                "Onboarded 30+ fragrance stores in two months as PIC Fragrances, checking product data, pricing and presentation against marketplace guidelines",
                "Grew 42 brand accounts +30% GMV QoQ through assortment, pricing and promotions, keeping catalogue data consistent",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Helped build the Retail vertical: onboarded brands and managed product listings, working with operations to resolve catalogue issues",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | Chocolate category",
            "bullets": [
                "Sell-in/sell-out, pricing and promotional analysis in Excel; built performance reports for the chocolate category",
            ],
        },
    ],
    "skills_brand": (  # -> "Catalogue"
        "product setup, specs & descriptions, images, pricing updates, listing audits, data quality"
    ),
    "skills_ecommerce": (  # -> "Platforms"
        "Shopify, Miravia (Alibaba), Careem, talabat, Noon, Glovo: marketplace listing requirements"
    ),
    "skills_commercial": (  # -> "Research"
        "product & market research, trend tracking, sell-in/sell-out and pricing analysis"
    ),
    "skills_data": (  # -> "AI & Data"
        "Claude, ChatGPT, AI-assisted research and data extraction, workflow automation, dashboards"
    ),
    "skills_tools": (
        "Excel (advanced), Word, PowerPoint, Google Sheets, Power BI, Looker, SAP, Salesforce"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Catalogue",
    "E-Commerce & Digital": "Platforms",
    "Commercial": "Research",
    "Data & Analytics": "AI & Data",
}

CL_PARAGRAPHS = {
    "opening_paragraph": (
        "I would like to apply for the Ecommerce Product Catalog & Research Analyst role. Keeping product "
        "data accurate is already part of my week: at DoFreeze, a UAE food group, I run our Shopify "
        "catalogue and our listings on Careem, talabat and Noon. I am based in the UAE and work well "
        "independently, which suits a remote role."
    ),
    "body_paragraph_1": (
        "On Shopify I set up products with pricing, descriptions, images and collections, and I keep "
        "our quick-commerce listings current as products and prices change. At Alibaba's Miravia I onboarded more than 30 fragrance stores in two months, checking "
        "product data, pricing and presentation against marketplace guidelines,, and grew 42 brand accounts by 30% GMV quarter on quarter. I also built an automation with Claude that "
        "gathers market research and produces our KPI reports, cutting manual work by "
        "about 40%. Using AI tools to research, extract and check information is how I work every day."
    ),
    "body_paragraph_2": (
        "To be transparent: I have not used Magento, and my catalogues have been food, beauty and fashion "
        "rather than IT hardware, software or cloud products. What transfers is the method: reading "
        "a source carefully, turning specifications into clear product content, and auditing a catalogue "
        "for errors. I learn new platforms and product categories quickly, and I follow AI products "
        "closely because I use them."
    ),
    "closing_paragraph": (
        "I would welcome the chance to discuss the role. Thank you for your consideration."
    ),
}


def make_job() -> Job:
    return Job(
        id="jobgether-ecommerce-catalog-research-analyst-2026-09",
        title=TITLE,
        company=COMPANY,
        location="United Arab Emirates (Remote)",
        url="https://jobgether.com/",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={"team": "Partner company (undisclosed) via Jobgether",
             "workplace": "Fully remote",
             "note": "Jobgether AI shortlist. Gaps: Magento, technology-product catalogue."},
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not posted (Analyst band, remote — likely below 20k AED/month floor)",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 50, "ai_tier": "Possible",
        "skills_match": [
            "Shopify catalogue: product setup, pricing, descriptions, images",
            "Marketplace listings on Careem/talabat/Noon; Miravia onboarding of 30+ stores",
            "AI-powered research automation with Claude (~40% less manual work)",
        ],
        "missing_skills": [
            "Magento (Shopify only)",
            "Technology products: IT hardware, software, cloud, AI products",
        ],
        "sector_fit": "weak — tech-product catalogue vs food/beauty/fashion",
        "seniority_fit": "below level — Analyst role vs Manager title",
        "red_flags": [
            "Analyst band: salary likely under her 20k AED/month floor",
            "Undisclosed partner company via Jobgether",
        ],
        "ats_keywords": ATS,
        "reasoning": "Catalogue and AI-research skills are real; tech-product and Magento gaps plus seniority make it a long shot.",
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Applied",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _pages(pdf: Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(pdf)).pages)


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    final_dir = settings.output_dir / DATE_FOLDER / f"{COMPANY} - {TITLE}"
    dest = final_dir / "01_CV_y_Carta"
    dest.mkdir(parents=True, exist_ok=True)

    cv_docx = cv._fill_template(CV_CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    cl_docx = cl._fill_template(CL_PARAGRAPHS, job, CONTACT)
    cl_pdf = _to_pdf_soffice(cl_docx)

    for src, short in ((cv_pdf, "Paula De Francisco - CV.pdf"),
                       (cl_pdf, "Paula De Francisco - Cover Letter.pdf")):
        moved = dest / src.name
        shutil.move(str(src), str(moved))
        shutil.copy(str(moved), str(dest / short))
        print("OK", moved.name, "PAGES", _pages(moved))

    stray = cv_pdf.parent.parent
    if stray.exists() and stray.resolve() != final_dir.resolve():
        shutil.rmtree(stray, ignore_errors=True)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
