"""One-off: CV de una página para **Performance Marketing Manager** en **izil Beauty**
(Dubai; skincare/haircare premium marroquí, e-commerce + servicios + retail).

Pide el JD: estrategia de performance alineada con objetivos comerciales; adquisición,
lead gen, remarketing y retención; KPIs ROAS, CAC, CPA, LTV, CR y crecimiento de revenue;
paid media en Meta, Google (Search, Display, Shopping, PMax, YouTube), TikTok y Snapchat;
gestión de presupuestos, pacing y forecasts; CRO y A/B testing sobre landings, ofertas y
funnels; CRM y retención (email, SMS, WhatsApp, loyalty); dashboards de spend/revenue/ROAS/
CAC/CPA; atribución, píxeles, conversion API, GA4, GTM, Looker Studio, Merchant Center y
feeds; colaboración con contenido, e-commerce, retail y agencias; guiar a juniors y agencias.

Ángulo honesto de Paula: 5 años; hoy lleva Meta Ads y Google Ads hands-on para tres marcas
en DoFreeze con Shopify, UTMs, píxeles y reporting semanal/mensual; Miravia (Alibaba) — 42
cuentas beauty/fragancias con inversión promocional leída contra tráfico, conversión, ROI y
ROAS (+30% GMV QoQ) y el canal Flash Sales contra objetivos de P&L; gestiona un equipo de dos
y 4 agencias. Sector beauty real: Miravia (42 marcas de beauty/fragancias, KIKO Milano).

Guardarraíles de honestidad:
- Plataformas hands-on: **Meta Ads y Google Ads (Search & Display)**. NO se reclama TikTok
  Ads, Snapchat Ads, YouTube, Performance Max, Shopping ni Merchant Center. TikTok/Instagram
  solo como contenido orgánico.
- GA4 / Google Tag Manager: certificada (Google Digital Marketing & E-Commerce) y uso de
  UTMs, píxeles y analítica; no se reclama implementación de Conversion API ni modelado de
  atribución avanzado.
- Forecast: contribuye con Sales/Finance, no es dueña.
- Visa: solo "UAE Residence Visa"; nunca "no sponsorship needed".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "izil Beauty"
TITLE = "Performance Marketing Manager"
DATE_FOLDER = "2026-09-22"

JOB_DESCRIPTION = """\
Performance Marketing Manager — izil Beauty, Dubai, UAE. On-site, full time. Premium Moroccan
skincare and haircare brand.
Role overview: develop, execute and optimize digital marketing campaigns that drive e-commerce sales,
service bookings, retail traffic and customer acquisition, combining analytical expertise, commercial
judgment and creative thinking to deliver measurable growth through paid media, conversion optimization
and data-driven marketing strategies.
Performance marketing strategy: align performance strategy with commercial objectives, annual plans and
growth targets; manage acquisition, lead generation, remarketing and retention campaigns; define and
improve KPIs including ROAS, CAC, CPA, LTV, conversion rate and revenue growth; translate business
priorities into media plans, channel targets and testing roadmaps.
Paid media: plan, launch and optimize campaigns across Meta Ads, Google Ads, TikTok Ads, Snapchat Ads;
manage Google Search, Display, Shopping, Performance Max and YouTube; optimize audiences, bidding,
placements, feeds and creatives for ROAS; manage budgets, pacing, monthly and annual forecasts; maintain
campaign governance, tracking accuracy, QA and platform compliance.
E-commerce and conversion optimization: drive qualified traffic, conversions and revenue through the
e-commerce platform; monitor website performance, merchandising journeys and funnel drop-off; run
structured A/B testing on landing pages, offers, creatives, audiences and funnels; partner with
e-commerce and technology teams on UX, conversion rate and checkout performance.
CRM and retention: improve retention, repeat purchase rate and customer lifetime value; support email,
SMS, WhatsApp, loyalty and lifecycle marketing; build segmentation, remarketing and lookalike strategies
from first-party data.
Analytics and reporting: analyze channel, campaign and funnel performance; build dashboards covering
spend, revenue, ROAS, CAC, CPA and conversion; monitor competitors, market trends and platform
developments; present results, risks, forecasts and recommendations to management; maintain reliable
attribution, tracking and reporting.
Collaboration: work with content creators, designers, e-commerce, CRM, retail teams, service operations
and external agencies; guide creative briefs and content testing; coordinate campaign calendars;
contribute to annual plans, commercial forecasts and budget preparation; coach junior team members and
agency partners.
Requirements: bachelor's degree in Marketing, Business Administration or Digital Marketing; minimum five
years in performance marketing, digital marketing or paid media; experience managing significant
advertising budgets and delivering revenue growth; beauty, skincare, wellness, luxury retail, consumer
services or e-commerce preferred; UAE or GCC markets preferred; advanced Meta Business Manager, Google
Ads, GA4, Google Tag Manager, Looker Studio and TikTok Ads Manager; attribution models, pixels, conversion
APIs, tracking implementation; product feeds, Google Merchant Center, remarketing and audience management
preferred; SEO, CRM and marketing automation an advantage.
"""

ATS = [
    "performance marketing", "paid media", "customer acquisition", "remarketing", "retention",
    "lead generation", "ROAS", "CAC", "CPA", "LTV", "conversion rate", "revenue growth",
    "Meta Ads", "Meta Business Manager", "Google Ads", "Google Search", "Display", "paid social",
    "audiences", "bidding", "creatives", "budget management", "pacing", "forecasts",
    "conversion optimization", "CRO", "A/B testing", "landing pages", "funnel", "checkout",
    "e-commerce", "Shopify", "merchandising journeys", "first-party data", "segmentation",
    "lookalike", "CRM", "email", "WhatsApp", "lifecycle marketing", "loyalty",
    "dashboards", "attribution", "tracking", "pixels", "UTM", "Google Analytics", "GA4",
    "Google Tag Manager", "Looker Studio", "reporting", "competitor analysis",
    "agencies", "creative briefs", "campaign calendar", "beauty", "skincare", "wellness",
    "luxury retail", "UAE", "GCC", "Dubai", "coaching",
]

CONTENT = {
    "headline": "Performance Marketing · Meta & Google Ads · E-Commerce CRO · Beauty & Retail",
    "professional_summary": (
        "Performance and e-commerce marketer with 5 years across beauty, FMCG and marketplaces. Runs Meta and "
        "Google Ads hands-on for three brands in Dubai against ROI/ROAS, and owns the Shopify funnel, tracking "
        "and reporting behind the click."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG group (Befit, Eurocake, Flair) | Shopify D2C + Noon, talabat, Careem | 50+ markets",
            "bullets": [
                "Plan, launch and optimise paid media on Meta (Facebook & Instagram) and Google Ads for three brands — campaign structure, budget allocation and pacing, audience building, remarketing and creative A/B testing — managed daily against ROI and ROAS",
                "Own the Shopify store and its tracking end-to-end: UTM-tagged campaigns, platform pixels and analytics read from click to checkout, testing landing pages, offers and creatives to lift conversion rate and AOV",
                "Build KPI dashboards on spend, revenue, conversion and ROAS, present monthly performance and recommendations to management, and contribute to the annual plan and forecast with Sales and Finance",
                "Lead a team of two (designer + social executive) and brief 4 external agencies — turning campaign insight into creative briefs, content testing and a shared campaign calendar; run EDM and lifecycle sends alongside paid",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's lifestyle marketplace | Beauty & fragrance categories | 100K+ employees",
            "bullets": [
                "Grew 42 beauty, fragrance and fashion accounts +30% GMV QoQ by allocating promotional and media investment where it paid back — testing offers, pricing and campaigns weekly against traffic, conversion, ROI and ROAS",
                "Owned the Flash Sales channel reporting to the CEO: paced investment against P&L targets and presented performance, risks and next steps monthly",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Ran in-app promo mechanics and bespoke activations for XL partners (KFC, Taco Bell, Sushi Shop), reading order, conversion and GMV data to grow volume profitably",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | Category planning & analytics",
            "bullets": [
                "Measured promotional effectiveness and sell-in/sell-out with Nielsen — the measurement discipline behind deciding which spend earns more budget",
            ],
        },
    ],
    "skills_brand": (  # label -> "Performance Marketing"
        "acquisition & remarketing, budget allocation & pacing, audience building, creative A/B testing, media plans"
    ),
    "skills_ecommerce": (  # label -> "Platforms & Channels"
        "Meta Ads (Facebook & Instagram), Google Ads (Search & Display), Shopify, EDM & lifecycle, marketplaces"
    ),
    "skills_commercial": (  # label -> "Tracking & CRO"
        "UTM frameworks, platform pixels, conversion tracking, Google Analytics & Tag Manager, funnel & drop-off, landing-page testing"
    ),
    "skills_data": (  # label -> "Analysis & Reporting"
        "ROAS, ROI, CPA, conversion rate, AOV, GMV, KPI dashboards, monthly reporting, forecasting support"
    ),
    "skills_tools": (  # label -> "Tools"
        "Meta Ads Manager, Google Ads, Google Analytics, Shopify, Power BI, Looker, Tableau, Canva, Claude (AI)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Performance Marketing",
    "E-Commerce & Digital": "Platforms & Channels",
    "Commercial": "Tracking & CRO",
    "Data & Analytics": "Analysis & Reporting",
}


def make_job() -> Job:
    return Job(
        id="izil-beauty-performance-marketing-manager-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE",
        url="https://www.linkedin.com/jobs/search/?keywords=izil%20Beauty%20Performance%20Marketing%20Manager%20Dubai",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "izil Beauty Performance Marketing Manager Dubai",
            "note": "Beauty premium marroquí (skincare/haircare), 106 empleados. Encaje de sector fuerte, "
                    "pero el JD es performance puro con TikTok/Snapchat Ads, PMax, Shopping, Merchant Center, "
                    "GA4/GTM/Looker Studio y Conversion API — Paula es hands-on solo en Meta y Google "
                    "(Search & Display). 987 candidatos, +100 solicitudes.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 62, "ai_tier": "Warm",
        "skills_match": [
            "Meta Ads y Google Ads hands-on a diario para tres marcas, optimizando por ROI/ROAS",
            "Shopify end-to-end con UTMs, píxeles y analítica: CRO, AOV, testeo de landings y ofertas",
            "Sector beauty real: 42 cuentas de beauty y fragancias en Miravia (+30% GMV QoQ), KIKO Milano",
            "Dashboards de KPIs y reporting mensual a dirección; aporta al plan anual y al forecast",
            "Gestiona equipo de dos y 4 agencias: briefs creativos, calendario de campañas, EDM/lifecycle",
        ],
        "missing_skills": [
            "TikTok Ads y Snapchat Ads (explícitos en el JD) — sin gestión de campañas",
            "Google Shopping, Performance Max, YouTube, Merchant Center y feeds de producto",
            "GA4 / GTM / Looker Studio a nivel avanzado, Conversion API y modelos de atribución",
            "Presupuestos publicitarios de gran volumen gestionados en solitario",
        ],
        "sector_fit": "alto — beauty/skincare premium y e-commerce, su terreno",
        "seniority_fit": "encaje — Manager con 5 años, justo el mínimo que piden",
        "red_flags": [
            "987 candidatos y +100 solicitudes",
            "Stack técnico del JD (PMax, Shopping, TikTok/Snapchat Ads, CAPI, Looker Studio) por encima de su experiencia real",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Encaje de sector muy bueno (beauty premium, e-commerce, Dubai) y de nivel (Manager, 5 años). "
            "El riesgo es técnico: el JD pide un performance marketer de stack completo — TikTok y Snapchat Ads, "
            "Shopping/PMax, Merchant Center, GA4/GTM/Looker Studio y Conversion API — y ella es hands-on en Meta "
            "y Google Search/Display con Shopify. El CV apoya el caso en resultados comerciales y CRO, sin "
            "inventar plataformas; conviene reconocer la brecha y ofrecer rampa de aprendizaje."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
