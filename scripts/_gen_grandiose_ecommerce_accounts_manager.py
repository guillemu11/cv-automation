"""One-off: CV de una página para **eCommerce Accounts Manager** en **Grandiose** (departamento de
eCommerce; marcas de comida, tartas y flores). Dubái. Oferta pegada desde LinkedIn; sin datos de
solicitudes ni sueldo.

Qué es el puesto: dueño de punta a punta del negocio online (ventas y rentabilidad) de una cartera
F&B y floral en agregadores (canal principal) y en la app propia: 5 marcas de restaurante con 30
locales, 15 marcas de cloud kitchen en 3 cocinas, 3 marcas de flores a domicilio y 2 de tartas, y
creciendo. Pide el JD: comerciales con agregadores (precio, promociones, comisiones, fees, inversión
en marketing), revisiones de cuenta y negociación, visibilidad y campañas patrocinadas; KPIs
operativos (disponibilidad, aceptación, cancelaciones, tiempos de preparación, exactitud de pedido,
valoraciones, entrega) con visitas a locales y cocinas; exactitud y optimización de menú y catálogo
(bundles, add-ons, AOV); lanzamientos online de marcas y locales nuevos (alta, menú, contenido,
precio, formación). Requisitos: experiencia en alimentación **esencial**, historial demostrable
haciendo crecer marcas de comida y gestión hands-on de agregadores. Flores: ventaja, no requisito.
No pide árabe.

Ángulo honesto de Paula:
- **Los dos lados del agregador**: en Glovo fue Account Manager de cadenas de restaurante XL (KFC,
  Taco Bell, La Tagliatella, Sushi Shop), con planes de crecimiento de GMV, activaciones en la app y
  negociación de condiciones; hoy gestiona sus marcas en talabat, Noon y Careem desde el lado marca.
- **Marcas de comida y tarta que crecen con métrica**: Eurocake es una marca de tartas. Lanzamiento de
  SMASH (marca nueva de Eurocake) en talabat: venta diaria +165% frente a +4,7% de la marca de
  control. Befit × Noon "New Year, New Me": +4.176 uds incrementales, +31% sobre baseline.
- **Apoyo promocional del agregador**: sorteo ligado a compra con tarjetas de cashback de talabat y
  samplings estacionales con talabat.
- **Disponibilidad**: hace ella la revisión semanal de riesgo de stock por dark store (cobertura de
  PO frente a demanda). Aporta al forecast mensual de sell-in.
- **App propia ≈ su Shopify**: es suya de punta a punta (catálogo, fichas, colecciones, descuentos,
  checkout) con merchandising sobre conversión y AOV; definió el MSL de la web.
- **Comerciales de plataforma**: Miravia (Alibaba), 42 cuentas, +30% GMV QoQ con precio, surtido y
  promociones; canal Flash Sales reportando al CEO contra objetivos de P&L.

Guardarraíles de honestidad:
- **Operación de restaurante desde el lado marca** (30 locales, cloud kitchens, tiempos de
  preparación, aceptación, cancelaciones, exactitud de pedido, visitas a local, formación de
  equipos de tienda): sin evidencia. No se menciona. Pendiente preguntar a Paula si en Glovo
  trabajaba esos KPIs con los restaurantes; si es así, se puede añadir.
- **Comisiones y fees**: "commercial terms", sin afirmar que negoció tasas de comisión concretas.
- **Bundles, add-ons, modifiers**: sin evidencia; no se afirman.
- **Flores**: cero experiencia (es ventaja, no requisito).
- **Deliveroo**: DoFreeze NO está. Plataformas: talabat, Noon, Careem. Forecast: "contribute".
- Visa: "UAE Employment Visa (employer-sponsored)".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Grandiose"
TITLE = "eCommerce Accounts Manager"
DATE_FOLDER = "2026-10-04"

JOB_DESCRIPTION = """\
eCommerce Accounts Manager. Grandiose. Dubai, UAE.
We are looking for eCommerce Accounts Manager for eCommerce dept. to manage F&B and Floral Brands of Grandiose.
Role overview: Accounts Manager will take end-to-end ownership of the online business performance and commercials of
our food, cake and flower brands across aggregators and our own app. The role is responsible for a substantial business
portfolio, with aggregators representing the primary sales channel, followed by our own app. It combines commercial
ownership, partner relationship management, operational performance improvement, marketing coordination and new brand
and location launches. Current portfolio: 5 brick-and-mortar food brands operating across 30 branches; 15 cloud
kitchen food brands operating from 3 cloud kitchens; 3 flower delivery brands; 2 cake brands. The portfolio is
continuously growing. Strong food industry experience and a proven track record of growing food brands are essential.
Prior flower delivery experience is an advantage but is not required.
Business performance and commercial ownership: end-to-end ownership of online performance and commercials across all
brands, locations and platforms; drive sales growth and profitability; monitor daily performance by brand, location
and platform; manage pricing, promotions and platform commercials; assess the impact of commissions, platform fees,
discounts and marketing investment; present performance reviews, forecasts and recommendations to management.
Aggregator and platform management: own day-to-day relationships with aggregators; lead account reviews, commercial
negotiations, campaign planning and issue resolution; secure visibility and promotional support; manage the food,
cake and flower offering across all platforms including our own app; identify and close new partnerships.
Operational performance and site engagement: monitor and improve availability, order acceptance, cancellations,
preparation times, order accuracy, customer ratings and delivery performance; regular visits to branches, cloud
kitchens and fulfilment sites; work with site teams, operations, technical teams and last-mile partners; coordinate
training on platform processes, order handling, menu updates and service standards.
Menu, pricing and assortment optimization: own menu and catalogue accuracy (availability, descriptions, images,
modifiers, pricing); optimize menu structure, product selection, bundles and add-ons to improve conversion, average
order value and profitability; competitor and pricing research.
Marketing, visibility and conversion: promotional plans aligned with seasonal opportunities; aggregator campaigns and
sponsored placements; improve the path from visibility to completed orders; monitor campaign performance, conversion
and return on marketing spend; ensure campaigns are supported by product availability and operational readiness.
New brand and location launches: lead online launch coordination for new brands, branches and cloud kitchen
locations; platform onboarding and readiness (listings, menus, content, pricing, operational setup, training, launch
promotions); monitor initial performance.
Stakeholder coordination: commercial, marketing, operations, content, technical and last-mile teams, aggregators and
external partners; scalable processes for a growing portfolio.
Required: proven track record of managing and growing food brands with measurable results in online sales,
profitability and operational performance; hands-on experience managing aggregators (relationships, commercial
negotiations, promotions, performance improvement); strong understanding of the online food business and e-commerce,
ideally across multiple brands, locations and cloud kitchens; aggregator marketing, visibility and conversion
(sponsored placements, platform campaigns, menu optimization, pricing, promotions, customer ratings); strong commercial
judgment; strong analytical skills with platform dashboards, spreadsheets and performance reports; excellent
communication, negotiation and relationship management; launch coordination and platform training; willingness to
conduct regular site visits and partner meetings. Food industry experience is essential.
"""

ATS = [
    "eCommerce", "e-commerce", "account management", "accounts manager", "aggregators", "food delivery",
    "online food business", "food industry", "F&B", "food brands", "cake brands", "talabat", "Noon", "Careem",
    "Glovo", "own app", "D2C", "Shopify", "commercials", "commercial negotiations", "commercial terms",
    "pricing", "promotions", "platform campaigns", "visibility", "sponsored placements", "conversion",
    "average order value", "AOV", "sales growth", "profitability", "P&L", "GMV", "ROI", "ROAS",
    "return on marketing spend", "performance reviews", "forecasts", "availability", "stock", "dark stores",
    "menu and catalogue accuracy", "listings", "product content", "assortment", "new brand launches",
    "platform onboarding", "launch promotions", "partner relationship management", "partnerships",
    "restaurant chains", "QSR", "seasonal campaigns", "Ramadan", "dashboards", "Excel", "cross-functional",
    "stakeholder management", "marketing", "operations", "UAE", "Dubai",
]

CONTENT = {
    "headline": "eCommerce Account Management · Food Aggregators & Quick-Commerce · F&B Brand Growth · UAE",
    "professional_summary": (
        "F&B e-commerce and account professional in Dubai who grows food and cake brands (Befit, Eurocake) on "
        "talabat, Noon and Careem and an own Shopify store. Before that, on the aggregator side: XL restaurant "
        "chains at Glovo (KFC, Taco Bell, La Tagliatella) and 42 accounts at Miravia (Alibaba), +30% GMV QoQ."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Homegrown UAE F&B group (Befit, Eurocake, Flair) | talabat, Noon & Careem + Shopify D2C | 50+ markets",
            "bullets": [
                "Grow the food and cake brands on talabat, Noon and Careem — listings, pricing, promotional mechanics and joint activations with the platforms (purchase-linked talabat cashback giveaway, seasonal sampling); Befit × Noon \"New Year, New Me\" added +4,176 incremental units, +31% over baseline",
                "Launched SMASH, a new Eurocake brand, on talabat with campaigns built to drive orders: daily sales +165% vs +4.7% for a control brand on the same platform; 6 product launches end-to-end (brief, packaging, pricing, go-to-market)",
                "Run the weekly stock-risk review by dark store on talabat, Noon and Careem, checking PO coverage against current and upcoming demand to protect availability; contribute to the monthly sell-in forecast",
                "Own the Shopify store end-to-end — catalogue, product content, collections, discounts and checkout — merchandised on conversion and AOV; plan Meta and Google Ads on ROAS; lead a team of two (designer and social media executive)",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees | platform side",
            "bullets": [
                "Managed 42 brand accounts on pricing, assortment, product content and promotions, delivering +30% GMV QoQ; onboarded 30+ new stores in two months as category lead",
                "Owned the Flash Sales channel for Beauty, Fashion & Home reporting to the CEO — trading plans against P&L targets, using conversion, traffic and ROI data to reallocate promotional investment",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Food-delivery & quick-commerce platform | €500M+ revenue | aggregator side",
            "bullets": [
                "Managed XL restaurant chains (KFC, Taco Bell, La Tagliatella, Sushi Shop) with data-led GMV growth plans and in-app marketing activations, negotiating commercial terms that kept both the partner and Glovo profitable",
                "Coordinated marketing, logistics and customer support to deliver campaigns and lift order volume; helped build the Retail vertical, onboarding new brands and setting up their listings",
            ],
        },
        {
            "company": "Mondelez International · Massimo Dutti (Inditex)",
            "role": "Trainee, Category Planning · Sales Associate",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG (chocolate category) | Premium fashion retail",
            "bullets": [
                "At Mondelez, sell-in/sell-out and promotional-effectiveness analysis for chocolate, contributing to NPD launches (Milka Spread, Mini Suchard); at Inditex, store operations and customer service",
            ],
        },
    ],
    "skills_brand": (  # -> "Aggregators"
        "talabat, Noon, Careem, Glovo · account management, negotiation, onboarding, joint campaigns"
    ),
    "skills_ecommerce": (  # -> "Catalogue & Ops"
        "listings & product content, assortment & MSL, availability & stock risk, launches, Shopify"
    ),
    "skills_commercial": (  # -> "Commercial"
        "pricing, promotions, sales & GMV growth, P&L targets, promo ROI, forecast input"
    ),
    "skills_data": (  # -> "Marketing & CVR"
        "platform campaigns, sampling, creators, Meta & Google Ads, ROAS, conversion, AOV"
    ),
    "skills_tools": (  # -> "Tools"
        "Excel (Advanced), Power BI, Tableau, Looker, Salesforce, SAP, Shopify, Claude (AI)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Aggregators",
    "E-Commerce & Digital": "Catalogue & Ops",
    "Commercial": "Commercial",
    "Data & Analytics": "Marketing & CVR",
}


def make_job() -> Job:
    return Job(
        id="grandiose-ecommerce-accounts-manager-dubai-2026-10",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE",
        url="https://www.linkedin.com/jobs/search/?keywords=Grandiose%20eCommerce%20Accounts%20Manager",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Grandiose eCommerce Accounts Manager Dubai",
            "note": "Dpto. eCommerce de Grandiose: marcas de comida (5 con 30 locales + 15 cloud kitchen), 3 de flores "
                    "y 2 de tartas, en agregadores y app propia. Encaje: Glovo (cadenas de restaurante XL, lado "
                    "agregador) + DoFreeze (marcas de comida y tarta en talabat, Noon y Careem; SMASH +165%). "
                    "Gaps: operación multi-local de restaurante y cloud kitchen, visitas a local, flores.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 75, "ai_tier": "Warm",
        "skills_match": [
            "Account Manager en Glovo de cadenas de restaurante XL (KFC, Taco Bell, La Tagliatella, Sushi Shop)",
            "Gestión hands-on de agregadores desde el lado marca: talabat, Noon y Careem",
            "Marcas de comida y tarta (Befit, Eurocake) con resultados medibles: SMASH +165%, Befit × Noon +31%",
            "Apoyo promocional de la plataforma: sorteo con cashback de talabat y samplings estacionales",
            "Revisión semanal de riesgo de stock por dark store (disponibilidad)",
            "Dueña del canal propio (Shopify): catálogo, fichas, descuentos, conversión y AOV",
            "Comerciales de plataforma en Miravia: 42 cuentas, +30% GMV QoQ, Flash Sales reportando al CEO",
            "Lanzamientos: marca nueva SMASH y 6 lanzamientos de producto",
        ],
        "missing_skills": [
            "Operación multi-local de restaurante y cloud kitchen desde el lado marca",
            "KPIs operativos de restaurante (aceptación, tiempos de preparación, cancelaciones, exactitud) sin evidencia",
            "Visitas a locales y formación de equipos de tienda en procesos de plataforma",
            "Optimización de menú con bundles, add-ons y modifiers",
            "Flores a domicilio (ventaja, no requisito)",
        ],
        "sector_fit": "muy bueno — F&B y tartas (Eurocake); agregadores de comida desde los dos lados (Glovo y marca)",
        "seniority_fit": "acorde — Manager, igual que su puesto actual",
        "red_flags": [
            "Cartera grande (25 marcas, 30 locales, 3 cloud kitchens): pedirán operación de restaurante que no tiene",
            "Sueldo no publicado: confirmar la banda frente al suelo de 20.000 AED/mes",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Encaje fuerte en lo esencial: experiencia en alimentación y gestión hands-on de agregadores, y desde los "
            "dos lados. En Glovo llevó cadenas de restaurante XL como Account Manager y hoy gestiona marcas de comida "
            "y tarta en talabat, Noon y Careem con resultados medibles (SMASH +165% frente a control, Befit × Noon "
            "+31%). El canal propio encaja con su Shopify y los comerciales con Miravia. El hueco está en la "
            "operación de restaurante multi-local (cloud kitchens, tiempos de preparación, visitas a local), que el "
            "JD pesa mucho; si en Glovo trabajaba esos KPIs con los restaurantes, conviene añadirlo."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
