"""One-off: generate Paula's CV for **Merchant - Fragrances** at **Sephora (LVMH)**
(Dubai, UAE; full-time; posted via the TALENTMATE aggregator portal but the role
is 100% Sephora — JD + inside-sephora.com).

Category/merchant role for the FRAGRANCE category at a prestige beauty retailer:
assortment management & merchandising (SKU-level assortments, stock planning,
planograms, category calendar, SKU sales forecasts, inventory analysis), brand
management (end-to-end brand lifecycle, brand-partner relationships & business
reviews, brand strategies driving sales/share/equity, performance reporting,
business presentations), omnichannel & marketing (activations across retail &
e-com, samples/gifts for CRM, PR/influencer/seeding, digital campaigns, social
content), e-commerce assortment (online strategy, hero products, exclusives,
conversion), and reporting & financial management (monthly KPI reports, budget
tracking by country/store). Asks 3-4 yrs in Merchandising/Category/Brand/Buying,
preferably beauty/fragrance/FMCG/retail; strong analytics; brand & assortment &
commercial-partnership experience; Advanced Excel & PowerPoint. NO Arabic required.

This is one of Paula's strongest, most on-target fits and fully honest:
- HERO: Key Account Manager for **Beauty, Fragrances & Fashion** at Alibaba's
  Miravia — 42 brand accounts, +30% GMV QoQ, assortment/pricing/promotions; as
  **PIC Fragrances** onboarded the leading Arabian & oud houses (Arabian Oud,
  Lattafa, Swiss Arabian, Ajmal); ran the Flash Sales channel (to CEO).
- DoFreeze: brand lifecycle, assortment/listings across retail + e-commerce
  (Shopify + quick-commerce), NPD launches, activations, sampling & seeding,
  influencer/PR, A&P budget & KPI reporting.
- Mondelez: category planning — SKU sell-in/sell-out & inventory analysis,
  promo-effectiveness, forecasting, NPD.

Honesty guardrails: fragrance-category depth is real (Miravia); planograms/visual
merchandising grounded in Miravia assortment + Inditex/Massimo Dutti retail; NO
Arabic claimed (not required); factual "UAE Residence Visa" only.

ONE-PAGE standard (since 2026-08-27): tight summary, 3 bullets current / 1-2 older,
~6 skills per row, SHORT skills-row labels (≤~20 chars) so they don't wrap. Fills
the real CV template, relabels the skills rows for a category/merchant profile
(template stays pristine), converts to PDF via LibreOffice, registers the job,
verifies 1 page, and lands the package under output/2026-08-30/. Also drops a
short-named 'Paula De Francisco - CV.pdf' copy for portals.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Sephora"
TITLE = "Merchant - Fragrances"
DATE_FOLDER = "2026-08-30"

JOB_DESCRIPTION = """\
Merchant - Fragrances — Sephora (LVMH), Dubai, UAE (posted via TALENTMATE portal).

The Fragrance Merchant supports the execution of merchandising strategies at
regional level by adapting central assortment, allocation and activation plans to
local market needs — ensuring product relevance, inventory optimization and
consistent execution across stores and e-commerce.

Assortment Management & Merchandising: support category strategies incl. sales
forecasting, stock planning and product lifecycle management; manage SKU-level
assortments across fragrance brands; monitor stock across stores and e-commerce;
maintain/optimize planograms; own the fragrance category calendar (launches,
expansions, store openings, discontinuations); prepare SKU-level sales forecasts by
country; analyze inventory performance (overstock/slow-movers) and recommend
actions; run assortment reviews; coordinate regulatory/compliance for launches;
plan activations/launches/promos with brand partners; handle shelving plans.

Brand Management: manage end-to-end lifecycle of assigned fragrance brands; build
strong brand-partner relationships; develop and execute brand strategies driving
sales, share and equity; lead business reviews; identify commercial opportunities;
monitor and report brand performance with actionable insights; develop compelling
business presentations from data.

Omnichannel & Marketing: coordinate brand activations across retail and e-commerce;
secure brand assets for campaigns/launches; manage samples & gifts for e-commerce,
CRM and promotions; lead PR & influencer activations; coordinate PR/seeding/gifting
plans; support digital marketing campaigns; align content calendars with Social.

E-commerce Management: define online assortment strategy; ensure hero-product &
new-launch representation; secure assets/content; identify online exclusives &
pre-launch; monitor e-commerce performance (visibility, conversion, sales).

Reporting & Financial: monthly category/brand KPI reports; track brand budgets by
country/store; business reviews; accurate budget & cost tracking.

Requirements: 3-4 years in Merchandising, Category Management, Brand Management,
Buying or a related commercial role, preferably beauty, fragrance, FMCG or retail;
strong analytical skills; proven experience managing brands, assortments and
commercial partnerships; excellent communication, presentation and stakeholder
management; strong organization in a fast-paced environment; Advanced Microsoft
Excel and PowerPoint; retail planning/reporting tools an advantage.
"""

ATS = [
    "merchant", "merchandising", "category management", "category manager",
    "assortment management", "SKU management", "assortment planning",
    "product lifecycle", "planogram", "shelving", "stock planning",
    "inventory optimization", "sales forecasting", "category calendar",
    "brand management", "brand lifecycle", "brand partner", "commercial partnerships",
    "buying", "brand strategy", "market share", "brand equity", "business reviews",
    "fragrance", "fragrances", "beauty", "prestige beauty", "FMCG", "retail",
    "omnichannel", "activations", "product launches", "NPD", "PR", "influencer",
    "seeding", "gifting", "samples", "CRM", "digital campaigns", "social content",
    "e-commerce", "online assortment", "hero products", "online exclusives",
    "conversion", "KPI reporting", "budget tracking", "P&L", "data-driven",
    "Advanced Excel", "PowerPoint", "Nielsen", "stakeholder management", "Dubai", "GCC",
]

CONTENT = {
    "headline": (
        "Category & Brand Merchant · Fragrances & Beauty · Assortment, "
        "Merchandising & E-Commerce · Brand-Partner Management"
    ),
    "professional_summary": (
        "Beauty & fragrance category and brand professional with 4+ years in merchandising, assortment "
        "and brand management across beauty, fragrances, FMCG and e-commerce. As Key Account Manager for "
        "Beauty, Fragrances & Fashion at Alibaba's Miravia I managed 42 brand accounts (+30% GMV QoQ) and "
        "onboarded the leading Arabian & oud houses; today I lead brand and assortment across retail and "
        "e-commerce at DoFreeze. Advanced Excel/PowerPoint, data-driven, Dubai-based."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Multi-brand FMCG group | Brands: Befit, Eurocake, Flair | 50+ countries",
            "bullets": [
                "Manage the end-to-end lifecycle of a multi-brand portfolio (Befit, Eurocake, Flair) — brand strategy, 6 NPD launches, assortment and go-to-market across 50+ GCC/MENA and international markets",
                "Own assortment and listings across retail and e-commerce — a Shopify D2C store plus UAE quick-commerce (Noon, Talabat, Careem, Deliveroo) — managing SKUs, hero products, shelf/planogram presence and launch availability",
                "Run brand activations, influencer/PR and sampling & seeding, own the A&P budget, and report category and brand KPIs to leadership — with an AI (Claude/GPT) layer automating reporting",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Miravia — Alibaba's marketplace | Top-5 global e-commerce | 100K+ employees",
            "bullets": [
                "Managed 42 brand accounts across beauty, fragrances and fashion, driving +30% GMV QoQ through assortment optimization, pricing and promotional strategy — with continuous assortment reviews and performance analysis",
                "As PIC Fragrances, onboarded 30+ stores in two months — including the leading Arabian & oud houses (Arabian Oud, Lattafa, Swiss Arabian, Ajmal) — and ran the Flash Sales channel reporting to the CEO",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | Retail vertical build-out",
            "bullets": [
                "Built assortment and onboarded brand catalogues onto Glovo's new Retail vertical, coordinating cross-functionally to launch listings and grow order volume and GMV",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG multinational | €36B revenue | 90K+ employees",
            "bullets": [
                "Category planning: SKU-level sell-in/sell-out and inventory analysis, promotional-effectiveness reviews and sales forecasting, plus NPD support (Milka Spread, Mini Suchard)",
            ],
        },
    ],
    # Skills-block ROW LABELS relabelled for a category/merchant profile (SHORT so they don't wrap).
    "skills_brand": (  # label -> "Category & Merch."
        "fragrance & beauty category management, SKU-level assortment planning, planograms & shelving, "
        "product lifecycle, sales forecasting by country, inventory & stock optimization, category calendar"
    ),
    "skills_ecommerce": (  # label -> "E-Com & Activations"
        "online assortment strategy, hero-product & new-launch merchandising, brand activations (retail "
        "& e-com), PR / influencer / seeding & gifting, samples for CRM, digital campaign support, social alignment"
    ),
    "skills_commercial": (  # label -> "Brand & Commercial"
        "brand lifecycle management, brand-partner relationships & business reviews, commercial partnerships, "
        "pricing & promotions, negotiation, cross-functional stakeholder management"
    ),
    "skills_data": (  # label -> "Data & Reporting"
        "data-driven decisions, monthly category/brand KPI reports, budget tracking by country/store, "
        "forecasting, business presentations, Nielsen/Kantar, Power BI, Tableau, Looker"
    ),
    "skills_tools": (  # label -> "Tools" (unchanged)
        "Microsoft Excel (Advanced), PowerPoint (Advanced), SAP, Salesforce, Nielsen, Kantar, Power BI, "
        "Tableau, Looker, Shopify, Generative AI (Claude, ChatGPT)"
    ),
}

# Skills-block row-label overrides (keep SHORT — long labels wrap to 2 lines).
ROLE_LABELS = {
    "Brand & Marketing": "Category & Merch.",
    "E-Commerce & Digital": "E-Com & Activations",
    "Commercial": "Brand & Commercial",
    "Data & Analytics": "Data & Reporting",
}


def make_job() -> Job:
    return Job(
        id="sephora-merchant-fragrances-2026-08",
        title=TITLE,
        company=COMPANY,
        location="Dubai, United Arab Emirates",
        url="http://www.inside-sephora.com/",
        source="linkedin",
        description=JOB_DESCRIPTION,
        raw={
            "query": "Merchant Fragrances Sephora Dubai category merchandising brand",
            "via": "LinkedIn via TALENTMATE portal (role is Sephora / LVMH)",
            "note": "STRONG honest fit. Fragrance category/merchant role at Sephora — maps directly to Paula's "
                    "Miravia Beauty/Fragrances/Fashion KAM (42 accounts, +30% GMV, onboarded Arabian Oud/Lattafa/"
                    "Swiss Arabian/Ajmal) + DoFreeze brand/assortment/e-com + Mondelez category planning. 3-4 yrs "
                    "asked (Paula 4+); Advanced Excel/PPT; NO Arabic required; named employer, Dubai.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id,
        "title": job.title,
        "company": job.company,
        "location": job.location,
        "url": job.url,
        "source": job.source,
        "description": job.description,
        "salary_raw": None,
        "salary_aed_min": None,
        "salary_aed_max": None,
        "posted_date": DATE_FOLDER,
        "raw": job.raw,
        "ai_score": 91,
        "ai_tier": "Hot",
        "skills_match": [
            "HERO fit: KAM for Beauty, Fragrances & Fashion at Miravia (Alibaba) — 42 brand accounts, +30% GMV QoQ, assortment/pricing/promotions",
            "PIC Fragrances: onboarded the leading Arabian & oud houses (Arabian Oud, Lattafa, Swiss Arabian, Ajmal) — direct fragrance-category depth",
            "Brand lifecycle + assortment across retail & e-commerce (DoFreeze: Shopify + quick-commerce, SKUs, hero products, launches)",
            "Category planning at Mondelez: SKU sell-in/sell-out, inventory analysis, forecasting, NPD",
            "Activations, PR/influencer, sampling & seeding, budget & KPI reporting (DoFreeze)",
            "3-4 yrs asked → Paula 4+; Advanced Excel & PowerPoint; data-driven; strong stakeholder management",
            "Named prestige employer (Sephora/LVMH), Dubai; NO Arabic required",
        ],
        "missing_skills": [
            "No prior Sephora/retailer-buyer title specifically — but category/merchant substance (assortment, planograms, forecasting, brand-partner mgmt) is all there via Miravia + Mondelez",
            "Planogram/visual-merchandising grounded in Miravia assortment + Inditex/Massimo Dutti retail (honest, not overclaimed)",
        ],
        "sector_fit": "excellent (fragrance & beauty category/brand merchandising — her strongest domain via Miravia Beauty/Fragrances KAM)",
        "seniority_fit": "on-band (3-4 yrs asked; Paula 4+) — clean, no stretch",
        "red_flags": [
            "Posted via TALENTMATE aggregator ('respuestas gestionadas fuera de LinkedIn') — apply on Sephora careers directly where possible; verify the live req",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "One of Paula's strongest, most on-target fits and fully honest. The role is a fragrance-category "
            "merchant at Sephora (LVMH) covering assortment/SKU management, planograms, forecasting, brand "
            "lifecycle & brand-partner management, omnichannel activations (PR/influencer/seeding/samples), "
            "e-commerce assortment and KPI/budget reporting — asking 3-4 yrs in merchandising/category/brand/"
            "buying, preferably beauty/fragrance/FMCG, with Advanced Excel/PowerPoint and no Arabic requirement. "
            "Paula's Miravia role (KAM Beauty, Fragrances & Fashion — 42 accounts, +30% GMV QoQ, assortment/"
            "pricing/promotions, and onboarding the leading Arabian & oud houses as PIC Fragrances) maps almost "
            "one-to-one; DoFreeze adds brand lifecycle, assortment across retail + e-commerce, NPD, activations, "
            "sampling & seeding, budget & KPI reporting; Mondelez adds category planning (SKU sell-in/sell-out, "
            "inventory analysis, forecasting). She is 4+ yrs (asked 3-4), Advanced Excel/PPT, data-driven, and "
            "Dubai-based. Only note: posted via the TALENTMATE aggregator — apply on Sephora careers directly "
            "where possible. No 'no-sponsorship' claim; factual UAE Residence Visa only."
        ),
        "scored_by": "manual:claude",
        "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00",
        "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    """Relabel skills rows for this category/merchant CV (template stays pristine)."""
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    """Convert DOCX -> PDF headlessly via LibreOffice (docx2pdf/Word fails on this Mac)."""
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    """Move output/<Company> - <Role>/ under output/<DATE_FOLDER>/."""
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS — trim a bullet/skills")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
