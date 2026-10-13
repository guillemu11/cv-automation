"""One-off: CV for Delivery Hero "Specialist, Performance Marketing - instashop", Dubai, hybrid
(found via Jack & Jill, 2026-10-13; SmartRecruiters 744000140440854). CV only for now.

Individual contributor on instashop's performance team (dotted line to talabat's performance lead):
own paid campaigns on Meta and TikTok (Advantage+ named), budget pacing to CAC targets, full-funnel
analysis and reports for UAE + Egypt, incentive tracking, structured A/B tests, AI-first workflow
(Claude/ChatGPT automation explicitly asked). 3–5 yrs performance, 2+ on paid social; MMPs
(Adjust/AppsFlyer); q-commerce / food delivery a strong plus; incrementality vs last-touch.

Paula's honest angle:
- DoFreeze: Meta Advantage+ catalogue retargeting (3.7x first-week ROAS, AED 34.5 per purchase), Aug
  account 829K impressions / 2.01% CTR / AED 4.64 CPM, budget moved on break-even ROAS net of
  marketplace fees, incrementality vs a control brand (SMASH x talabat), weekly creative-hook tests,
  briefs a designer + social exec, Claude automation system (~40% less manual work).
- Glovo: 5 XL restaurant accounts on a food-delivery app: promo spend, in-app ads/visibility, choosing
  activation days from campaign returns, weekly refund-fraud review (memory glovo-real-role).
- Miravia & AliExpress: on-site visibility, Flash Sales, L'Oréal brands (memory miravia-scope-all-cvs).

Honesty guardrails:
- No TikTok Ads, Apple Ads, Snapchat or MMP (Adjust/AppsFlyer) claims: she has none. TikTok is organic only.
- No app-install / LTV claims. Paid budget is small (~AED 4–5K/month): don't quote it; quote efficiency.
- Glovo: junior XL Accounts role; never the Retail vertical or deal negotiation.
- NIVEA only with "three months" (left out here). No Arabic. DoFreeze is NOT on Deliveroo.
- Salary: Jack estimates AED 120–180K/yr (10–15K/month), under her current 20K/month — flagged.

ONE-PAGE standard. Lands under output/2026-10-13/Delivery Hero - Specialist, Performance Marketing - instashop/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Delivery Hero"
TITLE = "Specialist, Performance Marketing - instashop"
DATE_FOLDER = "2026-10-13"

JOB_DESCRIPTION = """\
Specialist, Performance Marketing — instashop (Delivery Hero), Dubai, hybrid. Individual contributor.
instashop is the leading online local marketplace in the UAE and Egypt (Delivery Hero since 2020).
Own day-to-day execution of paid campaigns across channels, manage budgets and drive customer acquisition.
Works with the Head of Marketing at instashop on Marcomms strategy, incentive tracking and business
performance analysis; dotted line to talabat's performance marketing lead. Expected to run lean by embedding
AI into the workflow.
Campaign execution: end-to-end setup, bidding and optimisation on Meta, TikTok and other paid channels; take
creatives live, test formats, iterate on data; daily budget allocation and pacing to CAC targets; monitor health.
Performance analysis: full-funnel analysis from impression to conversion; regular reports for the Head of
Marketing across UAE & Egypt; incentive performance; CAC efficiency recommendations.
Experimentation: A/B tests on creatives, audiences, bids and landing experiences; documented hypotheses.
Strategy: marcomms calendar, incentive design and tracking, quarterly planning.
Qualifications: Meta Ads Manager depth (Advantage+, bidding, creative); TikTok Ads; short-form paid vs organic;
Google, Apple Ads, Snapchat; attribution and MMPs (Adjust, AppsFlyer); AI-first (ChatGPT, Claude, prompt
engineering, workflow automation); incrementality vs last-touch; CAC, LTV, ROAS, CPM. 3–5 years hands-on
performance marketing, 2+ on paid social; budgets across channels and markets; Meta and TikTok track record;
MMP experience; q-commerce, food delivery or e-commerce a strong plus; Bachelor's in Marketing or Business.
Nice to have: brand + performance objectives together; briefing creative teams on short-form video; AI tools.
"""

ATS = [
    "performance marketing", "paid social", "Meta Ads Manager", "Advantage+", "TikTok Ads", "Google Ads",
    "bidding", "budget pacing", "CAC", "ROAS", "CPM", "CTR", "LTV", "incrementality", "last-touch attribution",
    "A/B testing", "creative testing", "full funnel", "customer acquisition", "incentives", "q-commerce",
    "food delivery", "AI-first", "Claude", "ChatGPT", "prompt engineering", "automation", "reporting",
    "short-form video", "Adjust", "AppsFlyer",
]

CV_CONTENT = {
    "headline": "Performance & Paid Social · Q-Commerce · AI-First",
    "professional_summary": (
        "Performance-minded marketer with 5 years across e-commerce and quick-commerce (Glovo, Alibaba's "
        "Miravia, DoFreeze). In Dubai, runs Meta and Google Ads for brands sold on talabat, Noon and Careem "
        "and measures campaigns on incremental sales against a control. AI-first: automates planning, "
        "content and reporting with Claude (~40% less manual work)."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE F&B / FMCG group | 7 brands incl. Befit, Eurocake, SMASH | sold on talabat, Noon, Careem",
            "bullets": [
                "Run paid social and search on Meta and Google Ads: built an Advantage+ catalogue retargeting campaign at 3.7x ROAS and AED 34.5 per purchase in week one; August account at 829K impressions, 2.01% CTR, AED 4.64 CPM",
                "Allocate budget across brands and channels on break-even ROAS net of marketplace commissions; moved Google spend to Meta when its return did not cover the cost",
                "Measure incrementality, not last touch: SMASH x talabat lifted daily sales +165% vs +4.7% for a control brand; Befit x Noon added 4,176 incremental units (+31% vs baseline)",
                "Run weekly creative-hook tests across 4 social accounts and brief a designer and a social media exec on short-form video; a SMASH giveaway added followers at AED 0.73 each, 8x cheaper than a promoted reel",
                "Built the influencer programme from zero: 103 creator activations over 3 waves at AED 81 per content piece; negotiated −30% and audited the agency's reach reporting",
                "Web orders +86% and revenue +57% (Q1 vs Jun–Aug 2026) after adding reviews, basket-based bundles and a 30-day replenishment flow; landing pages built with AI agents",
            ],
        },
        {
            "company": "Miravia & AliExpress (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Beauty & Health department",
            "bullets": [
                "Grew 42 beauty and fragrance brands (incl. L'Oréal Paris, Maybelline, NYX, Redken, KIKO Milano) +30% GMV QoQ through pricing, promotions and on-site visibility; worked with marketing, product and ops",
                "Owned Flash Sales (up to ~30% of B&H sales), reporting to the CEO; grew Arabian perfume into a top-growing sub-category",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Food & quick-commerce delivery app | 10K+ employees",
            "bullets": [
                "Co-managed 5 XL restaurant accounts (KFC, Taco Bell, Vicio) on GMV, orders, promo spend and in-app ads; read campaign returns to pick activation days; weekly refund-fraud reviews",
            ],
        },
        {
            "company": "Mondelez International · Línea Directa Aseguradora",
            "role": "Early Career",
            "dates": "2019 – 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | Spanish direct insurer",
            "bullets": [
                "Category Planning Trainee, Mondelez (Aug 2021 – Aug 2022): sell-in/sell-out and promo-effectiveness analysis, NPD support",
                "Accounting Intern, Línea Directa (Jan – Dec 2019): accounting department of an insurer, working in SAP and Looker",
            ],
        },
    ],
    "skills_brand": (
        "Meta Ads (Advantage+, catalogue retargeting), Google Ads, budget allocation, creative testing"
    ),
    "skills_ecommerce": (
        "talabat, Noon, Careem, Amazon, Glovo, Miravia, AliExpress, Shopify, CRO, email segmentation"
    ),
    "skills_commercial": (
        "incrementality vs control, break-even ROAS, ROAS, CPA, CPM, CTR, promo uplift, KPI reporting"
    ),
    "skills_data": (
        "Excel (Expert), Looker, Power BI, SAP, Meta Ads reporting, Shopify analytics"
    ),
    "skills_tools": (
        "Claude / Claude Code (workflow automation), ChatGPT, prompt engineering, Canva, Photoshop"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Paid Media",
    "E-Commerce & Digital": "Q-Commerce & E-Com",
    "Commercial": "Measurement",
    "Tools": "AI & Tools",
}


def make_job() -> Job:
    return Job(
        id="deliveryhero-instashop-performance-marketing-specialist-2026-10",
        title=TITLE,
        company=COMPANY,
        location="Dubai, United Arab Emirates",
        url="https://jobs.smartrecruiters.com/DeliveryHero/744000140440854-specialist-performance-marketing-instashop",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={"via": "Jack & Jill AI (tracked job, 2026-10-13)",
             "workplace": "Dubai, hybrid",
             "level": "Specialist (individual contributor)"},
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not posted (Jack & Jill estimate AED 120–180K/yr ≈ 10–15K/month)",
        "salary_aed_min": 10000, "salary_aed_max": 15000,
        "posted_date": None, "raw": job.raw,
        "ai_score": 62, "ai_tier": "Possible",
        "skills_match": [
            "Meta Advantage+ catalogue retargeting: 3.7x first-week ROAS, AED 34.5 per purchase; 2.01% CTR, AED 4.64 CPM",
            "Incrementality vs control (SMASH x talabat +165% vs +4.7%), budget moved on break-even ROAS",
            "AI-first: Claude system automating planning, content and reporting (~40% less manual work)",
            "Q-commerce / food delivery from both sides: Glovo XL accounts, sells via talabat, Noon, Careem",
            "Weekly creative-hook tests; briefs a designer + social exec on short-form video",
        ],
        "missing_skills": [
            "MMPs (Adjust / AppsFlyer) and app-install attribution",
            "TikTok Ads, Apple Ads, Snapchat (TikTok organic only)",
            "~1 year of hands-on paid media (3–5 asked) on a small budget",
        ],
        "sector_fit": "excellent (q-commerce marketplace; talabat dotted line)",
        "seniority_fit": "step down (Specialist IC; she is a Manager today)",
        "red_flags": [
            "Salary estimate AED 10–15K/month, below her current 20K/month",
            "No MMP or TikTok Ads experience, both listed as requirements",
        ],
        "ats_keywords": ATS,
        "reasoning": "Strong on Advantage+, incrementality, AI-first and q-commerce; gaps on MMP, TikTok Ads and years of paid; pay likely below her current salary.",
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _pages(pdf: Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(pdf)).pages)


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    final_dir = settings.output_dir / DATE_FOLDER / f"{COMPANY} - {TITLE}"
    dest = final_dir / "01_CV_y_Carta"
    dest.mkdir(parents=True, exist_ok=True)

    cv_docx = cv._fill_template(CV_CONTENT, job)
    _relabel_for_role(cv_docx)
    stray = cv_docx.parent.parent

    clean_docx = dest / "Paula De Francisco - CV.docx"
    shutil.move(str(cv_docx), str(clean_docx))
    # The template writes to output/<Company> - <Role>/ at the root; only ever remove that.
    if stray.exists() and stray.parent.resolve() == settings.output_dir.resolve():
        shutil.rmtree(stray, ignore_errors=True)

    cv_pdf = _to_pdf_soffice(clean_docx)
    # The dashboard only indexes CV_Paula_<Company>_<Title>.pdf; the clean name is the one to send.
    dashboard_pdf = dest / f"CV_Paula_{cv._sanitize_filename(COMPANY)}_{cv._sanitize_filename(TITLE)}.pdf"
    shutil.copy(str(cv_pdf), str(dashboard_pdf))
    print("OK", cv_pdf, "PAGES", _pages(cv_pdf), "| dashboard copy:", dashboard_pdf.name)


if __name__ == "__main__":
    main()
