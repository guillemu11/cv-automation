"""One-off: generate Paula's ONE-PAGE CV for **Partnerships Marketing Manager** (Partnerships
Growth Marketing Manager) at **Airalo** (world's first eSIM store; fully remote; UAE via Deel).
Reports to the Director of Growth Strategy & Operations.

JD ask: global partnerships growth strategy; identify/prioritise co-marketing, affiliate, platform,
loyalty/travel partnerships; partner GTM (joint value props, launch plans, success metrics);
end-to-end partner campaigns across owned/paid/partner channels; KPIs (incremental revenue, CAC,
conversion); scalable partner ops (briefs, timelines, tracking, post-mortems); cross-functional
with BD/Product/Legal/Finance/Data; path to people leadership. 5+ yrs. Nice: marketplaces/apps,
affiliate programmes, A/B testing, integrated paid+owned+earned, people leadership.

Paula's honest angle: platform co-marketing with talabat (in-app giveaway, seasonal samplings) and
Noon measured on incremental sell-out vs control group; built an affiliate programme from scratch
(commission, AI-built landing page); community partnerships (running/padel/yoga clubs); partner
side at Glovo (XL partners) and Miravia (42 brands, 30+ onboarded); team of two.

Honesty guardrails:
- 5 yrs = exactly the floor. No travel/eSIM/telecom claim. No CAC figure, no affiliate-network tool.
- "Global" not claimed: GCC + Spain. Forecast not owned. Factual "UAE Residence Visa" only.

ONE-PAGE standard. Lands under output/2026-09-13/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Airalo"
TITLE = "Partnerships Marketing Manager"
DATE_FOLDER = "2026-09-13"

JOB_DESCRIPTION = """\
Partnerships (Growth) Marketing Manager — Airalo (world's first eSIM store, 400+ people, 60+ countries),
fully remote from Spain, UK or UAE (UAE via Deel). Reports to the Director of Growth Strategy & Operations.

Do: build and execute a global partnerships growth marketing strategy aligned to Growth objectives and
priority markets; identify, evaluate and prioritise partnership opportunities (co-marketing,
channel/affiliate, ecosystem/platform, loyalty/travel); own partner GTM planning (joint value
propositions, campaign concepts, launch and distribution plans, success metrics); lead end-to-end partner
campaigns across owned, paid, partner channels and PR; define KPIs and measurement (incremental revenue,
CAC efficiency, conversion, activation, engagement, retention); build scalable processes (partner briefs,
timelines, approvals, assets, tracking, post-mortems); partner with Marketing, BD, Product, Brand/Comms,
Legal, Finance, Data; grow into people leadership.

Bring: 5+ years in growth marketing, partner marketing or partnerships with outcomes; hands-on marketing
partnerships driving measurable growth (not just awareness); strategy to execution; cross-functional
stakeholder management remote; analytical rigor with KPIs, tracking/attribution; excellent English.
Nice: travel, eSIM, telecom, fintech, marketplaces, subscription/mobile apps; affiliate networks or
partner programmes; A/B and offer testing; integrated paid + owned + earned campaigns with partners;
people leadership. UAE salary AED 342,400 - 513,600 / year.
"""

ATS = [
    "partnerships", "partner marketing", "growth marketing", "co-marketing", "affiliate programme",
    "platform partnerships", "loyalty", "go-to-market", "GTM", "joint value proposition",
    "launch plans", "partner campaigns", "integrated campaigns", "owned", "paid", "earned",
    "KPIs", "incremental revenue", "conversion", "measurement", "attribution", "A/B testing",
    "partner briefs", "post-mortems", "cross-functional", "stakeholder management",
    "marketplaces", "mobile apps", "people leadership", "remote",
]

CONTENT = {
    "headline": "Partnerships & Growth Marketing · Co-marketing · Affiliate · Marketplaces",
    "professional_summary": (
        "Growth and partnerships marketer with 5 years across marketplaces and delivery apps (Alibaba, Glovo) and FMCG, "
        "now building co-marketing and affiliate partnerships in Dubai that are measured on incremental sales, not awareness."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG snacking group (Befit, Eurocake, Smash) | DTC + delivery platforms | GCC",
            "bullets": [
                "Built co-marketing partnerships with talabat and Noon from joint proposition to launch: an in-app giveaway where each purchase earned an entry to win talabat cashback, plus seasonal samplings",
                "Measured partners on incrementality: SMASH x talabat lifted daily sales +165% vs +4.7% for a control brand; Befit x Noon added 4,176 incremental units (+31% over baseline)",
                "Launched an affiliate programme from scratch (commission on sales, AI-built landing page, multilingual brief) and community partnerships with running, padel and yoga clubs",
                "Run briefs, timelines, approvals and post-mortems across 4 agencies (-30% negotiated, reporting audited) with Sales and Finance; lead a team of two (designer + social media executive)",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's eCommerce marketplace | 100K+ employees",
            "bullets": [
                "Grew 42 brand partners +30% GMV QoQ with joint promotions, tracking conversion, traffic and retention; onboarded 30+ fragrance brands in two months",
                "Created the Beauty Club and Hot on Social partner programmes; owned the Flash Sales channel, reporting to the CEO",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Delivery & quick-commerce app | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Ran joint marketing activations with XL partners (KFC, Taco Bell, Sushi Shop) across marketing, ops and CX; negotiated deals and helped build the Retail vertical",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | Chocolate category (Milka, Suchard) | €36B revenue",
            "bullets": [
                "Built promotional-effectiveness and sell-in/sell-out reports with Nielsen for the chocolate category",
            ],
        },
    ],
    "skills_brand": (  # label -> "Partnerships"
        "co-marketing, platform partners, affiliate programme, communities, loyalty mechanics"
    ),
    "skills_ecommerce": (  # label -> "Growth & GTM"
        "partner GTM, launch plans, integrated paid + owned + earned, A/B tests, landing pages"
    ),
    "skills_commercial": (  # label -> "Stakeholders"
        "cross-functional leadership, negotiation, 4 agencies, team of two, CEO reporting"
    ),
    "skills_data": (  # label -> "Measurement"
        "KPIs, incremental sales, control groups, conversion, ROAS / ROI, post-mortems"
    ),
    "skills_tools": (  # label -> "Tools"
        "Meta & Google Ads, Shopify, Looker, Power BI, Excel, Canva, Claude (AI agents)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Partnerships",
    "E-Commerce & Digital": "Growth & GTM",
    "Commercial": "Stakeholders",
    "Data & Analytics": "Measurement",
}


def make_job() -> Job:
    return Job(
        id="airalo-partnerships-marketing-manager-2026-09",
        title=TITLE,
        company=COMPANY,
        location="UAE (Remote)",
        url="https://www.linkedin.com/jobs/search/?keywords=Airalo%20Partnerships%20Marketing%20Manager",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw="AED 342,400 - 513,600 / year",
        salary_aed_min=342400 / 12,
        salary_aed_max=513600 / 12,
        raw={
            "query": "Airalo Partnerships Marketing Manager",
            "note": "Good fit: platform co-marketing (talabat, Noon) with control-group incrementality, affiliate "
                    "programme built from scratch, marketplace partner side (Glovo, Miravia). Gaps: global scope, travel/eSIM, CAC.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": job.salary_raw,
        "salary_aed_min": job.salary_aed_min, "salary_aed_max": job.salary_aed_max,
        "posted_date": "2026-09-06", "raw": job.raw,
        "ai_score": 72, "ai_tier": "Hot",
        "skills_match": [
            "Platform co-marketing with talabat and Noon (in-app giveaway, samplings) measured on incremental sales",
            "Affiliate programme built from scratch (commission, AI-built landing page)",
            "Partner-side experience at marketplaces/apps: Glovo XL partners, Miravia 42 brands",
            "End-to-end ops across agencies, Sales and Finance; team of two",
        ],
        "missing_skills": [
            "Global, multi-market partnerships scope (hers is GCC + Spain)",
            "Travel / eSIM / telecom sector",
            "CAC-based measurement, affiliate networks and attribution tooling",
        ],
        "sector_fit": "adjacent — marketplaces/delivery apps + FMCG vs travel-tech",
        "seniority_fit": "good — 5 yrs meets the 5+ floor; Manager title today",
        "red_flags": ["100+ applicants", "Fully remote global role"],
        "ats_keywords": ATS,
        "reasoning": (
            "Strong partnership-marketing evidence rarely seen at her level: platform co-marketing with control-group "
            "incrementality and a self-built affiliate programme, plus marketplace partner management. Gaps are global "
            "scope and travel/eSIM; no Arabic needed (English-only role). Salary well above floor."
        ),
        "scored_by": "manual:claude", "freshness": "recent",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
