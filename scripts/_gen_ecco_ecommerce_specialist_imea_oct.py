"""One-off: **versión de octubre** del CV para **E-Commerce Specialist – IMEA** en **ECCO** (Dubái, presencial,
LinkedIn Easy Apply, publicado por Menna Elhogaraty, Head of People & Culture MEA). Misma oferta que
`_gen_ecco_ecommerce_specialist_imea.py` (2026-09-23), que se reutiliza para el JD, los ATS y el job; a 2026-10-06
lleva 1.638 solicitudes y Paula aún no ha aplicado.

Por qué rehacerlo en vez de mandar el del 23-sep:
- Se generó durante el bug de la foto recortada (2026-08-27 → 2026-10-01).
- Cabecera con el visado antiguo; ahora "UAE Employment Visa (employer-sponsored)".
- Decía "across the GCC": solo hay evidencia de EAU.
- Le faltaban los datos verificados después: SMASH × talabat +165% vs +4,7% del control, Befit × Noon +4.176 uds
  (+31%), la revisión semanal de riesgo de stock por dark store (el JD pide "stock availability") y el equipo de dos.

Ángulo honesto (igual que la v1): operaciones diarias en Noon, talabat y Careem + Shopify de punta a punta (es
literalmente el JD); dos años del lado marketplace en Miravia con 42 cuentas de beauty, fragancias y moda, y el
canal de Flash Sales = online trading con calendario; moda y lifestyle reales (cuentas de moda en Miravia, vertical
Retail de Glovo, Massimo Dutti).

Guardarraíles: sin footwear; sin CMS enterprise (solo Shopify + marketplaces); sin sell-through por tallas; solo EAU,
nada de GCC/KSA; forecast "contribute"; DoFreeze NO está en Deliveroo; sin árabe (ventaja aquí); visa según
profile.yaml.
"""
from __future__ import annotations

import importlib.util
import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.generators import cv_generator as cv

spec = importlib.util.spec_from_file_location(
    "_gen_ecco_ecommerce_specialist_imea_v1",
    ROOT / "scripts" / "_gen_ecco_ecommerce_specialist_imea.py",
)
v1 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(v1)

DATE_FOLDER = "2026-10-06"

CONTENT = {
    "headline": "E-Commerce Operations & Marketplaces · Online Trading · Fashion, Beauty & FMCG · UAE",
    "professional_summary": (
        "E-commerce professional with 5 years across marketplaces, fashion, beauty and FMCG. In Dubai, runs day-to-day "
        "e-commerce for three brands on Noon, talabat, Careem and Shopify; before, managed 42 beauty, fragrance and "
        "fashion accounts at Miravia (Alibaba) at +30% GMV QoQ and owned its Flash Sales trading calendar."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG group (Befit, Eurocake, Flair) | Noon, talabat & Careem + Shopify D2C | 50+ markets",
            "bullets": [
                "Run day-to-day operations on Noon, talabat and Careem — listings and content, pricing, promotional mechanics and the partner relationship on each platform, including joint activations (purchase-linked talabat cashback giveaway, seasonal sampling)",
                "Own the Shopify store end-to-end — catalogue, product content and imagery, collections, discounts and checkout — merchandised on conversion and AOV",
                "Launches and seasonal campaigns read on results: SMASH launch on talabat lifted daily sales +165% vs +4.7% for a control brand; Befit × Noon \"New Year, New Me\" added +4,176 incremental units, +31% over baseline",
                "Run the weekly stock-risk review by dark store with a dashboard built with AI, checking PO coverage against demand; contribute to the monthly sell-in forecast with e-commerce and finance; lead a team of two",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees | platform side",
            "bullets": [
                "Managed 42 beauty, fragrance and fashion accounts on the marketplace — assortment, pricing, product content and promotional calendars — delivering +30% GMV QoQ; onboarded 30+ new stores in two months as fragrance lead",
                "Owned the Flash Sales channel for Beauty, Fashion & Home reporting to the CEO: trading plans against P&L targets, reading traffic, conversion, ROI and ROAS to adjust promotions and investment",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce & delivery platform | €500M+ revenue | platform side",
            "bullets": [
                "Helped build the Retail vertical, onboarding fashion and lifestyle brands to the platform; grew XL accounts (KFC, Taco Bell, Sushi Shop) through catalogue work and in-app promotional mechanics",
            ],
        },
        {
            "company": "Massimo Dutti (Inditex) · Mondelez International",
            "role": "Sales Associate · Trainee, Category Planning",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Premium fashion retail | Global FMCG (chocolate category)",
            "bullets": [
                "At Inditex, fashion retail operations, visual merchandising standards and product flow; at Mondelez, sell-in/sell-out and promotional-effectiveness analysis with Nielsen data",
            ],
        },
    ],
    "skills_brand": (  # -> "E-Commerce Ops"
        "listings & product content, imagery, online merchandising, pricing & promotions, stock availability"
    ),
    "skills_ecommerce": (  # -> "Platforms"
        "Shopify (owner) · Noon, talabat, Careem (brand side) · Miravia, Glovo (platform side)"
    ),
    "skills_commercial": (  # -> "Trading"
        "seasonal & promo calendar, launches, assortment, partner management, P&L targets"
    ),
    "skills_data": (  # -> "Analytics & KPIs"
        "conversion, traffic, AOV, sell-out, GMV, ROI, ROAS, control-group reads, dashboards"
    ),
    "skills_tools": (  # -> "Tools"
        "Excel (Advanced), Power BI, Tableau, Looker, Shopify, Meta & Google Ads, SAP, Claude (AI)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "E-Commerce Ops",
    "E-Commerce & Digital": "Platforms",
    "Commercial": "Trading",
    "Data & Analytics": "Analytics & KPIs",
}


def refresh_dashboard_entry() -> None:
    """Keep the single v1 entry, corrected to the October facts (UAE only, applicant count)."""
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    job_id = v1.make_job().id
    entry = next((j for j in jobs if j.get("id") == job_id), None)
    if entry is None:
        v1.register_in_dashboard(v1.make_job())
        return refresh_dashboard_entry()
    entry["skills_match"] = [
        "Operaciones diarias en Noon, talabat y Careem: listings, contenido, precio, promociones y relación con la plataforma",
        "Shopify de punta a punta: catálogo, contenido e imágenes, colecciones, descuentos y checkout",
        "Lanzamientos y campañas con resultado medido: SMASH × talabat +165% vs control, Befit × Noon +31%",
        "Disponibilidad: revisión semanal de riesgo de stock por dark store",
        "Lado marketplace en Miravia (Alibaba): 42 cuentas de beauty, fragancias y moda, +30% GMV QoQ",
        "Online trading: canal Flash Sales de Miravia contra objetivos de P&L",
        "Moda y lifestyle: cuentas de moda en Miravia, vertical Retail de Glovo, Massimo Dutti",
        "3+ años pedidos; tiene 5",
    ]
    entry["red_flags"] = [
        "1.638 solicitudes a 2026-10-06: el Easy Apply solo se pierde; escribir a Menna Elhogaraty",
        "Título Specialist frente a su Manager actual — posible paso atrás de banda y de scope",
        "Presencial, no híbrido",
        "ECCO lleva dos años de pérdidas según la ficha de LinkedIn — preguntar por la salud del negocio en MEA",
    ]
    entry["raw"] = {**entry.get("raw", {}), "cv_version": f"{DATE_FOLDER} (v2: foto, visado y datos actualizados)"}
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = v1.make_job()
    refresh_dashboard_entry()

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
