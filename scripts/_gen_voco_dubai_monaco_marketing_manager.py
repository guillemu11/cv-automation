"""One-off: CV de una página para **Marketing Manager** en **voco Dubai Monaco** ("The Heart of Europe",
World Islands; hotel IHG para adultos con beach club, piscina, nightlife, eventos y F&B). Dubái,
presencial, jornada completa, LinkedIn Easy Apply, promocionado por técnico de selección, publicado
hace 3 días. **823 solicitudes** (121 desde ayer). Sueldo no publicado. Paula la tenía guardada.

Pide el JD: estrategia de marketing y calendario anual del hotel y la isla; posicionamiento frente a
los hoteles tradicionales de Dubái; campañas comerciales para habitaciones, F&B, pool & beach,
nightlife, eventos y alquiler de espacios; staycations y demanda de domingo a jueves; marketing ligado
a ingresos y ROI; digital, social y contenido (foto, vídeo, reels); influencers, creators, PR y medios;
festivales, fiestas y activaciones de F&B con Eventos y Operaciones; partnerships con marcas de
hospitality, lifestyle y consumo; CRM, base de datos y fidelización para repetición y reserva directa;
presupuesto, agencias y proveedores creativos; informes de ROI, footfall, reservas y captación.
Requisitos: 5–7+ años de marketing, preferiblemente hospitality, lifestyle, F&B, entretenimiento o lujo;
grado en Marketing/ADE; red en medios, influencers y lifestyle de Dubái como ventaja. **Pregunta de
filtro del anunciante: más de 4 años de experiencia en Hospitality.**

Ángulo honesto de Paula:
- Calendario de campañas por temporada (Ramadán, back to school, Fitness Month, New Year) para un
  grupo F&B homegrown de los EAU con varias marcas (Befit, Eurocake con SMASH, Flair).
- Creators y comunidades de Dubái desde cero: 103 activaciones en 3 campañas (vía alist, negociando -30%:
  75 contratadas, 103 entregadas); comunidades de running, pádel y yoga. Las otras agencias (Yamammi,
  Amplify, Terrier) están en curso o evaluadas, así que no se escribe "con 4 agencias".
- Partnerships y activaciones: sorteo in-app con talabat (cada compra, una participación para tarjetas
  de cashback) y samplings de temporada con talabat.
- Ingresos y ROI con grupo de control: SMASH × talabat +165% de venta diaria frente a +4,7% del control
  (~AED 32.000 incrementales sobre ~AED 8.700 de coste); Befit × Noon +31% sobre baseline.
- Contenido: dirige a un diseñador y a una social media executive; escribe los briefs y aprueba todo
  el output de foto y vídeo.
- Fidelización: Beauty Club y Hot on Social en Miravia.
- Restauración: cuentas XL de restaurante en Glovo (KFC, Taco Bell, La Tagliatella, Sushi Shop) con
  activaciones de marketing conjuntas, coordinando marketing, operaciones y atención al cliente.

Guardarraíles de honestidad:
- **Hospitality: CERO años.** Ni hotel, ni resort, ni beach club, ni nightlife. No se escribe
  "hospitality" como experiencia suya; solo aparece como el sector al que quiere ir. La pregunta de
  filtro (4+ años en Hospitality) hay que contestarla con la verdad, y probablemente la corte.
- **Años de marketing**: piden 5–7+; ella tiene ~5 de carrera en total. No se escribe ningún número.
- **PR y medios**: no ha llevado prensa ni relación con periodistas. Solo creators y agencias de
  influencers. No se escribe "PR".
- **Revenue management, ocupación, reservas, venue hire, staycations**: sin experiencia; no aparecen.
- **Rodajes**: no hay evidencia de que dirija shoots en set; briefa y aprueba foto y vídeo.
- **CRM de hotel / reserva directa**: no. El paralelo es la fidelización de Miravia y el Shopify D2C.
- **Forecast**: la dueña es la e-commerce manager; no aparece.
- **Deliveroo**: DoFreeze NO está en Deliveroo. Plataformas: talabat, Noon, Careem.
- Creators: "103 activaciones en 3 campañas", nunca "una campaña con 103".
- El JD no pide árabe. Visa: "UAE Employment Visa (employer-sponsored)"; nunca "no sponsorship".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "voco Dubai Monaco"
TITLE = "Marketing Manager"
DATE_FOLDER = "2026-10-02"

JOB_DESCRIPTION = """\
Marketing Manager - voco Dubai Monaco. voco Dubai Monaco - The Heart of Europe. Dubai, UAE. On-site, full time.
We are seeking a dynamic, commercially driven Marketing Manager to lead the marketing strategy for voco Dubai Monaco,
a unique island lifestyle destination on Dubai's World Islands. The role will be responsible for positioning voco
Monaco as a distinctive destination for hotel stays, staycations, F&B, pool and beach experiences, nightlife,
entertainment, events and activations. The ideal candidate will combine strong strategic thinking with hands-on
execution and a deep understanding of the Dubai hospitality and lifestyle market. The successful candidate will work
closely with Sales, Hospitality Operations, Real Estate, Commercial, F&B, Revenue, Engineering, and external partners
to deliver integrated marketing campaigns that drive revenue, occupancy, lead generation, brand awareness, and
customer engagement.
Marketing Strategy & Brand Positioning: develop and execute the overall marketing strategy and annual marketing
calendar for the hotel and island destination; build strong market positioning and brand awareness, differentiating
voco Monaco from traditional Dubai hotels and lifestyle destinations; identify market trends, competitor activity and
new opportunities to increase visitation and revenue.
Commercial & Revenue Marketing: develop commercially focused campaigns to drive room, F&B, pool & beach, nightlife,
events, venue hire and overall island revenue; create targeted campaigns to increase local staycation demand and
Sunday-Thursday business; connect marketing activity to measurable commercial outcomes, revenue contribution and ROI.
Digital, Social Media & Content: lead digital marketing, social media, content and the property's online presence
across relevant platforms; develop compelling photography, video, reels and other campaign assets; drive digital
engagement, customer acquisition and conversion.
PR, Influencer & Media: manage influencer, creator, PR and media relationships to generate awareness, engagement and
measurable commercial returns; build strong visibility across Dubai's lifestyle, hospitality, entertainment and
consumer markets.
Events, Entertainment & Activations: develop and execute marketing campaigns for festivals, parties, entertainment,
F&B activations and seasonal events; work closely with Events, F&B, Entertainment and Operations to maximise
attendance, footfall and commercial performance.
Partnerships & Customer Engagement: develop strategic partnerships and collaborations with hospitality, lifestyle,
entertainment, corporate and consumer brands; support CRM, database and loyalty initiatives to drive repeat visitation
and direct bookings.
Stakeholder & Budget Management: partner closely with Sales, Revenue, F&B, Events, Entertainment and Operations to
align marketing activity with business objectives; manage the marketing budget, agencies, creative partners and
suppliers; ensure all marketing communications and guest-facing collateral are aligned with brand standards.
Performance & ROI: track and report marketing performance across campaign ROI, revenue contribution, footfall,
bookings, engagement, database growth and customer acquisition; use insights and performance data to continuously
optimise marketing investment and activity.
Qualifications & Experience: 5-7+ years of marketing experience, preferably within hospitality, lifestyle, F&B,
entertainment or luxury. Hospitality marketing experience, with exposure to real estate, mixed-use developments,
lifestyle brands, or destination marketing considered an advantage. Strong understanding of the Dubai hospitality and
consumer market. Bachelor's degree in Marketing, Business Administration, Communications, or a related field. Proven
experience across digital marketing, social media, content, PR and influencer marketing. Experience marketing hotels,
resorts, beach clubs, restaurants, nightlife or lifestyle destinations is highly desirable. Strong commercial mindset
with the ability to link marketing activity to measurable revenue and ROI. Creative, energetic and highly hands-on,
with strong project management and communication skills. Experience managing agencies, suppliers, budgets and
external creative partners. A strong network across Dubai's media, influencer, lifestyle, entertainment and
hospitality sectors would be an advantage. Excellent project management, organizational, communication, and
stakeholder management skills. Hands-on, proactive, execution-oriented, and able to work in a fast-paced environment.
Screening requirement added by the poster: 4+ years of experience in Hospitality.
"""

ATS = [
    "marketing manager", "marketing strategy", "annual marketing calendar", "brand positioning", "brand awareness",
    "market trends", "competitor activity", "commercial", "revenue", "ROI", "revenue contribution", "campaigns",
    "integrated marketing campaigns", "seasonal campaigns", "seasonal events", "activations", "F&B", "F&B activations",
    "lifestyle", "Dubai", "UAE", "digital marketing", "social media", "content", "photography", "video", "reels",
    "campaign assets", "customer acquisition", "conversion", "engagement", "influencer marketing", "creators",
    "influencer", "partnerships", "brand collaborations", "consumer brands", "CRM", "loyalty", "repeat visitation",
    "community", "events", "footfall", "agencies", "agency management", "creative partners", "suppliers",
    "marketing budget", "brand standards", "stakeholder management", "Sales", "Operations", "performance",
    "KPI reporting", "Meta Ads", "Google Ads", "hands-on", "project management", "fast-paced", "restaurants",
    "Business Administration",
]

CONTENT = {
    "headline": "Marketing Manager · F&B & Lifestyle Brands · Social, Creators & Activations · Revenue-led Campaigns",
    "professional_summary": (
        "Dubai-based marketing manager leading brand, social, creator and partnership marketing for a homegrown UAE "
        "F&B group, with every campaign tied to sales. Built creator and community activations from zero in "
        "Dubai. Earlier ran restaurant accounts at Glovo and lifestyle brands at Alibaba."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Homegrown UAE F&B / FMCG group | Befit, Eurocake, SMASH, Flair | GCC + 50+ markets",
            "bullets": [
                "Plan and run a year-round seasonal campaign calendar (Ramadan, back to school, Fitness Month, New Year) for a multi-brand F&B portfolio across social, talabat, Noon, Careem, retail and our Shopify store",
                "Lead an in-house designer and social media executive: write the creative briefs, set the art direction and approve all social, photo and video content",
                "Built creator and community marketing from zero: 103 creator activations over 3 campaigns (agency fees negotiated -30%), a talabat in-app sweepstakes, seasonal samplings and running, padel and yoga community activations",
                "Measure every campaign on revenue: SMASH × talabat lifted daily sales +165% vs +4.7% for the control brand (~AED 32K incremental on ~AED 8.7K spend); Befit × Noon +31% over baseline; run Meta & Google Ads on ROI and ROAS",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | beauty, fashion & lifestyle",
            "bullets": [
                "Managed 42 beauty, fragrance and fashion brands (incl. KIKO Milano) on pricing and promotional calendars, delivering +30% GMV QoQ; owned the Flash Sales channel, reporting to the CEO",
                "Created the Beauty Club and Hot on Social projects to build customer loyalty and position Miravia as a beauty and lifestyle destination",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Food delivery & quick-commerce leader | €500M+ revenue",
            "bullets": [
                "Managed XL restaurant accounts (KFC, Taco Bell, La Tagliatella, Sushi Shop) on GMV and joint marketing activations, coordinating marketing, operations and customer support",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | chocolate & confectionery (Milka, Suchard)",
            "bullets": [
                "Ran sell-in/sell-out and promotional-effectiveness analysis with Nielsen and supported NPD launches (Milka Spread, Mini Suchard)",
            ],
        },
    ],
    "skills_brand": (  # -> "Strategy & Brand"
        "marketing strategy, campaign calendar, brand positioning, trend & competitor research, launches"
    ),
    "skills_ecommerce": (  # -> "Digital & Content"
        "social media, content & art direction, photo & video briefs, Meta & Google Ads, Shopify, EDM"
    ),
    "skills_commercial": (  # -> "Creators & Partners"
        "influencer & creator marketing, agency management, brand partnerships, community activations"
    ),
    "skills_data": (  # -> "Performance"
        "revenue-linked campaigns, control groups, ROI & ROAS, budget management, KPI reporting"
    ),
    "skills_tools": (  # -> "Tools"
        "Adobe (Photoshop, Illustrator), Canva, Meta Business Suite, Shopify, Power BI, Excel, Claude (AI)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Strategy & Brand",
    "E-Commerce & Digital": "Digital & Content",
    "Commercial": "Creators & Partners",
    "Data & Analytics": "Performance",
}


def make_job() -> Job:
    return Job(
        id="voco-dubai-monaco-marketing-manager-2026-10",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE (On-site) — World Islands",
        url="https://www.linkedin.com/jobs/search/?keywords=voco%20Dubai%20Monaco%20Marketing%20Manager",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "voco Dubai Monaco Marketing Manager",
            "note": "Hotel voco (IHG) para adultos en World Islands: beach club, piscina, nightlife, eventos y F&B. "
                    "LinkedIn Easy Apply, 823 solicitudes, publicado hace 3 días. El anunciante añade como filtro "
                    "4+ años en Hospitality: Paula tiene cero, y hay que contestarlo con la verdad. Encaje bueno "
                    "en lo que hace hoy: calendario de campañas, creators, activaciones, partnerships y campañas "
                    "medidas en venta.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": "2026-09-29", "raw": job.raw,
        "ai_score": 45, "ai_tier": "Warm",
        "skills_match": [
            "Calendario de campañas por temporada (Ramadán, back to school, Fitness Month, New Year)",
            "Creators y comunidades de Dubái desde cero: 103 activaciones en 3 campañas",
            "Partnerships y activaciones: sorteo in-app y samplings con talabat; running, pádel y yoga",
            "Campañas medidas en ingresos con grupo de control: SMASH × talabat +165% vs +4,7%",
            "Gestión de agencias y presupuesto: negoció -30% con las agencias de creators",
            "Dirige a un diseñador y a una social media executive; aprueba foto y vídeo",
            "Restauración desde la plataforma: cuentas XL de restaurante en Glovo",
            "Grado en ADE (CUNEF)",
        ],
        "missing_skills": [
            "Hospitality: cero años; el anunciante filtra con 4+ años",
            "Marketing de hotel, resort, beach club o nightlife: no",
            "PR y relación con medios: no ha llevado prensa",
            "Revenue de hotel (ocupación, reservas, staycations, venue hire): no",
            "Años de marketing: piden 5–7+; tiene ~5 de carrera en total",
        ],
        "sector_fit": "parcial — F&B y lifestyle sí, hospitality no",
        "seniority_fit": "en el límite — Manager con equipo de dos, pero menos años de los pedidos",
        "red_flags": [
            "La pregunta de filtro (4+ años en Hospitality) puede descartarla automáticamente",
            "823 solicitudes en 3 días",
            "Isla de fiesta para adultos en World Islands: desplazamiento en barco y horarios de nightlife",
            "Sueldo no publicado: verificar pronto la banda (suelo 20.000 AED/mes)",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Lo que el puesto hace día a día encaja con lo que Paula ya hace en DoFreeze: un calendario de "
            "campañas por temporada, creators y comunidades de Dubái montados desde cero, partnerships y "
            "activaciones con talabat, un equipo de contenido propio y campañas medidas en ventas con grupo de "
            "control. El problema es el sector: no tiene ni un año en hospitality y el anunciante ha puesto 4+ "
            "años como pregunta de filtro, así que lo más probable es que el sistema la descarte. Tampoco ha "
            "llevado PR ni revenue de hotel. Merece la pena aplicar porque es Easy Apply y cuesta poco, pero la "
            "vía con opciones es un mensaje directo a quien contrata."
        ),
        "scored_by": "manual:claude", "freshness": "recent",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
