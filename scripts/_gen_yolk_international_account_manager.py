"""One-off: CV de una página para **International Account Manager** en **Yolk Brands** (grupo F&B
de Dubái, 201-500 empleados; marcas Pickl, BonBird, SouthPour). Dubái, presencial, jornada completa,
LinkedIn Easy Apply, promocionado por técnico de selección. **698 solicitudes**. Sueldo no publicado.

Qué es el puesto: el enlace entre el franquiciador (equipos de Emiratos) y los **franquiciados
internacionales**. Pide el JD: estrategia y KPIs por región, relación con franquiciados, estándares
de marca y calidad, informes semanales/mensuales/trimestrales a dirección, cumplimiento de contratos
de franquicia, crecimiento (nuevas tiendas, mercados); KPIs financieros (ventas, royalties, COGS,
labor), presupuestos, pricing y margen; cadena de suministro y proveedores aprobados, inventario
(evitar mermas y roturas), lanzamientos de menú y LTOs con stock y formación; apoyo a marketing local,
campañas regionales, reseñas online y aperturas; SOPs, auditorías y visitas a tienda, plataformas
tecnológicas, coaching; y el sistema de inventario (Edify): datos de SKUs, recetas, proveedores y
precios, soporte de primer nivel y formación de usuarios.

Ángulo honesto de Paula:
- **F&B y QSR de verdad**: DoFreeze es un grupo F&B/FMCG de Emiratos que vende en 50+ mercados; en
  Glovo llevó cuentas XL de cadenas QSR multi-local (KFC, Taco Bell, La Tagliatella, Sushi Shop).
- **Relación con partners en varios mercados**: 6 lanzamientos en GCC, MENA, Asia, Europa, EE. UU. y
  África con ventas y distribuidores; 42 cuentas como punto de contacto único en Miravia (+30% GMV QoQ).
- **Inventario y roturas**: hace ella la revisión semanal de riesgo de stock por dark store en Noon,
  talabat y Careem (cobertura de PO frente a la demanda). Aporta al forecast mensual de sell-in.
- **Informes a dirección**: canal Flash Sales de Miravia reportando al CEO, contra objetivos de P&L.
- **Sistemas y coaching**: montó el sistema de automatización con IA (~40% menos trabajo manual);
  dirige a dos personas (diseñador y social media executive).
- **Estándares de marca en tienda**: Massimo Dutti (Inditex), visual merchandising.

Guardarraíles de honestidad:
- **Franquicias**: NO ha gestionado una red franquiciador–franquiciado. Lo más cercano son
  distribuidores y cuentas clave. No se escribe "franchise" como experiencia suya.
- **Operación de restaurante** (COGS, labor, recetas, SOPs, auditorías, visitas a tienda): sin
  evidencia; no se menciona. **Royalties**: no. **Edify** u otro software de inventario: no.
- **Forecast**: "contribute", nunca "own". La revisión semanal de stock sí es suya.
- **Distribuidores**: trabaja con ellos en lanzamientos y trade plans; no es su dueña comercial.
- **Deliveroo**: DoFreeze NO está. Plataformas: talabat, Noon, Careem.
- El JD no pide árabe. Visa: "UAE Employment Visa (employer-sponsored)".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Yolk Brands"
TITLE = "International Account Manager"
DATE_FOLDER = "2026-10-02"

JOB_DESCRIPTION = """\
International Account Manager. Yolk Brands (Pickl | BonBird | SouthPour). Dubai, UAE. On-site, full time.
The International Account Manager is responsible for overseeing franchise operations across international markets,
ensuring brand consistency, operational excellence, financial performance, and strong franchisee relationships.
Serving as a seamless communication funnel between UAE teams (franchisor), international franchisees, and all
internal departments, this role translates strategic initiatives into local execution. The ideal candidate brings
exceptional communication, organizational, planning, and cultural sensitivity skills to provide ongoing coaching,
operational support, and performance analysis that drive global growth, compliance, and service excellence.
Strategic & Regional Oversight: develop and execute strategies aligned with corporate goals and brand standards;
translate strategic initiatives into actionable plans across franchise locations; monitor regional KPIs, analyse
performance data, and identify trends or underperformance; maintain strong franchisee relationships and provide
ongoing support; ensure consistent brand and quality standards across all locations; provide regular performance
reports and updates to corporate leadership; ensure compliance with legal, regulatory, and franchise agreements;
identify opportunities for growth, including new stores, market expansion, and franchisee development.
Reporting & Business Reviews: deliver weekly, monthly, and quarterly performance reports to executive leadership;
regular contact with franchise partners to review performance and align on actions.
Financial Management & Profitability: track financial KPIs (sales, royalties, COGS, labor, etc) for each franchise
and region; support franchisees in budgeting, financial planning, and cost control; benchmark financial performance
and propose improvements for profitability; collect and consolidate regional financial data for corporate review;
advise on pricing, margin optimization, and operational efficiencies.
Supply Chain & Product Quality: coordinate with the corporate supply chain and ensure franchisees use approved
suppliers; guide inventory management practices to reduce waste and prevent stockouts; address vendor performance
and product quality issues; work with the internal team and franchise partners to ensure seamless rollout of new
menu items or LTOs, including stock readiness and training.
Marketing & Brand Presence: support execution of local marketing plans aligned with brand standards; collaborate on
regional campaigns and evaluate marketing effectiveness; monitor online reviews and assist franchisees in managing
their digital reputation; provide opening support for new franchises and grand opening campaigns.
Operational Excellence & System Building: monitor compliance with SOPs and identify areas for improvement; conduct
audits and site visits with action plans for improvement; support the implementation of technology platforms across
locations; encourage innovation and share best practices among franchisees; provide coaching on leadership, business
management, and succession planning; assist with new store development, site readiness, and launch success.
Inventory Management: ensure proper usage of our designated inventory system across the region with our partners;
ensure data accuracy (SKUs, recipes, suppliers, pricing) and system configuration; provide first-level support and
escalate technical issues to Edify; train users on Yolk Brands approved inventory software, best practices and
ensure SOPs for inventory are followed; ensure integration with other platforms and troubleshoot issues.
"""

ATS = [
    "international account manager", "account management", "international markets", "franchise partners",
    "franchisee relationships", "partner management", "multi-market", "regional KPIs", "performance analysis",
    "performance reports", "business reviews", "executive leadership", "brand standards", "brand consistency",
    "quality standards", "market expansion", "growth opportunities", "financial KPIs", "sales", "budgeting",
    "cost control", "pricing", "margin", "profitability", "benchmarking", "supply chain", "inventory management",
    "stockouts", "stock readiness", "new product rollout", "LTO", "product launches", "data accuracy", "SKUs",
    "local marketing plans", "regional campaigns", "marketing effectiveness", "digital", "technology platforms",
    "coaching", "best practices", "cross-functional", "communication", "cultural sensitivity",
    "F&B", "QSR", "food and beverage", "restaurant brands", "Excel", "Power BI", "UAE", "GCC", "Dubai",
]

CONTENT = {
    "headline": "International Account Management · Multi-Market Partner Growth · F&B & QSR · KPI Reporting",
    "professional_summary": (
        "Account and commercial professional in Dubai, 4+ years across F&B, FMCG and e-commerce. Brand & "
        "Marketing Manager at a homegrown UAE F&B group sold in 50+ markets; before that, XL QSR accounts at "
        "Glovo (KFC, Taco Bell, La Tagliatella) and 42 partner accounts at Miravia (Alibaba), +30% GMV QoQ."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Homegrown UAE F&B / FMCG group (Befit, Eurocake, Flair) | sold in 50+ markets",
            "bullets": [
                "Lead 6 product launches end-to-end across GCC, MENA, Asia, Europe, USA and Africa — brief, packaging, pricing and go-to-market — aligning sales, distributors and platforms so every market launches to the same brand standard",
                "Run the weekly stock-risk review by dark store on Noon, talabat and Careem, checking purchase-order coverage against current and upcoming demand to prevent stock-outs; contribute to the monthly sell-in forecast",
                "Plan seasonal and local campaigns by channel (Ramadan, back to school, New Year) with 25–50 creators per campaign and paid social, and measure them on sell-out, ROI and ROAS: Befit × Noon delivered +31% over baseline",
                "Built an AI-powered reporting and automation system that cut manual workload ~40%; lead and coach a team of two (designer and social media executive)",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees",
            "bullets": [
                "Single point of contact for 42 partner brands, reviewing performance with each on assortment, pricing and promotions to deliver +30% GMV QoQ; onboarded 30+ fragrance stores in two months",
                "Owned the Flash Sales channel for Beauty, Fashion & Home reporting directly to the CEO — commercial plans against P&L targets, tracking conversion, traffic, ROI and retention to reallocate investment",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce platform | €500M+ revenue",
            "bullets": [
                "Managed multi-site QSR chains (KFC, Taco Bell, La Tagliatella, Sushi Shop) on GMV, pricing and promotional mechanics, coordinating marketing, logistics and customer support; negotiated terms that kept both sides profitable",
                "Helped build out the Retail vertical, onboarding new brands and setting up their listings and order flow",
            ],
        },
        {
            "company": "Mondelez International · Massimo Dutti (Inditex)",
            "role": "Trainee, Category Planning · Sales Associate",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG, Nielsen | Premium fashion retail",
            "bullets": [
                "At Mondelez, sell-in/sell-out and promotional-effectiveness analysis with Nielsen; at Inditex, store operations, visual merchandising standards and product flow",
            ],
        },
    ],
    "skills_brand": (  # -> "Partner Management"
        "partners & key accounts, performance reviews, multi-market coordination, negotiation, coaching"
    ),
    "skills_ecommerce": (  # -> "Operations & Stock"
        "stock-risk review, PO coverage, launch readiness, listings & data accuracy, talabat, Noon, Careem"
    ),
    "skills_commercial": (  # -> "Financial KPIs"
        "sales & GMV, sell-in/sell-out, P&L targets, pricing & margin, A&P budgets, forecast input"
    ),
    "skills_data": (  # -> "Brand & Marketing"
        "brand standards, local & seasonal campaigns, product launches, influencer & social, ROI/ROAS"
    ),
    "skills_tools": (  # -> "Tools"
        "Excel (Advanced), Power BI, Tableau, Looker, SAP, Salesforce, Shopify, Notion, Claude (AI)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Partner Management",
    "E-Commerce & Digital": "Operations & Stock",
    "Commercial": "Financial KPIs",
    "Data & Analytics": "Brand & Marketing",
}


def make_job() -> Job:
    return Job(
        id="yolk-brands-international-account-manager-dubai-2026-10",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE (On-site)",
        url="https://www.linkedin.com/jobs/search/?keywords=Yolk%20Brands%20International%20Account%20Manager",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Yolk Brands International Account Manager Dubai",
            "note": "Grupo F&B de Dubái (Pickl, BonBird, SouthPour). El puesto es el enlace con los franquiciados "
                    "internacionales: KPIs, informes a dirección, stock y proveedores, lanzamientos de menú/LTOs, "
                    "marketing local, SOPs, auditorías y sistema de inventario Edify. Easy Apply, promocionado por "
                    "técnico de selección, 698 solicitudes. Encaje: F&B, cuentas QSR en Glovo, partners en 50+ "
                    "mercados, revisión de stock. Gaps: franquicias y operación de restaurante.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": "2026-10-01", "raw": job.raw,
        "ai_score": 62, "ai_tier": "Warm",
        "skills_match": [
            "F&B: DoFreeze es un grupo F&B/FMCG de Emiratos que vende en 50+ mercados",
            "QSR: cuentas XL de KFC, Taco Bell, La Tagliatella y Sushi Shop en Glovo",
            "Partners en varios mercados: 6 lanzamientos en GCC, MENA, Asia, Europa, EE. UU. y África",
            "Inventario y roturas: hace la revisión semanal de riesgo de stock por dark store",
            "Informes a dirección: canal Flash Sales de Miravia reportando al CEO",
            "Punto de contacto único de 42 cuentas, +30% GMV QoQ",
            "Sistemas y coaching: automatización con IA y equipo de dos personas",
        ],
        "missing_skills": [
            "Gestión de franquicias (franquiciador–franquiciado): no la ha hecho",
            "Operación de restaurante: COGS, labor, recetas, SOPs, auditorías y visitas a tienda",
            "Royalties y contratos de franquicia",
            "Software de inventario Edify",
        ],
        "sector_fit": "bueno — F&B y QSR, aunque del lado marca/plataforma y no del de operación de restaurante",
        "seniority_fit": "lateral — Account Manager frente a su Manager actual",
        "red_flags": [
            "Es un puesto de operaciones de franquicia; su perfil es de marca y comercial",
            "698 solicitudes: conviene entrar por alguien de Yolk Brands",
            "Sueldo no publicado: confirmar suelo de 20.000 AED/mes",
            "Las auditorías y visitas a franquicias internacionales suponen viajar",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Encaja en F&B y en la parte de partners: DoFreeze es un grupo F&B de Emiratos que vende en 50+ "
            "mercados, en Glovo llevó cadenas QSR como KFC y Taco Bell, y en Miravia fue punto de contacto de 42 "
            "cuentas reportando al CEO. También encaja la parte de stock, por la revisión semanal por dark store. "
            "Pero el puesto es sobre todo operación de franquicias de restaurante (royalties, COGS, SOPs, "
            "auditorías, Edify), y eso no lo ha hecho. Con 698 solicitudes, sin una vía directa es Warm."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
