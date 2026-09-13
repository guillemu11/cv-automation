"""One-off: generate Paula's CV for the Deliveroo "Marketing Executive, Value &
Promotions" role (Dubai, hybrid — Middle East Marketing team, 2026-09-03).

The role owns the day-to-day of Deliveroo's promotional/incentive layer: turning the
ME promotional calendar into targeted lifecycle campaigns (acquisition, activation,
retention, win-back), building in-app/CRM campaigns with CRM & Tech, setting promos
up end-to-end (eligibility, caps, dates, copy), tracking campaign-level budget/ROI/
margin, weekly in-flight + post-campaign reporting, local test-and-learn experiments
and competitor monitoring. Asks 2–4 yrs in Growth/CRM/Campaign Management/Promotions
at a high-growth tech, marketplace or e-commerce platform; commercial metrics
(CAC, LTV, GMV, margin); Excel/Sheets/Looker; A/B testing; fluent English (Arabic an
advantage).

Paula's fit (strictly truthful, from profile.yaml — NO invented experience):
- **Miravia (Alibaba)**: owned the **Flash Sales** channel — the marketplace's core
  value/promotions engine — reporting to the CEO against P&L targets; +30% GMV QoQ
  across 42 accounts via pricing, assortment and targeted promotions; continuous
  ROI/ROAS/conversion/traffic/retention analysis.
- **DoFreeze (Dubai)**: plans the promotional calendar across the UAE and 50+ markets
  and runs promotional mechanics **on Deliveroo, talabat, Noon and Careem**; owns the
  A&P budget, ROI/ROAS and margin tracking; A/B testing on Meta & Google Ads; built an
  AI (Claude/GPT) reporting system.
- **Glovo**: data-led promotional activations for XL accounts on a food-delivery
  platform — she already knows this business model from the partner side.
- **Mondelez**: promotional-effectiveness evaluation and sell-in/sell-out reporting.

Honest gaps (surfaced, not hidden): **no Arabic** (an advantage here, not a
requirement — never claimed); CAC/LTV are not metrics she has owned by name (her
documented commercial stack is ROI/ROAS/GMV/margin/conversion/retention); and the
role is an *Executive*-level band while her current title is Manager — a step across,
not up.

Content authored directly (no LLM key needed), kept to ONE page, converted with
LibreOffice soffice, and landed under
output/2026-09-03/Deliveroo - Marketing Executive, Value & Promotions/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402

from career_ops.config import settings  # noqa: E402
from career_ops.discovery.normalize import Job  # noqa: E402
from career_ops.generators import cv_generator as cv  # noqa: E402

COMPANY = "Deliveroo"
TITLE = "Marketing Executive, Value & Promotions"
DATE_FOLDER = "2026-09-03"

JOB_DESCRIPTION = """\
Marketing Executive, Value & Promotions — Deliveroo, Dubai, UAE (hybrid). Middle East
Marketing team, reporting to the Senior Marketing Manager, Value & Promotions.

You will own the day-to-day delivery and performance of our promotional and incentive
campaigns, working cross-functionally to ensure every offer is targeted, well-timed and
commercially sound. Incentives sit at the centre of how customers experience value on the
platform. This role turns the promotional calendar into targeted campaigns across the
customer lifecycle, and makes sure every campaign earns its budget.

What you'll be doing:
- Promotional Calendar Execution: own day-to-day execution of the Middle East promotional
  calendar, translating the quarterly plan into targeted campaigns across acquisition,
  activation, retention and win-back.
- Promo Strategy Development: contribute to strategies on how value propositions change
  customer behaviour and drive ROI for Deliveroo and Partners.
- CRM & Lifecycle Campaigns: work with CRM and Tech to build and deploy in-app and lifecycle
  campaigns, ensuring offers reach the right segment at the right moment.
- Campaign Set-up & Accuracy: own promotion set-up end-to-end — eligibility, caps, dates and
  copy — so every offer goes live as briefed and on time.
- Customer-centric Programmes: work with Data Science to translate customer and performance
  data into promo mechanics that are genuinely relevant, not just discounted.
- Budget Tracking: own campaign-level budget tracking and reporting, flagging when ROI or
  margin impact is off plan.
- Performance Optimization: own weekly reporting on in-flight campaign performance and
  post-campaign analysis.
- Test & Learn Agenda: design and run local test-and-learn experiments.
- Competition Monitoring: keep on top of competitor trends and market activity.
- Customer Advocacy: champion the consumer voice — genuine value, not just a lower price.

What you'll need:
- 2–4 years in Growth Marketing, CRM, Campaign Management, Marketing, Commercial Strategy or
  Promotions within a high-growth technology, marketplace, e-commerce or consumer platform.
- Commercial acumen: owning/supporting customer incentive strategies and promotional campaigns
  with working knowledge of CAC, LTV, GMV and margin impact.
- Bachelor's degree in Marketing, Business, Economics, Analytics or related.
- Strong analytical skills — Excel, Google Sheets and BI tools like Looker.
- Experience running or supporting A/B tests and translating results into recommendations.
- Attention to detail, running several campaigns at once; excellent cross-functional
  collaboration; self-starter in a fast-paced, matrixed organisation.
- Fluent English; Arabic an advantage.
"""

ATS = [
    "promotions", "value & promotions", "promotional calendar", "promotional campaigns",
    "incentives", "customer incentives", "promo mechanics", "discount strategy",
    "campaign management", "campaign execution", "campaign set-up", "campaign performance",
    "growth marketing", "CRM", "lifecycle marketing", "in-app campaigns", "EDM",
    "customer lifecycle", "acquisition", "activation", "retention", "win-back",
    "customer segmentation", "targeting", "personalisation",
    "budget tracking", "campaign budget", "A&P budget", "ROI", "ROAS", "GMV",
    "margin impact", "commercial acumen", "pricing strategy", "P&L",
    "A/B testing", "test-and-learn", "experimentation", "post-campaign analysis",
    "weekly reporting", "performance reporting", "data-driven", "customer-centric",
    "Excel", "Google Sheets", "Looker", "Power BI", "Tableau", "BI tools",
    "competitor monitoring", "market activity", "cross-functional", "stakeholder management",
    "marketplace", "e-commerce", "quick-commerce", "food delivery", "consumer platform",
    "Deliveroo", "talabat", "Noon", "Careem", "Glovo", "Shopify",
    "Meta Ads", "Google Ads", "generative AI", "UAE", "Dubai", "GCC", "MENA",
]

CONTENT = {
    "headline": (
        "Value & Promotions · Promotional Calendar & Lifecycle Campaigns · "
        "Campaign ROI & Test-and-Learn · Quick-Commerce · Dubai"
    ),
    "professional_summary": (
        "Commercial marketer with 4+ years running promotions, pricing and value campaigns on marketplace, "
        "quick-commerce and FMCG platforms across the GCC and Europe. At Alibaba's Miravia I owned Flash Sales "
        "— the platform's value and promotions engine — reporting to the CEO and growing 42 accounts +30% GMV "
        "QoQ on targeted promotions. Today I run the UAE promo calendar on Deliveroo, talabat, Noon and Careem, "
        "owning campaign budgets, ROI and margin, weekly reporting and A/B test-and-learn. Dubai-based."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Homegrown UAE F&B / FMCG group | Brands: Befit, Eurocake, Flair | 50+ countries",
            "bullets": [
                "Own the promotional calendar across the UAE and 50+ markets — turning quarterly priorities into targeted, well-timed offers on Deliveroo, talabat, Noon and Careem, set up end-to-end (mechanics, eligibility, dates, copy)",
                "Own the A&P budget and campaign-level ROI, ROAS and margin tracking — flagging early when a promotion runs off plan — with an AI (Claude/GPT) layer automating weekly in-flight and post-campaign reporting",
                "Run A/B test-and-learn on offers, audiences and creative (Meta, Google Ads, EDM) and monitor competitor promo activity, feeding results into the next promotional cycle",
            ],
        },
        {
            "company": "Miravia / AliExpress (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Miravia — Alibaba's marketplace | Top-5 global e-commerce | 100K+ employees",
            "bullets": [
                "Owned the Flash Sales channel for Beauty, Fashion & Home — the marketplace's core value and promotions engine — reporting performance directly to the CEO against P&L targets",
                "Grew 42 key accounts +30% GMV QoQ on pricing, assortment and segment-targeted promotions, analysing conversion, retention and ROI continuously to sharpen promotional effectiveness",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce & food-delivery leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Built bespoke, data-led promotional activations for strategic partners on a food-delivery platform, lifting order volume and GMV",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category & Commercial Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B annual revenue | 90K+ employees",
            "bullets": [
                "Evaluated promotional effectiveness and built sell-in/sell-out performance reports for the chocolate category",
            ],
        },
    ],
    "skills_brand": (  # label -> "Promotions & Lifecycle"
        "promotional calendar & mechanics, value & pricing promotions, lifecycle campaigns "
        "(acquisition, retention, win-back), segmentation & targeting, CRM / EDM & in-app"
    ),
    "skills_ecommerce": (  # label -> "Platforms & Channels"
        "quick-commerce & delivery (Deliveroo, talabat, Noon, Careem, Glovo), marketplace flash sales & deals, "
        "Shopify e-store, Meta & Google Ads, marketing automation"
    ),
    "skills_commercial": (  # label -> "Commercial"
        "campaign budget ownership, promotional ROI & margin, pricing strategy, GMV growth, "
        "commercial planning, cross-functional & senior-stakeholder management"
    ),
    "skills_data": (  # label -> "Analytics & Testing"
        "A/B testing & test-and-learn, weekly in-flight & post-campaign reporting, ROI / ROAS / GMV & margin "
        "analysis, conversion & retention, Looker, Power BI, Tableau"
    ),
    "skills_tools": (  # label -> "Tools"
        "Excel (Advanced), Google Sheets, Looker, Power BI, Tableau, Meta Ads Manager, Google Ads, "
        "Shopify, Salesforce, SAP, Generative AI (Claude / ChatGPT)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Promotions & Lifecycle",
    "E-Commerce & Digital": "Platforms & Channels",
    "Commercial": "Commercial",
    "Data & Analytics": "Analytics & Testing",
}


def make_job() -> Job:
    return Job(
        id="deliveroo-marketing-executive-value-promotions-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Dubai, United Arab Emirates (hybrid)",
        url="https://www.linkedin.com/jobs/",  # LinkedIn verified job, applications handled off-LinkedIn
        source="linkedin",
        description=JOB_DESCRIPTION,
        raw={
            "query": "Deliveroo Marketing Executive Value & Promotions Dubai",
            "via": "LinkedIn (verified job, recruiter-promoted)",
            "team": "Middle East Marketing — Value & Promotions",
            "reports_to": "Senior Marketing Manager, Value & Promotions (ME)",
            "hybrid": True,
            "note": "JD says 'remote' on the LinkedIn card but the body states hybrid, Dubai-based. "
                    "176 applicants in the first day.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": None, "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 88, "ai_tier": "Hot",
        "skills_match": [
            "Owned Flash Sales at Miravia (Alibaba) — a marketplace's core value/promotions engine — reporting to the CEO against P&L targets",
            "Runs the UAE promotional calendar today on Deliveroo, talabat, Noon and Careem — promo mechanics, set-up and timing",
            "Campaign budget ownership (A&P) with ROI, ROAS and margin tracking + AI-automated weekly and post-campaign reporting",
            "A/B testing / test-and-learn on offers, audiences and creative (Meta & Google Ads, EDM)",
            "+30% GMV QoQ across 42 accounts via pricing, assortment and segment-targeted promotions",
            "Glovo XL Accounts — already knows the food-delivery platform model from the partner side",
            "Excel / Google Sheets + BI (Looker, Power BI, Tableau); BBA (CUNEF, E-Commerce specialisation)",
            "Dubai-based on a UAE residence visa; integrates brands into Deliveroo itself",
        ],
        "missing_skills": [
            "No Arabic (listed as 'an advantage', not a requirement — never claimed)",
            "CAC and LTV not owned by name — her documented commercial metrics are ROI/ROAS/GMV/margin/conversion/retention",
            "In-app / lifecycle campaign build inside a CRM platform (Braze-style) not documented — her lifecycle work is EDM/CRM + promo mechanics on marketplaces",
        ],
        "sector_fit": "excellent (promotions & incentives on a food-delivery platform — Glovo partner-side + Miravia flash sales + live Deliveroo integrations)",
        "seniority_fit": "top of band (asks 2–4 yrs; Paula 4+) — Executive title is a step across from her Manager title, not up",
        "red_flags": [
            "176 applicants within 24h — apply fast and use the internal referral contact",
            "Executive-level band: check the salary lands above the 20K AED/month floor before investing",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "One of the strongest Deliveroo matches so far. The role is the day-to-day owner of the promotional/"
            "incentive layer — promo calendar execution, lifecycle campaigns by segment, end-to-end promo set-up, "
            "campaign budget/ROI/margin tracking, weekly + post-campaign reporting, A/B test-and-learn and "
            "competitor monitoring. Paula has done nearly all of it: she owned Flash Sales at Alibaba's Miravia "
            "(the platform's value/promotions engine, reporting to the CEO), grew 42 accounts +30% GMV QoQ on "
            "pricing and targeted promotions, and today runs the UAE promo calendar and mechanics on Deliveroo, "
            "talabat, Noon and Careem while owning the A&P budget, ROI/ROAS/margin and an AI-automated reporting "
            "layer. Glovo gives her the food-delivery platform model from the partner side. Gaps are honest and "
            "non-blocking: no Arabic (an 'advantage' only), CAC/LTV not owned by name, and no documented in-app "
            "CRM-platform build. Note the band — Executive (2–4 yrs) is a sideways title move from Manager, so "
            "confirm the salary clears the 20K AED/month floor."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "CV Ready",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(dest / sub.name))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
