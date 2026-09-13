"""One-off: generate Paula's CV for **GCC Marketing Associate** at **IFF**
(International Flavors & Fragrances — global B2B ingredient house; Dubai, UAE).

Full JD was behind LinkedIn's login wall; a snippet indicated the role
maintains/updates marketing databases & BI tools and asks ~2 yrs brand marketing.
Content is authored from role + company: a regional B2B marketing associate for a
fragrance/FMCG ingredient supplier — market & consumer intelligence, category and
fragrance trends, competitive monitoring, BI databases/dashboards, brand & customer
marketing support, and data-driven business presentations for the GCC.

Paula's fit is genuine and honest:
- FRAGRANCE-category depth: KAM for Beauty, Fragrances & Fashion at Alibaba's
  Miravia — incl. the leading Arabian & oud houses (Arabian Oud, Lattafa, Swiss
  Arabian, Ajmal) — assortment, pricing, promotions, trend/competitive monitoring.
- FMCG market intelligence: Mondelez category planning (sell-in/sell-out, market/
  consumer analysis, promo effectiveness, forecasting, NPD).
- DoFreeze: brand & marketing across 50+ markets — market research, positioning,
  data-driven presentations, BI, AI-assisted analysis.
- Advanced Excel/PowerPoint + BI (Power BI, Tableau, Looker, Nielsen, Kantar).

Honesty guardrails: B2B ingredient-side is a pivot from her brand-side work — framed
on her real fragrance-category knowledge + market intelligence/BI, not overclaimed;
NO Arabic claimed (VERIFY whether the role needs it); factual "UAE Residence Visa".

ONE-PAGE standard: tight summary, 3/2/1/1 bullets, ~6 skills/row, SHORT labels.
Fills the real CV template, relabels the skills rows, converts to PDF via
LibreOffice, registers the job, verifies 1 page, lands under output/2026-08-30/.
Also drops a short-named 'Paula De Francisco - CV.pdf' copy for portals.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "IFF"
TITLE = "GCC Marketing Associate"
DATE_FOLDER = "2026-08-30"

JOB_DESCRIPTION = """\
GCC Marketing Associate — IFF (International Flavors & Fragrances), Dubai, UAE.
Regional marketing support for a global B2B fragrance/taste ingredient house:
maintain and continuously update marketing databases and business-intelligence
tools; track market, category and consumer/fragrance trends and the competitive
landscape; support brand and customer marketing initiatives; prepare data-driven
business presentations and insights for commercial/creative teams and management;
coordinate marketing projects, samples and briefs across the GCC. ~2 years in brand
marketing (fragrance/beauty/FMCG a plus); strong analytical and presentation skills;
Advanced Excel/PowerPoint and BI tools.
"""

ATS = [
    "marketing associate", "brand marketing", "market intelligence", "consumer insights",
    "market research", "category trends", "fragrance", "fragrances", "beauty", "FMCG",
    "competitive monitoring", "competitive landscape", "business intelligence", "BI",
    "databases", "dashboards", "data-driven", "presentations", "business presentations",
    "KPI reporting", "forecasting", "sell-in/sell-out", "assortment", "brand support",
    "customer marketing", "commercial", "stakeholder support", "Advanced Excel",
    "PowerPoint", "Power BI", "Tableau", "Looker", "Nielsen", "Kantar", "GCC", "MENA",
    "Dubai", "UAE", "generative AI",
]

CONTENT = {
    "headline": (
        "Marketing Associate · Fragrance & FMCG · Market Intelligence & Consumer "
        "Trends · Brand Support, Insights & Presentations"
    ),
    "professional_summary": (
        "Marketing professional with 4+ years across fragrances, beauty and FMCG — market intelligence, "
        "consumer and category trends, brand support and data-driven presentations. Fragrance-category "
        "depth from Alibaba's Miravia (Beauty, Fragrances & Fashion, incl. the leading Arabian & oud "
        "houses) and FMCG market grounding at Mondelez; now lead brand & marketing at DoFreeze across 50+ "
        "markets. Advanced Excel/PowerPoint and BI, AI-first, Dubai-based."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Multi-brand FMCG group | Brands: Befit, Eurocake, Flair | 50+ countries",
            "bullets": [
                "Lead brand & marketing across a multi-brand FMCG portfolio and 50+ GCC/MENA and international markets — market research, positioning, go-to-market and 6 NPD launches",
                "Turn market, category and consumer data into clear recommendations and business presentations for leadership — with an AI (Claude/GPT) layer automating research, reporting and BI",
                "Own campaign KPIs and the A&P budget, tracking performance and trends to steer brand and commercial decisions",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Miravia — Alibaba's marketplace | Top-5 global e-commerce | 100K+ employees",
            "bullets": [
                "Deep fragrance-category exposure — incl. the leading Arabian & oud houses (Arabian Oud, Lattafa, Swiss Arabian, Ajmal) — running assortment, pricing and promotions to +30% GMV QoQ",
                "Ran continuous trend and competitive monitoring across fragrances and beauty to shape assortment and commercial strategy",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Data-led account management on a quick-commerce platform — used market and performance data to build activations and growth plans that lifted order volume and GMV",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category & Commercial Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG multinational | €36B revenue | 90K+ employees",
            "bullets": [
                "Market intelligence & BI: sell-in/sell-out and market/consumer analysis, promotional-effectiveness reviews, performance reporting and sales forecasting for the chocolate category, plus NPD support (Milka Spread, Mini Suchard)",
            ],
        },
    ],
    "skills_brand": (  # label -> "Brand & Marketing"
        "brand marketing, positioning & go-to-market, NPD & product launches, campaigns & activations, "
        "PR & influencer, content, A&P budget management"
    ),
    "skills_ecommerce": (  # label -> "Category & Trends"
        "fragrance & beauty category knowledge, consumer & market trends, competitive monitoring, "
        "assortment & product mix, e-commerce & quick-commerce, social & digital"
    ),
    "skills_commercial": (  # label -> "Commercial"
        "key account & partner management, pricing & promotions, negotiation, commercial planning, "
        "cross-functional stakeholder support"
    ),
    "skills_data": (  # label -> "Market Intelligence"
        "market & consumer research, BI databases & dashboards, sell-in/sell-out, KPI reporting, "
        "forecasting, business presentations, Nielsen, Kantar, Power BI, Tableau, Looker"
    ),
    "skills_tools": (  # label -> "Tools"
        "Microsoft Excel (Advanced), PowerPoint (Advanced), Power BI, Tableau, Looker, Nielsen, Kantar, "
        "SAP, Salesforce, Generative AI (Claude, ChatGPT)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Brand & Marketing",
    "E-Commerce & Digital": "Category & Trends",
    "Commercial": "Commercial",
    "Data & Analytics": "Market Intelligence",
}


def make_job() -> Job:
    return Job(
        id="iff-gcc-marketing-associate-2026-08",
        title=TITLE,
        company=COMPANY,
        location="Dubai, United Arab Emirates",
        url="https://careers.iff.com/",
        source="linkedin",
        description=JOB_DESCRIPTION,
        raw={
            "query": "IFF GCC Marketing Associate Dubai fragrance market intelligence BI",
            "note": "B2B fragrance/FMCG marketing associate (market intelligence/BI). Honest fit via Paula's "
                    "fragrance-category depth (Miravia + Arabian oud houses) + FMCG market intelligence (Mondelez) "
                    "+ BI. ~2 yrs asked (Paula 4+). Full JD behind LinkedIn login — authored from role/company. "
                    "VERIFY whether Arabic is required. Factual UAE Residence Visa only.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": None, "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 80, "ai_tier": "Warm",
        "skills_match": [
            "Fragrance-category depth: Miravia KAM Beauty/Fragrances/Fashion incl. Arabian Oud, Lattafa, Swiss Arabian, Ajmal",
            "FMCG market intelligence & BI: Mondelez category planning (sell-in/sell-out, market/consumer analysis, forecasting)",
            "Turns market/consumer data into business presentations & recommendations (DoFreeze)",
            "BI + Advanced Excel/PowerPoint (Power BI, Tableau, Looker, Nielsen, Kantar)",
            "Trend & competitive monitoring across fragrances & beauty",
            "~2 yrs asked → Paula 4+; Dubai-based; AI-first for research/analysis",
        ],
        "missing_skills": [
            "B2B ingredient-side is a pivot from brand-side — framed on real fragrance-category + market-intelligence experience",
            "Arabic status unknown for this GCC B2B role — VERIFY before relying on the fit",
            "Full JD behind LinkedIn login — authored from role/company",
        ],
        "sector_fit": "good (fragrance + FMCG market intelligence — genuine domain overlap on the brand/insights side)",
        "seniority_fit": "over-band friendly (asks ~2 yrs; Paula 4+)",
        "red_flags": [
            "B2B ingredient supplier vs Paula's brand-side experience — pitch fragrance knowledge + market intelligence, not brand ownership",
            "Verify Arabic requirement (some GCC B2B roles ask for it)",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Good, honest fit on the fragrance + market-intelligence axis. IFF's GCC Marketing Associate supports "
            "regional marketing for a B2B fragrance/taste ingredient house — market/consumer intelligence, "
            "category & fragrance trends, BI databases/dashboards, brand/customer-marketing support and "
            "data-driven presentations, asking ~2 yrs brand marketing. Paula brings genuine fragrance-category "
            "depth (Miravia KAM for Beauty, Fragrances & Fashion incl. the leading Arabian & oud houses), FMCG "
            "market intelligence and BI (Mondelez category planning; Power BI/Tableau/Looker/Nielsen/Kantar), and "
            "turns market/consumer data into presentations at DoFreeze. It's a pivot to the ingredient/B2B side, "
            "so the CV leans on real fragrance knowledge + market intelligence rather than brand ownership. Full "
            "JD was behind LinkedIn's login (authored from role/company) and the Arabic requirement is unknown — "
            "verify. Factual UAE Residence Visa only."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
