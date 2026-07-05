#!/usr/bin/env python3
"""One-off runner: CV + cover letter (all 10) + deliverable + outreach (top 4)
for the Indeed jobs found 2026-06-15.

The LLM provider has no API key on this machine, so instead of calling the API
we feed pre-authored content straight into the generators via
``llm.override(...)``. The author of that content is Claude (this session) —
exactly the model the generator prompts target. Output is redirected into
``output/2026-06-15/<Company> - <Role>/`` by overriding ``output_dir``; the
generators further file each artefact under the standard 01_/02_/03_ subfolders.

All content below is drawn faithfully from profile.yaml — no invented metrics.

Top-4 (CV + CL + deliverable + outreach): SPACE, Henkel, Talabat, Huda Beauty.
Rest (CV + CL only): Justlife, PlaceUp, PureBorn, Farzana, DeFi, Solutions Leisure.
"""
from __future__ import annotations

import html
import json
import logging
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from career_ops import llm
from career_ops.analyzer import JobAnalysis
from career_ops.config import settings
from career_ops.discovery.normalize import Job, job_hash
from career_ops.generators import (
    generate_cover_letter,
    generate_cv,
    generate_deliverable,
)
from career_ops.generators._paths import job_subdir
from career_ops.generators.outreach import generate_outreach

logging.basicConfig(level=logging.INFO, format="%(asctime)s  %(levelname)-5s  %(message)s", datefmt="%H:%M:%S")
logger = logging.getLogger("run_2026_06_15")

DATE = "2026-06-15"
settings.output_dir = ROOT / "output" / DATE
settings.output_dir.mkdir(parents=True, exist_ok=True)

# Real experience context lines (from profile.yaml), reused verbatim.
CTX = {
    "DoFreeze": "Global FMCG distributor | Brands: Befit, Eurocake, Flair | 50+ countries",
    "Miravia": "Top 5 global e-commerce | Alibaba Group | 100K+ employees",
    "Glovo": "Quick-commerce leader | EUR 500M+ revenue | 10K+ employees",
    "Mondelez": "Global FMCG | EUR 36B annual revenue | 90K+ employees",
    "Massimo Dutti": "Inditex Group flagship premium fashion brand | Hands-on retail floor",
}


def dofreeze(bullets):
    return {"company": "DoFreeze LLC", "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present", "location": "Dubai, UAE",
            "context": CTX["DoFreeze"], "bullets": bullets}

def miravia(role, bullets):
    return {"company": "Miravia (Alibaba Group)", "role": role,
            "dates": "Nov 2023 – Oct 2025", "location": "Madrid, Spain",
            "context": CTX["Miravia"], "bullets": bullets}

def glovo(bullets):
    return {"company": "Glovo", "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023", "location": "Madrid, Spain",
            "context": CTX["Glovo"], "bullets": bullets}

def mondelez(bullets):
    return {"company": "Mondelez International", "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022", "location": "Madrid, Spain",
            "context": CTX["Mondelez"], "bullets": bullets}

def massimo(bullets):
    return {"company": "Massimo Dutti (Inditex)", "role": "Sales Associate – Las Rozas Village",
            "dates": "Jun 2018 – Jun 2019", "location": "Madrid, Spain",
            "context": CTX["Massimo Dutti"], "bullets": bullets}


# ===================================================================
# Outreach pack writer (simple, self-contained — email + LinkedIn copy)
# ===================================================================

def _html_to_paragraphs(html_str: str) -> list[str]:
    s = re.sub(r"<br\s*/?>", "\n", html_str, flags=re.I)
    parts = re.split(r"</p\s*>", s, flags=re.I)
    out = []
    for p in parts:
        text = re.sub(r"<[^>]+>", "", p)
        text = html.unescape(text)
        text = re.sub(r"[ \t]+", " ", text).strip()
        if text:
            out.append(text)
    return out


def build_outreach_pack(job: Job, content) -> Path:
    from docx import Document
    from docx.shared import Pt, RGBColor

    slug = re.sub(r"[^A-Za-z0-9]+", "_", job.company).strip("_")
    out_path = job_subdir(job, "outreach") / f"{slug}_Outreach_Pack.docx"

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    doc.add_heading(f"{job.company} — Outreach Pack", level=0)
    sub = doc.add_paragraph()
    r = sub.add_run(f"Outreach copy for an application to the {job.title} role at {job.company}.")
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.add_heading("Email", level=2)
    s = doc.add_paragraph()
    s.add_run("Subject:  ").bold = True
    s.add_run(content.email_subject)
    for para in _html_to_paragraphs(content.email_body):
        doc.add_paragraph(para)

    doc.add_heading("LinkedIn connection note (≤300 chars)", level=2)
    doc.add_paragraph(content.linkedin_connection)

    doc.add_heading("LinkedIn InMail", level=2)
    doc.add_paragraph(content.linkedin_inmail)

    doc.save(str(out_path))
    return out_path


# ===================================================================
# JOB DEFINITIONS
# ===================================================================

JOBS = []

# -------------------------------------------------------------------
# 1. SPACE — Brand Manager - Fragrances   [TOP-4: deliverable + outreach]
# -------------------------------------------------------------------
JOBS.append({
    "tier4": True,
    "deliverable_type": "brand_analysis",
    "apply_url": "https://to.indeed.com/aaz9fvt4nr4q",
    "job": Job(
        id=job_hash("Brand Manager - Fragrances", "SPACE", "Dubai, UAE"),
        title="Brand Manager - Fragrances",
        company="SPACE",
        location="Dubai, UAE",
        url="https://to.indeed.com/aaz9fvt4nr4q",
        source="indeed",
        description=(
            "Lead development and execution of brand strategies for a perfume portfolio across "
            "African markets to drive sales growth, market share, brand equity, profitability and "
            "successful product launches. Develop annual and long-term brand strategies; define "
            "positioning, target segments and growth opportunities. Analyse consumer behaviour, "
            "fragrance trends and competitor activity. Recommend portfolio expansion, SKU "
            "rationalisation and pricing. Coordinate NPD and launches. Manage ATL, BTL, digital, "
            "influencer, PR and trade marketing. Manage marketing budgets and track ROI. 5-8+ years "
            "brand management, preferably beauty, fragrances, cosmetics or luxury."
        ),
    ),
    "analysis": JobAnalysis(score=92, tier="Hot",
        skills_match=["brand strategy", "NPD end-to-end", "influencer marketing", "trade marketing", "pricing strategy", "P&L management", "fragrances"],
        missing_skills=[], sector_fit="Beauty & Fragrances", seniority_fit="Manager — strong fit (role asks 5-8 yrs; Paula 4+ with deep fragrance exposure)",
        red_flags=[], ats_keywords=["brand strategy", "fragrance", "NPD", "positioning", "influencer", "trade marketing", "P&L", "ROI", "go-to-market", "pricing"],
        reasoning="Fragrances + brand strategy is her exact sweet spot — Miravia PIC Fragrances + DoFreeze NPD."),
    "cv": {
        "headline": "Brand Manager · Fragrances & Beauty · NPD · Influencer & Trade Marketing",
        "professional_summary": (
            "Brand and marketing professional with 4+ years across beauty, fragrances, FMCG and "
            "e-commerce. Currently leads Brand & Marketing for DoFreeze across 50+ markets, running NPD "
            "end-to-end and an influencer programme scaled from zero to 25–50 creators per campaign. As "
            "PIC Fragrances at Alibaba's Miravia she grew 42 accounts +30% GMV QoQ and onboarded 30+ "
            "fragrance stores in two months, including leading Arabian & oud houses (Arabian Oud, Lattafa, Swiss Arabian, Ajmal) — a fragrance-category brand builder."
        ),
        "experience": [
            dofreeze([
                "Lead NPD end-to-end for 6 product launches — brief, packaging, pricing and go-to-market — across GCC, MENA, Asia, Europe, USA and Africa",
                "Define brand positioning and annual marketing calendars, building and scaling an influencer programme from zero (25–50 creators per campaign) plus sampling and seeding",
                "Plan ATL/BTL, digital, influencer and trade marketing activity, and manage A&P budget allocation to ROI across markets",
                "Develop trade and shopper marketing plans by channel and integrate brands into UAE quick-commerce (Noon, Talabat, Careem, Deliveroo)",
            ]),
            miravia("Key Account Manager & PIC Fragrances – Beauty, Fragrances & Fashion", [
                "Led the fragrance category as PIC Fragrances, onboarding 30+ stores in two months — including the official distributors of leading Arabian & oud houses (Arabian Oud, Lattafa, Swiss Arabian, Ajmal) — via trend-driven assortment and promotions",
                "Managed 42 beauty, fragrance and fashion accounts, achieving +30% GMV growth QoQ through pricing strategy, assortment optimisation and targeted promotions",
                "Created and led the Beauty Club and Hot on Social projects, boosting brand visibility, loyalty and positioning Miravia as a beauty destination",
                "Analysed fragrance trends, ROI, ROAS, conversion and retention to optimise category performance and forecasting",
            ]),
            glovo([
                "Managed strategic key accounts and drove GMV growth through data-led planning and bespoke marketing activations",
                "Onboarded beauty and lifestyle brands as Glovo built out its retail vertical beyond food delivery",
            ]),
            mondelez([
                "Conducted sell-in/sell-out analysis and evaluated promotional effectiveness for the chocolate category",
                "Contributed to NPD launches including Milka Spread and Mini Suchard",
            ]),
        ],
        "skills_brand": "brand strategy, brand positioning, NPD end-to-end, fragrance category management, influencer marketing, ATL/BTL, trade marketing, A&P budget management, go-to-market",
        "skills_ecommerce": "e-store management, quick-commerce, Noon, Talabat, Careem, Deliveroo, Instagram, TikTok, EDM, digital campaigns",
        "skills_commercial": "key account management, assortment planning, pricing strategy, SKU rationalisation, distributor management, modern trade, negotiation, forecasting",
        "skills_data": "P&L management, ROI, ROAS, market share, sell-in/sell-out, KPI tracking, Nielsen, Kantar",
        "skills_tools": "Canva, Meta Ads Manager, Google Ads, Salesforce, SAP, Looker, Nielsen, Kantar, Microsoft Office (Expert)",
    },
    "cl": {
        "opening_paragraph": (
            "Fragrance is the category where I have spent the most rewarding part of my career — first as "
            "PIC Fragrances at Alibaba's Miravia, now leading NPD and brand for a multi-market FMCG portfolio "
            "at DoFreeze. SPACE's focus on building perfume brands across African markets is exactly the kind "
            "of challenge I want to take on."
        ),
        "body_paragraph_1": (
            "This role lives on brand strategy, positioning and launches — my day-to-day. At Miravia I led "
            "fragrance category expansion as PIC Fragrances, onboarding 30+ stores in two months (including "
            "Arabian fragrances) and growing 42 accounts +30% GMV QoQ through pricing, assortment and "
            "promotions. At DoFreeze I lead NPD end-to-end for six launches across GCC, MENA, Asia, Europe, "
            "USA and Africa, managing A&P budgets to ROI."
        ),
        "body_paragraph_2": (
            "Beyond strategy, I bring hands-on ATL/BTL, influencer and trade marketing execution — I built an "
            "influencer programme from zero to 25–50 creators per campaign — plus genuine fluency in fragrance "
            "trends and consumer insight. I am already based in Dubai on a UAE residence visa, work bilingually "
            "in Spanish and English, and am comfortable coordinating distributors and country managers across markets."
        ),
        "closing_paragraph": (
            "I would welcome the chance to discuss how my fragrance brand-building experience can drive SPACE's "
            "growth across African markets. I am available to interview at your convenience and can start "
            "immediately. Thank you for your consideration."
        ),
    },
    "deliverable": {
        "title": "Fragrance Brand & Market Lens — SPACE across African Markets",
        "subtitle": "Three observations on positioning a perfume portfolio for growth and share, with concrete moves.",
        "sections": [
            {
                "heading": "Positioning across heterogeneous African markets",
                "finding": "A single perfume portfolio has to win in markets with very different price ceilings, scent preferences and channel structures (modern trade vs. open-market distribution).",
                "insight": "Treating Africa as one market dilutes positioning; the brands that win segment by scent family and price tier per country rather than pushing one global story everywhere.",
                "recommendation": "Build a per-country positioning grid (scent family × price tier × hero SKU) and concentrate launch and A&P firepower on 2–3 lead markets first — the same focus-then-scale approach I used leading NPD across 50+ markets at DoFreeze.",
            },
            {
                "heading": "Fragrance assortment & SKU productivity",
                "finding": "Perfume portfolios accumulate long tails of slow SKUs that absorb working capital and shelf space without driving share.",
                "insight": "Disciplined SKU rationalisation plus a few well-timed trend-led launches usually moves share more than broad range extension.",
                "recommendation": "Run a quarterly assortment review (bestseller protection + tail pruning + trend gaps such as Arabian/oud profiles) — I onboarded 30+ fragrance stores in two months at Miravia by leaning into exactly this trend-driven assortment discipline.",
            },
            {
                "heading": "Influencer, sampling & trade activation as the growth engine",
                "finding": "Fragrance is discovery-led: people buy what they smell and what they see worn — channels where ATL alone underperforms.",
                "insight": "A coordinated influencer + sampling + trade-activation motion compounds: creators drive trial intent, sampling converts it, and trade execution captures it at the shelf.",
                "recommendation": "Stand up a creator + sampling programme tied to retail activation calendars per market — I built one from zero to 25–50 creators per campaign at DoFreeze, with sampling/seeding across modern trade and quick-commerce.",
            },
        ],
        "closing": (
            "These are first-pass observations from the outside; with access to SPACE's market and sell-out data I "
            "would pressure-test them fast. They reflect how I actually run fragrance brands — focus markets, "
            "disciplined assortment, and a discovery-led activation engine."
        ),
    },
    "outreach_mode": "applied",
    "outreach": {
        "email_subject": "Fragrance brand-building for SPACE's Africa portfolio",
        "email_body": (
            "<p>Hello,</p>"
            "<p>I'm applying for the Brand Manager – Fragrances role. Fragrance is the category I know best: at "
            "Alibaba's Miravia I was PIC Fragrances, onboarding 30+ stores in two months and growing 42 accounts "
            "<strong>+30% GMV QoQ</strong>; at DoFreeze I now lead NPD end-to-end across 50+ markets including Africa.</p>"
            "<p>For a multi-market perfume portfolio I'd focus on per-country positioning, disciplined SKU productivity, "
            "and a discovery-led influencer + sampling + trade engine — I built one from zero to 25–50 creators per campaign.</p>"
            "<p>I'm Dubai-based on a residence visa and can start immediately. I'd welcome a short conversation about how "
            "I can support SPACE's growth across African markets.</p>"
            "<p>Best regards,<br>Paula De Francisco</p>"
        ),
        "linkedin_connection": (
            "Hi — I'm applying for your Brand Manager – Fragrances role. Fragrance is my home category (PIC Fragrances at "
            "Miravia, NPD lead at DoFreeze across 50+ markets). I'd love to connect and share how I'd grow the Africa portfolio."
        ),
        "linkedin_inmail": (
            "Hi — I've applied for the Brand Manager – Fragrances position and wanted to introduce myself. I led fragrance "
            "category expansion as PIC Fragrances at Alibaba's Miravia (30+ stores onboarded in two months, 42 accounts at "
            "+30% GMV QoQ) and now run NPD end-to-end across 50+ markets at DoFreeze. I'm Dubai-based on a residence visa and "
            "would value a short conversation about SPACE's perfume portfolio across African markets."
        ),
    },
})

# -------------------------------------------------------------------
# 2. Henkel — Trade & Shopper Marketing Manager - GCC   [TOP-4]
# -------------------------------------------------------------------
JOBS.append({
    "tier4": True,
    "deliverable_type": "action_plan",
    "apply_url": "https://to.indeed.com/aap2pgrfwfb4",
    "job": Job(
        id=job_hash("Trade & Shopper Marketing Manager - GCC", "Henkel", "Dubai, UAE"),
        title="Trade & Shopper Marketing Manager - GCC",
        company="Henkel",
        location="Dubai, UAE",
        url="https://to.indeed.com/aap2pgrfwfb4",
        source="indeed",
        description=(
            "Develop and execute GCC trade and shopper marketing strategies aligned with brand and "
            "commercial objectives. Translate category and shopper insights into channel and customer "
            "plans. Drive Perfect Store execution. Lead annual trade activation-calendar planning. Work "
            "closely with Key Account Managers and Sales. Design shopper activation programs and "
            "promotions. Support NPD launches via go-to-market planning. Manage trade marketing budgets "
            "with focus on ROI. 4-6 years in Trade/Shopper/Category/Sales within FMCG. Strong "
            "understanding of GCC Modern Trade and Traditional Trade channels."
        ),
    ),
    "analysis": JobAnalysis(score=90, tier="Hot",
        skills_match=["trade marketing", "shopper marketing", "category management", "key account management", "NPD", "modern trade", "ROI"],
        missing_skills=[], sector_fit="FMCG", seniority_fit="Manager — strong fit",
        red_flags=[], ats_keywords=["trade marketing", "shopper marketing", "GCC", "modern trade", "Perfect Store", "key account", "NPD", "ROI", "category"],
        reasoning="Near-exact match: trade & shopper marketing FMCG in GCC."),
    "cv": {
        "headline": "Trade & Shopper Marketing Manager · FMCG · GCC · Key Account & Category",
        "professional_summary": (
            "Trade and shopper marketing professional with 4+ years across FMCG, beauty and e-commerce, "
            "currently leading Brand & Marketing for DoFreeze across 50+ GCC, MENA and global markets. "
            "Proven at turning shopper insight into channel and customer plans, driving modern-trade and "
            "quick-commerce execution, and managing A&P budgets to ROI alongside key-account and sales teams."
        ),
        "experience": [
            dofreeze([
                "Develop trade and shopper marketing plans by channel across 50+ markets, supporting key-account strategy and category management with the sales team",
                "Integrated brands into UAE modern trade and quick-commerce (Noon, Careem, Talabat, Deliveroo) — onboarding, listings, promotional mechanics and retail execution",
                "Own A&P budget allocation and commercial planning, analysing ROI and ROAS across channels to maximise efficiency",
                "Lead NPD end-to-end for 6 launches (brief, packaging, pricing, go-to-market) across GCC, MENA, Asia, Europe, USA and Africa",
                "Represent brands at international trade fairs, opening distributor relationships and expanding GCC market coverage",
            ]),
            miravia("Key Account Manager – Beauty, Fragrances & Fashion", [
                "Managed 42 key accounts, achieving +30% GMV growth QoQ via pricing strategy, assortment optimisation and targeted promotions",
                "Led category expansion as PIC Fragrances, onboarding 30+ new stores in two months through trend-driven assortment and promotions",
                "Owned the Flash Sales channel for Beauty, Fashion & Home, executing commercial plans aligned to P&L targets",
                "Analysed ROI, ROAS, conversion and retention to optimise channel performance and forecasting accuracy",
            ]),
            glovo([
                "Managed strategic key accounts and drove GMV growth through data-led planning and bespoke marketing activations",
                "Helped build Glovo's retail vertical, onboarding partners and expanding the marketplace beyond food delivery",
                "Negotiated and closed high-impact commercial deals maximising profitability for partners and platform",
            ]),
            mondelez([
                "Conducted sell-in/sell-out analysis and evaluated promotional effectiveness for the chocolate category",
                "Identified growth opportunities and contributed to NPD launches including Milka Spread and Mini Suchard",
                "Built performance reports informing category and trade decisions",
            ]),
        ],
        "skills_brand": "trade marketing, shopper marketing, go-to-market, NPD end-to-end, A&P budget management, brand strategy, omnichannel campaigns",
        "skills_ecommerce": "quick-commerce, Noon, Talabat, Careem, Deliveroo, retail execution, e-store management, process automation",
        "skills_commercial": "key account management, category management, modern trade, general trade, distributor management, pricing strategy, assortment planning, planogram, negotiation, forecasting",
        "skills_data": "sell-in/sell-out, ROI, ROAS, P&L management, KPI tracking, Nielsen, Kantar, Looker",
        "skills_tools": "SAP, Salesforce, Nielsen, Kantar, Looker, Planorama, Microsoft Office (Expert)",
    },
    "cl": {
        "opening_paragraph": (
            "Henkel's strength in Modern Trade and its Perfect Store ambition across the GCC are exactly the "
            "kind of channel and shopper challenges I work on every day. As Brand & Marketing Manager at a "
            "Dubai-based FMCG distributor operating in 50+ markets, I would welcome the chance to bring that "
            "hands-on GCC trade experience to your portfolio."
        ),
        "body_paragraph_1": (
            "I build trade and shopper marketing plans by channel and translate them into customer plans with "
            "the sales team — the core of this role. At Miravia (Alibaba) I managed 42 key accounts and grew "
            "GMV +30% QoQ through pricing, assortment and promotions, and led category expansion as PIC "
            "Fragrances, onboarding 30+ stores in two months. At DoFreeze I run A&P budgets to ROI and lead "
            "NPD go-to-market end-to-end across six launches."
        ),
        "body_paragraph_2": (
            "What sets me apart for the GCC specifically: hands-on integration of brands into UAE modern trade "
            "and quick-commerce (Noon, Talabat, Careem, Deliveroo), a category-planning foundation from "
            "Mondelez (sell-in/sell-out, promo effectiveness), and the fact that I am already based in Dubai on "
            "a UAE residence visa — no relocation or sponsorship needed. I work bilingually in Spanish and English."
        ),
        "closing_paragraph": (
            "I would love to discuss how my GCC trade and shopper experience can support Henkel's category "
            "growth and retailer partnerships. I am available for an interview at your convenience and can start "
            "immediately. Thank you for your consideration."
        ),
    },
    "deliverable": {
        "title": "90-Day Trade & Shopper Plan — Henkel GCC",
        "subtitle": "How I would land the first 90 days against Perfect Store and the trade activation calendar.",
        "sections": [
            {
                "heading": "Weeks 1–2 — Channel & shopper audit",
                "finding": "A new trade & shopper manager inherits live customer plans, a Perfect Store standard and an activation calendar already in motion across GCC modern and traditional trade.",
                "insight": "The fastest way to add value is to map where current execution diverges from the Perfect Store picture-of-success by channel and key account, rather than rewriting strategy on day one.",
                "recommendation": "Run a structured audit with the KAM and Sales teams — Perfect Store compliance by channel, top-account gaps, promo ROI history — and turn it into a one-page diagnostic. This mirrors how I open new markets at DoFreeze before committing A&P.",
            },
            {
                "heading": "Month 1 — Quick wins on promo ROI and shopper activation",
                "finding": "Trade budgets often carry under-performing mechanics that persist because they are calendar habits, not ROI decisions.",
                "insight": "Reallocating spend from low-ROI mechanics to a few high-converting shopper activations usually lifts efficiency without extra budget.",
                "recommendation": "Pick 2–3 key accounts, re-cut their promo calendar by ROI, and pilot one upgraded shopper activation per channel — measured on sell-out. At Miravia I grew 42 accounts +30% GMV QoQ doing exactly this pricing/assortment/promo discipline.",
            },
            {
                "heading": "Months 2–3 — NPD go-to-market & trade calendar for the year",
                "finding": "NPD launches and the annual trade activation calendar are where trade & shopper marketing either accelerates the brand or leaves money on the shelf.",
                "insight": "Launches that bundle listing, Perfect Store execution and a shopper mechanic out of the gate outperform staggered ones.",
                "recommendation": "Build the annual trade activation calendar around the NPD pipeline, with a go-to-market checklist per launch (listing → planogram → shopper mechanic → measurement). I lead NPD end-to-end for six launches at DoFreeze across GCC and MENA.",
            },
        ],
        "closing": (
            "This is a first-pass plan from the outside — with Henkel's account and Nielsen data I would sharpen the "
            "priorities quickly. It reflects how I actually operate: audit first, win on ROI, then build the calendar around NPD."
        ),
    },
    "outreach_mode": "applied",
    "outreach": {
        "email_subject": "Trade & shopper for Henkel GCC — Perfect Store focus",
        "email_body": (
            "<p>Hello,</p>"
            "<p>I'm applying for the Trade & Shopper Marketing Manager – GCC role. I currently lead Brand & Marketing "
            "for DoFreeze across 50+ markets, building trade and shopper plans by channel and driving modern-trade and "
            "quick-commerce execution alongside the sales team.</p>"
            "<p>For the first 90 days I'd audit Perfect Store execution by channel and key account, win quickly on promo ROI, "
            "and build the annual trade calendar around the NPD pipeline. At Alibaba's Miravia I grew 42 accounts "
            "<strong>+30% GMV QoQ</strong> with exactly this pricing/assortment/promo discipline.</p>"
            "<p>I'm Dubai-based on a residence visa — no relocation or sponsorship needed — and can start immediately. "
            "I'd welcome a short conversation.</p>"
            "<p>Best regards,<br>Paula De Francisco</p>"
        ),
        "linkedin_connection": (
            "Hi — I've applied for your Trade & Shopper Marketing Manager – GCC role. I run trade & shopper plans across "
            "50+ FMCG markets at DoFreeze and grew 42 accounts +30% GMV QoQ at Miravia. Dubai-based on a visa. I'd love to connect."
        ),
        "linkedin_inmail": (
            "Hi — I've applied for the Trade & Shopper Marketing Manager – GCC position. I lead trade and shopper marketing "
            "by channel across 50+ markets at DoFreeze, with hands-on GCC modern-trade and quick-commerce execution, and I "
            "grew 42 key accounts +30% GMV QoQ at Alibaba's Miravia. I'm already in Dubai on a residence visa and would value "
            "a short conversation about Henkel's Perfect Store and trade activation priorities."
        ),
    },
})

# -------------------------------------------------------------------
# 3. Talabat (Delivery Hero) — Key Account Manager   [TOP-4]
# -------------------------------------------------------------------
JOBS.append({
    "tier4": True,
    "deliverable_type": "ecommerce_teardown",
    "apply_url": "https://to.indeed.com/aa68bmxgnnm4",
    "job": Job(
        id=job_hash("Key Account Manager", "Talabat (Delivery Hero)", "Dubai, UAE"),
        title="Key Account Manager",
        company="Talabat (Delivery Hero)",
        location="Dubai, UAE",
        url="https://to.indeed.com/aa68bmxgnnm4",
        source="indeed",
        description=(
            "Talabat (Delivery Hero) seeks a dynamic, analytical Key Account Manager to maximise the value of "
            "key restaurant partnerships. Define and execute account strategies aligned with Talabat's business "
            "objectives and revenue growth. Lead rigorous analysis of sales data, customer insights and "
            "operational metrics. Lead cross-functional projects (marketing, product, operations, finance) to "
            "optimise account performance. Build long-term partnerships with key stakeholders. Establish KPIs "
            "and dashboards. Manage contract negotiations, renewals and commercial terms. 4-6 years in account "
            "management, strategic partnerships or commercial roles. Experience in e-commerce, food delivery or "
            "tech a strong advantage. Proficiency in CRM (Salesforce) and data dashboards."
        ),
    ),
    "analysis": JobAnalysis(score=88, tier="Hot",
        skills_match=["key account management", "quick-commerce", "data-led planning", "negotiation", "pricing strategy", "cross-functional projects", "Salesforce"],
        missing_skills=[], sector_fit="Quick-commerce / E-commerce", seniority_fit="4-6 yrs — strong fit",
        red_flags=[], ats_keywords=["key account", "quick-commerce", "GMV", "Salesforce", "dashboards", "negotiation", "partnerships", "data analysis"],
        reasoning="Quick-commerce KAM — direct overlap with Talabat/Noon/Careem/Deliveroo experience + 42 accounts at Miravia + Glovo."),
    "cv": {
        "headline": "Key Account Manager · Quick-Commerce & E-Commerce · Data-Led Growth",
        "professional_summary": (
            "Commercial and key-account professional with 4+ years across quick-commerce, e-commerce and FMCG. "
            "Grew 42 key accounts +30% GMV QoQ at Alibaba's Miravia and helped build Glovo's retail vertical, with "
            "hands-on integration of brands into UAE quick-commerce platforms (Noon, Talabat, Careem, Deliveroo). "
            "Analytical, negotiation-led and fluent in the data dashboards that drive account decisions."
        ),
        "experience": [
            miravia("Key Account Manager – Beauty, Fragrances & Fashion", [
                "Owned day-to-day relationships with 42 key accounts, achieving +30% GMV growth QoQ through pricing strategy, assortment optimisation and targeted promotions",
                "Reported directly to the CEO on the Flash Sales channel, executing strategic commercial plans aligned to P&L targets",
                "Led category expansion as PIC Fragrances, onboarding 30+ new stores in two months via data-led assortment and promotions",
                "Continuously analysed ROI, ROAS, conversion, traffic and retention to optimise account performance and forecasting accuracy",
            ]),
            glovo([
                "Managed strategic key accounts (KFC, Taco Bell, La Tagliatella, Sushi Shop) on a quick-commerce platform, driving GMV growth through data-led planning and bespoke activations",
                "Part of the team that built Glovo's retail vertical — onboarding partners and expanding the marketplace beyond food delivery",
                "Led cross-functional teams across marketing, logistics and customer support to deliver campaigns and increase order volume",
                "Negotiated and closed high-impact commercial deals maximising profitability for both Glovo and partners",
            ]),
            dofreeze([
                "Integrated brands into key UAE quick-commerce platforms (Noon, Careem, Talabat, Deliveroo) — onboarding, listings, promotional mechanics and retail execution",
                "Own commercial planning and A&P budget allocation, analysing ROI and ROAS across channels",
                "Develop trade plans by channel, supporting key-account strategy across 50+ markets",
            ]),
            mondelez([
                "Conducted sell-in/sell-out analysis and built performance reports informing account and category decisions",
                "Evaluated promotional effectiveness and contributed to NPD launches",
            ]),
        ],
        "skills_brand": "go-to-market, promotional planning, trade marketing, bespoke activations",
        "skills_ecommerce": "quick-commerce, Talabat, Noon, Careem, Deliveroo, marketplace growth, e-store management, listings & promotions",
        "skills_commercial": "key account management, strategic partnerships, negotiation, contract renewals, pricing strategy, assortment planning, distributor management, forecasting",
        "skills_data": "GMV growth, P&L management, ROI, ROAS, conversion, KPI dashboards, sell-in/sell-out, Salesforce",
        "skills_tools": "Salesforce, Looker, SAP, Nielsen, Kantar, Microsoft Office (Expert)",
    },
    "cl": {
        "opening_paragraph": (
            "Talabat is the platform I work alongside from the brand side every week, so the chance to drive key "
            "restaurant partnerships from within is genuinely exciting. With 4+ years managing key accounts across "
            "quick-commerce and e-commerce — including at Glovo — I would bring an analytical, partner-first approach "
            "to your accounts."
        ),
        "body_paragraph_1": (
            "Maximising account value through data is exactly how I work. At Alibaba's Miravia I owned 42 key accounts "
            "and grew GMV +30% QoQ through pricing, assortment and promotions, reporting to the CEO on the Flash Sales "
            "channel. At Glovo I managed strategic accounts (KFC, Taco Bell, Sushi Shop) and led cross-functional teams "
            "across marketing, logistics and operations to grow order volume."
        ),
        "body_paragraph_2": (
            "I also bring the quick-commerce fluency this role rewards: at DoFreeze I integrate brands directly into "
            "Talabat, Noon, Careem and Deliveroo, so I understand both sides of the partnership. I am data-led on KPIs "
            "and dashboards, comfortable in Salesforce, already based in Dubai on a UAE residence visa, and bilingual in "
            "Spanish and English."
        ),
        "closing_paragraph": (
            "I would welcome the chance to discuss how I can grow and retain Talabat's key accounts. I am available to "
            "interview at your convenience and can start immediately. Thank you for considering my application."
        ),
    },
    "deliverable": {
        "title": "Key Account Growth Lens — Talabat Restaurant Partnerships",
        "subtitle": "Three observations on maximising value from key restaurant accounts, from someone who works the platform brand-side.",
        "sections": [
            {
                "heading": "Account segmentation drives where the growth actually is",
                "finding": "Key restaurant partners are not homogeneous — a few accounts drive a disproportionate share of GMV, and they each have different growth levers (menu, pricing, visibility, ops reliability).",
                "insight": "A flat 'grow every account' plan spreads effort thin; segmenting accounts by GMV potential and primary lever concentrates the manager's time where it compounds.",
                "recommendation": "Build an account tiering grid (GMV potential × primary lever) and a quarterly plan per top account — the same 42-account discipline I used at Miravia to deliver +30% GMV QoQ.",
            },
            {
                "heading": "Data and dashboards as the negotiation currency",
                "finding": "Partner conversations stall when they are opinion-led; they accelerate when the manager brings the partner's own sell-out, conversion and operational data.",
                "insight": "A shared dashboard reframes renewals and commercial terms from haggling into a joint growth conversation backed by evidence.",
                "recommendation": "Stand up a simple per-account KPI dashboard (GMV, conversion, prep time, promo ROI) and lead every review from it — I ran exactly this ROI/ROAS/conversion analysis cadence at Miravia and Glovo.",
            },
            {
                "heading": "Cross-functional execution is where account plans live or die",
                "finding": "Account growth depends on marketing, product, operations and finance moving together — a promo without ops readiness or visibility falls flat.",
                "insight": "The KAM who can orchestrate cross-functional delivery turns account plans into realised GMV rather than slideware.",
                "recommendation": "Run each major account initiative as a small cross-functional project with clear owners and a measurement plan — I led cross-functional teams across marketing, logistics and support at Glovo to lift order volume.",
            },
        ],
        "closing": (
            "These are outside-in observations; with Talabat's account data I'd refine them quickly. They reflect how I run "
            "key accounts — segment by potential, lead with data, and orchestrate cross-functional delivery."
        ),
    },
    "outreach_mode": "applied",
    "outreach": {
        "email_subject": "KAM application — I work Talabat from the brand side",
        "email_body": (
            "<p>Hello,</p>"
            "<p>I'm applying for the Key Account Manager role. I currently integrate FMCG brands directly into Talabat, "
            "Noon, Careem and Deliveroo at DoFreeze — so I understand the partnership from the other side of the table.</p>"
            "<p>On the account side, I grew 42 key accounts <strong>+30% GMV QoQ</strong> at Alibaba's Miravia and managed "
            "strategic accounts (KFC, Taco Bell, Sushi Shop) at Glovo, leading cross-functional teams to grow order volume. "
            "I lead account reviews from data and dashboards, and I'm comfortable in Salesforce.</p>"
            "<p>I'm Dubai-based on a residence visa and can start immediately. I'd welcome a short conversation about your "
            "key restaurant partnerships.</p>"
            "<p>Best regards,<br>Paula De Francisco</p>"
        ),
        "linkedin_connection": (
            "Hi — I've applied for your Key Account Manager role. I integrate brands into Talabat/Noon/Careem/Deliveroo at "
            "DoFreeze and grew 42 accounts +30% GMV QoQ at Miravia (plus KAM at Glovo). Dubai-based. I'd love to connect."
        ),
        "linkedin_inmail": (
            "Hi — I've applied for the Key Account Manager position at Talabat. I work the platform from the brand side at "
            "DoFreeze (integrating brands into Talabat, Noon, Careem, Deliveroo) and grew 42 key accounts +30% GMV QoQ at "
            "Alibaba's Miravia, with strategic-account and cross-functional experience from Glovo. I'm Dubai-based on a "
            "residence visa and would value a short conversation about your restaurant partnerships."
        ),
    },
})

# -------------------------------------------------------------------
# 4. Huda Beauty — Assistant Manager, Global Brand Marketing   [TOP-4]
# -------------------------------------------------------------------
JOBS.append({
    "tier4": True,
    "deliverable_type": "brand_analysis",
    "apply_url": "https://to.indeed.com/aaswzkyrrdkf",
    "job": Job(
        id=job_hash("Assistant Manager, Global Brand Marketing", "Huda Beauty", "Dubai, UAE"),
        title="Assistant Manager, Global Brand Marketing",
        company="Huda Beauty",
        location="Dubai, UAE",
        url="https://to.indeed.com/aaswzkyrrdkf",
        source="indeed",
        description=(
            "Support development and execution of 360 marketing campaigns — messaging, paid media assets, "
            "in-store assets, e-commerce content, press materials and influencer activations for hero and "
            "sub-hero product launches. Coordinate cross-functional teams. Collaborate with Creative, Product "
            "and Commercial. Track campaign performance and post-campaign analysis. Support global teams. "
            "5+ years in a similar role; beauty industry an advantage. Passion for beauty and cosmetics a must."
        ),
    ),
    "analysis": JobAnalysis(score=85, tier="Hot",
        skills_match=["360 campaigns", "influencer marketing", "e-commerce content", "NPD launches", "omnichannel", "beauty"],
        missing_skills=[], sector_fit="Beauty & Fragrances", seniority_fit="Assistant Manager — brand+sector ideal",
        red_flags=[], ats_keywords=["brand marketing", "360 campaign", "influencer", "product launch", "e-commerce content", "beauty", "go-to-market"],
        reasoning="Beauty global brand — her sweet spot; influencer + launches match strongly."),
    "cv": {
        "headline": "Brand & Marketing Manager · Beauty & Fragrances · Influencer & E-Commerce",
        "professional_summary": (
            "Brand and marketing professional with 4+ years across beauty, fragrances, FMCG and e-commerce, "
            "passionate about building brands consumers love. Currently leading Brand & Marketing for DoFreeze, "
            "running 360 launches end-to-end and an influencer programme scaled from zero to 25–50 creators per "
            "campaign. Trained on beauty and fragrance at Alibaba's Miravia, with global, omnichannel reach."
        ),
        "experience": [
            dofreeze([
                "Lead NPD end-to-end for 6 product launches — brief, packaging, messaging, pricing and go-to-market — across 50+ markets",
                "Built and scaled the influencer marketing programme from zero, managing 25–50 influencers per campaign to drive awareness and sell-out",
                "Manage e-commerce content, UX and promotional mechanics across the store and UAE quick-commerce platforms (Noon, Talabat, Careem, Deliveroo)",
                "Develop omnichannel campaigns and analyse KPIs, ROI and ROAS across paid media, social and EDM",
            ]),
            miravia("Key Account Manager – Beauty, Fragrances & Fashion", [
                "Created and led the Beauty Club and Hot on Social projects, boosting brand visibility, loyalty and positioning Miravia as a beauty and lifestyle destination",
                "Managed 42 beauty, fragrance and fashion accounts, achieving +30% GMV growth QoQ through assortment and targeted promotions",
                "Led category expansion as PIC Fragrances, onboarding 30+ new stores in two months, including Arabian fragrances",
                "Analysed ROI, ROAS, conversion and retention to optimise campaign and channel performance",
            ]),
            glovo([
                "Onboarded fashion, beauty and lifestyle brands as Glovo built out its retail vertical beyond food delivery",
                "Managed strategic key accounts and drove GMV growth through data-led planning and bespoke marketing activations",
            ]),
            mondelez([
                "Evaluated promotional effectiveness and contributed to NPD launches including Milka Spread and Mini Suchard",
                "Conducted sell-in/sell-out analysis and built performance reports for the chocolate category",
            ]),
        ],
        "skills_brand": "brand strategy, 360 campaigns, influencer marketing, NPD end-to-end, go-to-market, shopper marketing, A&P budget management",
        "skills_ecommerce": "e-store management, e-commerce content, quick-commerce, Noon, Talabat, Careem, Deliveroo, Instagram, TikTok, EDM, UX optimisation",
        "skills_commercial": "key account management, category management, assortment planning, pricing strategy, modern trade, negotiation",
        "skills_data": "ROI, ROAS, KPI tracking, conversion analysis, P&L management, sell-in/sell-out",
        "skills_tools": "Canva, Meta Ads Manager, Google Ads, Salesforce, Looker, Kantar, Nielsen, Microsoft Office (Expert)",
    },
    "cl": {
        "opening_paragraph": (
            "Huda Beauty redefined what a beauty brand born in Dubai could become globally, and its purpose-led, "
            "community-first approach is precisely the kind of brand I want to help grow. With four years building "
            "beauty and fragrance brands across e-commerce and FMCG, supporting your Global Brand Marketing team "
            "would be a natural fit."
        ),
        "body_paragraph_1": (
            "This role lives on 360 campaigns, launches and influencer activations — my day-to-day. At DoFreeze I "
            "lead NPD end-to-end for six launches and built an influencer programme from zero to 25–50 creators per "
            "campaign. At Miravia (Alibaba) I created the Beauty Club and Hot on Social projects and grew 42 beauty, "
            "fragrance and fashion accounts +30% GMV QoQ, giving me a sharp sense of what converts in beauty."
        ),
        "body_paragraph_2": (
            "Beyond campaign execution, I bring e-commerce content and UAE quick-commerce fluency (Noon, Talabat, "
            "Careem, Deliveroo), a data-led approach to ROI and ROAS, and genuine passion for the category. I am "
            "already in Dubai on a UAE residence visa and work bilingually in Spanish and English — useful for a brand "
            "operating across global markets."
        ),
        "closing_paragraph": (
            "I would be thrilled to contribute to Huda Beauty's hero and sub-hero launches and global campaign "
            "calendar. I am available to interview whenever suits and can start without notice constraints. Thank you "
            "for considering my application."
        ),
    },
    "deliverable": {
        "title": "Brand & Launch Lens — Huda Beauty Global Campaigns",
        "subtitle": "Three observations on 360 launch execution for a community-first global beauty brand.",
        "sections": [
            {
                "heading": "Hero vs. sub-hero launch architecture",
                "finding": "A 360 calendar with both hero and sub-hero launches competes for the same creative, paid and influencer resources.",
                "insight": "Sub-hero launches often underperform not because of the product but because they inherit a scaled-down hero playbook instead of a purpose-built, lower-cost motion.",
                "recommendation": "Define a distinct sub-hero playbook (tighter creator set, owned + e-commerce content led, lighter paid) so sub-heroes get a fit-for-purpose launch — I run exactly this tiered launch logic across six NPD launches at DoFreeze.",
            },
            {
                "heading": "Influencer activation as the conversion layer, not just awareness",
                "finding": "Beauty lives on creators, but influencer activity is frequently measured on reach rather than on sell-out and e-commerce conversion.",
                "insight": "Tying creator briefs to specific launch SKUs and tracking through to e-commerce conversion turns influencer spend from awareness into measurable revenue.",
                "recommendation": "Brief creators against launch SKUs with trackable links and post-campaign sell-out analysis — I built an influencer programme from zero to 25–50 creators per campaign at DoFreeze and led the Beauty Club / Hot on Social projects at Miravia.",
            },
            {
                "heading": "Consistency across a global, cross-functional calendar",
                "finding": "Hero and sub-hero campaigns ship across messaging, paid, in-store, e-commerce and PR — coordinated with Creative, Product and Commercial and global teams.",
                "insight": "The biggest execution risk is drift: the same launch looking and saying different things across touchpoints and markets.",
                "recommendation": "Run each launch from a single brief and asset tracker with a cross-functional check before go-live — the cross-functional coordination I do daily leading omnichannel campaigns at DoFreeze.",
            },
        ],
        "closing": (
            "These are outside-in observations from a beauty marketer who admires the brand; with Huda Beauty's campaign "
            "data I'd sharpen them fast. They reflect how I build launches — tiered playbooks, conversion-tracked influencer, and one consistent brief."
        ),
    },
    "outreach_mode": "applied",
    "outreach": {
        "email_subject": "Global Brand Marketing application — beauty launches",
        "email_body": (
            "<p>Hello,</p>"
            "<p>I'm applying for the Assistant Manager, Global Brand Marketing role. 360 launches and influencer "
            "activations are my day-to-day: at DoFreeze I lead NPD end-to-end for six launches and built an influencer "
            "programme from zero to <strong>25–50 creators per campaign</strong>.</p>"
            "<p>At Alibaba's Miravia I created the Beauty Club and Hot on Social projects and grew 42 beauty, fragrance and "
            "fashion accounts +30% GMV QoQ — so I have a sharp sense of what actually converts in beauty, not just what reaches.</p>"
            "<p>I'm Dubai-based on a residence visa and can start immediately. I'd love to contribute to Huda Beauty's hero "
            "and sub-hero launch calendar.</p>"
            "<p>Best regards,<br>Paula De Francisco</p>"
        ),
        "linkedin_connection": (
            "Hi — I've applied for your Assistant Manager, Global Brand Marketing role. I lead 360 beauty launches and a "
            "25–50 creator influencer programme at DoFreeze, after building beauty projects at Alibaba's Miravia. Dubai-based. Let's connect."
        ),
        "linkedin_inmail": (
            "Hi — I've applied for the Assistant Manager, Global Brand Marketing position. I lead 360 launches and an "
            "influencer programme (25–50 creators per campaign) at DoFreeze, and at Alibaba's Miravia I created the Beauty "
            "Club and Hot on Social projects while growing 42 beauty/fragrance accounts +30% GMV QoQ. I'm Dubai-based on a "
            "residence visa and would value a short conversation about Huda Beauty's global launch calendar."
        ),
    },
})

# -------------------------------------------------------------------
# 5. Justlife — Senior Growth Marketing Manager   [CV + CL only]
# -------------------------------------------------------------------
JOBS.append({
    "tier4": False,
    "apply_url": "https://to.indeed.com/aafhg7dbtk7t",
    "job": Job(
        id=job_hash("Senior Growth Marketing Manager", "Justlife", "Dubai, UAE"),
        title="Senior Growth Marketing Manager",
        company="Justlife",
        location="Dubai, UAE",
        url="https://to.indeed.com/aafhg7dbtk7t",
        source="indeed",
        description=(
            "Justlife, the GCC's leading super app for home services, seeks a hands-on Senior Growth "
            "Marketing Manager to co-own demand generation across UAE and KSA. Personally execute "
            "performance campaigns; treat CAC, LTV and incrementality as product surfaces. Design and run "
            "holdout tests, geo experiments and lift studies. App-first / marketplace growth. AI in "
            "production — design or deploy automation that improves performance outcomes. 5+ years hands-on "
            "performance marketing. SQL/Python a plus."
        ),
    ),
    "analysis": JobAnalysis(score=80, tier="Hot",
        skills_match=["performance marketing", "Meta Ads", "Google Ads", "ROAS", "AI automation", "quick-commerce", "marketplace growth"],
        missing_skills=["incrementality testing (geo/holdout) not explicit in profile"],
        sector_fit="Quick-commerce / Super app", seniority_fit="Senior — strong fit on hands-on + AI",
        red_flags=[], ats_keywords=["growth marketing", "performance", "CAC", "LTV", "ROAS", "AI automation", "Meta Ads", "marketplace"],
        reasoning="Her AI-first differentiator + paid media (Meta/Google) + quick-commerce marketplace fit Justlife's hands-on growth role."),
    "cv": {
        "headline": "Senior Growth & Performance Marketing · AI-First · Quick-Commerce & Marketplace",
        "professional_summary": (
            "Growth and performance marketer with 4+ years across quick-commerce, e-commerce and FMCG, and an "
            "AI-first operating style. At DoFreeze she built an AI-powered marketing automation system (Claude / "
            "generative AI) cutting manual workload ~40%, and runs Meta and Google Ads to ROI/ROAS. Marketplace "
            "growth credentials from Alibaba's Miravia (+30% GMV QoQ) and Glovo's quick-commerce platform."
        ),
        "experience": [
            dofreeze([
                "Built an AI-powered marketing automation system (Claude / generative AI) automating campaign planning, content, market research and KPI reporting — cutting manual workload ~40%",
                "Plan and optimise paid media on Meta Ads (Facebook & Instagram) and Google Ads — audience building, creative A/B testing, social and EDM — analysing ROI and ROAS for continuous improvement",
                "Own and optimise the Shopify e-store end-to-end (UX, collections, discounts, checkout), lifting conversion rate (CRO) and average order value through data-led merchandising",
                "Integrated brands into UAE quick-commerce (Noon, Talabat, Careem, Deliveroo), managing onboarding, listings and promotional mechanics",
            ]),
            miravia("Key Account Manager – Beauty, Fragrances & Fashion", [
                "Grew 42 accounts +30% GMV QoQ on a top-5 global e-commerce platform through pricing, assortment and promotion strategy",
                "Continuously analysed ROI, ROAS, conversion, traffic and retention to optimise channel performance and forecasting accuracy",
                "Owned the Flash Sales channel for Beauty, Fashion & Home, executing commercial plans aligned to P&L targets",
            ]),
            glovo([
                "Drove GMV growth on a quick-commerce marketplace through data-led planning and bespoke activations",
                "Part of the team that built Glovo's retail vertical, onboarding partners and expanding the marketplace beyond food delivery",
            ]),
            mondelez([
                "Conducted sell-in/sell-out analysis and built data-led performance reports",
            ]),
        ],
        "skills_brand": "generative AI campaigns, marketing automation, go-to-market, omnichannel campaigns, creative A/B testing",
        "skills_ecommerce": "performance marketing, Meta Ads, Google Ads, conversion rate optimisation (CRO), Shopify, quick-commerce, marketplace growth, EDM, UX optimisation",
        "skills_commercial": "key account management, pricing strategy, assortment planning, customer acquisition & retention",
        "skills_data": "CAC/LTV thinking, ROAS, ROI, conversion & retention analysis, KPI tracking, P&L management, AI-assisted analysis & forecasting, Looker",
        "skills_tools": "Generative AI (Claude, ChatGPT), Meta Ads Manager, Google Ads, Shopify, Looker, Salesforce, Microsoft Office (Expert)",
    },
    "cl": {
        "opening_paragraph": (
            "Justlife's bet — that the next 15 million sessions come from AI-enabled performance systems, not more of "
            "the same — is exactly how I work. I build generative-AI automation into marketing and run paid media "
            "hands-on, so co-owning growth across UAE and KSA is a role I'd be genuinely excited about."
        ),
        "body_paragraph_1": (
            "I'm an operator, not a deck-from-a-distance strategist. At DoFreeze I built an AI-powered marketing "
            "automation system (Claude) that cut manual workload ~40%, and I personally run Meta and Google Ads — "
            "audience building, creative A/B testing — analysing ROI and ROAS to improve performance. I also own a "
            "Shopify store end-to-end, lifting conversion and AOV through data-led merchandising."
        ),
        "body_paragraph_2": (
            "On the marketplace side, I grew 42 accounts +30% GMV QoQ at Alibaba's Miravia and drove GMV on Glovo's "
            "quick-commerce platform, so I understand install-to-repeat and marketplace economics. I'm candid that "
            "formal incrementality testing (geo/holdout) is where I'd ramp fastest — but the analytical instinct and "
            "AI-in-production track record are already there. I'm Dubai-based on a residence visa and bilingual."
        ),
        "closing_paragraph": (
            "I'd welcome a conversation about co-owning Justlife's growth engine. I'm available to interview at your "
            "convenience and can start immediately. Thank you for your consideration."
        ),
    },
})

# -------------------------------------------------------------------
# 6. PlaceUp — Category Manager   [CV + CL only]
# -------------------------------------------------------------------
JOBS.append({
    "tier4": False,
    "apply_url": "https://to.indeed.com/aal6k8y2yp6w",
    "job": Job(
        id=job_hash("Category Manager", "PlaceUp", "Dubai, UAE"),
        title="Category Manager",
        company="PlaceUp",
        location="Dubai, UAE",
        url="https://to.indeed.com/aal6k8y2yp6w",
        source="indeed",
        description=(
            "Manage and grow assigned product categories across FMCG retail and e-commerce. Develop category "
            "plans for sales, margin and market share across traditional retail and e-commerce. Optimise "
            "assortment, pricing and promotions offline and online (marketplaces, D2C, e-retailers). Coordinate "
            "with e-commerce teams on listings, content and digital merchandising. Ensure brand visibility at POS "
            "via planograms and displays. Negotiate with suppliers. 4-6 years in category management, trade "
            "marketing or product management in FMCG and/or e-commerce."
        ),
    ),
    "analysis": JobAnalysis(score=78, tier="Hot",
        skills_match=["category management", "trade marketing", "assortment", "pricing", "promotions", "FMCG", "e-commerce", "planogram"],
        missing_skills=[], sector_fit="FMCG / Consumer Goods", seniority_fit="4-6 yrs — good fit",
        red_flags=[], ats_keywords=["category management", "assortment", "pricing", "promotions", "planogram", "FMCG", "e-commerce", "market share", "supplier negotiation"],
        reasoning="Direct target title; FMCG + e-commerce category role matching Mondelez + Miravia background."),
    "cv": {
        "headline": "Category Manager · FMCG & E-Commerce · Trade & Shopper Marketing",
        "professional_summary": (
            "Category and commercial professional with 4+ years across FMCG, beauty and e-commerce, blending "
            "category planning at Mondelez with key-account growth at Alibaba's Miravia (+30% GMV QoQ across 42 "
            "accounts). Now leads trade, assortment and channel plans at DoFreeze across 50+ markets, fluent in both "
            "modern trade and online marketplaces."
        ),
        "experience": [
            dofreeze([
                "Develop trade and shopper marketing plans by channel, supporting category management and key-account strategy across 50+ markets",
                "Manage assortment, pricing and promotional mechanics across the e-store and UAE quick-commerce (Noon, Talabat, Careem, Deliveroo)",
                "Lead NPD end-to-end (brief, pricing, go-to-market) and own A&P budget allocation with ROI focus",
            ]),
            miravia("Key Account Manager – Beauty, Fragrances & Fashion", [
                "Led category expansion as PIC Fragrances, onboarding 30+ new stores in two months via trend-driven assortment and promotions",
                "Grew 42 accounts +30% GMV QoQ through pricing strategy, assortment optimisation and targeted promotions",
                "Analysed conversion, ROI and ROAS to optimise category performance and forecasting on a top-5 e-commerce platform",
            ]),
            mondelez([
                "Conducted sell-in/sell-out analysis and evaluated promotional effectiveness for the chocolate category",
                "Identified growth opportunities and contributed to NPD launches including Milka Spread and Mini Suchard",
                "Built category performance reports informing assortment and trade decisions",
            ]),
            glovo([
                "Managed strategic key accounts and drove GMV growth through data-led planning",
                "Helped expand the retail marketplace beyond food into new categories",
            ]),
        ],
        "skills_brand": "trade marketing, shopper marketing, go-to-market, NPD end-to-end, promotional planning",
        "skills_ecommerce": "online marketplaces, e-store management, listings & content, digital merchandising, quick-commerce, Noon, Talabat, Careem, Deliveroo",
        "skills_commercial": "category management, assortment planning, pricing strategy, planogram, supplier negotiation, key account management, modern trade, distributor management, forecasting",
        "skills_data": "sell-in/sell-out, market share, ROI, ROAS, KPI tracking, P&L management, Nielsen, Kantar",
        "skills_tools": "Nielsen, Kantar, Looker, SAP, Salesforce, Planorama, Microsoft Office (Expert)",
    },
    "cl": {
        "opening_paragraph": (
            "Growing categories across both FMCG retail and e-commerce is precisely where my experience sits — from "
            "category planning at Mondelez to fragrance category expansion at Alibaba's Miravia. PlaceUp's dual "
            "offline-and-online focus is a strong match for how I work."
        ),
        "body_paragraph_1": (
            "I build category and assortment plans, optimise pricing and promotions, and drive market share. At "
            "Miravia I led category expansion as PIC Fragrances, onboarding 30+ stores in two months and growing 42 "
            "accounts +30% GMV QoQ. At Mondelez I owned sell-in/sell-out analysis and promotional effectiveness for "
            "the chocolate category, and at DoFreeze I manage assortment and trade plans across 50+ markets."
        ),
        "body_paragraph_2": (
            "I bring both sides PlaceUp needs: modern-trade and planogram fundamentals plus hands-on online "
            "marketplace and UAE quick-commerce experience (Noon, Talabat, Careem, Deliveroo). I am Dubai-based on a "
            "UAE residence visa, data-led on ROI and market share, and bilingual in Spanish and English."
        ),
        "closing_paragraph": (
            "I would welcome the opportunity to discuss how I can grow PlaceUp's categories across retail and "
            "e-commerce. I am available to interview at your convenience and can start immediately. Thank you for your "
            "consideration."
        ),
    },
})

# -------------------------------------------------------------------
# 7. PureBorn — Trade Marketing Manager   [CV + CL only]
# -------------------------------------------------------------------
JOBS.append({
    "tier4": False,
    "apply_url": "https://to.indeed.com/aa6bkkhwb8h6",
    "job": Job(
        id=job_hash("Trade Marketing Manager", "PureBorn", "Jebel Ali, UAE"),
        title="Trade Marketing Manager",
        company="PureBorn",
        location="Jebel Ali, UAE",
        url="https://to.indeed.com/aa6bkkhwb8h6",
        source="indeed",
        description=(
            "PureBorn (baby care) seeks a Trade Marketing Manager to develop and execute trade marketing plans "
            "across modern and traditional trade. Drive in-store visibility, promotions and category growth. Work "
            "with sales and distributors. Manage trade budgets and analyse ROI. Support NPD listings and "
            "go-to-market. FMCG trade marketing experience required."
        ),
    ),
    "analysis": JobAnalysis(score=74, tier="Warm",
        skills_match=["trade marketing", "modern trade", "promotions", "distributor management", "NPD", "ROI", "FMCG"],
        missing_skills=[], sector_fit="FMCG (baby care)", seniority_fit="Manager — good fit",
        red_flags=["posting is a few months old — may be filled"], ats_keywords=["trade marketing", "modern trade", "promotions", "distributor", "NPD", "ROI", "FMCG", "category"],
        reasoning="Pure FMCG trade marketing — direct match to her DoFreeze trade & shopper work; posting older so freshness flagged."),
    "cv": {
        "headline": "Trade Marketing Manager · FMCG · Modern Trade · Distributor & Category",
        "professional_summary": (
            "Trade marketing professional with 4+ years across FMCG, beauty and e-commerce, currently leading trade "
            "and shopper plans for DoFreeze across 50+ markets. Hands-on with modern trade, distributor networks, "
            "promotions and NPD go-to-market, with a category-planning foundation from Mondelez and +30% GMV QoQ "
            "across 42 accounts at Alibaba's Miravia."
        ),
        "experience": [
            dofreeze([
                "Develop trade and shopper marketing plans by channel across 50+ markets, driving in-store visibility, promotions and category growth with the sales team",
                "Manage distributor networks and modern-trade execution, integrating brands into UAE retail and quick-commerce (Noon, Talabat, Careem, Deliveroo)",
                "Own A&P / trade budget allocation, analysing ROI to maximise promotional efficiency",
                "Lead NPD end-to-end (listings, pricing, go-to-market) for 6 launches across GCC, MENA and beyond",
            ]),
            miravia("Key Account Manager – Beauty, Fragrances & Fashion", [
                "Grew 42 accounts +30% GMV QoQ through pricing strategy, assortment optimisation and targeted promotions",
                "Led category expansion as PIC Fragrances, onboarding 30+ new stores in two months",
                "Analysed ROI, ROAS and conversion to optimise promotional and category performance",
            ]),
            mondelez([
                "Conducted sell-in/sell-out analysis and evaluated promotional effectiveness for the chocolate category",
                "Contributed to NPD launches including Milka Spread and Mini Suchard",
                "Built performance reports informing trade and category decisions",
            ]),
            glovo([
                "Managed strategic key accounts and drove GMV growth through data-led planning and bespoke activations",
            ]),
        ],
        "skills_brand": "trade marketing, shopper marketing, in-store visibility, promotional planning, go-to-market, NPD end-to-end",
        "skills_ecommerce": "quick-commerce, Noon, Talabat, Careem, Deliveroo, retail execution, listings & promotions",
        "skills_commercial": "modern trade, general trade, distributor management, category management, assortment planning, pricing strategy, planogram, negotiation, forecasting",
        "skills_data": "sell-in/sell-out, ROI, ROAS, KPI tracking, P&L management, Nielsen, Kantar",
        "skills_tools": "SAP, Salesforce, Nielsen, Kantar, Looker, Planorama, Microsoft Office (Expert)",
    },
    "cl": {
        "opening_paragraph": (
            "Trade marketing in FMCG is the core of what I do every day, so PureBorn's focus on in-store visibility, "
            "promotions and category growth across modern and traditional trade is a natural fit. As Brand & Marketing "
            "Manager at a Dubai-based FMCG distributor, I'd bring hands-on GCC trade experience to your baby-care portfolio."
        ),
        "body_paragraph_1": (
            "I build trade plans by channel, manage distributor execution and run promotions to ROI. At DoFreeze I do "
            "this across 50+ markets and lead NPD listings and go-to-market end-to-end. At Mondelez I owned "
            "sell-in/sell-out analysis and promotional effectiveness, and at Miravia I grew 42 accounts +30% GMV QoQ "
            "through pricing, assortment and promotions."
        ),
        "body_paragraph_2": (
            "I bring modern-trade and planogram fundamentals plus distributor-network management, and I'm data-led on "
            "trade ROI. I'm already based in Dubai on a UAE residence visa — no relocation or sponsorship needed — and "
            "work bilingually in Spanish and English."
        ),
        "closing_paragraph": (
            "I'd welcome the chance to discuss how I can grow PureBorn's categories across trade channels. I'm available "
            "to interview at your convenience and can start immediately. Thank you for your consideration."
        ),
    },
})

# -------------------------------------------------------------------
# 8. Farzana Trading — E-Commerce Manager   [CV + CL only]
# -------------------------------------------------------------------
JOBS.append({
    "tier4": False,
    "apply_url": "https://to.indeed.com/aar67s6zvtrv",
    "job": Job(
        id=job_hash("E-Commerce Manager", "Farzana Trading", "Dubai, UAE"),
        title="E-Commerce Manager",
        company="Farzana Trading",
        location="Dubai, UAE",
        url="https://to.indeed.com/aar67s6zvtrv",
        source="indeed",
        description=(
            "Manage and grow the company's e-commerce presence. Own online catalogue, listings, content and "
            "promotions across the e-store and marketplaces. Drive traffic, conversion and revenue. Coordinate "
            "digital marketing, paid media and merchandising. Analyse performance and optimise. Experience "
            "managing e-commerce / online retail required."
        ),
    ),
    "analysis": JobAnalysis(score=66, tier="Warm",
        skills_match=["e-commerce management", "Shopify", "CRO", "listings & content", "paid media", "marketplace"],
        missing_skills=["company scope/scale unclear"], sector_fit="E-commerce / Trading",
        seniority_fit="Manager — target title match", red_flags=["small/unknown company — scope to confirm"],
        ats_keywords=["e-commerce", "Shopify", "conversion", "listings", "paid media", "marketplace", "digital merchandising"],
        reasoning="Her exact target title; strong e-commerce/Shopify fit though company scope is unknown."),
    "cv": {
        "headline": "E-Commerce Manager · Shopify & Marketplaces · CRO & Paid Media",
        "professional_summary": (
            "E-commerce professional with 4+ years scaling online channels across beauty, fragrance and FMCG. "
            "Owns the DoFreeze Shopify store end-to-end — catalogue, UX, conversion and promotions — and grew 42 "
            "accounts +30% GMV QoQ on Alibaba's Miravia marketplace. Data-led on paid media ROI/ROAS with deep UAE "
            "quick-commerce experience."
        ),
        "experience": [
            dofreeze([
                "Own and optimise the Shopify e-store end-to-end (catalogue, UX, collections, discounts, checkout), lifting conversion rate (CRO) and average order value through data-led merchandising",
                "Plan and optimise paid media on Meta Ads and Google Ads, analysing ROI and ROAS for continuous performance improvement",
                "Integrated brands into UAE quick-commerce (Noon, Talabat, Careem, Deliveroo), managing listings, content and promotional mechanics",
                "Built AI-powered automation (Claude) for content, campaign planning and reporting, cutting manual workload ~40%",
            ]),
            miravia("Key Account Manager (E-Commerce) – Beauty, Fragrances & Fashion", [
                "Grew 42 accounts +30% GMV QoQ on a top-5 global e-commerce platform through pricing, assortment and promotions",
                "Owned the Flash Sales channel for Beauty, Fashion & Home, executing commercial plans to P&L targets",
                "Analysed conversion, traffic, retention, ROI and ROAS to optimise channel performance",
            ]),
            glovo([
                "Drove GMV growth on a quick-commerce platform through data-led planning and bespoke activations",
                "Helped build the retail vertical, onboarding brands beyond food delivery",
            ]),
            massimo([
                "Inditex retail operations and visual merchandising foundation underpinning an omnichannel view of retail",
            ]),
        ],
        "skills_brand": "digital merchandising, promotional planning, go-to-market, omnichannel campaigns, generative AI campaigns",
        "skills_ecommerce": "Shopify, e-store management, conversion rate optimisation (CRO), UX optimisation, online marketplaces, Meta Ads, Google Ads, EDM, quick-commerce, Noon, Talabat, Careem, Deliveroo",
        "skills_commercial": "pricing strategy, assortment planning, key account management, customer acquisition & retention",
        "skills_data": "ROAS, ROI, conversion & traffic analysis, KPI tracking, P&L management, Looker",
        "skills_tools": "Shopify, Meta Ads Manager, Google Ads, Looker, Salesforce, Generative AI (Claude), Microsoft Office (Expert)",
    },
    "cl": {
        "opening_paragraph": (
            "Running an e-commerce operation end-to-end — catalogue, conversion, paid media and promotions — is exactly "
            "what I do at DoFreeze today, so Farzana Trading's E-Commerce Manager role caught my attention. I'd be glad "
            "to bring that hands-on online-retail experience to your business."
        ),
        "body_paragraph_1": (
            "I own a Shopify store end-to-end, lifting conversion and average order value through data-led "
            "merchandising, and I run Meta and Google Ads to ROI and ROAS. At Alibaba's Miravia I grew 42 accounts "
            "+30% GMV QoQ on a top-5 marketplace, optimising conversion, traffic and retention — so I'm fluent across "
            "both own-store and marketplace e-commerce."
        ),
        "body_paragraph_2": (
            "I'm also Google E-Commerce and Digital Marketing certified, fluent in UAE quick-commerce (Noon, Talabat, "
            "Careem, Deliveroo), and I build AI automation that scales content and reporting. I'm Dubai-based on a UAE "
            "residence visa and bilingual in Spanish and English."
        ),
        "closing_paragraph": (
            "I'd welcome a conversation about growing Farzana Trading's e-commerce revenue. I'm available to interview "
            "at your convenience and can start immediately. Thank you for your consideration."
        ),
    },
})

# -------------------------------------------------------------------
# 9. DeFi Technologies — Marketing & Brand Growth Manager   [CV + CL only]
# -------------------------------------------------------------------
JOBS.append({
    "tier4": False,
    "apply_url": "https://to.indeed.com/aaxmk8g9fvgs",
    "job": Job(
        id=job_hash("Marketing & Brand Growth Manager", "DeFi Technologies LLC", "Dubai, UAE"),
        title="Marketing & Brand Growth Manager",
        company="DeFi Technologies LLC",
        location="Dubai, UAE",
        url="https://to.indeed.com/aaxmk8g9fvgs",
        source="indeed",
        description=(
            "Lead brand growth and marketing for a fast-scaling business in Dubai. Own brand strategy, "
            "positioning and go-to-market. Plan and run digital and social campaigns, paid media and content. "
            "Build the brand's audience and community. Analyse performance and optimise for growth. Manage "
            "marketing budget. Experience leading brand and growth marketing required."
        ),
    ),
    "analysis": JobAnalysis(score=58, tier="Warm",
        skills_match=["brand strategy", "go-to-market", "paid media", "content", "community", "generative AI"],
        missing_skills=["crypto/DeFi sector experience"], sector_fit="Crypto / Fintech (outside core sweet spot)",
        seniority_fit="Manager — title fit, sector stretch", red_flags=["sector outside her FMCG/beauty/e-commerce focus"],
        ats_keywords=["brand growth", "brand strategy", "go-to-market", "paid media", "content", "community", "digital marketing"],
        reasoning="Brand & growth title fits, but crypto sector is outside her core sectors — transferable skills, weaker sector match."),
    "cv": {
        "headline": "Marketing & Brand Growth Manager · Brand Strategy · Digital & Paid Media · AI-First",
        "professional_summary": (
            "Brand and growth marketer with 4+ years building brands and scaling digital channels across FMCG, beauty "
            "and e-commerce. Leads brand strategy, go-to-market and paid media at DoFreeze, where she built AI-powered "
            "marketing automation cutting manual workload ~40%. Strong on audience building, content and performance "
            "analysis, with a +30% GMV QoQ commercial track record."
        ),
        "experience": [
            dofreeze([
                "Own brand strategy, positioning and go-to-market for a multi-brand portfolio across 50+ markets",
                "Plan and run digital and social campaigns and paid media (Meta Ads, Google Ads), analysing ROI and ROAS to drive growth",
                "Built and scaled an influencer / community programme from zero to 25–50 creators per campaign, driving awareness and UGC",
                "Built AI-powered marketing automation (Claude) for content, campaign planning and analytics, cutting manual workload ~40%",
            ]),
            miravia("Key Account Manager – Beauty, Fragrances & Fashion", [
                "Created the Beauty Club and Hot on Social projects, building audience, loyalty and brand positioning",
                "Grew 42 accounts +30% GMV QoQ through pricing, assortment and targeted promotions",
                "Analysed ROI, ROAS, conversion and retention to optimise channel performance",
            ]),
            glovo([
                "Drove GMV growth through data-led planning and bespoke marketing activations",
                "Helped build and grow a new marketplace vertical",
            ]),
            mondelez([
                "Evaluated promotional effectiveness and contributed to NPD launches",
            ]),
        ],
        "skills_brand": "brand strategy, brand positioning, go-to-market, audience & community building, influencer marketing, content, generative AI campaigns",
        "skills_ecommerce": "paid media, Meta Ads, Google Ads, social, EDM, conversion optimisation, marketing automation, Shopify",
        "skills_commercial": "pricing strategy, key account management, customer acquisition & retention, negotiation",
        "skills_data": "ROI, ROAS, KPI tracking, conversion & retention analysis, P&L management, AI-assisted analysis",
        "skills_tools": "Generative AI (Claude, ChatGPT), Meta Ads Manager, Google Ads, Shopify, Canva, Looker, Microsoft Office (Expert)",
    },
    "cl": {
        "opening_paragraph": (
            "Building a brand and scaling its growth from the ground up — strategy, audience, content and paid media — "
            "is what I do at DoFreeze, so DeFi Technologies' Marketing & Brand Growth Manager role is an interesting "
            "fit. I'd bring a hands-on, AI-first approach to growing your brand."
        ),
        "body_paragraph_1": (
            "I own brand strategy, positioning and go-to-market, and I run digital, social and paid-media campaigns to "
            "ROI and ROAS. I built an audience-and-community programme from zero to 25–50 creators per campaign at "
            "DoFreeze, and at Alibaba's Miravia I created the Beauty Club and Hot on Social projects while growing 42 "
            "accounts +30% GMV QoQ."
        ),
        "body_paragraph_2": (
            "My differentiator is being AI-first: I built marketing automation with Claude that cut manual workload "
            "~40%, scaling content, campaign planning and analytics. While my background sits in FMCG, beauty and "
            "e-commerce rather than crypto, the brand-growth playbook is highly transferable. I'm Dubai-based on a UAE "
            "residence visa and bilingual in Spanish and English."
        ),
        "closing_paragraph": (
            "I'd welcome a conversation about how I can build and grow DeFi Technologies' brand. I'm available to "
            "interview at your convenience and can start immediately. Thank you for your consideration."
        ),
    },
})

# -------------------------------------------------------------------
# 10. Solutions Leisure Group — Assistant Marketing Manager   [CV + CL only]
# -------------------------------------------------------------------
JOBS.append({
    "tier4": False,
    "apply_url": "https://to.indeed.com/aavx8l4rjzhy",
    "job": Job(
        id=job_hash("Assistant Marketing Manager", "Solutions Leisure Group", "Dubai, UAE"),
        title="Assistant Marketing Manager",
        company="Solutions Leisure Group",
        location="Dubai, UAE",
        url="https://to.indeed.com/aavx8l4rjzhy",
        source="indeed",
        description=(
            "Solutions Leisure Group (hospitality / F&B and lifestyle venues) seeks an Assistant Marketing "
            "Manager to support brand and marketing across its venues. Plan and execute campaigns, social media, "
            "influencer activations, events and promotions. Coordinate with creative and operations. Track "
            "performance. Marketing experience, ideally in hospitality / lifestyle, required."
        ),
    ),
    "analysis": JobAnalysis(score=55, tier="Warm",
        skills_match=["campaign management", "social media", "influencer marketing", "events & promotions", "brand"],
        missing_skills=["hospitality/F&B sector experience"], sector_fit="Hospitality / F&B / Lifestyle",
        seniority_fit="Assistant Manager — slight step down", red_flags=["sector adjacent, not core; assistant level"],
        ats_keywords=["marketing", "brand", "social media", "influencer", "events", "promotions", "campaigns"],
        reasoning="Brand/campaign/influencer skills transfer to lifestyle marketing, but hospitality sector and assistant level make it a softer fit."),
    "cv": {
        "headline": "Marketing Manager · Brand, Social & Influencer · Campaigns & Activations",
        "professional_summary": (
            "Brand and marketing professional with 4+ years running campaigns, social and influencer activations "
            "across beauty, FMCG and e-commerce. At DoFreeze she leads omnichannel campaigns and built an influencer "
            "programme from zero to 25–50 creators per campaign; at Alibaba's Miravia she created social-led brand "
            "projects. Creative and data-led, with a strong sense of what drives engagement and footfall."
        ),
        "experience": [
            dofreeze([
                "Develop and run omnichannel campaigns across social, paid media (Meta Ads) and EDM, analysing engagement and ROI",
                "Built and scaled an influencer marketing programme from zero — sourcing, briefing and managing 25–50 creators per campaign — driving awareness and UGC",
                "Plan promotions and activations across markets, coordinating creative and execution",
                "Lead NPD go-to-market and brand campaigns end-to-end across 50+ markets",
            ]),
            miravia("Key Account Manager – Beauty, Fragrances & Fashion", [
                "Created and led the Beauty Club and Hot on Social projects, boosting brand visibility, engagement and loyalty",
                "Ran targeted promotions and activations, growing 42 accounts +30% GMV QoQ",
                "Analysed conversion, engagement and ROI to optimise campaign performance",
            ]),
            glovo([
                "Delivered bespoke marketing activations with strategic partners, driving order volume and engagement",
                "Coordinated cross-functional teams across marketing, logistics and operations",
            ]),
            massimo([
                "Customer-facing premium retail and brand-experience foundation at an Inditex flagship outlet",
            ]),
        ],
        "skills_brand": "brand marketing, campaign management, influencer marketing, social media, events & activations, promotional planning, go-to-market",
        "skills_ecommerce": "Meta Ads, Instagram, TikTok, EDM, social content, e-store promotions",
        "skills_commercial": "key account management, partner activations, promotions, negotiation",
        "skills_data": "engagement & ROI tracking, ROAS, KPI tracking, conversion analysis",
        "skills_tools": "Canva, Meta Ads Manager, Meta Business Suite, Google Ads, Instagram, TikTok, Microsoft Office (Expert)",
    },
    "cl": {
        "opening_paragraph": (
            "Solutions Leisure Group's portfolio of standout Dubai venues lives on brand, social buzz and memorable "
            "activations — the kind of marketing I most enjoy. With four years running campaigns and influencer "
            "programmes, I'd be glad to support your marketing team across the group's brands."
        ),
        "body_paragraph_1": (
            "Campaigns, social and influencer activations are my day-to-day. At DoFreeze I run omnichannel campaigns "
            "and built an influencer programme from zero to 25–50 creators per campaign, and at Alibaba's Miravia I "
            "created the Beauty Club and Hot on Social projects to drive engagement and loyalty. I plan promotions and "
            "activations and coordinate creative and execution end-to-end."
        ),
        "body_paragraph_2": (
            "While my background is in beauty, FMCG and e-commerce rather than hospitality, the brand, social and "
            "activation skills transfer directly to lifestyle and F&B marketing. I'm a creative, hands-on operator, "
            "Dubai-based on a UAE residence visa, and bilingual in Spanish and English."
        ),
        "closing_paragraph": (
            "I'd welcome the chance to discuss how I can support Solutions Leisure Group's brands and activations. I'm "
            "available to interview at your convenience and can start immediately. Thank you for your consideration."
        ),
    },
})


# ===================================================================
# RUN
# ===================================================================

def run():
    results = []
    for spec in JOBS:
        job = spec["job"]
        analysis = spec["analysis"]
        logger.info("=== %s @ %s ===", job.title, job.company)

        rec = {"company": job.company, "title": job.title, "cv": None,
               "cl": None, "deliverable": None, "outreach": None}

        try:
            with llm.override({"submit_cv_content": spec["cv"]}):
                rec["cv"] = generate_cv(job, analysis)
        except Exception as exc:
            logger.error("CV failed for %s: %s", job.company, exc)

        try:
            with llm.override({"submit_cover_letter": spec["cl"]}):
                rec["cl"] = generate_cover_letter(job, analysis)
        except Exception as exc:
            logger.error("CL failed for %s: %s", job.company, exc)

        if spec.get("tier4"):
            try:
                with llm.override({"submit_deliverable": spec["deliverable"]}):
                    rec["deliverable"] = generate_deliverable(
                        job, analysis, deliverable_type=spec["deliverable_type"],
                    )
            except Exception as exc:
                logger.error("Deliverable failed for %s: %s", job.company, exc)

            try:
                with llm.override({"submit_outreach": spec["outreach"]}):
                    content = generate_outreach(job, analysis)
                if content:
                    rec["outreach"] = build_outreach_pack(job, content)
            except Exception as exc:
                logger.error("Outreach failed for %s: %s", job.company, exc)

        results.append(rec)

    print("\n\n========== RESULT ==========")
    for r in results:
        print(f"\n{r['company']} — {r['title']}")
        print(f"  CV:          {r['cv']}")
        print(f"  CL:          {r['cl']}")
        if r["deliverable"] or r["outreach"]:
            print(f"  Deliverable: {r['deliverable']}")
            print(f"  Outreach:    {r['outreach']}")
    return results


if __name__ == "__main__":
    run()

