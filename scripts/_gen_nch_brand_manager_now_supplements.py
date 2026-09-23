"""One-off: CV de una página para **Brand Manager, NOW Nutritional Supplements** en
**New Country Healthcare** (NCH, distribuidor sanitario de Dubái fundado en 1978, ~209
empleados en LinkedIn, +1.500 productos). Dubái, presencial, jornada completa, Solicitud
sencilla. Publicado por Anas Jannoud (Head of HR). Ya con 141 solicitudes el primer día.

El JD pide: estrategia y plan anual de NOW (nowfoods.com) en el GCC; **accountability directa
del P&L de marca** (revenue, gross margin, inversión de marketing, trade spend, rentabilidad);
estrategia de distribución por canal y cuentas clave; presupuesto de marketing y trade con
governance y compliance; activación omnicanal (medical, trade, digital); **forecasting y demand
planning** con supply chain, salud del inventario y near-expiry; **field excellence** (medical
detailing, segmentación, coaching de la fuerza de ventas, CRM); y gestión de **principals
internacionales**. Requisitos: grado en Farmacia o Nutrición **con MBA**, **8+ años** en brand
management en wellness/healthcare incluyendo farmacias, tender y e-commerce, y experiencia
gestionando equipos en nutracéuticos / suplementos / consumer healthcare.

Ángulo honesto de Paula:
- DoFreeze: Befit es una marca de nutrición better-for-you / fitness — estrategia de marca,
  presupuesto de A&P y trade, 6 lanzamientos NPD end-to-end, red de distribuidores, modern trade,
  quick-commerce (Noon, talabat, Careem) y Shopify D2C en GCC y 50+ mercados, con equipo de dos.
- Miravia (Alibaba): 42 cuentas clave en beauty, fragancias y skincare (entre ellas KIKO Milano),
  planes comerciales contra objetivos de P&L, pricing, surtido y calendario promocional, +30% GMV
  QoQ; onboarding de distribuidores oficiales (Arabian Oud, Lattafa, Swiss Arabian, Ajmal).
- Mondelez: category planning en FMCG con Nielsen — sell-in/sell-out y efectividad promocional.

Guardarraíles de honestidad (brechas reales, se declaran, no se inventan):
- **Farmacia / Nutrición + MBA**: NO los tiene. Grado en ADE (CUNEF). No se disfraza.
- **8+ años**: tiene ~5 de carrera comercial. Las fechas no se tocan.
- **Canal farmacia, tender/institucional y medical detailing con fuerza de campo**: NO los ha
  trabajado. No aparecen en el CV. Es la brecha de fondo del puesto.
- **P&L de marca completo**: gestiona A&P y trade spend y trabaja contra objetivos de P&L, pero
  nunca "owned the brand P&L" con gross margin.
- **Forecast**: en DoFreeze el forecast mensual es de la e-commerce manager; Paula aporta inputs.
  Se dice "support demand planning", nunca "own forecasting".
- Sin árabe. Visa solo "UAE Residence Visa"; nunca "no sponsorship needed".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "New Country Healthcare"
TITLE = "Brand Manager, NOW Nutritional Supplements"
DATE_FOLDER = "2026-09-23"

JOB_DESCRIPTION = """\
Brand Manager, NOW Nutritional Supplements — New Country Healthcare (NCH). Dubai, UAE. On-site, full time.
The NOW (nowfoods.com) Brand Manager is responsible for the end-to-end management of the assigned
nutraceutical brand across GCC markets, covering strategy development, commercial performance and
operational execution, accountable for revenue growth, profitability and market share while ensuring
compliance with healthcare regulations, internal governance and brand guidelines.
Brand strategy, business planning and P&L: own short and long-term strategy for NOW across the GCC; lead the
annual business planning process (revenue, profitability, market share, distribution, portfolio, brand
investment); take direct accountability for the brand P&L including gross margin, marketing investment and
trade spend; manage trade-offs between growth, margin, investment, pricing and inventory; lead business
reviews and corrective action.
Distribution and channel development: define distribution strategy across key accounts and channels; expand
high-potential accounts and rationalise underperformers; monitor channel contribution; contribute to tender
and institutional business planning; align with Sales on execution.
Marketing strategy and budget governance: manage the annual brand marketing and trade budget offline and
online against ROI targets; allocate across medical, trade and digital based on performance data; ensure
compliance with brand guidelines, regulatory requirements and financial controls; manage accruals and
variances with Finance; partner with principals to localise global strategies.
Brand activation and omnichannel execution: integrated plans across medical promotion, trade and digital;
consistent messaging and visual identity; campaign effectiveness and optimisation; digital platforms and
data analytics for engagement and omnichannel growth.
Forecasting, supply and inventory: lead brand-level forecasting and demand planning; work with Supply Chain
on availability and service levels; manage slow-moving and near-expiry stock; support launches, relaunches
and phase-outs; identify portfolio gaps and contribute to NPD and line extensions.
Field excellence and customer engagement: define medical detailing priorities, segments and core messages;
refine customer segmentation and targeting; equip field teams with compliant materials; selective field
visits and coaching; leverage CRM insights.
Team contribution and cross-functional collaboration: build a high-performance, accountable team environment;
support team development through coaching and feedback; drive cross-functional alignment and execution.
Qualifications: Bachelor's degree in Pharmacy, Nutrition or related field, with an MBA; 8+ years in brand
management and business development in the wellness/healthcare industry across channels including pharmacies,
tender and e-commerce; strong experience managing teams within nutraceuticals, dietary supplements, consumer
healthcare or another relevant consumer-led sector; proven responsibility for a significant brand, category,
portfolio or business unit with hands-on P&L, revenue, gross margin, pricing, budget and profitability
ownership; experience managing international brand principals, manufacturers or regional partners strongly
preferred; strong critical thinking, commercial negotiation, stakeholder management and influencing.
"""

ATS = [
    "brand manager", "nutraceuticals", "dietary supplements", "nutritional supplements",
    "consumer healthcare", "wellness", "brand strategy", "annual business planning", "P&L",
    "gross margin", "profitability", "revenue growth", "market share", "brand investment",
    "trade spend", "marketing budget", "budget governance", "ROI", "distribution strategy",
    "channel development", "key accounts", "modern trade", "e-commerce", "GCC", "MENA",
    "pricing", "assortment", "portfolio development", "NPD", "line extensions", "launches",
    "relaunch", "phase-out", "omnichannel", "brand activation", "trade marketing",
    "shopper marketing", "digital campaigns", "data analytics", "campaign effectiveness",
    "demand planning", "forecasting", "supply chain", "inventory", "near-expiry",
    "distributors", "principals", "regional partners", "customer segmentation", "CRM",
    "cross-functional", "stakeholder management", "commercial negotiation", "business reviews",
    "corrective action", "sell-in", "sell-out", "Nielsen", "Excel", "PowerPoint", "Power BI",
]

CONTENT = {
    "headline": "Brand Manager · Better-for-You Nutrition · GCC Distribution, Trade & E-Commerce",
    "professional_summary": (
        "Brand manager for a better-for-you nutrition brand (Befit) in Dubai — brand strategy, A&P and trade "
        "budgets, 6 NPD launches and distributor-led distribution across the GCC and 50+ markets, leading a team of "
        "two. Previously 42 beauty and skincare key accounts at Miravia (Alibaba), +30% GMV QoQ, and FMCG category "
        "planning at Mondelez."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE F&B / FMCG group | Befit (better-for-you nutrition), Eurocake, Flair | principal side | GCC + 50+ markets",
            "bullets": [
                "Own brand strategy and the annual plan for Befit, a better-for-you nutrition brand — positioning, portfolio, pricing architecture and channel roles across the GCC and 50+ markets, leading a team of two",
                "Manage the A&P and trade budget across retail, trade and digital — promotional mechanics, spend tracked against sell-out and ROI reported with Finance",
                "Set distribution priorities with the distributor network and sales across modern trade, quick-commerce (Noon, talabat, Careem) and Shopify D2C — listings, visibility and in-store execution",
                "Lead 6 NPD launches end-to-end plus relaunches and phase-outs, support demand planning with supply chain on availability and near-expiry stock, and run business reviews on sell-out",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees",
            "bullets": [
                "Managed 42 key accounts in beauty, skincare and fragrances (incl. KIKO Milano) — assortment, pricing and promotional calendars — delivering +30% GMV growth QoQ and negotiating investment with brand principals and official distributors",
                "Owned the Flash Sales channel for Beauty, Fashion & Home reporting to the CEO, building commercial plans against P&L targets and reading promotional ROI to reallocate investment",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Managed XL accounts on catalogue, promotional mechanics and GMV, negotiating joint investment, and helped build out the Retail vertical onboarding non-food brands",
            ],
        },
        {
            "company": "Mondelez International · Massimo Dutti (Inditex)",
            "role": "Trainee, Category Planning · Sales Associate",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG, chocolate category (Milka, Suchard) | Premium fashion retail floor",
            "bullets": [
                "Ran sell-in/sell-out and promotional-effectiveness analysis with Nielsen and supported NPD launches (Milka Spread, Mini Suchard); Inditex grounding in retail execution",
            ],
        },
    ],
    "skills_brand": (  # -> "Brand & Portfolio"
        "brand strategy, annual business planning, portfolio & pricing architecture, NPD end-to-end, "
        "relaunches & phase-outs, team leadership (2 reports)"
    ),
    "skills_ecommerce": (  # -> "Channels & Distribution"
        "distributor management, modern trade, key accounts, quick-commerce (Noon, talabat, Careem), "
        "Shopify D2C, trade activation"
    ),
    "skills_commercial": (  # -> "Investment & Commercial"
        "A&P and trade budgets, promotional ROI, spend tracking with Finance, pricing strategy, assortment, "
        "commercial negotiation"
    ),
    "skills_data": (  # -> "Planning & Insights"
        "sell-in/sell-out, distribution tracking, business reviews, demand planning support, near-expiry "
        "monitoring, category & shopper insights, Nielsen"
    ),
    "skills_tools": (  # -> "Tools"
        "Excel (Advanced), PowerPoint (Advanced), Power BI, Nielsen, SAP, Salesforce, Shopify, Meta & Google Ads, "
        "Claude (AI automation)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Brand & Portfolio",
    "E-Commerce & Digital": "Channels & Distribution",
    "Commercial": "Investment & Commercial",
    "Data & Analytics": "Planning & Insights",
}


def make_job() -> Job:
    return Job(
        id="new-country-healthcare-brand-manager-now-supplements-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE (On-site)",
        url="https://www.linkedin.com/jobs/search/?keywords=New%20Country%20Healthcare%20Brand%20Manager%20NOW%20Dubai",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "New Country Healthcare Brand Manager NOW Nutritional Supplements Dubai",
            "note": "Distribuidor sanitario de Dubái (fundado 1978, ~209 empleados LinkedIn, +1.500 productos) "
                    "para la marca NOW Foods en GCC. Publicado por Anas Jannoud (Head of HR). 141 solicitudes "
                    "el primer día; 40% de los solicitantes con MBA. Encaje real por marca de nutrición "
                    "better-for-you (Befit), distribuidores, modern trade, e-commerce y presupuestos de A&P. "
                    "Brechas duras: grado en Farmacia/Nutrición + MBA (no los tiene), 8+ años (tiene ~5), y "
                    "canal farmacia, tender/institucional y medical detailing con fuerza de campo (nunca los "
                    "ha trabajado). Sin árabe.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 57, "ai_tier": "Warm",
        "skills_match": [
            "Befit: marca de nutrición better-for-you que lleva hoy — lo más cercano a nutracéuticos en su CV",
            "Presupuesto de A&P y trade con mecánicas promocionales y ROI contra sell-out",
            "Red de distribuidores y modern trade en GCC + 50+ mercados desde el lado del principal",
            "E-commerce real: Shopify D2C y quick-commerce (Noon, talabat, Careem), uno de los canales del JD",
            "6 lanzamientos NPD end-to-end, relanzamientos y gestión de surtido",
            "Miravia: 42 cuentas clave en beauty y skincare contra objetivos de P&L, +30% GMV QoQ",
            "Gestión de equipo: dos reportes directos (diseño + social)",
            "Mondelez: category planning FMCG con Nielsen, sell-in/sell-out y efectividad promocional",
        ],
        "missing_skills": [
            "Grado en Farmacia o Nutrición Y MBA — requisito explícito que no cumple (40% de los solicitantes tiene MBA)",
            "8+ años en brand management de wellness/healthcare (tiene ~5 de carrera comercial)",
            "Canal farmacia y negocio de tender/institucional: nunca los ha trabajado",
            "Medical detailing y gestión de fuerza de campo médica con CRM y coaching en visitas",
            "Compliance regulatorio sanitario (nutracéuticos, claims médicos)",
            "P&L de marca completo con gross margin (gestiona A&P y trade spend, no el P&L entero)",
            "Árabe",
        ],
        "sector_fit": "medio — nutrición better-for-you y beauty/skincare son adyacentes, pero healthcare/pharma no es su sector",
        "seniority_fit": "por encima — el JD pide 8+ años y MBA para un rol de dueño de P&L de marca",
        "red_flags": [
            "141 solicitudes el primer día y requisitos duros que ella no cumple (Farmacia/Nutrición + MBA)",
            "El rol es mitad healthcare de verdad: medical detailing, tender y compliance regulatorio",
            "Presencial, no híbrido",
            "Banda salarial sin publicar — verificar contra el suelo de 20.000 AED/mes",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "El puesto tiene dos mitades. La comercial-de-marca (estrategia, plan anual, presupuesto de A&P y "
            "trade, distribución por canal, e-commerce, NPD, business reviews) encaja bien con lo que Paula hace "
            "hoy en DoFreeze sobre una marca de nutrición better-for-you, con el respaldo de 42 cuentas clave en "
            "Miravia y category planning en Mondelez. La mitad sanitaria (medical detailing con fuerza de campo, "
            "canal farmacia, tender institucional, compliance regulatorio) no la ha tocado nunca, y los "
            "requisitos formales son excluyentes en papel: Farmacia o Nutrición más MBA, y 8+ años. Con 141 "
            "solicitudes en un día y 40% de ellas con MBA, la Solicitud sencilla sola no va a pasar el filtro. "
            "La vía realista es escribir a Anas Jannoud (Head of HR, el que publica) posicionándola como brand "
            "manager de nutrición y e-commerce en GCC y preguntando si el canal e-commerce/trade pesa lo "
            "suficiente para compensar la falta de perfil farmacéutico. Si la respuesta es no, no merece más tiempo."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
