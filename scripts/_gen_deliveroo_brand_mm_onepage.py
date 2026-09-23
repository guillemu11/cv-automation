"""One-off: generate Paula's ONE-PAGE CV for **Marketing Manager, Brand (Middle East)** at
**Deliveroo**, Dubai (on-site). Brand & Experience team, reports to the Marketing team lead,
covers UAE + Kuwait.

Replaces the Aug-2026 2-page version (_gen_deliveroo_brand_mm.py) with the current facts:
team of two, real influencer/sell-out numbers with control group, 4 agencies + report audit,
seasonal cultural calendar, forecast NOT owned.

JD ask: own campaigns end-to-end as project lead (timelines, contributors, on budget);
integrated coordination with media, CRM, social, PR, central teams across ATL/OOH/digital/
social/in-app; insight-led creative briefs + agency feedback/approvals; KPI tracking and
post-campaign reports; budget & timeline ownership; market/competitor/cultural moments;
stakeholder "connective tissue". 6-8 yrs, ideally tech; UAE & Kuwait familiarity.

Honesty guardrails:
- 5 yrs vs 6-8 asked — not papered over.
- No ATL/TV/OOH buying, no PR ownership, no Kuwait-specific claim, no National Days campaign.
- No Arabic (not requested in this JD). Factual "UAE Residence Visa" only.
- "Own" only where true (campaigns, Shopify web); forecast is contributed to, not owned.

ONE-PAGE standard. Lands under output/2026-09-13/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Deliveroo"
TITLE = "Marketing Manager, Brand"
DATE_FOLDER = "2026-09-13"

JOB_DESCRIPTION = """\
Marketing Manager, Brand, Middle East — Deliveroo, Dubai (Main Office, on-site). Brand &
Experience team. Covers marketing projects and campaign management across the Middle East,
focus on UAE and Kuwait. Reports into the Marketing team lead. Own projects end-to-end:
creative sensibility, strong project management, data-driven mindset.

Planning & Execution: own end-to-end delivery of marketing projects and campaigns as project
lead — timelines, contributors, accountability — on time, on budget, highest creative standard.
Integrated Campaign Coordination: cross-functional with media, CRM, social, PR/comms and central
vertical and marketing teams; activation across ATL, OOH, digital, social, in-app.
Briefing & Stakeholder Management: insight-led creative briefs; feedback and approvals with
internal creative teams and external agencies; brand guidelines.
Performance Tracking & Reporting: campaign KPIs, post-campaign reports, recommendations.
Budget & Timeline Management: own campaign budgets and timelines, flag risks proactively.
Market & Competitor Intelligence: trends, competitor campaigns, cultural moments.
Stakeholder Management: connective tissue across local and central teams, agencies, partners.

Needs: 6-8 years in brand marketing, campaign and project management, ideally in tech; integrated
brand campaigns across channels and markets; strong creative eye and agency feedback; excellent
project management across multiple workstreams; analytically minded; strong communication with
senior stakeholders; self-starter in a matrixed environment; brand campaigns, partner activations,
tentpole seasonal moments, cross-functional commercial initiatives; familiarity with UAE and Kuwait
markets, consumer behaviour, cultural moments (Ramadan, Eid, National Days), competitive landscape.
"""

ATS = [
    "brand marketing", "integrated campaigns", "campaign management", "project management",
    "end-to-end", "creative briefs", "agency management", "feedback and approvals",
    "brand guidelines", "cross-functional", "media", "CRM", "social", "in-app", "digital",
    "campaign KPIs", "post-campaign reports", "budget management", "timelines",
    "market intelligence", "competitor campaigns", "cultural moments", "Ramadan",
    "partner activations", "seasonal", "stakeholder management", "UAE", "Middle East",
    "GCC", "quick-commerce", "delivery", "tech",
]

CONTENT = {
    "headline": (
        "Brand & Marketing Manager · Integrated Campaigns · Partner Activations"
    ),
    "professional_summary": (
        "Brand marketer with 5 years across FMCG and delivery/e-commerce tech (Glovo, Alibaba), now running integrated "
        "campaigns end-to-end in Dubai: seasonal and partner activations on talabat and Noon with proven sell-out, 4 agencies, a team of two."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG snacking group (Befit, Eurocake, Smash) | GCC + 50+ markets",
            "bullets": [
                "Project-lead integrated campaigns from insight-led brief to post-campaign report across social, influencer, paid, CRM/EDM, in-app and in-store, built around the cultural calendar (Ramadan, Back to School, Fitness Month, New Year)",
                "Run partner activations with delivery platforms: an in-app talabat giveaway and seasonal samplings; SMASH x talabat lifted daily sales +165% vs +4.7% for a control brand, Befit x Noon added 4,176 incremental units (+31%)",
                "Brief and manage 4 agencies on feedback and approvals (negotiated -30%, 103 creator activations at AED 81 per content piece) and audited their reporting, catching reach estimated from followers rather than measured",
                "Lead a team of two (designer + social media executive); own campaign budgets and timelines across Sales and Finance",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's lifestyle marketplace | 100K+ employees",
            "bullets": [
                "Created and led the Beauty Club and Hot on Social brand projects with social, content and commercial teams; owned the Flash Sales channel, reporting to the CEO",
                "Grew 42 brand accounts +30% GMV QoQ through campaign-led promotions; onboarded 30+ fragrance brands in two months",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Delivery & quick-commerce tech leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Delivered bespoke marketing activations for XL partners (KFC, Taco Bell, Sushi Shop) with marketing, ops and CX; helped build the Retail vertical",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | Chocolate category (Milka, Suchard) | €36B revenue",
            "bullets": [
                "Built promotional-effectiveness and sell-in/sell-out reports with Nielsen; supported Milka Spread and Mini Suchard launches",
            ],
        },
    ],
    "skills_brand": (  # label -> "Brand & Campaigns"
        "integrated campaigns, creative briefs, cultural moments, partner activations, influencer"
    ),
    "skills_ecommerce": (  # label -> "Channels"
        "in-app (talabat, Noon, Careem), social, CRM & EDM, Meta & Google Ads"
    ),
    "skills_commercial": (  # label -> "Project & People"
        "project management, 4 agencies, team of two, senior stakeholders, budgets"
    ),
    "skills_data": (  # label -> "Performance"
        "campaign KPIs, post-campaign reports, control groups, ROI / ROAS, competitors"
    ),
    "skills_tools": (  # label -> "Tools"
        "Meta Ads Manager, Power BI, Looker, Excel & PowerPoint, Canva, Claude"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Brand & Campaigns",
    "E-Commerce & Digital": "Channels",
    "Commercial": "Project & People",
    "Data & Analytics": "Performance",
}


def make_job() -> Job:
    return Job(
        id="deliveroo-brand-mm-me-2026-08",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE (On-site)",
        url="https://careers.deliveroo.co.uk/",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={"query": "Deliveroo Marketing Manager Brand Middle East Dubai"},
    )


def register_in_dashboard(job: Job) -> None:
    # Already tracked in data/scored_jobs.json since 2026-08 (same id) — nothing to add.
    return


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
