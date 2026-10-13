"""One-off: CV for Qureos "Senior Growth Marketer", UAE (found via Jack & Jill, 2026-10-13; job 8440517).

AI-native hiring platform ($5M seed led by Prosus Ventures), two-sided marketplace (companies + talent). Find and
scale new channels (paid, affiliates, marketplaces, communities, search, content, email, lifecycle, referral),
run an experiment engine (hypothesis, budget, time box, bar to clear), own conversion with product/engineering,
work with founders. Wants someone who took acquisition from nothing to scale (PLG/self-serve, B2B SaaS demand
gen or a two-sided marketplace), builds it hands-on (ad account, analytics, page builder, briefs, sequences),
designs tests that settle arguments and pulls their own analysis.

Paula's honest angle: channels built from zero at DoFreeze (influencer to 103 activations, Befit Crew affiliates,
communities), tests with a control group (SMASH x talabat), scale/stop calls on break-even ROAS (Google spend
moved to Meta), conversion work on Shopify (+86% orders, bundles from basket analysis, replenishment flow, email
segmentation), hands-on building with Claude and AI agents; supply side of two marketplaces (Miravia seller
acquisition, Glovo XL accounts).

Honesty guardrails: no B2B SaaS, PLG, product-analytics (Mixpanel/Amplitude) or SEO-depth claims. "From zero" only
for what she started (influencer, affiliates). Glovo is NOT the Retail vertical. Miravia accounts were beauty and
fragrance. No Arabic.

ONE-PAGE standard. Lands under output/2026-10-13/Qureos - Senior Growth Marketer/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Qureos"
TITLE = "Senior Growth Marketer"
DATE_FOLDER = "2026-10-13"

JOB_DESCRIPTION = """\
Senior Growth Marketer — Qureos, UAE. AI-native hiring platform: AI agents source, screen, message and interview;
roles distributed to 2,000+ job boards. $5M seed led by Prosus Ventures. Join a growth team that runs the core
channels; own the next set and the conversion of everything that arrives. New channels, paid and organic: paid
media, affiliates, marketplaces, communities, search, content, email, lifecycle, referral; decide what to test,
make the case with numbers, scale or stop. Both sides of the marketplace: companies hiring and talent. Experiment
engine: hypothesis, budget, time box, bar to clear; document learnings. Conversion: analytics to find where people
stall; work with product and engineering. Work with founders, product, engineering, sales. Fixed + variable +
equity. After 12 months: one new channel proven and scaling; self-serve revenue up; more qualified inbound for
sales. Who: took acquisition from nothing to scale in PLG/self-serve, B2B SaaS demand gen or a two-sided
marketplace; builds it hands-on (ad account, analytics, page builder, creative briefs, sequences); designs tests
that settle arguments within a fortnight; pulls own analysis; changes mind with data. Not: people far from the
tools, who need a staffed team, or who mistake activity for growth.
"""

ATS = [
    "growth marketing", "acquisition", "new channels", "experimentation", "A/B testing", "hypothesis",
    "affiliates", "referral", "communities", "marketplaces", "two-sided marketplace", "paid media", "email",
    "lifecycle", "conversion rate optimisation", "funnel", "landing pages", "unit economics", "CAC", "ROAS",
    "self-serve", "AI", "automation", "hands-on",
]

CV_CONTENT = {
    "headline": "Growth · Channel Experiments · Marketplaces · AI-First",
    "professional_summary": (
        "Hands-on growth marketer who builds channels from zero and tests them against a control. At DoFreeze: "
        "an influencer channel to 103 creator activations, an affiliate programme, communities and Meta "
        "Advantage+, all judged on incremental sales. Before that, worked the supply side of two marketplaces: "
        "30+ new sellers and 42 brands at Alibaba's Miravia, XL accounts at Glovo."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE F&B / FMCG group | 7 brands incl. Befit, Eurocake, SMASH | sold on talabat, Noon, Careem",
            "bullets": [
                "Built channels from zero: influencer programme to 103 creator activations over 3 waves at AED 81 per content piece; Befit Crew affiliate programme (personal codes, tiered commissions); running, padel and yoga communities",
                "Tests that settle arguments: the SMASH x talabat giveaway lifted daily sales +165% vs +4.7% for a control brand; weekly creative-hook tests; a giveaway added followers at AED 0.73 each, 8x cheaper than a promoted reel",
                "Scale or stop on unit economics: computed break-even ROAS per brand and channel net of marketplace fees and moved Google spend to Meta; built an Advantage+ retargeting campaign at 3.7x ROAS in week one",
                "Conversion, not just traffic: web orders +86% and revenue +57% (Q1 vs Jun–Aug 2026) after reviews, bundles rebuilt from basket analysis and a 30-day replenishment flow; segmented a 16K email base",
                "Build it hands-on: a Claude-based system automating campaign planning, content and KPI reporting (~40% less manual work), landing pages built with AI agents; brief a designer and a social media exec",
            ],
        },
        {
            "company": "Miravia & AliExpress (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Two-sided marketplace | Beauty & Health department",
            "bullets": [
                "Supply-side growth: spotted the Arabian-perfume trend and onboarded 30+ sellers in two months (incl. Lattafa and Arabian Oud distributors), making it a top-growing sub-category",
                "Grew 42 beauty and fragrance brands (incl. L'Oréal Paris, Maybelline, NYX, Redken, KIKO Milano) +30% GMV QoQ via pricing, promotions and on-site visibility; worked with marketing, product and ops",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Three-sided delivery marketplace | 10K+ employees",
            "bullets": [
                "Co-managed 5 XL restaurant accounts (KFC, Taco Bell, Vicio) on GMV, orders, promo spend and in-app ads; read campaign returns to pick activation days",
            ],
        },
        {
            "company": "Mondelez International · Línea Directa Aseguradora",
            "role": "Early Career",
            "dates": "2019 – 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | Spanish direct insurer",
            "bullets": [
                "Category Planning Trainee, Mondelez (Aug 2021 – Aug 2022): sell-in/sell-out and promo-effectiveness analysis, NPD support",
                "Accounting Intern, Línea Directa (Jan – Dec 2019): accounting department of an insurer, working in SAP and Looker",
            ],
        },
    ],
    "skills_brand": (
        "influencer, affiliates & referral, communities, paid social & search, email & lifecycle, marketplaces"
    ),
    "skills_ecommerce": (
        "Shopify, CRO, AI-built landing pages, bundles, reviews, email flows, Meta & Google Ads"
    ),
    "skills_commercial": (
        "control-group tests, creative testing, break-even ROAS, unit economics, scale-or-stop decisions"
    ),
    "skills_data": (
        "ROAS, CPA, CTR, CVR, AOV, incrementality, basket analysis, Looker, Power BI, Excel (Expert)"
    ),
    "skills_tools": (
        "Claude / Claude Code (agents, automation, dashboards), ChatGPT, Canva, Photoshop"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Growth Channels",
    "E-Commerce & Digital": "Conversion & Build",
    "Commercial": "Experimentation",
    "Tools": "AI & Tools",
}


def make_job() -> Job:
    return Job(
        id="qureos-senior-growth-marketer-2026-10",
        title=TITLE,
        company=COMPANY,
        location="United Arab Emirates",
        url="https://jobs.qureos.com/jobs/8440517-senior-growth-marketer",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={"via": "Jack & Jill AI (tracked job, 2026-10-13)",
             "sector": "AI hiring platform (two-sided marketplace), seed-stage"},
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not posted (Jack & Jill estimate AED 240–360K/yr + equity ≈ 20–30K/month)",
        "salary_aed_min": 20000, "salary_aed_max": 30000,
        "posted_date": None, "raw": job.raw,
        "ai_score": 66, "ai_tier": "Warm",
        "skills_match": [
            "Channels from zero: influencer (103 activations), Befit Crew affiliates, communities",
            "Control-group tests (SMASH x talabat +165% vs +4.7%); scale/stop on break-even ROAS",
            "Conversion: web orders +86%, bundles from basket analysis, replenishment flow, email segmentation",
            "Hands-on builder: Claude automation, AI-built landing pages; supply side of Miravia and Glovo",
        ],
        "missing_skills": [
            "B2B SaaS / product-led or self-serve growth",
            "Product analytics tools and funnel work with engineering",
            "SEO/content at scale",
        ],
        "sector_fit": "partial (consumer/marketplace → B2B HR-tech)",
        "seniority_fit": "good (Senior IC, hiring on proof not years)",
        "red_flags": [
            "B2B pivot: both funnels (companies and talent) are new territory",
        ],
        "ats_keywords": ATS,
        "reasoning": "Builder profile with channels from zero, control tests and conversion wins; B2B SaaS/PLG is the gap.",
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _pages(pdf: Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(pdf)).pages)


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    final_dir = settings.output_dir / DATE_FOLDER / f"{COMPANY} - {TITLE}"
    dest = final_dir / "01_CV_y_Carta"
    dest.mkdir(parents=True, exist_ok=True)

    cv_docx = cv._fill_template(CV_CONTENT, job)
    _relabel_for_role(cv_docx)
    stray = cv_docx.parent.parent

    clean_docx = dest / "Paula De Francisco - CV.docx"
    shutil.move(str(cv_docx), str(clean_docx))
    # The template writes to output/<Company> - <Role>/ at the root; only ever remove that.
    if stray.exists() and stray.parent.resolve() == settings.output_dir.resolve():
        shutil.rmtree(stray, ignore_errors=True)

    cv_pdf = _to_pdf_soffice(clean_docx)
    # The dashboard only indexes CV_Paula_<Company>_<Title>.pdf; the clean name is the one to send.
    dashboard_pdf = dest / f"CV_Paula_{cv._sanitize_filename(COMPANY)}_{cv._sanitize_filename(TITLE)}.pdf"
    shutil.copy(str(cv_pdf), str(dashboard_pdf))
    print("OK", cv_pdf, "PAGES", _pages(cv_pdf), "| dashboard copy:", dashboard_pdf.name)


if __name__ == "__main__":
    main()
