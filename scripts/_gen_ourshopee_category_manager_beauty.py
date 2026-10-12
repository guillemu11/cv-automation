"""One-off: CV + application email for OurShopee "Category Manager – Beauty", Dubai (LinkedIn hiring
post by Mansi Deo, seen 2026-10-12). Apply by DM or by email to mansi.deo@ourshopee.com.

The brief: 5+ years in Category Management; UAE experience mandatory; strong Beauty / Cosmetics /
Personal Care background; e-commerce / retail / trading preferred; strong commercial, buying and
vendor management skills. OurShopee is a Dubai value / deals e-commerce marketplace (founded 2015,
Al Qusais, 201-500 employees).

Paula's honest angle:
- Beauty + vendors: Miravia (Alibaba) 42 beauty / fragrance / fashion brands incl. KIKO Milano and
  NIVEA (three months), +30% GMV QoQ; PIC Fragrances, 30+ stores onboarded in two months incl. the
  distributors of Arabian Oud, Lattafa, Swiss Arabian, Ajmal.
- Deals model: owned the Flash Sales channel for Beauty, Fashion & Home vs P&L targets, CEO reporting.
- UAE: DoFreeze sells into Noon, talabat and Careem (listings, MSL, promo mechanics, weekly stock-risk
  review by dark store vs platform POs). Knows how UAE platforms buy, from the vendor side.
- Category grounding: Mondelez category planning (sell-in/sell-out, promo effectiveness).

Honesty guardrails:
- GAP: never held a Category Manager or Buyer title; ~2 years of beauty category work (Miravia), not
  5+. The CV never claims "5+ years category management" nor a buying / procurement role.
- UAE experience is real but ~1 year (Oct 2025 – now). No Arabic (not asked for here).
- Never "own the forecast" (contributes). DoFreeze is NOT on Deliveroo. No KIKO / NIVEA metrics: the
  +30% GMV QoQ is the whole 42-brand portfolio. Visa line in CV from profile; none in the email.

Nothing is sent: the email and the LinkedIn DM land in 03_Outreach and in the dashboard outreach store
for Paula to send from her hotmail. ONE-PAGE CV. Lands under output/2026-10-12/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops import outreach_store
from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "OurShopee"
TITLE = "Category Manager – Beauty"
DATE_FOLDER = "2026-10-12"
JOB_ID = "ourshopee-category-manager-beauty-2026-10"
CONTACT_NAME = "Mansi Deo"
CONTACT_EMAIL = "mansi.deo@ourshopee.com"

JOB_DESCRIPTION = """\
Category Manager – Beauty — OurShopee, Dubai, UAE. LinkedIn hiring post by Mansi Deo (2026-10-12).
Apply: DM the post author or share the profile with mansi.deo@ourshopee.com.
- 5+ years of experience in Category Management
- UAE experience is mandatory
- Strong background in Beauty / Cosmetics / Personal Care
- Experience in E-commerce / Retail / Trading is highly preferred
- Strong commercial, buying & vendor management skills
Context: OurShopee is a Dubai-based value / deals e-commerce marketplace (founded 2015, Al Qusais).
"""

ATS = [
    "Category Manager", "category management", "beauty", "cosmetics", "personal care", "skincare",
    "makeup", "fragrances", "e-commerce", "marketplace", "retail", "trading", "buying",
    "vendor management", "vendor onboarding", "supplier negotiation", "commercial terms", "assortment",
    "pricing", "promotions", "flash sales", "deals", "GMV", "P&L", "sell-in", "sell-out",
    "purchase orders", "stock availability", "Noon", "talabat", "Careem", "Alibaba", "KIKO Milano",
    "NIVEA", "Dubai", "UAE", "GCC",
]

CV_CONTENT = {
    "headline": (
        "Category & Key Accounts · Beauty, Fragrances & Personal Care · "
        "Vendor Management · UAE E-Commerce"
    ),
    "professional_summary": (
        "Commercial and e-commerce professional with 5 years across marketplaces and FMCG. At Alibaba's "
        "Miravia I ran 42 beauty, fragrance and fashion brands to +30% GMV QoQ and the Flash Sales channel; "
        "in Dubai I sell into Noon, talabat and Careem."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE F&B / FMCG group | Befit, Eurocake, Flair | Noon, talabat, Careem + Shopify",
            "bullets": [
                "Sell into UAE e-commerce from the vendor side: onboarded our brands onto Noon, talabat and Careem and run listings, channel assortment (MSL) and promo mechanics with e-commerce sales",
                "Run the weekly stock-risk review by dark store, checking platform POs against demand, and contribute to the monthly sell-in forecast",
                "Read every promotion against a control: SMASH x talabat lifted daily sales +165% vs +4.7% for the control brand",
                "Led 6 NPD launches (brief, packaging, pricing, go-to-market); lead a team of two and built a Claude-based system that automates KPI reporting (~40% less manual work)",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace in Spain | 100K+ employees",
            "bullets": [
                "Managed 42 beauty, fragrance and fashion brands, incl. KIKO Milano and NIVEA (three months), to +30% GMV QoQ through assortment, pricing and promotions",
                "As PIC Fragrances, onboarded 30+ new stores in two months, incl. the official distributors of Arabian Oud, Lattafa, Swiss Arabian and Ajmal",
                "Owned the Flash Sales channel for Beauty, Fashion & Home, reporting directly to the CEO against P&L targets; created the Beauty Club loyalty programme",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Helped build the Retail vertical, onboarding fashion, beauty and lifestyle brands; managed XL accounts (KFC, Taco Bell, Sushi Shop) and negotiated their commercial deals",
            ],
        },
        {
            "company": "Mondelez International · Massimo Dutti (Inditex)",
            "role": "Category Planning Trainee · Sales Associate",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG category planning | Premium fashion retail floor",
            "bullets": [
                "Sell-in/sell-out and promo-effectiveness analysis for the chocolate category at Mondelez, supporting NPD (Milka Spread, Mini Suchard); Inditex retail-floor grounding",
            ],
        },
    ],
    "skills_brand": (  # -> "Category & Trading"
        "assortment, pricing & promotions, flash sales & deals, promo uplift vs control, NPD, P&L targets"
    ),
    "skills_ecommerce": (  # -> "Beauty & Marketplace"
        "makeup, skincare, fragrances incl. oud houses; Miravia (Alibaba), Noon, talabat, Careem, Shopify"
    ),
    "skills_commercial": (  # -> "Vendor Management"
        "vendor onboarding, commercial negotiation, brand & distributor relationships, POs & availability"
    ),
    "skills_data": (  # -> "Data & Team"
        "GMV, sell-in/sell-out, conversion, ROI, stock-risk dashboards; team of two; reporting to the CEO"
    ),
    "skills_tools": (
        "SAP, Salesforce, Power BI, Looker, Nielsen, Excel (Expert), Shopify, Claude / Claude Code"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Category & Trading",
    "E-Commerce & Digital": "Beauty & Marketplace",
    "Commercial": "Vendor Management",
    "Data & Analytics": "Data & Team",
    "Tools": "Tools & AI",
}

EMAIL_SUBJECT = "Category Manager – Beauty | Paula De Francisco"
EMAIL_BODY = (
    "Hi Mansi,\n\n"
    "I saw your post about the Category Manager – Beauty role at OurShopee and would like to be "
    "considered. My CV is attached.\n\n"
    "How my experience lines up with what you're looking for:\n\n"
    "- Beauty: at Alibaba's Miravia marketplace I managed 42 beauty, fragrance and fashion brands, "
    "including KIKO Milano and, for three months, NIVEA. Through assortment, pricing and promotions the "
    "portfolio grew 30% in GMV quarter on quarter.\n"
    "- Vendor management: as the fragrance lead I onboarded 30+ new stores in two months, including the "
    "official distributors of Arabian Oud, Lattafa, Swiss Arabian and Ajmal.\n"
    "- Trading: I owned the Flash Sales channel for Beauty, Fashion & Home, reporting directly to the CEO "
    "against P&L targets.\n"
    "- UAE: I'm based in Dubai, where I sell DoFreeze's brands into Noon, talabat and Careem. That covers "
    "listings and promotions, plus a weekly check that each platform's purchase orders cover demand in "
    "its dark stores. So I know how UAE e-commerce buys, from the vendor's side of the table.\n\n"
    "I'd be glad to talk whenever suits you.\n\n"
    "Best regards,\n"
    "Paula De Francisco\n"
    "+971 50 386 3656 · paulich98@hotmail.com\n"
    "linkedin.com/in/paula-de-francisco-perez-92a13119a"
)

LINKEDIN_DM = (
    "Hi Mansi, I saw your post for the Category Manager – Beauty role. I'm in Dubai. At Alibaba's "
    "Miravia I managed 42 beauty, fragrance and fashion brands (incl. KIKO Milano) to +30% GMV QoQ and "
    "onboarded the distributors of Arabian Oud, Lattafa, Swiss Arabian and Ajmal; today I sell into "
    "Noon, talabat and Careem. I've just emailed you my CV and am happy to share it here too. Paula"
)


def make_job() -> Job:
    return Job(
        id=JOB_ID,
        title=TITLE,
        company=COMPANY,
        location="Dubai, United Arab Emirates",
        url=f"mailto:{CONTACT_EMAIL}?subject=Category%20Manager%20-%20Beauty",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={"via": "LinkedIn hiring post by Mansi Deo (2026-10-12)",
             "apply": f"DM Mansi Deo or email {CONTACT_EMAIL}",
             "level": "Manager"},
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not posted (value marketplace: check it clears her 20k AED/month floor)",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 64, "ai_tier": "Warm",
        "skills_match": [
            "Miravia (Alibaba): 42 beauty / fragrance / fashion brands incl. KIKO Milano and NIVEA, +30% GMV QoQ",
            "Vendor onboarding: 30+ stores in two months as PIC Fragrances (Arabian Oud, Lattafa, Swiss Arabian, Ajmal)",
            "Flash Sales channel for Beauty, Fashion & Home vs P&L targets, CEO reporting (fits a deals marketplace)",
            "UAE e-commerce from the vendor side: Noon, talabat, Careem listings, promos, POs and dark-store stock",
            "Category planning grounding at Mondelez (sell-in/sell-out, promo effectiveness)",
        ],
        "missing_skills": [
            "5+ years in Category Management: ~2 years of beauty category work (Miravia), no Category Manager title",
            "No buyer / procurement role: has not bought own inventory for a retailer",
            "UAE experience is ~1 year (DoFreeze, since Oct 2025), and on the supplier side",
        ],
        "sector_fit": "strong (beauty, fragrances and personal care on a marketplace)",
        "seniority_fit": "slight stretch (asks for 5+ years in category management)",
        "red_flags": [
            "Salary not posted; value e-commerce player may sit below her 20k AED/month floor",
        ],
        "ats_keywords": ATS,
        "reasoning": "Beauty marketplace trading and vendor onboarding are a close match and she has real UAE e-commerce experience, but the post wants 5+ years of category management and buying she has not held.",
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "CV Ready",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def register_outreach() -> None:
    cpath = settings.data_dir / "contacts" / f"{JOB_ID}.json"
    cpath.parent.mkdir(parents=True, exist_ok=True)
    if not cpath.exists():
        cpath.write_text(json.dumps([{
            "name": CONTACT_NAME,
            "title": "Posted the hiring post (likely Talent Acquisition / HR)",
            "company": COMPANY,
            "email": CONTACT_EMAIL,
            "linkedin_url": None,
            "role_type": "Recruiter",
            "source": "LinkedIn hiring post (2026-10-12)",
        }], indent=2, ensure_ascii=False), encoding="utf-8")

    cid = outreach_store.contact_id(CONTACT_NAME, None)
    if cid not in outreach_store.load(JOB_ID):
        outreach_store.add_version(JOB_ID, cid, "email", f"Subject: {EMAIL_SUBJECT}\n\n{EMAIL_BODY}")


def write_outreach_file(final_dir: Path, cv_name: str) -> Path:
    out_dir = final_dir / "03_Outreach"
    out_dir.mkdir(parents=True, exist_ok=True)
    md = (
        "# OurShopee — Category Manager – Beauty: email + LinkedIn DM\n\n"
        "Send both from Paula's own accounts. Nothing has been sent.\n\n"
        "## 1. Email (from paulich98@hotmail.com)\n\n"
        f"**To:** {CONTACT_EMAIL}\n"
        f"**Subject:** {EMAIL_SUBJECT}\n"
        f"**Attach:** `../01_CV_y_Carta/{cv_name}`\n\n"
        "---\n\n"
        f"{EMAIL_BODY}\n\n"
        "---\n\n"
        "## 2. LinkedIn DM to Mansi Deo\n\n"
        "> The post says \"Please DM me\". Send after the email.\n\n"
        f"{LINKEDIN_DM}\n"
    )
    p = out_dir / "Email_y_DM_OurShopee.md"
    p.write_text(md, encoding="utf-8")
    return p


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _pages(pdf: Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(pdf)).pages)


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    final_dir = settings.output_dir / DATE_FOLDER / f"{COMPANY} - Category Manager - Beauty"
    dest = final_dir / "01_CV_y_Carta"
    dest.mkdir(parents=True, exist_ok=True)

    cv_docx = cv._fill_template(CV_CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)

    moved = dest / cv_pdf.name
    shutil.move(str(cv_pdf), str(moved))
    shutil.copy(str(moved), str(dest / "Paula De Francisco - CV.pdf"))
    print("OK", moved.name, "PAGES", _pages(moved))

    stray = cv_pdf.parent.parent
    if stray.exists() and stray.resolve() != final_dir.resolve():
        shutil.rmtree(stray, ignore_errors=True)

    register_outreach()
    print("OK_OUTREACH", write_outreach_file(final_dir, "Paula De Francisco - CV.pdf"))
    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
