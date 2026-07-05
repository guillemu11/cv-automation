"""One-off: build Paula's outreach pack (.docx) for the Colgate Ecommerce
Manager application. Copy authored directly (no LLM). Mirrors the Henkel pack.
"""
from __future__ import annotations

from docx import Document
from docx.shared import Pt, RGBColor

from career_ops.config import settings

OUT = (settings.output_dir / "2026-06-11" /
       "Colgate-Palmolive - Ecommerce Manager" / "03_Outreach" /
       "Colgate-Palmolive_Outreach_Pack.docx")

RED = RGBColor(0xC4, 0x16, 0x1C)
INK = RGBColor(0x0E, 0x17, 0x26)
GREY = RGBColor(0x5A, 0x64, 0x73)

P = settings.profile["personal"]

EMAIL_SUBJECT = "Ecommerce Manager (174114) — already running this playbook in the UAE"

EMAIL_BODY = """Dear Hiring Manager,

I've applied for the Ecommerce Manager role in Dubai (Job #174114), and I wanted to reach out directly because this is exactly the work I do today.

I currently lead Brand & E-Commerce for DoFreeze in Dubai, where I own the digital shelf end-to-end: a Shopify store plus product listings, content and promotions across Noon, Amazon.ae and quick-commerce (Noon Minutes, Talabat, Careem). Before that, trained at Alibaba (Miravia/AliExpress), I managed 42 key accounts to +30% GMV growth QoQ and reported on Flash Sales P&L directly to the CEO. Digital-shelf excellence, content that converts and ROI/ROAS analytics are my daily remit — and I'm already in Dubai on a UAE residence visa, so there's no relocation or sponsorship needed.

To show rather than tell, I put together a short concept of how I'd approach Colgate's digital shelf in the UAE — "Always Shelf-Ready" — as a one-page landing and a deck. I'd love to share it and hear how you're thinking about the role.

Thank you for your time — I look forward to connecting.

Best regards,
{name}
{phone} · {email}
{linkedin}""".format(
    name=P["name"], phone=P["phone"], email=P["email"],
    linkedin=P.get("linkedin", ""),
)

LINKEDIN_CONNECTION = (
    "Hi — I've just applied for Colgate's Ecommerce Manager role in Dubai (174114). "
    "I run the digital shelf for an FMCG brand here (Noon, Amazon.ae, quick-commerce) "
    "and prepared a short concept for Colgate. Would love to connect."
)

LINKEDIN_INMAIL = """Hi [name],

I've applied for the Ecommerce Manager role at Colgate-Palmolive in Dubai (174114) and wanted to introduce myself.

I lead Brand & E-Commerce for DoFreeze here in Dubai — owning the digital shelf across Noon, Amazon.ae and quick-commerce (Noon Minutes, Talabat, Careem), content that converts, and full-funnel media. Previously, at Alibaba (Miravia/AliExpress), I managed 42 key accounts to +30% GMV QoQ with direct P&L ownership of the Flash Sales channel.

I built a short "Always Shelf-Ready" concept for how I'd grow Colgate's digital shelf in the UAE (a one-page landing + deck) — happy to share it. Would you be open to a quick chat?

Best,
{name}""".format(name=P["name"])


def _h(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RED
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)


def _label(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(0)


def _body(doc, text):
    for para in text.split("\n"):
        p = doc.add_paragraph()
        r = p.add_run(para)
        r.font.size = Pt(10.5)
        r.font.color.rgb = INK
        p.paragraph_format.space_after = Pt(0)


def main() -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"

    title = doc.add_paragraph()
    tr = title.add_run("Outreach Pack — Colgate-Palmolive · Ecommerce Manager (Dubai)")
    tr.bold = True
    tr.font.size = Pt(16)
    tr.font.color.rgb = INK

    sub = doc.add_paragraph()
    sr = sub.add_run("Drafts for Paula to review and send manually. Nothing is sent automatically.")
    sr.italic = True
    sr.font.size = Pt(9.5)
    sr.font.color.rgb = GREY

    _h(doc, "1 · Email")
    _label(doc, "SUBJECT")
    _body(doc, EMAIL_SUBJECT)
    doc.add_paragraph()
    _label(doc, "BODY")
    _body(doc, EMAIL_BODY)

    _h(doc, "2 · LinkedIn — connection request  (<=300 chars)")
    _body(doc, LINKEDIN_CONNECTION)
    note = doc.add_paragraph()
    nr = note.add_run(f"[{len(LINKEDIN_CONNECTION)} characters]")
    nr.italic = True
    nr.font.size = Pt(8.5)
    nr.font.color.rgb = GREY

    _h(doc, "3 · LinkedIn — InMail / longer message")
    _body(doc, LINKEDIN_INMAIL)

    _h(doc, "Tips")
    _body(doc,
          "- Attach the CV + Cover Letter (01_CV_y_Carta) and link the landing + deck (02_Deliverables).\n"
          "- Best target: the recruiter on the post (#LI-FS1) or a Colgate GCC e-commerce/category lead.\n"
          "- Send the connection request first; once accepted, follow up with the InMail text.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print("OK_DOCX", OUT)


if __name__ == "__main__":
    main()
