"""One-off: CV de una página para **Brand Manager** en **Cereal Partners Worldwide** (joint venture
Nestlé & General Mills: NESQUIK, FITNESS, CHEERIOS, CHOCAPIC; 40+ marcas, 130+ países, HQ Suiza).
Dubái, híbrido, jornada completa. LinkedIn "Promocionado por técnico de selección", respuestas
gestionadas fuera de LinkedIn. **499 clics en «Solicitar» en un día.** CPW tiene publicado a la vez
un Senior Brand Manager en Dubái.

Pide el JD: traducir estrategias globales de marca y segmento a planes de clúster multi-país;
calendario anual de activación (campañas, iniciativas, lanzamientos de innovación); oportunidades
de crecimiento por ocasiones de consumo, need states y cohortes; toolkits, assets y guías de
activación para que los mercados ejecuten; ejecución impecable de campañas integradas (a tiempo,
en presupuesto); trabajo con el ecosistema de agencias ("ONE agency") en contenido, creatividad y
medios; aportar a la estrategia de medios del clúster; seguimiento de brand health, eficacia de
medios y KPIs comerciales; post-campaign reviews; consolidar insights de consumidor, shopper,
cliente y competencia; punto de contacto del segmento con los mercados; trabajo con Sales, Trade
Marketing, Category y Commercial; materiales de sell-in, presentaciones a retailers y category stories.
Requisitos: marketing en FMCG con brand building; proyectos de la planificación a la ejecución;
creative briefs y adaptación de campañas globales a lo local; social media y activación digital;
coordinación multi-mercado, multi-función y multi-stakeholder. (El JD está cortado en "… más":
no se ve si pide árabe ni años de experiencia.)

Ángulo honesto de Paula:
- DoFreeze es FMCG de alimentación (Befit, Eurocake, Flair) vendida en 50+ mercados: es dueña del
  calendario de activación atado a los picos de temporada (Ramadán, back to school, Fitness Month,
  New Year New Me) — back to school y "better-for-you" son territorio natural del cereal.
- Agencias y creators con resultado de venta medido: 4 agencias, 103 activaciones de creator en
  3 campañas; Befit × Noon +31% sobre baseline (+4.176 uds incrementales); SMASH × talabat +165% de
  venta diaria contra +4,7% del grupo de control (Befit). Auditó el informe de la agencia — es
  exactamente el "post-campaign review" y la eficacia de medios del JD.
- Briefs creativos de verdad: dirige a un diseñador y a una social media executive.
- 6 lanzamientos NPD end-to-end = "innovation launches" del calendario.
- Base FMCG multinacional en Mondelez (sell-in/sell-out, eficacia promocional con Nielsen, NPD
  Milka Spread y Mini Suchard).
- Lado cliente: 42 cuentas en Miravia y XL accounts en Glovo — sabe cómo se compra un sell-in.

Guardarraíles de honestidad:
- **Cereal / desayuno**: no lo ha trabajado. Befit es nutrición better-for-you, no cereal; no se
  insinúa categoría desayuno.
- **Brand health tracking** (Kantar, BHT, funnel de marca): NO hay evidencia. Se habla de eficacia
  de campaña sobre sell-out con grupo de control, que sí es suyo; nunca de "brand health trackers".
- **Media ATL / agencia de medios / planificación de inversión de clúster**: su media es paid social
  y search (Meta, Google) más creators. No se sugiere TV ni compra con agencia de medios.
- **Matriz global "global → local"**: nunca ha recibido estrategia de una central global; en DoFreeze
  ella ES la central y adapta por mercado y canal. Se dice "adapting plans and assets by market and
  channel", no "localised global campaigns".
- **Toolkits para mercados**: no consta que haya hecho toolkits formales; no se menciona.
- **Forecast**: la dueña es la e-commerce manager; no aparece.
- **Deliveroo**: DoFreeze NO está en Deliveroo. Plataformas: talabat, Noon, Careem.
- Sin árabe (no se ve en el JD visible — confirmar en el "… más"). Visa solo "UAE Residence Visa";
  nunca "no sponsorship".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Cereal Partners Worldwide"
TITLE = "Brand Manager"
DATE_FOLDER = "2026-10-01"

JOB_DESCRIPTION = """\
Brand Manager. Cereal Partners Worldwide (Nestlé & General Mills). Dubai, UAE. Hybrid, full time.
Position Summary: We are looking for a passionate and commercially minded Brand Manager to help drive growth
across a dynamic cluster of markets. In this role, you'll turn global strategies into impactful local activation
plans, uncover growth opportunities through consumer insights, and bring brands to life with excellence in
execution. You will play a pivotal role in driving demand creation, accelerating brand performance, and
delivering winning marketing plans across multiple markets. Working closely with cross-functional teams, you
will transform insights into action and ensure our brands connect with consumers in meaningful and consistent way.
A day in the life: translate global brand and consumer segment strategies into locally relevant cluster plans
across multiple countries; build and manage the annual segment activation calendar, coordinating key campaigns,
initiatives and innovation launches; identify, evaluate and prioritize growth opportunities by understanding
consumer occasions, need states, behaviours and target cohorts; develop, localize and deploy marketing toolkits,
assets and activation guidelines; drive flawless execution of integrated marketing campaigns across all markets,
on time, on budget and in line with brand standards; partner with the ONE agency ecosystem and external partners
to create consumer-facing content, creative assets and media activations; contribute to the cluster media
strategy (channel selection, investment allocation, audience targeting); monitor and analyze brand health, media
effectiveness and commercial performance KPIs; support post-campaign reviews; consolidate consumer, shopper,
customer and competitive insights; stay ahead of consumer trends and competitive activity; provide market
feedback to global teams; act as primary point of contact for the segment across local market teams; work
cross-functionally with Sales, Trade Marketing, Category and Commercial teams; support customer-facing
initiatives through sell-in materials, retailer presentations, category stories and commercial activation plans.
What will make you successful: strong marketing experience within the FMCG industry, with a solid understanding
of brand building and consumer-focused marketing; proven experience managing marketing projects from planning
through execution; experience developing clear and inspiring creative briefs and adapting global campaigns for
local markets; hands-on experience planning and delivering social media and digital activation campaigns;
strong coordination skills across multiple markets, functions and stakeholders.
"""

ATS = [
    "brand manager", "brand management", "brand building", "FMCG", "food", "consumer-focused marketing",
    "consumer insights", "shopper insights", "competitive insights", "consumer occasions", "need states",
    "growth opportunities", "demand creation", "brand performance", "marketing plans", "annual plan",
    "activation calendar", "activation plans", "integrated marketing campaigns", "campaign execution",
    "innovation launches", "NPD", "product launches", "multi-market", "cluster", "GCC", "MENA",
    "creative briefs", "creative assets", "content", "agency management", "agencies", "media activation",
    "media strategy", "investment allocation", "audience targeting", "social media", "digital activation",
    "influencer marketing", "Meta Ads", "Google Ads", "media effectiveness", "post-campaign review",
    "campaign effectiveness", "KPIs", "ROI", "ROAS", "sell-in", "sell-out", "Nielsen", "trade marketing",
    "shopper marketing", "category", "sales", "commercial", "retailer presentations", "key accounts",
    "modern trade", "distributors", "cross-functional", "stakeholder management", "on time on budget",
    "team leadership",
]

CONTENT = {
    "headline": "Brand Manager · FMCG Food · Multi-Market Activation, Agencies & Digital · GCC",
    "professional_summary": (
        "FMCG brand manager in Dubai running brand, activation and NPD for a food group sold in 50+ markets "
        "(Befit, Eurocake, Flair), leading a team of two. Builds the seasonal activation calendar with agencies "
        "and creators and measures it on sell-out — +31% over baseline on Noon. Grounded in Mondelez category "
        "planning and 42 brand accounts at Miravia (Alibaba)."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE F&B / FMCG group | Befit (better-for-you nutrition), Eurocake, Flair | GCC + 50+ markets",
            "bullets": [
                "Own the annual activation calendar across three brands — Ramadan, back to school, Fitness Month and New Year peaks plus launches and promotions — adapting plans and assets by market and channel across the GCC and export markets",
                "Write the creative briefs and run four agencies and 103 creator activations over three campaigns: Befit × Noon delivered +31% over baseline (+4,176 incremental units); SMASH × talabat lifted daily sales +165% vs +4.7% for the control brand",
                "Lead 6 NPD launches end-to-end (brief, packaging, pricing, go-to-market) and a team of two — designer and social media executive; plan Meta & Google Ads against ROI and ROAS",
                "Run post-campaign reviews on sell-out — including auditing agency reach reporting — and work with sales and distributors on trade and shopper plans across modern trade and quick-commerce",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees",
            "bullets": [
                "Managed 42 brand accounts (incl. KIKO Milano) on assortment, pricing and promotional calendars, delivering +30% GMV QoQ, and created the Beauty Club and Hot on Social projects to lift brand visibility",
                "Owned the Flash Sales channel for Beauty, Fashion & Home reporting to the CEO — commercial plans against P&L targets, reading promotional ROI, conversion and retention to reallocate investment",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Managed XL accounts (KFC, Taco Bell, La Tagliatella) on GMV and joint marketing activations, coordinating marketing, logistics and customer support to deliver campaigns",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | chocolate category (Milka, Suchard)",
            "bullets": [
                "Ran sell-in/sell-out and promotional-effectiveness analysis with Nielsen and supported NPD launches (Milka Spread, Mini Suchard)",
            ],
        },
    ],
    "skills_brand": (  # -> "Brand & Activation"
        "brand strategy, activation calendar, integrated campaigns, creative briefs, NPD, team leadership"
    ),
    "skills_ecommerce": (  # -> "Media & Digital"
        "Meta & Google Ads, social media, creator activation, agency management, quick-commerce, EDM"
    ),
    "skills_commercial": (  # -> "Commercial & Trade"
        "trade & shopper marketing, distributors & modern trade, key accounts, A&P budgets, negotiation"
    ),
    "skills_data": (  # -> "Insights & Effectiveness"
        "consumer & shopper insights, post-campaign reviews, control groups, sell-out, ROI & ROAS"
    ),
    "skills_tools": (  # -> "Tools"
        "Excel & PowerPoint (Advanced), Power BI, Nielsen, SAP, Salesforce, Canva & Adobe, Claude (AI)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Brand & Activation",
    "E-Commerce & Digital": "Media & Digital",
    "Commercial": "Commercial & Trade",
    "Data & Analytics": "Insights & Effectiveness",
}


def make_job() -> Job:
    return Job(
        id="cpw-brand-manager-dubai-2026-10",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE (Hybrid)",
        url="https://www.linkedin.com/jobs/search/?keywords=Cereal%20Partners%20Worldwide%20Brand%20Manager%20Dubai",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Cereal Partners Worldwide Brand Manager Dubai",
            "note": "Joint venture Nestlé & General Mills (NESQUIK, FITNESS, CHEERIOS, CHOCAPIC). Rol de clúster "
                    "multi-mercado desde Dubái, híbrido. Promocionado por técnico de selección; las respuestas se "
                    "gestionan fuera de LinkedIn. 499 clics en «Solicitar» en un día. CPW publica a la vez un "
                    "Senior Brand Manager en Dubái. JD cortado: confirmar si pide árabe y años de experiencia. "
                    "Encaje FMCG alimentación + calendario de activación + agencias y creators con sell-out "
                    "medido. Gaps: brand health tracking, media ATL y matriz global → local.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 78, "ai_tier": "Hot",
        "skills_match": [
            "FMCG de alimentación hoy mismo: Befit, Eurocake y Flair, vendidas en 50+ mercados",
            "Dueña del calendario de activación por picos de temporada (Ramadán, back to school, Fitness Month, New Year)",
            "Agencias y creators con venta medida: Noon +31% sobre baseline, talabat +165% frente a +4,7% del control",
            "Post-campaign review real: auditó el informe de reach de la agencia y midió con grupo de control",
            "Briefs creativos y equipo propio: diseñador y social media executive",
            "6 lanzamientos NPD end-to-end, que encajan con los innovation launches del calendario",
            "Base FMCG multinacional en Mondelez con Nielsen, sell-in/sell-out y eficacia promocional",
            "Conoce el lado cliente: 42 cuentas en Miravia y XL accounts en Glovo",
        ],
        "missing_skills": [
            "Categoría cereal / desayuno (Befit es nutrición better-for-you, no cereal)",
            "Brand health tracking formal (Kantar, funnel de marca); su medición es sobre sell-out",
            "Media ATL y agencia de medios: su media es Meta, Google y creators",
            "Matriz global → local de multinacional: en DoFreeze ella es la central",
            "Árabe: no aparece en la parte visible del JD, sin confirmar",
        ],
        "sector_fit": "muy bueno — FMCG de alimentación, brand building y activación multi-mercado",
        "seniority_fit": "en línea — Brand Manager frente a su Brand & Marketing Manager actual; CPW tiene además un Senior BM abierto",
        "red_flags": [
            "499 solicitudes en un día: hace falta una vía directa además de solicitar",
            "JD cortado: confirmar requisito de árabe y años de experiencia",
            "Entorno de multinacional en matriz (global → clúster → mercado), que no ha vivido desde el lado marketing",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Es de los encajes más naturales para su perfil de brand: FMCG de alimentación, calendario de "
            "activación por temporada, agencias, creators y campañas digitales en varios mercados, que es lo "
            "que Paula hace hoy en DoFreeze con resultados de venta medidos con grupo de control. Le faltan "
            "tres cosas de multinacional grande: tracking formal de brand health, media ATL con agencia de "
            "medios y la experiencia de recibir una estrategia global y bajarla al clúster. El sector y el "
            "título cuadran. El riesgo real es el volumen (499 en un día) y que el JD está cortado, así que hay "
            "que confirmar si pide árabe antes de invertir en outreach."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
