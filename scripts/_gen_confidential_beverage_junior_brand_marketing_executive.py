"""One-off: CV for a confidential UAE beverage brand — "Junior Brand Marketing Executive".
Content-led role: content calendar, Instagram/TikTok/LinkedIn content, photo/video capture & editing,
Canva/Adobe graphics, launches & activations, influencers/agencies/partners, social trends.

Honest angle: Brand & Marketing Manager at DoFreeze (UAE FMCG food) — leads a designer + social media
exec and reviews all creative/social output; year-round seasonal campaigns; 103 creator activations in
3 campaigns, 4 agencies; SMASH x talabat +165% (control +4.7%); samplings & community activations;
6 launches; Adobe (Photoshop, Illustrator) + Canva; AI content tooling.

Guardrails: NO video-editing tool claimed (CapCut/Premiere not in her record — flagged). No Arabic, no
sponsorship claim. Junior band (0-2 yrs): overqualified + salary far below 20k AED floor (flagged).
CV only, ONE page. Lands under output/2026-09-13/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Beverage Brand (Confidential)"
TITLE = "Junior Brand Marketing Executive"
DATE_FOLDER = "2026-09-13"

JOB_DESCRIPTION = """\
Junior Brand Marketing Executive — beverage brand, UAE. 0-2 yrs / fresh graduates.
Plan and manage the brand content calendar; create content for Instagram, TikTok, LinkedIn; capture
and edit photos and videos; design simple graphics in Canva or Adobe Creative Suite; assist with
campaigns, product launches and brand activations; coordinate influencers, agencies and external
partners; monitor social media trends and suggest content ideas; support execution of marketing.
Needs: Bachelor's in Marketing, Communications, Multimedia or related; basic video editing (CapCut,
Premiere Pro, Final Cut); basic graphic design (Canva, Photoshop, Illustrator); passion for social
media, branding and content; communication and organisation; creative, proactive, eager to learn.
"""

ATS = [
    "Brand Marketing", "brand marketing executive", "content calendar", "content creation",
    "social media", "Instagram", "TikTok", "LinkedIn", "photo", "video", "graphic design", "Canva",
    "Adobe Creative Suite", "Photoshop", "Illustrator", "campaigns", "product launches",
    "brand activations", "influencers", "agencies", "external partners", "social media trends",
    "creative ideas", "branding", "beverage", "FMCG", "communication", "organisation", "UAE", "Dubai",
]

CV_CONTENT = {
    "headline": (
        "Brand & Social Content · Content Calendar · Influencers & Creators · Launches & Activations"
    ),
    "professional_summary": (
        "Hands-on brand marketer for UAE food brands, passionate about social media and content. "
        "I plan content calendars, brief and review creative in Canva and Adobe, run creator campaigns "
        "and activations, and track what content actually sells."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG food group (Befit, Eurocake, Smash) | Modern trade & quick-commerce | 50+ markets",
            "bullets": [
                "Plan the year-round seasonal content and campaign calendar (Ramadan, Back to School, Fitness Month, New Year) and work hands-on with a designer and a social media executive on every Instagram and TikTok post",
                "Built influencer marketing from zero: 103 creator activations across 3 campaigns and 4 agencies; SMASH x talabat lifted daily sales +165% (control brand +4.7%)",
                "Launched 6 products and ran samplings and brand activations with running, padel and yoga communities; built an affiliate landing page and multilingual content briefs with AI",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Top-5 global e-commerce | 100K+ employees",
            "bullets": [
                "Managed 42 brand partners, growing them +30% GMV QoQ through campaign calendars, promotions and visibility",
                "Onboarded 30+ fragrance stores in two months, incl. Arabian Oud, Lattafa and Ajmal distributors",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Coordinated marketing activations with XL partners (KFC, Taco Bell, Sushi Shop)",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | Chocolate category",
            "bullets": [
                "Supported the Milka Spread and Mini Suchard launches with promo and sell-out analysis",
            ],
        },
    ],
    "skills_brand": (  # -> "Content & Brand"
        "content calendar, social content ideas, art direction, launches, brand activations, sampling"
    ),
    "skills_ecommerce": (  # -> "Social"
        "Instagram, TikTok, LinkedIn, trend monitoring, UGC, Meta Ads, Shopify"
    ),
    "skills_commercial": (  # -> "Partners"
        "influencers & creators, agencies, communities, platform partners (talabat, Noon, Careem)"
    ),
    "skills_data": (  # -> "Insights"
        "content & campaign performance, control-group measurement, ROI, dashboards"
    ),
    "skills_tools": (
        "Canva, Adobe Photoshop & Illustrator, Meta Business Suite, Shopify, Generative AI (Claude, ChatGPT)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Content & Brand",
    "E-Commerce & Digital": "Social",
    "Commercial": "Partners",
    "Data & Analytics": "Insights",
}


def make_job() -> Job:
    return Job(
        id="confidential-beverage-junior-brand-marketing-executive-2026-09",
        title=TITLE,
        company=COMPANY,
        location="United Arab Emirates",
        url="",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={"note": "Company not named in the ad; beverage brand. Junior 0-2 yrs."},
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not posted (Junior band — well below 20k AED/month floor)",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 40, "ai_tier": "Possible",
        "skills_match": [
            "Seasonal content & campaign calendar; works with designer + social media exec",
            "Influencers: 103 creator activations, 4 agencies",
            "Launches, samplings, community activations",
            "Canva, Adobe Photoshop & Illustrator",
        ],
        "missing_skills": ["Video editing tools (CapCut / Premiere / Final Cut) — not in record"],
        "sector_fit": "adjacent — FMCG food, not beverage",
        "seniority_fit": "well below level — junior 0-2 yrs vs Manager",
        "red_flags": ["Junior salary band", "Overqualified: may be screened out"],
        "ats_keywords": ATS,
        "reasoning": "Content/influencer skills match fully; seniority and pay are the mismatch.",
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _pages(pdf: Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(pdf)).pages)


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    final_dir = settings.output_dir / DATE_FOLDER / f"{COMPANY} - {TITLE}"
    dest = final_dir / "01_CV"
    dest.mkdir(parents=True, exist_ok=True)

    cv_docx = cv._fill_template(CV_CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)

    moved = dest / cv_pdf.name
    shutil.move(str(cv_pdf), str(moved))
    shutil.copy(str(moved), str(dest / "Paula De Francisco - CV.pdf"))
    print("OK", moved.name, "PAGES", _pages(moved))

    stray = cv_pdf.parent.parent
    if stray.exists() and stray.resolve() != final_dir.resolve():
        shutil.rmtree(stray, ignore_errors=True)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
