"""One-off: CV + cover letter for an unnamed brand's "Event Manager" (Brand Marketing team, Dubai, GCC).
Fashion / e-commerce brand (pop-ups, fashion shows, showrooms, PR/media, partnership & sponsorship
events; KOL Marketing team; Mandarin preferred -> likely a Chinese fashion e-commerce player, but the JD
does not name it, so the package stays "Confidential").

Paula's honest angle: brand marketing in the UAE at DoFreeze — offline and seasonal activations with
talabat (samplings, in-app giveaway partnership), community events (running, padel, yoga clubs),
6 launches with timelines/budgets, agencies and suppliers (4 agencies, negotiated -30%), leads designer +
social exec, post-campaign reporting with a control group (SMASH x talabat +165% vs +4.7%). Fashion:
Miravia Beauty/Fragrances/Fashion KAM (Beauty Club, Hot on Social), Glovo retail vertical (fashion
brands), Massimo Dutti retail floor, CUNEF specialisation in Fashion Industry.

Honesty guardrails:
- NO claim of 5 years as a dedicated event manager, nor pop-up stores / fashion shows / showrooms /
  large-scale productions: not in her record. The letter says so and frames how she would run them.
- NO Mandarin, no Arabic, no sponsorship claim. GCC travel willingness not stated as fact beyond UAE.
ONE-PAGE standard. Lands under output/2026-09-13/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cover_letter as cl
from career_ops.generators import cv_generator as cv

COMPANY = "Confidential"
TITLE = "Event Manager"
DATE_FOLDER = "2026-09-13"
CONTACT = None

JOB_DESCRIPTION = """\
Event Manager — Brand Marketing team, Dubai (company not named). Plan, manage and execute brand events
across the GCC: pop-up stores, fashion shows, showrooms, PR and media events, partnership and
sponsorship events. Develop event concepts, timelines, budgets and execution plans aligned with
marketing objectives. Coordinate with PR, Creative, Social Media, KOL Marketing, Merchandising and
Operations. Manage external agencies, production vendors, venues, suppliers and event partners.
Oversee logistics, production, venue setup, staffing and on-site execution. Monitor timelines and
budgets; identify risks and build contingency plans. Measure event performance and prepare post-event
reports with insights and recommendations. Bring innovative event ideas.
Needs: Bachelor's in Marketing/Events/Communications/Business; 5+ yrs event planning or event marketing
(fashion, retail, e-commerce, lifestyle or consumer brands); large-scale events end-to-end; project
management; stakeholder management; agencies and cross-functional teams; multiple projects under
pressure; frequent GCC travel. Preferred: Mandarin; international fast-paced e-commerce; GCC markets and
regional event operations; overtime/weekends in peak periods. UAE work permit required.
"""

ATS = [
    "Event Manager", "event marketing", "event planning", "brand events", "brand experience",
    "pop-up", "fashion shows", "showrooms", "PR and media events", "partnership events", "sponsorship",
    "event concepts", "timelines", "budgets", "execution plans", "on-site execution", "logistics",
    "production vendors", "venues", "suppliers", "agencies", "contingency planning", "risk management",
    "post-event reports", "KPIs", "cross-functional", "PR", "Creative", "Social Media", "KOL", "influencers",
    "Merchandising", "Operations", "project management", "stakeholder management", "fashion", "retail",
    "e-commerce", "lifestyle", "GCC", "UAE", "Dubai",
]

CV_CONTENT = {
    "headline": (
        "Brand Marketing & Activations · Launches, Partnerships & Community Events · "
        "Agencies & Vendors · Fashion, Retail & E-commerce"
    ),
    "professional_summary": (
        "Brand marketer with 5 years across FMCG, fashion e-commerce and quick-commerce, now in Dubai. "
        "I take activations from concept, budget and timeline to on-site delivery and post-event report, "
        "working with agencies, suppliers, platform partners and creators."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG food group (Befit, Eurocake, Smash) | Modern trade & quick-commerce | 50+ markets",
            "bullets": [
                "Plan and run seasonal brand activations (Ramadan, Back to School, Fitness Month) with talabat samplings, an in-app giveaway partnership and community events with running, padel and yoga clubs",
                "Manage 4 agencies and suppliers on briefs, budgets and timelines (negotiated -30%, 103 creator activations delivered vs 75 contracted); lead a designer and a social media executive",
                "Launched 6 products from brief to go-to-market; report every campaign against a control group: SMASH x talabat lifted daily sales +165% (control +4.7%)",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Top-5 global e-commerce | 100K+ employees",
            "bullets": [
                "Created and led the Beauty Club and Hot on Social brand projects; ran Flash Sales for Beauty, Fashion & Home, reporting to the CEO",
                "Managed 42 brand partners (+30% GMV QoQ); onboarded 30+ stores in two months as PIC Fragrances",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Ran partner activations with XL brands (KFC, Taco Bell, Sushi Shop), coordinating marketing, logistics and support; helped build the fashion & Retail vertical",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | Chocolate category",
            "bullets": [
                "Promotional-effectiveness analysis; supported the Milka Spread and Mini Suchard launches",
            ],
        },
    ],
    "skills_brand": (  # -> "Events"
        "activation concepts, launches, samplings, community & partner events, brand experience, fashion retail (Inditex)"
    ),
    "skills_ecommerce": (  # -> "Delivery"
        "timelines & budgets, vendor & supplier coordination, logistics, risk & contingency planning"
    ),
    "skills_commercial": (  # -> "Partners"
        "agencies (4), platform partnerships (talabat, Noon), creators/KOLs, cross-functional teams, negotiation"
    ),
    "skills_data": (  # -> "Reporting"
        "post-campaign reports, control-group measurement, KPIs, sell-out impact, dashboards"
    ),
    "skills_tools": (
        "Excel (advanced), Google Sheets, Power BI, Looker, Canva, Adobe (Photoshop, Illustrator), Generative AI (Claude)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Events",
    "E-Commerce & Digital": "Delivery",
    "Commercial": "Partners",
    "Data & Analytics": "Reporting",
}

CL_PARAGRAPHS = {
    "opening_paragraph": (
        "I would like to apply for the Event Manager role in your Brand Marketing team in Dubai. I am "
        "Brand & Marketing Manager at DoFreeze, a UAE food group, and before that I spent two years on "
        "beauty, fragrances and fashion at Alibaba's Miravia. Events that bring a fashion and e-commerce "
        "brand into the real world across the GCC are where I want to focus next."
    ),
    "body_paragraph_1": (
        "In Dubai I turn brand moments into activations people can take part in. I plan our seasonal "
        "campaigns around Ramadan, Back to School and Fitness Month, with product samplings and an "
        "in-app giveaway run in partnership with talabat, and community events with running, padel and "
        "yoga clubs. I work across concept, budget, timeline and suppliers, and I manage four agencies: "
        "I negotiated a 30% saving and they delivered 103 creator activations against 75 contracted. I "
        "lead a designer and a social media executive, so creative, social and creators move as one "
        "plan. After each campaign I report on what it achieved. Our SMASH x talabat campaign lifted "
        "daily sales by 165% while a control brand grew 4.7%. My fashion grounding comes from Miravia, "
        "Glovo's retail vertical and my start on the shop floor at Massimo Dutti."
    ),
    "body_paragraph_2": (
        "To be transparent: my experience is in brand activations and partnership events rather than "
        "five years as a dedicated event manager, and I have not yet produced a fashion show or a "
        "pop-up store. I also do not speak Mandarin. What I would bring is the discipline I already "
        "use: a clear run-of-show and budget for each event, contingency plans agreed with vendors "
        "before the day, and a post-event report that ties footfall, content and sales back to the "
        "brief."
    ),
    "closing_paragraph": (
        "I would welcome the chance to discuss how I could help deliver your events across the GCC. "
        "Thank you for your consideration."
    ),
}


def make_job() -> Job:
    return Job(
        id="confidential-event-manager-gcc-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Dubai, United Arab Emirates",
        url="",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={"note": "Company not named; fashion/e-commerce brand with KOL team, Mandarin preferred. UAE work permit required."},
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not posted",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 45, "ai_tier": "Possible",
        "skills_match": [
            "Seasonal brand activations, samplings and talabat giveaway partnership in the UAE",
            "Community events (running, padel, yoga clubs); 6 launches with budgets/timelines",
            "4 agencies + suppliers managed; leads designer and social exec; creators/KOLs",
            "Post-campaign reporting with control group; fashion e-commerce (Miravia) + Inditex retail",
        ],
        "missing_skills": [
            "5+ years dedicated event planning / large-scale event production",
            "Pop-up stores, fashion shows, showrooms",
            "Mandarin (preferred); GCC regional event operations beyond UAE",
        ],
        "sector_fit": "adjacent — fashion e-commerce background, but FMCG activations rather than events",
        "seniority_fit": "on level",
        "red_flags": [
            "Core requirement is 5+ yrs event management — Paula has activations, not dedicated events",
            "Frequent GCC travel, overtime and weekends",
        ],
        "ats_keywords": ATS,
        "reasoning": "Activation, agency and partner skills transfer; lacks dedicated large-scale event production and Mandarin.",
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _pages(pdf: Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(pdf)).pages)


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    final_dir = settings.output_dir / DATE_FOLDER / f"{COMPANY} - {TITLE}"
    dest = final_dir / "01_CV_y_Carta"
    dest.mkdir(parents=True, exist_ok=True)

    cv_docx = cv._fill_template(CV_CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    cl_docx = cl._fill_template(CL_PARAGRAPHS, job, CONTACT)
    cl_pdf = _to_pdf_soffice(cl_docx)

    for src, short in ((cv_pdf, "Paula De Francisco - CV.pdf"),
                       (cl_pdf, "Paula De Francisco - Cover Letter.pdf")):
        moved = dest / src.name
        shutil.move(str(src), str(moved))
        shutil.copy(str(moved), str(dest / short))
        print("OK", moved.name, "PAGES", _pages(moved))

    stray = cv_pdf.parent.parent
    if stray.exists() and stray.resolve() != final_dir.resolve():
        shutil.rmtree(stray, ignore_errors=True)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
