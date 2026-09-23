"""One-off: CV de una pagina para **Growth Manager (Remote)** en **Hired** (agencia de
seleccion; cliente sin nombrar, sector "Software Development / Technology / IT Services"),
UAE remoto. Publicada hace 1 semana, 873 candidatos.

Pide el JD: estrategias de crecimiento para aumentar adquisicion, retencion y monetizacion
de usuarios; analizar metricas de plataforma y comportamiento de usuario; colaborar con
producto, marketing e ingenieria; gestionar campanas de adquisicion de pago y optimizar el
spend por canal digital; seguimiento de KPIs y reporting a stakeholders. Piden experiencia
en growth marketing o user acquisition, herramientas de analisis (SQL, Google Analytics o
similares), A/B testing y frameworks de experimentacion, project management, y conocimiento
de canales digitales (SEO, SEM, social).

Angulo honesto de Paula: DoFreeze — dueña de Meta Ads y Google Ads (audiencias, A/B testing
de creatividades, ROI/ROAS), CRO sobre Shopify, EDM y ciclo de vida, programa de influencers
de cero a 25-50 creadores por campana con afiliacion; Miravia (Alibaba) — analisis de trafico,
conversion, retencion y ROI sobre 42 cuentas en un marketplace, +30% GMV QoQ; Glovo —
quick-commerce, crecimiento de GMV por partner.

Guardarrailes de honestidad:
- **NO sabe SQL.** El JD lo lista como ejemplo ("SQL, Google Analytics, or similar"). Se
  declara lo que si tiene: Google Analytics, Power BI, Tableau, Looker y automatizacion en
  Python. Nunca se escribe SQL en el CV.
- NO tiene experiencia de growth en producto software/SaaS ni product-led growth. El angulo
  es growth de e-commerce y consumo, que es real.
- SEO/SEM: tiene Google Ads (SEM) y AEO/GEO (schema markup, Wikidata, llms.txt). SEO organico
  clasico no es su fuerte — se menciona solo lo que ha hecho.
- Visa: solo "UAE Residence Visa" (patrocinada por el empleador); nunca "no sponsorship needed".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Hired"
TITLE = "Growth Manager (Remote)"
DATE_FOLDER = "2026-09-22"

JOB_DESCRIPTION = """\
Growth Manager (Remote) — Hired (hiring for a client; Software Development / Technology / IT Services), UAE, remote.
Drive scalable growth initiatives across the platform and optimize user acquisition channels. Strategic planning,
cross-functional collaboration and data-driven decision making to enhance platform performance.
Responsibilities: develop and execute growth strategies to increase user acquisition, retention and monetization;
analyze platform metrics and user behaviour to identify opportunities for optimization; collaborate with product,
marketing and engineering teams to implement growth initiatives; manage paid acquisition campaigns and optimize
spend across digital channels; track KPIs and report on growth performance to stakeholders regularly.
Required: experience in growth marketing or user acquisition with a track record of success; proficiency in data
analysis tools such as SQL, Google Analytics or similar platforms; familiarity with A/B testing methodologies and
experimentation frameworks; strong project management skills with the ability to prioritize initiatives; knowledge
of digital marketing channels including SEO, SEM and social media.
"""

ATS = [
    "growth marketing", "user acquisition", "retention", "monetization", "growth strategy",
    "paid acquisition", "performance marketing", "media spend", "budget optimization",
    "Meta Ads", "Facebook Ads", "Instagram Ads", "Google Ads", "SEM", "social media", "EDM",
    "lifecycle marketing", "A/B testing", "experimentation", "conversion rate optimisation", "CRO",
    "funnel", "Google Analytics", "KPI tracking", "dashboards", "ROI", "ROAS", "CAC", "AOV", "GMV",
    "cohort analysis", "user behaviour", "cross-functional", "project management", "stakeholder reporting",
    "affiliate", "influencer marketing", "e-commerce", "Shopify", "marketplaces", "data-driven",
]

CONTENT = {
    "headline": "Growth & Performance Marketing · Paid Acquisition · CRO & Lifecycle",
    "professional_summary": (
        "Growth and performance marketer with 5 years driving acquisition, conversion and revenue — owning "
        "Meta and Google Ads spend, CRO on Shopify and a creator/affiliate channel built from zero, with "
        "+30% GMV growth QoQ at Alibaba's Miravia."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE consumer group | Shopify D2C + talabat, Noon, Careem, Deliveroo | 50+ markets",
            "bullets": [
                "Own paid acquisition on Meta Ads and Google Ads — audiences, creative A/B testing and budget reallocation across channels — optimising ROAS and cost per acquisition",
                "Run CRO on the Shopify store (UX, collections, checkout), lifting conversion rate and AOV, and drive retention through EDM and lifecycle campaigns",
                "Built a creator and affiliate acquisition channel from zero to 25–50 creators per campaign, turning UGC into measurable sell-out",
                "Track the full funnel in KPI dashboards (traffic, conversion, AOV, ROI, ROAS) and report growth performance weekly to leadership, shipping initiatives across design, social and agency teams",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba Group marketplace platform | 100K+ employees",
            "bullets": [
                "Grew a 42-account portfolio +30% GMV QoQ through pricing, assortment and promotional experimentation — testing mechanics and scaling what converted",
                "Owned the Flash Sales growth channel for Beauty, Fashion & Home reporting to the CEO — campaign calendar and P&L-aligned plans on a high-traffic platform",
                "Analysed platform traffic, conversion, retention and ROI to find where the funnel leaked, and rebuilt category plans around the findings",
                "Launched the Beauty Club and Hot on Social engagement programmes, driving repeat purchase and loyalty",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce platform | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Grew GMV for strategic partners (KFC, Taco Bell, Sushi Shop) through data-led promotional planning and in-app activations, and helped launch the Retail vertical",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | 90K+ employees",
            "bullets": [
                "Measured promotional effectiveness and sell-in/sell-out with Nielsen data to steer where commercial investment went next",
            ],
        },
    ],
    "skills_brand": (  # -> "Growth & Acquisition"
        "user acquisition, retention & lifecycle, creator & affiliate programmes, funnel optimisation, go-to-market"
    ),
    "skills_ecommerce": (  # -> "Paid & Channels"
        "Meta Ads (Facebook & Instagram), Google Ads / SEM, social, EDM, AEO/GEO, budget & spend allocation"
    ),
    "skills_commercial": (  # -> "Experimentation & CRO"
        "A/B testing of creatives, offers & landing pages, CRO, pricing & promotional testing, Shopify"
    ),
    "skills_data": (  # -> "Analytics & Reporting"
        "Google Analytics, Power BI, Tableau, Looker, Excel (Expert), ROI & ROAS, funnel & cohort analysis, Python"
    ),
    "skills_tools": (  # -> "Tools & Languages"
        "Shopify, Meta Business Suite, Google Ads, Salesforce, SAP, Notion, Claude (AI) | "
        "Spanish (native), English (C1)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Growth & Acquisition",
    "E-Commerce & Digital": "Paid & Channels",
    "Commercial": "Experimentation & CRO",
    "Data & Analytics": "Analytics & Reporting",
}


def make_job() -> Job:
    return Job(
        id="hired-growth-manager-remote-2026-09",
        title=TITLE,
        company=COMPANY,
        location="UAE (Remote)",
        url="https://www.linkedin.com/jobs/search/?keywords=Hired%20Growth%20Manager%20Remote%20UAE",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Hired Growth Manager Remote UAE",
            "note": "Agencia de seleccion, cliente sin nombrar. 873 candidatos. El anuncio usa la MISMA "
                    "plantilla literal que Quik Hire Staffing ('We are hiring for one of our clients', "
                    "'regardless of background, experience, or prior employment history') — publicacion "
                    "en masa. Sector cliente: software/tech, no es el vertical de Paula. Piden SQL.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 52, "ai_tier": "Warm",
        "skills_match": [
            "Meta Ads y Google Ads end-to-end: audiencias, A/B testing de creatividades, ROI/ROAS",
            "CRO sobre Shopify (UX, colecciones, checkout) + EDM y ciclo de vida para retencion",
            "Canal de afiliacion/creadores construido de cero a 25-50 creadores por campana",
            "Analitica de funnel: trafico, conversion, AOV, retencion, ROI en Miravia (42 cuentas, +30% GMV QoQ)",
            "Reporting semanal de KPIs a direccion y trabajo cross-funcional con diseno, social, agencias y supply chain",
        ],
        "missing_skills": [
            "SQL (el JD lo lista explicitamente; Paula no lo tiene)",
            "Growth en producto software/SaaS y product-led growth",
            "SEO organico clasico (si tiene SEM via Google Ads y AEO/GEO)",
            "Frameworks formales de experimentacion (Optimizely, VWO, growth squads)",
        ],
        "sector_fit": "medio — growth de e-commerce y consumo si; plataforma software/tech no",
        "seniority_fit": "encaje de titulo (Manager) pero cliente y banda sin definir",
        "red_flags": [
            "Agencia con cliente sin nombrar y plantilla de anuncio identica a la de Quik Hire Staffing",
            "873 candidatos, 'work from anywhere'",
            "Piden SQL y experiencia de plataforma tech: dos huecos reales",
            "Sin banda salarial: confirmar >= 20k AED antes de invertir tiempo",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "El JD es growth marketing con paid acquisition, A/B testing, CRO y reporting de KPIs — la mitad "
            "que Paula cubre de verdad desde DoFreeze y Miravia. Los dos huecos son SQL y que el cliente es "
            "una plataforma de software, no consumo. Sirve como opcion remota, pero el mismo perfil de CV "
            "aplica mucho mejor a Growth Marketing Manager (Performance Marketing) de Revolut, remoto UAE."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
