"""One-off: CV de una pagina para **Marketing Manager** en **Confidential** (Dubai),
oferta de LinkedIn publicada hace 1 dia, Easy Apply, +600 solicitudes.

JD: marketing de RETAIL puro - estrategias para atraer trafico a tienda y subir ventas;
campanas y promociones en tienda, publicidad local y digital; analisis de datos de venta,
insights de cliente y tendencias de mercado; briefing a equipos creativos (posters, flyers,
contenido digital); consistencia de marca y identidad visual en todos los canales retail;
social, newsletters y web; displays, senaletica y material en tienda; social ads, email y
publicidad online; analitica de web, engagement y ventas online; programas de fidelidad,
eventos, lanzamientos y aperturas de tienda; market research y seguimiento de competencia;
reporting de performance, ventas y ROI a direccion; coordinacion con store managers, equipos
de venta y agencias externas; desarrollo y gestion del presupuesto de marketing retail.
Requisitos: grado en Marketing/ADE/Retail Management (MBA preferido); minimo 4-5 anos de
marketing con foco retail; track record de campanas retail que generaron crecimiento de ventas.

Angulo de Paula: 4+ anos de marketing (5+ contando el ano de tienda en Massimo Dutti) y el
puesto pide 4-5 - por primera vez esta DENTRO de banda. Hoy en DoFreeze lleva marca, promos,
material de tienda, paid, Shopify, presupuesto A&P y reporting mensual; en Miravia llevo
calendarios promocionales de 42 marcas retail con +30% GMV QoQ; en Glovo activaciones con
partners; y en Massimo Dutti (Inditex) piso de tienda real - visual merchandising y cadencia
de tienda, que en un JD de retail marketing es un activo, no relleno.

Guardarrailes de honestidad:
- Paid hands-on: Meta Ads y Google Ads. TikTok/Instagram solo organico - NO TikTok Ads.
- Arabe: NO lo tiene. El JD no lo pide; no insinuarlo.
- Forecast: contribuye con Sales/Finance, no es duena del forecast mensual.
- Presupuesto: gestiona el A&P de marketing de sus marcas; no inflar a "P&L ownership".
- Visa: solo "UAE Residence Visa"; nunca "no sponsorship needed".
- MBA: no lo tiene (el JD lo marca como "preferred", no requisito) - no se menciona.
- Empresa confidencial (agencia de seleccion) -> imposible personalizar por marca; el CV se
  escribe en lenguaje retail generico pero concreto.

Estandar UNA PAGINA: summary corto, bullets 4/2/1/1, ~6 skills por fila, labels cortos.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Confidential Retail"
TITLE = "Marketing Manager (Retail)"
DATE_FOLDER = "2026-09-22"

JOB_DESCRIPTION = """\
Marketing Manager - Confidential (retail employer via recruitment partner), Dubai, UAE. On-site, full time.
We are looking for a Marketing Manager who will be responsible for developing and executing marketing
strategies tailored to enhance the retail brand's presence, drive foot traffic, and increase sales. This role
involves managing marketing campaigns, optimizing store promotions, and analyzing market trends to align
marketing initiatives with retail goals.
Develop and implement marketing strategies specific to the retail environment to drive store traffic and sales.
Create promotional plans and campaigns that align with retail objectives and target customer demographics.
Analyze sales data, customer insights, and market trends to inform marketing strategies and decisions.
Plan, execute, and monitor retail marketing campaigns, including in-store promotions, local advertising, and
digital marketing efforts. Collaborate with creative teams to produce marketing materials, including posters,
flyers, and digital content. Evaluate the effectiveness of campaigns and adjust strategies based on performance
metrics and feedback. Ensure consistent brand messaging and visual identity across all retail marketing
channels. Oversee the creation and distribution of brand-related content, including social media posts, email
newsletters, and website updates. Manage store displays, signage, and in-store promotions to enhance the
customer shopping experience. Develop and manage digital marketing efforts, including social media campaigns,
email marketing, and online advertising. Analyze website traffic, social media engagement, and online sales to
optimize digital marketing strategies. Utilize social media platforms to engage with customers, promote events,
and respond to inquiries. Implement strategies to build and maintain customer loyalty and engagement, including
loyalty programs, special events, and promotions. Analyze customer feedback and preferences to tailor marketing
efforts and improve the retail experience. Organize and execute events such as product launches, sales events,
and store openings to drive traffic and sales. Conduct market research to understand consumer behavior,
competitor activities, and industry trends. Monitor competitor promotions and strategies to identify
opportunities and threats in the retail market. Prepare and present reports on marketing performance, sales
trends, and ROI to senior management. Collaborate with store managers, sales teams, and other departments to
align marketing efforts with retail objectives. Coordinate with external vendors and agencies to execute
marketing campaigns and projects. Budget Management: develop and manage the marketing budget for retail
initiatives, monitor expenses and ensure cost-effective use of resources, evaluate the financial impact of
marketing campaigns and initiatives.
Educational Qualifications: Bachelor's degree in Marketing, Business Administration, Retail Management, or a
related field; Master's degree or MBA preferred.
Experience: minimum of 4-5 years of experience in marketing, with a focus on retail marketing preferred.
Proven track record in managing retail campaigns and driving sales growth.
"""

ATS = [
    "retail marketing", "marketing strategy", "foot traffic", "store traffic", "sales growth",
    "promotional plans", "store promotions", "in-store promotions", "campaign management",
    "local advertising", "digital marketing", "creative briefs", "marketing materials",
    "posters", "flyers", "digital content", "brand messaging", "visual identity",
    "brand guidelines", "social media", "email newsletters", "EDM", "website updates",
    "store displays", "signage", "POSM", "visual merchandising", "customer experience",
    "social media campaigns", "email marketing", "online advertising", "Meta Ads",
    "Google Ads", "website traffic", "engagement", "online sales", "conversion rate",
    "customer loyalty", "loyalty programs", "CRM", "retention", "events", "product launches",
    "store openings", "activations", "market research", "consumer behaviour", "competitor analysis",
    "sales data analysis", "customer insights", "performance reporting", "ROI", "ROAS",
    "store managers", "sales teams", "agencies", "vendors", "budget management",
    "marketing budget", "A&P", "cost control", "FMCG", "retail", "UAE", "Dubai",
]

CONTENT = {
    "headline": "Retail Marketing Manager · Campaigns & In-Store Promotions · Traffic & Sales Growth · Dubai",
    "professional_summary": (
        "Marketing manager with 5+ years across retail, FMCG and e-commerce in Dubai and Europe, building "
        "campaigns that move store traffic and sales - in-store promotions and displays, local and digital "
        "advertising, loyalty and events - owning the marketing budget and reporting ROI to management."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 - Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG group (Befit, Eurocake, Flair) | 1,000+ retail doors, HORECA, Shopify D2C, talabat & Careem",
            "bullets": [
                "Build and run the promotional calendar across 1,000+ retail doors - in-store promotions, displays, signage and POSM, local activations and sampling - briefing store and distributor teams to convert footfall into sell-out",
                "Plan and optimise digital and local advertising day to day: Meta and Google Ads, social content, EDM newsletters and website updates, keeping brand messaging and visual identity consistent across every channel",
                "Own the marketing budget for three brands - allocate spend by channel, track cost against return and present monthly reports on campaign performance, sales trends and ROI to senior management",
                "Brief a designer, a social executive and 4 external agencies on posters, flyers, in-store material and digital content; read sell-out data, customer feedback and competitor promotions to adjust the plan",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager - Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 - Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's lifestyle marketplace | 42 retail & fashion brands | 100K+ employees",
            "bullets": [
                "Planned and executed promotional campaigns and seasonal events with 42 retail and fashion brands - offers, pricing and assortment - delivering +30% GMV growth QoQ",
                "Created and led the Beauty Club and Hot on Social loyalty and community programmes, lifting repeat purchase and engagement, and analysed traffic, conversion and ROI to steer investment",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager - XL Accounts",
            "dates": "Sep 2022 - Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | retail & F&B partners | 10K+ employees",
            "bullets": [
                "Ran bespoke promotions and marketing activations with retail and F&B partners (KFC, Taco Bell, Sushi Shop), using order and conversion data to grow volume profitably",
            ],
        },
        {
            "company": "Massimo Dutti (Inditex)",
            "role": "Sales Associate - Las Rozas Village",
            "dates": "Jun 2018 - Jun 2019",
            "location": "Madrid, Spain",
            "context": "Inditex premium fashion retail | shop floor, visual merchandising & customer experience",
            "bullets": [
                "Front-line retail experience at Inditex - visual merchandising standards, store cadence and customer experience - the shop-floor view behind every campaign and display decision",
            ],
        },
    ],
    "skills_brand": (  # label -> "Retail Marketing"
        "retail marketing strategy, promotional plans & campaigns, in-store promotions, displays, signage & POSM, "
        "visual merchandising, local advertising, events & store activations"
    ),
    "skills_ecommerce": (  # label -> "Digital & Content"
        "Meta Ads, Google Ads, social media campaigns, email newsletters & EDM, website & Shopify updates, "
        "brand messaging & visual identity, creative briefs (posters, flyers, digital content)"
    ),
    "skills_commercial": (  # label -> "Customer & Commercial"
        "customer loyalty programmes, retention & engagement, market research & competitor tracking, "
        "store and sales team collaboration, agency & vendor management, negotiation"
    ),
    "skills_data": (  # label -> "Budget & Reporting"
        "marketing budget ownership (A&P), cost control, sales data & customer insights, footfall-to-sell-out analysis, "
        "website traffic & conversion, ROI/ROAS reporting to senior management"
    ),
    "skills_tools": (  # label -> "Tools"
        "Meta Ads Manager, Google Ads, Google Analytics, Shopify, Power BI, Nielsen, Canva, Adobe (Photoshop), Claude (AI)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Retail Marketing",
    "E-Commerce & Digital": "Digital & Content",
    "Commercial": "Customer & Commercial",
    "Data & Analytics": "Budget & Reporting",
}


def make_job() -> Job:
    return Job(
        id="confidential-retail-marketing-manager-dubai-2026-09-22",
        title="Marketing Manager",
        company="Confidential",
        location="Dubai, United Arab Emirates",
        url="https://www.linkedin.com/jobs/",
        source="linkedin",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Confidential Marketing Manager Dubai retail",
            "note": "Oferta publicada hace 1 dia via LinkedIn Easy Apply, empresa confidencial (partner de "
                    "seleccion, 3.455 empleados en LinkedIn). Retail marketing puro: trafico a tienda, promos "
                    "en tienda, displays y senaletica, publicidad local y digital, fidelidad, eventos, market "
                    "research, presupuesto y ROI. Pide 4-5 anos con foco retail - Paula esta EN banda. "
                    "+600 solicitudes en 1 dia: competencia altisima, la Easy Apply es una loteria. "
                    "Empleador desconocido -> sin personalizacion posible por marca.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 72, "ai_tier": "Warm",
        "skills_match": [
            "4-5 anos pedidos: Paula esta en banda (4+ de marketing, 5+ contando Inditex)",
            "Promos en tienda, displays y POSM en 1.000+ puntos de venta en EAU hoy mismo",
            "Publicidad digital y local: Meta Ads, Google Ads, social, EDM y web",
            "Presupuesto de marketing propio y reporting mensual de ROI a direccion",
            "Fidelidad y comunidad: Beauty Club y Hot on Social en Miravia",
            "Piso de tienda real en Massimo Dutti (Inditex): visual merchandising y cadencia de tienda",
        ],
        "missing_skills": [
            "MBA (preferido en el JD, no requisito) - no lo tiene",
            "No ha llevado marketing de una cadena de tiendas propia con store managers a su cargo",
            "Arabe - no lo tiene (el JD no lo pide)",
        ],
        "sector_fit": "alto - retail/FMCG, su categoria; aunque el empleador es confidencial",
        "seniority_fit": "en banda - piden 4-5 anos con foco retail y ella esta justo ahi",
        "red_flags": [
            "+600 solicitudes en 24 horas (Easy Apply) - probabilidad estadistica baja",
            "Empleador confidencial: no se puede investigar marca, cultura ni salario",
            "JD larguisimo y generico: puede ser un rol de todo-en-uno mal pagado",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Es el JD de retail marketing mas alineado que le ha salido en banda de experiencia: trafico a "
            "tienda, promos y displays, publicidad local y digital, fidelidad, eventos, research de "
            "competencia, presupuesto y ROI - todo lo que hace hoy en DoFreeze mas el calendario promocional "
            "de 42 marcas en Miravia. El diferencial honesto frente a 600 candidatos es la combinacion de "
            "marketing de marca + ejecucion en punto de venta + un ano real de piso de tienda en Inditex. "
            "Lastre: empleador confidencial y volumen brutal de solicitudes; conviene aplicar pero sin "
            "invertir mas que el CV."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    # Etiquetas de carpeta/archivo legibles para el paquete
    job.company = COMPANY
    job.title = TITLE

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
