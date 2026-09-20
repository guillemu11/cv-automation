"""One-off: generate Paula's ONE-PAGE CV for **Digital Marketing Executive** at **HR Experts**
(recruitment agency, Dubai; hybrid, PART-TIME; mid-to-senior; must already be UAE-based).

JD ask: plan/execute/optimise digital campaigns; coordinate content calendars across web, social,
email; support lead generation and monitor incoming leads; manage website content, landing pages and
digital assets; coordinate designers, content creators, sales and external partners; monitor traffic,
engagement, conversions; prepare performance reports; assist paid advertising (targeting, setup,
monitoring); support SEO (keyword research, on-page, content optimisation); email marketing (lists,
scheduling, tracking); competitor/trend research; test landing pages and workflows. Tools: Google
Analytics, Google Ads, Meta Ads, SEO tools, CRM, marketing automation.

Paula's honest angle: she already runs exactly this stack hands-on at DoFreeze (Meta + Google Ads,
Shopify site/landing pages and product content, EDM, content calendar, KPI reporting, agency and
designer coordination) and has the analytics/commercial depth from Miravia and Glovo.

Honesty guardrails:
- SEO framed as on-page/product-content optimisation on Shopify + keyword-led content — no agency SEO claims.
- Lead generation framed as campaign-to-conversion tracking and CRM-adjacent work, not B2B demand gen ownership.
- Forecast wording avoided (she contributes, doesn't own). No Arabic implied. Visa = UAE Residence Visa only.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "HR Experts"
TITLE = "Digital Marketing Executive"
DATE_FOLDER = "2026-09-20"

JOB_DESCRIPTION = """\
Digital Marketing Executive — HR Experts, Dubai, UAE. Hybrid, part-time, mid-level to senior.
Applicants must currently be based in the United Arab Emirates.
Support the planning, execution and optimisation of digital marketing campaigns across online channels.
Coordinate campaigns, promotional activities, content schedules and marketing initiatives. Create and
coordinate digital content for website, social media and email. Support lead-generation campaigns and
monitor incoming digital leads. Manage and update website content, landing pages, promotional materials
and digital assets. Coordinate with designers, content creators, sales teams and external marketing
partners. Monitor campaign performance, traffic, engagement and conversions. Prepare digital marketing
reports and performance summaries. Assist with paid advertising campaigns, audience targeting, setup and
monitoring. Support SEO activities including keyword research, on-page improvements and content
optimisation. Assist with email marketing campaigns, audience lists, scheduling and tracking. Research
digital marketing trends and competitor activity. Maintain records of campaigns, content, assets, leads
and performance data. Monitor deadlines, approvals, budgets and deliverables. Test marketing content,
landing pages and digital workflows. Analyse results and identify improvement opportunities. Ensure
content follows brand guidelines. Requirements: 3+ years in digital marketing / performance marketing;
Bachelor's in Marketing, Communications, Business or Digital Media; understanding of digital channels,
campaign management, content marketing, SEO and lead generation; social media, email marketing, websites,
landing pages and digital advertising platforms; analytics, reporting, marketing automation and CMS tools;
Google Analytics, Google Ads, Meta Ads, SEO tools, CRM or marketing automation an advantage.
"""

ATS = [
    "digital marketing", "digital campaigns", "campaign management", "campaign optimisation",
    "performance marketing", "paid advertising", "Meta Ads", "Facebook Ads", "Instagram Ads",
    "Google Ads", "audience targeting", "A/B testing", "SEO", "keyword research", "on-page",
    "content optimisation", "content marketing", "content calendar", "social media", "email marketing",
    "EDM", "audience lists", "lead generation", "leads", "landing pages", "website content", "CMS",
    "Shopify", "digital assets", "brand guidelines", "traffic", "engagement", "conversions",
    "conversion rate optimisation", "CRO", "Google Analytics", "reporting", "performance reports",
    "KPIs", "ROI", "ROAS", "marketing automation", "CRM", "competitor research", "budgets",
    "agency coordination", "designers", "sales teams", "UAE", "Dubai",
]

CONTENT = {
    "headline": "Digital Marketing · Paid Media & SEO · Content, Email & Landing Pages",
    "professional_summary": (
        "Digital marketer with 5 years across FMCG, beauty and e-commerce, now running the full digital stack in Dubai: "
        "Meta and Google Ads, Shopify site and landing pages, SEO-led content, email and KPI reporting — briefing designers, "
        "content creators and agencies end-to-end."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG group (Befit, Eurocake, Flair) | Shopify D2C, talabat, Noon, Careem | 50+ markets",
            "bullets": [
                "Plan, launch and optimise paid campaigns on Meta Ads and Google Ads — audience targeting, creative A/B testing and budget pacing — reporting ROI and ROAS weekly",
                "Own the website end-to-end on Shopify: page and landing-page content, product copy, on-page SEO and keyword-led content, testing changes to lift conversion rate and AOV",
                "Run the content calendar across web, social and email (EDM): brief and review a designer and a social media executive plus 4 external agencies, keeping everything on brand guidelines and on deadline",
                "Build the KPI dashboards behind it — traffic, engagement, conversions, campaign results — and turn them into monthly performance summaries with clear next actions",
                "Built an AI-powered automation layer (Claude / generative AI) for campaign planning, content and reporting, cutting manual workload ~40%",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees",
            "bullets": [
                "Ran always-on promotions and campaign calendars for 42 beauty, fragrance and fashion brands, analysing traffic, conversion and retention to reallocate spend (+30% GMV QoQ)",
                "Owned the Flash Sales channel reporting to the CEO and created the Beauty Club and Hot on Social projects, growing visibility and repeat purchase",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce app | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Planned bespoke marketing activations and promotions for XL partners (KFC, Taco Bell, Sushi Shop), coordinating marketing, ops and support to grow orders",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | Chocolate category (Milka, Suchard) | €36B revenue",
            "bullets": [
                "Built sell-in/sell-out and promotional-effectiveness reports with Nielsen, measuring which activations actually paid back",
            ],
        },
    ],
    "skills_brand": (  # label -> "Digital Campaigns"
        "campaign planning & execution, content calendar, promotions, brand guidelines, agency & designer briefing"
    ),
    "skills_ecommerce": (  # label -> "Paid Media & SEO"
        "Meta Ads, Google Ads, audience targeting, A/B testing, on-page SEO, keyword research, content optimisation"
    ),
    "skills_commercial": (  # label -> "Web, Content & Email"
        "Shopify CMS, landing pages, website content, digital assets, EDM & audience lists, social (IG, TikTok)"
    ),
    "skills_data": (  # label -> "Analytics & Reporting"
        "traffic, engagement, conversions, CRO, ROI & ROAS, performance reports, competitor research"
    ),
    "skills_tools": (  # label -> "Tools"
        "Google Analytics, Google Ads, Meta Business Suite, Shopify, Looker, Power BI, Canva, Adobe, Claude (AI)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Digital Campaigns",
    "E-Commerce & Digital": "Paid Media & SEO",
    "Commercial": "Web, Content & Email",
    "Data & Analytics": "Analytics & Reporting",
}


def make_job() -> Job:
    return Job(
        id="hr-experts-digital-marketing-executive-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE",
        url="https://www.linkedin.com/jobs/search/?keywords=HR%20Experts%20Digital%20Marketing%20Executive%20Dubai",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "HR Experts Digital Marketing Executive Dubai",
            "job_poster": "Ahmad AlRefaai (job advertiser on LinkedIn)",
            "note": "Recruitment agency posting — end client not named. PART-TIME, hybrid. Executive level: "
                    "likely below the AED 20k/month floor. 63 applicants in 3 hours.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 48, "ai_tier": "Warm",
        "skills_match": [
            "Runs Meta Ads + Google Ads hands-on today (targeting, A/B testing, ROI/ROAS)",
            "Owns a Shopify site end-to-end: landing pages, content, on-page SEO, CRO",
            "Content calendar across web, social and email; briefs designers, creators and 4 agencies",
            "Builds the KPI dashboards and monthly performance reports the JD asks for",
            "Already UAE-based with residence visa (JD hard requirement)",
        ],
        "missing_skills": [
            "No dedicated agency-side SEO tooling experience (Ahrefs/SEMrush) beyond on-page",
            "No B2B lead-generation / demand-gen ownership",
            "No formal CRM / marketing-automation platform ownership (HubSpot, Klaviyo-level)",
        ],
        "sector_fit": "generic — agency posting, end client unnamed",
        "seniority_fit": "step down — Executive vs her Manager title",
        "red_flags": [
            "Part-time role — likely well below the AED 20k/month floor",
            "Recruitment agency posting, end client undisclosed",
            "63 applicants within 3 hours",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Function fit is excellent — this is literally the digital stack she runs daily — and she meets the "
            "UAE-based hard requirement. The problem is level and terms: part-time Executive at an unnamed client, "
            "almost certainly under her salary floor. Worth a low-effort Easy Apply, not a priority."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
