"""One-off: CV de una página para **Key Account Manager - Middle East** en **Puig** (Dubái, equipo Commercial &
Business Development, indefinido). Oferta pegada por Guille el 2026-10-04; sin sueldo ni fecha.

Qué es el puesto: KAM lado marca que reporta al Middle East Commercial Director y lleva una cartera de retailers
de belleza: **Al Tayer, Ulta y Gold Apple**, entre otros. Tres bloques: (1) cuentas: plan anual alineado con la
visión a tres años, relación con el retailer, negociación de condiciones, inversiones y visibilidad junto al
director, espacios fijos y temporales en tienda y animaciones; (2) operación y lanzamientos: punto de contacto
operativo, NLTs, plantillas del retailer y altas a tiempo, coordinación con Supply Chain, Demand Planning, Marketing
y Brand, avisar de retrasos, seguir el lanzamiento hasta que está listado, pedido, entregado y activado; (3)
forecast y stock: forecast con cada retailer, rolling mensual de lanzamientos y trimestral de core, fill rate,
cobertura y reposición. Requisitos: experiencia comercial con exposición a KAM, **idealmente en belleza o lujo**;
inglés; negociación; analítica; detalle con datos; conciencia cultural en la región. **No pide árabe.** Puig es
empresa familiar de Barcelona: el español nativo de Paula suma.

Ángulo honesto de Paula — ha estado **en los dos lados de la mesa**:
- **Lado retailer, en belleza y fragancia**: en Miravia (Alibaba) fue KAM de 42 marcas de beauty, fragancia y moda,
  KIKO Milano (maquillaje y skincare) incluida; +30% GMV QoQ en la cartera con surtido, precio, promociones y
  visibilidad; canal Flash Sales reportando al CEO contra objetivos de P&L. Como PIC Fragrances dio de alta 30+
  casas en dos meses (distribuidores oficiales de Arabian Oud, Lattafa, Swiss Arabian y Ajmal). Sabe qué pide un
  retailer como Gold Apple o Ulta a una marca.
- **Lado marca, vendiendo a retailers que compran por PO**: en DoFreeze, talabat, Noon y Careem compran con
  previsión + PO a su dark store principal. Paula hace las altas, fichas y mecánicas promocionales, las
  activaciones conjuntas (sorteo con cashback de talabat, samplings estacionales) y la **revisión semanal de riesgo
  de stock por dark store** (cobertura de PO frente a demanda). Aporta al forecast mensual de sell-in. Revisó el MSL
  de plataformas junto a e-commerce sales.
- **Lanzamientos de punta a punta**: 6 lanzamientos (brief, packaging, precio, salida al mercado) y la marca SMASH
  en talabat: venta diaria +165% frente a +4,7% de la marca de control.
- **Negociación**: condiciones comerciales en Glovo con cadenas XL; fees de agencias -30% en DoFreeze.
- **Sell-in/sell-out**: Mondelez (análisis de categoría y eficacia promocional). Retail premium en Massimo Dutti.

Guardarraíles de honestidad:
- **KAM lado marca en retail físico de belleza** (Al Tayer, Ulta, Gold Apple, grandes almacenes): sin experiencia.
  No se nombran esos retailers ni se insinúa relación previa.
- **Espacios en tienda, animaciones, VMS, ranking, NLTs y plantillas de retailer**: sin evidencia. Lo más cercano son
  las altas y fichas en plataformas; se dice "listings", no "NLTs".
- **Forecast**: lo cierra la e-commerce manager; Paula "contribute". Nada de rolling forecasts propios ni fill rate.
- **Marcas de Puig en Miravia**: sin confirmar; no se afirma.
- **Supply Chain / Demand Planning**: no se afirma que trabaje con esos equipos en DoFreeze; se habla de la revisión
  de stock y de los avisos a ventas.
- **Deliveroo**: DoFreeze NO está. Plataformas: talabat, Noon, Careem. Visa: la de profile.yaml.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Puig"
TITLE = "Key Account Manager - Middle East"
DATE_FOLDER = "2026-10-04"

JOB_DESCRIPTION = """\
Key Account Manager - Middle East. Puig. Location: Dubai, DU, AE. Team: Commercial & Business Development.
Job type: Permanent. Puig is a major player in the worldwide fashion and beauty industry, with a portfolio of
luxury brands across fashion, fragrance, makeup, skincare and wellness; family-owned, founded more than 100 years ago.
The Opportunity: reporting to the Middle East Commercial Director, you will manage a portfolio of key customers,
including Al Tayer, Ulta and Gold Apple, among others. Develop account plans, build strong customer relationships and
identify opportunities to drive sustainable business growth. Strategic account management with day-to-day
coordination across Commercial, Marketing and Supply Chain teams; ensure plans and product launches are delivered
smoothly while contributing to sell-out, VMS and ranking targets.
Key Account Management: develop the annual key account strategy and account plan, aligned with the three-year vision;
build trusted customer relationships and partner with retailers to maximize sell-in, sell-out and investment
opportunities; lead key account negotiations alongside the Commercial Director (commercial terms, investments,
visibility, retailer support); identify growth opportunities and monitor delivery of the annual account plan;
negotiate permanent and temporary in-store spaces, including animations and visibility opportunities; share account
and market trends with the Commercial and Marketing teams.
Account Operations & Launch Execution: manage day-to-day operations of assigned accounts and serve as main point of
contact for retailer operational matters; ensure NLTs, retailer templates, listings and documentation are submitted
accurately and on time; coordinate with Supply Chain, Demand Planning, Marketing and Brand teams so new products arrive
on time and launches are fully prepared; proactively communicate delays, supply risks or changes; oversee launches
through to completion (listed, ordered, delivered, activated); resolve ordering, listing, replenishment and logistics
issues promptly.
Forecasting & Stock Management: lead the forecasting process with assigned retailers with Supply Chain and each
retailer's supply chain function; monthly rolling forecast for new products and launches; quarterly rolling forecast
for core and best-selling SKUs; improve forecast accuracy and fill rates and secure stock availability; monitor stock
cover and replenishment; identify supply and inventory risks early and drive corrective action.
We'd love to meet you if you have: relevant sales or commercial experience, with exposure to key account management,
ideally within the beauty or luxury industry; fluency in English; collaborative approach and ability to work across
functions as a trusted business partner; excellent communication and interpersonal skills; strong influencing and
negotiation skills; strong analytical and commercial acumen; high integrity and accountability; strong attention to
detail and data-management skills; hands-on mindset; highly organised and structured; strong cultural awareness and
ability to work across diverse markets and stakeholders in the region.
"""

ATS = [
    "Key Account Manager", "key account management", "account plan", "annual account plan", "account strategy",
    "beauty", "luxury", "fragrance", "makeup", "skincare", "retailers", "retail partners", "customer relationships",
    "sell-in", "sell-out", "negotiation", "commercial terms", "investments", "visibility", "retailer support",
    "growth opportunities", "market trends", "listings", "product launches", "launch execution", "new products",
    "go-to-market", "point of contact", "ordering", "replenishment", "forecasting", "forecast", "stock availability",
    "stock cover", "inventory risk", "PO", "assortment", "MSL", "pricing", "promotions", "P&L", "GMV",
    "cross-functional", "Commercial", "Marketing", "Supply Chain", "business partner", "analytical",
    "commercial acumen", "data management", "Excel", "SAP", "Nielsen", "attention to detail", "cultural awareness",
    "Middle East", "GCC", "UAE", "Dubai", "Spanish",
]

CONTENT = {
    "headline": "Key Account Manager · Beauty & Fragrance · Retailer Partnerships · Launches & Stock",
    "professional_summary": (
        "Key account professional in Dubai who has sat on both sides of the table: retailer-side KAM for 42 beauty, "
        "fragrance and fashion brands at Alibaba's Miravia (incl. KIKO Milano, +30% GMV QoQ); now brand-side, "
        "running launches, listings, joint activations and stock-risk reviews with talabat, Noon and Careem."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG group (Befit, Eurocake, Flair) | talabat, Noon, Careem + modern trade | 50+ markets",
            "bullets": [
                "Lead product launches end-to-end — 6 launches (brief, packaging, pricing, go-to-market) and the new SMASH brand on talabat, coordinating listings, content and timing: daily sales +165% vs +4.7% for a control brand",
                "Partner with talabat, Noon and Careem — which buy on PO into their dark stores — on listings, promotional mechanics, MSL by channel and joint activations (talabat cashback giveaway, seasonal sampling); Befit × Noon \"New Year, New Me\": +4,176 incremental units, +31% over baseline",
                "Run the weekly stock-risk review by dark store, checking PO coverage against current and upcoming demand and flagging gaps early to protect availability; contribute to the monthly sell-in forecast",
                "Lead a team of two (designer and social media executive) and four agencies, with fees renegotiated -30%",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees | retailer side",
            "bullets": [
                "Managed 42 beauty, fragrance and fashion brands — incl. KIKO Milano makeup & skincare — planning assortment, pricing, promotions and visibility for +30% GMV QoQ; owned the Flash Sales channel, reporting to the CEO against P&L targets",
                "As PIC Fragrances, onboarded 30+ fragrance houses in two months, incl. the official distributors of Arabian Oud, Lattafa, Swiss Arabian and Ajmal; created the Beauty Club and Hot on Social programmes to lift brand visibility",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Managed XL chains (KFC, Taco Bell, La Tagliatella, Sushi Shop) on data-led growth plans, negotiating commercial terms and coordinating marketing, logistics and customer support to deliver campaigns",
            ],
        },
        {
            "company": "Mondelez International · Massimo Dutti (Inditex)",
            "role": "Trainee, Category Planning · Sales Associate",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG (chocolate category) | Premium fashion retail",
            "bullets": [
                "At Mondelez, sell-in/sell-out and promotional-effectiveness analysis for chocolate, supporting NPD launches (Milka Spread, Mini Suchard); at Massimo Dutti, premium retail floor and visual merchandising standards",
            ],
        },
    ],
    "skills_brand": (  # -> "Key Accounts"
        "account plans, negotiation, commercial terms, joint activations, visibility, P&L targets"
    ),
    "skills_ecommerce": (  # -> "Launches & Ops"
        "NPD launches, listings & product content, assortment & MSL, cross-functional coordination"
    ),
    "skills_commercial": (  # -> "Forecast & Stock"
        "stock-risk reviews, PO coverage, availability, replenishment, sell-in forecast input"
    ),
    "skills_data": (  # -> "Category & Data"
        "beauty & fragrance, sell-in/sell-out, Nielsen, pricing & promo analysis, ROI, conversion"
    ),
    "skills_tools": (  # -> "Tools"
        "Excel (Advanced), PowerPoint, Power BI, SAP, Salesforce, Tableau, Looker, Claude (AI)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Key Accounts",
    "E-Commerce & Digital": "Launches & Ops",
    "Commercial": "Forecast & Stock",
    "Data & Analytics": "Category & Data",
}


def make_job() -> Job:
    return Job(
        id="puig-key-account-manager-middle-east-dubai-2026-10",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE",
        url="https://www.linkedin.com/jobs/search/?keywords=Puig%20Key%20Account%20Manager%20Middle%20East",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Puig Key Account Manager Middle East Dubai",
            "note": "KAM lado marca para retailers de belleza (Al Tayer, Ulta, Gold Apple), reporta al Commercial "
                    "Director ME. Encaje: KAM de beauty y fragancia en Miravia (lado retailer, KIKO incluida) + "
                    "DoFreeze (lanzamientos, altas y revisión de stock por dark store con talabat, Noon y Careem). "
                    "Gaps: retail físico de belleza lado marca, espacios en tienda y animaciones, forecast propio.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 80, "ai_tier": "Hot",
        "skills_match": [
            "KAM de 42 marcas de beauty, fragancia y moda en Miravia (KIKO Milano incluida), +30% GMV QoQ",
            "Ha estado en el lado retailer: sabe qué piden Gold Apple o Ulta a una marca",
            "PIC Fragrances: 30+ casas dadas de alta en dos meses (Arabian Oud, Lattafa, Swiss Arabian, Ajmal)",
            "Canal Flash Sales reportando al CEO contra objetivos de P&L",
            "Lanzamientos de punta a punta: 6 productos y la marca SMASH (+165% frente a control)",
            "Retailers que compran por PO (talabat, Noon, Careem): altas, promociones, activaciones conjuntas",
            "Revisión semanal de riesgo de stock por dark store; aporta al forecast mensual de sell-in",
            "Negociación: condiciones con cadenas XL en Glovo; agencias -30%",
            "Española nativa en una empresa familiar de Barcelona",
        ],
        "missing_skills": [
            "KAM lado marca en retail físico de belleza (Al Tayer, Ulta, Gold Apple, grandes almacenes)",
            "Negociación de espacios fijos y temporales en tienda y animaciones",
            "NLTs y plantillas de alta de retailers de belleza",
            "Rolling forecasts propios (mensual de lanzamientos, trimestral de core) y fill rate",
            "Trabajo directo con Supply Chain y Demand Planning en una multinacional",
        ],
        "sector_fit": "muy bueno — belleza y fragancia como KAM en Miravia (KIKO incluida); Puig es fragancia y maquillaje de lujo",
        "seniority_fit": "acorde — KAM que reporta al Commercial Director; ella ya fue KAM y hoy es Manager",
        "red_flags": [
            "Su belleza es de marketplace, no de retail físico: preguntarán por espacios en tienda y grandes almacenes",
            "El bloque de forecast es pesado y en DoFreeze ella aporta, no cierra la cifra",
            "Sueldo no publicado: confirmar la banda frente al suelo de 20.000 AED/mes",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Encaje fuerte: el JD pide experiencia comercial con exposición a KAM, idealmente en belleza o lujo, y "
            "Paula fue KAM de 42 marcas de beauty y fragancia en Miravia (KIKO incluida), con +30% GMV QoQ y el canal "
            "Flash Sales ante el CEO. Ha estado en el lado retailer, que es justo con quien negociará en Puig. En "
            "DoFreeze vende a retailers que compran por PO (talabat, Noon, Careem), lleva lanzamientos de punta a "
            "punta y hace la revisión semanal de stock, lo que cubre buena parte de operación y lanzamientos. El "
            "hueco es el retail físico de belleza lado marca (espacios en tienda, animaciones, NLTs) y el forecast "
            "propio. No pide árabe y Puig es de Barcelona."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
