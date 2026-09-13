"""One-off: generate Paula's CV for the **Shopify eCommerce Specialist** role at
**Flat 101** (Zaragoza-based CRO / digital-analytics agency, Spain — remote /
hybrid / on-site).

The role wants a technically-autonomous eCommerce profile to lead Shopify
projects **end-to-end**: store setup, theme configuration, apps/modules,
light HTML/CSS + basic Liquid, base tracking (GTM, GA4, conversion pixels),
and — very desirable — E2E setup & integrations (payment gateways, shipping,
taxes, connecting the store to ERPs/CRMs/logistics). They *explicitly say they
are NOT looking for a pure developer*, but someone who "no se asusta a la hora
de meterse en el código para pequeños ajustes." Plus valued specialty tracks:
SEO/WPO, Retention/CRO & automation, the broader eCommerce ecosystem
(marketplaces: Amazon, **Miravia**, AliExpress, ManoMano…) and Paid Media.

Paula genuinely covers a strong, honest slice of this — so the CV is reframed
toward Shopify + technical eCommerce while staying strictly truthful:
- Owns DoFreeze's Shopify store **end-to-end** (implementation, theme config &
  customisation, apps/modules, collections, discounts, checkout, UX, CRO/AOV).
- Configures the store's commercial setup (payments, shipping, taxes, discounts,
  checkout) and coordinates technical connections to third-party tools
  (logistics, quick-commerce platforms, CRM/ERP) — the "E2E setup & integrations"
  track, framed as real config + coordination (not deep dev).
- Implements **base eCommerce tracking** (Meta Pixel, Google Ads conversion
  tracking, GA4) — exactly the JD's "medición básica / tracking de base".
- **Marketplace ecosystem track, verbatim**: Key Account Manager at **Miravia
  (Alibaba)** — one of the marketplaces the JD names — plus quick-commerce
  (Noon, Talabat, Careem, Deliveroo, Glovo).
- **Paid Media track**: hands-on Meta Ads + Google Ads.
- **AI-first**: builds generative-AI (Claude/GPT) automation — a direct fit for
  Flat 101's stated "integramos la IA en nuestro día a día" culture.
- Google **E-Commerce** + Digital Marketing certificates; CUNEF BBA with an
  **E-Commerce & Fashion** specialisation (thesis 9.5/10).

Honesty guardrails (do NOT cross):
- NO "developer" claims. Technical autonomy is framed exactly as the JD frames
  it — comfortable in the theme/code editor for *small* HTML/CSS adjustments and
  *basic* Liquid/templating logic; not a pure developer.
- NO invented deep technical-SEO / WPO audits (not her specialty) — the SEO/WPO
  track is left out; her honest specialty tracks are Marketplaces + Paid + CRO.
- NO invented WooCommerce / Magento / Prestashop hands-on — omitted; Shopify +
  marketplaces are the real ground.
- NO invented automation-tool names (Klaviyo / Connectif / HubSpot / SFMC) she
  hasn't used — framed generically as email/EDM & automation flows.
- Real DoFreeze job title kept, with a Shopify/eCommerce parenthetical.

Spain context: Paula is a **Spanish national** (currently in Dubai). For a
Spain-based role the relevant, truthful fact is her full right to work in
Spain / the EU — surfaced in the header (EU citizenship, NOT the UAE
"no-sponsorship" claim, which is never made). Skills-block row labels are
relabelled for this output only (template stays pristine).

Fills the real CV template, relabels the skills block + work-authorization line
for this output, converts to PDF via LibreOffice, registers the job for the
dashboard, and lands the package under output/2026-08-22/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Flat 101"
TITLE = "Shopify eCommerce Specialist"
DATE_FOLDER = "2026-08-22"

JOB_DESCRIPTION = """\
Shopify eCommerce Specialist — Flat 101 (CRO / digital-analytics agency,
Zaragoza, Spain; remote / hybrid / on-site).

Shopify is the CMS of the moment, so we are reinforcing the team. We are looking
for a profile with the technical maturity and autonomy to lead our Shopify
projects end-to-end (E2E) — not only implementing a store, but bringing
judgement, spotting improvement opportunities and becoming a reference within
projects, collaborating with multidisciplinary teams.

Mission: not just building the store or activating modules, but contributing
your expertise to audience acquisition and development — from the initial
technical setup and analytics implementation through organic growth (SEO) or
retention (automation), depending on your specialty.

Requirements (someone who, without being a pure developer, isn't afraid to get
into the code for small adjustments and has technical autonomy):
- Shopify mastery: demonstrable experience in implementation, theme
  configuration, app/module activation & management, etc.
- Technical autonomy: able to touch HTML/CSS and understand basic Liquid logic.
- Base measurement: comfortable implementing base eCommerce tracking (Google
  Tag Manager, GA4, conversion pixels and API).
- (Very desirable) E2E setup & integrations: configure payment gateways,
  shipping rules, taxes, and lead the technical connection of the store to
  third-party tools (ERPs, CRMs, logistics tools).

Valued positively (specialty tracks):
- Traffic & SEO: technical audits, WPO optimisation, data markup, web
  architecture for eCommerce.
- Retention & CRO: automation, funnel analysis, segmentation and conversion-flow
  optimisation in tools like Connectif, Klaviyo, HubSpot or Salesforce Marketing
  Cloud.
- eCommerce ecosystem: experience with other CMS (WooCommerce, Magento,
  Prestashop…) or marketplace account management (Amazon, Miravia, AliExpress,
  ManoMano…).
- Paid acquisition: creating and optimising Paid Media campaigns (Google Ads,
  Meta Ads…) and advanced web analytics.
- Languages: English (and others) valued — for international project management
  and technical documentation / the Shopify ecosystem.
"""

ATS = [
    "Shopify", "Shopify eCommerce Specialist", "eCommerce", "e-commerce",
    "CMS", "end-to-end", "E2E", "implementation", "theme configuration",
    "themes", "apps", "modules", "app activation", "HTML", "CSS", "Liquid",
    "technical autonomy", "technical setup", "analytics implementation",
    "tracking", "base tracking", "Google Tag Manager", "GTM", "GA4",
    "Google Analytics", "conversion pixels", "Meta Pixel", "conversion tracking",
    "measurement", "web analytics", "payment gateways", "checkout",
    "shipping rules", "taxes", "integrations", "third-party tools", "ERP", "CRM",
    "logistics", "CRO", "conversion rate optimisation", "funnel", "A/B testing",
    "retention", "automation", "marketing automation", "email", "EDM",
    "segmentation", "SEO", "organic growth", "WPO", "web architecture",
    "marketplaces", "Amazon", "Miravia", "AliExpress", "ManoMano",
    "account management", "listings", "catalogue", "Paid Media", "Google Ads",
    "Meta Ads", "Facebook Ads", "Instagram Ads", "traffic", "acquisition",
    "audience development", "generative AI", "AI automation", "Claude", "GPT",
    "KPI", "ROI", "ROAS", "GMV", "AOV", "quick-commerce", "Noon", "Talabat",
    "Careem", "Deliveroo", "Glovo", "Spain", "EU", "English", "remote", "hybrid",
    "multidisciplinary", "cross-functional",
]

CONTENT = {
    "headline": (
        "Shopify eCommerce Specialist · End-to-End Store Setup, Themes, Apps & Checkout · "
        "Marketplaces (Miravia / Alibaba) & Quick-Commerce · CRO, Tracking, Paid Media & AI Automation"
    ),
    "professional_summary": (
        "eCommerce specialist with 4+ years across Shopify, marketplaces and quick-commerce — comfortable "
        "owning a store end-to-end and, without being a pure developer, getting into the code for small "
        "adjustments. Currently own DoFreeze's Shopify store end-to-end: implementation, theme configuration "
        "and customisation (theme editor, sections, light HTML/CSS, basic Liquid), apps & modules, "
        "collections, discounts, checkout and UX — lifting conversion and AOV through data-led CRO. "
        "Configure the commercial setup (payments, shipping, taxes, discounts) and coordinate technical "
        "connections between the store and third-party tools (logistics, quick-commerce platforms, CRM/ERP), "
        "and implement base eCommerce tracking (Meta Pixel, Google Ads conversion tracking, GA4). Deep "
        "marketplace ecosystem background — Key Account Manager at Miravia (Alibaba), one of the marketplaces "
        "named for this role — plus quick-commerce (Noon, Talabat, Careem, Deliveroo, Glovo) and hands-on "
        "Paid Media (Meta Ads, Google Ads). Early adopter of generative AI (Claude/GPT) for automation and "
        "analytics — a strong fit for Flat 101's AI-first way of working. Google E-Commerce & Digital "
        "Marketing certified; CUNEF BBA specialised in E-Commerce & Fashion. Spanish national with full "
        "right to work in Spain / the EU, English C1, happy working remote, hybrid or on-site."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager (Shopify & eCommerce — End-to-End)",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Global FMCG / D2C distributor | Brands: Befit, Eurocake, Flair | 50+ countries",
            "bullets": [
                "Own the D2C Shopify store end-to-end — implementation, theme configuration and customisation (theme editor, sections, settings, light HTML/CSS edits and basic Liquid/templating logic), activation and management of apps & modules, collections, discounts, checkout and UX — running the store as a single point of ownership from setup through growth",
                "Configure the store's commercial setup end-to-end (payment methods, shipping rules, taxes, discounts and checkout) and coordinate the technical connection between the store and third-party tools — logistics, quick-commerce platforms and CRM/ERP systems (Salesforce, SAP)",
                "Implement base eCommerce tracking and measurement — Meta Pixel, Google Ads conversion tracking and GA4 eCommerce analytics — to attribute performance and steer optimisation",
                "Lift conversion rate (CRO) and average order value through data-led merchandising, funnel and UX optimisation and A/B testing",
                "Plan and optimise Paid Media on Meta Ads and Google Ads (audience building, creative A/B testing, retargeting) to drive qualified traffic and sales, analysing ROI/ROAS",
                "Built an AI-powered (Claude/GPT) automation layer for content, campaign planning, analytics and reporting — cutting manual workload ~40% — directly aligned with Flat 101's AI-first way of working",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion (Marketplace eCommerce)",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Miravia — Alibaba's marketplace in Spain | Top 5 global e-commerce group | 100K+ employees",
            "bullets": [
                "Managed 42 marketplace stores/accounts on Miravia (Alibaba Group) — one of the marketplaces named for this role — owning listings, catalogue, assortment, pricing and promotions, and driving +30% GMV QoQ",
                "Onboarded 30+ new stores in two months as PIC Fragrances through a standardised listing and promotion process — hands-on marketplace catalogue and merchandising operations",
                "Owned the Flash Sales channel (Beauty, Fashion & Home) reporting to the CEO — planning promotional mechanics and merchandising against P&L targets",
                "Continuously analysed conversion, traffic, retention and ROI/ROAS to optimise listing and channel performance and forecasting accuracy",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts (Quick-Commerce / Retail eCommerce)",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees | Retail vertical build-out",
            "bullets": [
                "Part of the team that built Glovo's Retail vertical — onboarding brands onto the platform and expanding the marketplace beyond food delivery into apparel, beauty and non-food categories",
                "Managed strategic accounts (KFC, Taco Bell, La Tagliatella, Sushi Shop), coordinating catalogue, listings and bespoke activations to grow GMV",
                "Led cross-functional coordination across marketing, logistics and customer support to deliver campaigns and increase order volume",
            ],
        },
        {
            "company": "Mondelez International · earlier: Massimo Dutti (Inditex)",
            "role": "FMCG Category Planning (Trainee) & Premium Fashion Retail",
            "dates": "2018 – 2022",
            "location": "Madrid, Spain",
            "context": "Mondelez global FMCG (€36B) + Inditex premium fashion retail",
            "bullets": [
                "Mondelez: built sell-in/sell-out and promo-effectiveness analysis, performance reporting and NPD support (Milka Spread, Mini Suchard) — analytical and data foundation for eCommerce measurement",
                "Massimo Dutti (Inditex): hands-on premium fashion retail and visual-merchandising standards — the product and customer-experience foundation behind data-led online merchandising",
            ],
        },
    ],
    # NOTE: skills-block ROW LABELS are relabelled for Shopify/eCommerce in
    # _relabel_for_role(); the values below are authored to sit under those labels.
    "skills_brand": (  # label -> "Shopify & Technical"
        "Shopify implementation & store setup (end-to-end), theme configuration & customisation "
        "(theme editor, sections, settings), light HTML/CSS edits, basic Liquid / templating logic, "
        "apps & modules activation and management, collections, discounts & checkout configuration, "
        "UX optimisation"
    ),
    "skills_ecommerce": (  # label -> "Marketplaces & Digital Growth"
        "Marketplaces (Miravia / Alibaba, marketplace listings & account management), quick-commerce "
        "(Noon, Talabat, Careem, Deliveroo, Glovo), catalogue & listings, Paid Media (Meta Ads, "
        "Google Ads), email / EDM & marketing-automation flows, generative-AI automation, "
        "organic content basics"
    ),
    "skills_commercial": (  # label -> "eCommerce Setup & Integrations"
        "payment gateways, shipping rules, taxes & checkout setup, third-party integrations "
        "(logistics, quick-commerce platforms, CRM/ERP — Salesforce, SAP), onboarding & platform "
        "coordination, pricing & assortment, distributor & modern-trade management"
    ),
    "skills_data": (  # label -> "Analytics, Tracking & CRO"
        "base eCommerce tracking (Meta Pixel, Google Ads conversion tracking, GA4), CRO & funnel "
        "optimisation, A/B testing, KPI dashboards, ROI, ROAS, GMV, AOV, forecasting, "
        "AI-assisted analysis, Power BI, Tableau, Looker"
    ),
    "skills_tools": (  # label -> "Tools" (unchanged)
        "Shopify, Meta Ads Manager, Google Ads, Google Analytics (GA4), Salesforce, SAP, Power BI, "
        "Tableau, Looker, Generative AI (Claude, ChatGPT), Canva, Microsoft Office (Expert)"
    ),
}

# Skills-block row-label overrides for this Shopify/eCommerce CV (template kept pristine).
ROLE_LABELS = {
    "Brand & Marketing": "Shopify & Technical",
    "E-Commerce & Digital": "Marketplaces & Digital Growth",
    "Commercial": "eCommerce Setup & Integrations",
    "Data & Analytics": "Analytics, Tracking & CRO",
}

# Work-authorisation line: Spain role, Spanish national -> EU right to work.
# (This is EU citizenship, NOT the UAE "no-sponsorship" claim, which is never made.)
WORK_AUTH_OLD = "UAE Residence Visa"
WORK_AUTH_NEW = "Full right to work in Spain / EU (Spanish national)"


def make_job() -> Job:
    return Job(
        id="flat101-shopify-ecommerce-specialist-2026-08",
        title=TITLE,
        company=COMPANY,
        location="Spain (Remote / Hybrid / On-site — Zaragoza)",
        url="https://www.flat101.es/",
        source="linkedin",
        description=JOB_DESCRIPTION,
        raw={
            "query": "Shopify eCommerce Specialist Flat 101 Spain",
            "brand": "Flat 101 (CRO / digital-analytics agency)",
            "division": "Shopify eCommerce — end-to-end store projects for national & international clients",
            "note": "Spain-based role (remote/hybrid/on-site); Paula is a Spanish national (EU right to work). Technical-eCommerce pivot from her brand/marketing core, grounded in real Shopify end-to-end ownership + marketplace (Miravia) experience.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id,
        "title": job.title,
        "company": job.company,
        "location": job.location,
        "url": job.url,
        "source": job.source,
        "description": job.description,
        "salary_raw": None,
        "salary_aed_min": None,
        "salary_aed_max": None,
        "posted_date": DATE_FOLDER,
        "raw": job.raw,
        "ai_score": 82,
        "ai_tier": "Warm",
        "skills_match": [
            "Owns a Shopify store end-to-end (implementation, theme config & customisation, apps/modules, collections, discounts, checkout, UX) — the JD's core requirement",
            "Technical autonomy without being a pure developer — theme/code editor for small HTML/CSS edits + basic Liquid/templating logic (exactly as the JD frames it)",
            "Configures payments, shipping, taxes & checkout and coordinates integrations with third-party tools (logistics, quick-commerce platforms, CRM/ERP) — the 'very desirable' E2E setup & integrations track",
            "Implements base eCommerce tracking (Meta Pixel, Google Ads conversion tracking, GA4) — the JD's 'medición básica'",
            "Marketplace ecosystem track, verbatim: Key Account Manager at Miravia (Alibaba) — a marketplace the JD names — plus quick-commerce (Noon, Talabat, Careem, Deliveroo, Glovo)",
            "Paid Media track: hands-on Meta Ads + Google Ads",
            "AI-first — builds generative-AI (Claude/GPT) automation, matching Flat 101's stated 'integramos la IA en nuestro día a día' culture",
            "Google E-Commerce + Digital Marketing certified; CUNEF BBA specialised in E-Commerce & Fashion",
            "Spanish national — full right to work in Spain / the EU; English C1; happy remote/hybrid/on-site",
        ],
        "missing_skills": [
            "Not a pure developer — Liquid/HTML/CSS positioned as small-adjustment autonomy (which is exactly what the JD asks for), not deep theme development",
            "Technical SEO / WPO is not her specialty track — omitted rather than fabricated; her honest specialty tracks are Marketplaces + Paid Media + Retention/CRO",
            "No hands-on WooCommerce / Magento / Prestashop — Shopify + marketplaces are the real ground; other CMS not claimed",
            "No named experience in specific automation suites (Klaviyo/Connectif/HubSpot/SFMC) — framed generically as email/EDM & automation flows",
        ],
        "sector_fit": "strong (Shopify end-to-end ownership + genuine marketplace ecosystem incl. Miravia; CRO, Paid Media, base tracking; AI-first — the exact profile Flat 101 describes)",
        "seniority_fit": "strong (4+ yrs across Shopify, marketplace & quick-commerce eCommerce; the JD wants maturity + autonomy over a specific seniority band)",
        "red_flags": [
            "Technical-eCommerce pivot from Paula's brand/marketing core — frame around the genuine Shopify end-to-end ownership, setup/integrations and tracking she does, not brand strategy",
            "Spain-based role while Paula is currently in Dubai — mitigated by Spanish nationality / EU right to work and the role's remote option",
            "'Pure developer' expectations must be managed — she is technically autonomous (small code adjustments, basic Liquid), which is precisely what the JD asks for, not more",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Strong, honestly-pivoted fit. Flat 101 wants a technically-autonomous eCommerce specialist to "
            "lead Shopify projects end-to-end (store setup, themes, apps/modules, small code adjustments, "
            "base tracking with GTM/GA4/pixels) and — very desirable — E2E setup & integrations (payments, "
            "shipping, taxes, connecting the store to ERPs/CRMs/logistics), plus valued specialty tracks in "
            "SEO/WPO, Retention/CRO & automation, the marketplace ecosystem (Amazon, Miravia, AliExpress, "
            "ManoMano) and Paid Media. Paula covers a genuine, strong slice: she owns DoFreeze's Shopify "
            "store end-to-end (implementation, theme config & customisation, apps/modules, collections, "
            "discounts, checkout, UX, CRO/AOV), configures the commercial setup (payments/shipping/taxes) "
            "and coordinates third-party integrations, and implements base tracking (Meta Pixel, Google Ads "
            "conversion tracking, GA4). Her marketplace track is verbatim to the JD — Key Account Manager at "
            "Miravia (Alibaba), a marketplace the JD explicitly names — plus quick-commerce and Glovo Retail; "
            "her Paid Media track is real (Meta Ads, Google Ads); and her AI-first automation matches Flat "
            "101's stated culture. Google E-Commerce cert and a CUNEF E-Commerce/Fashion specialisation "
            "reinforce it. Softeners handled truthfully: she is technically autonomous but not a pure "
            "developer (framed exactly as the JD frames it), technical SEO/WPO and other CMS (Woo/Magento/"
            "Presta) are omitted rather than fabricated, and specific automation suites aren't named. Spain "
            "angle covered by Spanish nationality / EU right to work and the role's remote option."
        ),
        "scored_by": "manual:claude",
        "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00",
        "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    """Relabel skills rows + work-auth line for this Shopify CV (template stays pristine)."""
    doc = Document(str(docx_path))

    # Skills-block row labels (exact first-cell match).
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)

    # Work-authorisation line (unique to the header cell).
    cv._replace_in_paragraphs(doc, WORK_AUTH_OLD, WORK_AUTH_NEW)

    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    """Convert DOCX -> PDF headlessly via LibreOffice (docx2pdf/Word is unreliable here)."""
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    """Move output/<Company> - <Role>/ under output/<DATE_FOLDER>/."""
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)
    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
