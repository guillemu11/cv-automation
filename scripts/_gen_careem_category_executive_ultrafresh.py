"""One-off: CV + cover letter for Careem "Category Executive" (Careem Groceries, Ultrafresh
category), Dubai. Applied via Careem's Greenhouse form, 2026-09-13.

Paula's honest angle: she sits on the SUPPLIER side of exactly this table. At DoFreeze she works
with Careem, talabat and Noon dark stores: platform POs, listings, promo calendars, a weekly
stock-risk review by dark store (hers), MSL/assortment review with e-commerce sales, and she
contributes to the monthly sell-in forecast. Plus category lineage (Mondelez Category Planning,
Miravia category expansion, +30% GMV QoQ) and Glovo quick-commerce.

Honesty guardrails:
- NO Ultrafresh experience (fruit & veg, meat, fish, dairy, bakery). JD asks 2–3 yrs in it: the
  main gap. Not claimed; the letter names it.
- Never "own the forecast" (e-commerce manager signs off the number; Paula contributes).
- MSL review on platforms was joint with e-commerce sales. Team of two stated exactly.
- No Arabic, no sponsorship claim ("UAE Residence Visa" only).
- Executive-level role: likely below her 20k AED/month floor; flagged in the dashboard entry.

ONE-PAGE standard. Lands under output/2026-09-13/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cover_letter as cl
from career_ops.generators import cv_generator as cv

COMPANY = "Careem"
TITLE = "Category Executive - Ultrafresh"
DATE_FOLDER = "2026-09-13"
CONTACT = None

JOB_DESCRIPTION = """\
Category Executive — Careem Groceries (Q-commerce), Dubai, UAE. Ultrafresh category.
Support the Category Manager in executing category strategies; source new products and maintain the
Ultrafresh assortment; coordinate with suppliers on product details, pricing and promotional offers;
track category performance (sales, pricing, stock availability) and flag risks/opportunities; plan
and execute weekly and monthly promotions with suppliers and internal stakeholders; ensure accurate,
timely listing of products, content and images with the content team; liaise with supply chain on
availability, new-SKU onboarding and replenishment; market research on trends and price benchmarks;
maintain supplier relationships and prepare periodic business reviews.
Needs: 2–3 years in the Ultrafresh segment (retail or quick/e-commerce preferred); understanding of the
UAE Ultrafresh market (supplier landscape, seasonality, quality standards); communication and
coordination; commercial mindset; detail-oriented; Excel/Google Sheets; commercial KPIs and category
performance metrics; passion for grocery retail. 4 days office / 1 day home.
"""

ATS = [
    "Category Executive", "category management", "Careem Groceries", "Q-commerce", "quick commerce",
    "grocery", "Ultrafresh", "assortment", "sourcing", "new products", "SKU onboarding", "pricing",
    "price benchmarking", "promotions", "weekly promotions", "monthly promotions", "suppliers",
    "supplier management", "business reviews", "stock availability", "replenishment", "supply chain",
    "dark stores", "purchase orders", "listings", "content", "market research", "trends",
    "commercial KPIs", "category performance", "sales tracking", "Excel", "Google Sheets",
    "detail-oriented", "UAE", "Dubai",
]

CV_CONTENT = {
    "headline": (
        "Category & Q-Commerce · Assortment, Pricing & Promotions · Supplier Management · "
        "Dark-store Availability"
    ),
    "professional_summary": (
        "Category and commercial professional with 5 years in FMCG and quick-commerce. Today I work "
        "with Careem, talabat and Noon from the supplier side: listings, promotions, POs and weekly "
        "stock-risk reviews by dark store. Category background at Mondelez and Alibaba's Miravia."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG food group (Befit, Eurocake, Smash) | Careem, talabat, Noon, Deliveroo | 50+ markets",
            "bullets": [
                "Run the weekly stock-risk review by dark store for Careem, talabat and Noon: check whether platform POs cover upcoming demand and flag gaps before products go out of stock",
                "Manage listings, content and images, and plan promotions with the platforms; campaigns measured against control, e.g. SMASH x talabat lifted daily sales +165% (control brand +4.7%)",
                "Reviewed assortment and defined must-stock lists (MSL) by channel with the e-commerce sales team; contribute to the monthly sell-in forecast and launched 6 new products",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Top-5 global e-commerce | 100K+ employees",
            "bullets": [
                "Grew 42 brand accounts +30% GMV QoQ through pricing, assortment and promotions; ran the Flash Sales channel reporting to the CEO",
                "Led category expansion as PIC Fragrances: sourced and onboarded 30+ new stores in two months",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Managed XL partners and in-app promotions; helped build the Retail vertical, coordinating with operations and logistics",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | Chocolate category",
            "bullets": [
                "Sell-in/sell-out and promotional-effectiveness analysis and performance reports for the chocolate category",
            ],
        },
    ],
    "skills_brand": (  # -> "Category"
        "assortment & MSL, new-product sourcing, pricing & benchmarking, promotion calendars, market research"
    ),
    "skills_ecommerce": (  # -> "Q-Commerce"
        "Careem, talabat, Noon, Deliveroo, Glovo: listings & content, dark-store availability, POs"
    ),
    "skills_commercial": (  # -> "Suppliers"
        "supplier & key account management (42), negotiation, business reviews, cross-functional coordination"
    ),
    "skills_data": (  # -> "Data & KPIs"
        "sales, availability, price, promo ROI, sell-in/sell-out, GMV, forecast input, dashboards"
    ),
    "skills_tools": (
        "Excel (advanced), Google Sheets, Power BI, Looker, SAP, Salesforce, Generative AI (Claude)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Category",
    "E-Commerce & Digital": "Q-Commerce",
    "Commercial": "Suppliers",
    "Data & Analytics": "Data & KPIs",
}

CL_PARAGRAPHS = {
    "opening_paragraph": (
        "I would like to apply for the Category Executive role in the Ultrafresh team at Careem Groceries. "
        "I already work with Careem every week, from the other side of the table: as Brand & Marketing "
        "Manager at DoFreeze, a UAE food group, I handle our listings, promotions and stock on Careem, "
        "talabat and Noon. I want to move to the platform side and build a career in grocery category "
        "management."
    ),
    "body_paragraph_1": (
        "Much of this role is already part of my week. I run a weekly stock-risk review by dark store, "
        "checking whether platform purchase orders cover upcoming demand, and flag gaps before products "
        "go out of stock. I keep listings, content and images accurate, plan promotions with the platforms "
        "and measure them properly: a SMASH x talabat campaign lifted daily sales by 165% while a control "
        "brand grew 4.7%. I reviewed our must-stock lists by channel with the e-commerce sales team and "
        "contribute to our monthly sell-in forecast. Before Dubai, I grew 42 brand accounts by 30% GMV "
        "quarter on quarter at Alibaba's Miravia through pricing, assortment and promotions, and "
        "started my career in Category Planning at Mondelez, tracking sell-in, sell-out and promotional "
        "effectiveness."
    ),
    "body_paragraph_2": (
        "To be transparent: I have not yet worked in Ultrafresh. My categories have been packaged food, "
        "beauty and fashion, so the fresh supplier landscape, seasonality and quality standards are what I "
        "would learn first. I would bring what transfers directly: I know how suppliers think about "
        "Careem, how POs and dark-store replenishment work, and how to run a promotion calendar against "
        "sales and availability data in Excel and Google Sheets."
    ),
    "closing_paragraph": (
        "I am based in Dubai and would welcome the chance to discuss how I could support the Category "
        "Manager from day one. Thank you for your consideration."
    ),
}


def make_job() -> Job:
    return Job(
        id="careem-category-executive-ultrafresh-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Dubai, United Arab Emirates",
        url="https://www.careem.com/en-AE/careers/",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={"team": "Careem Groceries — Category Management (Ultrafresh)",
             "workplace": "Dubai, 4 days office / 1 home",
             "note": "Greenhouse form. Gap: no Ultrafresh experience (asks 2–3 yrs)."},
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not posted (Executive band — likely below 20k AED/month floor)",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 55, "ai_tier": "Possible",
        "skills_match": [
            "Supplier-side Q-commerce: listings, promotions, POs with Careem/talabat/Noon",
            "Weekly stock-risk review by dark store (availability, replenishment)",
            "Category lineage: Mondelez Category Planning, Miravia +30% GMV QoQ (42 accounts)",
            "Promotions measured vs control (SMASH x talabat +165%)",
        ],
        "missing_skills": [
            "2–3 years in Ultrafresh (fresh produce, meat, dairy, bakery) — none",
            "UAE fresh supplier landscape, seasonality, quality standards",
        ],
        "sector_fit": "strong on Q-commerce grocery; no fresh category",
        "seniority_fit": "below level — Executive role vs Manager title",
        "red_flags": [
            "No Ultrafresh experience (explicit requirement)",
            "Executive band: salary likely under her 20k AED/month floor",
        ],
        "ats_keywords": ATS,
        "reasoning": "Foot in the door at Careem Groceries; supplier-side experience is real, fresh-category gap is the blocker.",
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Applied",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _pages(pdf: Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(pdf)).pages)


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    final_dir = settings.output_dir / DATE_FOLDER / f"{COMPANY} - {TITLE}"
    dest = final_dir / "01_CV_y_Carta"
    dest.mkdir(parents=True, exist_ok=True)

    cv_docx = cv._fill_template(CV_CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    cl_docx = cl._fill_template(CL_PARAGRAPHS, job, CONTACT)
    cl_pdf = _to_pdf_soffice(cl_docx)

    for src, short in ((cv_pdf, "Paula De Francisco - CV.pdf"),
                       (cl_pdf, "Paula De Francisco - Cover Letter.pdf")):
        moved = dest / src.name
        shutil.move(str(src), str(moved))
        shutil.copy(str(moved), str(dest / short))
        print("OK", moved.name, "PAGES", _pages(moved))

    stray = cv_pdf.parent.parent
    if stray.exists() and stray.resolve() != final_dir.resolve():
        shutil.rmtree(stray, ignore_errors=True)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
