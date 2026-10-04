"""One-off: CV de una página para **Brand Marketing Lead** en **AZADEA** (grupo libanés de retail de moda y
lifestyle: 50+ franquicias internacionales, 650+ tiendas en 14 países de Oriente Medio y África, EAU incluido).
Oferta pegada por Guille el 2026-10-04 desde LinkedIn; el texto no dice ni la marca ni la ciudad. Se asume EAU.

Pide el JD: liderar el equipo de marketing de una marca para subir awareness y **llevar tráfico a las tiendas**;
implementar la estrategia de marca adaptada al país con iniciativas locales; conceptos creativos que generen
venta y tráfico; coordinar a todas las partes en planificación, ejecución y seguimiento; **supervisar el
presupuesto de marketing de la marca** con el Group Marketing Manager; campañas alineadas con el **calendario de
marketing**; estar al día de tendencias de mercado y online; medir el impacto y el ROI de campañas y promociones
con análisis e informes; **reclutar, formar, motivar y evaluar al equipo**.
Requisitos: grado en Marketing o equivalente (máster suma); **4–6 años de marketing online y offline**;
conocimiento de herramientas y técnicas de **investigación de mercado**; **experiencia en retail de moda o en
beauty: OBLIGATORIA**; inglés fluido; MS Office. **No pide árabe.**

Ángulo honesto de Paula:
- Moda y beauty (el "MUST"): 2 años en Miravia (Alibaba) con 42 marcas de beauty, fragancia y moda, KIKO
  Milano incluida, +30% GMV QoQ y canal Flash Sales reportando al CEO; un año en tienda en **Massimo Dutti
  (Inditex)**; en Glovo ayudó a montar el vertical Retail con marcas de moda y lifestyle; grado en ADE por CUNEF
  con especialización en moda y e-commerce (TFG 9,5).
- Lead con equipo: lleva hoy a un diseñador gráfico y a un social media executive (briefs, prioridades,
  desarrollo).
- Calendario de marketing con momentos locales (Ramadán, back to school, Fitness Month, New Year) para varias
  marcas = "local relevant initiatives" + "marketing calendar".
- Online y offline: creators (103 activaciones en 3 campañas), samplings y sorteo in-app con talabat,
  comunidades de running, pádel y yoga. Todas orientadas a venta.
- Presupuesto y ROI: negoció -30% con agencias (AED 81 por pieza de contenido); SMASH × talabat +165% de venta
  diaria frente a +4,7% del grupo de control (~AED 32K incrementales sobre ~AED 8,7K); Befit × Noon +31%.
- Investigación de mercado: Nielsen en Mondelez (sell-in/sell-out, eficacia promocional).

Guardarraíles de honestidad:
- **Tráfico a tienda física de moda**: no lo ha llevado desde el lado marca. Su retail de moda es de dependienta
  (Massimo Dutti) y de plataforma (Miravia, Glovo). No se escribe "footfall" ni "store traffic" como logro.
- **Reclutar y evaluar**: no consta que haya contratado. Se dice "lead and coach", nunca "hired/recruited".
- **Presupuesto**: no consta que sea dueña del A&P entero. Se dice que controla el gasto de campañas y agencias.
- **Forecast**: lo cierra la e-commerce manager; no aparece.
- **Deliveroo**: DoFreeze NO está. Plataformas: talabat, Noon, Careem.
- Creators: "103 activaciones en 3 campañas", nunca "una campaña con 103".
- Sin árabe (aquí no se pide). Visa: "UAE Employment Visa (employer-sponsored)", nunca "no sponsorship".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "AZADEA"
TITLE = "Brand Marketing Lead"
DATE_FOLDER = "2026-10-04"
JOB_ID = "azadea-brand-marketing-lead-2026-10"

JOB_DESCRIPTION = """\
Brand Marketing Lead. AZADEA. Location not stated in the posting (assumed UAE).
Role Purpose: The Brand Marketing Lead is responsible for leading the marketing team to attain set goals, improve
brand awareness and drive traffic to the stores.
Key Responsibilities: Implement brand strategies and ensure adherence with local country requirements including
local relevant initiatives. Develop and implement creative marketing concepts that will support the brand, generate
revenue, and drive traffic to the stores. Liaise with all concerned parties to guarantee adequate planning, execution
and monitoring of marketing activities. Supervise the brand marketing budget expenses and coordinate with the Group
Marketing Manager on measures to be taken accordingly. Ensure all marketing plans/ campaigns are in line with the
marketing calendar and strategy in order to ensure brand image standardization. Stay up to date with the market and
online trends relevant to the brand to ensure that campaigns and promotions are relevant within each market with an
optimal commercial return. Monitor the implementation/ impact of marketing campaigns and promotions to ensure maximum
return on investments followed by analysis and reporting. Assist in recruiting, training, motivating, and evaluating
his / her team to ensure that the department has the necessary skill base and that staff are optimally motivated and
enabled to maximize their potential and contribution to the company.
The Qualifications: Bachelor's Degree in Marketing or equivalent; Master's Degree is a plus. 4-6 years of experience
in a offline and online marketing field. Strong knowledge of market research tools and techniques. Retail Fashion
experience or Beauty experience is a MUST. Fluency in English. Proficiency in MS Office.
AZADEA is committed to equal employment opportunity for all individuals. We will only get in touch if you have been
shortlisted for the role.
"""

ATS = [
    "Brand Marketing Lead", "brand marketing", "brand strategy", "brand awareness", "brand image",
    "brand guidelines", "drive traffic", "retail", "fashion retail", "beauty", "fragrances", "fashion",
    "Inditex", "Massimo Dutti", "KIKO Milano", "local market", "local initiatives", "creative concepts",
    "marketing concepts", "generate revenue", "planning", "execution", "monitoring", "marketing budget",
    "budget control", "Group Marketing Manager", "marketing calendar", "seasonal campaigns", "Ramadan",
    "campaigns", "promotions", "market trends", "online trends", "commercial return", "ROI", "analysis",
    "reporting", "team leadership", "coaching", "training", "motivating", "online marketing",
    "offline marketing", "activations", "sampling", "influencer marketing", "community", "agencies",
    "market research", "Nielsen", "sell-in/sell-out", "Meta Ads", "Google Ads", "social media",
    "MS Office", "Excel", "PowerPoint", "Business Administration", "UAE", "Dubai",
]

CONTENT = {
    "headline": "Brand Marketing · Beauty & Fashion Retail · Online & Offline Campaigns · Team Leadership",
    "professional_summary": (
        "Brand marketer with 4+ years across beauty, fashion retail and consumer brands. Leads brand and marketing "
        "for a homegrown UAE group with a team of two, running local, seasonal online and offline campaigns measured "
        "on sales. Previously managed 42 beauty and fashion brands at Alibaba; started in Inditex retail."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Homegrown UAE F&B / FMCG group | Befit, Eurocake, SMASH, Flair | GCC + 50+ markets",
            "bullets": [
                "Implement brand strategy for a multi-brand portfolio (Befit, Eurocake, the new SMASH brand, Flair) through a year-round marketing calendar built on local moments (Ramadan, back to school, Fitness Month, New Year) across social, retail, talabat, Noon, Careem and our Shopify store",
                "Lead a team of two (graphic designer and social media executive): set priorities, write the creative briefs, coach their work and approve all social, photo and video content",
                "Built creator and community marketing from zero — 103 creator activations over 3 campaigns, talabat samplings and an in-app sweepstakes, running, padel and yoga community events — every concept built to drive purchase",
                "Control campaign and agency spend (fees negotiated -30%) and report ROI on sales: SMASH × talabat lifted daily sales +165% vs +4.7% for the control brand (~AED 32K incremental on ~AED 8.7K spend); Befit × Noon +31% over baseline",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | beauty, fashion & lifestyle",
            "bullets": [
                "Managed 42 beauty, fragrance and fashion brands (incl. KIKO Milano) on assortment, pricing and promotional calendars, delivering +30% GMV QoQ; owned the Flash Sales channel for Beauty, Fashion & Home, reporting to the CEO",
                "Created the Beauty Club and Hot on Social projects to raise brand visibility and loyalty, and onboarded 30+ fragrance stores in two months with trend-driven promotions",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue",
            "bullets": [
                "Helped build the Retail vertical, onboarding fashion and lifestyle brands, and grew XL accounts (KFC, Taco Bell) through joint marketing activations with marketing, operations and customer support",
            ],
        },
        {
            "company": "Mondelez International · Massimo Dutti (Inditex)",
            "role": "Trainee, Category Planning · Sales Associate",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG, Nielsen | Premium fashion retail, Las Rozas Village store",
            "bullets": [
                "At Mondelez, market research with Nielsen (sell-in/sell-out, promotional effectiveness) and NPD launches (Milka Spread, Mini Suchard); at Massimo Dutti, shop-floor sales, styling and Inditex visual merchandising standards",
            ],
        },
    ],
    "skills_brand": (  # -> "Brand & Strategy"
        "brand strategy & image standards, marketing calendar, local & seasonal campaigns, creative concepts"
    ),
    "skills_ecommerce": (  # -> "Online & Offline"
        "social media, creators & communities, sampling & activations, promotions, Meta & Google Ads, EDM"
    ),
    "skills_commercial": (  # -> "Team & Budget"
        "team leadership & coaching, creative briefs, agency management, campaign budget control"
    ),
    "skills_data": (  # -> "Research & ROI"
        "market & trend research, Nielsen, sell-in/sell-out, control groups, ROI analysis & reporting"
    ),
    "skills_tools": (  # -> "Tools"
        "MS Office (Excel, PowerPoint), Adobe (Photoshop, Illustrator), Canva, Power BI, Claude (AI)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Brand & Strategy",
    "E-Commerce & Digital": "Online & Offline",
    "Commercial": "Team & Budget",
    "Data & Analytics": "Research & ROI",
}


def make_job() -> Job:
    return Job(
        id=JOB_ID,
        title=TITLE,
        company=COMPANY,
        location="UAE (assumed — not stated in the posting)",
        url="https://www.linkedin.com/jobs/search/?keywords=AZADEA%20Brand%20Marketing%20Lead",
        source="linkedin",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "AZADEA Brand Marketing Lead",
            "note": "Grupo de retail de moda y lifestyle (50+ franquicias, 650+ tiendas, 14 países). El JD no dice "
                    "ni la marca ni la ciudad. Retail de moda o beauty es OBLIGATORIO; no pide árabe. CV del "
                    "2026-10-04.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    existing = next((j for j in jobs if j.get("id") == job.id), None)
    entry = existing if existing is not None else {}
    entry.update({
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": entry.get("posted_date", DATE_FOLDER), "raw": job.raw,
        "ai_score": 78, "ai_tier": "Warm",
        "skills_match": [
            "Años en rango: el JD pide 4–6 de marketing online y offline, Paula tiene 4+",
            "Beauty y moda (el requisito obligatorio): 42 marcas de beauty, fragancia y moda en Miravia, KIKO Milano incluida, +30% GMV QoQ",
            "Retail de moda: un año en tienda en Massimo Dutti (Inditex) y el vertical Retail de Glovo",
            "Lidera equipo: diseñador gráfico y social media executive",
            "Calendario de marketing con momentos locales (Ramadán, back to school, Fitness Month, New Year)",
            "Online y offline: 103 activaciones de creator en 3 campañas, samplings y sorteo con talabat, comunidades",
            "ROI medido en venta: SMASH × talabat +165% frente a +4,7% del control; agencias -30%",
            "Investigación de mercado con Nielsen (Mondelez); grado en ADE con especialización en moda",
        ],
        "missing_skills": [
            "Tráfico a tienda física desde el lado marca: no lo ha llevado; su retail de moda es de tienda y de plataforma",
            "Reclutar y evaluar equipo: lidera a dos personas, pero no consta que haya contratado",
            "Presupuesto de marca completo: controla el gasto de campañas y agencias, no el A&P entero",
            "Máster: no tiene (es un plus, no requisito)",
        ],
        "sector_fit": "bueno — retail de moda y lifestyle; su beauty y moda vienen del lado plataforma y de tienda",
        "seniority_fit": "en línea — Lead con equipo, 4–6 años pedidos frente a sus 4+",
        "red_flags": [
            "El JD no dice ciudad: si no es EAU, queda fuera del filtro de ubicación",
            "Salario no publicado: confirmar que llega al suelo de 20.000 AED/mes",
            "No dice qué marca: el encaje cambia mucho si es moda (Inditex, Mango) o beauty",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Encaje bueno: AZADEA busca a alguien que lidere el marketing de una marca de retail de moda o beauty "
            "con equipo, calendario, presupuesto y ROI. Paula lidera hoy un equipo de dos, lleva un calendario de "
            "campañas locales online y offline para varias marcas y mide todo en venta. El requisito obligatorio de "
            "moda o beauty lo cubre con Miravia (42 marcas de beauty, fragancia y moda, KIKO incluida), un año en "
            "Massimo Dutti y el vertical Retail de Glovo. No pide árabe. Huecos: no ha llevado tráfico a tienda "
            "física desde una marca, no ha contratado y no es dueña de un presupuesto de marca completo. Si la marca "
            "es española (Inditex, Mango, etc.), su español nativo y su paso por Inditex suman. Falta confirmar "
            "ciudad y salario."
        ),
        "scored_by": "manual:claude", "freshness": entry.get("freshness", "fresh"),
        "discovered_at": entry.get("discovered_at", f"{DATE_FOLDER}T00:00:00+00:00"),
        "status": entry.get("status", "Reviewing"),
    })
    if existing is None:
        jobs.insert(0, entry)
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
