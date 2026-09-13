"""One-off: generate Paula's CV for **Marketing Specialist** at **Zimmermann**
(Australian luxury fashion house; Dubai, UAE — hybrid; LinkedIn Easy Apply).

Full JD was behind LinkedIn's login wall, so content is authored from the role +
company: a regional brand/retail-marketing execution role for a luxury fashion
house — brand campaigns & activations across retail and e-commerce, events &
clienteling, social & influencer, CRM/EDM lifecycle, PR coordination, visual-
merchandising alignment and brand-consistency, with performance reporting.

Paula's fit is strong and honest — she has genuine fashion + luxury-adjacent DNA:
- Inditex (Massimo Dutti): premium fashion retail, visual merchandising & clienteling.
- Miravia (Alibaba): Key Account Manager for Beauty, Fragrances & FASHION — brand
  campaigns (Beauty Club, Hot on Social), +30% GMV QoQ, trend/competitive monitoring.
- DoFreeze: leads brand & marketing — campaigns end-to-end, influencer 0→25-50/
  campaign, social/UGC, CRM/EDM, sampling & seeding, activations, budget & KPIs.
- CUNEF BBA with a Fashion-Industry specialisation; creative eye (Adobe/Canva).

Honesty guardrails: fashion depth is real (Inditex + Miravia Fashion); NO Arabic
claimed (luxury houses run in English); factual "UAE Residence Visa" only.

ONE-PAGE standard: tight summary, 3/2/1/1 bullets, ~6 skills/row, SHORT skills-row
labels (≤~20 chars). Fills the real CV template, relabels the skills rows for a
luxury-fashion marketing profile (template stays pristine), converts to PDF via
LibreOffice, registers the job, verifies 1 page, lands under output/2026-08-30/.
Also drops a short-named 'Paula De Francisco - CV.pdf' copy for portals.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Zimmermann"
TITLE = "Marketing Specialist"
DATE_FOLDER = "2026-08-30"

JOB_DESCRIPTION = """\
Marketing Specialist — Zimmermann (Australian luxury fashion house), Dubai, UAE
(hybrid). Regional brand & retail-marketing execution for a luxury fashion brand:
plan and execute brand campaigns and activations across retail and e-commerce;
coordinate events, clienteling and VIP experiences; run social, content and
influencer/PR; manage CRM/EDM lifecycle communications; align visual merchandising
and brand consistency across touchpoints; support e-commerce/omnichannel launches;
track campaign performance and report. Requires brand/marketing experience
(preferably fashion, luxury, beauty or retail), creative sensibility, strong
organisation and stakeholder skills, and a data-aware, hands-on approach.
"""

ATS = [
    "marketing specialist", "brand marketing", "luxury", "luxury fashion", "fashion",
    "retail marketing", "brand campaigns", "campaign execution", "activations",
    "events", "clienteling", "VIP", "customer experience", "social media",
    "content", "influencer", "UGC", "PR", "seeding", "gifting", "CRM", "EDM",
    "lifecycle", "omnichannel", "e-commerce", "visual merchandising",
    "brand consistency", "brand guidelines", "creative direction", "art direction",
    "go-to-market", "product launches", "KPI", "ROI", "reporting", "budget management",
    "stakeholder management", "premium", "beauty", "Inditex", "Dubai", "GCC", "UAE",
]

CONTENT = {
    "headline": (
        "Marketing Specialist · Luxury Fashion & Beauty · Brand Campaigns, "
        "Social & Influencer · Clienteling, CRM & Retail Activations"
    ),
    "professional_summary": (
        "Brand & marketing specialist with 4+ years across luxury-adjacent fashion, beauty and fragrances — "
        "brand campaigns, social & influencer, CRM and retail/e-commerce activations. Fashion roots at "
        "Inditex (Massimo Dutti) and Alibaba's Miravia (Beauty, Fragrances & Fashion); now lead brand & "
        "marketing at DoFreeze across 50+ markets. Creative eye (art direction & design), data-led and "
        "AI-first. Dubai-based, bilingual Spanish/English."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Multi-brand FMCG group | Brands: Befit, Eurocake, Flair | 50+ countries",
            "bullets": [
                "Lead brand campaigns end-to-end — concept, creative direction, social, influencer, CRM/EDM and events/activations — across retail and e-commerce, keeping the brand consistent at every touchpoint",
                "Built and scale the influencer programme from zero to 25–50 creators per campaign, plus UGC, sampling & seeding — driving brand awareness and engagement",
                "Own the A&P budget and report campaign KPIs (reach, engagement, sell-out, ROI) to leadership — with an AI (Claude/GPT) layer automating content and reporting",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Miravia — Alibaba's marketplace | Top-5 global e-commerce | 100K+ employees",
            "bullets": [
                "Created and led brand campaigns (Beauty Club, Hot on Social) and grew 42 accounts +30% GMV QoQ across a luxury-adjacent beauty, fragrance and fashion portfolio",
                "Ran trend and competitive monitoring across beauty, fragrances and fashion to shape positioning, assortment and creative decisions",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | Retail vertical build-out",
            "bullets": [
                "Delivered bespoke brand activations for strategic partners, coordinating marketing, creative and operations to lift order volume and GMV",
            ],
        },
        {
            "company": "Massimo Dutti (Inditex) · earlier: Mondelez International",
            "role": "Premium Fashion Retail & FMCG Category Planning",
            "dates": "2018 – 2022",
            "location": "Madrid, Spain",
            "context": "Inditex premium fashion retail + Mondelez global FMCG (€36B)",
            "bullets": [
                "Massimo Dutti (Inditex): premium fashion retail, visual merchandising, styling and clienteling at a flagship store — the luxury retail-floor and client experience behind brand marketing; earlier, Mondelez category-planning trainee (sell-in/sell-out, promo analysis, NPD)",
            ],
        },
    ],
    "skills_brand": (  # label -> "Brand & Campaigns"
        "brand campaigns end-to-end, creative & art direction, brand strategy & guidelines, events & "
        "activations, go-to-market, NPD & product launches, storytelling"
    ),
    "skills_ecommerce": (  # label -> "Social & Influencer"
        "influencer & UGC (0→25–50/campaign), social content & calendars, PR / seeding & gifting, "
        "CRM / EDM lifecycle, paid social (Meta), e-commerce & Shopify, digital campaigns"
    ),
    "skills_commercial": (  # label -> "Retail & Clienteling"
        "retail & e-commerce activations, visual merchandising, clienteling & customer experience, "
        "luxury/premium fashion & beauty, assortment, cross-functional stakeholder management"
    ),
    "skills_data": (  # label -> "Data & Reporting"
        "campaign KPIs (reach, engagement, sell-out, ROI/ROAS), reporting & dashboards, market & trend "
        "analysis, budget management, AI-assisted analysis, Power BI, Tableau, Looker"
    ),
    "skills_tools": (  # label -> "Tools"
        "Adobe Creative Suite, Canva, Meta Ads Manager, Google Ads, Shopify, Salesforce, Power BI, "
        "Nielsen, Kantar, Generative AI (Claude, ChatGPT), Microsoft Office (Expert)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Brand & Campaigns",
    "E-Commerce & Digital": "Social & Influencer",
    "Commercial": "Retail & Clienteling",
    "Data & Analytics": "Data & Reporting",
}


def make_job() -> Job:
    return Job(
        id="zimmermann-marketing-specialist-2026-08",
        title=TITLE,
        company=COMPANY,
        location="Dubai, United Arab Emirates (Hybrid)",
        url="https://www.zimmermann.com/careers",
        source="linkedin",
        description=JOB_DESCRIPTION,
        raw={
            "query": "Zimmermann Marketing Specialist Dubai luxury fashion brand",
            "note": "Luxury fashion brand-marketing role. Strong honest fit via Paula's fashion DNA (Inditex/"
                    "Massimo Dutti + Miravia Fashion), brand campaigns, social/influencer, CRM & activations. "
                    "Full JD behind LinkedIn login — content authored from role/company. No Arabic required. "
                    "Factual UAE Residence Visa only.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": None, "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 82, "ai_tier": "Warm-Hot",
        "skills_match": [
            "Fashion DNA: Inditex/Massimo Dutti (premium fashion retail, visual merchandising, clienteling) + Miravia Fashion KAM",
            "Brand campaigns end-to-end + creative/art direction (DoFreeze); Beauty Club / Hot on Social (Miravia)",
            "Social & influencer (0→25-50 creators/campaign), UGC, PR, sampling & seeding",
            "CRM/EDM lifecycle, e-commerce/Shopify, retail activations",
            "Luxury-adjacent beauty & fragrance breadth; +30% GMV QoQ; CUNEF BBA (Fashion-Industry specialisation)",
            "Dubai-based; bilingual ES/EN; AI-first; creative eye (Adobe/Canva)",
        ],
        "missing_skills": [
            "Full JD behind LinkedIn login — tailored from role/company; verify any must-haves",
            "No Arabic (luxury houses run in English; typically not required)",
        ],
        "sector_fit": "strong (luxury fashion brand marketing — fashion + beauty + campaigns + social/influencer + clienteling)",
        "seniority_fit": "on-band (4+ yrs; specialist-level execution)",
        "red_flags": ["Full JD not visible pre-application — skim for any specific must-have"],
        "ats_keywords": ATS,
        "reasoning": (
            "Strong, honest fit for a luxury-fashion brand-marketing specialist. Paula has genuine fashion DNA "
            "(Inditex/Massimo Dutti premium retail + visual merchandising & clienteling; Miravia KAM for Beauty, "
            "Fragrances & Fashion with brand campaigns and +30% GMV QoQ), leads brand campaigns end-to-end at "
            "DoFreeze with social/influencer (0→25-50 creators), UGC, PR, sampling & seeding, CRM/EDM and "
            "activations, and has a creative eye (art direction & design). CUNEF BBA with a Fashion-Industry "
            "specialisation reinforces it. Full JD was behind LinkedIn's login so content is authored from the "
            "role/company; no Arabic required for a luxury house; factual UAE Residence Visa only."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
