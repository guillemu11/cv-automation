"""One-off: CV for "Performance Marketing Manager", Handoff (recruiting for a US clean-ingredient snack
brand: candied nuts, organic dried fruit), remote (found via Jack & Jill, 2026-10-13; Ashby aa95ed13).

Owns paid media for profit, not awareness: Amazon PPC (SP/SB/SD/DSP), Walmart Connect, Meta, TikTok, Google,
YouTube, Reddit, Instacart; marketplaces Amazon, Walmart.com, TikTok Shop, Faire. TACOS, ROAS, contribution
margin, CAC, LTV; ad spend aligned with inventory; weekly dashboards; launches, promotions, listings with
leadership, ops and creative. 3+ yrs performance in e-commerce/CPG. Nice: food/snack CPG, Shopify brands,
structured creative testing, landing-page optimisation, marketplace expansion.

Paula's honest angle:
- DoFreeze is a snack/bakery CPG (Befit, Eurocake, SMASH crisps) sold on talabat, Noon, Careem, Amazon and
  its Shopify store: Advantage+ retargeting 3.7x first-week ROAS, break-even ROAS net of marketplace fees
  (contribution-margin logic), incrementality vs control, weekly dark-store stock-risk review (AI dashboard
  built as a team — memory dofreeze-ecommerce-forecast-ownership), Shopify Q3 CR 2.2% / AOV ~AED 326
  (memory dofreeze-shopify-store-kpis), web orders +86%.
- Retail media from the platform side: Miravia on-site visibility; Glovo in-app ads for XL restaurants.
- Mondelez: chocolate CPG, sell-in/sell-out and promo effectiveness.

Honesty guardrails:
- NO Amazon PPC, Walmart Connect, TikTok Ads/Shop, YouTube, Reddit, Instacart, Helium10/Pacvue claims.
  "Amazon" appears only as a channel SMASH is listed on, never as a skill, until Paula confirms ad work.
- Never "own the forecast"; the stock dashboard is "built with the team".
- Glovo is NOT the Retail vertical. Miravia accounts were beauty and fragrance (Flash Sales covered fashion).
- Remote for a US brand: time zone and work eligibility from Dubai unconfirmed — flagged.

ONE-PAGE standard. Lands under output/2026-10-13/Handoff - Performance Marketing Manager/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Handoff"
TITLE = "Performance Marketing Manager"
DATE_FOLDER = "2026-10-13"

JOB_DESCRIPTION = """\
Performance Marketing Manager — via Handoff for a fast-growing US consumer snack brand (clean-ingredient
candied nuts, organic dried fruit) sold across retail, marketplaces and e-commerce. Remote.
Own and optimise paid media across Amazon, Walmart, Meta, TikTok, Google, Reddit, Instacart and other channels,
focused on profitable growth and contribution margin, not awareness or vanity metrics.
Manage Amazon PPC, Walmart Connect, Meta, TikTok, Google Ads, YouTube, Reddit, Instacart and retail media; grow
Amazon, Walmart.com, TikTok Shop, Faire; keywords, bids, audiences, budgets, structures. Improve TACOS, ROAS,
contribution margin, CAC, LTV; scale spend within efficiency targets; test creatives, audiences, landing pages
and promotions; align ad spend with inventory. Weekly dashboards; analyse campaigns, customers, competitors and
marketplace trends; partner with leadership, operations and creative on launches, promotions, listings,
conversion and marketplace expansion.
3+ years performance marketing for e-commerce or CPG; deep Amazon PPC (SP, SB, SD or DSP); Walmart Connect,
Meta, Google, YouTube, TikTok Ads or TikTok Shop; TACOS, ROAS, contribution margin, CAC, LTV; profitable budget
management; Excel/Sheets; keyword research, bid optimisation, creative testing. Nice: Reddit/Instacart ads,
food/snack CPG, Shopify brands, Helium10/Pacvue/Perpetua/Triple Whale, retail velocity, landing pages.
"""

ATS = [
    "performance marketing", "Amazon PPC", "Sponsored Products", "Walmart Connect", "retail media",
    "Meta Ads", "Google Ads", "TikTok Ads", "TikTok Shop", "TACOS", "ROAS", "contribution margin", "CAC", "LTV",
    "bid optimisation", "keyword research", "creative testing", "landing page optimisation", "marketplaces",
    "inventory", "dashboards", "CPG", "snacks", "Shopify", "promotions", "product launches",
]

CV_CONTENT = {
    "headline": "Performance & Marketplace Marketing · CPG Snacks",
    "professional_summary": (
        "Performance and marketplace marketer with 5 years across CPG and e-commerce. At DoFreeze, runs "
        "Meta and Google Ads for a snack and bakery portfolio sold on talabat, Noon, Careem, Amazon and "
        "Shopify, budgeting on break-even ROAS net of marketplace fees and measuring incrementality against "
        "a control. Previously grew 42 brands +30% GMV QoQ at Alibaba's Miravia."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE snack & bakery CPG | Befit, Eurocake, SMASH | talabat, Noon, Careem, Amazon, Shopify",
            "bullets": [
                "Run paid social and search on Meta and Google Ads: built an Advantage+ catalogue retargeting campaign at 3.7x ROAS and AED 34.5 per purchase in week one; August account at 829K impressions, 2.01% CTR, AED 4.64 CPM",
                "Profit before volume: computed break-even ROAS per brand and channel net of marketplace commissions and moved Google spend to Meta when its return did not cover the cost",
                "Measure incrementality, not last touch: SMASH x talabat lifted daily sales +165% vs +4.7% for a control brand; Befit x Noon added 4,176 incremental units (+31% vs baseline)",
                "Own the Shopify store: web orders +86% and revenue +57% (Q1 vs Jun–Aug 2026) after reviews, basket-based bundles and a 30-day replenishment flow; Q3 conversion 2.2%, AOV ~AED 326",
                "Run the weekly dark-store stock-risk review for talabat, Noon and Careem on an AI dashboard built with the team; 6 product launches, incl. SMASH, now on talabat, Noon and Amazon",
                "Weekly creative-hook tests across 4 social accounts; influencer programme from zero (103 activations at AED 81 per content piece); Claude-based reporting automation (~40% less manual work)",
            ],
        },
        {
            "company": "Miravia & AliExpress (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Marketplace | Beauty & Health department",
            "bullets": [
                "Grew 42 beauty and fragrance brands (incl. L'Oréal Paris, Maybelline, NYX, Redken, KIKO Milano) +30% GMV QoQ through pricing, assortment, promotions and on-site visibility; worked with marketing, product and ops",
                "Owned the Flash Sales channel for Beauty, Fashion & Home (up to ~30% of B&H sales), reporting to the CEO; grew Arabian perfume into a top-growing sub-category",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Food & quick-commerce delivery app | 10K+ employees",
            "bullets": [
                "Co-managed 5 XL restaurant accounts (KFC, Taco Bell, Vicio) on GMV, orders, promo spend and in-app ads; read campaign returns to pick activation days",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Category Planning Trainee",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global snacking CPG",
            "bullets": [
                "Sell-in/sell-out and promo-effectiveness analysis for chocolate; supported NPD (Milka Spread, Mini Suchard)",
            ],
        },
    ],
    "skills_brand": (
        "Meta Ads (Advantage+, catalogue retargeting), Google Ads, budget allocation, creative testing"
    ),
    "skills_ecommerce": (
        "talabat, Noon, Careem, Miravia, AliExpress, Glovo, Shopify, on-site visibility, promotions, bundles"
    ),
    "skills_commercial": (
        "break-even ROAS, contribution margin, incrementality vs control, ROAS, CPA, CVR, AOV, promo uplift"
    ),
    "skills_data": (
        "Excel (Expert), Looker, Power BI, SAP, Shopify analytics, weekly KPI dashboards"
    ),
    "skills_tools": (
        "Claude / Claude Code (automation, dashboards), ChatGPT, Canva, Photoshop"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Paid Media",
    "E-Commerce & Digital": "Marketplaces",
    "Commercial": "Profitability",
    "Tools": "AI & Tools",
}


def make_job() -> Job:
    return Job(
        id="handoff-performance-marketing-manager-snacks-2026-10",
        title=TITLE,
        company=COMPANY,
        location="Remote (US consumer snack brand)",
        url="https://jobs.ashbyhq.com/yourhandoff/aa95ed13-67d4-45c8-a5a1-5efba186b9c0",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={"via": "Jack & Jill AI (tracked job, 2026-10-13)",
             "workplace": "Remote (US brand — time zone / eligibility from Dubai unconfirmed)",
             "sector": "CPG snacks (candied nuts, dried fruit)"},
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not posted (Jack & Jill estimate $80–110K/yr ≈ AED 24.5–33.7K/month)",
        "salary_aed_min": 24500, "salary_aed_max": 33700,
        "posted_date": None, "raw": job.raw,
        "ai_score": 50, "ai_tier": "Stretch",
        "skills_match": [
            "Snack/bakery CPG today (DoFreeze), on talabat, Noon, Careem, Amazon and Shopify",
            "Meta Advantage+ retargeting 3.7x first-week ROAS; break-even ROAS net of marketplace fees",
            "Incrementality vs control; Shopify Q3 CR 2.2%, AOV ~AED 326; web orders +86%",
            "Weekly dark-store stock-risk review (inventory-aware), Mondelez chocolate CPG background",
        ],
        "missing_skills": [
            "Amazon PPC (Sponsored Products/Brands/Display, DSP) — the core of the role",
            "Walmart Connect, TikTok Ads/Shop, YouTube, Reddit, Instacart; Helium10/Pacvue-type tools",
            "~1 year of hands-on paid media (3+ asked); no US market experience",
        ],
        "sector_fit": "excellent (snack CPG across marketplaces)",
        "seniority_fit": "ok (Manager title; light on years of paid)",
        "red_flags": [
            "Amazon PPC is the core requirement and she has none (unless she ran ads for SMASH on Amazon)",
            "Remote for a US brand: likely US hours from Dubai (8–9 h behind)",
        ],
        "ats_keywords": ATS,
        "reasoning": "Snack CPG and profit-first measurement fit; Amazon PPC, Walmart and TikTok Ads are hard gaps; good pay.",
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _pages(pdf: Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(pdf)).pages)


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    final_dir = settings.output_dir / DATE_FOLDER / f"{COMPANY} - {TITLE}"
    dest = final_dir / "01_CV_y_Carta"
    dest.mkdir(parents=True, exist_ok=True)

    cv_docx = cv._fill_template(CV_CONTENT, job)
    _relabel_for_role(cv_docx)
    stray = cv_docx.parent.parent

    clean_docx = dest / "Paula De Francisco - CV.docx"
    shutil.move(str(cv_docx), str(clean_docx))
    # The template writes to output/<Company> - <Role>/ at the root; only ever remove that.
    if stray.exists() and stray.parent.resolve() == settings.output_dir.resolve():
        shutil.rmtree(stray, ignore_errors=True)

    cv_pdf = _to_pdf_soffice(clean_docx)
    # The dashboard only indexes CV_Paula_<Company>_<Title>.pdf; the clean name is the one to send.
    dashboard_pdf = dest / f"CV_Paula_{cv._sanitize_filename(COMPANY)}_{cv._sanitize_filename(TITLE)}.pdf"
    shutil.copy(str(cv_pdf), str(dashboard_pdf))
    print("OK", cv_pdf, "PAGES", _pages(cv_pdf), "| dashboard copy:", dashboard_pdf.name)


if __name__ == "__main__":
    main()
