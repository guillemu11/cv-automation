"""One-off: CV de una página para **B2B Digital Marketing & Brand Specialist** en
**Quantum Talent Group** (Dubai, remoto; staffing tecnológico para programas de
transformación digital en UAE / Arabia Saudí).

Pide el JD, y el acento está en **ejecutar personalmente**, no en estrategia delegada:
rediseñar y construir la web; refrescar identidad visual, mensajes y assets digitales;
poseer el LinkedIn corporativo; crear contenido técnico creíble (AI & Data, Cloud,
Ciberseguridad, Software Engineering, Transformación Digital); convertir proyectos de
cliente en case studies anonimizados; colateral de servicio y capacidades; campañas de
outreach B2B a CTOs, CIOs, Heads of Engineering, líderes de Data/AI y Procurement;
lead generation vía LinkedIn, email, web y campañas; sales enablement (capability decks,
propuestas, one-pagers); analítica y SEO.

Ángulo honesto de Paula: es la marketer que **hace las cosas con sus manos** — construye y
opera la tienda Shopify, su propio portfolio web, landing pages de campaña y decks listos
para cliente; dirige identidad visual y diseño (Adobe, Canva) briefando a su diseñador;
escribe y programa contenido social; y ha construido un sistema de automatización con IA
(Claude) que genera research, contenido, reporting y landings. Su lado B2B real es
comercial: vende A negocios y A TRAVÉS de negocios — 42 cuentas en Miravia (Alibaba),
cuentas XL en Glovo, distribuidores y modern trade — tratando con decisores y cerrando
acuerdos.

Guardarraíles de honestidad:
- Sector: B2C (FMCG, beauty, e-commerce, quick-commerce). NO es marketing B2B de
  tecnología ni de staffing — no se insinúa experiencia en AI/Cloud/Cyber como dominio
  propio. Lo que se afirma es capacidad de construir contenido sobre temas técnicos
  apoyada en su uso real de IA y automatización.
- Web: construye y opera Shopify, landing pages de campaña y su portfolio en producción;
  NO se reclama rediseño corporativo completo en WordPress/Webflow ni desarrollo a medida.
- SEO: trabaja SEO de producto/contenido y AEO/GEO (schema, llms.txt); NO se reclama
  ownership de una estrategia SEO técnica de sitio corporativo.
- Lead generation: campañas de captación y outreach; NO se reclama pipeline B2B
  MQL→SQL ni HubSpot/Marketo (usa Salesforce).
- Visa: solo "UAE Residence Visa"; nunca "no sponsorship needed".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Quantum Talent Group"
TITLE = "B2B Digital Marketing & Brand Specialist"
DATE_FOLDER = "2026-09-22"

JOB_DESCRIPTION = """\
B2B Digital Marketing & Brand Specialist — Quantum Talent Group, Dubai, UAE. Remote, full time.
Quantum provides specialist technology professionals and delivery teams to enterprise organisations,
helping clients access the technical capability required to deliver complex technology programmes.
Hands-on role taking ownership of the digital presence and helping reposition and grow the brand across
the UAE, Saudi Arabia and the wider region. Content and brand should position Quantum around technology
expertise, project delivery and specialist technical capability.
What you'll own: Website — redesign and build a modern website positioning Quantum as a technology
staffing and delivery partner. Brand positioning — refresh visual identity, messaging and digital assets
to create a modern technology-led brand. LinkedIn — own the LinkedIn presence with content on technology,
projects, market intelligence and specialist capabilities. Technology content — credible content across
AI & Data, Cloud & Platform, Cybersecurity, Software Engineering and Digital Transformation. Case studies
— turn successful client projects into compelling anonymised case studies and capability stories. Service
and capability content — digital collateral around technology capabilities, delivery models and specialist
contractor networks. B2B outreach — targeted campaigns reaching CTOs, CIOs, Heads of Engineering, Data/AI
leaders, Procurement and other technology decision-makers. Lead generation — use LinkedIn, email, website
content and targeted campaigns to generate qualified conversations with potential clients. Sales enablement
— capability decks, proposals, one-pagers and collateral that helps the team win new business. Analytics
and SEO — grow relevant website traffic and measure which campaigns and content generate engagement and
commercial opportunities.
We need someone who can personally execute, not simply develop a marketing strategy and outsource
everything. Experience across: B2B technology marketing, website design/build, LinkedIn, content creation,
brand development, lead generation, email outreach, SEO, graphic design, analytics.
"""

ATS = [
    "B2B marketing", "digital marketing", "brand specialist", "brand positioning", "visual identity",
    "messaging", "digital assets", "website design", "website build", "landing pages", "LinkedIn",
    "content creation", "content strategy", "case studies", "capability stories", "collateral",
    "sales enablement", "capability decks", "proposals", "one-pagers", "B2B outreach",
    "targeted campaigns", "decision-makers", "lead generation", "email outreach", "EDM",
    "SEO", "AEO", "analytics", "website traffic", "engagement", "graphic design", "Adobe", "Canva",
    "hands-on", "AI", "data", "digital transformation", "automation", "campaign reporting",
    "stakeholder management", "agencies", "UAE", "Saudi Arabia", "GCC", "remote",
]

CONTENT = {
    "headline": "Digital Marketing & Brand · Hands-On Web, Content & Design · Lead Generation · AI Automation",
    "professional_summary": (
        "Hands-on marketer who builds what she plans: websites and landing pages, brand identity and decks, "
        "LinkedIn and email content, and the analytics behind them — plus an AI automation system (Claude) "
        "that produces research, content and client-ready collateral."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG group (Befit, Eurocake, Flair) | Shopify D2C + marketplaces | 50+ markets",
            "bullets": [
                "Build and run the websites personally — Shopify store end-to-end plus campaign landing pages: structure, copy, design, SEO-ready product content and analytics, lifting conversion rate and AOV",
                "Own brand positioning and visual identity across markets — messaging, digital assets and art direction (Adobe, Canva) — briefing and reviewing all output from a designer and a social media executive",
                "Create and schedule content across LinkedIn, Instagram and TikTok plus EDM campaigns, and run targeted paid campaigns on Meta and Google Ads to generate qualified traffic and leads",
                "Built an AI automation system (Claude) that produces market research, campaign content, KPI reporting and client-ready pitch decks and landing pages — cutting manual workload ~40%",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | B2B partner management | 100K+ employees",
            "bullets": [
                "Sold to and grew 42 business accounts — pitching commercial plans to brand decision-makers, negotiating terms and delivering +30% GMV growth QoQ",
                "Ran outreach that onboarded 30+ new partner stores in two months, and created the Beauty Club and Hot on Social brand projects to build visibility and positioning",
                "Owned the Flash Sales channel reporting to the CEO, presenting performance, insight and recommendations monthly",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Managed and pitched to XL business partners (KFC, Taco Bell, Sushi Shop), closing commercial deals and helping build a new Retail vertical from the ground up",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | Category planning & analytics",
            "bullets": [
                "Built performance and promotional-effectiveness reporting with Nielsen for the chocolate category",
            ],
        },
    ],
    "skills_brand": (  # label -> "Brand & Content"
        "brand positioning, visual identity & messaging, art direction (Adobe, Canva), copywriting, "
        "case studies & decks, social content"
    ),
    "skills_ecommerce": (  # label -> "Web & Channels"
        "website & landing-page build (Shopify, no-code, HTML), LinkedIn, EDM & email outreach, "
        "Meta & Google Ads, SEO / AEO content"
    ),
    "skills_commercial": (  # label -> "B2B & Lead Generation"
        "selling to business decision-makers, partner outreach & onboarding, negotiation, "
        "sales collateral & proposals, key accounts"
    ),
    "skills_data": (  # label -> "Analytics"
        "website traffic & conversion, campaign ROI/ROAS, KPI dashboards, Google Analytics, "
        "Power BI, Looker"
    ),
    "skills_tools": (  # label -> "Tools"
        "Claude / Claude Code (AI automation), Shopify, Canva, Adobe, Meta Ads, Google Ads, "
        "Salesforce, Notion, Power BI"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Brand & Content",
    "E-Commerce & Digital": "Web & Channels",
    "Commercial": "B2B & Lead Generation",
    "Data & Analytics": "Analytics",
}


def make_job() -> Job:
    return Job(
        id="quantum-talent-b2b-digital-marketing-brand-specialist-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE (Remote)",
        url="https://www.linkedin.com/jobs/search/?keywords=Quantum%20Talent%20Group%20B2B%20Digital%20Marketing%20Brand%20Specialist",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Quantum Talent Group B2B Digital Marketing & Brand Specialist Dubai remote",
            "note": "Staffing tecnológico, equipo pequeño (12 en LinkedIn). Rol de una sola persona que "
                    "ejecuta todo a mano: web, marca, LinkedIn, contenido técnico, outreach B2B, SEO y "
                    "diseño. Paula encaja en el 'hazlo tú misma'; el sector (B2B tech/staffing) y el "
                    "contenido técnico son la brecha. 930 candidatos, +100 solicitudes.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 55, "ai_tier": "Warm",
        "skills_match": [
            "Perfil hands-on real: construye web y landings (Shopify, portfolio propio), diseña y escribe",
            "Marca end-to-end: identidad visual, mensajes, dirección de arte y decks listos para cliente",
            "Contenido y outreach: LinkedIn, social, EDM y campañas de captación en Meta y Google Ads",
            "Automatización con IA (Claude): research, contenido, reporting y landings — encaja con contenido técnico",
            "Lado B2B comercial: vende a decisores de negocio (42 cuentas Miravia, cuentas XL en Glovo), negocia y cierra",
        ],
        "missing_skills": [
            "Marketing B2B de tecnología / staffing como sector",
            "Contenido técnico creíble en AI & Data, Cloud, Ciberseguridad, Software Engineering",
            "Rediseño de web corporativa (WordPress/Webflow) y SEO técnico de sitio",
            "Pipeline B2B formal (MQL→SQL), HubSpot/Marketo",
        ],
        "sector_fit": "bajo — B2C FMCG/beauty/e-commerce vs staffing tecnológico B2B",
        "seniority_fit": "por debajo — Specialist vs su título de Manager; ojo al salario",
        "red_flags": [
            "930 candidatos y +100 solicitudes",
            "Sector y contenido técnico son brecha real; el JD insiste en ejecutar personalmente",
            "Nivel Specialist en una agencia de staffing: riesgo de quedar por debajo del suelo de AED 20k",
            "Remoto en una empresa de 12 personas en LinkedIn: verificar estructura y presupuesto",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "El JD busca a alguien que lo haga todo con sus manos — web, marca, LinkedIn, contenido, outreach, "
            "SEO, diseño y analítica — y ese perfil generalista y ejecutor es exactamente Paula, además con una "
            "ventaja poco común: automatiza research y contenido con IA. La brecha es de dominio: el contenido "
            "tiene que sonar creíble en AI/Cloud/Cyber para CTOs y CIOs, y su background es B2C de consumo. "
            "Merece la pena aplicar apoyándose en las piezas construidas (portfolio, landings, decks), "
            "pero comprobando nivel y salario antes de invertir mucho."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
