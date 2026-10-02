"""One-off: CV de una página para **Ecommerce Specialist** en **Rituals** (B Corp, Ámsterdam;
cosmética de baño, cuerpo y hogar; ~1.000 tiendas en 36 países). Dubái, presencial, jornada
completa, LinkedIn Easy Apply, promocionado por técnico de selección. **652 solicitudes (76 desde
ayer)**. Requisito añadido por el anunciante: **más de 3 años de experiencia en Retail**.

Contexto: Rituals deja el modelo de franquicia en Oriente Medio y pasa a una **Joint Venture**
desde 2026 para construir su omnicanal con presencia local. El puesto reporta al **Senior
E-commerce Manager** y es **individual contributor sin reportes** al principio.

Pide el JD: trading diario de la web y site merchandising en varios mercados (productos,
categorías, colecciones, navegación, sorting, ranking, búsqueda, recomendaciones, cross-sell/upsell,
contenido y promociones); calidad de catálogo y PDP, taxonomía y datos de producto, búsqueda
onsite; KPIs diarios/semanales (ventas, tráfico, conversión, AOV, UPT, producto/categoría,
promociones); calendario de trading, lanzamientos y campañas con Marketing, Retail y Merchandising;
apoyo a SEO (metadata, enlazado interno, categorías); apoyo a CRM desde el calendario comercial;
**disponibilidad de stock online** y escalado de huecos a Retail, Merchandising, Supply Chain y
Operaciones; seguimiento de fulfilment y cancelaciones; UAT, mejoras de web, UX/CRO con A/B testing;
informes de trading y **contribuir al forecasting** y business reviews.
Requisitos: grado en ADE / E-commerce; **4–6 años de e-commerce**, idealmente retail, beauty,
luxury o consumer goods; trading/site merchandising hands-on y CMS/plataformas de e-commerce;
**GA4**, dashboards y Excel; nociones de SEO, CRM, UX/CRO y fulfilment omnicanal; varios mercados
con poca supervisión; inglés fuerte; **árabe como ventaja** (no obligatorio).

Ángulo honesto de Paula:
- La tienda Shopify de DoFreeze es suya de principio a fin (catálogo, colecciones, navegación,
  descuentos, UX y checkout), con CRO y AOV: es el site merchandising del JD.
- Disponibilidad de stock online: **hace ella la revisión semanal de riesgo de stock por dark store**
  en Noon, talabat y Careem — comprueba si la PO cubre el riesgo actual o próximo. Encaja con
  "online stock availability" y "reduce stock-related cancellations".
- Surtido: definió el **MSL de la web** ella sola. Aporta al **forecast mensual** de sell-in
  (lo cierra la e-commerce manager) = "contribute to forecasting", literal del JD.
- Calendario de trading por temporada (Ramadán, back to school, Fitness Month, New Year) y 6
  lanzamientos; los creators mandan tráfico a talabat, Noon, Careem y a la web propia.
- **Beauty de verdad**: 2 años en Miravia (Alibaba) del lado plataforma con 42 cuentas de belleza,
  fragancia y moda, incluida KIKO Milano; +30% GMV QoQ; canal Flash Sales reportando al CEO.
- Retail: Massimo Dutti (Inditex) en tienda con visual merchandising; Glovo montando el vertical
  Retail. Grado en ADE por CUNEF con especialización en e-commerce (TFG 9,5).

Guardarraíles de honestidad:
- **GA4**: no consta. Su analítica es Shopify, marketplaces, Power BI/Tableau/Looker y Excel.
  No se escribe GA4.
- **CMS enterprise** (Salesforce Commerce Cloud, SAP Commerce, Magento): no. Solo Shopify.
- **A/B testing en web**: no consta; el A/B testing que tiene es de creatividades en Meta Ads.
- **UPT, búsqueda onsite, recomendaciones, UAT**: sin evidencia; no se mencionan.
- **Fulfilment y cancelaciones**: no son suyos; sí la revisión de riesgo de stock.
- **Dashboard de stock**: se hizo en equipo con Claude ("we built"); la revisión semanal es de ella.
- **Forecast**: "contribute", nunca "own".
- **Nivel**: Specialist sin reportes frente a su Manager actual con dos reportes. Verificar banda
  pronto (suelo 20.000 AED/mes).
- **Deliveroo**: DoFreeze NO está. Plataformas: talabat, Noon, Careem.
- Sin árabe (ventaja, no requisito). Visa: "UAE Employment Visa (employer-sponsored)".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Rituals"
TITLE = "Ecommerce Specialist"
DATE_FOLDER = "2026-10-02"

JOB_DESCRIPTION = """\
Ecommerce Specialist. Rituals (B Corp). Dubai, UAE. On-site, full time.
About Rituals: Rituals is entering a new phase in the Middle East. Starting in 2026, we are transitioning from a
franchise model to a new Joint Venture, to further build our omnichannel brand with a strong local presence.
Reporting to the Senior E-commerce Manager, you will support the day-to-day commercial management and continuous
optimisation of Rituals' e-commerce business across the Middle East, with primary ownership of site merchandising,
online trading, catalogue quality, performance analysis and coordination of online stock availability.
Your role: own day-to-day website trading and site merchandising across markets (products, categories,
collections, navigation, sorting, ranking, search, recommendations, cross-sell/upsell, content and promotions);
maintain catalogue and PDP quality and product discoverability (taxonomy, categorisation, product data
completeness, content consistency, onsite search optimisation); monitor daily/weekly e-commerce KPIs (sales,
traffic, conversion, AOV, UPT, product/category performance, promotional performance) and recommend actions;
coordinate the e-commerce trading calendar, product launches, campaigns and onsite activations with Marketing,
Retail, Merchandising and other teams; support SEO execution (content hygiene, metadata, internal linking,
category optimisation); support CRM planning and execution from the commercial calendar; monitor online stock
availability and fulfilment-location readiness, identify stock gaps and coordinate with Retail, Merchandising,
Supply Chain and Operations; support online fulfilment performance (operational KPIs, cancellations,
customer-impacting exceptions); support UAT, website enhancements, new features and UX/CRO improvements using
data, customer insights and A/B testing; maintain recurring trading reports and contribute to forecasting,
business reviews and ad-hoc commercial analysis.
What you'll bring: Bachelor's degree in Business Administration, Ecommerce or related; 4-6 years of relevant
e-commerce experience, ideally in retail, beauty, luxury or consumer goods; strong hands-on experience in
e-commerce trading/site merchandising and confident use of CMS/e-commerce platforms; commercial and analytical
mindset, comfortable with GA4, dashboards and Excel; understanding of SEO, CRM, UX/CRO and omnichannel fulfilment
fundamentals; highly organised, proactive, able to manage multiple markets and stakeholders with limited
supervision; strong English; Arabic is an advantage.
Indicative success measures: conversion rate and commercial performance of onsite initiatives; catalogue accuracy,
product discoverability and content readiness; online stock availability / reduction of avoidable stock-related
cancellations; timeliness and quality of launches, promotions and site updates; quality and actionability of
trading analysis and reporting. Strong individual-contributor role, no direct reports initially.
Added by the job poster: 3+ years of experience in Retail.
"""

ATS = [
    "e-commerce", "ecommerce specialist", "online trading", "website trading", "site merchandising",
    "online merchandising", "catalogue", "PDP", "product data", "product discoverability", "categories",
    "collections", "navigation", "promotions", "cross-sell", "upsell", "trading calendar", "product launches",
    "campaigns", "onsite activations", "KPIs", "sales", "traffic", "conversion", "AOV",
    "category performance", "promotional performance", "CRO", "UX", "SEO", "metadata", "CRM", "EDM",
    "stock availability", "stock gaps", "supply chain", "forecasting", "trading reports", "business reviews",
    "dashboards", "Excel", "Shopify", "CMS", "e-commerce platforms", "marketplaces", "omnichannel",
    "multi-market", "Middle East", "GCC", "beauty", "fragrance", "retail", "consumer goods",
    "Business Administration", "stakeholder management",
]

CONTENT = {
    "headline": "E-Commerce · Online Trading, Site Merchandising & Stock Availability · Beauty · GCC",
    "professional_summary": (
        "E-commerce professional in Dubai who owns a Shopify store end-to-end — catalogue, merchandising, "
        "promotions, UX and CRO — and runs the weekly online stock-risk review across Noon, talabat and "
        "Careem. Two years platform side at Miravia (Alibaba), trading 42 beauty, fragrance and fashion "
        "accounts, incl. KIKO Milano, to +30% GMV QoQ."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG group (Befit, Eurocake, Flair) | Shopify D2C + Noon, talabat & Careem | GCC + 50+ markets",
            "bullets": [
                "Own the Shopify store end-to-end — catalogue and product pages, collections, navigation, discounts and checkout; merchandise it on conversion and AOV data and set the web must-stock list (MSL)",
                "Run the weekly online stock-risk review by dark store on Noon, talabat and Careem, checking purchase-order coverage against current and upcoming demand to flag gaps before they turn into stock-outs; contribute to the monthly sell-in forecast",
                "Coordinate the trading calendar — Ramadan, back to school, Fitness Month, New Year — with 6 product launches, onsite promotions, EDM and creator traffic to the web and platforms: Befit × Noon delivered +31% over baseline",
                "Maintain recurring performance reports on AI-assisted dashboards across 50+ markets and turn them into prioritised actions",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees | platform side",
            "bullets": [
                "Managed 42 beauty, fragrance and fashion accounts (incl. KIKO Milano) on assortment, product content, pricing and promotional calendars, delivering +30% GMV QoQ; onboarded 30+ fragrance stores in two months",
                "Owned the Flash Sales channel for Beauty, Fashion & Home reporting to the CEO — trading plans against P&L targets, reading conversion, traffic, ROI and retention to reallocate investment",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce platform | €500M+ revenue",
            "bullets": [
                "Helped build out the Retail vertical, onboarding fashion and lifestyle brands, and grew GMV on XL accounts through catalogue work and in-app promotional mechanics",
            ],
        },
        {
            "company": "Mondelez International · Massimo Dutti (Inditex)",
            "role": "Trainee, Category Planning · Sales Associate",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG, Nielsen | Premium fashion retail & visual merchandising",
            "bullets": [
                "At Mondelez, sell-in/sell-out and promotional-effectiveness analysis with Nielsen; at Inditex, store retail, visual merchandising standards and product flow",
            ],
        },
    ],
    "skills_brand": (  # -> "Online Trading"
        "site merchandising, catalogue & product pages, collections & navigation, promotions, trading calendar"
    ),
    "skills_ecommerce": (  # -> "Platforms & CRM"
        "Shopify (owner), Noon, talabat, Careem, Miravia (Alibaba), Glovo, EDM, Meta & Google Ads"
    ),
    "skills_commercial": (  # -> "Stock & Assortment"
        "online stock availability, PO coverage, dark-store stock risk, MSL & assortment, forecast input"
    ),
    "skills_data": (  # -> "Analytics & CRO"
        "conversion, traffic, AOV, category & promo performance, CRO, trading reports, dashboards"
    ),
    "skills_tools": (  # -> "Tools"
        "Excel (Advanced), Shopify, Power BI, Tableau, Looker, Salesforce, SAP, Canva & Adobe, Claude (AI)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Online Trading",
    "E-Commerce & Digital": "Platforms & CRM",
    "Commercial": "Stock & Assortment",
    "Data & Analytics": "Analytics & CRO",
}


def make_job() -> Job:
    return Job(
        id="rituals-ecommerce-specialist-dubai-2026-10",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE (On-site)",
        url="https://www.linkedin.com/jobs/search/?keywords=Rituals%20Ecommerce%20Specialist%20Dubai",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Rituals Ecommerce Specialist Dubai",
            "note": "Rituals pasa de franquicia a Joint Venture en Oriente Medio desde 2026. Reporta al Senior "
                    "E-commerce Manager; individual contributor sin reportes. Easy Apply, promocionado por técnico "
                    "de selección, 652 solicitudes. Requisito añadido: 3+ años en retail. Encaje: Shopify propio, "
                    "revisión semanal de stock por dark store, beauty en Miravia (KIKO). Gaps: GA4, CMS enterprise, "
                    "A/B testing en web; paso atrás de título (Specialist).",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": "2026-09-29", "raw": job.raw,
        "ai_score": 76, "ai_tier": "Warm",
        "skills_match": [
            "Dueña de la tienda Shopify de DoFreeze: catálogo, colecciones, navegación, promociones, UX y CRO",
            "Disponibilidad de stock online: hace la revisión semanal de riesgo por dark store y cobertura de PO",
            "Aporta al forecast mensual, que es justo lo que pide el JD (contribute to forecasting)",
            "Calendario de trading por temporada y 6 lanzamientos con Marketing y plataformas",
            "Beauty y fragancia: 42 cuentas en Miravia, incluida KIKO Milano, +30% GMV QoQ",
            "Trading puro: canal Flash Sales de Miravia reportando al CEO",
            "Retail en Massimo Dutti (Inditex) y en el vertical Retail de Glovo",
            "Grado en ADE (CUNEF) con especialización en e-commerce",
        ],
        "missing_skills": [
            "GA4: no consta en su experiencia",
            "CMS enterprise (Salesforce Commerce Cloud o similar): solo Shopify",
            "A/B testing en la web, UAT y búsqueda onsite: sin evidencia",
            "SEO y CRM solo a nivel básico (EDM)",
        ],
        "sector_fit": "muy bueno — beauty y bienestar, con su experiencia de belleza en Miravia",
        "seniority_fit": "por debajo — Specialist sin reportes frente a su Manager actual con dos reportes",
        "red_flags": [
            "Paso atrás de título y posiblemente de banda: confirmar sueldo (suelo 20.000 AED/mes)",
            "652 solicitudes: conviene una vía directa al Senior E-commerce Manager",
            "Pide GA4 y CMS con soltura, que no puede acreditar",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Funcionalmente es de los encajes más directos: el puesto es trading y merchandising de la web, "
            "disponibilidad de stock online y reporting, que es lo que Paula hace con la tienda Shopify y la "
            "revisión semanal de stock por dark store en DoFreeze. Suma belleza real en Miravia (KIKO Milano) y "
            "base de retail en Inditex. Le restan GA4, un CMS enterprise y A/B testing en web, que no puede "
            "acreditar. El mayor riesgo no es el encaje sino el nivel: es Specialist sin reportes, así que hay "
            "que confirmar la banda salarial pronto."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
