"""One-off: CV for Fiskl AI Accounting "Growth Marketing Manager", remote within 3 h of UK time (found via
Jack & Jill, 2026-10-13; fiskl.com/company/careers/growth-marketing-manager).

AI-accounting SaaS for SMEs, 80% automated. Wants an AI-first, hands-on experimenter: LinkedIn as a performance
channel, AEO/SEO and "prompt hacking" to be the default AI recommendation, MCP servers running experiments 24/7,
viral content, communities, 10+ experiments a week, awareness for outbound. Non-negotiable: daily Claude/GPT use
to build campaigns, analysis and marketing infrastructure. Proof over years; "don't send a CV full of 'increased
MQLs by 30%'". Ideal: did your own books / accounting software.

Paula's honest angle: Claude Code + MCP integrations (profile.yaml ai_automation), Claude system (~40% less manual
work), AI-built landing pages and a multilingual AI avatar that briefs affiliates (memory paula-influencer-real-work),
AEO/GEO (schema markup, Wikidata, llms.txt), giveaway mechanics with a control group, agency-report audit (micro
creators 10–40x the campaign average), weekly hook tests, communities; accounting internship (SAP, Looker).

Honesty guardrails: no B2B, LinkedIn Ads, outbound or SaaS claims. MCP listed as a skill, not tied to a specific
DoFreeze build. Glovo is NOT the Retail vertical. Miravia accounts were beauty and fragrance.
Flags: remote for a UK company (Dubai is 3 h ahead in UK summer, 4 h in winter); a foreign employer cannot
sponsor her UAE visa.

ONE-PAGE standard. Lands under output/2026-10-13/Fiskl - Growth Marketing Manager/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Fiskl"
TITLE = "Growth Marketing Manager"
DATE_FOLDER = "2026-10-13"

JOB_DESCRIPTION = """\
Growth Marketing Manager — Fiskl AI Accounting. Remote (within 3 hours of UK timezone). Mid level (2–5 years).
80% automation today, building toward 95% autonomous accounting for SMEs. Run unconventional B2B campaigns; turn
LinkedIn into a performance channel; make Fiskl the default AI recommendation through SEO and prompt hacking;
build MCP servers that run experiments 24/7; viral accounting content; infiltrate SME communities; ship 10+
experiments weekly; create awareness for outbound. Non-negotiable: use AI daily (Claude/GPT-4) for content,
automation, experimentation and marketing infrastructure. Proof over years: campaigns that worked, growth driven,
customers acquired. Required: shipped campaigns with data; can build automation; moves fast; uses AI to 10x
output. Ideal: did own books (freelancer/founder), accounting/financial software, SME/startup network. Metrics:
1,000+ trials/month in 6 months, £1M pipeline/quarter. Apply with: wildest campaign idea, proof weird ideas work,
how you use AI to scale marketing, why competitors' marketing sucks.
"""

ATS = [
    "growth marketing", "experiments", "AI-first", "Claude", "GPT-4", "MCP", "automation", "AEO", "GEO", "SEO",
    "prompt engineering", "viral content", "communities", "giveaways", "creators", "affiliates", "LinkedIn",
    "B2B", "SMEs", "trials", "pipeline", "accounting", "SAP",
]

CV_CONTENT = {
    "headline": "AI-First Growth Marketer · Experiments · Claude & MCP",
    "professional_summary": (
        "AI-first growth marketer who builds the stack hands-on with Claude Code and MCP integrations. At "
        "DoFreeze, a Claude-based system runs campaign planning, content and reporting (~40% less manual work), "
        "and every idea is judged on incremental sales: a purchase-to-enter talabat giveaway beat its control "
        "brand +165% to +4.7%. Started out in an accounting department (SAP, Looker)."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE F&B / FMCG group | 7 brands incl. Befit, Eurocake, SMASH | sold on talabat, Noon, Careem",
            "bullets": [
                "Built a Claude-based system that automates campaign planning, content, market research and KPI reporting (~40% less manual work); landing pages built with AI agents; an AI avatar briefs affiliates in several languages",
                "Giveaways measured like experiments: a purchase-to-enter cashback draw with talabat lifted SMASH daily sales +165% vs +4.7% for a control brand; another bought followers at AED 0.73 each, 8x cheaper than a promoted reel",
                "Influencer channel from zero: 103 creator activations at AED 81 per content piece (market AED 1,000–3,500); audited the agency report, found reach estimated not measured, and micro-creators at 10–40x the campaign average",
                "Weekly creative-hook tests across 4 social accounts; Befit Crew affiliate programme (personal codes, tiered commissions); running, padel and yoga communities; a Sun & Sand Sports x Fabletics session that led to a listing proposal",
                "Scale or stop on data: computed break-even ROAS net of marketplace fees and moved Google spend to Meta; Advantage+ retargeting at 3.7x ROAS in week one; web orders +86% after reviews, bundles and a replenishment flow",
            ],
        },
        {
            "company": "Miravia & AliExpress (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Marketplace | Beauty & Health department",
            "bullets": [
                "Spotted the Arabian-perfume trend and onboarded 30+ sellers in two months (incl. Lattafa and Arabian Oud distributors), making it a top-growing sub-category",
                "Grew 42 beauty and fragrance brands (incl. L'Oréal Paris, Maybelline, NYX, Redken, KIKO Milano) +30% GMV QoQ; created Hot on Social, a monthly trend showcase; worked with marketing, product and ops",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Food delivery marketplace | 10K+ employees",
            "bullets": [
                "Co-managed 5 XL restaurant accounts (KFC, Taco Bell, Vicio) on GMV, orders, promo spend and in-app ads",
            ],
        },
        {
            "company": "Mondelez International · Línea Directa Aseguradora",
            "role": "Early Career",
            "dates": "2019 – 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | Spanish direct insurer",
            "bullets": [
                "Category Planning Trainee, Mondelez (Aug 2021 – Aug 2022): sell-in/sell-out and promo-effectiveness analysis, NPD support",
                "Accounting Intern, Línea Directa (Jan – Dec 2019): a year inside an accounting department, working in SAP and Looker",
            ],
        },
    ],
    "skills_brand": (
        "growth experiments, giveaway mechanics, creators & affiliates, communities, paid social, email"
    ),
    "skills_ecommerce": (
        "Claude Code, MCP integrations (Shopify, Meta Ads, Gmail, Notion, Canva), AI agents, AI landing pages"
    ),
    "skills_commercial": (
        "AEO / GEO (schema markup, Wikidata, llms.txt), prompt engineering, content, Shopify, CRO"
    ),
    "skills_data": (
        "control-group tests, incrementality, break-even ROAS, ROAS, CPA, Looker, Power BI, Excel (Expert)"
    ),
    "skills_tools": (
        "Claude, ChatGPT, Higgsfield & Hailuo (AI video), Canva, Photoshop, Meta Ads Manager, SAP"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Growth",
    "E-Commerce & Digital": "AI & Build",
    "Commercial": "AI Search & Content",
    "Tools": "Tools",
}


def make_job() -> Job:
    return Job(
        id="fiskl-growth-marketing-manager-2026-10",
        title=TITLE,
        company=COMPANY,
        location="Remote (within 3 h of UK time)",
        url="https://fiskl.com/company/careers/growth-marketing-manager/",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={"via": "Jack & Jill AI (tracked job, 2026-10-13)",
             "workplace": "Remote, within 3 h of UK time (Dubai is +3 h in UK summer, +4 h in winter)",
             "sector": "AI accounting SaaS for SMEs (B2B)"},
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not posted (Jack & Jill estimate £45–65K + bonus ≈ AED 17.4–25.2K/month at 4.65)",
        "salary_aed_min": 17400, "salary_aed_max": 25200,
        "posted_date": None, "raw": job.raw,
        "ai_score": 64, "ai_tier": "Warm",
        "skills_match": [
            "AI-first builder: Claude Code, MCP integrations, Claude automation (~40% less manual work), AI agents",
            "AEO/GEO (schema markup, Wikidata, llms.txt) — 'default AI recommendation' ask",
            "Experiments with a control group; giveaway mechanics; weekly hook tests; communities",
            "Accounting internship (SAP, Looker) — the 'done your own books' nice-to-have",
        ],
        "missing_skills": [
            "B2B / SaaS growth, LinkedIn Ads, trials and pipeline metrics",
            "Outbound support and SME/startup network",
        ],
        "sector_fit": "partial (consumer → B2B SaaS), strong on AI-first",
        "seniority_fit": "good (2–5 yrs mid level)",
        "red_flags": [
            "Remote for a UK company: Dubai is 4 h ahead in UK winter (limit is 3 h)",
            "A foreign remote employer can't sponsor her UAE visa",
        ],
        "ats_keywords": ATS,
        "reasoning": "Unusually strong AI-builder match (Claude, MCP, AEO) plus control-group experiments; B2B SaaS and time zone are the risks.",
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _pages(pdf: Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(pdf)).pages)


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    final_dir = settings.output_dir / DATE_FOLDER / f"{COMPANY} - {TITLE}"
    dest = final_dir / "01_CV_y_Carta"
    dest.mkdir(parents=True, exist_ok=True)

    cv_docx = cv._fill_template(CV_CONTENT, job)
    _relabel_for_role(cv_docx)
    stray = cv_docx.parent.parent

    clean_docx = dest / "Paula De Francisco - CV.docx"
    shutil.move(str(cv_docx), str(clean_docx))
    # The template writes to output/<Company> - <Role>/ at the root; only ever remove that.
    if stray.exists() and stray.parent.resolve() == settings.output_dir.resolve():
        shutil.rmtree(stray, ignore_errors=True)

    cv_pdf = _to_pdf_soffice(clean_docx)
    # The dashboard only indexes CV_Paula_<Company>_<Title>.pdf; the clean name is the one to send.
    dashboard_pdf = dest / f"CV_Paula_{cv._sanitize_filename(COMPANY)}_{cv._sanitize_filename(TITLE)}.pdf"
    shutil.copy(str(cv_pdf), str(dashboard_pdf))
    print("OK", cv_pdf, "PAGES", _pages(cv_pdf), "| dashboard copy:", dashboard_pdf.name)


if __name__ == "__main__":
    main()
