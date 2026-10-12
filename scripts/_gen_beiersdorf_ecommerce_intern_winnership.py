"""One-off: CV de una página **SIN FOTO** para **WINNERSHIP: E-Commerce Intern** en **Beiersdorf Middle
East FZCO** (NIVEA, Eucerin, Hansaplast). Dubái. Contrato limitado de 6 meses, full-time, función Sales.
Contacto: Kartik Kulshrestha (TA Lead). Sexta candidatura a Beiersdorf (antes: Eucerin ABM, KAM Modern
Trade, Global Perf. Marketing & Retail Media, Precision Marketing, Ecommerce Shopper & Customer
Marketing Manager ME), todas en "Reviewing". CV solo (sin carta), a petición de Guille 2026-10-12.

**Sin foto**: la oferta pide subir el CV sin foto para evitar sesgos (mismo tratamiento que el CV del
6-oct para Beiersdorf).

Qué es el puesto: e-commerce de varias marcas en EAU sobre Amazon, Noon, Carrefour y talabat. Auditar
y mejorar contenido en retailers (imágenes, A+, keywords, SEO), métricas de digital shelf, seguimiento
semanal de ventas y dashboards de KPIs, tráfico/conversión/medios, calendario promocional y trackers de
campaña, activaciones y lanzamientos, presupuestos de publicidad y promociones, briefs y assets de
campañas, benchmarking de competidores y deck mensual de tendencias. Pide grado (o estudiándolo).

Ángulo honesto de Paula (hace hoy en Dubái la parte de ejecución que describe el JD):
- Marcas propias en talabat, Noon y Careem: alta de SKUs, contenido e imágenes, precio, mecánicas
  promocionales, calendario por ocasiones (Fitness Month, New Year New Me, Ramadán).
- Medición en ventas: SMASH × talabat +165% venta diaria vs +4,7% del control; Befit × Noon
  "New Year, New Me" +4.176 uds incrementales, +31% sobre baseline.
- Revisión semanal de riesgo de stock por dark store con dashboard hecho con IA (en equipo).
- Meta y Google Ads con ROI/ROAS; briefs de todo lo creativo a su equipo de dos; Shopify de punta a punta.
- Ya conoce Beiersdorf: llevó la cuenta de NIVEA 3 meses en Miravia (sin métricas propias de NIVEA).
- Mondelez: sell-in/sell-out y eficacia promocional.

Guardarraíles de honestidad:
- A+ content, SEO, keyword tagging, digital shelf / content compliance: sin evidencia, no se afirman.
- Amazon y Carrefour: no ha trabajado con ellos; solo talabat y Noon de los que cita el JD.
- Benchmarking de competidores: solo "market research" (sistema con Claude), que sí tiene.
- Presupuesto: Meta/Google Ads con ROI/ROAS; no "own the A&P budget". Forecast: "contribute".
- Deliveroo: DoFreeze NO está. Sin árabe. Visa: la de profile.yaml.

Red flags (anotadas en el dashboard): "intern" es rechazo duro en blacklist.yaml; 6 meses limitado;
estipendio de becario muy por debajo de 20k AED/mes; sobrecualificada (5 años, título Manager); dejar
DoFreeze pone fin a su visado patrocinado.

ONE-PAGE standard. Lands under output/2026-10-12/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Beiersdorf"
TITLE = "E-Commerce Intern (WINNERSHIP)"
DATE_FOLDER = "2026-10-12"

JOB_DESCRIPTION = """\
WINNERSHIP: E-Commerce Intern. Beiersdorf Middle East FZCO. Dubai, UAE. Limited, 6 months / Full-Time.
Job function: Sales. Queries: Kartik Kulshrestha, TA Lead.
Support the online business across Amazon, Noon, Carrefour and talabat for NIVEA, Eucerin and Hansaplast
in the UAE, with exposure to brand, supply chain and finance.
E-Commerce Brand Management: oversee end-to-end e-commerce operations and performance of a handful of
brands; collaborate with marketing, sales and supply chain.
Digital & Sales Excellence: audit and improve product content on retailer platforms; image and A+
content uploads, keyword tagging, SEO; track content compliance and digital shelf metrics; analytics
tools to track KPIs and deliver insights.
Data & Reporting: weekly sales tracking and KPI dashboards; traffic, conversion and media performance;
promotion calendars and campaign performance trackers.
Campaign & Launch Support: e-commerce activations and new product launches; manage budgets for
advertising, promotions and campaigns to maximise ROI; assets and briefs for online campaigns (brand
days, promotions, media banners); monitor execution and consolidate learnings.
Competitive Benchmarking: competitors' visibility, pricing, ratings/reviews and promotions; monthly
insights deck on category trends.
Profile: Bachelor's/Master's in Business, Marketing, Digital or related (or currently pursuing);
analytical, data-driven; interest in e-commerce, digital marketing and FMCG/retail; proactive,
organised; good communication and teamwork.
Note: Beiersdorf encourages candidates to upload the CV without a picture.
"""

ATS = [
    "eCommerce", "e-commerce", "FMCG", "retail", "NIVEA", "Eucerin", "Hansaplast", "Noon", "talabat",
    "Careem", "retailer platforms", "product content", "images", "listings", "pricing", "promotions",
    "promotion calendar", "campaign tracker", "weekly sales tracking", "KPI dashboards", "traffic",
    "conversion", "media performance", "ROI", "ROAS", "Meta Ads", "Google Ads", "budgets", "activations",
    "new product launches", "briefs", "assets", "brand days", "market research", "category trends",
    "insights", "supply chain", "finance", "sell-in", "sell-out", "analytics", "data-driven",
    "Excel", "Power BI", "Shopify", "AI", "Claude", "beauty", "skincare", "UAE", "Dubai",
]

CV_CONTENT = {
    "headline": (
        "E-Commerce · Noon, talabat & Careem · Product Content & Promotions · "
        "Weekly Sales & KPI Tracking · FMCG & Beauty"
    ),
    "professional_summary": (
        "E-commerce and FMCG professional with 5 years on marketplaces and quick commerce. In Dubai I manage "
        "Befit, Eurocake and Flair on talabat, Noon and Careem: listings, content, promotions and weekly "
        "tracking. At Alibaba's Miravia I managed 42 beauty brands, including NIVEA for three months, at "
        "+30% GMV QoQ."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG group (Befit, Eurocake, Flair) | talabat, Noon & Careem + Shopify D2C | 50+ markets",
            "bullets": [
                "Manage our brands on talabat, Noon and Careem: SKU onboarding, product content and images, pricing and promotional mechanics, planned on a seasonal calendar (Fitness Month, New Year New Me, Ramadan)",
                "Measure every activation on sales: SMASH × talabat lifted daily sales +165% (control brand +4.7%); Befit × Noon \"New Year, New Me\" added 4,176 incremental units, +31% over baseline",
                "Run the weekly stock-risk review by dark store with an AI dashboard we built, checking platform POs against demand; contribute to the monthly sell-in forecast",
                "Plan Meta and Google Ads against ROI and ROAS and brief all campaign assets to my team of two (designer + social media executive); own the Shopify store end-to-end; 6 product launches",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees | retailer side",
            "bullets": [
                "Managed 42 beauty, fragrance and fashion brands (incl. NIVEA for three months and KIKO Milano), planning promotions, pricing and assortment for +30% GMV QoQ",
                "Ran the Flash Sales channel reporting to the CEO, tracking traffic, conversion, ROI and ROAS to move promotional investment; onboarded 30+ new stores in two months",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce & delivery platform | €500M+ revenue | platform side",
            "bullets": [
                "Managed XL accounts (KFC, Taco Bell, La Tagliatella, Sushi Shop) with data-led plans and in-app promotions; helped build the Retail vertical",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | Chocolate category",
            "bullets": [
                "Sell-in/sell-out and promotional-effectiveness analysis and performance reports for chocolate; supported the Milka Spread and Mini Suchard launches",
            ],
        },
    ],
    "skills_brand": (  # -> "Retailer Platforms"
        "talabat, Noon, Careem, Miravia, Glovo, Shopify: listings, product content & images, pricing"
    ),
    "skills_ecommerce": (  # -> "Campaigns"
        "promotion calendars, online activations, launches, creative briefs, Meta & Google Ads"
    ),
    "skills_commercial": (  # -> "Commercial"
        "key accounts, promo mechanics, assortment, stock & PO coverage, sell-in forecast input"
    ),
    "skills_data": (  # -> "Data & Reporting"
        "weekly sales tracking, KPI dashboards, traffic & conversion, ROI/ROAS, market research"
    ),
    "skills_tools": (
        "Excel (Advanced), Power BI, Looker, Salesforce, SAP, Shopify, Claude / ChatGPT (AI workflows)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Retailer Platforms",
    "E-Commerce & Digital": "Campaigns",
    "Commercial": "Commercial",
    "Data & Analytics": "Data & Reporting",
}


def make_job() -> Job:
    return Job(
        id="beiersdorf-ecommerce-intern-winnership-2026-10",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE",
        url="https://www.beiersdorf.com/careers",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={
            "query": "Beiersdorf WINNERSHIP E-Commerce Intern Dubai",
            "note": "Becario 6 meses, Sales/e-commerce ME (NIVEA, Eucerin, Hansaplast; Amazon, Noon, Carrefour, "
                    "talabat). CV SIN FOTO (lo pide la oferta). Contacto: Kartik Kulshrestha (TA Lead).",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not posted (internship stipend — far below 20k AED/month floor)",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 40, "ai_tier": "Stretch",
        "skills_match": [
            "Hace hoy la ejecución del JD en talabat y Noon: contenido, imágenes, precio, promociones",
            "Calendario promocional y medición en ventas: SMASH × talabat +165% vs control, Befit × Noon +31%",
            "Seguimiento semanal (riesgo de stock por dark store) y dashboards con IA",
            "Meta y Google Ads con ROI/ROAS; briefs creativos; 6 lanzamientos",
            "Ya conoce Beiersdorf: cuenta de NIVEA 3 meses en Miravia",
            "Análisis sell-in/sell-out y eficacia promocional (Mondelez)",
        ],
        "missing_skills": [
            "A+ content, SEO y keyword tagging en retailers",
            "Digital shelf / content compliance como KPI",
            "Amazon y Carrefour",
            "Benchmarking de competidores (ratings, reviews, visibilidad)",
        ],
        "sector_fit": "muy bueno — FMCG y beauty/skincare, e-commerce en EAU",
        "seniority_fit": "muy por debajo — internship para alguien con 5 años y título de Manager",
        "red_flags": [
            "Internship: rechazo duro en blacklist.yaml (seniority_reject: Intern)",
            "Contrato limitado de 6 meses; estipendio de becario muy por debajo de 20k AED/mes",
            "Dejar DoFreeze termina su visado patrocinado; confirmar qué visado da un becario de 6 meses",
            "Programa WINNERSHIP probablemente orientado a estudiantes / recién graduados",
            "CV sin foto (lo pide la oferta) — usar la versión de este paquete",
            "Sexta candidatura a Beiersdorf; mantener el discurso coherente",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "El contenido del puesto encaja con lo que hace hoy en talabat y Noon, y conoce NIVEA desde Miravia, "
            "pero es una internship de 6 meses: nivel, sueldo y visado juegan en contra. Se prepara el CV porque "
            "Guille lo pidió explícitamente."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _remove_photo(docx_path: Path) -> None:
    """Drop the header photo and merge its cell into the name block (full width)."""
    doc = Document(str(docx_path))
    header = doc.tables[0]
    name_cell, photo_cell = header.rows[0].cells[0], header.rows[0].cells[1]
    for drawing in photo_cell._tc.xpath(".//w:drawing"):
        run = drawing.getparent()
        run.getparent().remove(run)
    merged = name_cell.merge(photo_cell)
    # merge() appends the (now empty) photo paragraph — drop it so the header keeps its height
    paras = merged.paragraphs
    if len(paras) > 1 and not paras[-1].text.strip():
        paras[-1]._p.getparent().remove(paras[-1]._p)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _pages(pdf: Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(pdf)).pages)


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    final_dir = settings.output_dir / DATE_FOLDER / f"{COMPANY} - {TITLE}"
    dest = final_dir / "01_CV_y_Carta"
    dest.mkdir(parents=True, exist_ok=True)

    cv_docx = cv._fill_template(CV_CONTENT, job)
    _relabel_for_role(cv_docx)
    _remove_photo(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)

    moved = dest / cv_pdf.name
    shutil.move(str(cv_pdf), str(moved))
    shutil.copy(str(moved), str(dest / "Paula De Francisco - CV.pdf"))
    print("OK", moved.name, "PAGES", _pages(moved))

    stray = cv_pdf.parent.parent
    if stray.exists() and stray.resolve() != final_dir.resolve():
        shutil.rmtree(stray, ignore_errors=True)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
