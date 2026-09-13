"""One-off: generate Paula's CV for **E-commerce Simulation Expert** at **micro1**
(remote contractor, USD 40–50/hour — AI data lab that turns domain experts into
training data and evaluations for frontier models).

Why this one is a genuinely good fit (no stretching required):
- The hard requirement is "minimum 3 years hands-on e-commerce operations, merchandising
  or catalog management". Paula has ~4 uninterrupted years of exactly that: Miravia
  (Alibaba Group) marketplace KAM for beauty/fragrance/fashion, Glovo's retail vertical,
  and today owning DoFreeze's Shopify store plus its listings on Noon, Talabat, Careem
  and Deliveroo.
- "Product catalog structuring, data normalization, content creation for online stores":
  onboarding 30+ fragrance stores in two months at Miravia meant building their catalogues
  from zero — taxonomy, attributes, product content, pricing — and pushing the same
  assortment across four quick-commerce platforms means normalising the same product data
  into four different schemas. That is the job description, almost literally.
- "AI/LLM prompting, evaluation or training is highly valued" (valued, NOT required):
  she built DoFreeze's AI marketing automation on Claude — writing and iterating prompts,
  and reviewing every output against source data before it ships.
- "Consumer behaviour and online shopping journeys", "analytical thinking and QA mindset":
  CRO and merchandising work on Shopify, sell-in/sell-out and promotional-effectiveness
  analysis at Mondelez.

Honesty guardrails: she has NEVER trained an AI model, never built an LLM evaluation
framework, never worked in simulation or synthetic-data design, and does not know named
eval harnesses. The CV claims prompt design, output QA against source data and AI
workflow automation — all true — and never claims model training, eval frameworks or
red-teaming. No Arabic. Factual "UAE Residence Visa" only, never "no sponsorship needed"
(the visa is paid by her current employer).

ONE-PAGE standard: tight summary, 3/2/1/1 bullets, short skills rows.
Fills the real CV template, relabels the skills rows for this role, converts to PDF via
LibreOffice, registers the job, verifies 1 page, lands under output/2026-09-03/.
Also drops a short-named 'Paula De Francisco - CV.pdf' for the portal upload, plus the
form answers and the two things to check before accepting (04_Aplicacion).
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "micro1"
TITLE = "E-commerce Simulation Expert"
DATE_FOLDER = "2026-09-03"

JOB_DESCRIPTION = """\
E-commerce Simulation Expert — micro1. Contractor. Remote. USD 40–50/hour.
About: micro1 is an AI data lab training frontier models and evaluating AI agents. Domain experts
contribute subject-matter knowledge that micro1 turns into high-quality training data, evaluations
and feedback loops. No prior AI experience required — domain knowledge is what matters.
Key responsibilities: design and build synthetic digital stores with realistic product assortments,
pricing and category structures; generate, structure and maintain comprehensive product catalogs,
ensuring data quality, logical consistency and real-world complexity; develop diverse, authentic
shopping scenarios and user prompts simulating varied consumer intents, preferences and challenging
edge cases; create evaluation rubrics and conduct rigorous QA to assess AI-generated responses for
accuracy, relevance and reasoning; continuously analyse model performance, identify failure patterns
and refine data and scenarios; collaborate with AI and data science teams to translate e-commerce
expertise into training datasets and insights; document processes and share findings.
Required: minimum 3 years hands-on experience in e-commerce operations, merchandising or catalog
management; proven expertise in product catalog structuring, data normalization and content creation
for online stores; experience with AI/LLM prompting, evaluation or training highly valued; strong
analytical thinking and rigorous QA mindset; demonstrated understanding of consumer behaviour and
online shopping journeys; exceptional written and verbal communication; ability to work independently
and deliver high-quality, scalable datasets.
Preferred: past experience training AI models or working in simulation/data design roles; knowledge
of emerging trends in e-commerce, retail analytics or digital merchandising; familiarity with LLM
evaluation frameworks and iterative data improvement methodologies.
Required skills listed: E-commerce Operations, Product Catalog Structuring, AI/LLM Prompting &
Evaluation, Scenario Design, User Simulation, Analytical Thinking, QA Mindset, Consumer Behaviour
Understanding.
"""

ATS = [
    "e-commerce operations", "ecommerce operations", "merchandising", "digital merchandising",
    "catalog management", "catalogue management", "product catalog structuring",
    "category structure", "taxonomy", "product data", "data normalization", "data quality",
    "logical consistency", "attributes", "SKU", "assortment", "pricing", "promotions",
    "content creation", "product listings", "online store", "storefront", "marketplace",
    "Shopify", "quick-commerce", "D2C", "Noon", "Talabat", "Careem", "Deliveroo",
    "AI", "LLM", "prompting", "prompt design", "prompt engineering", "evaluation",
    "AI-generated responses", "QA", "quality assurance", "accuracy", "relevance",
    "scenario design", "user simulation", "shopping scenarios", "edge cases",
    "consumer behaviour", "consumer behavior", "shopping journey", "customer journey",
    "conversion rate optimisation", "CRO", "AOV", "GMV", "analytical thinking",
    "sell-in", "sell-out", "performance analysis", "A/B testing", "documentation",
    "remote", "contractor", "independent delivery", "English",
]

CONTENT = {
    "headline": (
        "E-Commerce Operations & Catalogue Management · Marketplace Merchandising · "
        "AI/LLM Prompting & Output QA · 50+ Markets"
    ),
    "professional_summary": (
        "E-commerce operator with 4+ years hands-on in catalogue management, merchandising and "
        "assortment across marketplaces, quick-commerce and D2C — Alibaba's Miravia, Glovo and a "
        "Shopify storefront owned end-to-end. Builds AI workflows on Claude daily: prompt design, "
        "iteration and output QA against source data."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Consumer brands (Befit, Eurocake, Flair) | Shopify D2C + quick-commerce | 50+ markets",
            "bullets": [
                "Own the Shopify store end-to-end — category structure, product data, collections, imagery, pricing and offers — keeping the catalogue logically consistent across SKUs and lifting conversion and average order value through data-led merchandising",
                "Maintain the same assortment across four quick-commerce catalogues (Noon, Talabat, Careem, Deliveroo), each with its own schema, titles, attributes and promo mechanics — normalising product data per platform and QA-ing every listing before go-live",
                "Built the team's AI automation on Claude: designing and iterating prompts for content, market research and reporting, then reviewing each output against source data before use — cutting manual workload ~40%",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba Group's marketplace in Spain | Top-5 global e-commerce | 100K+ employees",
            "bullets": [
                "Onboarded 30+ new fragrance stores in two months as category PIC — building their catalogues from zero (taxonomy, product content, attributes, pricing) and auditing data quality and consistency before listings went live",
                "Grew 42 beauty, fragrance and fashion accounts +30% GMV QoQ through assortment optimisation, pricing strategy and promotional design, reading SKU-level performance to decide what to push, reprice or delist; owned the Flash Sales channel reporting to the CEO",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce marketplace | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Helped build the retail vertical from zero — structuring new non-food categories and onboarding fashion and lifestyle catalogues onto a marketplace originally designed for food",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | Category planning & analytics",
            "bullets": [
                "Ran sell-in/sell-out and promotional-effectiveness analysis for the chocolate category, building the performance reports behind assortment and launch decisions",
            ],
        },
    ],
    "skills_brand": (  # label -> "E-Commerce Operations"
        "catalogue structuring & taxonomy, product data normalisation, listing content creation, "
        "data quality audits, assortment & range planning, pricing & promo mechanics, "
        "digital merchandising, CRO"
    ),
    "skills_ecommerce": (  # label -> "Platforms & Storefronts"
        "Shopify (owner), Miravia / Alibaba marketplace, Noon, Talabat, Careem, Deliveroo, Glovo, "
        "D2C + marketplace + quick-commerce"
    ),
    "skills_commercial": (  # label -> "AI, Prompting & QA"
        "prompt design & iteration (Claude, ChatGPT), reviewing AI output against source data, "
        "AI workflow automation, structured briefs & checklists, process documentation, "
        "independent remote delivery"
    ),
    "skills_data": (  # label -> "Analysis & Consumer Insight"
        "online shopping journeys, consumer behaviour, SKU & category performance, sell-in/sell-out, "
        "promotional effectiveness, A/B testing, conversion, AOV, GMV, KPI reporting"
    ),
    "skills_tools": (  # label -> "Tools"
        "Shopify, Advanced Excel, Power BI, Tableau, Looker, Google Analytics, Salesforce, SAP, "
        "Claude & ChatGPT, Meta Ads Manager, Canva, Adobe Photoshop & Illustrator"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "E-Commerce Operations",
    "E-Commerce & Digital": "Platforms & Storefronts",
    "Commercial": "AI, Prompting & QA",
    "Data & Analytics": "Analysis & Consumer Insight",
}

FORM = """\
# Formulario micro1 — E-commerce Simulation Expert

Campos exactos del formulario ("Apply now"), listos para copiar:

| Campo | Valor |
|---|---|
| First name | Paula |
| Last name | De Francisco Pérez |
| Phone number | +971 50 386 3656 (país: United Arab Emirates, +971) |
| LinkedIn profile URL | https://www.linkedin.com/in/paula-de-francisco-perez-92a13119a/ |
| Upload your resume (in English, .pdf) | `01_CV_y_Carta/Paula De Francisco - CV.pdf` |

Notas del propio formulario:
- Solo acepta **.pdf** — sube el archivo de nombre corto, no el largo.
- Tras el proceso, te consideran también para otras posiciones que encajen con tu perfil,
  así que el CV que subas define para qué más te llaman: este va cargado de e-commerce
  operations, catálogo y QA, que es justo la familia de encargos que reparte micro1.
- Hay programa de referidos ($100) por si alguien de tu red ya está dentro.

## Qué esperar después

micro1 filtra con un **AI recruiter**: una entrevista por vídeo con un agente que pregunta y
repregunta sobre tu dominio. Para esta posición, prepara tres historias concretas:

1. **Construir un catálogo desde cero** — las 30+ tiendas de fragancias en Miravia: cómo
   decidiste la taxonomía, qué atributos exigías por producto, qué rechazabas antes de publicar.
2. **El mismo producto en cuatro esquemas distintos** — Noon vs Talabat vs Careem vs Deliveroo:
   qué cambia en título, imagen, atributos y mecánica promocional, y qué se rompe cuando no lo
   normalizas.
3. **Comportamiento de compra y casos límite** — qué busca de verdad un comprador de fragancia
   (nombre de casa, notas, tamaño, ¿regalo o reposición?), y dónde una recomendación automática
   se equivoca: tamaños, packs, duplicados, sustitutos, agotados.

Eso es exactamente lo que el puesto llama *scenario design* y *user simulation*.
"""

NOTES = """\
# micro1 — E-commerce Simulation Expert · notas antes de aplicar

## Encaje: el mejor de la semana en requisitos duros

El requisito obligatorio es "mínimo 3 años hands-on en e-commerce operations, merchandising o
catalog management". Paula lleva ~4 años seguidos haciendo justo eso, y el segundo requisito
—"product catalog structuring, data normalization y creación de contenido para tiendas online"—
describe literalmente dos cosas que ya ha hecho: montar de cero el catálogo de 30+ tiendas de
fragancias en Miravia, y mantener el mismo surtido en cuatro plataformas de quick-commerce con
cuatro esquemas de datos distintos.

Lo de IA está listado como *highly valued*, no obligatorio, y la oferta dice explícitamente
"no prior experience in AI is required — your domain knowledge is what matters". Aun así suma:
lleva un año construyendo automatizaciones con Claude, escribiendo e iterando prompts y revisando
cada salida contra la fuente.

## Lo que el CV NO dice (y no debe decir en la entrevista)

- Nunca ha **entrenado un modelo** de IA ni ha trabajado en simulación o diseño de datos sintéticos.
- Nunca ha construido un **framework de evaluación de LLMs** ni conoce las herramientas del sector.
- El CV reivindica diseño de prompts, QA de output contra fuente y automatización de flujos — todo
  cierto — y nada más. Esas tres cosas son *preferred*, no requisitos: no hay que inflarlas.

## Dos cosas que comprobar antes de aceptar (no antes de aplicar)

1. **Es contractor, remoto y por horas** — 40–50 USD/h ≈ 147–184 AED/h. A jornada completa serían
   ~26–32K AED/mes, pero estos encargos de data labs suelen ser por proyecto y horas variables:
   no es un sustituto del sueldo fijo, es un complemento. El suelo de 20K AED/mes del pipeline no
   aplica igual aquí.
2. **Trabajo freelance con visado patrocinado por su empresa actual** — en EAU, facturar a un
   tercero estando bajo visado de empleador normalmente exige permiso de freelance o NOC de la
   empresa. Merece una comprobación antes de firmar nada. No afecta a enviar la candidatura.

## Recomendación

Aplicar. Es rápido (formulario de cinco campos), el encaje en los requisitos duros es alto, y el
propio proceso te mete en su pool para otros encargos de e-commerce del mismo laboratorio.
"""


def make_job() -> Job:
    return Job(
        id="micro1-ecommerce-simulation-expert-remote-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Remote (contractor)",
        url="https://micro1.ai/",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={
            "query": "micro1 E-commerce Simulation Expert remote contractor",
            "note": "AI data lab. Contractor, remote, USD 40-50/hour. Hard requirement is 3+ years of "
                    "e-commerce operations / merchandising / catalog management, which Paula has. AI/LLM "
                    "experience is explicitly 'valued, not required'. Screening runs through micro1's AI "
                    "recruiter (video interview with an agent). Application form: first name, last name, "
                    "phone, LinkedIn URL, resume PDF.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "USD 40-50/hour", "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 78, "ai_tier": "Hot",
        "skills_match": [
            "4+ years hands-on e-commerce operations, merchandising and catalogue management vs the 3-year minimum asked",
            "Catalogue structuring from zero: onboarded 30+ fragrance stores at Miravia — taxonomy, attributes, product content, pricing, pre-launch data-quality audits",
            "Data normalisation in practice: same assortment maintained across four quick-commerce schemas (Noon, Talabat, Careem, Deliveroo)",
            "Owns a Shopify storefront end-to-end — category structure, product data, collections, pricing, offers, CRO",
            "AI/LLM prompting hands-on: built the team's Claude automation, iterating prompts and QA-ing every output against source data (listed as 'highly valued')",
            "Consumer behaviour and shopping journeys: CRO and merchandising decisions read off SKU-level performance",
            "Analytical/QA rigour from Mondelez category planning — sell-in/sell-out and promotional-effectiveness analysis",
        ],
        "missing_skills": [
            "Has never trained an AI model — listed as a preferred qualification",
            "No experience in simulation or synthetic data design roles",
            "No familiarity with named LLM evaluation frameworks or formal eval rubrics",
            "English is C1, not native, and the JD asks for exceptional written and verbal communication",
        ],
        "sector_fit": "excellent — the required domain is e-commerce operations and catalogue management, her core",
        "seniority_fit": "good — contractor/expert contribution, no management scope required",
        "red_flags": [
            "Contractor and hourly, remote: complements rather than replaces a salaried role, and hours are typically project-based",
            "Freelance income under an employer-sponsored UAE residence visa usually needs a freelance permit or an NOC — check before signing",
            "Screening runs through an AI recruiter video interview, not a human first round",
            "Applying also enrolls her in micro1's wider expert pool for other assignments",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Strongest hard-requirement match in the current pipeline: the role's mandatory ask is 3+ years "
            "hands-on e-commerce operations, merchandising or catalog management, and Paula has roughly four "
            "uninterrupted years of exactly that across Miravia (Alibaba), Glovo and DoFreeze. The second "
            "requirement — catalogue structuring, data normalisation and content creation for online stores — "
            "maps one-to-one onto building 30+ fragrance store catalogues at Miravia and maintaining one "
            "assortment across four quick-commerce schemas today. AI/LLM prompting is explicitly 'valued, not "
            "required', and she has real prompt-design and output-QA practice from the Claude automation she "
            "built. The gaps are all in the preferred column: no model training, no simulation/synthetic-data "
            "role, no eval frameworks — so the CV claims none of them. Practical caveats are commercial rather "
            "than professional: hourly contractor pay, and freelance work under an employer-sponsored visa "
            "normally requiring a permit or NOC."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    app_dir = final_dir / "04_Aplicacion"
    app_dir.mkdir(parents=True, exist_ok=True)
    (app_dir / "FORM_micro1.md").write_text(FORM, encoding="utf-8")
    (app_dir / "NOTAS_ANTES_DE_APLICAR.md").write_text(NOTES, encoding="utf-8")
    print("OK_FORM", app_dir / "FORM_micro1.md")

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
