"""One-off: generate Paula's CV for **Performance Marketing & e-Brand Executive**
("eComm Executive") at **Reckitt** — eRB (eCommerce) function, MENAP, Dubai.
Reports to the Performance Marketing Manager. Apply via Reckitt careers portal.

JD ask: execute & optimise retail media across Amazon Search Ads / DSP, Noon,
Carrefour, Talabat, Instashop, Careem (revenue, ROAS, share of search, visibility);
own digital shelf (e-content, brand stores, discoverability, ratings & reviews,
online merchandising); coordinate agencies / platform teams, budgets, reporting;
upper-funnel on platforms (sponsorships, homepage takeovers, influencer activations,
seasonal campaigns); insights from retail media + digital shelf analytics; plan
launches / promo moments with Brand, Sales and KAMs. 2 yrs, Excel/PPT, FMCG preferred.

Paula's honest angle:
- Brand-side activation of three FMCG brands on Noon / Talabat / Careem / Deliveroo:
  seasonal platform campaigns with real sell-out proof — SMASH x talabat +165% daily
  sales vs +4.7% control (Befit); Befit x Noon "New Year, New Me" +4,176 incremental
  units (+31% vs baseline). talabat giveaway + samplings (Ramadan, Back to School,
  Fitness Month). Influencer activations: 103 creator activations, AED 81/content
  piece, 4 agencies managed (negotiated -30%, audited agency reach reporting).
- Marketplace side: Miravia (Alibaba) 42 accounts, onsite visibility / Flash Sales /
  promotions, +30% GMV QoQ, ROI/ROAS/conversion reporting to the CEO.
- Digital shelf: Shopify D2C catalogue / content / merchandising / CRO; listings and
  content on platforms.
- FMCG: Mondelez promo effectiveness, sell-in/sell-out.

Honesty guardrails:
- NO Amazon Search Ads / Amazon DSP, NO Carrefour or Instashop retail media — not
  claimed. Paula should confirm whether she has run PAID sponsored placements on
  Noon/Talabat (CV says "visibility & promo activations", not "sponsored ads").
- No share-of-search tooling claimed. No Arabic (JD doesn't require it).
- Factual "UAE Residence Visa" only — never "no sponsorship needed".
- Seniority: "Executive", 2 years — Paula is 5 yrs at Manager level (over-qualified);
  pay likely below the 20k AED floor — flagged.

ONE-PAGE standard. Lands under output/2026-09-13/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Reckitt"
TITLE = "eComm Executive"
DATE_FOLDER = "2026-09-13"

JOB_DESCRIPTION = """\
eComm Executive (Performance Marketing & e-Brand Executive) — Reckitt eRB, MENAP. Dubai.
Reports to the Performance Marketing Manager. Supports growth of Reckitt brands across
e-commerce channels in MENAP: retail media, digital shelf excellence, e-content and
marketplace activations to drive visibility, conversion and profitable growth. Digital
extension of the brand team, partnering with Marketing, Sales and platform stakeholders.

Strategic: support e-commerce growth strategies for all Reckitt brands across
marketplaces, e-grocery and quick commerce; generate insights from retail media
performance, digital shelf analytics, consumer behaviour and competitive activity;
partner with Brand Marketing, Sales and KAMs on launches, promotional events and key
commercial moments.

Operational: execute and optimise retail media campaigns (Amazon Search Ads, Amazon DSP,
Noon, Carrefour, Talabat, Instashop, Careem) against revenue, ROAS, share of search and
visibility KPIs; own digital shelf excellence (e-content, brand stores, discoverability,
ratings & reviews, online merchandising); coordinate agencies and platform teams, manage
budgets, reporting and recommendations; support upper-funnel initiatives — sponsorships,
homepage takeovers, influencer activations, seasonal campaigns on e-commerce platforms.

Experience: Bachelor's in Marketing/Business/Economics/Engineering; 2 years in
E-commerce, Digital/Performance/Brand Marketing or commercial; marketplace advertising,
retail media, FMCG preferred; advanced Excel & PowerPoint; Amazon Ads, DSP, Google Ads,
Meta Ads or retailer media platforms an advantage; strong analytical, planning, project
and stakeholder management; cross-functional communication with Marketing, Sales,
Agencies and Retailer Partners.
"""

ATS = [
    "eCommerce", "e-commerce", "performance marketing", "retail media", "marketplace advertising",
    "retailer media platforms", "digital shelf", "digital shelf analytics", "e-content",
    "brand stores", "product discoverability", "ratings & reviews", "online merchandising",
    "marketplace activations", "e-grocery", "quick commerce", "Noon", "Talabat", "Careem",
    "Carrefour", "Instashop", "Amazon", "ROAS", "revenue", "share of search", "visibility",
    "conversion", "insights", "competitive activity", "consumer behaviour", "new product launches",
    "promotional events", "commercial moments", "seasonal campaigns", "homepage takeovers",
    "sponsorships", "influencer activations", "upper-funnel", "agencies", "budgets", "reporting",
    "Brand Marketing", "Sales", "Key Account Managers", "stakeholder management",
    "Meta Ads", "Google Ads", "Excel", "PowerPoint", "FMCG", "MENAP", "GCC", "UAE",
]

CONTENT = {
    "headline": (
        "eCommerce & Retail Media · Digital Shelf · Seasonal & Influencer Activations on "
        "Noon, Talabat, Careem"
    ),
    "professional_summary": (
        "FMCG e-commerce marketer with 5 years across brands and marketplaces: today activating three FMCG "
        "brands on Noon, Talabat, Careem and Deliveroo in Dubai with campaigns measured on sell-out, and "
        "previously two years inside Alibaba's Miravia marketplace driving visibility, promotions and "
        "+30% GMV QoQ for 42 brands."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG group (Befit, Eurocake, Smash) | 50+ markets | Noon, Talabat, Careem, Deliveroo + Shopify",
            "bullets": [
                "Run seasonal e-commerce campaigns on quick-commerce and marketplaces (Ramadan, Back to School, Fitness Month, New Year): SMASH x talabat lifted daily sales +165% vs +4.7% for a control brand; Befit x Noon 'New Year, New Me' added 4,176 incremental units (+31% vs baseline)",
                "Own platform activations with talabat and Noon teams — listings, e-content, visibility, promo mechanics, in-app giveaway and samplings — plus Meta and Google Ads and the Shopify store, reporting ROAS and sell-out",
                "Launched influencer activations from zero: 103 creators across 3 campaigns at AED 81 per content piece; manage 4 agencies (negotiated -30%, audited their reach reporting) and lead a team of two",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's lifestyle marketplace | Top-5 global e-commerce | 100K+ employees",
            "bullets": [
                "Managed 42 brand accounts from the marketplace side — onsite visibility, campaign slots, promotions, pricing, assortment and content — delivering +30% GMV QoQ",
                "Owned the Flash Sales channel (Beauty, Fashion & Home) reporting to the CEO; tracked traffic, conversion, ROI, ROAS and retention to steer investment",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Grew XL accounts (KFC, Taco Bell, Sushi Shop) through in-app promotions and data-led activation plans; helped build the Retail vertical onboarding lifestyle brands",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | Category planning & analytics",
            "bullets": [
                "Analysed promotional effectiveness and sell-in/sell-out for the chocolate category and supported NPD launches (Milka Spread, Mini Suchard)",
            ],
        },
    ],
    "skills_brand": (  # label -> "Retail Media & Activation"
        "platform activations (Noon, Talabat, Careem), promo mechanics, seasonal campaigns, "
        "launches, influencer activations, Meta & Google Ads"
    ),
    "skills_ecommerce": (  # label -> "Digital Shelf"
        "listings & e-content, online merchandising, discoverability, Shopify store, conversion (CRO), "
        "assortment & pricing"
    ),
    "skills_commercial": (  # label -> "Stakeholders"
        "platform teams, agencies (4), Sales & key accounts, distributors, budget management, team of two"
    ),
    "skills_data": (  # label -> "Insights & KPIs"
        "ROAS, ROI, GMV, conversion, incremental units vs control, sell-in / sell-out, promo effectiveness"
    ),
    "skills_tools": (  # label -> "Tools"
        "Excel & PowerPoint (advanced), Meta Ads Manager, Google Ads, Google Analytics, Power BI, "
        "Shopify, Generative AI (Claude)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Retail Media & Activation",
    "E-Commerce & Digital": "Digital Shelf",
    "Commercial": "Stakeholders",
    "Data & Analytics": "Insights & KPIs",
}


def make_job() -> Job:
    return Job(
        id="reckitt-ecomm-executive-menap-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE",
        url="https://careers.reckitt.com/",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Reckitt eComm Executive Performance Marketing e-Brand MENAP Dubai",
            "note": "Strong function fit: brand-side seasonal/influencer activations on Noon/Talabat/Careem "
                    "with sell-out proof + marketplace-side (Miravia) + FMCG. Gaps: Amazon Ads/DSP, "
                    "Carrefour/Instashop retail media, share of search. Executive level — likely below "
                    "20k AED floor.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed (Executive band — confirm vs 20k AED floor)",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 78, "ai_tier": "Strong",
        "skills_match": [
            "Seasonal campaigns on the exact platforms named (Noon, Talabat, Careem) with sell-out proof: SMASH x talabat +165% vs control; Befit x Noon +4,176 incremental units",
            "Influencer activations on e-commerce platforms: 103 creators, AED 81/content piece, 4 agencies managed",
            "Marketplace-side retail media view: Miravia 42 accounts, visibility & promotions, +30% GMV QoQ, ROAS/ROI reporting",
            "FMCG background (DoFreeze, Mondelez) + Meta/Google Ads + Shopify digital shelf",
        ],
        "missing_skills": [
            "No Amazon Search Ads / Amazon DSP",
            "No Carrefour or Instashop retail media; share-of-search tooling not used",
            "Paid sponsored-placement management on Noon/Talabat to be confirmed",
        ],
        "sector_fit": "strong — FMCG e-commerce in UAE quick-commerce and marketplaces",
        "seniority_fit": "below her level — Executive, 2 years; Paula is 5 years at Manager level",
        "red_flags": [
            "Executive title: pay likely below the 20k AED/month floor — confirm before investing",
            "Amazon Ads / DSP is the first platform listed in the JD",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Good functional fit: the role is brand-side e-commerce activation on Noon, Talabat and Careem, "
            "which Paula does today with measured sell-out results, plus influencer and seasonal campaigns on "
            "platforms — both explicit JD asks. Real gaps are Amazon Ads/DSP and Carrefour/Instashop. The main "
            "risk is level and pay: an Executive role reporting to a Performance Marketing Manager."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
