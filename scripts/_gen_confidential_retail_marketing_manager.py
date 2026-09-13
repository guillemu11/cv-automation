"""One-off: generate Paula's CV for **Marketing Manager** at **Confidential Careers**
(confidential retail employer, Dubai — LinkedIn, "Empleo verificado", 899 applicants).

JD (from the posting): commercially minded and highly creative Marketing Manager to
lead brand, digital, community and customer engagement across the Retail industry.
Blend of creativity, analytical thinking and execution excellence; building brands,
creating meaningful customer experiences, leveraging digital channels for growth and
developing community connections. Drives customer acquisition, retention, brand
awareness, traffic, sales growth and engagement. Enthusiasm for outdoor activities,
adventure sports and active lifestyles highly advantageous.
Experience asked: 5–8 yrs in retail / consumer / lifestyle / sports / fashion / outdoor
brands; digital-first environments; acquisition & retention; managing agencies,
influencers and partnerships; UAE market knowledge and network preferred.

Paula's fit is genuine and honest:
- RETAIL + LIFESTYLE/FASHION line: Miravia (Alibaba) KAM for Beauty, Fragrances &
  Fashion; Glovo's Retail vertical (onboarded fashion & lifestyle brands); Inditex/
  Massimo Dutti shop floor at the start of her career.
- DIGITAL-FIRST GROWTH: owns the Shopify store end-to-end (CRO, AOV), Meta/Google Ads,
  EDM, and UAE quick-commerce (Noon, Talabat, Careem, Deliveroo).
- COMMUNITY + INFLUENCERS + AGENCIES/PARTNERSHIPS: built the influencer programme from
  zero (25–50 creators/campaign, sampling & seeding); created the Beauty Club and
  Hot on Social community/loyalty projects at Miravia; negotiates partners/agencies.
- ACQUISITION & RETENTION ANALYTICS: ROI/ROAS, conversion, traffic and retention
  analysis at Miravia; KPI/A&P ownership at DoFreeze.

Honesty guardrails: asks 5–8 yrs and Paula has 4+ (5+ counting her Inditex retail
year) — summary says "4+ years", no inflation; the outdoor/adventure-sports affinity
is NOT claimed on the CV (unverified personal interest — it belongs in the cover
letter/interview if true); no Arabic claimed; factual "UAE Residence Visa" only,
never "no sponsorship needed".

ONE-PAGE standard: tight summary, 3/2/1/1 bullets, ~6 skills/row, SHORT labels.
Fills the real CV template, relabels the skills rows for this role, converts to PDF
via LibreOffice, registers the job, verifies 1 page, lands under output/2026-09-03/.
Also drops a short-named 'Paula De Francisco - CV.pdf' copy for portals / Easy Apply.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Confidential Careers"
TITLE = "Marketing Manager"
DATE_FOLDER = "2026-09-03"

JOB_DESCRIPTION = """\
Marketing Manager — Confidential Careers (confidential retail employer), Dubai, UAE. On-site, full time.
We are seeking a commercially minded and highly creative Marketing Manager to lead brand, digital,
community and customer engagement initiatives across the Retail industry. This role requires a unique
blend of creativity, analytical thinking and execution excellence. The successful candidate will be
passionate about building brands, creating meaningful customer experiences, leveraging digital channels
for growth, and developing strong community connections. An enthusiasm for outdoor activities, adventure
sports and active lifestyles would be highly advantageous. The role will be responsible for driving
customer acquisition, retention, brand awareness, traffic, sales growth and customer engagement.
Experience: 5–8 years of marketing experience within retail, consumer, lifestyle, sports, fashion or
outdoor brands. Proven experience in digital-first marketing environments. Strong understanding of
customer acquisition and retention strategies. Experience managing agencies, influencers and
partnerships. UAE market knowledge and network preferred.
"""

ATS = [
    "marketing manager", "brand", "brand awareness", "brand building", "retail",
    "consumer", "lifestyle", "sports", "fashion", "outdoor brands", "digital-first",
    "digital marketing", "customer acquisition", "customer retention", "customer engagement",
    "customer experience", "community", "community building", "loyalty", "traffic",
    "sales growth", "e-commerce", "Shopify", "CRO", "conversion", "performance marketing",
    "Meta Ads", "Google Ads", "paid social", "EDM", "CRM", "influencers", "influencer marketing",
    "UGC", "agencies", "agency management", "partnerships", "campaigns", "activations",
    "go-to-market", "omnichannel", "retail marketing", "in-store activation", "events",
    "budget management", "A&P", "ROI", "ROAS", "GMV", "KPI", "analytics", "UAE", "Dubai", "GCC",
]

CONTENT = {
    "headline": (
        "Marketing Manager · Retail, Lifestyle & Fashion · Brand, Digital & Community · "
        "Customer Acquisition, Retention & Growth"
    ),
    "professional_summary": (
        "Commercially minded, creative marketer with 4+ years across retail, lifestyle, fashion and "
        "consumer brands — brand building, digital-first growth, community and customer engagement. "
        "Drives acquisition and retention through Shopify/CRO, paid social, influencers, partnerships "
        "and agencies; +30% GMV QoQ. Dubai-based with hands-on UAE retail and quick-commerce network."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Homegrown UAE consumer group | Brands: Befit, Eurocake, Flair | 50+ countries",
            "bullets": [
                "Lead brand, digital and customer engagement end-to-end — positioning, go-to-market and 6 launches — owning the A&P budget and the KPIs for awareness, traffic, acquisition and sell-out",
                "Built the influencer and community programme from zero: 25–50 creators per campaign plus sampling and seeding across retail and quick-commerce, driving UGC, brand awareness and measurable sales growth",
                "Run digital-first growth — own the Shopify store (merchandising, UX, CRO, AOV), Meta and Google Ads, social and EDM — testing creative and optimising ROI/ROAS, and manage agencies and retail partners (Noon, Talabat, Careem, Deliveroo) for traffic and conversion",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Miravia — Alibaba's lifestyle marketplace | Top-5 global e-commerce | 100K+ employees",
            "bullets": [
                "Grew 42 retail and fashion accounts +30% GMV QoQ through assortment, pricing, promotions and campaign calendars, analysing traffic, conversion, ROI/ROAS and retention to steer the plan",
                "Created and led the Beauty Club and Hot on Social community projects — content, creators and customer engagement that lifted loyalty, repeat purchase and the brand's lifestyle positioning",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Helped build the Retail vertical — onboarding fashion and lifestyle brands to the platform — and ran bespoke marketing activations with partners that grew order volume and GMV",
            ],
        },
        {
            "company": "Massimo Dutti (Inditex)",
            "role": "Sales Associate – Las Rozas Village Factory Store",
            "dates": "Jun 2018 – Jun 2019",
            "location": "Madrid, Spain",
            "context": "Inditex Group premium fashion retail | Shop-floor retail experience",
            "bullets": [
                "Front-line retail experience at Inditex — customer experience, visual merchandising standards and store cadence that still ground how campaigns land in-store",
            ],
        },
    ],
    "skills_brand": (  # label -> "Brand & Creative"
        "brand building & positioning, creative direction & content, campaigns & activations, "
        "go-to-market & launches, in-store & retail marketing, A&P budget management"
    ),
    "skills_ecommerce": (  # label -> "Digital & Growth"
        "digital-first marketing, Shopify & e-store, CRO & AOV, Meta & Google Ads, "
        "social (Instagram, TikTok), EDM & CRM, quick-commerce (Noon, Talabat, Careem, Deliveroo)"
    ),
    "skills_commercial": (  # label -> "Community & Partnerships"
        "influencer marketing (25–50 creators/campaign), UGC & seeding, community & loyalty programmes, "
        "agency management, brand partnerships, key account & distributor management, negotiation"
    ),
    "skills_data": (  # label -> "Customer & Analytics"
        "customer acquisition & retention, traffic & conversion analysis, ROI, ROAS, GMV, "
        "KPI dashboards & reporting, sell-in/sell-out, AI-assisted analysis"
    ),
    "skills_tools": (  # label -> "Tools"
        "Shopify, Meta Ads Manager, Google Ads, Google Analytics, Power BI, Tableau, Looker, Salesforce, "
        "Adobe (Photoshop, Illustrator), Canva, Generative AI (Claude, ChatGPT)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Brand & Creative",
    "E-Commerce & Digital": "Digital & Growth",
    "Commercial": "Community & Partnerships",
    "Data & Analytics": "Customer & Analytics",
}


def make_job() -> Job:
    return Job(
        id="confidential-careers-marketing-manager-dubai-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Dubai, United Arab Emirates",
        url="https://www.linkedin.com/jobs/",
        source="linkedin",
        description=JOB_DESCRIPTION,
        raw={
            "query": "Confidential Careers Marketing Manager Dubai retail outdoor adventure lifestyle",
            "note": "Confidential retail employer via LinkedIn (verified job, Easy Apply, 899 applicants, "
                    "64 in the last day — very high competition). Brand + digital + community + customer "
                    "engagement for a retail brand with an outdoor/adventure/active-lifestyle tilt. Asks 5–8 "
                    "yrs; Paula has 4+ (5+ counting the Inditex retail year) — under-band, do not inflate. "
                    "Outdoor/adventure affinity NOT claimed on the CV; use it in the cover letter only if "
                    "genuinely true. Employer unknown → no company-specific tailoring possible. Factual UAE "
                    "Residence Visa only.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": None, "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 78, "ai_tier": "Warm",
        "skills_match": [
            "Retail + lifestyle/fashion line: Miravia KAM (Beauty, Fragrances & Fashion), Glovo Retail vertical, Inditex/Massimo Dutti shop floor",
            "Digital-first growth: Shopify + CRO/AOV, Meta & Google Ads, social, EDM, UAE quick-commerce (Noon, Talabat, Careem, Deliveroo)",
            "Community & engagement: Beauty Club and Hot on Social projects (loyalty, creators, repeat purchase)",
            "Influencers, agencies & partnerships: programme built from zero, 25–50 creators/campaign, sampling & seeding, agency and partner management",
            "Acquisition & retention analytics: traffic, conversion, ROI/ROAS, retention; +30% GMV QoQ across 42 accounts",
            "UAE market knowledge and network; Dubai-based with residence visa",
        ],
        "missing_skills": [
            "Asks 5–8 yrs; Paula has 4+ (5+ counting the Inditex retail year) — under the band, not inflated on the CV",
            "No outdoor / adventure-sports / active-lifestyle category experience — the JD calls that affinity 'highly advantageous'; only mention it in the cover letter if genuinely true",
            "Confidential employer: no brand, no company research, no tailored deliverable possible",
            "899 applicants (64 in the last day) — very high competition for an Easy Apply post",
        ],
        "sector_fit": "good (retail / consumer / lifestyle / fashion — genuine overlap; the outdoor-sports niche is the gap)",
        "seniority_fit": "slightly under-band (asks 5–8 yrs; Paula 4+, 5+ including retail floor experience)",
        "red_flags": [
            "Confidential poster — cannot verify the employer, the brand or the package before applying",
            "Very high applicant volume; Easy Apply means low differentiation — worth a LinkedIn recruiter note if the poster can be identified",
            "No salary published — confirm against the AED 20K/month floor",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Solid, honest fit. The role is a generalist retail Marketing Manager covering brand, digital, "
            "community and customer engagement with acquisition/retention and sales-growth ownership — which "
            "is close to what Paula does today at DoFreeze and did at Miravia. Her retail/lifestyle/fashion "
            "line (Miravia KAM, Glovo's Retail vertical, Inditex shop floor), digital-first toolkit "
            "(Shopify/CRO, Meta & Google Ads, EDM, quick-commerce) and hands-on influencer/community/agency "
            "management map directly onto the ask, and she already has the UAE market knowledge the posting "
            "prefers. Two honest gaps: she is slightly under the 5–8 year band, and she has no outdoor / "
            "adventure-sports category background — the JD flags that affinity as highly advantageous, so it "
            "should only be raised if genuinely true. The employer is confidential, so no company-specific "
            "tailoring or deliverable is possible; the CV therefore leads on retail + digital growth + "
            "community. Factual UAE Residence Visa only."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
