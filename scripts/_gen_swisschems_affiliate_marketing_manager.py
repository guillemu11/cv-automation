"""One-off: generate Paula's ONE-PAGE CV for **Affiliate Marketing Manager** at **SwissChems**
(supplements / biohacking DTC eCommerce; remote, UAE-based talent accepted).

JD ask: recruit/onboard/manage affiliates, influencers, creators; scale affiliate programme across
TikTok, IG, YouTube, FB, X, Reddit, Telegram, Discord; performance tracking (sales, conversions, ROI);
campaigns/promos/launches; landing pages with marketing; compliance with brand guidelines. 3+ yrs
affiliate/partnerships, eCom/DTC; affiliate platforms (Impact, Everflow, CJ...); GA/Shopify analytics;
supplements/health/fitness preferred. Nice: existing network, Shopify/DTC, TikTok Shop, Meta Ads.

Paula's honest angle: built Befit's affiliate programme from scratch (commission on sales, AI-built
landing page, multilingual AI-avatar brief); 103 creator activations in 3 campaigns at AED 81/piece;
sales-driven creator campaigns measured with control group (SMASH x talabat +165% vs +4.7%);
Befit = fitness/better-for-you snacks (Fitness Month, New Year New Me); Shopify DTC store; Meta Ads.

Honesty guardrails:
- No affiliate network tool (Impact/Everflow/CJ...) claimed. No Reddit/Telegram/Discord/X/YouTube
  programme claimed. No TikTok Shop. No supplements claim — fitness snacks only.
- Forecast not owned. Factual "UAE Residence Visa" only.

ONE-PAGE standard. Lands under output/2026-09-13/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "SwissChems"
TITLE = "Affiliate Marketing Manager"
DATE_FOLDER = "2026-09-13"

JOB_DESCRIPTION = """\
Affiliate Marketing Manager — SwissChems (supplements / health DTC eCommerce). Remote, full-time;
preferably Turkey, USA, UAE or India-based talent.

Do: recruit, onboard and manage affiliates, influencers, creators and brand partners; develop and scale
affiliate programmes across TikTok, Instagram, YouTube, Facebook, X, Reddit, Telegram, Discord; build
long-term relationships with high-performing affiliates and creators; launch and optimise affiliate
campaigns, promotions and product launches; track and analyse affiliate performance, sales, conversions
and ROI; identify new partnership opportunities for acquisition and revenue; collaborate with marketing
on promotional assets, landing pages and campaign strategies; ensure affiliates comply with brand
guidelines and platform policies.

Bring: 3+ years in affiliate marketing or partnership management; proven growth of affiliate programmes
for eCommerce/DTC brands; affiliate platforms (Impact, Everflow, CJ, ShareASale, Refersion, PartnerStack,
Rakuten, Awin or similar); influencer, creator and performance marketing; Google Analytics / Shopify
Analytics; communication, negotiation, relationship management; supplements, health, fitness, biohacking
or wellness preferred; excellent English for US partners.
Nice: existing health/fitness affiliate & influencer network; Shopify and DTC; TikTok Shop, Meta Ads,
creator marketplaces. Application must include a portfolio of the affiliate/influencer network.
"""

ATS = [
    "affiliate marketing", "affiliate programme", "affiliates", "influencers", "creators",
    "creator partnerships", "brand partners", "recruit", "onboard", "performance marketing",
    "commission", "TikTok", "Instagram", "YouTube", "Facebook", "sales", "conversions", "ROI",
    "product launches", "promotions", "landing pages", "brand guidelines", "negotiation",
    "relationship management", "eCommerce", "DTC", "Shopify", "Shopify Analytics",
    "Google Analytics", "Meta Ads", "health", "fitness", "wellness", "remote",
]

CONTENT = {
    "headline": "Affiliate & Influencer Marketing · Creator Partnerships · DTC",
    "professional_summary": (
        "Performance-driven marketer with 5 years across eCommerce marketplaces (Alibaba, Glovo) and FMCG. In Dubai I built "
        "a fitness snack brand's affiliate and creator programme from zero, with every campaign measured on sales, not reach."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG group | Befit (fitness / better-for-you snacks), Eurocake, Smash | DTC Shopify + delivery apps",
            "bullets": [
                "Launched Befit's affiliate programme from scratch: commission on sales, direct creator deals, an AI-built landing page and an AI avatar explaining the content brief in several languages",
                "Recruited, briefed and managed 103 creator activations across 3 sales-driven campaigns at AED 81 per content piece, tied to fitness peaks (Fitness Month, New Year New Me) and product launches",
                "Proved ROI with a control group: SMASH x talabat creators lifted daily sales +165% vs +4.7% for the control brand; Befit x Noon added 4,176 incremental units (+31% over baseline)",
                "Run creator communities (running, padel, yoga clubs), 4 agencies (-30% negotiated, reports audited) and the Shopify store and Meta Ads; lead a team of two",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's eCommerce marketplace | 100K+ employees",
            "bullets": [
                "Grew 42 brand partners +30% GMV QoQ with joint promotions, tracking conversion, traffic and retention; onboarded 30+ fragrance brands in two months",
                "Created the Beauty Club and Hot on Social creator-led programmes; owned the Flash Sales channel, reporting to the CEO",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Delivery & quick-commerce app | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Negotiated deals and ran joint marketing activations with XL partners (KFC, Taco Bell, Sushi Shop); helped build the Retail vertical",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | Chocolate category (Milka, Suchard) | €36B revenue",
            "bullets": [
                "Built promotional-effectiveness and sell-in/sell-out reports with Nielsen for the chocolate category",
            ],
        },
    ],
    "skills_brand": (  # label -> "Affiliates"
        "programme design, commission models, recruiting & onboarding, briefs, brand guidelines"
    ),
    "skills_ecommerce": (  # label -> "Creators"
        "influencer & UGC campaigns, communities, gifting & sampling, product launches, agencies"
    ),
    "skills_commercial": (  # label -> "Partnerships"
        "negotiation, long-term partner relationships, platform co-marketing (talabat, Noon)"
    ),
    "skills_data": (  # label -> "Performance"
        "sales & conversion tracking, ROI / ROAS, control groups, cost per content, audits"
    ),
    "skills_tools": (  # label -> "Tools"
        "Shopify, Meta Ads, Google Ads, Instagram, TikTok, Looker, Excel, Canva, Claude (AI agents)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Affiliates",
    "E-Commerce & Digital": "Creators",
    "Commercial": "Partnerships",
    "Data & Analytics": "Performance",
}


def make_job() -> Job:
    return Job(
        id="swisschems-affiliate-marketing-manager-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Remote (UAE-based)",
        url="https://www.linkedin.com/jobs/search/?keywords=SwissChems%20Affiliate%20Marketing%20Manager",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "SwissChems Affiliate Marketing Manager",
            "note": "Good fit: affiliate programme built from scratch, 103 creator activations, sales-measured creator "
                    "campaigns, fitness snack brand, Shopify DTC. Gaps: affiliate network tools, Reddit/Telegram/Discord, supplements, US network.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": job.salary_raw,
        "salary_aed_min": job.salary_aed_min, "salary_aed_max": job.salary_aed_max,
        "posted_date": "2026-09-13", "raw": job.raw,
        "ai_score": 68, "ai_tier": "Hot",
        "skills_match": [
            "Built Befit's affiliate programme from scratch (commission, AI landing page, multilingual AI brief)",
            "103 creator activations across 3 sales-driven campaigns at AED 81/content piece",
            "Creator ROI proven with control groups (SMASH x talabat +165% vs +4.7%; Noon +4,176 units)",
            "Fitness / better-for-you snack brand; Shopify DTC store and Meta Ads",
        ],
        "missing_skills": [
            "Affiliate network platforms (Impact, Everflow, CJ, Refersion...)",
            "Reddit / Telegram / Discord / X affiliate channels",
            "Supplements industry and existing US health/fitness affiliate network",
        ],
        "sector_fit": "adjacent — fitness snacks DTC vs supplements/biohacking",
        "seniority_fit": "good — 5 yrs vs 3+ required",
        "red_flags": ["Portfolio of affiliate/influencer network required", "Remote, US-partner facing; salary not stated", "Supplements/peptides category — check brand reputation"],
        "ats_keywords": ATS,
        "reasoning": (
            "Rare direct evidence: she built an affiliate programme and a sales-measured creator engine for a fitness "
            "brand on Shopify. Gaps are affiliate SaaS platforms, niche channels (Reddit/Telegram/Discord) and a "
            "supplements network. UAE-based accepted; English-only role, no Arabic needed."
        ),
        "scored_by": "manual:claude", "freshness": "recent",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
