"""One-off: CV for Tashas Group "Brand & Marketing Manager – Elevated Concepts" (Avli by tashas, Bungalo34,
Galaxy Bar, Ouzeri), Dubai, across UAE, Saudi Arabia (Red Sea) and Bahrain (found via Jack & Jill, 2026-10-13;
careers.tashasgroup.com job 8083312).

Day-to-day custodian of a luxury hospitality portfolio under the Head of Marketing: marketing calendar,
integrated campaigns (digital, social, PR, CRM, partnerships, experiential), creative briefs and agency/supplier
workflow, events, activations, venue launches, partnerships, PR and influencer/gifting/tastemakers, paid media,
SEO and social, reporting; works with Operations, Culinary, Beverage, Design, Guest Experience. KPIs: new guests,
retention, efficiency, awareness/PR, social growth, events. 5–7 yrs; luxury hospitality/lifestyle preferred;
GCC (UAE, KSA, Bahrain).

Paula's honest angle: same title today at DoFreeze (7 F&B brands, team of two), influencer gifting/seeding from
zero, events (Saudi Food Show 2026, Sun & Sand Sports x Fabletics, sports-camp sampling, clubs), talabat
partnership giveaway, CRM (16K base segmented, replenishment flow), agencies and suppliers; restaurant-side
knowledge from Glovo (KFC, Taco Bell, Vicio; kitchen visits); premium retail at Massimo Dutti; Miravia loyalty.

Honesty guardrails: no PR-agency, media-relations, venue-launch or luxury-hospitality claims (gaps). Glovo is NOT
the Retail vertical. Miravia accounts were beauty and fragrance. Massimo Dutti / Mondelez dates from her original
CV (to confirm on LinkedIn). No Arabic.

ONE-PAGE standard. Lands under output/2026-10-13/Tashas Group - Brand & Marketing Manager/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Tashas Group"
TITLE = "Brand & Marketing Manager"
DATE_FOLDER = "2026-10-13"

JOB_DESCRIPTION = """\
Brand & Marketing Manager – Elevated Concepts, Tashas Group, Dubai. Portfolio: Avli by tashas, Bungalo34, Galaxy
Bar, Ouzeri across the UAE, Saudi Arabia (Red Sea) and Bahrain, plus future expansion. Reports to the Head of
Marketing. Execute annual marketing plans; manage the marketing calendar; integrated campaigns across digital,
social media, PR, CRM, partnerships and experiential; brand standards and tone of voice; competitor and guest
insights. Coordinate content and campaign assets; creative briefs; workflow with agencies and suppliers; CRM and
email. Lead events, activations, venue launches and partnerships. Manage PR agencies and media partners; media
visits, tastemaker experiences, gifting and influencer collaborations. Coordinate paid media, website content,
SEO and social; reporting and optimisation. Manage agencies, suppliers, partners within budget; work with
Operations, Culinary, Beverage, Design and Guest Experience. KPIs: new customer acquisition, retention and
engagement, marketing efficiency, brand awareness and PR coverage, social growth, event success.
5–7 years in marketing, brand management or communications; luxury hospitality, lifestyle, retail or experiential
brands preferred; agencies and suppliers; digital, social and CRM; GCC markets (UAE, KSA, Bahrain).
"""

ATS = [
    "brand management", "marketing calendar", "integrated campaigns", "creative briefs", "content",
    "events", "activations", "partnerships", "influencer", "gifting", "tastemakers", "PR", "CRM",
    "email marketing", "social media", "paid media", "SEO", "agency management", "suppliers",
    "guest acquisition", "retention", "loyalty", "hospitality", "F&B", "GCC", "UAE", "Saudi Arabia",
]

CV_CONTENT = {
    "headline": "Brand & Marketing · F&B · Events, Influencer & CRM",
    "professional_summary": (
        "Brand & Marketing Manager with 5 years across F&B, beauty and e-commerce. At DoFreeze in Dubai, runs "
        "the marketing calendar for 7 F&B brands with a team of two: campaigns, influencer gifting, events and "
        "partnerships, CRM and paid social. Previously grew 42 beauty and fragrance brands +30% GMV QoQ at "
        "Alibaba's Miravia and co-managed KFC, Taco Bell and Vicio at Glovo."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE F&B group | 7 brands incl. Befit, Eurocake, SMASH | sold across the GCC and 50+ countries",
            "bullets": [
                "Run the marketing calendar for 7 F&B brands with a team of two (designer, social media exec): seasonal campaigns (Ramadan, New Year, back to school), 6 launches, creative briefs and sign-off on all content",
                "Built influencer gifting and seeding from zero: 103 creator activations over 3 waves at AED 81 per content piece (market AED 1,000–3,500); negotiated −30% and audited the agency's reporting",
                "Events and activations: Saudi Food Show 2026 (Jeddah), a Sun & Sand Sports x Fabletics yoga session that led to a listing proposal, kids sports-camp sampling with QR-tracked codes, running and padel clubs",
                "Partnerships measured on sales: a purchase-to-enter talabat giveaway lifted SMASH daily sales +165% vs +4.7% for a control brand; Befit x Noon 'New Year, New Me' added 4,176 incremental units",
                "CRM and digital: segmented a 16K email base, added a 30-day replenishment flow and reviews (web orders +86%); Meta Advantage+ retargeting at 3.7x ROAS in week one",
                "Manage agencies and suppliers (alist, Yamammi, BixGrow, Lovable for the website) within budget; built a Claude-based system automating planning, content and reporting (~40% less manual work)",
            ],
        },
        {
            "company": "Miravia & AliExpress (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Marketplace | Beauty & Health department",
            "bullets": [
                "Grew 42 beauty and fragrance brands (incl. L'Oréal Paris, Maybelline, NYX, Redken, KIKO Milano) +30% GMV QoQ through pricing, promotions and on-site visibility; worked with marketing, product and ops",
                "Created the Beauty Club loyalty programme and Hot on Social, a monthly trend showcase; grew Arabian perfume into a top-growing sub-category; owned Flash Sales, reporting to the CEO",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Food delivery app | 10K+ employees",
            "bullets": [
                "Co-managed 5 XL restaurant accounts (KFC, Taco Bell, Vicio): promotions, in-app visibility; kitchen visits to cut prep times",
            ],
        },
        {
            "company": "Mondelez International · Massimo Dutti (Inditex)",
            "role": "Early Career",
            "dates": "2018 – 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | Premium fashion retail",
            "bullets": [
                "Category Planning Trainee, Mondelez (Aug 2021 – Aug 2022): promo-effectiveness analysis and NPD support",
                "Sales Associate, Massimo Dutti (Jun 2018 – Jun 2019): premium customer experience on the Las Rozas Village store floor",
            ],
        },
    ],
    "skills_brand": (
        "marketing calendar, integrated campaigns, creative briefs, content, launches, brand standards"
    ),
    "skills_ecommerce": (
        "Meta Ads, Google Ads, Instagram, TikTok, email & segmentation, Shopify, website content"
    ),
    "skills_commercial": (
        "events & activations, partnerships, influencer & gifting, sampling, agency & supplier management"
    ),
    "skills_data": (
        "campaign reporting, ROI, ROAS, incremental uplift vs control, Excel (Expert), Looker, Power BI"
    ),
    "skills_tools": (
        "Claude (AI automation), ChatGPT, Canva, Photoshop, Illustrator, Notion"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Brand & Campaigns",
    "E-Commerce & Digital": "Digital & CRM",
    "Commercial": "Events & Partners",
    "Tools": "Tools & AI",
}


def make_job() -> Job:
    return Job(
        id="tashas-brand-marketing-manager-elevated-concepts-2026-10",
        title=TITLE,
        company=COMPANY,
        location="Dubai, United Arab Emirates",
        url="https://careers.tashasgroup.com/jobs/8083312-brand-marketing-manager",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={"via": "Jack & Jill AI (tracked job, 2026-10-13)",
             "portfolio": "Avli by tashas, Bungalo34, Galaxy Bar, Ouzeri (Elevated Concepts)",
             "markets": "UAE, Saudi Arabia (Red Sea), Bahrain"},
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not posted (Jack & Jill estimate AED 180–260K/yr ≈ 15–21.7K/month)",
        "salary_aed_min": 15000, "salary_aed_max": 21700,
        "posted_date": None, "raw": job.raw,
        "ai_score": 72, "ai_tier": "Warm",
        "skills_match": [
            "Same title today (Brand & Marketing Manager, DoFreeze F&B, team of two)",
            "Influencer gifting/seeding from zero (103 activations, AED 81 per piece), agencies and suppliers",
            "Events and partnerships: Saudi Food Show 2026, Sun & Sand x Fabletics, talabat giveaway (+165% vs control)",
            "CRM: 16K email base segmented, replenishment flow; restaurant-side knowledge from Glovo",
        ],
        "missing_skills": [
            "Luxury hospitality / venue launches",
            "PR agency and media-relations management",
        ],
        "sector_fit": "good (F&B and lifestyle; not hospitality venues)",
        "seniority_fit": "good (5–7 yrs asked; Paula 5, same title)",
        "red_flags": [
            "Salary estimate AED 15–21.7K/month, around her current 20K",
        ],
        "ats_keywords": ATS,
        "reasoning": "Same title and the same events/influencer/partnerships/CRM mix in F&B; gaps are luxury hospitality and PR.",
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _pages(pdf: Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(pdf)).pages)


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    final_dir = settings.output_dir / DATE_FOLDER / f"{COMPANY} - {TITLE}"
    dest = final_dir / "01_CV_y_Carta"
    dest.mkdir(parents=True, exist_ok=True)

    cv_docx = cv._fill_template(CV_CONTENT, job)
    _relabel_for_role(cv_docx)
    stray = cv_docx.parent.parent

    clean_docx = dest / "Paula De Francisco - CV.docx"
    shutil.move(str(cv_docx), str(clean_docx))
    # The template writes to output/<Company> - <Role>/ at the root; only ever remove that.
    if stray.exists() and stray.parent.resolve() == settings.output_dir.resolve():
        shutil.rmtree(stray, ignore_errors=True)

    cv_pdf = _to_pdf_soffice(clean_docx)
    # The dashboard only indexes CV_Paula_<Company>_<Title>.pdf; the clean name is the one to send.
    dashboard_pdf = dest / f"CV_Paula_{cv._sanitize_filename(COMPANY)}_{cv._sanitize_filename(TITLE)}.pdf"
    shutil.copy(str(cv_pdf), str(dashboard_pdf))
    print("OK", cv_pdf, "PAGES", _pages(cv_pdf), "| dashboard copy:", dashboard_pdf.name)


if __name__ == "__main__":
    main()
