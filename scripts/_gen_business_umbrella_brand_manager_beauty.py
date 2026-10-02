"""One-off: CV de una página para **Brand Manager** publicado por **Business Umbrella** (headhunter de
Dubái; el CEO, Mohummed Yasir Charakla, es el anunciante) para un cliente confidencial de **belleza**:
haircare, skincare y otros productos. EAU, en remoto, jornada completa, reporta al Brand Director.
LinkedIn "Promocionado por técnico de selección", Solicitud sencilla, 24 solicitudes en 36 minutos.

Pide el JD: apoyar al Brand Director en la estrategia y ejecución de marca; gestión diaria de marca,
identidad y guías visuales en todos los puntos de contacto; calendario de marca (campañas, lanzamientos,
activaciones, promociones); estudio de mercado y competencia; campañas en digital, social, retail, PR,
eventos e influencers; posicionamiento de producto y lanzamientos de la planificación a la ejecución;
apoyo a distribución y retail (trade marketing, POS, sell-through); creación de contenido con equipos
creativos y agencias, revisión y aprobación; comunicación de marca y PR; trabajo con ventas, educación,
operaciones, finanzas y desarrollo de producto; relación con consumidores, retailers, distribuidores e
influencers; KPIs y reporting al Brand Director; apoyo en presupuesto (seguimiento de gasto, POs,
facturas, ROI); formación de producto a equipos de ventas y educación.
Requisitos: grado en Marketing/Empresa; experiencia en brand management, preferiblemente belleza,
cosmética, haircare, skincare o FMCG; digital, social, influencer y contenido; analítica; conocer el
mercado de belleza es un plus. **No pide árabe** en el JD visible. Salario no publicado.

Ángulo honesto de Paula:
- Belleza real del lado comercial: en Miravia (Alibaba) llevó 42 marcas de belleza, fragancia y moda,
  incluida la cuenta de maquillaje y skincare de KIKO Milano; +30% GMV QoQ (cifra de toda la cartera,
  no de KIKO); creó Beauty Club y Hot on Social; como PIC Fragrances incorporó 30+ casas en dos meses
  (Arabian Oud, Lattafa, Swiss Arabian, Ajmal), que en un cliente de belleza en EAU pesan.
- El oficio de brand manager es el de hoy en DoFreeze (FMCG): calendario de marca por temporadas,
  6 lanzamientos NPD end-to-end, equipo de dos (diseñador + social media executive) y aprobación de
  todo el contenido.
- Influencer con venta medida: 4 agencias, 103 activaciones de creator en 3 campañas; Befit × Noon
  +31% sobre baseline (+4.176 uds incrementales); SMASH × talabat +165% de venta diaria contra +4,7%
  del grupo de control. Presupuesto: negoció -30% con agencias, AED 81 por pieza de contenido.
- Grado en Administración de Empresas (CUNEF): cubre el requisito.

Guardarraíles de honestidad:
- **Haircare**: no lo ha trabajado. No se menciona.
- **Brand-side en belleza**: no lo ha tenido. Su belleza es comercial/marketplace (Miravia); el oficio
  de marca viene de FMCG alimentación. No se insinúan años de brand management en belleza.
- **Canal profesional / salones / equipo de educación**: sin evidencia. No se menciona.
- **POS e in-store**: no hay prueba sólida de que lo diseñe; se habla de promociones con ventas y
  distribuidores en modern trade y quick-commerce, como en el CV de CPW.
- **Forecast**: la dueña es la e-commerce manager; no aparece.
- **Deliveroo**: DoFreeze NO está en Deliveroo. Plataformas: talabat, Noon, Careem.
- Sin árabe. Visa solo "UAE Residence Visa"; nunca "no sponsorship".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Business Umbrella"
TITLE = "Brand Manager"
DATE_FOLDER = "2026-10-01"

JOB_DESCRIPTION = """\
Brand Manager. Business Umbrella (recruiter) for a confidential beauty client. UAE, remote, full time.
Reports to: Brand Director.
Overview: The Brand Manager will support the Brand Director in managing and developing the company's portfolio of
beauty products, including haircare, skincare, and other beauty products. Responsible for the day-to-day management
and execution of brand strategies, marketing initiatives, product positioning, brand communications, and commercial
activities, with the objective of driving brand growth, awareness, and profitability. Works closely with internal
teams, distributors, retailers, suppliers, agencies, influencers and other external partners to ensure consistent
brand execution across all channels; monitors market trends, coordinates product launches, supports sales
activities, manages content and tracks brand performance.
Key responsibilities: 1. Brand management: support brand strategy for haircare, skincare and other beauty products;
manage day-to-day brand activities; ensure consistent brand identity, guidelines, messaging and visual standards;
maintain the brand calendar (campaigns, launches, activations, promotions). 2. Market research & competitive
analysis: industry trends, consumer needs, competitor activity, pricing, launches. 3. Marketing campaigns across
digital, social media, retail, PR, events and influencer partnerships; coordinate timelines and deliverables;
monitor performance and optimize. 4. Product positioning & launches: positioning strategies, value propositions,
launch plans from planning through execution with sales, marketing and operations; post-launch performance.
5. Distribution & retail support: product distribution and visibility across key retail and professional channels;
promotional and trade marketing initiatives; brand standards in retail, displays and POS; sell-through.
6. Content creation & brand materials with creative teams, agencies and suppliers: product descriptions, campaign
visuals, social content, presentations, brochures, POS; review and approve content. 7. Brand communication across
digital, social, packaging, advertising, events and retail; support PR, influencer, media and partnerships.
8. Cross-functional collaboration with sales, marketing, education, operations, finance and product development;
coordinate agencies, suppliers and distributors. 9. Consumer engagement & relationship building with consumers,
retailers, distributors, influencers and educators; monitor reviews and sentiment. 10. Performance tracking &
reporting: KPIs on brand, sales, campaigns, launches and engagement; regular reports for the Brand Director.
11. Budget & marketing investment: support budgets, track expenses, purchase orders and invoices, evaluate ROI.
12. Brand training & product knowledge for education and sales teams.
Skills & qualifications: Bachelor's degree in Marketing, Business, Communications or related; experience in brand
management or marketing, preferably within beauty, cosmetics, haircare, skincare or FMCG; brand strategy, consumer
behavior and commercial activities; project management; communication; cross-functional collaboration;
digital marketing, social media, influencer marketing and content development; analytical mindset; creative
thinking; presentation and reporting skills; knowledge of the beauty market is an advantage.
"""

ATS = [
    "brand manager", "brand management", "brand strategy", "beauty", "cosmetics", "skincare", "haircare",
    "fragrances", "FMCG", "brand portfolio", "brand growth", "brand awareness", "profitability",
    "brand identity", "brand guidelines", "visual standards", "brand calendar", "campaigns", "launches",
    "activations", "promotions", "market research", "competitive analysis", "consumer insights", "trends",
    "marketing campaigns", "digital marketing", "social media", "PR", "events", "influencer marketing",
    "influencer partnerships", "product positioning", "value proposition", "product launches", "launch plans",
    "NPD", "distribution", "retail", "trade marketing", "sell-through", "distributors", "retailers",
    "content creation", "content development", "creative briefs", "campaign visuals", "agencies", "suppliers",
    "brand communication", "cross-functional", "sales", "consumer engagement", "KPIs", "performance tracking",
    "reporting", "budget", "marketing investment", "ROI", "presentations", "project management",
]

CONTENT = {
    "headline": "Brand Manager · Beauty & FMCG · Launches, Campaigns, Influencer & Content · UAE",
    "professional_summary": (
        "Brand manager in Dubai running brand strategy, campaigns and launches for an FMCG group sold in 50+ "
        "markets, with a team of two. Beauty grounding from Alibaba's Miravia: 42 beauty, fragrance and fashion "
        "brands, incl. KIKO Milano makeup & skincare. Every campaign measured on sell-out."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG group | Befit, Eurocake, Flair | modern trade, talabat, Noon, Careem, Shopify | 50+ markets",
            "bullets": [
                "Own the brand calendar for three brands — seasonal campaigns (Ramadan, back to school, Fitness Month, New Year), activations and promotions — keeping identity, messaging and visual standards consistent across retail, e-commerce and social",
                "Lead 6 product launches end-to-end (brief and positioning, packaging, pricing, go-to-market) and a team of two — designer and social media executive — briefing and approving campaign visuals, social content and product copy",
                "Run influencer campaigns with four agencies — 103 creator activations over three campaigns: Befit × Noon +31% over baseline (+4,176 incremental units); SMASH × talabat lifted daily sales +165% vs +4.7% for the control brand",
                "Track spend and ROI per campaign — negotiated agency fees down 30%, AED 81 per content piece — audit agency reach reports, and plan trade promotions with sales and distributors across modern trade and quick-commerce",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees",
            "bullets": [
                "Managed 42 beauty, fragrance and fashion brands — incl. KIKO Milano makeup & skincare — on assortment, pricing and promotional calendars, delivering +30% GMV QoQ across the portfolio",
                "Created the Beauty Club and Hot on Social programmes to lift brand visibility; as PIC Fragrances, onboarded 30+ houses in two months (Arabian Oud, Lattafa, Swiss Arabian, Ajmal)",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Managed XL accounts (KFC, Taco Bell, La Tagliatella) on GMV and joint marketing activations, coordinating marketing, logistics and customer support to deliver campaigns",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | chocolate category (Milka, Suchard)",
            "bullets": [
                "Ran sell-in/sell-out and promotional-effectiveness analysis with Nielsen and supported NPD launches (Milka Spread, Mini Suchard)",
            ],
        },
    ],
    "skills_brand": (  # -> "Brand & Launches"
        "brand strategy & positioning, brand calendar, product launches (6 NPD), integrated campaigns, "
        "brand guidelines, team leadership (2 reports)"
    ),
    "skills_ecommerce": (  # -> "Digital & Content"
        "social media, influencer & creator marketing, content & creative briefs, agency management, "
        "Meta & Google Ads, Shopify"
    ),
    "skills_commercial": (  # -> "Retail & Trade"
        "trade promotions, distributors, modern trade, quick-commerce (talabat, Noon, Careem), "
        "key account management, negotiation"
    ),
    "skills_data": (  # -> "Insights & Reporting"
        "market & competitor research, campaign KPIs, ROI & ROAS, budget tracking, sell-in/sell-out, Nielsen"
    ),
    "skills_tools": (  # -> "Tools"
        "Excel (Advanced), PowerPoint (Advanced), Power BI, Nielsen, SAP, Salesforce, Meta Ads, Google Ads, "
        "Canva & Adobe, Claude (AI automation)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Brand & Launches",
    "E-Commerce & Digital": "Digital & Content",
    "Commercial": "Retail & Trade",
    "Data & Analytics": "Insights & Reporting",
}


def make_job() -> Job:
    return Job(
        id="business-umbrella-brand-manager-beauty-2026-10",
        title=TITLE,
        company=COMPANY,
        location="United Arab Emirates (Remote)",
        url="https://www.linkedin.com/jobs/search/?keywords=Business%20Umbrella%20Brand%20Manager",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Business Umbrella Brand Manager beauty UAE",
            "note": "Headhunter (Business Umbrella) para un cliente confidencial de belleza: haircare, skincare y "
                    "otros. Remoto en EAU, reporta al Brand Director. Anunciante: Mohummed Yasir Charakla (CEO de "
                    "Business Umbrella), contacto de 2º grado. Solicitud sencilla; 24 solicitudes en 36 minutos. "
                    "No pide árabe en el JD. Encaje: belleza comercial en Miravia (KIKO, fragancias árabes) + "
                    "oficio de marca en DoFreeze. Gaps: haircare, canal profesional/educación, brand-side en belleza.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 80, "ai_tier": "Hot",
        "skills_match": [
            "Belleza real: 42 marcas de belleza, fragancia y moda en Miravia, incluida KIKO Milano (maquillaje y skincare)",
            "Fragancias árabes: incorporó Arabian Oud, Lattafa, Swiss Arabian y Ajmal como PIC Fragrances",
            "Calendario de marca por temporadas (Ramadán, back to school, Fitness Month, New Year) en DoFreeze",
            "6 lanzamientos NPD end-to-end, que es el bloque de posicionamiento y lanzamientos del JD",
            "Influencer con venta medida: Noon +31% sobre baseline, talabat +165% frente a +4,7% del control",
            "Contenido y agencias: dirige a diseñador y social media executive, 4 agencias, AED 81 por pieza",
            "Presupuesto y ROI: negoció -30% con agencias y auditó su informe de reach",
            "Grado en Administración de Empresas (CUNEF) y sin requisito de árabe en el JD",
        ],
        "missing_skills": [
            "Haircare: no lo ha trabajado",
            "Brand management en belleza desde la marca: su belleza es del lado marketplace (Miravia)",
            "Canal profesional (salones) y formación a equipos de educación: sin experiencia",
            "Diseño de POS y visual merchandising en tienda: poca evidencia",
        ],
        "sector_fit": "bueno — belleza (skincare, fragancias) por Miravia y FMCG por DoFreeze",
        "seniority_fit": "en línea — Brand Manager que reporta a un Brand Director, frente a su Brand & Marketing Manager actual",
        "red_flags": [
            "Headhunter con cliente confidencial: no se sabe la marca ni el salario",
            "Dice remoto pero pide asistir a lanzamientos, eventos y activaciones: confirmar modalidad real",
            "Volumen rápido (24 en 36 minutos): conviene escribir al anunciante además de solicitar",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Encaje bueno y poco frecuente: es un Brand Manager de belleza sin requisito de árabe. Paula trae "
            "belleza de verdad del lado comercial (42 marcas en Miravia, KIKO, las fragancias árabes que más "
            "venden en EAU) y el oficio completo de brand manager de DoFreeze: calendario, lanzamientos, "
            "contenido, agencias e influencers con venta medida. Le falta haircare, el canal profesional con "
            "equipo de educación y haber llevado una marca de belleza desde dentro. El riesgo principal es que "
            "es un headhunter con cliente confidencial y no se sabe salario ni marca."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
