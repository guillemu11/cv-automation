"""One-off: CV de una página para **Influencer, Creator & Affiliate Manager** en **rite.**
(marca de wellness/suplementos funcionales nacida en EAU: gummies, hair, hydration, gut health,
protein, mushroom coffee). Dubái, presencial, jornada completa. LinkedIn Easy Apply, +100
solicitudes en 5 días (462 candidatos totales, 66 en un día).

Pide el JD: ownership de la estrategia de influencer, creator, embajadores y afiliación en EAU y
Arabia Saudí; sourcing y relación con creators de wellness, fitness, beauty, hair y lifestyle;
programa de seeding consistente; construir y escalar el programa de afiliación (reclutamiento,
comisiones, incentivos, performance); campañas de principio a fin — outreach, negociación,
activación y reporting; convertir creators top en embajadores y afiliados; pipeline continuo de
UGC para orgánico y paid; trabajar con Creative y Paid Media para escalar el contenido ganador;
medir revenue, CAC, ventas de afiliación y retorno del gasto en creators; gestión de presupuesto.
Perfil ideal: 3+ años en influencer/creator/afiliación/social, en DTC, wellness, beauty, fitness o
consumo; conocimiento del ecosistema creator de EAU y KSA; Instagram y TikTok; mentalidad
comercial y de datos; negociador; hands-on.

Ángulo de Paula — esta es LA vacante para su munición de influencer (todo real, DoFreeze):
- Construyó el programa de influencer desde cero (antes no existía nada).
- 103 activaciones de creador en 3 campañas (22 / 33 / 48), a AED 243 por creador y AED 81 por
  pieza de contenido frente a tarifas de mercado de 1.000-3.500.
- SMASH x talabat: venta diaria 355 -> 940 AED (+165%), con Befit como grupo de control (+4,7%).
- Befit x Noon "New Year, New Me": +4.176 unidades incrementales, +31% sobre baseline.
- Programa de afiliación creado por ella: comisión por ventas, trato directo sin agencia, landing
  construida con agentes de IA y avatar de IA que explica el brief en varios idiomas.
- Seeding y sampling estacional con talabat; sorteo con talabat (compra = participación).
- 4 agencias gestionadas (alist, Yamammi, Amplify, Terrier), -30% negociado y sobre-entrega.
- Auditó el informe de agencia: reach estimado desde seguidores, doble conteo de unique accounts;
  micro (<20K) rindiendo 10-40x la media.
- Shopify DTC propio + Meta/Google Ads -> encaje con el lado DTC y paid del rol.

Guardarraíles de honestidad:
- **KSA**: NO tiene ecosistema creator saudí trabajado de primera mano. DoFreeze exporta a KSA y
  ha corrido campañas en GCC, pero el CV no dice "KSA creator ecosystem". Es la brecha real.
- **Wellness/suplementos**: Befit es fitness/health snacking, no suplementos. Se dice "fitness and
  health-focused F&B", no "wellness supplements".
- **3+ años de influencer dedicado**: lo suyo es ~1 año de influencer intensivo dentro de un rol
  de brand manager, sobre una base comercial de 5 años. No se infla ninguna fecha ni se renombra
  ningún puesto.
- Sin árabe. Visa solo "UAE Residence Visa"; nunca "no sponsorship needed".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "rite."
TITLE = "Influencer, Creator & Affiliate Manager"
DATE_FOLDER = "2026-09-23"

JOB_DESCRIPTION = """\
Influencer, Creator & Affiliate Manager — rite. Dubai, UAE. On-site, full time.
Rite. is a fast-growing UAE-born wellness brand creating functional supplements — Hair, Hydration,
Gut Health, Protein and Mushroom Coffee — with a strong DTC business, growing retail presence and
expansion into Saudi Arabia. The brand is making creators, influencers and affiliates a core growth
channel.
Key responsibilities: own the influencer, creator, ambassador and affiliate strategy; source and
build relationships with relevant creators across wellness, fitness, beauty, hair and lifestyle;
develop a consistent and strategic product seeding programme; build and scale the affiliate
programme including recruitment, commissions, incentives and performance; manage influencer
campaigns from outreach and negotiation to activation and reporting; turn high-performing creators
into long-term ambassadors and affiliates; build a continuous pipeline of UGC and creator content
for organic and paid media; work closely with Creative and Paid Media teams to identify and scale
winning creator content; track revenue, CAC, affiliate sales and return on creator spend; manage
budgets and continuously optimise partnerships based on performance.
Ideal background: 3+ years in influencer, creator, affiliate or social marketing; experience within
DTC, wellness, beauty, fitness or consumer brands; strong knowledge of the UAE & KSA creator
ecosystem; strong understanding of Instagram, TikTok and creator-led content; commercial and
data-driven mindset comfortable connecting creator activity to sales and ROI; strong negotiator and
relationship builder; hands-on, proactive and highly organised.
"""

ATS = [
    "influencer marketing", "creator marketing", "affiliate marketing", "affiliate programme",
    "ambassadors", "product seeding", "gifting", "UGC", "user-generated content", "creator content",
    "outreach", "negotiation", "activation", "reporting", "campaign management", "commissions",
    "incentives", "recruitment", "DTC", "D2C", "wellness", "fitness", "beauty", "hair", "lifestyle",
    "consumer brands", "UAE", "KSA", "Saudi Arabia", "GCC", "MENA", "Instagram", "TikTok",
    "social media", "paid media", "Meta Ads", "Google Ads", "creative", "revenue", "CAC",
    "return on creator spend", "ROI", "ROAS", "sell-out", "incremental sales", "budget management",
    "performance marketing", "Shopify", "e-commerce", "quick-commerce", "talabat", "Noon", "Careem",
    "micro-influencers", "cost per content", "agency management", "relationship building", "data-driven",
]

CONTENT = {
    "headline": "Influencer, Creator & Affiliate Marketing · DTC & Quick-Commerce · UAE",
    "professional_summary": (
        "Built a creator programme from zero in the UAE: 103 creator activations at AED 81 per content "
        "piece, an in-house affiliate programme run direct with no agency, and campaigns measured on "
        "sell-out — +165% daily revenue against a control brand. Owns DTC Shopify and paid media."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE consumer brands — Befit (fitness & health snacking), SMASH, Eurocake | DTC + quick-commerce",
            "bullets": [
                "Built the influencer and creator programme from zero: 103 creator activations across 3 campaigns at AED 243 per creator and AED 81 per content piece, vs market rates of AED 1,000–3,500 — sourcing, briefing, negotiating and activating creators across fitness, food and lifestyle",
                "Launched the affiliate programme end-to-end: sales-based commissions, direct creator relationships with no agency margin, and an AI-built landing page with a multilingual AI avatar briefing content",
                "Ran campaigns to sell-out, not awareness: SMASH × talabat lifted daily revenue AED 355 → 940 (+165%) vs a control brand (+4.7%); Befit × Noon delivered +4,176 incremental units, +31% over baseline",
                "Managed 4 creator agencies, negotiating −30% on rate and over-delivering 103 creators vs 75 contracted; audited their reach reporting and shifted spend to micro creators outperforming campaign average 10–40×",
                "Built a seasonal seeding calendar with talabat, turned top creators into repeat ambassadors, and fed winning content into Meta and Google Ads alongside the Shopify DTC store; lead a designer and a social media executive",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees",
            "bullets": [
                "Created and led 'Beauty Club' and 'Hot on Social', social-led programmes building visibility and loyalty for beauty brands; managed 42 beauty, fragrance and fashion accounts, +30% GMV QoQ, negotiating joint investment",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue",
            "bullets": [
                "Managed XL accounts with bespoke marketing activations and promotional mechanics, negotiating joint investment and driving GMV",
            ],
        },
        {
            "company": "Mondelez International · Massimo Dutti (Inditex)",
            "role": "Trainee, Category Planning · Sales Associate",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | Premium fashion retail",
            "bullets": [
                "Sell-in/sell-out and promotional-effectiveness analysis supporting NPD launches; retail floor grounding in premium consumer brands",
            ],
        },
    ],
    "skills_brand": (  # -> "Influencer & Creator"
        "creator sourcing & outreach, negotiation, briefing, seeding & gifting, ambassadors, UGC pipeline, agency management, community partnerships"
    ),
    "skills_ecommerce": (  # -> "Affiliate & DTC"
        "affiliate programme build, commissions & incentives, Shopify DTC, CRO, quick-commerce (talabat, Noon, Careem), Instagram, TikTok"
    ),
    "skills_commercial": (  # -> "Performance & Budgets"
        "creator budget management, cost per content, return on creator spend, incremental sell-out, control-group testing, Meta Ads, Google Ads"
    ),
    "skills_data": (  # -> "Data & Insights"
        "revenue & ROI tracking, ROAS, reach & engagement auditing, micro vs macro performance, campaign reporting, AI-automated analysis"
    ),
    "skills_tools": (  # -> "Tools"
        "Shopify, Meta Business Suite, Meta Ads Manager, Google Ads, Canva, Adobe CS, Notion, Excel (Advanced), Claude (AI agents)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Influencer & Creator",
    "E-Commerce & Digital": "Affiliate & DTC",
    "Commercial": "Performance & Budgets",
    "Data & Analytics": "Data & Insights",
}


def make_job() -> Job:
    return Job(
        id="rite-influencer-creator-affiliate-manager-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE (On-site)",
        url="https://www.linkedin.com/jobs/search/?keywords=rite%20Influencer%20Creator%20Affiliate%20Manager%20Dubai",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "rite. Influencer, Creator & Affiliate Manager Dubai",
            "note": "Marca de wellness nacida en EAU (gummies y bebidas funcionales), 11-50 empleados, "
                    "DTC fuerte + retail creciente + expansión a KSA. El rol reporta de cerca a los "
                    "fundadores. LinkedIn marca a Paula como 'candidata destacada' y encaja con los "
                    "requisitos indispensables. Es la mejor coincidencia vista para su munición real de "
                    "influencer: programa desde cero, afiliación propia, sell-out con grupo de control y "
                    "coste por pieza de AED 81. Brecha real: ecosistema creator de Arabia Saudí y "
                    "suplementos como categoría. Competencia muy alta (462 solicitudes).",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 88, "ai_tier": "Hot",
        "skills_match": [
            "Programa de influencer construido desde cero en EAU — exactamente el encargo del rol",
            "Programa de afiliación creado por ella: comisiones por venta, trato directo y landing con IA",
            "103 activaciones de creador en 3 campañas, AED 243/creator y AED 81/pieza vs 1.000-3.500 de mercado",
            "Campañas medidas en venta con grupo de control: SMASH x talabat +165% diario frente a +4,7% del control",
            "Seeding y sampling estacional con talabat + sorteo de compra-participa",
            "4 agencias gestionadas, -30% negociado y auditoría del reporting de reach",
            "Shopify DTC propio, Meta Ads y Google Ads — el puente creator -> paid que pide el JD",
            "Micro-creators identificados como los de mejor rendimiento (10-40x la media)",
        ],
        "missing_skills": [
            "Ecosistema creator de Arabia Saudí (KSA) — no lo ha trabajado de primera mano",
            "Suplementos / wellness como categoría (lo más cercano es Befit, fitness & health snacking)",
            "3+ años de influencer DEDICADO: ~1 año intensivo dentro de un rol de brand manager",
            "Sin árabe",
        ],
        "sector_fit": "alto — consumo DTC con quick-commerce, muy cercano a su día a día",
        "seniority_fit": "bueno — Manager individual contributor, con fundadores cerca; encaja con su nivel actual",
        "red_flags": [
            "462 solicitudes, 66 en un solo día — competencia muy alta",
            "Empresa de 11-50 empleados: recursos y presupuesto por confirmar, y el suelo de 20K AED puede apretar",
            "Presencial, no híbrido",
            "Expansión a KSA como parte central del rol y ella no tiene ese ecosistema",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Es el encaje más directo que ha aparecido para lo que Paula realmente ha construido: un "
            "programa de creators desde cero, un programa de afiliación propio y campañas medidas en "
            "venta incremental con grupo de control, no en alcance. El JD pide literalmente seeding, "
            "afiliación, embajadores, pipeline de UGC para orgánico y paid, y conectar actividad de "
            "creator con revenue y CAC — y ella tiene cifra defendible para cada uno de esos puntos, "
            "incluida la auditoría del reporting de la agencia. Las dos brechas son KSA y la categoría "
            "de suplementos; ninguna de las dos se disimula en el CV. Con 462 solicitudes, la vía "
            "realista no es solo el Easy Apply: conviene un mensaje directo a los fundadores o al "
            "equipo de marketing con el caso de SMASH y el coste por pieza como gancho."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
