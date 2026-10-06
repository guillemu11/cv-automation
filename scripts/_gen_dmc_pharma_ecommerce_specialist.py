"""One-off: CV de una página para **E-commerce Specialist** que recluta **Decision Management
Consultants LLC** para un cliente farmacéutico "world-leading" con sede en Dubái. Contrato de 12
meses (proyecto). Salario publicado: 20.000 AED/mes (justo el suelo). Publicada el 2026-09-22.

Qué es el puesto: KAM de e-commerce del lado marca para cuentas digitales clave en KSA y EAU.
Relación diaria con los e-retailers, joint business plans, P&L de las cuentas asignadas,
stakeholders internos (finanzas, supply chain, marketing), negociación de listings, visibilidad,
retail media y condiciones comerciales; digital shelf (contenido, calidad de página, ratings y
reviews, búsqueda, precio, disponibilidad); campañas de retail media; presupuestos, facturación y
revisiones; lanzamientos y visibilidad promocional; scorecards y forecasts. Requisitos: 5-7 años en
e-commerce u omnicanal, cuentas e-commerce GCC (ideal KSA + EAU), **FMCG, beauty o consumer health
esencial** (OTC / e-pharmacy es ventaja), digital shelf y retail media, Excel y dashboards. Inglés
obligatorio, **árabe solo ventaja**. La cabecera dice "10+ years" pero los requisitos dicen 5-7.

Ángulo honesto de Paula:
- **Lado marca en las plataformas de EAU**: hoy gestiona sus marcas FMCG en talabat, Noon y Careem
  (listings, precio, mecánicas promocionales, activaciones conjuntas: sorteo con cashback de
  talabat y samplings estacionales). Resultados medibles: SMASH × talabat +165% de venta diaria vs
  +4,7% del grupo de control; Befit × Noon "New Year, New Me" +4.176 uds incrementales, +31%.
- **Lado plataforma**: Miravia (Alibaba), 42 cuentas de beauty, fragancias y moda (incluida KIKO
  Milano), +30% GMV QoQ; canal Flash Sales reportando al CEO contra objetivos de P&L. Glovo: cuentas
  XL con planes de crecimiento de GMV y negociación de condiciones.
- **FMCG + beauty** (lo esencial): DoFreeze y Mondelez (FMCG), Miravia (beauty y fragancias).
- **Disponibilidad y supply chain**: hace ella la revisión semanal de riesgo de stock por dark store
  (cobertura de PO frente a demanda) con un dashboard que montaron con IA; aporta al forecast
  mensual de sell-in.
- **Contenido y calidad de página**: su Shopify de punta a punta; contenido y surtido en Miravia.

Guardarraíles de honestidad:
- **KSA**: sin evidencia de cuentas en Arabia Saudí. Solo EAU. No se afirma GCC más allá de eso.
- **Retail media de pago** (sponsored ads con presupuesto propio en Noon/talabat/Amazon): sigue sin
  confirmar. Se habla de "platform campaigns", "campaign slots" y "visibility", no de gestionar
  inversión en retail media.
- **P&L**: el de Miravia es "against P&L targets" del canal; en DoFreeze no es dueña del P&L.
- **Forecast**: "contribute" (lo cierra la e-commerce manager). Facturación: sin evidencia.
- **Ratings, reviews, búsqueda**: sin evidencia específica; no se afirman como logro.
- **OTC / e-pharmacy / consumer health**: cero. Befit es repostería fitness (sin azúcar); no se
  vende como consumer health.
- Árabe: no lo habla (aquí es ventaja, no requisito). Deliveroo: DoFreeze NO está.
- Visa: "UAE Employment Visa (employer-sponsored)".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Decision Management Consultants"
TITLE = "E-commerce Specialist (Pharma contract)"
DATE_FOLDER = "2026-10-06"

JOB_DESCRIPTION = """\
E-commerce Specialist. Decision Management Consultants LLC (for a pharmaceutical client). Dubai, UAE. Contract.
Industry: Pharma. Work experience: 10+ years (header) / 5-7 years (requirements). Salary: 20000 AED. Posted 09/22/2026.
We are recruiting an E-Commerce Specialist for a 12-month project with one of our clients, a world-leading
pharmaceutical company based in Dubai, recognized for innovation, quality and patient-centric healthcare solutions,
with a strong regional presence and a data-driven culture.
Key responsibilities: day-to-day relationship management with key e-commerce customers across KSA and the UAE; drive
joint business plans with clear commercial priorities and disciplined execution; own the P&L of assigned accounts;
manage internal stakeholders across finance, supply chain, marketing and other functions; lead negotiations for product
listings, visibility, retail media and trade terms; optimize product content and page quality to improve visibility
and conversion; track ratings, reviews, search performance, pricing and product availability; plan, launch and optimize
retail media campaigns across leading e-commerce platforms; manage budgets, invoicing and performance reviews; manage
product submissions, launches and promotional visibility; track sales, visibility, stock levels and media investments;
maintain scorecards, forecasts and issue resolution with key stakeholders.
Requirements: 12-month project-based position; Bachelor's degree in Business, Management or related; minimum 5-7 years
of experience in e-commerce, digital commerce or omnichannel roles; proven experience managing GCC e-commerce accounts,
ideally across KSA and the UAE; experience in FMCG, beauty or consumer health is essential, OTC or e-pharmacy is an
advantage; strong knowledge of digital shelf management, retail media and performance-driven execution; proficiency in
platform analytics, Microsoft Excel and performance dashboards; strong stakeholder management, commercial acumen and
execution skills; fluency in English required, Arabic an advantage.
Benefits: as per UAE Labour Law; competitive salary.
"""

ATS = [
    "e-commerce", "eCommerce", "digital commerce", "omnichannel", "key accounts", "key account management",
    "e-commerce customers", "e-retailers", "GCC", "UAE", "talabat", "Noon", "Careem", "Glovo", "Miravia",
    "joint business plans", "growth plans", "P&L", "commercial acumen", "negotiation", "listings", "visibility",
    "trade terms", "commercial terms", "retail media", "platform campaigns", "digital shelf", "product content",
    "page quality", "conversion", "pricing", "availability", "stock levels", "out-of-stock", "PO coverage",
    "supply chain", "finance", "marketing", "stakeholder management", "budgets", "performance reviews",
    "product launches", "promotional visibility", "promotions", "scorecards", "forecasts", "sell-in", "sell-out",
    "incremental units", "GMV", "ROI", "ROAS", "platform analytics", "Excel", "dashboards", "Power BI",
    "FMCG", "beauty", "fragrances", "KIKO Milano", "Shopify", "Business Administration", "Dubai",
]

CONTENT = {
    "headline": "E-Commerce Key Accounts · Digital Shelf · Platform Campaigns · FMCG & Beauty · UAE",
    "professional_summary": (
        "E-commerce and key account professional with 5 years across FMCG, beauty and marketplaces. In Dubai, grows FMCG brands on "
        "talabat, Noon and Careem with campaigns measured on sell-out; before, managed 42 beauty and fragrance "
        "accounts at Miravia (Alibaba), including KIKO Milano, at +30% GMV QoQ."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Homegrown UAE FMCG group (Befit, Eurocake, Flair) | talabat, Noon & Careem + Shopify D2C | 50+ markets",
            "bullets": [
                "Manage day-to-day work with the talabat, Noon and Careem teams — listings, pricing, promotional mechanics and joint activations (purchase-linked talabat cashback giveaway, seasonal sampling); Befit × Noon \"New Year, New Me\" added +4,176 incremental units, +31% over baseline",
                "Launched SMASH, a new Eurocake brand, on talabat with campaigns built to drive orders: daily sales +165% vs +4.7% for a control brand on the same platform; 6 product launches end-to-end (brief, packaging, pricing, go-to-market)",
                "Run the weekly stock-risk review by dark store with a dashboard we built using AI, checking PO coverage against current and upcoming demand to protect availability; contribute to the monthly sell-in forecast with e-commerce and finance",
                "Own the Shopify store end-to-end — product content, page quality, collections, discounts and checkout — merchandised on conversion and AOV; manage 4 influencer agencies (negotiated -30%) and lead a team of two",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees | platform side",
            "bullets": [
                "Managed 42 beauty, fragrance and fashion accounts (incl. KIKO Milano) on pricing, assortment, product content, campaign slots and promotions, delivering +30% GMV QoQ; onboarded 30+ new stores in two months as fragrance lead",
                "Owned the Flash Sales channel for Beauty, Fashion & Home reporting to the CEO — commercial plans against P&L targets, using traffic, conversion, ROI and ROAS data to reallocate promotional investment",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce & delivery platform | €500M+ revenue | platform side",
            "bullets": [
                "Managed XL accounts (KFC, Taco Bell, La Tagliatella, Sushi Shop) with data-led GMV growth plans and in-app activations; negotiated commercial terms and aligned marketing, logistics and customer support on execution",
            ],
        },
        {
            "company": "Mondelez International · Massimo Dutti (Inditex)",
            "role": "Trainee, Category Planning · Sales Associate",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG (chocolate category) | Premium fashion retail",
            "bullets": [
                "At Mondelez, sell-in/sell-out and promotional-effectiveness analysis for chocolate, contributing to NPD launches (Milka Spread, Mini Suchard); at Inditex, store operations and customer service",
            ],
        },
    ],
    "skills_brand": (  # -> "Key Accounts"
        "talabat, Noon, Careem (brand side) · Miravia, Glovo (platform side) · listings, negotiation"
    ),
    "skills_ecommerce": (  # -> "Digital Shelf"
        "product content & page quality, assortment & MSL, availability & stock risk, launches, Shopify"
    ),
    "skills_commercial": (  # -> "Commercial"
        "pricing, promotions, growth plans, P&L targets, budgets, agency negotiation, forecast input"
    ),
    "skills_data": (  # -> "Campaigns & KPIs"
        "platform campaigns, sampling, creators, Meta & Google Ads · GMV, ROI, ROAS, sell-out vs control"
    ),
    "skills_tools": (  # -> "Tools"
        "Excel (Advanced), Power BI, Tableau, Looker, Salesforce, SAP, Shopify, Claude (AI dashboards)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Key Accounts",
    "E-Commerce & Digital": "Digital Shelf",
    "Commercial": "Commercial",
    "Data & Analytics": "Campaigns & KPIs",
}


def make_job() -> Job:
    return Job(
        id="dmc-pharma-ecommerce-specialist-dubai-2026-10",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE",
        url="https://www.google.com/search?q=Decision+Management+Consultants+E-commerce+Specialist+Dubai",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw="20,000 AED/month",
        salary_aed_min=20000,
        salary_aed_max=20000,
        raw={
            "query": "Decision Management Consultants E-commerce Specialist Dubai",
            "note": "Agencia de reclutamiento para un cliente farmacéutico en Dubái; contrato de 12 meses. KAM de "
                    "e-commerce lado marca, cuentas KSA + EAU, digital shelf y retail media. Encaje: talabat, Noon "
                    "y Careem (marca) + Miravia y Glovo (plataforma), FMCG y beauty. Gaps: KSA, retail media de "
                    "pago sin confirmar, OTC / e-pharmacy.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": job.salary_raw,
        "salary_aed_min": job.salary_aed_min, "salary_aed_max": job.salary_aed_max,
        "posted_date": "2026-09-22", "raw": job.raw,
        "ai_score": 74, "ai_tier": "Warm",
        "skills_match": [
            "Cuentas e-commerce de EAU desde el lado marca: talabat, Noon y Careem",
            "Key account management desde el lado plataforma: Miravia (42 cuentas, +30% GMV QoQ) y Glovo (cuentas XL)",
            "FMCG (DoFreeze, Mondelez) y beauty (Miravia, cuenta KIKO Milano), lo que el JD llama esencial",
            "Campañas en plataforma con sell-out medido: SMASH × talabat +165% vs control, Befit × Noon +31%",
            "Disponibilidad: revisión semanal de riesgo de stock por dark store con dashboard propio",
            "Contenido y calidad de página: Shopify de punta a punta, contenido y surtido en Miravia",
            "P&L: canal Flash Sales de Miravia contra objetivos de P&L, reportando al CEO",
            "Lanzamientos: marca nueva SMASH y 6 lanzamientos de producto",
        ],
        "missing_skills": [
            "Cuentas en Arabia Saudí (el JD pide KSA y EAU)",
            "Retail media de pago con presupuesto propio (sponsored ads en Noon, talabat o Amazon): sin confirmar",
            "OTC, consumer health o e-pharmacy (ventaja, no requisito)",
            "Ratings, reviews y share of search como KPI gestionado",
            "Facturación con cuentas",
            "Árabe (ventaja, no requisito)",
        ],
        "sector_fit": "bueno — FMCG y beauty cumplen lo esencial; pharma / consumer health sería nuevo",
        "seniority_fit": "acorde — 5 años frente a 5-7 pedidos; Specialist por debajo de su título de Manager",
        "red_flags": [
            "Contrato de 12 meses vía agencia: deja un puesto fijo y un visado de empleador por un proyecto temporal",
            "Salario publicado 20.000 AED/mes: justo en el suelo, sin margen",
            "La cabecera dice 10+ años de experiencia (los requisitos dicen 5-7)",
            "Cliente farmacéutico sin nombre",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Encaje sólido en el núcleo del puesto: lleva cuentas de e-commerce de EAU desde el lado marca (talabat, "
            "Noon, Careem) y antes desde el lado plataforma (Miravia, Glovo), con resultados medibles, y cumple el "
            "requisito esencial de FMCG o beauty. Los huecos son KSA, retail media de pago sin confirmar y "
            "consumer health / OTC. El riesgo real es el formato: contrato de 12 meses vía agencia, al suelo "
            "salarial, cambiando un puesto fijo con visado."
        ),
        "scored_by": "manual:claude", "freshness": "aging",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
