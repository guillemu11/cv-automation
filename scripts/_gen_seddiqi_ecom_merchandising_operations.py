"""One-off: CV de una página para **Executive - Ecommerce Merchandising & Operations**
en **Seddiqi Holding** (lujo: relojería y joyería, Ahmed Seddiqi, Dubai, presencial).

Pide el JD: trading diario de los storefronts asignados (volumen, conversión, revenue vs
forecast); diagnóstico de gaps a nivel producto y causa (sessions → CR → volumen → returns →
AIV → net sales); weekly trading read; disponibilidad digital, sell-through por categoría y
mix de full-price; merchandising on-site (estructura de colecciones, orden de PLP, hero
placement, navegación, featured slots, búsqueda interna); higiene de catálogo, calidad de
datos de producto, taxonomía/atributos, pricing y markdown; carga masiva de SKUs y launches;
contenido y campañas con marketing/omnichannel/brand; localización (árabe = ventaja);
stock diario (broken sizes, OOS de líneas clave), escalado de bloqueos con impacto cuantificado
y transferencias de stock entre mercados; conocimiento de order/fulfilment/returns.

Ángulo honesto de Paula: Miravia (Alibaba) — 42 cuentas de beauty/fragancias/moda con surtido,
pricing, promociones y **Flash Sales** reportando al CEO (+30% GMV QoQ), leyendo tráfico,
conversión y sell-through a diario; DoFreeze — Shopify end-to-end (catálogo, contenido,
colecciones, pricing, CRO) más listings en talabat/Noon/Careem y revisión semanal de riesgo de
stock por dark store; Glovo — catálogo y promociones de cuentas XL; Mondelez — sell-in/sell-out
con Nielsen.

Guardarraíles de honestidad:
- **Sin árabe** (ES nativo, EN C1). El JD lo marca como "advantage", no requisito: no se insinúa.
- Forecast: **contribuye** con Sales/Finance, no es dueña (la e-commerce manager lo es).
- Sin experiencia en lujo/relojería-joyería ni en retail de tienda física — no se insinúa.
- Visa: solo "UAE Residence Visa" (patrocinada por el empleador); nunca "no sponsorship needed".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Seddiqi Holding"
TITLE = "Executive - Ecommerce Merchandising & Operations"
DATE_FOLDER = "2026-09-22"

JOB_DESCRIPTION = """\
Executive - Ecommerce Merchandising & Operations — Seddiqi Holding, Dubai, UAE. On-site, full time.
Reports to the Ecommerce Operations Manager. Trading and merchandising led role: grow online sales by
reading performance data daily, diagnosing where a gap sits, and acting through on-site merchandising,
product focus, placement and campaign execution.
Trading Performance: own daily trading performance of allocated storefronts, reviewing volume, conversion
and revenue against forecast every working day; diagnose performance gaps to product and cause level
through the driver sequence (sessions, conversion rate, volume before returns, return rate, volume after
returns, average item value, net sales) to determine whether traffic, conversion, stock availability or
price is the issue; produce a weekly trading read on performance vs forecast, best and worst performing
products and actions; monitor and report weekly on digital availability, sell-through rate by category and
full-price sales mix; track competitor activity, local market events and seasonal demand patterns.
Digital Merchandising and Site Execution: own the trading and merchandising calendar (launches, promotions,
seasonal campaigns) sequenced to local demand; execute and optimise on-site merchandising — collection
structure, product listing page sequencing, hero placement, navigation, featured slots — responding to
weekly performance and internal search behaviour; manage product presentation and site hygiene including
catalogue accuracy, product data quality, categorisation, pricing updates and markdown execution; apply
taxonomy and attribute standards and naming conventions; gather and validate SKU and product data from
cross-functional teams and execute launches through product information and bulk upload tools; adapt
merchandising and assortment by market based on local calendars, customer behaviour and sell-through.
Content, Campaign and Localisation: deliver content updates (product grids, banners, campaign storytelling)
partnering with marketing, omnichannel and brand; support Arabic content quality (written Arabic an
advantage); contribute to creative briefs for digital merchandising assets.
Stock, Availability and Escalation: monitor stock availability and depth daily, identify broken sizes,
out-of-stock key lines and lost sales risk; escalate trading constraints early, naming the constraint,
quantifying commercial impact and requesting a decision; support cross-market stock transfer workflow.
Operational Support: maintain working awareness of order, fulfilment and returns processes; flag inventory
discrepancies between inventory system, WMS and the ecommerce platform.
Qualifications: Bachelor Degree or equivalent with Business, Ecommerce or Retail focus. Minimum 3 years site
merchandising experience in ecommerce, digital trading, retail merchandising or buying, with demonstrated
exposure to online sales performance.
"""

ATS = [
    "ecommerce", "e-commerce", "site merchandising", "digital merchandising", "online trading",
    "trading performance", "daily trading", "weekly trading read", "performance vs forecast",
    "conversion rate", "sessions", "traffic", "average item value", "net sales", "return rate",
    "sell-through rate", "digital availability", "full-price sales mix", "markdown", "pricing updates",
    "product listing page", "PLP", "collection structure", "hero placement", "navigation",
    "featured slots", "internal search", "product discoverability", "catalogue accuracy",
    "product data quality", "taxonomy", "attributes", "categorisation", "naming conventions",
    "SKU", "bulk upload", "PIM", "product launches", "site hygiene", "trading calendar",
    "promotions", "seasonal campaigns", "assortment", "campaign execution", "banners", "product grids",
    "localisation", "stock availability", "out of stock", "lost sales", "stock transfer",
    "order fulfilment", "returns", "WMS", "inventory accuracy", "competitor tracking",
    "site analytics", "Google Analytics", "Shopify", "luxury retail", "cross-functional",
]

CONTENT = {
    "headline": "E-Commerce Trading · Site Merchandising · Catalogue, Pricing & Availability",
    "professional_summary": (
        "E-commerce trading and site merchandising professional with 5 years across Alibaba's Miravia "
        "(42 accounts: assortment, pricing, placement and Flash Sales, +30% GMV QoQ), Glovo and FMCG, now "
        "running Shopify and UAE marketplace storefronts end-to-end in Dubai."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG group (Befit, Eurocake, Flair) | Shopify D2C + talabat, Noon, Careem | 50+ markets",
            "bullets": [
                "Own the Shopify storefront end-to-end — collection structure, product listing sequencing, hero banners, navigation, catalogue accuracy and product data — lifting conversion rate and average order value through weekly data-led merchandising",
                "Read sessions, conversion, basket and sell-out weekly to diagnose gaps to product level and act the same week on placement, price or content; escalate stock and pricing blockers with quantified commercial impact",
                "Own the trading calendar across web, talabat, Noon and Careem — 6 NPD launches, Ramadan and seasonal campaigns — validating SKU data, executing bulk listings on time and aligning marketing, brand and 4 agencies on assets",
                "Run a weekly availability and listing review by dark store to catch out-of-stock key lines and lost-sales risk, and contribute to the monthly forecast with Sales and Finance",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees",
            "bullets": [
                "Traded 42 storefronts daily — assortment, pricing, on-site placement and promotions against competitor benchmarks — delivering +30% GMV growth QoQ",
                "Owned the Flash Sales channel for Beauty, Fashion & Home reporting to the CEO: event calendar, item selection, quantities, full-price vs promotional mix and margin-aligned commercial plans",
                "Produced weekly trading reads on traffic, conversion, sell-through and best/worst performing products, and led category expansion as PIC Fragrances (30+ stores onboarded in two months)",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Managed XL partners (KFC, Taco Bell, Sushi Shop) on catalogue quality, menu merchandising, promotions and GMV, and helped build the Retail vertical onboarding fashion and lifestyle brands",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | Chocolate category (Milka, Suchard) | €36B revenue",
            "bullets": [
                "Built sell-in/sell-out and promotional-effectiveness reports with Nielsen and flagged category trends to the commercial team",
            ],
        },
    ],
    "skills_brand": (  # label -> "Trading & Merchandising"
        "daily trading, weekly trading read, assortment & range, pricing & markdown, promotions calendar"
    ),
    "skills_ecommerce": (  # label -> "Site Execution"
        "Shopify, PLP sequencing, collections & navigation, hero placement, catalogue accuracy, taxonomy & attributes, bulk SKU upload, CRO"
    ),
    "skills_commercial": (  # label -> "Stakeholders"
        "marketing, brand, omnichannel, buyers, supply chain, customer service, agencies"
    ),
    "skills_data": (  # label -> "Analytics & Availability"
        "site analytics, conversion & sessions, sell-through, digital availability, stock & OOS risk, competitor tracking"
    ),
    "skills_tools": (  # label -> "Tools"
        "Excel (Expert), Shopify, Google Analytics, Power BI, Tableau, Looker, SAP, Salesforce, Claude (AI)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Trading & Merchandising",
    "E-Commerce & Digital": "Site Execution",
    "Commercial": "Stakeholders",
    "Data & Analytics": "Analytics & Availability",
}


def make_job() -> Job:
    return Job(
        id="seddiqi-holding-ecommerce-merchandising-operations-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE",
        url="https://www.linkedin.com/jobs/search/?keywords=Seddiqi%20Holding%20Ecommerce%20Merchandising%20Operations",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Seddiqi Holding Executive Ecommerce Merchandising Operations Dubai",
            "note": "Lujo (relojería/joyería, Ahmed Seddiqi). Trading + site merchandising puro: encaje "
                    "funcional muy alto con Miravia y Shopify. Nivel 'Executive' (por debajo de su título "
                    "actual de Manager) y árabe escrito como ventaja (Paula no tiene árabe). "
                    "+100 solicitudes, 142 candidatos. Recruiter: Nick Sisnett.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 70, "ai_tier": "Warm",
        "skills_match": [
            "Trading diario de 42 storefronts en Miravia: surtido, pricing, placement y benchmarking (+30% GMV QoQ)",
            "Flash Sales channel owner (Beauty, Fashion & Home) reportando al CEO: calendario, selección e items y mix full-price",
            "Shopify end-to-end: colecciones, orden de PLP, banners, navegación, calidad de datos de producto y CRO",
            "Weekly trading read: sesiones, conversión, sell-through, top/worst sellers y acción correctiva la misma semana",
            "Revisión semanal de disponibilidad y riesgo de OOS por dark store; escalado con impacto cuantificado",
        ],
        "missing_skills": [
            "Árabe escrito (marcado como ventaja en el JD)",
            "Lujo / relojería y joyería como categoría",
            "Herramientas PIM específicas de retail de lujo",
        ],
        "sector_fit": "medio-alto — e-commerce trading y site merchandising sí; lujo/relojería no",
        "seniority_fit": "por debajo — nivel Executive frente a su título actual de Manager; decisión de Paula",
        "red_flags": [
            "Nivel 'Executive' y descrito como 'entry point into ecommerce trading': posible downgrade de título y banda salarial",
            "Árabe escrito como ventaja y localización de contenido árabe entre las funciones",
            "142 candidatos y +100 solicitudes en 6 días",
            "100% presencial en Dubái",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "El JD es trading y site merchandising puro: leer el embudo a diario (sesiones, CR, devoluciones, AIV, "
            "net sales), diagnosticar a nivel producto y actuar sobre placement, precio o contenido. Paula ha hecho "
            "exactamente eso en Miravia (42 cuentas, Flash Sales al CEO) y lo hace hoy en Shopify y quick-commerce "
            "en Dubái. Las dos brechas reales son el árabe escrito (ventaja, no requisito) y el lujo como categoría. "
            "La duda no es el encaje, es el nivel: 'Executive' suena a paso atrás desde Manager."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
