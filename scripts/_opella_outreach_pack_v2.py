"""Build Opella_Outreach_Pack.docx — 13 personalized outreach blocks.

For each of the 13 v2 contacts:
  - Generates a tailored email (subject + body) that references the two
    Vercel-hosted campaign landings:
       paula-pitch-a  →  Buscopan "Silent Co-Pilot"
       paula-pitch-b  →  Doliprane "La Pastille Francaise"
  - Generates a LinkedIn connection note (<=300 chars).
  - Generates a longer LinkedIn InMail variant.

Outputs:
  output/Opella_Outreach_Pack.docx       — copy-paste ready Word document
  data/outreach/f324f5af0ff41e06_v2.json  — structured backup of all 13 packs

The LLM provider is forced to anthropic for this run (the .env default is
chatqueue which would block waiting for a human reply).
"""
from __future__ import annotations

import html
import json
import logging
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from career_ops.config import settings

settings.llm_provider = "anthropic"

from career_ops import llm  # noqa: E402

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-5s  %(name)s  %(message)s",
    datefmt="%H:%M:%S",
)

OPELLA_JOB_ID = "f324f5af0ff41e06"
VERCEL_BUSCOPAN = "https://paula-pitch-a.vercel.app"
VERCEL_DOLIPRANE = "https://paula-pitch-b.vercel.app"


# ---------------------------------------------------------------------
# 1. Custom outreach generator (Vercel-aware, speculative-tone-aware)
# ---------------------------------------------------------------------

_SYSTEM_PROMPT = """\
You write concise, compelling speculative outreach for Paula De Francisco, a
Brand/Marketing Manager based in Dubai. The Opella AMET Brand Manager posting
she had targeted has just been pulled (likely paused for internal referral or
agency search), so every message must work as SPECULATIVE outreach — NOT as
"I saw your job posting and want to apply".

Tone: professional, intelligent, concrete, never sycophantic, never
self-impressed, never AI-slop (no "passionate", "innovative", "leverage",
"unleash", "synergy", "I came across your profile", "I hope this finds you well").

Each message must reference Paula's two speculative campaign landings already
deployed online:

  - Buscopan "Pharmacy closed. Relief isn't." — a late-night quick-commerce
    play in UAE. Live at: {vercel_buscopan}
  - Doliprane "The word still means paracetamol." — a MENA Francophone
    diaspora launch for the iconic French OTC. Live at: {vercel_doliprane}

CRITICAL: the email body must NATURALLY weave in BOTH URLs as clickable
hyperlinks, with one sentence describing each. They are two distinct
strategic angles, not two versions of the same idea. Don't just paste both
links at the bottom — embed them inline in a sentence each, like:

  ... I prototyped two angles speculatively: one on late-night quick-commerce
  for Buscopan ({vercel_buscopan}) and one on a MENA Francophone launch for
  Doliprane ({vercel_doliprane}) ...

Rules:
1. Email subject: under 60 chars, specific. NEVER use "Application for"
   (the posting was pulled — that frame is wrong). Try angles like
   "Two Opella AMET campaign concepts" or "Pitch: Buscopan late-night q-comm".
2. Email body: HTML with <p> tags. 5-7 sentences. Open with a SPECIFIC
   observation about Opella, the contact's domain, or AMET — not a
   compliment. Land the two landing URLs naturally. End with a soft ask
   (15 min call, opinion on the angles, an introduction).
3. LinkedIn connection note: MUST be under 300 characters. One sentence
   about why connecting + one about what you bring. Reference ONE landing
   only — pick the one most relevant to the contact's domain. Always
   include the URL.
4. LinkedIn InMail: 3-5 sentences, slightly more detailed than the
   connection note. Reference one or both URLs depending on contact type.
5. Tailor the angle to the contact:
   - A-tier (Hiring Manager) → pitch directly, both URLs prominent
   - B-tier (Country head) → frame as commercial opportunity per geography
   - C-tier (Peer Brand Mgr) → peer-to-peer 15-min ask, one URL as conversation seed
   - D-tier (Cross-functional) → angle to their function (digital/trade/etc)
   - E-tier (Referral node) → ask for a warm intro to Murali Rao
6. NEVER reference "the open Brand Manager position" or "your job posting"
   — it was pulled. Always frame as "I wanted to share two campaign concepts I
   built for AMET" or similar speculative framing.

Output via the submit_outreach tool. The email_body MUST be valid HTML with
<p> tags. The URLs MUST appear as <a href="...">...</a> tags inside it.
"""


_TOOL = {
    "name": "submit_outreach",
    "description": "Submit all outreach copy for this contact",
    "input_schema": {
        "type": "object",
        "properties": {
            "email_subject": {"type": "string"},
            "email_body": {"type": "string", "description": "HTML with <p> tags, MUST include both Vercel URLs as <a href> hyperlinks inline"},
            "linkedin_connection": {"type": "string", "description": "<=300 chars, must include the most relevant Vercel URL"},
            "linkedin_inmail": {"type": "string"},
        },
        "required": ["email_subject", "email_body", "linkedin_connection", "linkedin_inmail"],
    },
}


def generate_for_contact(contact: dict, job: dict) -> dict:
    """One LLM call per contact. Returns the submit_outreach dict."""
    first_name = contact["name"].split()[0]
    system = _SYSTEM_PROMPT.format(
        vercel_buscopan=VERCEL_BUSCOPAN,
        vercel_doliprane=VERCEL_DOLIPRANE,
    )
    user = f"""\
## Contact
- Name: {contact['name']}
- Use first name: {first_name}
- Title: {contact['title']}
- LinkedIn: {contact['linkedin_url']}
- Email (pattern-guess, unverified): {contact['primary_email']}
- Tier: {contact['tier']}  ({contact['role_type']})
- Strategic notes about this contact:
  {contact['notes']}

## Target context
- Opella AMET (Africa, Middle East, Turkey) — consumer healthcare/OTC
- The Brand Manager - Dubai posting Paula had targeted was just pulled.
- Speculative outreach. The two landings are the conversation seed.

## Paula's differentiators (use sparingly, only where relevant)
- 4 years FMCG + UAE quick-commerce execution (Noon, Talabat, Careem, Deliveroo)
- +30% GMV QoQ at Alibaba/Miravia across 42 key accounts
- Mondelez + Inditex foundation, Spanish/English C1, already in Dubai with visa

## Vercel landings to embed
- Buscopan campaign (late-night quick-commerce): {VERCEL_BUSCOPAN}
- Doliprane campaign (MENA Francophone diaspora): {VERCEL_DOLIPRANE}

Use the submit_outreach tool. The email body MUST be HTML and MUST contain
BOTH Vercel URLs as inline <a href> hyperlinks (one sentence each, woven in
naturally — not pasted at the end)."""

    data = llm.generate_structured(
        system=system,
        user=user,
        tool_schema=_TOOL,
        tier="sonnet",
        max_tokens=1500,
        temperature=0.4,
    )
    if not data:
        raise RuntimeError(f"empty LLM response for {contact['name']}")
    # Enforce 300-char limit on LinkedIn connection
    conn = data.get("linkedin_connection", "")
    if len(conn) > 300:
        data["linkedin_connection"] = conn[:297] + "..."
    return data


# ---------------------------------------------------------------------
# 2. HTML -> readable text for Word
# ---------------------------------------------------------------------

def html_to_paragraphs(html_str: str) -> list[str]:
    """Convert <p>-tagged HTML into a list of plaintext paragraphs.

    Anchor URLs become 'text (url)' so they remain clickable when copied.
    """
    s = html_str

    # Replace <a href="X">label</a> -> "label (X)"
    def _anchor_repl(m: re.Match) -> str:
        href = m.group(1).strip()
        label = re.sub(r"<[^>]+>", "", m.group(2)).strip()
        if label and label != href:
            return f"{label} ({href})"
        return href

    s = re.sub(r'<a[^>]*href="([^"]+)"[^>]*>(.*?)</a>', _anchor_repl, s, flags=re.S | re.I)
    s = re.sub(r"<br\s*/?>", "\n", s, flags=re.I)
    # Split on </p>
    parts = re.split(r"</p\s*>", s, flags=re.I)
    paras = []
    for p in parts:
        text = re.sub(r"<[^>]+>", "", p)
        text = html.unescape(text)
        text = re.sub(r"[ \t]+", " ", text)
        text = text.strip()
        if text:
            paras.append(text)
    return paras


# ---------------------------------------------------------------------
# 3. Word doc builder
# ---------------------------------------------------------------------

def build_docx(outreach_records: list[dict], out_path: Path) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # --- styling: tighten the default base ---
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # --- cover ---
    title = doc.add_heading("Opella AMET — Outreach Pack", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    sub = doc.add_paragraph()
    sub_run = sub.add_run(
        "Speculative outreach to 13 AMET contacts. The Brand Manager - Dubai posting "
        "was pulled in May 2026; the play is to insert directly into the org before a "
        "referral closes the gap."
    )
    sub_run.italic = True
    sub_run.font.size = Pt(10)
    sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # --- attached landings block ---
    doc.add_heading("Attached campaign landings", level=2)
    p1 = doc.add_paragraph()
    p1.add_run("Buscopan — \"Pharmacy closed. Relief isn't.\" (late-night quick-commerce):  ").bold = True
    p1.add_run(VERCEL_BUSCOPAN)
    p2 = doc.add_paragraph()
    p2.add_run("Doliprane — \"The word still means paracetamol.\" (MENA Francophone diaspora):  ").bold = True
    p2.add_run(VERCEL_DOLIPRANE)

    # --- outreach order block ---
    doc.add_heading("Suggested outreach sequence", level=2)
    seq = doc.add_paragraph()
    seq.add_run(
        "A-tier first (Murali, Duygu, Hossam) — they decide. "
        "C-tier (Olivia Stefanelli) is the strongest peer-to-peer warm intro. "
        "E-tier as referral pressure if A-tier is silent by week 3. "
        "One contact per day max — don't let them compare notes."
    )

    doc.add_page_break()

    # --- one section per contact ---
    for i, rec in enumerate(outreach_records, start=1):
        c = rec["contact"]
        o = rec["outreach"]
        head = doc.add_heading(f"{i:02d}. {c['name']}", level=1)
        head.paragraph_format.keep_with_next = True

        # subheading: title + tier + role_type
        meta = doc.add_paragraph()
        meta.add_run(f"{c['title']}").italic = True
        meta.add_run(f"   ·   Tier {c['tier']}   ·   {c['role_type']}").italic = True

        # contact lines
        link_p = doc.add_paragraph()
        link_p.add_run("LinkedIn:  ").bold = True
        link_p.add_run(c["linkedin_url"])
        email_p = doc.add_paragraph()
        email_p.add_run("Email:  ").bold = True
        email_p.add_run(f"{c['primary_email']}  (status: {c['primary_email_status']})")

        if c.get("notes"):
            n = doc.add_paragraph()
            n_run = n.add_run(f"Why this contact:  {c['notes']}")
            n_run.font.size = Pt(9)
            n_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

        # --- email ---
        doc.add_heading("Email", level=3)
        s = doc.add_paragraph()
        s.add_run("Subject:  ").bold = True
        s.add_run(o["email_subject"])
        for para in html_to_paragraphs(o["email_body"]):
            doc.add_paragraph(para)

        # --- LinkedIn connection ---
        doc.add_heading("LinkedIn connection note (<=300 chars)", level=3)
        doc.add_paragraph(o["linkedin_connection"])

        # --- LinkedIn InMail ---
        doc.add_heading("LinkedIn InMail", level=3)
        doc.add_paragraph(o["linkedin_inmail"])

        # divider
        sep = doc.add_paragraph()
        sep_run = sep.add_run("— · —")
        sep_run.font.size = Pt(8)
        sep_run.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
        sep.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.save(str(out_path))


# ---------------------------------------------------------------------
# 4. Driver
# ---------------------------------------------------------------------

def main() -> None:
    # Load contacts
    contacts_path = settings.data_dir / "contacts" / f"{OPELLA_JOB_ID}.emails.json"
    contacts = json.loads(contacts_path.read_text(encoding="utf-8"))
    print(f"[start] {len(contacts)} contacts. Generating personalized outreach via Claude...")

    # Load job (for context only)
    scored_path = settings.data_dir / "scored_jobs.json"
    scored = json.loads(scored_path.read_text(encoding="utf-8"))
    job = next(r for r in scored if r["id"].startswith(OPELLA_JOB_ID))

    records = []
    for i, c in enumerate(contacts, start=1):
        print(f"  [{i}/{len(contacts)}] {c['name']} (tier {c['tier']}) ...")
        try:
            o = generate_for_contact(c, job)
            records.append({"contact": c, "outreach": o})
            print(f"     subject: {o['email_subject']}")
        except Exception as exc:  # noqa: BLE001
            print(f"     FAILED: {exc}")
            continue

    # --- persist structured JSON for re-use ---
    outreach_dir = settings.data_dir / "outreach"
    outreach_dir.mkdir(parents=True, exist_ok=True)
    json_path = outreach_dir / f"{OPELLA_JOB_ID}_v2.json"
    json_path.write_text(
        json.dumps(records, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"\n[saved] {json_path}")

    # --- build the Word doc ---
    docx_path = settings.output_dir / "Opella_Outreach_Pack.docx"
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    build_docx(records, docx_path)
    print(f"[saved] {docx_path}")
    print("Done.")


if __name__ == "__main__":
    main()
