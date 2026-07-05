"""One-off: build a Word version of the Colgate application fill-in guide."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

from career_ops.config import settings

OUT = (settings.output_dir / "2026-06-11" /
       "Colgate-Palmolive - Ecommerce Manager" / "04_Aplicacion" /
       "RELLENAR_A_MANO_Colgate.docx")

RED = RGBColor(0xC0, 0x00, 0x00)
GREEN = RGBColor(0x1F, 0x7A, 0x1F)
BLUE = RGBColor(0x0B, 0x4F, 0x8A)


def main() -> None:
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    h = doc.add_heading("Colgate-Palmolive — Ecommerce Manager (Dubai) · Job #174114", level=0)
    p = doc.add_paragraph()
    p.add_run("Apply: ").bold = True
    p.add_run("https://jobs.colgate.com/job/Dubai-Ecommerce-Manager-DU/174114-en_US/")
    p = doc.add_paragraph()
    p.add_run("Subir: ").bold = True
    p.add_run("CV_Paula_Colgate-Palmolive_Ecommerce_Manager.pdf  +  CL_Paula_Colgate-Palmolive_Ecommerce_Manager.pdf")
    note = doc.add_paragraph()
    r = note.add_run("Verde = ya hecho   ·   ROJO = te falta rellenar")
    r.italic = True

    def section(title):
        doc.add_heading(title, level=1)

    def done(label, value=""):
        para = doc.add_paragraph(style="List Bullet")
        r = para.add_run("HECHO  ")
        r.bold = True
        r.font.color.rgb = GREEN
        para.add_run(label + (": " if value else ""))
        if value:
            rv = para.add_run(value)
            rv.bold = True

    def todo(label, value=""):
        para = doc.add_paragraph(style="List Bullet")
        r = para.add_run("FALTA  ")
        r.bold = True
        r.font.color.rgb = RED
        para.add_run(label + (" -> " if value else ""))
        if value:
            rv = para.add_run(value)
            rv.bold = True
            rv.font.color.rgb = RED

    # ---- My Documents ----
    section("My Documents")
    done("Resume / CV")
    done("Cover Letter")

    # ---- Profile Information ----
    section("Profile Information")
    done("Preferred First Name", "Paula")
    done("Preferred Last Name", "De Francisco Perez")
    done("Phone", "+971503863656")
    done("Email", "paulich98@hotmail.com")
    done("Country", "United Arab Emirates")
    done("Preferred Language", "English US")
    done("Company Name of Current/Last Employer", "Dofreeze")
    done("Title of Current/Last Job", "Brand Manager")
    todo("Previously employed by our company?", "No")

    # ---- Employment History ----
    section("Employment History  (3 filas — rellena las tres)")
    emp_headers = ["Start", "End", "Company Name", "Country", "Title", "Category"]
    emp_rows = [
        ["10/2025", "(actual)", "DoFreeze LLC", "United Arab Emirates",
         "Brand & Marketing Manager", "Marketing"],
        ["11/2023", "10/2025", "Miravia / AliExpress (Alibaba Group)", "Spain",
         "Key Account Manager – Beauty, Fragrances & Fashion", "Sales / Account Management"],
        ["09/2022", "11/2023", "Glovo", "Spain",
         "Account Manager – XL Accounts", "Sales / Account Management"],
        ["08/2021", "08/2022", "Mondelez International (opcional, 'Add')", "Spain",
         "Trainee – Category Planning", "Marketing"],
    ]
    t = doc.add_table(rows=1, cols=len(emp_headers))
    t.style = "Light Grid Accent 1"
    for i, htext in enumerate(emp_headers):
        c = t.rows[0].cells[i]
        c.paragraphs[0].add_run(htext).bold = True
    for row in emp_rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val

    # ---- Education ----
    section("Education  (1 fila — vacía)")
    edu = [
        ("Completed or In Progress?", "Completed"),
        ("Name of School", "CUNEF Universidad"),
        ("Country", "Spain"),
        ("School Type", "University"),
        ("Degree", "Bachelor's Degree"),
        ("Primary Branch of Study", "Business Administration"),
        ("Secondary Branch of Study", "E-Commerce"),
        ("Start Date at School", "09/2016"),
        ("End Date at School", "06/2020"),
    ]
    for label, val in edu:
        todo(label, val)

    # ---- Certifications ----
    section("Certifications/Licenses  (opcional — 'Add')")
    doc.add_paragraph("Google E-Commerce Certificate", style="List Bullet")
    doc.add_paragraph("Google Digital Marketing Certificate", style="List Bullet")

    # ---- Language Skills ----
    section("Language Skills")
    done("Spanish", "Native")
    done("English", "Excellent")

    # ---- Job-Specific ----
    section("Job-Specific Information")
    done("Currency", "Emirati Dirham")
    done("Before Tax Base Salary Expectations", "25000")
    done("Notice Period", "1 month and a half")
    todo("Related to a Colgate employee?", "No")

    # ---- Voluntary Self-ID ----
    section("Voluntary Self-Identification")
    done("Age", "27")
    done("Gender", "Female")
    done("Disability", "No, I do not have a disability...")

    # ---- Acknowledgment ----
    section("Applicant Acknowledgment  (ultimo paso, antes de Apply)")
    todo("Do you agree to these terms?", "I Agree")
    todo("Do you hold a bachelor's degree?", "Yes")
    todo("Minimum 4 years eCommerce / account management in multinational FMCG?", "Yes")

    # ---- Resumen ----
    section("Resumen: lo que te falta")
    for item in [
        "Employment History -> 3 filas (tabla arriba)",
        "Education -> 1 fila (datos arriba)",
        "Previously employed -> No",
        "Related to Colgate employee -> No",
        "Agree to terms -> I Agree",
        "Bachelor's degree -> Yes",
        "4+ anos eCommerce/FMCG -> Yes",
    ]:
        para = doc.add_paragraph(style="List Number")
        para.add_run(item)
    fin = doc.add_paragraph()
    fin.add_run("Cuando este todo -> Apply").bold = True

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print("OK_DOCX", OUT)


if __name__ == "__main__":
    main()
