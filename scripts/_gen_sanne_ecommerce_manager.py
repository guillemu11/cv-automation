"""One-off: CV de una página para **Ecommerce Manager** en **Sanne**
(casa de moda de lujo made-on-demand, Dubai HQ, presencial).

Pide el JD: ownership del rendimiento de venta de la web (CR, AOV, checkout) sobre un
**Shopify Plus** recién lanzado; operación diaria del site (listings, merchandising, pricing,
grid de tallas UK4–16 en medias tallas × Petite/Regular/Tall); **CRM end-to-end** (segmentación,
automatización, flows de retención, lifecycle); **Klaviyo end-to-end** incluyendo escribir el
copy en la voz de la marca, calendario, crecimiento de lista, testing y deliverability;
**programa de afiliados de creadores en ShopMy** (comisiones, links/códigos, reporting de
revenue y CPA contra la fee de plataforma); gestión del stack de apps de Shopify y de las
agencias (web, email, paid) bajo el CTO; briefing de contenido con Head of Content; funnel y
sistemas de back-end; reporting al CTO; y **crear y usar agentes de IA** para automatizar
procesos de e-commerce con aprobación humana en lo que toque pricing o datos de cliente.

Ángulo honesto de Paula — encaje inusualmente directo en los tres pilares:
- Site: lleva la tienda e-commerce de DoFreeze end-to-end (catálogo, merchandising, pricing,
  CRO, AOV).
- Email/CRM: planifica, **escribe** y envía el canal de email de DoFreeze (lanzamientos,
  promos, lifecycle) — el JD pide justo escribir copy que convierta.
- Creators: diseñó **Befit Crew**, un programa de afiliados de creadores en vivo — códigos
  únicos, comisión escalonada 8–10% (Rising / Elite), descuento de audiencia, seeding y
  funnel de registro self-serve.
- IA: construye agentes de IA que automatizan planificación, contenido y reporting (~40% de
  carga manual eliminada) — el JD lo pide explícitamente y casi nadie lo tiene.

Guardarraíles de honestidad:
- **Klaviyo**: NO se afirma. Su email/EDM ha corrido por otras herramientas; se posiciona como
  ownership real del canal, platform-agnostic. La carta lo dice de frente si se genera.
- **ShopMy**: NO se afirma. Su programa de afiliados es propio (códigos + funnel), no ShopMy.
- **Shopify Plus**: se dice "Shopify" a secas, que es lo cierto; nunca "Plus".
- **8+ años de digital marketing** (requisito añadido por el anunciante): Paula tiene ~5. No se
  infla ninguna fecha. Es la brecha dura de esta candidatura.
- Sin apparel de lujo DTC en plantilla; sin árabe; visa solo "UAE Residence Visa".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Sanne"
TITLE = "Ecommerce Manager"
DATE_FOLDER = "2026-09-22"

JOB_DESCRIPTION = """\
Ecommerce Manager — Sanne, Dubai, UAE (Sanne HQ). On-site, full time.
Sanne is a next-generation luxury house built around Made-on-Demand, entering a new phase of growth
following the launch of a new Shopify Plus website. This role owns the ecommerce engine: website sales
performance, day-to-day site operations, the CRM and email programme, and the creator affiliate channel.
Hands-on, execution-focused: build the operational foundation the business needs in order to grow.
Responsibilities: own website sales performance end to end (conversion rate, average order value, checkout
experience, overall commercial output); hands-on execution of the marketing roadmap across digital
touchpoints including new launches and international market expansion; manage day-to-day ecommerce
operations (product listings, merchandising, pricing, site maintenance, including a size and height grid
spanning UK 4 to 16 in half sizes across Petite, Regular and Tall); own and continuously improve the CRM
(segmentation, automation, retention flows, lifecycle marketing); own Klaviyo end to end — write campaign
and lifecycle email copy in the Sanne voice, plan the email calendar around launches, promotions and
seasonal moments, drive list growth, segmentation, testing and deliverability; own the creator affiliate
programme on ShopMy (commission structure, tracking links and codes, reporting on revenue and cost per
acquisition against the platform fee); manage the Shopify app and integration partners and the Shopify Plus
platform relationship (maintain the app stack, evaluate tools, track subscription costs, flag
recommendations to the CTO); work day to day with web development, email and digital marketing agencies;
work with the Head of Content to brief content supporting commercial priorities; build and refine the sales
funnel and back-end systems; track, report and act on site, CRM, email and affiliate performance data,
feeding regular updates to the CTO; create and use AI agents to design, automate and continuously improve
ecommerce processes, reducing manual work and producing faster reporting, with human approval built in for
anything touching live pricing or client data.
Looking for: proven ecommerce management / digital trading track record, ideally DTC or luxury retail;
hands-on Shopify Plus including apps, integrations and agency partners; strong CRM and marketing automation
knowledge, Klaviyo preferred; excellent written English and a feel for luxury brand voice; experience
running a creator, affiliate or influencer programme, ideally ShopMy; confident across systems, operations
and sales, hands-on rather than purely strategic; data-driven test-and-learn approach; experience briefing
content commercially; practical experience creating and using AI agents; based in Dubai.
Recruiter-added requirement: 8+ years of work experience in Digital Marketing.
"""

ATS = [
    "e-commerce", "ecommerce manager", "digital trading", "DTC", "luxury retail", "Shopify",
    "website sales performance", "conversion rate", "CRO", "average order value", "AOV", "checkout",
    "site operations", "product listings", "merchandising", "pricing", "site maintenance",
    "CRM", "marketing automation", "segmentation", "retention flows", "lifecycle marketing",
    "email marketing", "email copywriting", "campaign calendar", "list growth", "A/B testing",
    "deliverability", "creator affiliate programme", "affiliate", "influencer", "commission structure",
    "tracking links", "discount codes", "cost per acquisition", "CPA", "revenue reporting",
    "app stack", "integrations", "agency management", "content briefing", "sales funnel",
    "performance data", "reporting", "AI agents", "automation", "test-and-learn", "brand voice",
    "product launches", "international expansion", "GMV", "ROAS",
]

CONTENT = {
    "headline": "E-Commerce Manager · Site, CRM & Email, Creator Affiliates · AI-Automated",
    "professional_summary": (
        "Hands-on e-commerce manager owning the full engine — store, email/CRM and a live creator affiliate "
        "programme — and building AI agents behind the reporting. Fashion and beauty: Miravia (Alibaba, 42 "
        "accounts, +30% GMV QoQ), Glovo Retail, Inditex."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE consumer group, 6 brands | D2C store + Noon, Amazon.ae, talabat, Careem | 50+ markets",
            "bullets": [
                "Own the D2C store end-to-end — listings, merchandising, collections, pricing, bundles, promotions and checkout — lifting conversion rate and average order value through test-and-learn merchandising",
                "Plan, write and send the email and CRM channel: launch announcements, promotions and lifecycle newsletters across a six-brand portfolio, driving repeat purchase and list growth",
                "Designed and run Befit Crew, a live creator affiliate programme — unique codes, 8–10% tiered commission, audience discount, seeding and a self-serve registration funnel — tracked on revenue and acquisition cost",
                "Build AI agents automating campaign planning, content and KPI reporting (~40% less manual work, human approval kept on live pricing), run paid media on Meta and Google Ads, and brief four agencies plus a designer and social executive",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees",
            "bullets": [
                "Managed 42 brand accounts across beauty, fragrances and fashion — assortment, pricing, competitive benchmarking and promotions — delivering +30% GMV growth QoQ",
                "Owned the Flash Sales channel for Beauty, Fashion & Home reporting to the CEO — campaign calendar, item selection, P&L-aligned plans — and created the Beauty Club retention project",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Helped build the Retail vertical — onboarding fashion and lifestyle brands — and grew XL partner GMV through catalogue, promotions and bespoke activations",
            ],
        },
        {
            "company": "Mondelez International · Massimo Dutti (Inditex)",
            "role": "Category Planning Trainee · Sales Associate",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG category planning | Premium fashion retail floor",
            "bullets": [
                "Sell-in/sell-out and promotional-effectiveness analysis with Nielsen at Mondelez; Inditex grounding in fashion retail operations, visual merchandising and premium customer experience",
            ],
        },
    ],
    "skills_brand": (  # -> "Site & Trading"
        "Shopify ownership, listings, merchandising, pricing, bundles & promotions, CRO, AOV, checkout"
    ),
    "skills_ecommerce": (  # -> "CRM & Email"
        "email/CRM ownership, campaign & lifecycle copywriting, calendar planning, segmentation, list growth, A/B testing"
    ),
    "skills_commercial": (  # -> "Creators & Affiliates"
        "creator programme design, tiered commission, codes & tracking, seeding, CPA & revenue reporting"
    ),
    "skills_data": (  # -> "AI & Automation"
        "AI agents for workflow automation, automated KPI reporting, Claude / Claude Code, Python, human-in-the-loop controls"
    ),
    "skills_tools": (  # -> "Tools"
        "Shopify, Meta Ads Manager, Google Ads, Claude (AI), Power BI, Tableau, Looker, Salesforce, SAP, Excel (Expert)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Site & Trading",
    "E-Commerce & Digital": "CRM & Email",
    "Commercial": "Creators & Affiliates",
    "Data & Analytics": "AI & Automation",
}


def make_job() -> Job:
    return Job(
        id="sanne-ecommerce-manager-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE (On-site)",
        url="https://www.linkedin.com/jobs/search/?keywords=Sanne%20Ecommerce%20Manager%20Dubai",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Sanne Ecommerce Manager Dubai",
            "note": "Casa de moda de lujo made-on-demand, 11-50 empleados, HQ Dubai. Shopify Plus recién "
                    "lanzado. Encaje de contenido altísimo: los tres pilares del puesto (site, email/CRM, "
                    "programa de creadores) son literalmente lo que Paula lleva hoy en DoFreeze, y el JD pide "
                    "agentes de IA — su diferenciador. Brechas duras: requisito de 8+ años de digital marketing "
                    "(Paula ~5), Klaviyo y ShopMy por nombre. 946 solicitudes. Contacto: Lena McCroary, "
                    "Co-Founder/Creative Director, es quien publica la oferta y acepta mensajes.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 84, "ai_tier": "Hot",
        "skills_match": [
            "Los tres pilares del puesto son su día a día: tienda D2C end-to-end, canal de email/CRM que planifica y ESCRIBE, y programa de afiliados de creadores",
            "Befit Crew: programa de creadores diseñado por ella y en vivo — códigos únicos, comisión escalonada 8-10%, seeding y funnel self-serve",
            "Agentes de IA para automatizar planificación, contenido y reporting (~40% carga manual) — el JD lo pide explícitamente y es rarísimo en el mercado",
            "Shopify hands-on: listings, merchandising, pricing, bundles, CRO y AOV",
            "Moda y lujo por trayectoria: Miravia (beauty, fragancias y moda), Glovo Retail e Inditex",
            "Inglés escrito excelente y sensibilidad de marca — construye landings y campañas de concepto",
        ],
        "missing_skills": [
            "8+ años de digital marketing (requisito añadido por el anunciante) — Paula tiene ~5",
            "Klaviyo por nombre (su email/CRM ha corrido por otras herramientas)",
            "ShopMy por nombre (su programa de afiliados es propio, no en ShopMy)",
            "Shopify PLUS y gestión del app stack a nivel plataforma",
            "Apparel de lujo DTC en plantilla (grid de tallas Petite/Regular/Tall)",
        ],
        "sector_fit": "alto — moda y lujo, con e-commerce DTC como núcleo",
        "seniority_fit": "encaje funcional de Manager, pero el anunciante pide 8+ años y ella tiene ~5",
        "red_flags": [
            "946 solicitudes y 97 en un día: volumen brutal",
            "Requisito explícito de 8+ años de digital marketing — es filtro duro si lo aplican literal",
            "Empresa de 11-50 personas: menos estructura y probablemente menos banda salarial que el mínimo de 20K AED",
            "Presencial en HQ, no híbrido",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "De todo el pipeline, éste es el JD que más se parece a lo que Paula ya hace: el puesto es site + "
            "email/CRM + programa de creadores + agentes de IA, y ella lleva exactamente esos cuatro en DoFreeze, "
            "con el programa Befit Crew y el canal de email en vivo y documentados en su portfolio. La brecha no "
            "es de contenido sino de etiqueta y de años: piden Klaviyo y ShopMy por nombre y 8+ años de digital "
            "marketing frente a sus ~5. Merece la pena aplicar y, sobre todo, escribir a Lena McCroary "
            "(Co-Founder/Creative Director) enseñando el programa de creadores y el portfolio — el trabajo hace "
            "el argumento mejor que el CV."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
