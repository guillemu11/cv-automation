"""One-off: CV de una página para **Ecommerce Manager MENA** en **The Clorox Company** (Brita, Burt's Bees,
Clorox, Glad, Pine-Sol...). Dubái, híbrido. Oferta pegada desde LinkedIn: publicada hace ~1 semana, 549 clics en
"Solicitar" (21% nivel director, 50% senior). Se aplica en la web de Clorox (respuestas fuera de LinkedIn).

Qué es el puesto: líder regional de e-commerce. Dueño del P&L de e-commerce de MENA (ingresos, beneficio, cuota),
plan anual y estrategia plurianual, mix de canal (pure-play, quick commerce, bricks-and-clicks) — 30%. Relación con
los grandes e-retailers de MENA, JBP, negociaciones y renovaciones de contrato — 20%. Expansión de canal: nuevos
clientes, plataformas y canales emergentes, planes por mercado con equipos locales y distribuidores — 20%. Retail
media, digital shelf y presupuestos de inversión (eCommerce, ASP, retail media) — 20%. Equipo: el eCommerce Manager
de KSA y, si se aprueba, un nuevo eCommerce Operations Manager Gulf; capacidades de distribuidores — 10%.
Requisitos: **12+ años** de liderazgo comercial, mayoría en ventas de e-commerce, P&L de canal y negociación en una
multinacional FMCG/CPG. Grado. Inglés fluido; **árabe ventaja, no obligatorio**.

Encaje: **stretch claro de seniority** (5 años frente a 12+). Se presenta igualmente a petición de Paula; la carta
lo dice abiertamente y deja la puerta abierta al futuro eCommerce Operations Manager, Gulf que menciona el JD.

Ángulo honesto de Paula:
- **Mix de canal real**: quick commerce (talabat, Careem), pure play (Noon) y D2C (Shopify) desde el lado marca;
  Miravia (pure play) y Glovo (quick commerce) desde el lado plataforma.
- **Clientes e-commerce**: listings, precio, mecánicas promocionales, activaciones conjuntas (sorteo con cashback de
  talabat, samplings). SMASH × talabat +165% venta diaria vs +4,7% control; Befit × Noon +4.176 uds, +31%.
- **Comercial**: Miravia, 42 cuentas, +30% GMV QoQ; canal Flash Sales reportando al CEO contra objetivos de P&L;
  30+ tiendas nuevas en dos meses (lo más parecido a "onboard new customers"). Glovo: condiciones comerciales XL.
- **Exportación**: DoFreeze vende en 50+ mercados vía distribuidores (planes de trade/shopper por canal).
- **Disponibilidad**: revisión semanal de riesgo de stock por dark store; aporta al forecast mensual de sell-in.
- **Equipo**: lidera a dos (diseñadora + social media exec) y 4 agencias (negoció -30%).

Guardarraíles de honestidad:
- **P&L de e-commerce de una región**: no. En DoFreeze no es dueña del P&L; en Miravia es "against P&L targets".
- **Contratos y renovaciones con e-retailers, JBP desde el lado marca**: no se afirma; "joint activations".
- **Retail media de pago** con presupuesto propio: sin confirmar → "platform campaigns".
- **KSA / resto de MENA**: sin evidencia de cuentas fuera de EAU. Equipo de managers: no.
- Árabe no (aquí es ventaja). Deliveroo: DoFreeze NO está. Forecast: "contribute".
- Visa: "UAE Employment Visa (employer-sponsored)".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "The Clorox Company"
TITLE = "Ecommerce Manager MENA"
DATE_FOLDER = "2026-10-06"

JOB_DESCRIPTION = """\
Ecommerce Manager MENA. The Clorox Company. Dubai, UAE (Hybrid). Full-time. Posted ~1 week ago (LinkedIn, promoted by
recruiter, applications managed off LinkedIn). 549 clicked apply.
As the eCommerce Manager, MENA, you will lead the region's eCommerce agenda, driving commercial performance, customer
engagement and capability development across multiple digital channels. You will own the eCommerce P&L, shape growth
strategies and partner with key customers to unlock new opportunities across the MENA market.
eCommerce Strategy & P&L Ownership (30%): owns the full eCommerce P&L across MENA, delivering revenue, profit and market
share targets; develops and executes annual operating plans and multi-year growth strategies; drives channel mix
optimization across pure-play, quick commerce and bricks-and-clicks.
Customer Leadership & Joint Business Planning (20%): leads strategic relationships with top eCommerce retailers across
MENA; owns annual joint business plans, commercial negotiations and contract renewals; builds senior customer
partnerships.
Channel Expansion & Market Development (20%): defines and executes the eCommerce expansion roadmap across MENA;
identifies and onboards new customers, platforms and emerging channels; builds market-level plans with local teams and
distributors.
Retail Media, Digital Shelf & Investment Strategy (20%): owns eCommerce, ASP and retail media budgets across MENA; leads
annual budgeting and investment planning with retailers and agencies; drives retail media effectiveness, digital shelf
excellence and measurable ROI.
Team Leadership & Capability Building (10%): leads and develops the regional eCommerce team, including the eCommerce
Manager, KSA and, upon approval, a new eCommerce Operations Manager, Gulf; builds distributor and market eCommerce
capabilities; establishes a MENA community of practice.
What we look for: minimum 12+ years of commercial leadership experience, preferably majority in eCommerce Sales;
demonstrated channel P&L ownership, operational leadership and customer negotiation leadership within a multinational
FMCG/CPG environment. Minimum bachelor's degree. Fluent English; Arabic an advantage, not mandatory. Hybrid.
"""

ATS = [
    "eCommerce", "e-commerce", "eCommerce Sales", "commercial", "channel mix", "pure-play", "quick commerce",
    "bricks-and-clicks", "D2C", "marketplaces", "talabat", "Noon", "Careem", "Miravia", "Glovo", "Shopify",
    "eCommerce retailers", "key accounts", "customer partnerships", "joint activations", "growth plans",
    "commercial terms", "negotiation", "P&L targets", "revenue", "GMV", "market development", "onboarding",
    "new customers", "distributors", "50+ markets", "retail media", "platform campaigns", "digital shelf",
    "product content", "availability", "investment", "ROI", "ROAS", "budgets", "agencies", "team leadership",
    "sell-in", "sell-out", "forecast", "promotions", "pricing", "FMCG", "CPG", "multinational", "MENA", "UAE", "Dubai",
]

CONTENT = {
    "headline": "E-Commerce · Key Accounts · Quick Commerce & Marketplaces · FMCG · UAE",
    "professional_summary": (
        "E-commerce and key account professional with 5 years across FMCG, beauty and marketplaces, on both sides of "
        "the table. In Dubai, grows FMCG brands on talabat, Noon, Careem and Shopify D2C; before, ran 42 accounts at "
        "Miravia (Alibaba) at +30% GMV QoQ and its Flash Sales channel against P&L targets."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG group (Befit, Eurocake, Flair) | talabat, Noon & Careem + Shopify D2C | 50+ markets",
            "bullets": [
                "Manage three brands across quick commerce (talabat, Careem), pure play (Noon) and Shopify D2C — listings, pricing, promotional mechanics and joint activations with the platforms (purchase-linked talabat cashback giveaway, seasonal sampling)",
                "Launched SMASH, a new Eurocake brand, on talabat: daily sales +165% vs +4.7% for a control brand; Befit × Noon \"New Year, New Me\" added +4,176 incremental units, +31% over baseline; 6 product launches end-to-end",
                "Run the weekly stock-risk review by dark store with a dashboard built with AI, checking PO coverage against demand; contribute to the monthly sell-in forecast with e-commerce and finance",
                "Lead a team of two and 4 agencies (negotiated -30%); write channel trade and shopper plans for distributors in 50+ markets; built an AI workflow (Claude) that cut manual reporting and planning work ~40%",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees | platform side",
            "bullets": [
                "Managed 42 beauty, fragrance and fashion accounts (incl. KIKO Milano) on pricing, assortment, product content, campaign slots and promotions, delivering +30% GMV QoQ; onboarded 30+ new stores in two months as fragrance lead",
                "Owned the Flash Sales channel for Beauty, Fashion & Home reporting to the CEO — commercial plans against P&L targets, reading traffic, conversion, ROI and ROAS to reallocate promotional investment",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce & delivery platform | €500M+ revenue | platform side",
            "bullets": [
                "Managed XL accounts (KFC, Taco Bell, La Tagliatella, Sushi Shop) with data-led GMV growth plans and in-app activations; negotiated commercial terms and aligned marketing, logistics and customer support on execution",
            ],
        },
        {
            "company": "Mondelez International · Massimo Dutti (Inditex)",
            "role": "Trainee, Category Planning · Sales Associate",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG (chocolate category) | Premium fashion retail",
            "bullets": [
                "At Mondelez, sell-in/sell-out and promotional-effectiveness analysis for chocolate, contributing to NPD launches (Milka Spread, Mini Suchard); at Inditex, store operations and customer service",
            ],
        },
    ],
    "skills_brand": (  # -> "Key Accounts"
        "talabat, Noon, Careem (brand side) · Miravia, Glovo (platform side) · negotiation, onboarding"
    ),
    "skills_ecommerce": (  # -> "Channels"
        "quick commerce, pure play, Shopify D2C · product content, availability & stock risk, launches"
    ),
    "skills_commercial": (  # -> "Commercial"
        "growth plans, P&L targets, pricing, promotions, budgets, agency negotiation, forecast input"
    ),
    "skills_data": (  # -> "Performance"
        "platform campaigns, sell-in/sell-out, control-group reads · GMV, ROI, ROAS, conversion"
    ),
    "skills_tools": (  # -> "Tools"
        "Excel (Advanced), Power BI, Tableau, Looker, Salesforce, SAP, Shopify, Claude (AI dashboards)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Key Accounts",
    "E-Commerce & Digital": "Channels",
    "Commercial": "Commercial",
    "Data & Analytics": "Performance",
}


def make_job() -> Job:
    return Job(
        id="clorox-ecommerce-manager-mena-dubai-2026-10",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE",
        url="https://www.linkedin.com/jobs/search/?keywords=Clorox%20Ecommerce%20Manager%20MENA",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={
            "query": "The Clorox Company Ecommerce Manager MENA Dubai",
            "note": "Líder regional de e-commerce MENA: P&L, JBP con e-retailers, expansión de canal, retail media, "
                    "equipo (eCommerce Manager KSA). Pide 12+ años. Stretch de seniority; se aplica a petición de "
                    "Paula. Árabe ventaja, no obligatorio.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": None, "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": "2026-09-29", "raw": job.raw,
        "ai_score": 45, "ai_tier": "Stretch",
        "skills_match": [
            "Mix de canal real: quick commerce (talabat, Careem), pure play (Noon) y Shopify D2C",
            "Cuentas e-commerce desde los dos lados: marca (DoFreeze) y plataforma (Miravia, Glovo)",
            "Comercial: 42 cuentas en Miravia, +30% GMV QoQ; Flash Sales contra objetivos de P&L",
            "Activaciones con resultado medido: SMASH × talabat +165% vs control, Befit × Noon +31%",
            "Captación: 30+ tiendas nuevas en dos meses en Miravia",
            "Distribuidores en 50+ mercados; equipo de dos y 4 agencias",
            "FMCG (DoFreeze, Mondelez)",
        ],
        "missing_skills": [
            "12+ años de liderazgo comercial (tiene 5)",
            "P&L de e-commerce de una región como dueña",
            "JBP, negociación y renovación de contratos con e-retailers desde el lado marca",
            "Presupuestos de retail media propios",
            "Mercados MENA fuera de EAU (KSA)",
            "Gestionar managers (el puesto lleva al eCommerce Manager de KSA)",
            "Multinacional FMCG/CPG en el puesto actual",
        ],
        "sector_fit": "bueno — FMCG/CPG y e-commerce en EAU",
        "seniority_fit": "muy por encima — 12+ años y liderazgo regional frente a 5 años",
        "red_flags": [
            "Pide 12+ años; Paula tiene 5",
            "549 clics en solicitar, 21% nivel director",
            "Dueño del P&L de MENA y jefe de otros managers",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "El trabajo de fondo encaja (e-commerce de FMCG en EAU, cuentas desde los dos lados, mix de quick commerce, "
            "pure play y D2C), pero el nivel no: pide 12+ años, P&L regional y gestionar managers. Se aplica a petición "
            "de Paula, con una carta que reconoce el salto y apunta al futuro eCommerce Operations Manager, Gulf."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
