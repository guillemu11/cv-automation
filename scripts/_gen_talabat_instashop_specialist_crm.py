"""One-off: CV de una página para **Specialist CRM – instashop** (talabat / Delivery Hero), Dubái,
presencial, individual contributor. LinkedIn Easy Apply, 293 solicitudes.

instashop es el marketplace local líder en EAU y Egipto, parte de talabat/Delivery Hero desde 2020.

Pide el JD: ser dueño del motor de retención y engagement; diseñar el customer journey end-to-end
con triggers automatizados (onboarding, activación, retención, win-back, loyalty); **ownership de
Braze** como super-user y product owner interno (push, in-app, email, SMS); segmentación avanzada
por comportamiento e historial de compra para subir LTV; higiene de base de datos y cumplimiento
(GDPR/CCPA) con producto e ingeniería; A/B testing continuo de subject lines, send-times, contenido
y canal; colaboración con Brand, Commercial y Performance para lanzamientos de categoría y
promociones localizadas; reporting semanal y mensual de retention rate, churn, frecuencia y repeat
purchase rate. Requiere 4+ años en CRM y marketing automation, preferiblemente con Braze, y
familiaridad con Salesforce / HubSpot / Marketo.

Ángulo honesto de Paula:
- Quick-commerce por los dos lados: Glovo (dentro de la plataforma, cuentas XL) y DoFreeze
  (marca vendiendo en talabat, Noon y Careem). Conoce el negocio de instashop de primera mano.
- Retención y recompra como objetivo, no awareness: campañas medidas en venta incremental con
  grupo de control, sorteo de compra-participa con talabat, y el programa de afiliación.
- Miravia: creó 'Beauty Club' — programa de fidelidad y recurrencia — y analizó retención,
  conversión y tráfico como parte del día a día del canal.
- Shopify DTC propio: base de clientes, segmentación, EDM y CRO.
- A/B testing real en creatividades y audiencias de Meta Ads, con lectura de ROAS.
- Salesforce en su stack (Miravia/Glovo), que el JD acepta como alternativa a Braze.

Guardarraíles de honestidad (IMPORTANTES en esta vacante):
- **Braze: NO lo ha usado. No aparece en el CV bajo ningún concepto.** Tampoco HubSpot ni Marketo.
  Solo se lista Salesforce, que sí está en su stack real.
- **No es una CRM specialist.** No tiene 4+ años de CRM dedicado ni ha sido dueña de un journey
  de lifecycle con triggers automatizados en una plataforma de CRM. Lo más cercano es EDM,
  flujos de Shopify, audiencias de Meta y el programa de fidelidad de Miravia. El CV lo presenta
  como lo que es: marketing de ciclo de vida y recurrencia desde el lado de marca y comercial.
- **GDPR/CCPA**: no se reclama expertise de compliance; se dice exposición a datos de cliente en
  un marketplace europeo, que es lo real.
- **Push / in-app / SMS**: no los ha operado. No se mencionan como canales propios.
- Sin árabe. Visa solo "UAE Residence Visa"; nunca "no sponsorship needed".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "talabat (instashop)"
TITLE = "Specialist CRM - instashop"
DATE_FOLDER = "2026-09-23"

JOB_DESCRIPTION = """\
Specialist CRM — instashop (talabat / Delivery Hero). Dubai, UAE. On-site, full time, individual contributor.
instashop is the leading online local marketplace in the UAE and Egypt, part of the Delivery Hero family since 2020.
As Specialist CRM you will own the retention and engagement engine of instashop: design, build and scale
lifecycle marketing journeys to maximise user retention, reactivate lapsed customers and build brand loyalty,
using Braze to turn customer data into personalised multi-channel communication.
Responsibilities: own the end-to-end customer journey map with automated triggers for onboarding, activation,
retention, win-back and loyalty; act as internal super-user and product owner for Braze across push
notifications, in-app messages, email and SMS; analyse user behaviour and purchase history to build dynamic
segments driving LTV; oversee database hygiene and regional data compliance with product and engineering;
run continuous A/B tests on subject lines, send-times, content and channel to improve open rates, CTRs and
conversion; partner with Brand, Commercial and Performance teams on category launches, localised promotions
and co-marketing; deliver weekly and monthly reporting on retention rate, churn rate, frequency and repeat
purchase rate.
Qualifications: Bachelor's degree in Marketing or Business Administration; 4+ years in CRM and marketing
automation, preferably with Braze, ideally in e-commerce or technology; experience with CRM and automation
platforms such as Braze, Salesforce, HubSpot or Marketo; familiarity with GDPR/CCPA; budget management;
fluent English. Strong analytics, A/B testing methodology, creative eye and cross-functional collaboration.
"""

ATS = [
    "CRM", "lifecycle marketing", "customer journey", "retention", "churn", "win-back", "reactivation",
    "onboarding", "activation", "loyalty", "LTV", "repeat purchase rate", "frequency", "segmentation",
    "dynamic segments", "personalisation", "marketing automation", "automated triggers", "email", "EDM",
    "push notifications", "multi-channel", "A/B testing", "open rate", "CTR", "conversion rate", "CRO",
    "customer data", "purchase history", "user behaviour", "data analytics", "campaign reporting",
    "KPI reporting", "Salesforce", "Shopify", "Meta Ads", "Google Ads", "audiences", "e-commerce",
    "quick-commerce", "marketplace", "talabat", "Noon", "Careem", "Glovo", "UAE", "MENA", "GDPR",
    "customer database", "cross-functional", "brand", "commercial", "performance marketing",
    "category launches", "localised promotions", "co-marketing", "budget management", "ROI", "ROAS", "GMV",
]

CONTENT = {
    "headline": "Lifecycle & Retention Marketing · E-Commerce & Quick-Commerce · UAE",
    "professional_summary": (
        "Commercial marketer who builds repeat purchase, not reach: loyalty and recurrence programmes "
        "at an Alibaba marketplace, owned Shopify DTC customer base and EDM, and quick-commerce "
        "campaigns measured on incremental sell-out with control groups. Knows the instashop business "
        "from both sides — inside a q-commerce platform at Glovo, and as a brand selling on talabat, "
        "Noon and Careem."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE consumer brands (Befit, SMASH, Eurocake) | Shopify DTC + talabat, Noon, Careem",
            "bullets": [
                "Own the Shopify DTC store end-to-end — customer base, segmentation, EDM campaigns, collections and checkout — lifting conversion rate and average order value through data-led merchandising",
                "Run campaigns to repeat purchase, not awareness, and prove it: SMASH × talabat lifted daily revenue AED 355 → 940 (+165%) against a control brand (+4.7%), holding at double after the campaign; Befit × Noon delivered +4,176 incremental units, +31% over baseline",
                "Designed a purchase-to-enter cashback giveaway with talabat and a seasonal seeding calendar to drive in-app frequency, plus an affiliate programme with sales-based commissions",
                "A/B test creative, audiences and mechanics on Meta and Google Ads, reading ROAS and conversion to reallocate budget; built AI-automated KPI reporting that cut manual reporting workload ~40%",
                "Partner with commercial, sales and platform teams on category launches and localised promotional mechanics across GCC and 50+ markets; lead a designer and a social media executive",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees | Salesforce",
            "bullets": [
                "Created and led 'Beauty Club' and 'Hot on Social', loyalty and engagement programmes driving retention and repeat purchase; analysed conversion, traffic, retention and ROI to optimise the channel, managing 42 accounts to +30% GMV QoQ",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce platform | €500M+ revenue",
            "bullets": [
                "Worked platform-side on XL accounts: in-app promotional mechanics, bespoke activations and order-frequency growth",
            ],
        },
        {
            "company": "Mondelez International · Massimo Dutti (Inditex)",
            "role": "Trainee, Category Planning · Sales Associate",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | Premium fashion retail",
            "bullets": [
                "Sell-in/sell-out and promotional-effectiveness analysis with Nielsen, building category performance reports",
            ],
        },
    ],
    "skills_brand": (  # -> "Lifecycle & Retention"
        "customer journey thinking, loyalty & recurrence programmes, win-back & reactivation mechanics, EDM campaigns, personalised content, promotional mechanics"
    ),
    "skills_ecommerce": (  # -> "E-Commerce & Platforms"
        "Shopify DTC, CRO, customer database & segmentation, quick-commerce (talabat, Noon, Careem, Glovo), marketplace operations, Meta & Google Ads audiences"
    ),
    "skills_commercial": (  # -> "Testing & Optimisation"
        "A/B testing (creative, audience, mechanics), control-group measurement, send & offer optimisation, budget management, incremental sell-out"
    ),
    "skills_data": (  # -> "Data & Reporting"
        "repeat purchase & frequency analysis, conversion & traffic, ROI/ROAS, weekly & monthly KPI reporting, AI-automated analysis, Power BI, Tableau, Looker"
    ),
    "skills_tools": (  # -> "Tools"
        "Shopify, Salesforce, SAP, Meta Business Suite, Meta Ads Manager, Google Ads, Power BI, Tableau, Excel (Advanced), Claude (AI agents)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Lifecycle & Retention",
    "E-Commerce & Digital": "E-Commerce & Platforms",
    "Commercial": "Testing & Optimisation",
    "Data & Analytics": "Data & Reporting",
}


def make_job() -> Job:
    return Job(
        id="talabat-instashop-specialist-crm-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE (On-site)",
        url="https://www.linkedin.com/jobs/search/?keywords=talabat%20Specialist%20CRM%20instashop%20Dubai",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "talabat Specialist CRM instashop Dubai",
            "note": "instashop (marketplace local de EAU y Egipto, dentro de talabat/Delivery Hero). Rol de "
                    "individual contributor, dueño del motor de retención. LinkedIn marca 'idoneidad alta' "
                    "y Paula tiene 1 contacto dentro. PERO el JD pide 4+ años de CRM con Braze y ella NO "
                    "ha trabajado Braze ni ha tenido un rol de CRM dedicado: es su brecha más seria de "
                    "todas las vacantes recientes. Su baza es el conocimiento del negocio de q-commerce "
                    "por los dos lados (Glovo dentro, DoFreeze como marca en talabat/Noon/Careem) y el "
                    "haber medido campañas en recompra incremental con grupo de control.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 62, "ai_tier": "Warm",
        "skills_match": [
            "Quick-commerce por los dos lados: Glovo desde dentro de la plataforma, DoFreeze como marca en talabat/Noon/Careem",
            "Shopify DTC propio: base de clientes, segmentación, EDM y CRO",
            "Miravia 'Beauty Club': programa de fidelidad y recurrencia creado por ella",
            "Campañas medidas en recompra y venta incremental con grupo de control, no en awareness",
            "A/B testing real en creatividades y audiencias con lectura de ROAS",
            "Salesforce en su stack, que el JD acepta como alternativa a Braze",
            "Reporting semanal/mensual de KPIs, ahora automatizado con IA",
            "Inglés C1 y colaboración cross-funcional con brand, comercial y performance",
        ],
        "missing_skills": [
            "BRAZE — no lo ha usado nunca. Es el requisito central del puesto",
            "4+ años de CRM dedicado: no tiene un rol de CRM, sino lifecycle desde marca y comercial",
            "Journeys automatizados con triggers en una plataforma de CRM (onboarding, win-back)",
            "Push, in-app y SMS como canales operados por ella",
            "GDPR/CCPA como responsabilidad de compliance",
            "Sin árabe",
        ],
        "sector_fit": "muy alto — q-commerce en EAU es literalmente su ecosistema diario",
        "seniority_fit": "bueno — Specialist individual contributor; podría incluso quedarse corto de banda (verificar 20K AED)",
        "red_flags": [
            "El requisito de Braze es explícito y ella no lo cumple: filtro probable del ATS o del recruiter",
            "293 solicitudes",
            "Specialist (no Manager): riesgo de banda por debajo del suelo de 20.000 AED/mes",
            "Presencial",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Encaje de sector inmejorable y encaje de función flojo. instashop es q-commerce en EAU, que es "
            "el ecosistema donde Paula se mueve todos los días desde los dos lados del tablero, y su forma de "
            "trabajar — campañas atadas a recompra, grupo de control, lectura de conversión y frecuencia — es "
            "exactamente la mentalidad que pide el rol. Pero el puesto es CRM puro con Braze como columna "
            "vertebral, y ahí no hay nada que presentar: ni Braze, ni journeys con triggers, ni push/in-app/SMS. "
            "El CV no inventa ninguna de esas cosas. La candidatura tiene sentido como apuesta lateral, no como "
            "favorita, y la vía realista es el contacto de primer grado que ya tiene dentro de talabat: que "
            "alguien lea el perfil como 'conoce el negocio y aprende la herramienta' en vez de dejar que el "
            "filtro de Braze lo descarte. Conviene decidir si merece la pena frente a vacantes donde sí lidera."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
