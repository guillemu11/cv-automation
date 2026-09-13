"""One-off: generate Paula's ONE-PAGE CV for **Media Strategy & Planning (Budget & Agency Management)**
at **Piaget** (Richemont luxury watches & jewellery; MEIA regional team, Dubai).

JD ask: support annual MEIA media strategy aligned with global/regional brand objectives; campaign objectives,
audiences, channels, KPIs, investment plans; monitor media budgets across social, digital, print, OOH; track
spend/performance for ROI; lead media-agency briefing with media HQ (objectives, audiences, messages, budgets,
timelines, learnings); review and challenge agency plans and reports; asset adaptation, local content
production, approvals, brand guidelines; deploy campaigns; HQ/agency/regional coordination; 360 amplification
of brand campaigns, High Jewellery events, boutique openings, launches, cultural moments; consolidate
performance, insights, trends/competitors, updates to regional and HQ stakeholders.

Paula's honest angle: A&P budget + paid social/search (Meta, Google) across 50+ markets; seasonal 360 campaigns
tied to cultural moments (Ramadan, New Year, Fitness Month, back to school); manages 4 agencies, negotiated -30%
and audited agency reporting (reach estimated from followers, double counting); briefs and reviews all creative
with her designer; 6 launches; localised assets (digital + packaging/POS); measured incrementality with a control
group (SMASH x talabat +165% vs +4.7%); luxury/beauty exposure via Miravia fragrances (Arabian oud houses) and
Massimo Dutti retail; HQ-style reporting to Alibaba CEO.

Honesty guardrails:
- NO print/OOH media buying, NO media-agency (planning/buying) retainer claim, NO luxury watches/jewellery sector claim.
- NO event/boutique-opening ownership. Spend figures not disclosed. Forecast not owned.
- Factual "UAE Residence Visa" only.

ONE-PAGE standard. Lands under output/2026-09-13/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Piaget"
TITLE = "Media Strategy & Planning Manager"
DATE_FOLDER = "2026-09-13"

JOB_DESCRIPTION = """\
Piaget (Richemont) — Media Strategy & Planning, Budget & Agency Management, MEIA region.

Media Strategy & Planning, Budget & Agency Management: support development and execution of annual media strategies
across MEIA aligned with global and regional brand objectives; contribute to campaign objectives, target audiences,
channel strategies, KPIs and investment plans; monitor media budgets across social, digital, print and out-of-home;
track spending and campaign performance to optimise and maximise ROI; coordinate agency partners and timely delivery
of campaign assets.

Campaign Development & Execution: lead briefing with the media agency with the media HQ team for annual planning and
campaign activations (objectives, audiences, key messages, budgets, timelines, learnings); review and challenge agency
recommendations, media plans, content proposals and performance reports; oversee asset adaptation, local content
production, approvals and implementation consistent with Piaget brand guidelines; coordinate deployment across social,
digital, print and OOH; collaboration between HQ teams, agencies, regional stakeholders and media partners.

360 communication oversight: support head of communication in 360 amplification of brand campaigns, High Jewellery
events, boutique openings, product launches, events and cultural moments.

Performance Analysis & Market Intelligence: consolidate and analyse campaign performance across channels; translate
insights into recommendations; monitor industry trends, competitors and media best practices; share regular
performance updates with regional and HQ stakeholders.
"""

ATS = [
    "media strategy", "media planning", "annual media plan", "MEIA", "media budget", "investment plan", "ROI",
    "KPIs", "target audiences", "channel strategy", "social media", "digital media", "paid social", "print",
    "out-of-home", "OOH", "media agency", "agency briefing", "agency management", "media plans",
    "performance reports", "asset adaptation", "localisation", "local content production", "brand guidelines",
    "approvals", "campaign deployment", "360 campaigns", "product launches", "events", "cultural moments",
    "Ramadan", "HQ stakeholders", "regional stakeholders", "performance analysis", "insights",
    "competitor analysis", "luxury", "Meta Ads", "Google Ads",
]

CONTENT = {
    "headline": "Media Planning · Budget & Agency Management · 360° Campaigns",
    "professional_summary": (
        "Brand and marketing manager with 5 years across FMCG, beauty and eCommerce (Alibaba, Mondelez) in Europe and the UAE: "
        "plans media and budgets across 50+ markets, briefs and challenges agencies, and reports results with measured ROI."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG group (Befit, Eurocake, Smash) | GCC, MENA, Asia, Europe, Africa | 50+ countries",
            "bullets": [
                "Plan and manage the A&P budget and paid media (Meta, Google, creators) by market: objectives, audiences, channel mix and KPIs, optimising spend against ROI and ROAS",
                "Brief and manage 4 agencies; negotiated fees -30% and audited their performance reports, catching reach estimated from followers and double-counted audiences",
                "Run year-round 360° campaigns tied to cultural moments (Ramadan, New Year, back to school) and 6 product launches; brief, review and approve all creative and localised assets with the in-house designer",
                "Report incremental results, not vanity metrics: SMASH x talabat +165% daily sales vs +4.7% for a control brand",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's eCommerce marketplace | 100K+ employees",
            "bullets": [
                "Managed 42 beauty, fragrance and fashion brands (+30% GMV QoQ), incl. Arabian Oud, Lattafa and Ajmal; created the Beauty Club and Hot on Social visibility projects",
                "Owned the Flash Sales channel, consolidating performance and recommendations for the CEO",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Delivery & quick-commerce app | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Delivered co-marketing activations for XL partners (KFC, Taco Bell, Sushi Shop), coordinating marketing, logistics and support",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | Chocolate category (Milka, Suchard) | €36B revenue",
            "bullets": [
                "Evaluated promotional effectiveness with Nielsen and supported the Milka Spread and Mini Suchard launches",
            ],
        },
    ],
    "skills_brand": (  # label -> "Media & Budget"
        "media planning, A&P budget tracking, channel mix, KPIs, ROI, paid social & search"
    ),
    "skills_ecommerce": (  # label -> "Campaigns"
        "360° campaigns, cultural moments, launches, creators, asset adaptation, localisation"
    ),
    "skills_commercial": (  # label -> "Agencies"
        "agency briefs, challenging plans & reports, fee negotiation, approvals, brand guidelines"
    ),
    "skills_data": (  # label -> "Insights"
        "performance consolidation, control groups, competitor & trend analysis, CEO reporting"
    ),
    "skills_tools": (  # label -> "Tools"
        "Meta Ads Manager, Google Ads, Looker, Power BI, Nielsen, Excel, Canva, Adobe, Claude"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Media & Budget",
    "E-Commerce & Digital": "Campaigns",
    "Commercial": "Agencies",
    "Data & Analytics": "Insights",
}


def make_job() -> Job:
    return Job(
        id="piaget-media-strategy-planning-meia-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE (MEIA regional)",
        url="https://www.linkedin.com/jobs/search/?keywords=Piaget%20Media%20Strategy%20Planning",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Piaget Media Strategy & Planning MEIA",
            "note": "Fit on budget + agency management + 360 campaigns + reporting; gaps: luxury watches/jewellery, "
                    "print/OOH buying, media-agency (planning/buying) relationship, HQ matrix in luxury.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 64, "ai_tier": "Warm",
        "skills_match": [
            "A&P budget and paid media (Meta, Google, creators) across 50+ markets",
            "Manages 4 agencies; negotiated -30% and audited agency performance reports",
            "Seasonal 360 campaigns on cultural moments (Ramadan, New Year) and 6 launches",
            "Creative review/approval and asset localisation with her designer",
            "Incrementality measurement with control group; CEO-level reporting at Miravia",
        ],
        "missing_skills": [
            "Luxury watches & jewellery sector (closest: Miravia fragrances, Massimo Dutti)",
            "Print and OOH media planning/buying",
            "Working with a media planning/buying agency and a global media HQ",
            "Events / boutique-opening amplification",
        ],
        "sector_fit": "adjacent — FMCG/beauty/eCommerce vs luxury maison",
        "seniority_fit": "good — support/executive-to-manager level media role",
        "red_flags": ["Luxury houses often prefer luxury background", "Arabic may be requested for MEIA (not stated)"],
        "ats_keywords": ATS,
        "reasoning": (
            "Budget, agency briefing/challenging, 360 campaigns and performance reporting all map to her DoFreeze work, and "
            "the JD doesn't ask for Arabic. Main gaps are luxury sector and traditional media (print/OOH)."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
