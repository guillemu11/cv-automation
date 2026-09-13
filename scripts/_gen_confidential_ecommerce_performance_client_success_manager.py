"""One-off: generate Paula's ONE-PAGE CV for **Ecommerce Performance & Client Success Manager**
(company not named in the JD — portfolio operator/agency running 6-8 ecommerce brands across Europe,
USA & UAE; inventory-held + dropship; team of 5-6; min. 5 years).

JD ask: own P&L across 6-8 brands (revenue, gross margin, EBITDA), budgets, forecasts, targets; CVR, AOV,
CAC, ROAS/MER, COGS; Meta + Google Ads as primary channels (review accounts, challenge agencies, scale
spend, testing creative/audiences/feeds/landing pages); day-to-day trading (calendars, promotions,
pricing, launches, merchandising, CRO roadmap); client success (primary commercial contact, performance
meetings, retention); lead 5-6 people; oversee inventory/dropship suppliers, logistics, CS. Shopify
highly desirable; Europe/USA/ME advantageous; ChatGPT/Claude beneficial.

Paula's honest angle: in-house owner of a Shopify D2C store fed by Meta + Google Ads (creative A/B tests,
ROAS/ROI); trading calendar, promotions and 6 NPD launches across web + talabat/Noon/Careem; weekly
dark-store stock-risk review (inventory-held model); agency management (4 agencies, audited reports);
client-facing commercial roles at Miravia (42 brands, +30% GMV QoQ, Flash Sales channel reporting to CEO,
P&L-aligned targets) and Glovo XL accounts (ops issues with logistics + CS); markets: Spain/Europe + UAE,
brands sold in 50+ countries; heavy Claude/AI-agent user.

Honesty guardrails:
- NO EBITDA ownership, NO multi-brand P&L ownership claim, NO spend figures, NO dropship experience.
- Team = two direct reports (not 5-6). Forecast not owned ("contribute"). No MER/LTV/CAC ownership claim.
- ~5 yrs incl. Mondelez trainee year. Factual "UAE Residence Visa" only.

ONE-PAGE standard. Lands under output/2026-09-13/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Confidential"
TITLE = "Ecommerce Performance & Client Success Manager"
DATE_FOLDER = "2026-09-13"

JOB_DESCRIPTION = """\
Ecommerce Performance & Client Success Manager. Markets: Europe, USA & UAE. Portfolio: 6-8 ecommerce brands.
Team: 5-6 people. Business models: inventory-held & dropship. Experience: min. 5 years.

Commercially driven Ecommerce Performance & Trading Manager to take end-to-end ownership of a portfolio of 6-8
ecommerce brands, with direct P&L responsibility (revenue, profitability, growth), Meta Ads and Google Ads as
primary acquisition channels. Lead a team of 5-6, manage client relationships, coordinate internal specialists,
agencies and partners across trading, performance marketing, growth, development, operations, logistics, CS.

Commercial & P&L: own P&L (revenue, gross margin, EBITDA); budgets, forecasts, targets; CVR, AOV, CAC, ROAS/MER,
COGS, marketing spend; risks and growth opportunities; inventory-held vs dropship requirements.
Meta & Google Ads: own performance with in-house specialists and agencies; budgets vs revenue/CAC/ROAS/MER;
review accounts and challenge strategy; scale spend profitably; testing across creative, audiences, campaigns,
product feeds, landing pages; align paid media with trading and P&L.
Ecommerce Trading & Growth: day-to-day trading; trading calendars, promotions, pricing, launches, onsite
merchandising; analyse channels/products/markets/customers; conversion, AOV, acquisition, profitability; own the
website and CRO roadmap with development.
Client Success: primary commercial contact; relationships; regular performance meetings (P&L, trading, paid media,
forecasts); proactive communication; satisfaction and retention.
Team Leadership: lead and develop 5-6 people; coach; coordinate internal teams, agencies, partners.
Operations & CX: inventory, dropship suppliers, logistics, fulfilment, customer service; stock availability,
delivery, returns, cancellations.
Skills: ecommerce trading/performance marketing/growth; ecommerce P&L; Meta Ads and Google Ads essential; paid
media budgets across brands/markets; CVR, AOV, CAC, ROAS/MER, LTV, margin, EBITDA; client management; inventory-held
and/or dropship; Shopify highly desirable; CRO, analytics, web dev, logistics, CX; Europe/USA/Middle East
advantageous; AI tools such as ChatGPT and Claude beneficial.
"""

ATS = [
    "ecommerce", "P&L", "revenue", "gross margin", "budgets", "forecasts", "commercial targets", "CVR", "AOV",
    "CAC", "ROAS", "MER", "COGS", "Meta Ads", "Google Ads", "paid media", "agencies", "creative testing",
    "audiences", "landing pages", "trading", "trading calendar", "promotions", "pricing", "product launches",
    "merchandising", "CRO", "Shopify", "client success", "client relationships", "performance meetings",
    "retention", "team leadership", "stakeholder management", "inventory", "stock availability", "logistics",
    "customer service", "Europe", "UAE", "ChatGPT", "Claude", "multi-brand",
]

CONTENT = {
    "headline": "Ecommerce Performance · Meta & Google Ads · Client Success",
    "professional_summary": (
        "Ecommerce and commercial manager with 5 years across Alibaba, Glovo and FMCG in Europe and the UAE: runs a Shopify "
        "store fed by Meta and Google Ads, the trading calendar and agencies, and has grown 42 client brands (+30% GMV QoQ)."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG group (Befit, Eurocake, Smash) | Shopify D2C + quick-commerce | sold in 50+ countries",
            "bullets": [
                "Run Meta and Google Ads into the Shopify store: audiences, creative A/B tests and landing pages, optimised against ROAS, ROI and AOV; own the store end-to-end (catalogue, pricing, discounts, checkout, CRO)",
                "Own the trading calendar across web, talabat, Noon and Careem: promotions, pricing and 6 NPD launches; manage the A&P budget and contribute to the monthly forecast with Sales and Finance",
                "Run a weekly stock-risk review by dark store with an AI-built dashboard to protect availability; manage 4 agencies and audit their reporting, challenging results with data",
                "Lead a team of two (designer + social media executive); built Claude AI agents that cut manual reporting and planning work ~40%",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's eCommerce marketplace | 100K+ employees",
            "bullets": [
                "Primary commercial contact for 42 brands: regular performance reviews on pricing, assortment and promotions for +30% GMV QoQ, tracking ROAS, conversion and retention",
                "Owned the Flash Sales channel against P&L-aligned targets, reporting to the CEO; onboarded 30+ fragrance stores in two months",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Delivery & quick-commerce app | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Managed XL partners (KFC, Taco Bell, Sushi Shop), coordinating marketing, logistics and customer support to grow orders",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | Chocolate category (Milka, Suchard) | €36B revenue",
            "bullets": [
                "Built promotional-effectiveness and sell-in/sell-out reports with Nielsen for the chocolate category",
            ],
        },
    ],
    "skills_brand": (  # label -> "Performance"
        "Meta Ads, Google Ads, creative & audience testing, landing pages, agency management"
    ),
    "skills_ecommerce": (  # label -> "Trading"
        "Shopify, CRO, AOV, pricing, promotions, trading calendar, launches, marketplaces"
    ),
    "skills_commercial": (  # label -> "Client & Team"
        "client performance reviews, CEO reporting, negotiation, team of two, cross-functional"
    ),
    "skills_data": (  # label -> "Commercial KPIs"
        "ROAS / ROI, CVR, AOV, GMV, A&P budget, stock availability, sell-in / sell-out"
    ),
    "skills_tools": (  # label -> "Tools"
        "Meta Ads Manager, Google Ads, Shopify, Excel, Looker, Power BI, Claude, ChatGPT"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Performance",
    "E-Commerce & Digital": "Trading",
    "Commercial": "Client & Team",
    "Data & Analytics": "Commercial KPIs",
}


def make_job() -> Job:
    return Job(
        id="confidential-ecommerce-performance-client-success-manager-2026-09",
        title=TITLE,
        company=COMPANY,
        location="UAE (portfolio across Europe, USA & UAE)",
        url="https://www.linkedin.com/jobs/search/?keywords=Ecommerce%20Performance%20%26%20Client%20Success%20Manager",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Ecommerce Performance & Client Success Manager",
            "note": "Stretch: Shopify + Meta/Google + trading + client-facing KAM fit; gaps are multi-brand P&L/EBITDA, "
                    "team of 5-6 (she leads 2), dropship, paid-media spend scale.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 62, "ai_tier": "Warm",
        "skills_match": [
            "Runs Meta + Google Ads into a Shopify store she owns end-to-end (CRO, AOV, checkout)",
            "Trading calendar, promotions, pricing and 6 NPD launches across web + quick-commerce",
            "Client-facing commercial account management for 42 brands at Miravia (+30% GMV QoQ, CEO reporting)",
            "Agency management and audit; weekly stock-availability review",
            "Europe + UAE experience; heavy Claude/ChatGPT user",
        ],
        "missing_skills": [
            "Direct multi-brand P&L ownership incl. gross margin and EBITDA",
            "Leading a team of 5-6 (she leads two)",
            "Dropship model and supplier/fulfilment management",
            "Large paid-media budgets across multiple brands/markets; CAC/MER/LTV ownership",
        ],
        "sector_fit": "good — ecommerce D2C and marketplace brands",
        "seniority_fit": "stretch — Head-of-portfolio scope vs her single-brand Manager scope",
        "red_flags": ["Company not named in the JD", "USA market coverage (time zones)"],
        "ats_keywords": ATS,
        "reasoning": (
            "Strong overlap on Shopify, Meta/Google, trading and client success, no Arabic required. The P&L/EBITDA "
            "ownership, team size and dropship ask are above her current scope — worth applying, expect screening on those."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
