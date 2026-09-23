"""One-off: CV de una página para **E-Commerce Manager, Lacoste & Guess** en
**Chalhoub Group** (Dubai, híbrido).

Pide el JD: estrategia de e-commerce y presupuesto anual junto al eCommerce GM y al Brand GM;
trading (producto, precio, promociones) con sinergias entre Retail, eCommerce y e-Retail;
**ownership del P&L online** y optimización de cost drivers; gestión de relaciones internas,
agencias terceras y equipos de e-Retail; **online merchandising** (surtido, presentación,
categorización, tagging, search, contenido de producto, promos y campañas); **performance
marketing y web analytics** (PPC, display, social, afiliación; optimización de inversión,
creatividades y targeting; KPIs vía dashboards de e-commerce); product management y roadmap;
liderazgo cross-funcional en iniciativas omnicanal con equipos centrales, de producto y Retail.

Ángulo honesto de Paula: Miravia (Alibaba) — KAM de **Beauty, Fragrances & Fashion**, 42 cuentas
con pricing, surtido y promociones (+30% GMV QoQ) y dueña del canal Flash Sales reportando al CEO
con planes comerciales alineados a P&L; DoFreeze — Shopify end-to-end (catálogo, contenido,
colecciones, CRO/AOV) + Meta/Google Ads + quick-commerce (talabat, Noon, Careem, Deliveroo) y
gestión de 4 agencias; Glovo — construcción de la vertical Retail onboardeando marcas de moda
y lifestyle; Massimo Dutti (Inditex) — base en retail de moda premium y visual merchandising.

Guardarraíles de honestidad:
- P&L: ha ejecutado planes comerciales **alineados a P&L** y gestiona cost drivers de su canal;
  NO se afirma ownership pleno de un P&L de marca.
- Forecast: contribuye con Sales/Finance, no es dueña.
- Sin experiencia en retail de lujo/apparel a escala Chalhoub ni en Salesforce Commerce Cloud:
  no se insinúa.
- Idiomas: ES nativo + EN C1. **Sin árabe** — no se inventa.
- Visa: solo "UAE Residence Visa" (patrocinada por el empleador); nunca "no sponsorship needed".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Chalhoub Group"
TITLE = "E-Commerce Manager - Lacoste & Guess"
DATE_FOLDER = "2026-09-22"

JOB_DESCRIPTION = """\
E-Commerce Manager, Lacoste & Guess — Chalhoub Group, Dubai, UAE. Hybrid, full time.
The eCommerce Manager is responsible for driving the eCommerce strategy, managing the online P&L,
and optimizing profitable growth. This role involves leveraging data insights, performance marketing,
operations and online merchandising strategies to deliver a world-class customer experience while
collaborating with internal teams, suppliers, and third-party agencies.
Strategy & Planning: lead the annual eCommerce strategy and budget in collaboration with the eCommerce
General Manager and Brand General Manager; manage product, pricing and promotional strategies to drive
site performance and build synergies across all channels (Retail, eCommerce, e-Retail); own the online
brand P&L and leverage its insights to optimize cost drivers and improve profitability; cultivate internal
and external relationships, including third-party agencies and e-Retail teams.
Online Merchandising: develop and execute a compelling product assortment and presentation strategy based
on customer preferences, inventory, performance and trends; oversee product categorization, tagging and
search functionality for improved user experience and conversions; plan and execute promotions, campaigns
and update product content to reflect accurate pricing, images and descriptions.
Performance Marketing & Web Analytics: develop and execute performance marketing strategies across digital
channels (PPC, Display, Social, Affiliate) to drive traffic and conversions; continuously optimize marketing
spend, creative and targeting based on performance data; use web data to identify areas for improvement,
understand customer behaviour onsite, optimize marketing performance and manage KPIs via eCommerce dashboards.
Product Management: develop product strategies aligned to customer needs and brand objectives; work closely
with internal teams to ensure products deliver an exceptional customer experience.
Cross-Functional Leadership: lead innovation in omnichannel experiences with central, product and Retail
teams; manage the product roadmap and ensure alignment with business goals.
"""

ATS = [
    "e-commerce", "eCommerce strategy", "online P&L", "profitable growth", "budget", "trading",
    "pricing", "promotional strategy", "omnichannel", "e-Retail", "retail", "online merchandising",
    "product assortment", "categorization", "tagging", "site search", "product content", "conversion",
    "CRO", "customer experience", "performance marketing", "PPC", "display", "paid social", "affiliate",
    "Meta Ads", "Google Ads", "marketing spend", "ROAS", "ROI", "web analytics", "dashboards", "KPIs",
    "customer behaviour", "campaigns", "product roadmap", "cross-functional", "third-party agencies",
    "fashion", "luxury retail", "GMV", "marketplaces", "stakeholder management", "team leadership",
]

CONTENT = {
    "headline": "E-Commerce Manager · Online Trading, Merchandising & Performance Marketing",
    "professional_summary": (
        "E-commerce and brand professional with 5 years across fashion, beauty and FMCG — Alibaba's Miravia "
        "(42 beauty, fragrance and fashion accounts: assortment, pricing and promotions, +30% GMV QoQ), Glovo's "
        "Retail vertical and Inditex — now running a D2C store, paid media and marketplace channels in Dubai."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE consumer group (Befit, Eurocake, Flair) | Shopify D2C + talabat, Noon, Careem, Deliveroo | 50+ markets",
            "bullets": [
                "Own the Shopify store end-to-end — assortment and collections, categorisation, product content, pricing and checkout — lifting conversion rate and average order value through data-led online merchandising",
                "Plan and optimise performance marketing on Meta and Google Ads (audiences, creative A/B testing, social and EDM), steering spend, creative and targeting on ROAS and cost-per-acquisition data",
                "Run trading and promotional calendars across D2C and e-Retail partners (talabat, Noon, Careem, Deliveroo) — pricing, mechanics, campaign assets and stock availability — keeping brand execution consistent across channels",
                "Lead a team of two and brief four external agencies, and build weekly dashboards on traffic, conversion, sell-out and margin drivers to steer corrective pricing and range actions",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees",
            "bullets": [
                "Managed 42 brand accounts across beauty, fragrances and fashion — assortment strategy, pricing, online competition benchmarking and targeted promotions — delivering +30% GMV growth QoQ",
                "Owned the Flash Sales channel for Beauty, Fashion & Home reporting to the CEO: campaign calendar, item selection, quantities and commercial plans aligned to P&L targets",
                "Led category expansion as PIC Fragrances (30+ brands onboarded in two months) and analysed traffic, conversion, retention, ROI and ROAS to optimise channel performance",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Helped build the Retail vertical — onboarding fashion and lifestyle brands onto the platform — and grew XL partner GMV through catalogue, promotions and bespoke activations",
            ],
        },
        {
            "company": "Mondelez International · Massimo Dutti (Inditex)",
            "role": "Category Planning Trainee · Sales Associate",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG category planning | Premium fashion retail floor",
            "bullets": [
                "Sell-in/sell-out and promotional-effectiveness analysis with Nielsen at Mondelez; Inditex grounding in fashion retail operations, visual merchandising and premium customer experience",
            ],
        },
    ],
    "skills_brand": (  # -> "E-Commerce & Trading"
        "online trading, assortment & range, pricing & promotions, campaign calendars, budget planning, P&L-aligned commercial plans"
    ),
    "skills_ecommerce": (  # -> "Merchandising & Site"
        "Shopify, catalogue & collections, categorisation, tagging & site search, product content, CRO, UX, AOV"
    ),
    "skills_commercial": (  # -> "Performance Marketing"
        "Meta Ads, Google Ads (PPC), paid social, display, affiliate & influencer, EDM, spend & creative optimisation"
    ),
    "skills_data": (  # -> "Analytics & Stakeholders"
        "e-commerce dashboards, traffic & conversion analysis, ROI/ROAS, cost drivers, agencies, Retail & e-Retail teams"
    ),
    "skills_tools": (  # -> "Tools"
        "Shopify, Meta Ads Manager, Google Ads, Power BI, Tableau, Looker, Salesforce, SAP, Excel (Expert), Claude (AI)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "E-Commerce & Trading",
    "E-Commerce & Digital": "Merchandising & Site",
    "Commercial": "Performance Marketing",
    "Data & Analytics": "Analytics & Stakeholders",
}


def make_job() -> Job:
    return Job(
        id="chalhoub-group-ecommerce-manager-lacoste-guess-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE (Hybrid)",
        url="https://www.linkedin.com/jobs/search/?keywords=Chalhoub%20Group%20E-Commerce%20Manager%20Lacoste%20Guess",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Chalhoub Group E-Commerce Manager Lacoste Guess Dubai",
            "note": "Fashion e-commerce (Lacoste & Guess) dentro de Chalhoub. Encaje funcional alto "
                    "(trading, merchandising online, performance marketing, analytics) y sector moda/beauty "
                    "afín a Miravia e Inditex. Brechas: ownership pleno de P&L de marca y apparel a escala. "
                    "386 candidatos, +100 solicitudes en 4 días. Contacto de 2º grado: Jorge Lluch (VP Marketing, "
                    "alumni CUNEF) — posible entrada por LinkedIn.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 80, "ai_tier": "Hot",
        "skills_match": [
            "Miravia (Alibaba): 42 cuentas de beauty, fragancias y MODA — surtido, pricing y promos, +30% GMV QoQ",
            "Flash Sales channel owner reportando al CEO, con planes comerciales alineados a P&L",
            "Shopify end-to-end: catálogo, categorización, contenido de producto, CRO y AOV",
            "Performance marketing hands-on: Meta Ads y Google Ads, optimización de spend, creatividades y targeting por ROAS",
            "Omnicanal real en UAE: D2C + e-Retail (talabat, Noon, Careem, Deliveroo) y gestión de 4 agencias",
            "Base en moda premium: Massimo Dutti (Inditex) y onboarding de marcas de moda en la vertical Retail de Glovo",
        ],
        "missing_skills": [
            "Ownership pleno de un P&L de marca online (ha ejecutado planes alineados a P&L, no ownership)",
            "Apparel de lujo a escala Chalhoub (Salesforce Commerce Cloud / plataformas enterprise)",
            "Árabe",
        ],
        "sector_fit": "alto — moda y beauty, su terreno desde Miravia e Inditex",
        "seniority_fit": "encaje — Manager, en línea con su título actual",
        "red_flags": [
            "386 candidatos y +100 solicitudes en 4 días: mucha competencia",
            "34% de los candidatos tiene MBA; Paula tiene grado (CUNEF)",
            "El JD pide ownership explícito del P&L online: es la pregunta que hay que preparar",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Es de los encajes más limpios del pipeline: trading online, merchandising, performance marketing y "
            "analytics son exactamente lo que Paula hace hoy en Shopify y lo que hizo en Miravia con 42 cuentas de "
            "beauty, fragancias y moda. Suma la base de Inditex y el onboarding de marcas de moda en Glovo. La brecha "
            "real es el ownership pleno del P&L de marca y la escala de apparel de lujo — preparar respuesta concreta "
            "sobre cost drivers y margen. Vía caliente: Jorge Lluch (VP Marketing en Chalhoub, alumni CUNEF, 2º grado)."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
