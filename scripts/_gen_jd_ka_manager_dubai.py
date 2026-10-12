"""One-off: CV for JD.com (JINGDONG) "KA Manager", Dubai (Workday req JR101248, posted 30+ days
before 2026-10-12). CV only (no cover letter requested).

The posting sits under JD's International Logistics block (JINGDONG Logistics / JoyLogistics, which
opened a JAFZA warehouse in 2025): a B2B key-account role selling and running supply chain solutions
for brands and retailers. Five areas: customer development, solution design (cost modelling, pricing,
bidding, negotiation), project operations (quality, profitability, payment timelines), market insights
(trend / competitor reports) and team collaboration (roles, training, process improvement).

Paula's honest angle:
- Miravia (Alibaba) KAM: 42 accounts, +30% GMV QoQ, Flash Sales vs P&L targets reporting to the CEO,
  30+ stores onboarded in two months (incl. Arabian Oud, Lattafa, Swiss Arabian, Ajmal distributors).
- Glovo XL accounts: negotiated deals profitable for both sides; worked with logistics and support.
- DoFreeze: she is the kind of client JD Logistics sells to. Weekly stock-risk review by dark store
  across Careem, talabat and Noon (AI dashboard built as a team); 6 launches across 50+ markets;
  manages two reports; Claude automation (~40% less manual work) as process improvement.

Honesty guardrails:
- NO 3PL / warehousing sales, tenders or logistics cost modelling claimed: flagged as the gap.
- Never "own the forecast"; the stock-risk dashboard is "we built". DoFreeze is NOT on Deliveroo.
- No business-review / payment-collection claims; no Arabic or Mandarin; visa line from profile.yaml.

ONE-PAGE standard. Lands under output/2026-10-12/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "JD.com"
TITLE = "KA Manager"
DATE_FOLDER = "2026-10-12"

JOB_DESCRIPTION = """\
KA Manager — JD.com (JINGDONG), Dubai, UAE. Full time. Req JR101248 (posted 30+ days ago).
1. Customer Development & Maintenance: identify key client needs, build and nurture high-level
relationships, design account management strategies for client satisfaction and business continuity.
2. Solution Design: analyse industry trends and client challenges; develop tailored supply chain
solutions (cost modelling, pricing strategies); lead bidding and negotiations; ensure efficient
project execution.
3. Project Operations Management: monitor project quality, profitability and payment timelines;
coordinate cross-functional resources to resolve issues, optimise workflows and cost structures.
4. Market Insights & Strategy: gather industry trends, competitor intelligence and client feedback to
produce market analysis reports; drive product innovation and model enhancements.
5. Team Collaboration & Empowerment: oversee project team roles, improve efficiency through training,
incentives and process improvements, build a talent pipeline.
Context: JD.com International Logistics (JINGDONG Logistics / JoyLogistics): warehousing, line-haul,
last-mile, cold chain and cross-border networks; bonded and overseas warehouses incl. the UAE.
"""

ATS = [
    "Key Account Manager", "KA Manager", "key account management", "account management strategy",
    "client relationships", "customer development", "client needs", "business continuity",
    "solution design", "supply chain", "pricing strategy", "negotiation", "commercial deals",
    "project execution", "profitability", "P&L", "cross-functional", "logistics", "operations",
    "workflow optimisation", "process improvement", "market analysis", "competitor intelligence",
    "industry trends", "team leadership", "training", "e-commerce", "marketplace", "Alibaba",
    "GMV", "onboarding", "dark stores", "stock availability", "Careem", "talabat", "Noon",
    "AI", "automation", "Excel", "Power BI", "Dubai", "UAE",
]

CV_CONTENT = {
    "headline": (
        "Key Account Management · E-Commerce & Q-Commerce Partners · Commercial Negotiation · "
        "Data & AI Automation"
    ),
    "professional_summary": (
        "Key account manager with 5 years across e-commerce, delivery platforms and FMCG. At Alibaba's "
        "Miravia I grew 42 accounts +30% GMV QoQ; at Glovo, XL chains. In Dubai I sit on the client side "
        "of e-commerce supply (Careem, talabat, Noon dark stores) and build AI tools that automate reporting."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Homegrown UAE F&B / FMCG group | Befit, Eurocake, Flair | Careem, talabat, Noon | 50+ markets",
            "bullets": [
                "Client-side view of UAE e-commerce supply: run the weekly stock-risk review by dark store across Careem, talabat and Noon, using an AI dashboard we built with Claude to flag where platform POs will not cover demand",
                "Led 6 product launches end-to-end (brief, packaging, pricing, go-to-market) across GCC, MENA, Asia, Europe, USA and Africa",
                "Lead a team of two (designer, social media exec); built a Claude system that automates market research and KPI reporting, cutting manual work ~40%",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Top-5 global e-commerce | 100K+ employees",
            "bullets": [
                "Managed 42 key accounts to +30% GMV QoQ, turning sales, traffic and conversion data into account plans and pricing, assortment and promotion recommendations",
                "Owned the Flash Sales channel for Beauty, Fashion & Home against P&L targets, reporting directly to the CEO",
                "Onboarded 30+ new stores in two months as PIC Fragrances, incl. the official distributors of Arabian Oud, Lattafa, Swiss Arabian and Ajmal",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Food-delivery & quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Managed XL accounts (KFC, Taco Bell, La Tagliatella, Sushi Shop) and negotiated commercial deals profitable for both Glovo and the partner, growing orders and GMV",
                "Coordinated marketing, logistics and customer support to deliver partner campaigns; helped build the Retail vertical beyond food",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | Chocolate category",
            "bullets": [
                "Sell-in/sell-out and promotional-effectiveness analysis and performance reports for the chocolate category",
            ],
        },
    ],
    "skills_brand": (  # -> "Key Accounts"
        "account strategy, relationship management, commercial negotiation, onboarding, pricing"
    ),
    "skills_ecommerce": (  # -> "Platforms & Supply"
        "Careem, talabat, Noon, Miravia, Glovo: dark-store stock risk, platform POs, listings"
    ),
    "skills_commercial": (  # -> "Cross-functional"
        "Logistics, Operations, Customer Support, Marketing, distributors; team of two; CEO reporting"
    ),
    "skills_data": (  # -> "Data & Insight"
        "GMV, P&L targets, sell-in/sell-out, promo ROI, market & competitor analysis, dashboards"
    ),
    "skills_tools": (
        "Excel (advanced), Power BI, Looker, Salesforce, SAP, Claude / Claude Code, Python automation"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Key Accounts",
    "E-Commerce & Digital": "Platforms & Supply",
    "Commercial": "Cross-functional",
    "Data & Analytics": "Data & Insight",
    "Tools": "Tools & AI",
}


def make_job() -> Job:
    return Job(
        id="jd-com-ka-manager-dubai-jr101248",
        title=TITLE,
        company=COMPANY,
        location="Dubai, United Arab Emirates",
        url="https://www.bayt.com/en/uae/jobs/ka-manager-74538145/",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={"req": "JR101248 (JD.com Workday careers)",
             "via": "Workday posting pasted by Paula; URL is a Bayt mirror listing",
             "business": "JD.com International Logistics (JINGDONG Logistics / JoyLogistics)",
             "posted": "30+ days before 2026-10-12"},
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not posted",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": None, "raw": job.raw,
        "ai_score": 62, "ai_tier": "Warm",
        "skills_match": [
            "Miravia (Alibaba) KAM: 42 accounts, +30% GMV QoQ, Flash Sales vs P&L targets reporting to the CEO",
            "Glovo XL accounts: negotiated commercial deals profitable for both sides; worked with logistics",
            "Client-side view of UAE e-commerce supply: weekly dark-store stock-risk review on Careem, talabat, Noon",
            "Manages two reports; AI automation of market research and reporting (~40% less manual work)",
        ],
        "missing_skills": [
            "No 3PL / warehousing / logistics solution selling",
            "No tender / bidding experience or supply chain cost modelling",
            "Mandarin often preferred at JD (not stated in this JD)",
        ],
        "sector_fit": "partial (B2B logistics key accounts; her KAM experience is marketplace / platform side)",
        "seniority_fit": "good (Manager-level KAM)",
        "red_flags": [
            "Posted 30+ days ago: may be filled or near shortlist",
            "Core of the role is supply chain solution design and bidding, outside her track record",
        ],
        "ats_keywords": ATS,
        "reasoning": "Strong KAM and e-commerce track record (Alibaba, Glovo) and client-side UAE q-commerce supply view, but the role sells logistics solutions with cost modelling and bidding she has not done.",
        "scored_by": "manual:claude", "freshness": "stale",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _pages(pdf: Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(pdf)).pages)


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    final_dir = settings.output_dir / DATE_FOLDER / f"{COMPANY} - {TITLE}"
    dest = final_dir / "01_CV_y_Carta"
    dest.mkdir(parents=True, exist_ok=True)

    cv_docx = cv._fill_template(CV_CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)

    moved = dest / cv_pdf.name
    shutil.move(str(cv_pdf), str(moved))
    shutil.copy(str(moved), str(dest / "Paula De Francisco - CV.pdf"))
    print("OK", moved.name, "PAGES", _pages(moved))

    stray = cv_pdf.parent.parent
    if stray.exists() and stray.resolve() != final_dir.resolve():
        shutil.rmtree(stray, ignore_errors=True)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
