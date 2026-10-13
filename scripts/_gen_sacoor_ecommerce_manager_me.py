"""One-off: CV for Sacoor Brothers "E-commerce Manager - Middle East", Dubai (LinkedIn verified job,
Easy Apply, posted ~2026-10-12, 660 applicants). Job poster: Amina B., Talent Partner (2nd degree).
CV only.

The brief: own the GCC online shop (Shopify Plus, Arabic site) from first click to delivered parcel;
grow online revenue; run peaks (Ramadan, Eid, White Friday, DSS, National Days); Google/Meta Ads, SEO,
CRM; manage agencies; own the e-commerce budget; customer journey to repeat purchase; work with
Ops/Logistics and Retail (online + in-store); weekly/monthly reporting. Sacoor is a Portuguese premium
fashion brand (menswear/womenswear, also fragrances), competing with Massimo Dutti.

Paula's honest angle:
- Shopify: owns the DoFreeze web store end-to-end (confirmed 2026-09-13); orders +86%, revenue +57%
  (Jun-Aug vs Q1 2026) after reviews, replenishment automation and bundles from basket analysis.
- Paid + CRM: Meta/Google on break-even ROAS net of marketplace commissions; Advantage+ catalogue
  retargeting 3.7x ROAS in its first week; email base cleaned for deliverability.
- Peaks: Ramadan, New Year, Fitness Month, back to school across talabat, Noon, Careem and web,
  read against control (SMASH x talabat +165% vs +4.7%; Befit x Noon +4,176 units).
- Agencies: compared 4, negotiated -30%, audited the agency's reach reporting.
- Fashion: Miravia fashion/beauty/fragrance accounts; Massimo Dutti shop floor (online + in-store);
  CUNEF e-commerce & fashion specialisation. Fragrance + Arabian perfume is a GCC hook.

Honesty guardrails:
- Miravia per Paula 2026-10-13 (memory: miravia-scope-all-cvs): "Miravia & AliExpress (Alibaba
  Group)", L'Oréal brands ahead of KIKO, NIVEA "three months", visibility decisions, "worked with"
  marketing/product/ops. +30% GMV QoQ is her 42-account portfolio, never "the category grew".
- Glovo per Paula 2026-10-13 (memory: glovo-real-role): co-managed XL restaurant chains. NO Retail
  vertical, NO deal negotiation, NO team leadership.
- NOT claimed (not on record): Shopify Plus, Google Analytics, Arabic site, COD / BNPL, Portuguese,
  3 years of GCC e-commerce (she has one), Saudi e-commerce. Never "own the forecast".
- No Arabic. Visa line comes from profile.yaml.

ONE-PAGE standard. Lands under output/2026-10-13/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Sacoor Brothers"
TITLE = "E-commerce Manager - Middle East"
DATE_FOLDER = "2026-10-13"

JOB_DESCRIPTION = """\
E-commerce Manager - Middle East, Sacoor Brothers, Dubai, UAE. On-site, full-time. LinkedIn verified
job, Easy Apply. Posted by Amina B., Talent Partner (GCC & International Hiring).

Sacoor Brothers is an international premium fashion brand with Portuguese roots and stores across the
UAE, Saudi Arabia, Qatar, Kuwait, Bahrain and Oman.

We are looking for an E-commerce Manager - Middle East to grow our online sales across the GCC. You
will own the online shop from first click to delivered parcel. You will set the plan and also do the
work.

What you will do:
- Grow online revenue across GCC markets and hit monthly and yearly targets
- Run the online shop on Shopify Plus, including the Arabic site and local content for each market
- Plan and run peak seasons: Ramadan, Eid, White Friday, DSS and National Days
- Plan and manage paid ads (Google and Meta), SEO and CRM
- Manage our outside agencies and keep them on target
- Own the e-commerce budget and track what it returns
- Improve the customer journey, from first visit to repeat purchase
- Work with Operations and Logistics on delivery, returns and payment options
- Work with the Retail team to connect online and in-store shopping
- Report weekly and monthly on sales, traffic and results

Requirements:
- 5+ years in e-commerce, ideally in fashion or retail
- At least 3 years running e-commerce in the GCC, with Saudi experience a strong plus
- Hands-on with Shopify (Shopify Plus is a big plus)
- Good knowledge of Google Ads, Meta Ads, SEO and Google Analytics
- Experience with local payment and delivery options, such as cash on delivery, buy-now-pay-later and
  last-mile partners
- Experience managing agencies and budgets, and owning sales targets
- Comfortable with numbers, Excel and reports
- Able to set the plan and also do the work
- Bachelor's degree in Business, Marketing or a related field
- Fluent English. Arabic is a strong plus, and Portuguese is a bonus.

Benefits: international premium fashion brand; head office and luxury retail; training and development.
"""

ATS = [
    "E-commerce Manager", "e-commerce", "online shop", "Shopify", "online revenue", "GCC", "UAE",
    "Middle East", "fashion", "retail", "premium", "peak seasons", "Ramadan", "White Friday", "DSS",
    "Google Ads", "Meta Ads", "SEO", "CRM", "email marketing", "agencies", "budget", "ROAS",
    "customer journey", "repeat purchase", "retention", "conversion", "AOV", "omnichannel",
    "online and in-store", "Operations", "Logistics", "returns", "weekly reporting", "monthly reporting",
    "sales targets", "Excel", "Business Administration", "marketplaces", "Noon", "Amazon", "talabat",
]

CV_CONTENT = {
    "headline": (
        "E-Commerce · Shopify & Marketplaces · Meta & Google Ads · CRM · Fashion & Beauty"
    ),
    "professional_summary": (
        "E-commerce professional with 5 years in e-commerce and commercial roles across fashion, beauty "
        "and FMCG. In Dubai I run a Shopify store end to end (orders +86%, revenue +57%), Meta and Google "
        "Ads and UAE marketplaces; at Alibaba I grew 42 fashion and beauty accounts +30% GMV QoQ."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE F&B / FMCG group | Befit, Eurocake, Flair | Shopify web store + talabat, Noon, Careem, Amazon",
            "bullets": [
                "Own the Shopify web store end to end: orders +86% and revenue +57% (Jun–Aug vs Q1 2026) after adding reviews, a 30-day replenishment flow and bundles rebuilt from basket analysis",
                "Run Meta and Google Ads on break-even ROAS per brand and channel, net of marketplace commissions, moving budget from Google to Meta; a dynamic-catalogue retargeting campaign (Advantage+) returned 3.7x ROAS in its first week",
                "Plan seasonal peaks (Ramadan, New Year, Fitness Month) across web, talabat, Noon and Careem, measured on incremental sales: SMASH x talabat +165% vs +4.7% control; Befit x Noon +4,176 units",
                "Manage creator agencies: compared four, negotiated −30% and audited their reach reports; 103 activations at AED 81 per piece",
                "Lead a team of two (designer, social media executive); built a Claude-based KPI-reporting system, ~40% less manual work",
            ],
        },
        {
            "company": "Miravia & AliExpress (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplaces | Beauty & Health department",
            "bullets": [
                "Managed 42 fashion, beauty and fragrance accounts to +30% GMV QoQ through pricing, assortment and promotions, incl. L'Oréal Group (L'Oréal Paris, Maybelline, NYX, Redken), KIKO Milano and NIVEA (three months)",
                "Decided in-platform visibility (banners, placements, traffic) and owned the Flash Sales channel for Beauty, Fashion & Home, reporting to the CEO; worked with marketing, product and operations teams",
                "Spotted the Arabian-perfume trend and grew it into one of the fastest-growing sub-categories in Beauty & Health, onboarding 30+ stores in two months (Lattafa, Arabian Oud, Swiss Arabian, Ajmal)",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Food-delivery & quick-commerce marketplace | 10K+ employees",
            "bullets": [
                "Co-managed five XL restaurant chains (incl. KFC, Taco Bell, Vicio): promotions and promo budgets, in-app visibility and its return, delivery times, cancellations and refunds",
            ],
        },
        {
            "company": "Massimo Dutti (Inditex) · Mondelez International",
            "role": "Sales Associate · Category Planning Trainee",
            "dates": "Jun 2018 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Premium fashion retail floor | Global FMCG category planning",
            "bullets": [
                "Premium fashion on the shop floor at Massimo Dutti (styling, visual merchandising, store cadence); sell-in/sell-out and promo-effectiveness analysis at Mondelez",
            ],
        },
    ],
    "skills_brand": (  # -> "E-Commerce"
        "Shopify store management, merchandising, bundles, reviews, CRO, AOV, checkout, marketplaces"
    ),
    "skills_ecommerce": (  # -> "Acquisition & CRM"
        "Meta Ads (Advantage+), Google Ads, SEO & AEO, email & replenishment flows, affiliates"
    ),
    "skills_commercial": (  # -> "Trading & Budget"
        "peak-season calendars, promotions, pricing, assortment, agencies, break-even ROAS, P&L targets"
    ),
    "skills_data": (  # -> "Data & Reporting"
        "revenue, orders, conversion, ROAS, uplift vs control, weekly & monthly KPI reports, Excel (Expert)"
    ),
    "skills_tools": (
        "Shopify, Meta Ads Manager, Google Ads, Judge.me, BixGrow, Power BI, Looker, Claude / Claude Code"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "E-Commerce",
    "E-Commerce & Digital": "Acquisition & CRM",
    "Commercial": "Trading & Budget",
    "Data & Analytics": "Data & Reporting",
    "Tools": "Tools & AI",
}


def make_job() -> Job:
    return Job(
        id="sacoor-brothers-ecommerce-manager-middle-east-2026-10",
        title=TITLE,
        company=COMPANY,
        location="Dubai, United Arab Emirates",
        url="https://www.linkedin.com/jobs/",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={"team": "Sacoor Brothers — E-commerce, Middle East (GCC)",
             "via": "LinkedIn (verified job, Easy Apply, posted ~2026-10-12, 660 applicants)",
             "poster": "Amina B., Talent Partner (2nd degree)"},
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": None, "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": "2026-10-12", "raw": job.raw,
        "ai_score": 70, "ai_tier": "Warm",
        "skills_match": [
            "Owns a Shopify web store end to end (orders +86%, revenue +57%)",
            "Meta and Google Ads on break-even ROAS; Advantage+ catalogue retargeting 3.7x ROAS",
            "Seasonal peaks (Ramadan, New Year) on UAE marketplaces, measured against control",
            "Agency management: compared four, negotiated −30%, audited reporting",
            "Fashion grounding: Miravia fashion accounts, Massimo Dutti shop floor, CUNEF e-commerce & fashion",
        ],
        "missing_skills": [
            "1 year of GCC e-commerce (asks 3+); no Saudi e-commerce",
            "Shopify Plus, Arabic site, COD / BNPL not on record",
            "Google Analytics not confirmed",
            "No Arabic (strong plus) or Portuguese (bonus)",
        ],
        "sector_fit": "good (premium fashion e-commerce; her fashion work is marketplace-side and retail floor)",
        "seniority_fit": "stretch — asks 5+ yrs e-commerce, she has ~4",
        "red_flags": ["660 applicants in a day", "On-site"],
        "ats_keywords": ATS,
        "reasoning": "Shopify ownership, paid media, CRM and seasonal trading map well; GCC tenure and Arabic are the gaps.",
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "CV Ready",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _pages(pdf: Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(pdf)).pages)


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    final_dir = settings.output_dir / DATE_FOLDER / f"{COMPANY} - {TITLE}"
    dest = final_dir / "01_CV_y_Carta"
    dest.mkdir(parents=True, exist_ok=True)

    cv_docx = cv._fill_template(CV_CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)

    # Keep the DOCX next to the PDF (matches the other packages) so Paula can tweak it.
    for f in (cv_pdf, cv_docx):
        if f.exists():
            shutil.move(str(f), str(dest / f.name))
    moved = dest / cv_pdf.name
    shutil.copy(str(moved), str(dest / "Paula De Francisco - CV.pdf"))
    print("OK", moved.name, "PAGES", _pages(moved))

    stray = cv_pdf.parent.parent
    if stray.exists() and stray.resolve() != final_dir.resolve():
        shutil.rmtree(stray, ignore_errors=True)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
