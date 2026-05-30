#!/usr/bin/env python3
"""Bespoke Trade & Shopper Marketing Plan — Schwarzkopf Hair Care, GCC.

A send-ready, branded 2-3 page PDF for Paula's Henkel application. Authored
content (Claude as the strategist), rendered via python-docx and converted to
PDF with the same reliable win32com path the repo's deliverables generator uses.
Saved into the existing Henkel folder under output/2026-05-29/.
"""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from career_ops.config import settings

OUT_DIR = ROOT / "output" / "2026-05-29" / "Henkel - Trade & Shopper Marketing Manager - GCC"
OUT_DIR.mkdir(parents=True, exist_ok=True)

P = settings.profile["personal"]

INK = RGBColor(0x1A, 0x1A, 0x1A)        # near-black body
DARK = RGBColor(0x11, 0x11, 0x11)       # headings
RED = RGBColor(0xE1, 0x00, 0x0F)        # Henkel red accent
GREY = RGBColor(0x66, 0x66, 0x66)
LIGHT = RGBColor(0x8A, 0x8A, 0x8A)
HDR_FILL = "111111"
ACCENT_FILL = "E1000F"


# ------------------------------------------------------------------ helpers
def run(p, text, *, bold=False, italic=False, size=10.5, color=INK, font="Calibri"):
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.bold = bold
    r.italic = italic
    return r


def shade(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def left_border(paragraph, color="111111", sz="18", space="8"):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), sz)
    left.set(qn("w:color"), color)
    left.set(qn("w:space"), space)
    bdr.append(left)
    pPr.append(bdr)


def bottom_border(paragraph, color="111111", sz="12"):
    pPr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:color"), color)
    bdr.append(bottom)
    pPr.append(bdr)


def heading(doc, n, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(3)
    left_border(h)
    run(h, f"{n}.  {text}", bold=True, size=12.5, color=DARK)
    return h


def body(doc, lead, text, *, color=INK, after=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.left_indent = Cm(0.3)
    if lead:
        run(p, f"{lead}  ", bold=True, size=10.5, color=color)
    run(p, text, size=10.5, color=color)
    return p


def bullet(doc, lead, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.7)
    if lead:
        run(p, f"{lead}: ", bold=True, size=10.5, color=INK)
    run(p, text, size=10.5, color=INK)
    return p


def make_table(doc, headers, rows, widths_cm):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = False
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].width = Cm(widths_cm[i])
        para = hdr[i].paragraphs[0]
        run(para, htext, bold=True, size=9.5, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(para, HDR_FILL)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].width = Cm(widths_cm[i])
            para = cells[i].paragraphs[0]
            run(para, val, size=9, color=INK)
    return t


# ------------------------------------------------------------------ build
doc = Document()
s = doc.sections[0]
s.top_margin = Cm(1.6)
s.bottom_margin = Cm(1.6)
s.left_margin = Cm(2.0)
s.right_margin = Cm(2.0)

# Badge
badge = doc.add_paragraph()
shade(badge, ACCENT_FILL)
run(badge, "  TRADE & SHOPPER MARKETING PLAN  ", bold=True, size=8.5,
    color=RGBColor(0xFF, 0xFF, 0xFF))

# Title
title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(4)
title.paragraph_format.space_after = Pt(1)
run(title, "Schwarzkopf Hair Care — GCC", bold=True, size=20, color=DARK)

sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(4)
run(sub, "A 90-day channel & shopper activation blueprint for Modern Trade, "
         "pharmacy and e-/quick-commerce", italic=True, size=11, color=GREY)

meta = doc.add_paragraph()
run(meta, f"Prepared by {P['name']}   ·   {date.today().strftime('%d %B %Y')}   ·   "
          f"For: Trade & Shopper Marketing Manager – GCC, Henkel", size=9, color=LIGHT)
divp = doc.add_paragraph()
divp.paragraph_format.space_after = Pt(6)
bottom_border(divp)

# Intro
intro = doc.add_paragraph()
intro.paragraph_format.space_after = Pt(4)
run(intro, "Purpose.  ", bold=True, size=10.5, color=RED)
run(intro, "This is how I would land in the role: a fast read of the GCC hair-care "
           "shopper, a clear definition of Perfect Store by channel, a customer-led "
           "activation calendar, and the KPI/ROI discipline to prove it works — built "
           "to be executed with the Key Account and Sales teams from week one.", size=10.5)

# 1. Shopper & category insight
heading(doc, 1, "Shopper & category insight (GCC)")
bullet(doc, "Climate-driven need", "Heat, sun and humidity 8+ months a year make damage, "
       "frizz and scalp care year-round needs — Schwarzkopf can own \"repair & protect\" "
       "(Gliss) and premium care (BC Bonacure) against a humidity-led benefit story.")
bullet(doc, "Two shoppers, two missions", "High-frequency replenishment in grocery/hypermarket "
       "vs. advice-led discovery in pharmacy — each needs a different Perfect Store and mechanic.")
bullet(doc, "Premiumisation + value polarity", "A large premium expat segment and a value-driven "
       "mass segment coexist; the range and pack-price architecture must serve both without trading down the brand.")
bullet(doc, "Channels that matter", "Modern Trade (Carrefour/MAF, Lulu, Union Coop, Spinneys), "
       "pharmacy (BinSina, Aster, Life), e-tail (Noon, Amazon.ae) and quick-commerce "
       "(Noon Minutes, Talabat) — the latter now a real hair-care discovery channel for younger shoppers.")
bullet(doc, "Seasonal peaks", "Summer (anti-frizz/sun protection), Ramadan & Eid (gifting, premium), "
       "back-to-school and DSF/GITEX retail moments anchor the calendar.")

# 2. Perfect Store
heading(doc, 2, "Perfect Store standards by channel")
body(doc, "Definition:", "a single, measurable picture of success per channel tier — "
     "availability, visibility, pricing and activation — audited monthly and tied to KAM scorecards.")
make_table(
    doc,
    ["Channel", "Perfect Store priorities"],
    [
        ["Hypermarket / Modern Trade",
         "Hero SKU 100% availability · planogrammed block by benefit · 1 secondary display per quarter · price-pack compliance · promo ROI tracked"],
        ["Pharmacy",
         "Advice fixture for premium care · trained staff/sampling · BC Bonacure & repair range visible · GWP at till"],
        ["E-tail (Noon / Amazon.ae)",
         "A+ content on hero SKUs · 4.3★+ review health · search share on \"shampoo/hair repair\" · pack-shot & title compliance"],
        ["Quick-commerce (Talabat / Noon Minutes)",
         "Top-up bundles · hero-SKU availability · creator-led discovery tiles · impulse pricing"],
    ],
    [4.5, 12.5],
)

# 3. Key account & channel plan
heading(doc, 3, "Key account & channel plan")
body(doc, "Joint Business Plan (flagship — Carrefour/MAF):", "agree annual growth, distribution and "
     "share-of-shelf targets; build a quarterly promo grid and a retail-media plan; cluster stores "
     "(premium mall vs. community) for differentiated assortment.")
body(doc, "Traditional Trade & distributor coverage:", "drive numeric distribution of hero SKUs through "
     "the distributor, with a simple Perfect Store checklist and incentive for must-stock list compliance.")
body(doc, "Ways of working with KAMs:", "I act as the shopper/category engine behind the account teams — "
     "translating insight into customer plans, building the sell-in story, and owning post-evaluation.")

# 4. 90-day activation calendar
heading(doc, 4, "90-day trade activation calendar")
make_table(
    doc,
    ["Window", "Activation", "Channel", "Mechanic", "Objective"],
    [
        ["Wk 1–2", "Onboarding & Perfect Store audit", "All", "Baseline scorecard + KAM 1:1s", "Diagnose gaps"],
        ["Wk 3–4", "Hero-SKU availability fix", "MT + e-tail", "Must-stock list, A+ content refresh", "Distribution +X pts"],
        ["Month 2", "Summer anti-frizz burst", "MT + pharmacy", "Secondary display + GWP + sampling", "Sell-out uplift"],
        ["Month 2", "Creator-led q-commerce push", "Talabat/Noon Min.", "Influencer codes + bundle", "Trial & basket size"],
        ["Month 3", "Ramadan/premium gifting pre-build", "MT + pharmacy", "Gift packs + premium endcap", "Premium mix up"],
        ["Month 3", "Business review & scale", "All", "Promo ROI readout, scale winners", "Lock Q4 plan"],
    ],
    [1.8, 4.2, 3.0, 4.2, 3.3],
)

# 5. Shopper activation & NPD
heading(doc, 5, "Shopper activation mechanics & NPD go-to-market")
bullet(doc, "Mechanics toolkit", "GWP (gift-with-purchase), price-pack bundles, pharmacy sampling, "
       "endcaps/secondary displays, and creator-led q-commerce codes — chosen by shopper mission, not habit.")
bullet(doc, "NPD launch playbook", "listing & Perfect Store readiness → launch burst (display + sampling + "
       "creator seeding) → e-tail A+ & review seeding → 6-week post-evaluation and scale decision.")
bullet(doc, "Retail media", "fund always-on search and sponsored placements on Noon/Amazon.ae tied to the "
       "promo grid, so in-store and on-shelf-digital pull in the same direction.")

# 6. KPIs & ROI
heading(doc, 6, "KPIs, ROI & cadence")
make_table(
    doc,
    ["KPI", "What it proves", "Cadence"],
    [
        ["Numeric & weighted distribution (hero SKUs)", "Availability foundation", "Monthly"],
        ["Perfect Store compliance %", "Execution quality", "Monthly audit"],
        ["Sell-out uplift vs. baseline", "Activation effectiveness", "Per activity"],
        ["Promo ROI / ROAS", "Spend efficiency", "Per activity + QBR"],
        ["E-tail search share & review health", "Digital shelf strength", "Bi-weekly"],
        ["Premium mix %", "Value growth, not just volume", "Monthly"],
    ],
    [6.5, 6.5, 4.0],
)
wow = doc.add_paragraph()
wow.paragraph_format.space_before = Pt(4)
run(wow, "Cadence:  ", bold=True, size=10.5, color=RED)
run(wow, "monthly business reviews with Sales, Supply Chain, Finance and Marketing; "
         "quarterly JBP checkpoints with key accounts; A&P managed to ROI throughout.", size=10.5)

# Closing
close = doc.add_paragraph()
close.paragraph_format.space_before = Pt(12)
bottom_border(close, color="DDDDDD", sz="6")
close2 = doc.add_paragraph()
run(close2, "Why I can land this fast:  ", bold=True, size=10, color=DARK)
run(close2, "I already build trade & shopper plans by channel across 50+ markets at DoFreeze, "
            "ran 42 key accounts to +30% GMV QoQ at Alibaba's Miravia, and integrate brands into UAE "
            "modern trade and quick-commerce (Noon, Talabat, Careem, Deliveroo). This plan is the way "
            "I'd put that to work for Schwarzkopf in the GCC.", italic=True, size=10, color=RGBColor(0x44, 0x44, 0x44))

# Footer
fp = s.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(fp, f"{P['name']}   ·   {P['email']}   ·   {P.get('linkedin','')}   ·   {P['phone']}",
    size=8, color=LIGHT)

# Save + PDF
docx_path = OUT_DIR / "Deliverable_Paula_Henkel_Trade-Shopper-Plan_Schwarzkopf.docx"
doc.save(str(docx_path))
pdf_path = docx_path.with_suffix(".pdf")

import pythoncom  # noqa: E402
pythoncom.CoInitialize()
try:
    import win32com.client as win32
    word = win32.gencache.EnsureDispatch("Word.Application")
    word.Visible = False
    try:
        d = word.Documents.Open(str(docx_path))
        d.SaveAs(str(pdf_path), FileFormat=17)
        d.Close(False)
    finally:
        word.Quit()
    docx_path.unlink(missing_ok=True)
    print(f"PDF: {pdf_path}")
except Exception as exc:
    print(f"PDF conversion failed ({type(exc).__name__}); DOCX kept: {docx_path}")
finally:
    pythoncom.CoUninitialize()
