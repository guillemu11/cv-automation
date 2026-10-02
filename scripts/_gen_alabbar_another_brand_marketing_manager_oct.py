"""One-off: CV de una página para **Brand Marketing Manager | F&B** en **Alabbar Enterprises & ANOTHER**
(grupo de retail y F&B de EAU: Candylicious, Garrett Popcorn, Yogurtland, Ethan Allen; conceptos propios
Social House, Karak House, Angelina, Gia, Markette, Ganache Chocolatier, Parka, Krema, etc.).
Dubái, híbrido, jornada completa. LinkedIn: publicada hace 8 meses, sigue viva (2.368 clics en
«Solicitar», 29 en el último día); respuestas gestionadas fuera de LinkedIn.

Es la misma oferta que ya tuvo CV el 2026-08-15 y el 2026-08-27 (`_gen_alabbar_another_brand_marketing_manager.py`).
Aquellos salieron a 3 páginas, con la foto recortada, Deliveroo en skills y las cifras viejas de
creators ("25–50 por campaña"). Este es el rehecho con los estándares actuales.

Pide el JD: brand building y gestión de marca, desarrollo de negocio y de conceptos, estrategia y gestión
de social media, creación de contenido, campañas, posicionamiento y comunicación de marca en todos los
canales con un equipo interno, y medición de resultados, para varios conceptos de F&B y retail.
Requisitos: 4–5 años en brand management, social, comunicación de marca, activación online y offline,
desarrollo de conceptos y campañas; experiencia con marcas conocidas de F&B, retail y hospitality
(preferiblemente homegrown); coordinar equipos internos de digital, diseño gráfico, contenido y
operaciones; dirección de arte; dirección de contenido, foto y vídeo; plazos cortos y varios proyectos a
la vez; copywriting; entorno multicultural; grado en ADE/Marketing. **Árabe: "a definite advantage",
no requisito.**

Ángulo honesto de Paula:
- DoFreeze es un grupo F&B homegrown de EAU (así consta en profile.yaml) con varias marcas: Befit,
  Eurocake (con la nueva SMASH) y Flair — el paralelo directo de "varios conceptos".
- Equipo interno real: diseñador gráfico + social media executive. Ella escribe los briefs y aprueba
  todo el output creativo y social → dirección de arte y de contenido.
- Activación online y offline real: campañas por temporada (Ramadán, back to school, Fitness Month, New
  Year), 4 agencias, 103 activaciones de creator en 3 campañas, samplings con talabat, sorteo in-app con
  talabat, comunidades de running, pádel y yoga.
- Medido en venta: SMASH × talabat +165% de venta diaria contra +4,7% del grupo de control (Befit);
  Befit × Noon +31% sobre baseline.
- Desarrollo de conceptos: 6 NPD end-to-end.
- Marcas de F&B conocidas: en Glovo llevó KFC, Taco Bell, La Tagliatella y Sushi Shop (cuentas XL con
  activaciones de marketing conjuntas). Confitería en Mondelez (Milka, Suchard) — encaja con
  Candylicious, Ganache y Garrett.
- Grado en ADE (CUNEF): cubre el requisito.

Guardarraíles de honestidad:
- **Hospitality / restaurantes desde el lado marca**: no lo ha tenido. Las marcas de restaurante las
  llevó como plataforma (Glovo), no como su equipo de marketing; se dice "managed XL restaurant
  accounts ... joint marketing activations". No se habla de "guest experience" como skill.
- **Dirección de rodajes**: no hay evidencia de que dirija shoots en set. Se dice que briefa y aprueba el
  contenido de foto y vídeo (equipo, creators, agencias), que sí es suyo.
- **Forecast**: la dueña es la e-commerce manager; no aparece.
- **Deliveroo**: DoFreeze NO está en Deliveroo. Plataformas: talabat, Noon, Careem.
- Creators: "103 activaciones en 3 campañas", nunca "una campaña con 103".
- Sin árabe. Visa solo "UAE Residence Visa"; nunca "no sponsorship".
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Alabbar Enterprises & ANOTHER"
TITLE = "Brand Marketing Manager"
DATE_FOLDER = "2026-10-01"
JOB_ID = "alabbar-another-brand-marketing-manager-2026-08"  # mismo id que el CV de agosto: actualiza, no duplica

JOB_DESCRIPTION = """\
Brand Marketing Manager | F&B. Alabbar Enterprises & ANOTHER. Dubai, UAE. Hybrid, full time.
Alabbar Enterprises & ANOTHER is a UAE-based retail group that represents established brands like Candylicious,
Ethan Allen, Yogurtland and Garrett Popcorn. It is also the parent company of multiple F&B concepts like Social
House, Karak House, Angelina, Gia, Markette, Two Bistro, Caya, Parka, Krema, Carpo, Ganache Chocolatier, Parka
Bakehouse, The Good Hood Kitchen and Klay by Karak under its umbrella.
Job Summary: Brand Marketing Manager whose primary responsibility will lie in brand building and management,
business and concept development, social media strategy and management, content creation, developing marketing
campaigns, maintaining brand positioning, and brand communication across all channels with an in-house team and
the measurement of results for a number of our strategic F&B and Retail concepts. Directly accountable for
identifying opportunities, internally and externally, for growing brand presence and growing their image and
positioning through diligent creativity, precise planning, strong relationship building and meticulous coordination.
Job Requirements: at least 4-5 years' experience within brand management, social media, brand communications,
online and offline activation management, concept and campaign development. Must have experience with well-known
F&B, retail and hospitality brands (preferably homegrown). High level of creativity. Passion for F&B and the
lifestyle industry. Excellent communication and presentation skills in English. Arabic is a definite advantage.
Ability to manage and coordinate with in-house digital, graphic design, content creation, operations team and
wider stakeholders. Strong organizational skills. Ability to think on the feet, create innovative concepts and
show strategic thinking with the guest always in mind. Strong understanding of guest experience, up to date with
the market and the region and the latest trends in the F&B and lifestyle industry. Strong aesthetic familiarity
with any art direction skills. Previous experience with content creation and photography/videography direction,
planning, coordination. Ability to cope with short timelines and manage multiple projects at the same time.
Excellent copy writing and influencing skills. A people's person with the ability to manage relations within a
multicultural environment. University degree, specialization in business administration/marketing is advantageous.
"""

ATS = [
    "Brand Marketing Manager", "brand building", "brand management", "brand positioning", "brand communication",
    "brand identity", "concept development", "business development", "innovative concepts", "NPD", "launches",
    "social media strategy", "social media management", "content creation", "content direction", "art direction",
    "photography direction", "videography direction", "creative briefs", "copywriting", "marketing campaigns",
    "campaign development", "seasonal campaigns", "Ramadan", "online activation", "offline activation",
    "activation management", "sampling", "community", "influencer marketing", "creators", "UGC", "agencies",
    "partnerships", "talabat", "Noon", "Careem", "F&B", "food", "retail", "lifestyle", "homegrown", "UAE",
    "Dubai", "in-house team", "graphic design", "team leadership", "stakeholder management", "multicultural",
    "measurement of results", "KPIs", "ROI", "ROAS", "market trends", "Instagram", "TikTok", "Meta Ads",
    "Adobe", "Canva", "multiple projects", "short timelines", "business administration",
]

CONTENT = {
    "headline": "Brand Marketing Manager · F&B · Social, Content & Art Direction · Campaigns & Activations",
    "professional_summary": (
        "Brand marketing manager with 4+ years across F&B, retail and lifestyle brands, now leading brand and "
        "marketing for DoFreeze, a homegrown UAE F&B group, with an in-house designer and social media executive. "
        "Builds concepts, content, creator campaigns and online & offline activations, and measures them on sales."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "Homegrown UAE F&B / FMCG group | Befit, Eurocake, SMASH, Flair | GCC + 50+ markets",
            "bullets": [
                "Lead brand building, positioning and communication for a multi-brand F&B portfolio — Befit, Eurocake, the new SMASH brand and Flair — across social, e-commerce, quick-commerce (talabat, Noon, Careem) and retail",
                "Lead an in-house team of two (graphic designer and social media executive): write the creative briefs, set the art direction and approve all social, photo and video content and copy",
                "Built creator and community marketing from zero: seasonal campaigns (Ramadan, back to school, Fitness Month, New Year) with 4 agencies and 103 creator activations, plus talabat sampling, an in-app sweepstakes and running, padel and yoga communities",
                "Develop concepts from brief to launch — 6 NPD launches (packaging, pricing, go-to-market) — and measure every campaign on sales: SMASH × talabat lifted daily sales +165% vs +4.7% for the control brand; Befit × Noon +31% over baseline",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | beauty, fashion & lifestyle",
            "bullets": [
                "Managed 42 beauty, fragrance and fashion brands (incl. KIKO Milano) on assortment, pricing and promotional calendars (+30% GMV QoQ), and created the Beauty Club and Hot on Social projects to lift brand visibility",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Food delivery & quick-commerce leader | €500M+ revenue",
            "bullets": [
                "Managed XL restaurant accounts (KFC, Taco Bell, La Tagliatella, Sushi Shop) on GMV and joint marketing activations, coordinating marketing, operations and customer support to deliver campaigns",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | chocolate & confectionery (Milka, Suchard)",
            "bullets": [
                "Supported confectionery NPD launches (Milka Spread, Mini Suchard) and ran sell-in/sell-out and promotional-effectiveness analysis with Nielsen",
            ],
        },
    ],
    "skills_brand": (  # -> "Brand & Concepts"
        "brand building, positioning & communication, concepts & NPD, copywriting, team leadership"
    ),
    "skills_ecommerce": (  # -> "Social & Content"
        "social media strategy, content creation, art direction, creative briefs, creators & communities"
    ),
    "skills_commercial": (  # -> "Activation & Partners"
        "online & offline activations, sampling, platform partnerships, agency management, Meta & Google Ads"
    ),
    "skills_data": (  # -> "Measurement"
        "campaign results on sales, control groups, KPI tracking, ROI & ROAS, market & trend research"
    ),
    "skills_tools": (  # -> "Tools"
        "Adobe (Photoshop, Illustrator), Canva, Meta Business Suite, Shopify, Power BI, Excel, Claude (AI)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Brand & Concepts",
    "E-Commerce & Digital": "Social & Content",
    "Commercial": "Activation & Partners",
    "Data & Analytics": "Measurement",
}


def make_job() -> Job:
    return Job(
        id=JOB_ID,
        title=TITLE,
        company=COMPANY,
        location="Dubai, UAE (Hybrid)",
        url="https://www.linkedin.com/jobs/search/?keywords=Alabbar%20Enterprises%20ANOTHER%20Brand%20Marketing%20Manager",
        source="linkedin",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Alabbar Enterprises & ANOTHER Brand Marketing Manager F&B",
            "note": "Grupo retail/F&B de EAU (Candylicious, Garrett, Yogurtland; Social House, Karak House, Angelina, "
                    "Ganache...). Publicada hace 8 meses y sigue recibiendo solicitudes (2.368 total, 29 en el último "
                    "día); respuestas fuera de LinkedIn. Árabe es ventaja, no requisito. CV rehecho el 2026-10-01 a "
                    "una página (los de agosto salían a 3 páginas).",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    existing = next((j for j in jobs if j.get("id") == job.id), None)
    entry = existing if existing is not None else {}
    entry.update({
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": entry.get("posted_date", DATE_FOLDER), "raw": job.raw,
        "ai_score": 84, "ai_tier": "Hot",
        "skills_match": [
            "Título exacto y años en rango: el JD pide 4–5, Paula tiene 4+",
            "Lleva hoy la marca de un grupo F&B homegrown de EAU con varias marcas (Befit, Eurocake, SMASH, Flair)",
            "Equipo interno real: diseñador gráfico y social media executive; escribe los briefs y aprueba todo el contenido",
            "Activación online y offline: samplings y sorteo con talabat, comunidades de running, pádel y yoga",
            "Creators con venta medida: 103 activaciones en 3 campañas; SMASH × talabat +165% frente a +4,7% del control",
            "Desarrollo de conceptos: 6 lanzamientos NPD end-to-end",
            "Marcas de F&B conocidas en Glovo (KFC, Taco Bell, La Tagliatella, Sushi Shop) y confitería en Mondelez",
            "Grado en ADE (CUNEF)",
        ],
        "missing_skills": [
            "Árabe: es 'a definite advantage', no requisito; no se pone",
            "Hospitality / restaurante desde el lado marca: las marcas de restaurante las llevó como plataforma (Glovo)",
            "Dirección de rodajes en set: briefa y aprueba foto y vídeo, pero no consta que dirija shoots",
        ],
        "sector_fit": "muy bueno — brand marketing de F&B y lifestyle en un grupo homegrown de EAU",
        "seniority_fit": "en línea — Brand Marketing Manager, 4–5 años pedidos frente a sus 4+",
        "red_flags": [
            "Publicada hace 8 meses con 2.368 solicitudes: puede ser una vacante abierta permanentemente; buscar contacto directo",
            "Respuestas fuera de LinkedIn: hay que solicitar en la web del grupo",
            "Árabe como ventaja: resta frente a candidatos locales, aunque no descarta",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Encaje muy natural: grupo de F&B y retail de EAU que busca un Brand Marketing Manager para varios "
            "conceptos, con equipo interno de diseño y contenido, social, campañas y activación online y offline. "
            "Es lo que Paula hace hoy en DoFreeze, un grupo F&B homegrown de EAU con varias marcas, con su propio "
            "diseñador y social media executive y campañas medidas en venta. Los años cuadran (4–5 pedidos). "
            "Huecos: no ha hecho marketing de restaurante u hospitality desde el lado marca (en Glovo las llevó "
            "como plataforma), no consta que dirija rodajes en set y no habla árabe, que aquí solo es ventaja. "
            "El riesgo real es la oferta: 8 meses publicada y más de 2.300 solicitudes, así que conviene "
            "acompañar la solicitud con un mensaje directo a alguien de marketing del grupo."
        ),
        "scored_by": "manual:claude", "freshness": entry.get("freshness", "fresh"),
        "discovered_at": entry.get("discovered_at", f"{DATE_FOLDER}T00:00:00+00:00"),
        "status": entry.get("status", "Reviewing"),
    })
    if existing is None:
        jobs.insert(0, entry)
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
