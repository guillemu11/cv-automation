"""One-off: CV for "Digital Marketing Executive", posted by Impetus Executive (HR / recruitment firm) for a
confidential client: a fast-growing tech company selling an AI-driven app suite for guest experience,
visitor-footfall analytics and indoor tracking / IoT to large enterprise customers (malls, retail and other
physical spaces). Dubai, hybrid, full-time. LinkedIn Easy Apply, posted 2026-10-12, 176 applicants.
CV only (no cover letter requested).

The JD: a fully in-house, hands-on role (explicitly NOT agency management). One person runs the whole chain:
SEO -> Google/LinkedIn Ads -> content -> social -> website -> analytics -> lead generation -> CRO. Wants 3-6+
yrs, preferably agency-trained, B2B / enterprise understanding, tech or SaaS a plus; GA4, Search Console,
Google/Meta/LinkedIn Ads, SEMrush/Ahrefs; strong copywriting. No Arabic asked.

Paula's honest angle (she does this chain in-house at DoFreeze today):
- SEO: built the SEO/AEO strategy from scratch for dofreeze.ae: technical fixes (canonicals, redirects,
  404s, indexation), structured data, brand pages; Search Console via Claude Code; llms.txt, Wikidata,
  Google Business Profile for AI answers.
- Paid: runs Google Ads and Meta Ads hands-on (audiences, retargeting, catalogues, Pixel/CAPI, A/B tests).
- Website & CRO: owns the Shopify store and campaign landing pages end-to-end.
- Content & social: leads a designer + social media executive; creator programme measured vs control.
- Analytics & AI: daily executive dashboard (traffic, revenue by channel, paid campaigns); Claude system
  (~40% less manual work, incl. client-ready pitch decks).
- B2B: three years selling to and growing business partners at Miravia (42 accounts, +30% GMV QoQ) and Glovo.

Honesty guardrails:
- No agency background: not claimed anywhere.
- LinkedIn Ads, SEMrush and Ahrefs are not owned: left off the CV (GA4 + Search Console are claimed).
- B2B SaaS / tech lead-gen (MQL/CPL pipelines) not claimed; B2B framed as partner management.
- DoFreeze is NOT on Deliveroo. Never "own the forecast". No Arabic. Visa line comes from profile.yaml.
- Executive-level title: salary likely below her 20k AED/month floor (flagged in the dashboard entry).

ONE-PAGE standard. Lands under output/2026-10-12/.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Impetus Executive"
TITLE = "Digital Marketing Executive"
DATE_FOLDER = "2026-10-12"

JOB_DESCRIPTION = """\
Digital Marketing Executive — Impetus Executive (recruiter) for a confidential client. Dubai, hybrid, full-time.
Client: fast-growing technology and digital services company providing an AI-driven application suite for
guest experience, visitor footfall analytics and indoor tracking & IoT solutions for physical spaces, sold to
large enterprise customers across multiple markets.
Fully in-house, hands-on role (not agency management): develop, execute, manage and optimise all digital
marketing directly.
Strategy & execution: integrated digital strategy; own day-to-day activity; campaigns for awareness, traffic,
qualified leads; new channels and tools; full funnel from awareness to lead generation and conversion.
SEO: own strategy and execution; keyword research, competitor analysis, audits; on-page, off-page and
technical SEO; content, landing pages, metadata, internal linking, site structure; rankings, organic traffic,
leads; Google Search Console, GA4, SEMrush, Ahrefs.
Paid: plan, launch, manage and optimise Google Ads, Meta Ads, LinkedIn Ads; structures, audiences, targeting,
ad copy, landing pages; CPC, CPL, CPA, conversion rate, ROI; A/B testing; budgets.
Content & social: B2B content strategy for enterprise decision-makers; social channels and content calendar;
website, LinkedIn, blog, newsletter content; SEO copy; translate complex technology into clear messages.
Website & CRO: own the website from a marketing perspective; UX and conversion; landing pages; CTAs, forms,
journeys; work with developers and designers.
Analytics: KPIs; website, SEO, paid, social and lead-gen performance; GA4, Search Console; regular reports.
Requirements: 3–6+ years hands-on digital marketing, preferably agency; SEO, PPC/SEM, content, social,
analytics, lead gen; B2B / enterprise marketing; tech, SaaS, IT a plus; Google Ads, LinkedIn Ads, Meta Ads,
GA4, Search Console; SEMrush / Ahrefs; copywriting; analytical; commercial mindset; independent ownership;
excellent English.
"""

ATS = [
    "Digital Marketing Executive", "digital marketing", "hands-on", "in-house", "integrated digital strategy",
    "full funnel", "SEO", "technical SEO", "on-page SEO", "structured data", "keyword research", "site structure",
    "Google Search Console", "GA4", "Google Analytics 4", "organic traffic", "PPC", "SEM", "paid media",
    "Google Ads", "Meta Ads", "LinkedIn Ads", "audiences", "targeting", "retargeting", "ad copy", "A/B testing",
    "CPC", "CPA", "conversion rate", "ROI", "ROAS", "budget management", "content marketing", "content calendar",
    "social media", "copywriting", "newsletters", "website", "landing pages", "CTAs", "UX",
    "conversion rate optimisation", "CRO", "lead generation", "B2B", "enterprise", "KPIs", "reporting",
    "dashboards", "AI", "automation", "Shopify", "Dubai", "UAE",
]

CV_CONTENT = {
    "headline": (
        "Hands-on Digital Marketing · SEO · Google & Meta Ads · Content & Social · Website & CRO · AI Automation"
    ),
    "professional_summary": (
        "Hands-on digital marketer in Dubai with 5 years across e-commerce, marketplaces and FMCG (Alibaba, "
        "Glovo, Mondelez). At DoFreeze I execute SEO, Google and Meta Ads, the website and reporting myself, and "
        "lead content and social, on AI workflows I build. Three years managing B2B partners at Miravia and Glovo."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE F&B / FMCG group | Befit, Eurocake, Flair | Shopify D2C, talabat, Noon, Careem | 50+ markets",
            "bullets": [
                "SEO: built the SEO/AEO strategy from scratch for dofreeze.ae: technical fixes (canonicals, redirects, 404s, indexation), structured data and brand pages tracked in Search Console; llms.txt, Wikidata and Google Business Profile for AI answers",
                "Paid media: plan, launch and optimise Google Ads and Meta Ads (campaign structure, audiences, retargeting, catalogue ads, Pixel/CAPI, creative A/B tests), managing budget against CPA and ROAS",
                "Website & CRO: own the Shopify store and campaign landing pages end-to-end (offers, UX, CTAs, checkout) to lift conversion rate and average order value",
                "Content & social: lead a designer and a social media executive on the content calendar; built the creator programme from zero, measured on incremental sales (SMASH × talabat +165% vs +4.7% control)",
                "Analytics & AI: built a daily executive dashboard (traffic, revenue by channel, every paid campaign) and a Claude system that automates planning, content, KPI reports and client pitch decks, cutting manual work ~40%",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's e-commerce marketplace | 100K+ employees",
            "bullets": [
                "Grew 42 B2B brand accounts (incl. KIKO Milano) +30% GMV QoQ, turning traffic, conversion and sales data into promotion and assortment plans; reported the Flash Sales channel to the CEO",
                "Onboarded 30+ new stores in two months as PIC Fragrances; created the Beauty Club and Hot on Social programmes to turn social content into marketplace demand",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Managed XL accounts (KFC, Taco Bell, La Tagliatella) on GMV, negotiated commercial deals and ran joint marketing activations; helped build the Retail vertical, onboarding new brand partners",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | Chocolate category",
            "bullets": [
                "Sell-in/sell-out and promotional-effectiveness analysis and performance reports for the chocolate category",
            ],
        },
    ],
    "skills_brand": (  # -> "SEO & Content"
        "technical SEO, structured data, AEO/GEO, EN/ES web & ad copy, newsletters & EDM"
    ),
    "skills_ecommerce": (  # -> "Paid Media"
        "Google Ads, Meta Ads, audiences & retargeting, catalogue ads, Pixel/CAPI, A/B testing, budget pacing"
    ),
    "skills_commercial": (  # -> "Website & CRO"
        "Shopify (end-to-end), landing pages, CTAs & customer journeys, UX, conversion rate, AOV"
    ),
    "skills_data": (  # -> "Analytics"
        "GA4, Google Search Console, Merchant Center, Looker, Power BI, CPC, CPA, ROAS, lift vs control"
    ),
    "skills_tools": (
        "Claude / Claude Code (MCP), ChatGPT, Canva & Adobe, Notion, Excel (advanced), Salesforce"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "SEO & Content",
    "E-Commerce & Digital": "Paid Media",
    "Commercial": "Website & CRO",
    "Data & Analytics": "Analytics",
    "Tools": "Tools & AI",
}


def make_job() -> Job:
    return Job(
        id="impetus-digital-marketing-executive-2026-10",
        title=TITLE,
        company=COMPANY,
        location="Dubai, United Arab Emirates",
        url="https://www.linkedin.com/jobs/search/?keywords=Digital%20Marketing%20Executive%20Impetus%20Executive",
        source="manual",
        description=JOB_DESCRIPTION,
        raw={"client": "Confidential: AI guest-experience, footfall analytics and indoor tracking / IoT tech company",
             "via": "LinkedIn Easy Apply, Impetus Executive (posted 2026-10-12, 176 applicants)",
             "workplace": "Dubai, hybrid",
             "level": "Executive"},
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not posted (Executive band — likely below 20k AED/month floor)",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": "2026-10-12", "raw": job.raw,
        "ai_score": 64, "ai_tier": "Warm",
        "skills_match": [
            "Runs the in-house chain at DoFreeze: SEO, Google & Meta Ads, Shopify site & landing pages, content, social, reporting",
            "Technical SEO + structured data + AEO/GEO on dofreeze.ae; GA4 and Search Console",
            "Daily executive dashboard and Claude automation (~40% less manual work)",
            "B2B partner management at Miravia (42 accounts, +30% GMV QoQ) and Glovo",
            "Hybrid Dubai, 3–6 yrs asked (Paula ~5), no Arabic required",
        ],
        "missing_skills": [
            "No agency background (JD prefers agency-trained)",
            "LinkedIn Ads, SEMrush / Ahrefs not used hands-on",
            "B2B SaaS / enterprise lead-gen (CPL, MQL) not owned; sector is consumer",
        ],
        "sector_fit": "stretch (B2B enterprise tech vs her consumer / e-commerce background)",
        "seniority_fit": "below her level (Executive, solo IC; she is Manager with a team of two)",
        "red_flags": [
            "Executive band: salary likely under her 20k AED/month floor",
            "176 applicants; recruiter-filtered, agency background preferred",
        ],
        "ats_keywords": ATS,
        "reasoning": "The hands-on SEO → paid → web → content → analytics chain is what she does in-house; B2B tech sector, missing LinkedIn Ads / SEO tooling and the Executive salary band are the risks.",
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _pages(pdf: Path) -> int:
    from pypdf import PdfReader
    return len(PdfReader(str(pdf)).pages)


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    final_dir = settings.output_dir / DATE_FOLDER / f"{COMPANY} - {TITLE}"
    dest = final_dir / "01_CV_y_Carta"
    dest.mkdir(parents=True, exist_ok=True)

    cv_docx = cv._fill_template(CV_CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)

    moved = dest / cv_pdf.name
    shutil.move(str(cv_pdf), str(moved))
    shutil.copy(str(moved), str(dest / "Paula De Francisco - CV.pdf"))
    print("OK", moved.name, "PAGES", _pages(moved))

    stray = cv_pdf.parent.parent
    if stray.exists() and stray.resolve() != final_dir.resolve():
        shutil.rmtree(stray, ignore_errors=True)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
