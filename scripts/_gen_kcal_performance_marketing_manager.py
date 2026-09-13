"""One-off: generate Paula's CV for **Performance Marketing Manager** at
**Kcal Healthy Fast Food** (UAE healthy F&B chain — restaurants + delivery + meal plans;
Dubai, in person — Indeed, AED 18,000–20,000/month).

JD ask: own paid media end-to-end across Meta, Google Ads (Search, Display, YouTube, PMax),
TikTok and Snapchat — campaign/ad-set structure, naming conventions, budget allocation and
pacing, daily optimisation across brands. Define a structured UTM framework, ensure correct
tracking via Google Analytics, GTM and platform pixels. Turn performance data into weekly
updates and monthly deep-dive reports. Work with Brand, Creative, Content, CRM and tech.
Proactively bring creative/testing ideas. Manage budgets for ROI. Manage and mentor a Junior
Performance Marketing Executive. 5 years performance marketing preferred.

Paula's honest angle:
- PAID MEDIA HANDS-ON: runs Meta (FB/IG) and Google Ads end-to-end at DoFreeze — structure,
  budgets, audience building, creative A/B testing, daily ROI/ROAS optimisation across 3 brands.
- F&B + DELIVERY, TWICE: DoFreeze is UAE F&B (Befit, Eurocake, Flair) integrated into Talabat,
  Deliveroo, Careem and Noon; at Glovo she ran XL restaurant accounts (KFC, Taco Bell,
  La Tagliatella, Sushi Shop). Kcal lives exactly there — own channel + aggregators.
- FUNNEL + TRACKING: owns the Shopify store end-to-end, so she reads the whole path from ad
  to checkout (traffic, conversion, AOV) rather than stopping at platform metrics.
- REPORTING CADENCE: weekly/monthly performance readouts to stakeholders — and at Miravia she
  presented Flash Sales performance to the CEO against P&L targets.
- CREATIVE PIPELINE: briefs 25–50 creators per campaign, so paid social never runs out of
  fresh UGC to test.

Honesty guardrails (do NOT overstate):
- Never held a "Performance Marketing Manager" title — paid media sits inside a broader brand /
  e-commerce remit. CV claims hands-on paid media, not a pure-performance career.
- Hands-on platforms are Meta Ads and Google Ads (Search & Display). Do NOT claim TikTok Ads,
  Snapchat Ads, YouTube or Performance Max campaign management; TikTok/Instagram appear as
  content channels only.
- Google Analytics / Tag Manager: certified (Google Digital Marketing & E-Commerce) and used
  for reporting — do NOT claim she architected GTM containers or a dataLayer.
- No direct reports ever. Mentoring claim limited to briefing/QC of creators and freelancers.
- No Arabic. Factual "UAE Residence Visa" only — never "no sponsorship needed".
- Salary band 18–20K sits AT/BELOW her 20K floor — flagged in the dashboard entry.

ONE-PAGE standard: tight summary, 3/2/1/1 bullets, ~6-9 short skills per row.
Fills the real CV template, relabels the skills rows for this role, converts to PDF via
LibreOffice, registers the job, verifies 1 page, lands under output/2026-09-03/.
Also drops a short-named 'Paula De Francisco - CV.pdf' copy for portals / Easy Apply.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Kcal Healthy Fast Food"
TITLE = "Performance Marketing Manager"
DATE_FOLDER = "2026-09-03"

JOB_DESCRIPTION = """\
Performance Marketing Manager — Kcal Healthy Fast Food. Dubai, UAE. Permanent, in person.
AED 18,000 - 20,000 per month.
Duties and Responsibilities: Plan, create, launch and manage paid media campaigns across Meta
(Facebook & Instagram), Google Ads (Search, Display, YouTube, Performance Max), TikTok Ads,
Snapchat Ads and other relevant platforms. Own the full campaign structure, including campaign,
ad set and ad setup, naming conventions, budget allocation and pacing, and platform-specific
best practices. Actively manage and optimise campaigns on a daily basis to drive growth,
efficiency and scale across different brands and objectives. Define and implement a clear and
structured UTM framework across all paid channels to ensure accurate tracking and full
transparency for stakeholders. Ensure correct implementation of tracking events using tools such
as Google Analytics, Google Tag Manager and platform pixels. Analyse performance data to identify
trends, insights, risks and opportunities, translating them into clear recommendations and
actions. Produce and present weekly performance updates and monthly in-depth reports with
insights and next steps. Work closely with Brand Managers, Creative, Content, CRM and technical
teams to ensure alignment between media strategy, creatives and business goals. Proactively
provide content and creative ideas, trends and out-of-the-box testing concepts to continuously
improve performance. Allocate and manage budgets efficiently across platforms to maximise ROI
and performance impact. Stay up to date with platform updates, trends, new ad formats and
industry best practices. Manage, mentor and guide a Junior Performance Marketing Executive,
ensuring quality, structure and learning development.
Experience: Performance Marketing 5 years (Preferred). Work Location: In person.
"""

ATS = [
    "performance marketing", "paid media", "paid social", "paid search",
    "Meta Ads", "Facebook Ads", "Instagram Ads", "Google Ads", "Search", "Display",
    "TikTok", "Snapchat", "campaign structure", "ad set", "naming conventions",
    "budget allocation", "budget pacing", "daily optimisation", "audience building",
    "creative A/B testing", "creative testing", "UTM", "UTM framework", "tracking",
    "conversion tracking", "Google Analytics", "Google Tag Manager", "pixels",
    "attribution", "ROAS", "ROI", "CPA", "CPC", "CTR", "CPM", "conversion rate",
    "AOV", "GMV", "full-funnel", "CRO", "landing pages", "Shopify",
    "weekly performance updates", "monthly reports", "KPI reporting", "dashboards",
    "Power BI", "Looker", "stakeholder management", "CRM", "EDM", "lifecycle",
    "influencer marketing", "UGC", "content ideas", "F&B", "food delivery",
    "quick-commerce", "Talabat", "Deliveroo", "Careem", "Noon",
    "multi-brand", "mentoring", "Dubai", "UAE",
]

CONTENT = {
    "headline": (
        "Performance Marketing · Paid Media on Meta & Google Ads · UTM, Tracking & Reporting · "
        "F&B, Delivery Platforms & E-Commerce"
    ),
    "professional_summary": (
        "Marketer with 5 years across F&B, FMCG and marketplace e-commerce, running paid media hands-on. "
        "Plans and optimises Meta and Google Ads daily for three brands — structure, budget pacing, creative "
        "testing, ROI/ROAS — and owns the Shopify funnel and tracking behind the click. Dubai-based, with "
        "F&B delivery experience (Talabat, Deliveroo, Careem, Noon)."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE F&B / FMCG group (Befit, Eurocake, Flair) | Shopify D2C + Talabat, Deliveroo, Careem, Noon",
            "bullets": [
                "Plan, launch and manage paid media on Meta (Facebook & Instagram) and Google Ads for three brands — campaign and ad-set structure, naming conventions, budget allocation and pacing, audience building and creative A/B testing — optimised daily against ROI and ROAS",
                "Own the Shopify store and its tracking end-to-end — UTM-tagged campaigns, platform pixels, analytics — reading traffic, conversion and AOV from click to checkout to reallocate budget and lift CRO",
                "Turn performance data into weekly updates and monthly deep-dives for stakeholders, and feed Creative and Content the concepts worth testing — including 25–50 briefed creators per campaign supplying UGC for paid social",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba's lifestyle marketplace | Top-5 global e-commerce | 100K+ employees",
            "bullets": [
                "Grew 42 accounts +30% GMV QoQ by allocating promo and media investment where it paid back — testing pricing, offers and campaigns against traffic, conversion, ROI and ROAS data every week",
                "Owned the Flash Sales channel for Beauty, Fashion & Home reporting directly to the CEO — pacing investment against P&L targets and presenting performance, insights and next steps monthly",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce / food delivery leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Managed XL restaurant accounts (KFC, Taco Bell, La Tagliatella, Sushi Shop) — running in-app promo mechanics and bespoke activations, and reading order, conversion and GMV data to grow volume profitably",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | Category planning & analytics",
            "bullets": [
                "Measured promotional effectiveness and sell-in/sell-out performance for the chocolate category — the measurement discipline behind deciding which spend earns more budget",
            ],
        },
    ],
    "skills_brand": (  # label -> "Paid Media & Performance"
        "campaign & ad-set structure, naming conventions, budget allocation & pacing, daily optimisation, "
        "audience building, creative A/B testing, multi-brand budgets"
    ),
    "skills_ecommerce": (  # label -> "Platforms & Channels"
        "Meta Ads (Facebook & Instagram), Google Ads (Search & Display), paid social, Shopify, EDM / CRM, "
        "food delivery & quick-commerce (Talabat, Deliveroo, Careem, Noon), TikTok content"
    ),
    "skills_commercial": (  # label -> "Tracking & Measurement"
        "UTM frameworks & naming taxonomy, platform pixels, conversion tracking, Google Analytics, "
        "Google Tag Manager, funnel & drop-off diagnosis, attribution reading"
    ),
    "skills_data": (  # label -> "Analysis & Reporting"
        "ROAS, ROI, CPA, CPC, CTR, conversion rate, AOV, GMV, weekly performance updates, "
        "monthly in-depth reports, Power BI, Looker, Tableau"
    ),
    "skills_tools": (  # label -> "Tools"
        "Meta Ads Manager, Google Ads, Google Analytics, Shopify, Power BI, Looker, Canva, Adobe, "
        "Generative AI (Claude, ChatGPT)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Paid Media & Performance",
    "E-Commerce & Digital": "Platforms & Channels",
    "Commercial": "Tracking & Measurement",
    "Data & Analytics": "Analysis & Reporting",
}


def make_job() -> Job:
    return Job(
        id="kcal-performance-marketing-manager-dubai-2026-09",
        title=TITLE,
        company=COMPANY,
        location="Dubai, United Arab Emirates (In person)",
        url="https://ae.indeed.com/",
        source="indeed",
        description=JOB_DESCRIPTION,
        salary_raw="AED 18,000 - AED 20,000 a month",
        salary_aed_min=18000,
        salary_aed_max=20000,
        raw={
            "query": "Kcal Healthy Fast Food Performance Marketing Manager Dubai",
            "note": "UAE healthy F&B chain (restaurants, delivery, meal plans). Pure paid-media role: "
                    "Meta, Google, TikTok, Snapchat + UTM/GA/GTM tracking + weekly/monthly reporting + "
                    "one junior direct report. Real gaps: no TikTok/Snapchat Ads or YouTube/PMax "
                    "campaign management, no GTM implementation ownership, no direct reports ever. "
                    "Salary band 18–20K AED/month sits AT/BELOW Paula's 20K floor — negotiate at the top "
                    "of the band. On-site role in Dubai, which fits.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": job.salary_raw, "salary_aed_min": 18000, "salary_aed_max": 20000,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 70, "ai_tier": "Warm",
        "skills_match": [
            "Hands-on Meta (FB/IG) and Google Ads across three brands — structure, budget pacing, audience building, creative A/B testing, daily ROI/ROAS optimisation",
            "F&B + delivery twice over: DoFreeze on Talabat/Deliveroo/Careem/Noon, and Glovo XL restaurant accounts (KFC, Taco Bell, Sushi Shop) — Kcal's exact channel mix",
            "Owns the Shopify funnel and its tracking (UTMs, pixels, analytics) — reads ad click → checkout, not just platform metrics",
            "Weekly/monthly reporting cadence to senior stakeholders, including Flash Sales performance to Miravia's CEO against P&L targets",
            "Creative pipeline for paid social: 25–50 creators briefed per campaign generating testable UGC",
            "Multi-brand budget allocation experience (+30% GMV QoQ across 42 accounts)",
            "5 years experience, matching the 5-year preference; already in Dubai for an in-person role",
        ],
        "missing_skills": [
            "No TikTok Ads or Snapchat Ads campaign management — both are explicitly in the JD's platform list",
            "No YouTube or Performance Max experience within Google Ads (Search & Display only)",
            "Google Analytics/GTM at certified + reporting level, not implementation — has not architected containers, dataLayers or event taxonomies",
            "No direct reports ever — the JD includes managing and mentoring a Junior Performance Marketing Executive (she has briefed creators, agencies and freelancers, not employees)",
            "No pure-performance job title — paid media has always sat inside a broader brand/e-commerce remit",
        ],
        "sector_fit": "strong — UAE F&B with own-channel + aggregator delivery is exactly DoFreeze + Glovo territory",
        "seniority_fit": "borderline-good — 5 years total matches the preference, but as a broad marketer rather than 5 years of pure performance",
        "red_flags": [
            "AED 18,000–20,000/month: the band tops out AT Paula's 20K floor, so anything below the ceiling is a downgrade",
            "Manager title but a heavily executional remit (daily campaign management) plus one junior",
            "Platform gap is visible on paper: TikTok and Snapchat Ads are in the JD and not on her CV",
            "In-person role — no hybrid mentioned",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Good-but-honest fit. The channel and category context could hardly be closer: Kcal is a UAE "
            "healthy F&B brand selling through its own site, meal plans and the delivery aggregators, and "
            "Paula currently runs paid media for three F&B/FMCG brands on Shopify plus Talabat, Deliveroo, "
            "Careem and Noon — after a year managing XL restaurant accounts at Glovo. She does the core of "
            "this job today: campaign structure, budget pacing, creative testing, daily ROI/ROAS "
            "optimisation, UTM discipline and weekly/monthly stakeholder reporting. Two gaps are real and "
            "should not be papered over: she has not run TikTok or Snapchat Ads (nor YouTube/PMax), and she "
            "has never had a direct report, while this role includes mentoring a junior. Both are learnable "
            "and worth addressing head-on in outreach — paid-social mechanics transfer platform to platform, "
            "and she has managed creator and freelance workflows. The commercial catch is the band: "
            "18–20K AED/month means the ceiling equals her floor, so this is only worth pursuing at 20K "
            "with the package (or if the title/scope justifies it)."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
