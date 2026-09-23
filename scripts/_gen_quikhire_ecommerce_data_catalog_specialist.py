"""One-off: CV de una pagina para **E-commerce Specialist - Data & Catalog (Remote)**
en **Quik Hire Staffing** (agencia; cliente sin nombrar), UAE remoto. Publicada hace
1 semana, 761 candidatos, 47% sin experiencia.

Pide el JD: analizar y limpiar datos de producto para precision y completitud en catalogos;
organizar y estructurar la informacion de producto para busqueda y presentacion optimas;
colaborar con equipos cross-funcionales para resolver discrepancias y mejorar calidad de dato;
implementar y mantener politicas de gobernanza de dato; monitorizar y reportar metricas de
rendimiento del catalogo. Piden experiencia con plataformas de e-commerce y sistemas de gestion
de catalogo, analisis y limpieza de datos, estandares de dato de producto (GTIN/UPC/ISBN),
politicas de gobernanza, atencion al detalle y trabajo con datasets grandes.

Angulo honesto de Paula: DoFreeze — Shopify end-to-end (catalogo, contenido de producto,
colecciones, precios) mas listings en talabat/Noon/Careem/Deliveroo, con revision semanal de
listings activos vs stock; Miravia (Alibaba) — 42 cuentas con surtido, pricing y ficha de
producto en un marketplace de 100K+ empleados; Mondelez — analisis sell-in/sell-out con Nielsen
sobre datasets grandes. Excel a nivel experto y Power BI/Tableau/Looker.

Guardarrailes de honestidad:
- NO tiene experiencia formal con GTIN / UPC / ISBN ni con un PIM dedicado. No se menciona
  ninguno de esos estandares como experiencia propia — se habla de "product data standards"
  solo en la forma en que realmente los ha tocado (integridad de ficha, SKU, barcodes de FMCG).
- NO ha implementado politicas de data governance. Se dice "mantiene consistencia de catalogo",
  que es lo que de verdad hace.
- Visa: solo "UAE Residence Visa" (patrocinada por el empleador); nunca "no sponsorship needed".
- Sin arabe. El JD no lo pide.
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document

from career_ops.config import settings
from career_ops.discovery.normalize import Job
from career_ops.generators import cv_generator as cv

COMPANY = "Quik Hire Staffing"
TITLE = "E-commerce Specialist - Data & Catalog (Remote)"
DATE_FOLDER = "2026-09-22"

JOB_DESCRIPTION = """\
E-commerce Specialist - Data & Catalog (Remote) — Quik Hire Staffing (hiring for a client), UAE, remote, full time.
Managing and optimizing product data to ensure accuracy and consistency across e-commerce platforms; maintaining
high-quality catalogs that drive sales and improve customer experience.
Responsibilities: analyze and cleanse product data to ensure accuracy and completeness in catalogs; organize and
structure product information for optimal searchability and presentation; collaborate with cross-functional teams
to resolve discrepancies and improve data quality; implement and maintain data governance policies to ensure
compliance with standards; monitor and report on catalog performance metrics to identify trends and areas for
improvement.
Required: experience with e-commerce platforms and catalog management systems; proficiency in data analysis and
cleansing techniques; familiarity with product data standards such as GTIN, UPC or ISBN; ability to interpret and
apply data governance policies; strong attention to detail and problem-solving skills; experience working with
large datasets and identifying inconsistencies.
"""

ATS = [
    "e-commerce", "product data", "catalog management", "catalogue", "data quality", "data cleansing",
    "data accuracy", "completeness", "consistency", "product information", "product content",
    "searchability", "merchandising", "SKU", "assortment", "listings", "marketplaces", "Shopify",
    "large datasets", "data analysis", "discrepancies", "cross-functional", "reporting",
    "catalog performance", "metrics", "KPI dashboards", "Excel", "Power BI", "Tableau", "Looker",
    "attention to detail", "problem solving", "conversion", "CRO", "GMV", "sell-in", "sell-out",
]

CONTENT = {
    "headline": "E-Commerce · Product Data & Catalogue · Marketplaces & Shopify",
    "professional_summary": (
        "E-commerce professional with 5 years owning product catalogues end-to-end — Shopify and UAE "
        "marketplaces today, 42 accounts across Alibaba's Miravia before that — with a record of keeping "
        "listings accurate, complete and consistent at scale, and Expert-level Excel and BI reporting."
    ),
    "experience": [
        {
            "company": "DoFreeze LLC",
            "role": "Brand & Marketing Manager",
            "dates": "Oct 2025 – Present",
            "location": "Dubai, UAE",
            "context": "UAE FMCG house (Befit, Eurocake, Flair) | Shopify D2C + talabat, Noon, Careem, Deliveroo",
            "bullets": [
                "Own the Shopify catalogue end-to-end — product records, titles, descriptions, imagery, attributes, collections and pricing — structuring product information so items are findable and correctly presented on site",
                "Maintain product listings across talabat, Noon, Careem and Deliveroo, running a weekly review of active listings against live stock to catch missing, duplicated or out-of-date records before they cost sales",
                "Resolve data discrepancies with supply chain, sales and platform account managers, keeping SKU, pack-size, barcode and price information consistent across every channel the brand sells on",
                "Report on catalogue and conversion performance in Excel and BI dashboards, tracing drops back to listing quality and acting on what the data shows",
            ],
        },
        {
            "company": "Miravia (Alibaba Group)",
            "role": "Key Account Manager – Beauty, Fragrances & Fashion",
            "dates": "Nov 2023 – Oct 2025",
            "location": "Madrid, Spain",
            "context": "Alibaba Group marketplace | 100K+ employees | Beauty, Fragrance & Fashion verticals",
            "bullets": [
                "Managed the catalogues of 42 accounts on a large marketplace — assortment, product listings, pricing and promotional set-up — delivering +30% GMV growth QoQ",
                "Onboarded 30+ new stores in two months as PIC Fragrances, auditing and correcting their product data at intake so listings went live complete and compliant with platform requirements",
                "Owned the Flash Sales channel for Beauty, Fashion & Home reporting to the CEO: item selection, quantities and campaign set-up across a high-volume event calendar",
                "Analysed traffic, conversion, ROI and retention across large datasets to find where listing quality was holding categories back",
            ],
        },
        {
            "company": "Glovo",
            "role": "Account Manager – XL Accounts",
            "dates": "Sep 2022 – Nov 2023",
            "location": "Madrid, Spain",
            "context": "Quick-commerce leader | €500M+ revenue | 10K+ employees",
            "bullets": [
                "Managed partner menus and catalogues (KFC, Taco Bell, Sushi Shop) and helped onboard fashion and lifestyle brands to the Retail vertical, setting up their product data on the platform",
            ],
        },
        {
            "company": "Mondelez International",
            "role": "Trainee – Category Planning",
            "dates": "Aug 2021 – Aug 2022",
            "location": "Madrid, Spain",
            "context": "Global FMCG | €36B revenue | 90K+ employees",
            "bullets": [
                "Built sell-in/sell-out and promotional-effectiveness reports from large Nielsen datasets, reconciling inconsistencies across sources before analysis",
            ],
        },
    ],
    "skills_brand": (  # -> "Catalogue & Product Data"
        "catalogue management, product content & attributes, listing accuracy & completeness, "
        "assortment structuring, SKU & pack data consistency"
    ),
    "skills_ecommerce": (  # -> "Platforms"
        "Shopify, Noon, talabat, Careem, Deliveroo, Miravia (Alibaba), Glovo, collections & merchandising, CRO"
    ),
    "skills_commercial": (  # -> "Data & Analysis"
        "data cleansing & reconciliation, large datasets, discrepancy resolution, pricing data, "
        "sell-in/sell-out, forecasting support"
    ),
    "skills_data": (  # -> "Reporting"
        "Excel (Expert), Power BI, Tableau, Looker, Nielsen, KPI dashboards, catalogue performance reporting"
    ),
    "skills_tools": (  # -> "Tools & Languages"
        "Shopify, SAP, Salesforce, Google Sheets, Python automation, Claude (AI) | "
        "Spanish (native), English (C1)"
    ),
}

ROLE_LABELS = {
    "Brand & Marketing": "Catalogue & Product Data",
    "E-Commerce & Digital": "Platforms",
    "Commercial": "Data & Analysis",
    "Data & Analytics": "Reporting",
}


def make_job() -> Job:
    return Job(
        id="quik-hire-ecommerce-specialist-data-catalog-2026-09",
        title=TITLE,
        company=COMPANY,
        location="UAE (Remote)",
        url="https://www.linkedin.com/jobs/search/?keywords=Quik%20Hire%20Staffing%20E-commerce%20Specialist%20Data%20Catalog",
        source="manual",
        description=JOB_DESCRIPTION,
        salary_raw=None,
        salary_aed_min=None,
        salary_aed_max=None,
        raw={
            "query": "Quik Hire Staffing E-commerce Specialist Data Catalog Remote UAE",
            "note": "Agencia de staffing, cliente sin nombrar. 761 candidatos, 47% sin experiencia. "
                    "Senales de baja calidad: 'work from anywhere', lenguaje de captacion masiva "
                    "('regardless of background, experience, or prior employment history'). "
                    "Nivel Specialist de data/catalogo: probable por debajo del suelo de 20k AED.",
        },
    )


def register_in_dashboard(job: Job) -> None:
    path = settings.data_dir / "scored_jobs.json"
    if not path.exists():
        return
    jobs = json.loads(path.read_text(encoding="utf-8"))
    if any(j.get("id") == job.id for j in jobs):
        return
    jobs.insert(0, {
        "id": job.id, "title": job.title, "company": job.company, "location": job.location,
        "url": job.url, "source": job.source, "description": job.description,
        "salary_raw": "Not disclosed",
        "salary_aed_min": None, "salary_aed_max": None,
        "posted_date": DATE_FOLDER, "raw": job.raw,
        "ai_score": 48, "ai_tier": "Warm",
        "skills_match": [
            "Shopify end-to-end: catalogo, contenido de producto, colecciones y pricing",
            "Listings en talabat, Noon, Careem y Deliveroo con revision semanal activos vs stock",
            "42 cuentas en Miravia (Alibaba): surtido, listings y correccion de ficha en el onboarding de 30+ tiendas",
            "Datasets grandes y reconciliacion: Nielsen en Mondelez, analitica de trafico y conversion en Miravia",
            "Excel nivel experto + Power BI / Tableau / Looker",
        ],
        "missing_skills": [
            "Estandares formales de dato de producto (GTIN / UPC / ISBN)",
            "PIM dedicado o catalog management system especifico",
            "Implementacion de politicas de data governance",
        ],
        "sector_fit": "medio — e-commerce y catalogo si; data governance pura no es su perfil",
        "seniority_fit": "por debajo — rol Specialist de data/catalogo frente a su titulo actual de Manager",
        "red_flags": [
            "Agencia de staffing con cliente sin nombrar",
            "761 candidatos y 47% sin experiencia: puesto tasado a nivel entry",
            "'Work from anywhere' y lenguaje de captacion masiva en el anuncio",
            "Salario probablemente por debajo del suelo de 20k AED/mes",
            "Rol de ejecucion de dato, no comercial: no construye su carrera hacia Brand/Marketing Manager",
        ],
        "ats_keywords": ATS,
        "reasoning": (
            "Paula cubre bien la mitad de e-commerce y catalogo del JD — Shopify end-to-end, listings en "
            "marketplaces de UAE, 42 cuentas en Miravia y analisis de datasets grandes. Lo que no cubre es la "
            "mitad de gobernanza de dato: GTIN/UPC/ISBN y politicas formales no estan en su experiencia. "
            "Ademas el puesto esta tasado por debajo de su nivel y la oferta viene de una agencia con cliente "
            "sin nombrar. Vale como opcion remota de respaldo, no como objetivo."
        ),
        "scored_by": "manual:claude", "freshness": "fresh",
        "discovered_at": f"{DATE_FOLDER}T00:00:00+00:00", "status": "Reviewing",
    })
    path.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")


def _relabel_for_role(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    for table in doc.tables:
        for row in table.rows:
            first = row.cells[0]
            new = ROLE_LABELS.get(first.text.strip())
            if not new:
                continue
            para = first.paragraphs[0]
            if para.runs:
                para.runs[0].text = new
                for r in para.runs[1:]:
                    r.text = ""
            else:
                para.add_run(new)
    doc.save(str(docx_path))


def _to_pdf_soffice(docx_path: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    pdf_path = docx_path.with_suffix(".pdf")
    if Path(soffice).exists():
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir",
                 str(docx_path.parent), str(docx_path)],
                check=True, capture_output=True, timeout=180,
            )
            if pdf_path.exists():
                docx_path.unlink(missing_ok=True)
                return pdf_path
        except Exception as exc:  # noqa: BLE001
            print("soffice conversion failed, falling back to docx2pdf:", exc)
    return cv._to_pdf(docx_path)


def _relocate_to_dated_folder(pos_dir: Path) -> Path:
    dated_parent = settings.output_dir / DATE_FOLDER
    dated_parent.mkdir(parents=True, exist_ok=True)
    dest = dated_parent / pos_dir.name
    if pos_dir.resolve() == dest.resolve():
        return dest
    if dest.exists():
        for sub in pos_dir.iterdir():
            target = dest / sub.name
            if sub.is_dir():
                target.mkdir(parents=True, exist_ok=True)
                for f in sub.iterdir():
                    shutil.move(str(f), str(target / f.name))
            else:
                shutil.move(str(sub), str(target))
        shutil.rmtree(pos_dir, ignore_errors=True)
    else:
        shutil.move(str(pos_dir), str(dest))
    return dest


def main() -> None:
    job = make_job()
    register_in_dashboard(job)

    cv_docx = cv._fill_template(CONTENT, job)
    _relabel_for_role(cv_docx)
    cv_pdf = _to_pdf_soffice(cv_docx)
    print("OK_CV", cv_pdf)

    pos_dir = cv_pdf.parent.parent
    final_dir = _relocate_to_dated_folder(pos_dir)

    short = final_dir / "01_CV_y_Carta" / "Paula De Francisco - CV.pdf"
    final_cv = final_dir / "01_CV_y_Carta" / cv_pdf.name
    if final_cv.exists():
        shutil.copy(str(final_cv), str(short))
        print("OK_SHORT", short)

    try:
        from pypdf import PdfReader
        n = len(PdfReader(str(final_cv)).pages)
        print(f"PAGES {n}", "OK" if n == 1 else "!! SPILLS")
    except Exception as exc:  # noqa: BLE001
        print("page-count check skipped:", exc)

    print("OK_PACKAGE", final_dir)


if __name__ == "__main__":
    main()
